#!/usr/bin/env python3
"""Budget Planner Spreadsheet, version 1 (v3 dashboard design, Oct 10, 2026).

A plain monthly budget planner for Excel and Google Sheets:
  1 Setup    take-home income (any pay schedule turned into a monthly amount), bills with due days,
             everyday spending and savings, each with a monthly plan; what is left to plan.
  2 Actuals  one row per line from Setup, one column per month: type what came in and what went out.
  3 Month    pick a month: left to spend, planned vs actual for every line, bills in due-date order.
  4 Year     twelve months side by side, planned vs actual, with in-cell bars.

Built to design/WORKBOOK_STANDARD.md section 0 with tools/pnf_dash.py and tools/wb_style.py.
No macros, no TEXT() format strings, every tab protected without a password.

  python3 build_xlsx.py Budget-Planner-EXAMPLE.xlsx           (the worked example)
  python3 build_xlsx.py Budget-Planner.xlsx --blank           (line names kept, every amount empty)
"""
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font  # noqa: E402

PRODUCT = "Budget Planner Spreadsheet"
VERSION = "1"
CHECKED = "Oct 10, 2026"
CHECKED_LONG = "October 10, 2026"
AS_OF = f"Checked {CHECKED}"
HERE = os.path.dirname(os.path.abspath(__file__))
args = sys.argv[1:]
BLANK = "--blank" in args
args = [a for a in args if a != "--blank"]
OUT = args[0] if args else "Budget-Planner-EXAMPLE.xlsx"
MONEY = '"$"#,##0.00;-"$"#,##0.00;"$"0.00'
MONEY_T = '"$"#,##0.00;-"$"#,##0.00;"-"'     # tables: a dash for zero so empty months stay quiet
USD0 = '"$"#,##0;-"$"#,##0;"-"'
PCT = '0%;-0%;0%'
INT = S.FMT["int"]
HIDE = ";;;"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH_NAMES = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
               "October", "November", "December"]
ARR_MON = "{" + ",".join(f'"{m}"' for m in MONTHS) + "}"
ARR_NAME = "{" + ",".join(f'"{m}"' for m in MONTH_NAMES) + "}"
OFTEN = ["Monthly", "Twice a month", "Every 2 weeks", "Weekly", "Every 3 months", "Yearly"]
OFTEN_ARR = "{" + ",".join(f'"{o}"' for o in OFTEN) + "}"
FACTOR = {"Monthly": 1, "Twice a month": 2, "Every 2 weeks": 26 / 12, "Weekly": 52 / 12, "Every 3 months": 1 / 3,
          "Yearly": 1 / 12}

# ---------------------------------------------------------------- the example: a made-up two-earner household, 2026
# (name, how often, amount)
INCOME = [("Paycheck 1", "Every 2 weeks", 1950), ("Paycheck 2", "Twice a month", 1400), ("Side income", "Monthly", 250),
          ("", "Monthly", None)]
# (name, due day, monthly amount)
BILLS = [("Rent", 1, 1850), ("Car payment", 5, 385), ("Streaming", 9, 32), ("Car insurance", 12, 142), ("Gym", 15, 40),
         ("Electric", 18, 120), ("Water", 20, 55), ("Internet", 22, 70), ("Phone", 25, 95), ("Student loan", 28, 260),
         ("", None, None), ("", None, None)]
SPEND = [("Groceries", "Monthly", 900), ("Gas", "Monthly", 260), ("Eating out", "Monthly", 250),
         ("Household items", "Monthly", 120), ("Kids and school", "Monthly", 200), ("Personal care", "Monthly", 80),
         ("Pets", "Monthly", 90), ("Gifts", "Monthly", 75), ("Fun", "Monthly", 150), ("Clothing", "Monthly", 100),
         ("Car registration", "Yearly", 240), ("Other", "Monthly", 100), ("", "Monthly", None), ("", "Monthly", None),
         ("", "Monthly", None)]
SAVE = [("Emergency fund", "Monthly", 500), ("Retirement, extra", "Monthly", 200), ("Holiday fund", "Monthly", 100),
        ("Extra debt payment", "Monthly", 150), ("", "Monthly", None), ("", "Monthly", None)]
EX_YEAR, EX_MONTH = 2026, 9      # actuals typed for January to September; the Month tab shows September

# Actuals, January to September (index 0 = January). Paycheck 1 is every other Friday from Jan 2, 2026,
# so January and July hold three paychecks.
ACT = {
    "Paycheck 1": [5850, 3900, 3900, 3900, 3900, 3900, 5850, 3900, 3900],
    "Paycheck 2": [2800] * 9,
    "Side income": [180, 240, 310, 250, 0, 275, 320, 190, 260],
    "Rent": [1850] * 9, "Car payment": [385] * 9, "Streaming": [32] * 9, "Car insurance": [142] * 9, "Gym": [40] * 9,
    "Electric": [148.62, 131.40, 112.08, 96.55, 101.27, 134.90, 162.35, 158.11, 127.44],
    "Water": [52.10, 49.80, 54.25, 55.90, 58.40, 61.75, 66.20, 63.10, 52.35],
    "Internet": [70] * 9, "Phone": [95] * 9, "Student loan": [260] * 9,
    "Groceries": [884.31, 912.58, 867.20, 941.05, 896.77, 925.40, 958.12, 903.66, 846.40],
    "Gas": [238.40, 251.12, 247.85, 266.30, 279.14, 284.60, 301.25, 288.42, 247.30],
    "Eating out": [212.45, 268.30, 241.10, 305.75, 226.80, 289.95, 334.60, 247.15, 198.15],
    "Household items": [96.20, 134.85, 108.40, 117.65, 142.30, 99.75, 121.10, 88.90, 104.60],
    "Kids and school": [145.00, 180.00, 165.50, 210.00, 175.25, 95.00, 120.00, 386.40, 186.25],
    "Personal care": [64.30, 71.85, 88.20, 59.40, 92.15, 77.60, 68.45, 81.30, 74.95],
    "Pets": [82.15, 79.40, 146.80, 85.25, 91.60, 78.35, 88.10, 84.70, 82.40],
    "Gifts": [40.00, 95.50, 30.00, 120.00, 85.00, 60.00, 45.00, 70.00, 55.00],
    "Fun": [118.25, 162.40, 140.00, 175.80, 131.55, 196.30, 214.75, 158.20, 112.60],
    "Clothing": [62.40, 0, 148.75, 89.90, 112.35, 54.20, 0, 186.45, 71.30],
    "Car registration": [0, 0, 0, 0, 0, 0, 0, 0, 240],
    "Other": [74.20, 91.35, 112.60, 66.15, 104.80, 88.45, 73.90, 119.25, 64.60],
    "Emergency fund": [500] * 9, "Retirement, extra": [200] * 9, "Holiday fund": [100] * 9,
    "Extra debt payment": [150, 150, 150, 150, 150, 150, 150, 150, 150],
}


def monthly(how, amt):
    return 0 if amt is None else amt * FACTOR.get(how or "Monthly", 1)


def ex_numbers(m):
    """Every headline the example shows for month m (1 to 12), worked out here so listing copy can be checked."""
    i = m - 1
    inc_plan = round(sum(monthly(h, a) for _, h, a in INCOME), 2)
    out_plan = round(sum(a or 0 for _, _, a in BILLS) + sum(monthly(h, a) for _, h, a in SPEND + SAVE), 2)
    inc_act = round(sum(ACT[n][i] for n, _, _ in INCOME if n), 2) if i < 9 else 0
    out_act = round(sum(ACT[n][i] for n, _, _ in BILLS + SPEND + SAVE if n), 2) if i < 9 else 0
    over = [n for n, h, a in SPEND + SAVE if n and i < 9 and ACT[n][i] > monthly(h, a) + 1e-9]
    over += [n for n, _, a in BILLS if n and i < 9 and ACT[n][i] > (a or 0) + 1e-9]
    return dict(inc_plan=inc_plan, out_plan=out_plan, left_to_plan=round(inc_plan - out_plan, 2), inc_act=inc_act,
                out_act=out_act, left_to_spend=round(out_plan - out_act, 2), left_after=round(inc_act - out_act, 2),
                over=over)


EXAMPLE = {m: ex_numbers(m) for m in range(1, 10)}
YEAR_INC = round(sum(EXAMPLE[m]["inc_act"] for m in EXAMPLE), 2)
YEAR_OUT = round(sum(EXAMPLE[m]["out_act"] for m in EXAMPLE), 2)
if BLANK:
    INCOME = [(n, h, None) for n, h, _ in INCOME]
    BILLS = [(n, None, None) for n, _, _ in BILLS]
    SPEND = [(n, h, None) for n, h, _ in SPEND]
    SAVE = [(n, h, None) for n, h, _ in SAVE]

SET_N, ACT_N, MON_N, YR_N = "1 Setup", "2 Actuals", "3 Month", "4 Year"
SET, ACTS, MON, YR = (f"'{n}'" for n in (SET_N, ACT_N, MON_N, YR_N))
GROUPS = [("inc", "INCOME", "Take-home income", INCOME), ("bill", "BILLS", "Bills", BILLS),
          ("spend", "EVERYDAY SPENDING", "Everyday spending", SPEND),
          ("save", "SAVINGS AND EXTRA DEBT PAYMENTS", "Savings and extra debt payments", SAVE)]


def rich(head, body, size=9):
    return CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=size, b=True, color=D.T["ink"]), head + ". "),
                        TextBlock(InlineFont(rFont=D.FONT, sz=size, color=D.T["ink2"]), body))


def sign_rule(ws, ref):
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'AND(ISNUMBER({ref}),{ref}<0)'], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'AND(ISNUMBER({ref}),{ref}>0)'], font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))


def notes_card(ws, top, left, right, c0, c1, notes, title="Notes"):
    r = top
    D.h(ws, r, 12); r += 1
    t = ws[f"{c0}{r}"]; t.value = title; t.font = D.f(12, True); ws.merge_cells(f"{c0}{r}:{c1}{r}"); D.h(ws, r, 24); r += 1
    width = sum(S.width_of(ws, k) for k in D.cols(c0, c1))
    for head, body in notes:
        cell = ws[f"{c0}{r}"]; cell.value = rich(head, body)
        cell.alignment = Alignment(vertical="top", wrap_text=True); ws.merge_cells(f"{c0}{r}:{c1}{r}")
        D.h(ws, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
    D.h(ws, r, 12)
    D.frame(ws, top, r, left, right)
    return r


def fill_heights(ws, last):
    for rr in range(1, last + 1):
        if ws.row_dimensions[rr].height is None:
            D.h(ws, rr, D.GRID_ROW)


wb = S.new_workbook(PRODUCT, version=VERSION)

# =====================================================================  1 Setup
st = wb.create_sheet(SET_N)
GRID = {"A": 3, "B": 2, "C": 30, "D": 16, "E": 15, "F": 2, "G": 2, "H": 15, "I": 12, "J": 24, "K": 2, "L": 3}
S.set_widths(st, GRID)
D.page_header(st, "Your monthly plan", "Type your take-home pay and every bill, spend and saving once. The other tabs follow.",
              AS_OF, span=("B", "K"), side_cols=3)

# ---- tables first (their rows feed the cards above them)
T0 = 23
ROWS = {}           # group key -> (first, last) on Setup
r = T0
for key, eyebrow, title, items in GROUPS:
    if key == "bill":
        mid = dict(col="D", head="Due day", kind="input", fmt="0", align="right", values=[x[1] for x in items],
                   validation=("whole", 1, 31), error="Type a whole day of the month from 1 to 31.",
                   prompt="The day of the month this bill is due, 1 to 31. Use 31 for the last day; shorter months use their last day.")
        amt_prompt = "What this bill costs in a month, for example 1850. For a bill that changes, type a typical month."
    else:
        mid = dict(col="D", head="How often", kind="input", values=[x[1] for x in items], validation=("list", OFTEN),
                   error="Pick one from the list: " + ", ".join(OFTEN) + ".",
                   prompt="Pick how often this amount comes. The Per month column turns it into a monthly amount.")
        amt_prompt = ("The amount each time, for example 1950 for one paycheck." if key == "inc" else
                      "The amount each time, for example 240 for a yearly bill. Per month shows a twelfth of it.")
    cols_t = [
        dict(col="C", head="Name", kind="input", values=[x[0] or None for x in items], validation=("textLength", 0, 40),
             prompt="A short name you will recognise, up to 40 characters. Clear it to drop the line.",
             error="Up to 40 characters."),
        mid,
        dict(col="E", head="Amount", kind="input", fmt=MONEY, align="right", values=[x[2] for x in items],
             validation=("decimal", 0, None), prompt=amt_prompt, error="Type 0 or more, as a positive number."),
        dict(col="H", head="Per month", kind="calc", fmt=MONEY_T, align="right",
             formula=(lambda rr: f'=IF(OR(C{rr}="",E{rr}=""),0,E{rr})') if key == "bill" else
             (lambda rr: f'=IF(OR(C{rr}="",E{rr}=""),0,E{rr}*IFERROR(CHOOSE(MATCH(D{rr},{OFTEN_ARR},0),1,2,26/12,52/12,1/3,1/12),1))')),
        dict(col="I", head="Of income", kind="calc", fmt=PCT, align="right",
             formula=lambda rr: f'=IF(OR(H{rr}=0,$E$__INC__=0),"",H{rr}/$E$__INC__)'),
        dict(col="J", head="", kind="calc", formula=lambda rr: f'=IF(I{rr}="","",REPT("{D.BAR_CHAR}",MIN(18,ROUND(I{rr}*16,0))))'),
    ]
    f1, f2, end = D.table_card(st, r, "B", "K", cols_t, len(items), title=title,
                               sub={"inc": "After tax and deductions. Pick how often each one comes; Per month does the maths.",
                                    "bill": "Fixed bills with a due day. The Month tab lists them in due-date order.",
                                    "spend": "Groceries, gas and the rest. A yearly cost such as car registration shows a twelfth each month.",
                                    "save": "Money you set aside on purpose: savings, investing, paying a debt down faster."}[key],
                               wb=wb)
    ROWS[key] = (f1, f2)
    for rr in range(f1, f2 + 1):
        st[f"C{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if key != "bill":
            st[f"D{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        st[f"I{rr}"].font = D.f(9, False, "ink2")
        st[f"J{rr}"].font = D.f(9, False, {"inc": "teal", "bill": "ink", "spend": "blue", "save": "gold"}[key])
        st[f"J{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if key == "bill":   # sort key for the bill calendar: due day plus row/1000 (32 = no due day), hidden
            k = st[f"F{rr}"]
            k.value = f'=IF(OR(C{rr}="",E{rr}=""),"",IF(D{rr}="",32,D{rr})+ROW()/1000)'
            k.number_format = HIDE; k.font = D.f(8, False, "muted")
    if key == "inc":
        st[f"I{f1 - 1}"].value = "SHARE"
    r = end + 2
SET_LAST_TABLE = r - 2

# ---- summary cards (rows 11+)
P = D.Card(st, 11, "B", "C", "E", "F", wb=wb)
P.title("Your month at a glance", "Totals from the tables below, per month.")
r_year = P.input("Budget year", None if BLANK else EX_YEAR, "0", name="BudgetYear", validation=("whole", 2000, 2100),
                 allow_blank=True, prompt_title="Budget year",
                 prompt="The year this plan is for, for example 2026. Leave blank for this year.",
                 error="Type a year from 2000 to 2100, or leave it blank.")
st[f"E{r_year}"].number_format = "0"
SUMS = {}
for key, _, title, _ in GROUPS:
    a, b = ROWS[key]
    SUMS[key] = P.calc(title if key != "save" else "Savings and extra debt", f"=SUM(H{a}:H{b})", MONEY)
r_out = P.calc("Planned out each month", f"=E{SUMS['bill']}+E{SUMS['spend']}+E{SUMS['save']}", MONEY, total=True)
r_left = P.calc("Left to plan", f"=E{SUMS['inc']}-E{r_out}", MONEY, total=True, tint="total_tint")
sign_rule(st, f"E{r_left}")
P.close()
INC = f"{SUMS['inc']}"
for rr in range(T0, SET_LAST_TABLE + 1):
    c = st[f"I{rr}"]
    if isinstance(c.value, str) and "__INC__" in c.value:
        c.value = c.value.replace("__INC__", INC)
YEAR = f"{SET}!$F${r_year}"
hy = st[f"F{r_year}"]; hy.value = f"=IF(ISNUMBER(E{r_year}),E{r_year},YEAR(TODAY()))"; hy.number_format = HIDE

Q = D.Card(st, 11, "G", "H", "I", "K", extra="J", wb=wb)
Q.title("Where each dollar goes", "Share of your monthly take-home.")
Q.eyebrow("SHARE OF TAKE-HOME")
SHARE = {}
for key, color, label in (("bill", "ink", "Bills"), ("spend", "blue", "Spending"), ("save", "gold", "Savings")):
    SHARE[key] = Q.bar(label, f'=IF($E${INC}=0,"",$E${SUMS[key]}/$E${INC})', PCT, "1", color, width=20)
    st[f"J{SHARE[key]}"].value = f'=IF(OR(I{SHARE[key]}="",I{SHARE[key]}<=0),"",REPT("{D.BAR_CHAR}",MIN(16,ROUND(I{SHARE[key]}*14,0))))'
r_lshare = Q.bar("Left to plan", f'=IF($E${INC}=0,"",$E${r_left}/$E${INC})', PCT, "1", "teal", width=20, bold=True)
st[f"J{r_lshare}"].value = f'=IF(OR(I{r_lshare}="",I{r_lshare}<=0),"",REPT("{D.BAR_CHAR}",MIN(16,ROUND(I{r_lshare}*14,0))))'
q_status = Q.text(f'=IF($E${INC}=0,"Add your take-home pay to see the shares",IF($E${r_left}<0,"Warning: the plan spends more than you bring home",'
                  f'IF($E${r_left}=0,"Every dollar has a job","Give the left-to-plan dollars a job: savings or a line below")))',
                  size=9, color="ink2", bold=True)
Q.text("Savings includes extra debt payments.", size=9, color="muted")
Q.close()
D.status_rule(st, f"H{q_status}", f"$E${r_left}<0", fill_tint=False, anchor=f"H{q_status}")

# ---- tiles
D.tile(st, 6, "B", "F", "LEFT TO PLAN EACH MONTH", f"=$E${r_left}", MONEY,
       f'=IF($E${INC}=0,"Start with your take-home pay in the Income table",IF($E${r_left}<0,"Your plan is more than your income. Trim a line below.",'
       f'IF($E${r_left}=0,"Every dollar of income has a job","Income not yet given a job. Add it to savings or a line.")))', dark=True)
v2, s2, _ = D.tile(st, 6, "G", "K", "PLANNED FOR SAVINGS EACH MONTH", f"=$E${SUMS['save']}", MONEY,
                   f'=IF($E${INC}=0,"Share of take-home shows once income is in",ROUND($E${SUMS["save"]}/$E${INC}*100,0)&"% of take-home goes to savings and extra debt payments")')
D.h(st, 10, 14)

N_TOP = SET_LAST_TABLE + 2
n_end = notes_card(st, N_TOP, "B", "K", "C", "J", [
    ("Pay schedules", "Per month turns each amount into a monthly figure: twice a month is times 2, every 2 weeks is times 26 "
     "and divided by 12, weekly is times 52 and divided by 12, every 3 months is divided by 3, yearly is divided by 12. A month "
     "with a third paycheck shows up on the Actuals tab, where you type what really came in."),
    ("Take-home", "Use pay after tax and deductions, the amount that lands in your account."),
    ("Changing lines", "Rename, add or clear any yellow cell. The Actuals, Month and Year tabs follow the names here, line for line."),
    ("Any currency", "The maths works in any currency. Select the money cells and pick your symbol under Format, Number."),
])
SET_LAST = n_end + 1
D.h(st, SET_LAST, 14)
fill_heights(st, SET_LAST)
D.paint_canvas(st, SET_LAST, "Z")
S.finish_sheet(st, PRODUCT, SET_LAST, span=("A", "L"), freeze="A11", tab_color="accent")
for key in ("bill", "spend", "save"):
    S.page_break_before(st, ROWS[key][0] - 5)

# =====================================================================  2 Actuals
ac = wb.create_sheet(ACT_N)
MCOLS = D.cols("E", "P")
S.set_widths(ac, {"A": 3, "B": 2, "C": 26, "D": 12, **{k: 10.5 for k in MCOLS}, "Q": 12, "R": 2, "S": 3})
D.page_header(ac, "What actually happened", "One column per month. Type what came in and what went out; blanks count as nothing yet.",
              "Example: January to September" if not BLANK else "Type over the yellow cells", span=("B", "R"), side_cols=4)
A_TOP = 11
r = A_TOP
D.h(ac, r, 12); r += 1
t = ac[f"C{r}"]; t.value = f'="Planned and actual, "&{YEAR}'; t.font = D.f(12, True); ac.merge_cells(f"C{r}:Q{r}"); D.h(ac, r, 24); r += 1
t = ac[f"C{r}"]; t.value = "Line names and plans come from the Setup tab. Type positive amounts: money in for income, money out for the rest."
t.font = D.f(9, False, "muted"); ac.merge_cells(f"C{r}:Q{r}"); D.h(ac, r, 18); r += 1
D.h(ac, r, 6); r += 1
AHEAD = r
for k, txt in zip(["C", "D", *MCOLS, "Q"], ["", "PLAN / MONTH", *[m.upper() for m in MONTHS], "YEAR"]):
    c = ac[f"{k}{AHEAD}"]; c.value = txt; c.font = D.f(8, True, "muted")
    c.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
for k in D.cols("C", "Q"):
    ac[f"{k}{AHEAD}"].border = Border(bottom=D.side("ink"))
D.h(ac, AHEAD, 26)
r += 1
AROW = {}          # group key -> (first, last) on Actuals
ATOT = {}          # group key -> total row
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
dv.showErrorMessage = True; dv.errorStyle = "stop"; dv.errorTitle = "Check this cell"
dv.error = "Type 0 or more, as a positive number, cents included."
dv.showInputMessage = True; dv.promptTitle = "Actual amount"
dv.prompt = "What came in (income) or went out (everything else) this month, for example 912.58. Leave blank until it happens."
ac.add_data_validation(dv)


def a_eyebrow(text):
    global r
    c = ac[f"C{r}"]; c.value = text; c.font = D.f(8, True, "muted"); c.alignment = Alignment(vertical="bottom")
    for k in D.cols("C", "Q"):
        ac[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    D.h(ac, r, D.GRID_ROW); r += 1


def a_total(label, fn, plan, tint=None, signed=False):
    global r
    c = ac[f"C{r}"]; c.value = label; c.font = D.f(10, True); c.alignment = Alignment(vertical="center")
    p = ac[f"D{r}"]; p.value = plan; p.number_format = MONEY_T; p.font = D.f(10, True, "ink2")
    p.alignment = Alignment(horizontal="right", vertical="center")
    for m, k in enumerate(MCOLS, 1):
        v = ac[f"{k}{r}"]; v.value = fn(k, m); v.number_format = MONEY_T; v.font = D.f(10, True)
        v.alignment = Alignment(horizontal="right", vertical="center")
        if signed:
            sign_rule(ac, f"{k}{r}")
    y = ac[f"Q{r}"]; y.value = f"=SUM(E{r}:P{r})"; y.number_format = USD0; y.font = D.f(10, True)
    y.alignment = Alignment(horizontal="right", vertical="center")
    if signed:
        sign_rule(ac, f"Q{r}")
    for k in D.cols("C", "Q"):
        cell = ac[f"{k}{r}"]
        cell.border = Border(top=D.side("ink"), bottom=D.side("hair"))
        if tint:
            cell.fill = D.fill(tint)
    D.h(ac, r, D.GRID_ROW + (3 if tint else 0))
    out = r; r += 1
    return out


for key, eyebrow, title, items in GROUPS:
    if key == "spend":
        S.page_break_before(ac, r)
    a_eyebrow(eyebrow)
    s1, s2 = ROWS[key]
    first = r
    for i in range(len(items)):
        sr = s1 + i
        c = ac[f"C{r}"]; c.value = f'=IF({SET}!$C${sr}="","",{SET}!$C${sr})'; c.font = D.f(10); c.alignment = Alignment(vertical="center")
        p = ac[f"D{r}"]; p.value = f'=IF(C{r}="","",{SET}!$H${sr})'; p.number_format = MONEY_T; p.font = D.f(10, False, "ink2")
        p.alignment = Alignment(horizontal="right", vertical="center")
        name = items[i][0]
        for m, k in enumerate(MCOLS):
            v = ac[f"{k}{r}"]
            if not BLANK and name and m < 9:
                v.value = ACT[name][m] if ACT[name][m] != 0 else 0
            v.number_format = MONEY_T; v.fill = D.fill("input_fill"); v.font = D.f(10, True, "input_text")
            v.protection = v.protection.copy(locked=False)
            v.alignment = Alignment(horizontal="right", vertical="center")
            v.border = Border(bottom=D.side("FFFFFF"), right=D.side("FFFFFF"))
        y = ac[f"Q{r}"]; y.value = f'=IF(COUNT(E{r}:P{r})=0,"",SUM(E{r}:P{r}))'; y.number_format = USD0; y.font = D.f(10, True)
        y.alignment = Alignment(horizontal="right", vertical="center")
        for k in ("C", "D", "Q"):
            ac[f"{k}{r}"].border = Border(bottom=D.side("hair"))
        D.h(ac, r, D.GRID_ROW)
        r += 1
    last = r - 1
    AROW[key] = (first, last)
    dv.add(f"E{first}:P{last}")
    ATOT[key] = a_total(f"Total {title.lower()}" if key != "inc" else "Total income", lambda k, m, a=first, b=last: f'=SUMPRODUCT(($C${a}:$C${b}<>"")*({k}{a}:{k}{b}))',
                        f"=SUM(D{first}:D{last})")
a_eyebrow("THE MONTH IN ONE LINE")
A_OUT = a_total("Total out",
                lambda k, m: f"={k}{ATOT['bill']}+{k}{ATOT['spend']}+{k}{ATOT['save']}",
                f"=D{ATOT['bill']}+D{ATOT['spend']}+D{ATOT['save']}")
A_LEFT = a_total("Income minus everything out", lambda k, m: f"={k}{ATOT['inc']}-{k}{A_OUT}", f"=D{ATOT['inc']}-D{A_OUT}",
                 tint="total_tint", signed=True)
D.h(ac, r, 12)
A_END = r
D.frame(ac, A_TOP, A_END, "B", "R")
# headline strip (rows 6 to 9)
D.stat_strip(ac, 7, [
    ("C:D", "Income logged, year", f"=$Q${ATOT['inc']}", USD0),
    ("E:H", "Total out, year", f"=$Q${A_OUT}", USD0),
    ("I:L", "Left over, year", f"=$Q${A_LEFT}", USD0),
    ("M:Q", "Months with entries", f'=SUMPRODUCT((COUNT(E{AROW["inc"][0]}:E{AROW["save"][1]})>0)*1)', INT),
])
# months with entries: count columns that hold any number
def has_col(k, sheet=""):
    return "((" + "+".join(f"COUNT({sheet}{k}{a}:{k}{b})" for a, b in AROW.values()) + ")>0)"


mc = ac["M8"]
mc.value = "=" + "+".join(f"{has_col(k)}*1" for k in MCOLS)
D.h(ac, 9, 10)
D.frame(ac, 6, 9, "B", "R")
sign_rule(ac, "I8")
D.h(ac, 10, 14)
AN_TOP = A_END + 2
an_end = notes_card(ac, AN_TOP, "B", "R", "C", "Q", [
    ("Blanks", "A blank cell means nothing yet. Type 0 when a bill was skipped on purpose."),
    ("Three-paycheck months", "Paid every 2 weeks? Two months a year hold a third paycheck. Type what really came in each month; "
     "the plan on Setup stays the average."),
    ("Rows", "Each row here is the line with the same name on the Setup tab. Rename it there and the name changes here too."),
])
AC_LAST = an_end + 1
D.h(ac, AC_LAST, 14)
fill_heights(ac, AC_LAST)
D.paint_canvas(ac, AC_LAST, "Z")
S.finish_sheet(ac, PRODUCT, AC_LAST, span=("A", "S"), freeze=f"E{AHEAD + 1}", tab_color="gold")
ac.page_setup.orientation = "landscape"
ac.print_title_rows = f"{AHEAD}:{AHEAD}"

# =====================================================================  3 Month
mo = wb.create_sheet(MON_N)
S.set_widths(mo, GRID)
D.page_header(mo, "Your month", "Pick a month. See what is left, every line against its plan, and the bills by due date.",
              AS_OF, span=("B", "K"), side_cols=3)
K = D.Card(mo, 11, "B", "C", "E", "F", wb=wb)
K.title("The month", "Leave the month blank to see this month.")
r_m = K.input("Month to show", None if BLANK else MONTHS[EX_MONTH - 1], "@", name="MonthShown", validation=("list", MONTHS),
              allow_blank=True, prompt_title="Month",
              prompt="Pick a month from the list. Leave blank for this month (January when the budget year is not this year).",
              error="Pick a month from the list: Jan to Dec.")
mo[f"E{r_m}"].alignment = Alignment(horizontal="center", vertical="center")
M = f"$F${r_m}"
hm = mo[f"F{r_m}"]
hm.value = (f'=IF($E${r_m}="",IF({YEAR}=YEAR(TODAY()),MONTH(TODAY()),1),'
            f'IFERROR(MATCH($E${r_m},{ARR_MON},0),IF({YEAR}=YEAR(TODAY()),MONTH(TODAY()),1)))')
hm.number_format = HIDE


def act(key_row_ref):
    """The actual for the shown month on an Actuals row."""
    return f"INDEX({ACTS}!$E${key_row_ref}:$P${key_row_ref},{M})"


r_mshow = K.calc("Showing", f'=INDEX({ARR_NAME},{M})&" "&{YEAR}', "@")
r_ip = K.calc("Income, planned", f"={SET}!$E${SUMS['inc']}", MONEY)
r_ia = K.calc("Income, logged", f"={act(ATOT['inc'])}", MONEY)
r_op = K.calc("Planned out", f"={SET}!$E${r_out}", MONEY)
r_oa = K.calc("Actually out so far", f"={act(A_OUT)}", MONEY, total=True)
K.close()
for rr in (r_mshow,):
    mo[f"E{rr}"].alignment = Alignment(horizontal="right", vertical="center")
LOGGED = "(" + "+".join(f"COUNT(INDEX({ACTS}!$E${a}:$P${b},0,{M}))" for a, b in AROW.values()) + ")>0"
INC_USED = f"IF($E${r_ia}>0,$E${r_ia},$E${r_ip})"

# right card: spent of plan by group
G = D.Card(mo, 11, "G", "H", "I", "K", extra="J", wb=wb)
G.title("Spent of plan", "Each group's actual as a share of its plan.")
GROW = {}
for key, color, label in (("bill", "ink", "Bills"), ("spend", "blue", "Spending"), ("save", "gold", "Savings")):
    GROW[key] = G.bar(label, f'=IF({SET}!$E${SUMS[key]}=0,"",{act(ATOT[key])}/{SET}!$E${SUMS[key]})', PCT, "1", color, width=20)
r_gall = G.bar("All of it", f'=IF($E${r_op}=0,"",$E${r_oa}/$E${r_op})', PCT, "1", "teal", width=20, bold=True)
for rr in (*GROW.values(), r_gall):
    mo[f"J{rr}"].value = f'=IF(OR(I{rr}="",I{rr}<=0),"",REPT("{D.BAR_CHAR}",MIN(16,MAX(1,ROUND(I{rr}*14,0)))))'
    mo[f"I{rr}"].number_format = '0.0%;-0.0%;0.0%'
    if rr != GROW["save"]:
        mo.conditional_formatting.add(f"J{rr}", FormulaRule(formula=[f'AND(ISNUMBER($I{rr}),$I{rr}>1)'], font=Font(color=D.T["accent"]), stopIfTrue=True))
        D.status_rule(mo, f"I{rr}", f"$I${rr}>1", fill_tint=False)
G.text("A full-length bar is 100% of the plan. Red means that group went over it; saving more than planned stays gold.", size=9, color="muted",
       height=D.GRID_ROW)
G.close()

# ---- line by line table
L_TOP = max(K.end, G.end) + 2
r = L_TOP
D.h(mo, r, 12); r += 1
t = mo[f"C{r}"]; t.value = f'="Line by line, "&INDEX({ARR_NAME},{M})&" "&{YEAR}'; t.font = D.f(12, True); mo.merge_cells(f"C{r}:J{r}")
D.h(mo, r, 24); r += 1
t = mo[f"C{r}"]; t.value = "Plan from Setup, actual from the Actuals tab. Left is plan minus actual; a minus means over."
t.font = D.f(9, False, "muted"); mo.merge_cells(f"C{r}:J{r}"); D.h(mo, r, 18); r += 1
D.h(mo, r, 6); r += 1
LHEAD = r
for k, txt, al in (("C", "LINE", "left"), ("D", "PLANNED", "right"), ("E", "ACTUAL", "right"), ("H", "LEFT", "right"),
                   ("I", "USED", "right"), ("J", "STATUS", "left")):
    c = mo[f"{k}{LHEAD}"]; c.value = txt; c.font = D.f(8, True, "muted"); c.alignment = Alignment(horizontal=al, vertical="bottom", indent=1 if k == "J" else 0)
for k in D.cols("C", "J"):
    mo[f"{k}{LHEAD}"].border = Border(bottom=D.side("ink"))
D.h(mo, LHEAD, 24)
r += 1
OVER_CELLS = []
LINE_RANGES = []
LTOT = {}


def l_eyebrow(text):
    global r
    c = mo[f"C{r}"]; c.value = text; c.font = D.f(8, True, "muted"); c.alignment = Alignment(vertical="bottom")
    for k in D.cols("C", "J"):
        mo[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    D.h(mo, r, D.GRID_ROW); r += 1


def l_row(name, plan, actual, total=False, tint=None, save=False):
    global r
    c = mo[f"C{r}"]; c.value = name; c.font = D.f(10, total); c.alignment = Alignment(vertical="center")
    for k, v in (("D", plan), ("E", actual)):
        cell = mo[f"{k}{r}"]; cell.value = v; cell.number_format = MONEY_T; cell.font = D.f(10, total, "ink2" if k == "D" else "ink")
        cell.alignment = Alignment(horizontal="right", vertical="center")
    lf = mo[f"H{r}"]; lf.value = f'=IF(C{r}="","",D{r}-E{r})'; lf.number_format = MONEY_T; lf.font = D.f(10, True)
    lf.alignment = Alignment(horizontal="right", vertical="center")
    sign_rule(mo, f"H{r}")
    us = mo[f"I{r}"]; us.value = f'=IF(OR(C{r}="",D{r}=0),"",E{r}/D{r})'; us.number_format = PCT; us.font = D.f(9, False, "ink2")
    us.alignment = Alignment(horizontal="right", vertical="center")
    sw = mo[f"J{r}"]
    OVER_WORD = "Saved more than planned" if save else "Over plan"
    sw.value = (f'=IF(C{r}="","",IF(AND(D{r}=0,E{r}>0),"Not in the plan",IF(E{r}>D{r},"{OVER_WORD}",'
                f'IF(E{r}=0,"Nothing yet",REPT("{D.BAR_CHAR}",MIN(18,MAX(1,ROUND(E{r}/D{r}*16,0))))))))')
    sw.font = D.f(9, total, "teal"); sw.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    mo.conditional_formatting.add(f"J{r}", FormulaRule(formula=[f'OR($J{r}="Over plan",$J{r}="Not in the plan")'],
                                                      font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
    mo.conditional_formatting.add(f"J{r}", FormulaRule(formula=[f'$J{r}="Nothing yet"'], font=Font(color=D.T["muted"]), stopIfTrue=True))
    for k in D.cols("C", "J"):
        cell = mo[f"{k}{r}"]
        cell.border = Border(top=D.side("ink") if total else None, bottom=D.side("hair"))
        if tint:
            cell.fill = D.fill(tint)
    D.h(mo, r, D.GRID_ROW)
    out = r; r += 1
    return out


for key, eyebrow, title, items in GROUPS[1:]:
    l_eyebrow(eyebrow)
    a1, a2 = AROW[key]
    s1, _ = ROWS[key]
    first = r
    for i in range(len(items)):
        ar = a1 + i
        rr = l_row(f"={ACTS}!$C${ar}", f'=IF({ACTS}!$C${ar}="","",{SET}!$H${s1 + i})',
                   f'=IF({ACTS}!$C${ar}="","",N({act(ar)}))', save=key == "save")
        OVER_CELLS.append(rr)
    last = r - 1
    if key != "save":
        LINE_RANGES.append((first, last))
    LTOT[key] = l_row("Total savings" if key == "save" else f"Total {title.lower()}", f"=SUM(D{first}:D{last})", f"=SUM(E{first}:E{last})",
                      total=True, save=key == "save")
L_ALL = l_row("Everything out", f"=D{LTOT['bill']}+D{LTOT['spend']}+D{LTOT['save']}", f"=E{LTOT['bill']}+E{LTOT['spend']}+E{LTOT['save']}",
              total=True, tint="total_tint")
D.h(mo, r, 12)
L_END = r
D.frame(mo, L_TOP, L_END, "B", "K")
S.page_break_before(mo, L_TOP)
N_OVER = "(" + "+".join(f'COUNTIF($J${a}:$J${b},"Over plan")+COUNTIF($J${a}:$J${b},"Not in the plan")' for a, b in LINE_RANGES) + ")"

# ---- bills by due date
B_TOP = L_END + 2
r = B_TOP
D.h(mo, r, 12); r += 1
t = mo[f"C{r}"]; t.value = f'="Bills by due date, "&INDEX({ARR_NAME},{M})'; t.font = D.f(12, True); mo.merge_cells(f"C{r}:J{r}")
D.h(mo, r, 24); r += 1
t = mo[f"C{r}"]; t.value = "Sorted by the due day on Setup. A bill counts as paid once its actual is in for the month."
t.font = D.f(9, False, "muted"); mo.merge_cells(f"C{r}:J{r}"); D.h(mo, r, 18); r += 1
D.h(mo, r, 6); r += 1
BHEAD = r
for k, txt, al in (("C", "BILL", "left"), ("D", "DUE", "right"), ("E", "AMOUNT", "right"), ("H", "PAID", "right"), ("J", "STATUS", "left")):
    c = mo[f"{k}{BHEAD}"]; c.value = txt; c.font = D.f(8, True, "muted"); c.alignment = Alignment(horizontal=al, vertical="bottom", indent=1 if k == "J" else 0)
for k in D.cols("C", "J"):
    mo[f"{k}{BHEAD}"].border = Border(bottom=D.side("ink"))
D.h(mo, BHEAD, 24)
r += 1
b1, b2 = ROWS["bill"]
KEYS = f"{SET}!$F${b1}:$F${b2}"
BILL_ROWS = []
for n in range(1, b2 - b1 + 2):
    key = f"$I${r}"
    kc = mo[f"I{r}"]; kc.value = f'=IF(COUNT({KEYS})<{n},"",SMALL({KEYS},{n}))'; kc.number_format = HIDE
    srow = f"ROUND(({key}-INT({key}))*1000,0)"
    arow = f"({srow}-{b1}+{AROW['bill'][0]})"
    cells = {
        "C": f'=IF({key}="","",INDEX({SET}!$C$1:$C${b2},{srow}))',
        "D": f'=IF({key}="","",IF(INT({key})=32,"No date","Day "&INT({key})))',
        "E": f'=IF({key}="","",INDEX({SET}!$H$1:$H${b2},{srow}))',
        "H": f'=IF({key}="","",INDEX({ACTS}!$E$1:$P${AROW["bill"][1]},{arow},{M}))',
        "J": f'=IF({key}="","",IF(N(H{r})=0,"Not paid yet",IF(H{r}>E{r},"Paid, more than planned","Paid")))',
    }
    for k, v in cells.items():
        c = mo[f"{k}{r}"]; c.value = v
        c.font = D.f(10, False, "ink2" if k == "D" else "ink")
        c.alignment = Alignment(horizontal="left" if k in ("C", "J") else "right", vertical="center", indent=1 if k == "J" else 0)
        if k in ("E", "H"):
            c.number_format = MONEY_T
    mo[f"H{r}"].value = f'=IF({key}="","",N(INDEX({ACTS}!$E$1:$P${AROW["bill"][1]},{arow},{M})))'
    mo[f"J{r}"].font = D.f(9, True, "teal")
    mo.conditional_formatting.add(f"J{r}", FormulaRule(formula=[f'$J{r}="Not paid yet"'], font=Font(color=D.T["ink2"], bold=True), stopIfTrue=True))
    mo.conditional_formatting.add(f"J{r}", FormulaRule(formula=[f'$J{r}="Paid, more than planned"'],
                                                      font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
    for k in D.cols("C", "J"):
        mo[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    D.h(mo, r, D.GRID_ROW)
    BILL_ROWS.append(r)
    r += 1
c = mo[f"C{r}"]
c.value = (f'=IF(COUNT({KEYS})=0,"No bills with an amount on Setup yet",COUNTIF(J{BILL_ROWS[0]}:J{BILL_ROWS[-1]},"Not paid yet")&" of "'
           f'&COUNT({KEYS})&" bills not paid yet this month")')
c.font = D.f(10, True, "ink2"); c.alignment = Alignment(vertical="center"); mo.merge_cells(f"C{r}:J{r}"); D.h(mo, r, D.GRID_ROW + 4)
B_STAT = r
r += 1
D.h(mo, r, 12)
B_END = r
D.frame(mo, B_TOP, B_END, "B", "K")
S.page_break_before(mo, B_TOP)

# ---- tiles (rows 6 to 9)
D.tile(mo, 6, "B", "F", f'="LEFT TO SPEND, "&UPPER(INDEX({ARR_NAME},{M}))&" "&{YEAR}', f"=$E${r_op}-$E${r_oa}", MONEY,
       f'=IF($E${r_op}=0,"Plan your month on the Setup tab first",IF(NOT({LOGGED}),"Nothing logged for "&INDEX({ARR_NAME},{M})&" yet: type it on the Actuals tab",'
       f'IF($E${r_oa}>$E${r_op},"Over the plan. Spent "&ROUND($E${r_oa}/$E${r_op}*100,0)&"% of it.",'
       f'"Spent "&ROUND($E${r_oa}/$E${r_op}*100,0)&"% of the plan so far")))', dark=True)
v_r, s_r, _ = D.tile(mo, 6, "G", "K", "INCOME LEFT AFTER EVERYTHING OUT", f"={INC_USED}-$E${r_oa}", MONEY,
                     f'=IF({INC_USED}=0,"Add your take-home pay on the Setup tab",IF($E${r_ia}=0,"Using planned income until you log what came in",'
                     f'IF({N_OVER}=0,"Every line within its plan",{N_OVER}&IF({N_OVER}=1," line over its plan"," lines over their plan"))))')
sign_rule(mo, v_r.coordinate)
D.h(mo, 10, 14)

MN_TOP = B_END + 2
mn_end = notes_card(mo, MN_TOP, "B", "K", "C", "J", [
    ("Left to spend", "Planned out for the month (bills, everyday spending and savings from Setup) minus what the Actuals tab shows "
     "went out in that month. A minus means the month went over the plan."),
    ("Income left after everything out", "What came in that month minus everything that went out. Until you log income for the "
     "month, it uses your planned income."),
    ("Not in the plan", "A line with no plan amount but money logged against it. Add a plan on Setup if it comes back."),
    ("Before you rely on a number", S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab")),
], title="How the numbers work")
MO_LAST = mn_end + 1
D.h(mo, MO_LAST, 14)
fill_heights(mo, MO_LAST)
D.paint_canvas(mo, MO_LAST, "Z")
S.finish_sheet(mo, PRODUCT, MO_LAST, span=("A", "L"), freeze="A11", tab_color="teal")
# print: the line-by-line card on one Letter or A4 page (page breaks sit before it and before the bills card)
_lines_pt = sum(mo.row_dimensions[rr].height or D.GRID_ROW for rr in range(L_TOP, L_END + 1))
mo.sheet_properties.pageSetUpPr.fitToPage = False
mo.page_setup.scale = max(40, min(100, int(680 / _lines_pt * 100)))

# =====================================================================  4 Year
yr = wb.create_sheet(YR_N)
S.set_widths(yr, {"A": 3, "B": 2, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14, "I": 20, "J": 20, "K": 2, "L": 3})
D.page_header(yr, "Your year", "Twelve months side by side: planned against actual, and what was left over.", AS_OF,
              span=("B", "K"), side_cols=3)
Y_TOP = 11
r = Y_TOP
D.h(yr, r, 12); r += 1
t = yr[f"C{r}"]; t.value = f'="Month by month, "&{YEAR}'; t.font = D.f(12, True); yr.merge_cells(f"C{r}:J{r}"); D.h(yr, r, 24); r += 1
t = yr[f"C{r}"]; t.value = "Bars show how far each month landed from the plan: teal under it, accent over it. Months with nothing logged stay empty."
t.font = D.f(9, False, "muted"); yr.merge_cells(f"C{r}:J{r}"); D.h(yr, r, 18); r += 1
D.h(yr, r, 6); r += 1
YHEAD = r
for k, txt in (("C", "MONTH"), ("D", "PLANNED OUT"), ("E", "ACTUAL OUT"), ("F", "LEFT OF PLAN"), ("G", "INCOME"),
               ("H", "LEFT OVER"), ("I", "UNDER OR OVER PLAN"), ("J", "")):
    c = yr[f"{k}{YHEAD}"]; c.value = txt; c.font = D.f(8, True, "muted")
    c.alignment = Alignment(horizontal="left" if k in ("C", "I", "J") else "right", vertical="bottom", wrap_text=True, indent=1 if k == "I" else 0)
for k in D.cols("C", "J"):
    yr[f"{k}{YHEAD}"].border = Border(bottom=D.side("ink"))
D.h(yr, YHEAD, 24)
r += 1
Y1 = r
SCALE = f"MAX(MAX($F${Y1}:$F${Y1 + 11}),-MIN($F${Y1}:$F${Y1 + 11}),1)"
for m, k in enumerate(MCOLS):
    has = has_col(k, ACTS + "!")
    vals = {
        "C": MONTH_NAMES[m],
        "D": f"={SET}!$E${r_out}",
        "E": f'=IF({has},{ACTS}!{k}{A_OUT},"")',
        "F": f'=IF(E{r}="","",D{r}-E{r})',
        "G": f'=IF({has},{ACTS}!{k}{ATOT["inc"]},"")',
        "H": f'=IF({has},{ACTS}!{k}{A_LEFT},"")',
        "I": f'=IF(OR(F{r}="",F{r}=0),"",REPT("{D.BAR_CHAR}",MAX(1,ROUND(ABS(F{r})/{SCALE}*24,0))))',
    }
    for kk, v in vals.items():
        c = yr[f"{kk}{r}"]; c.value = v
        if kk in "DEFGH":
            c.number_format = MONEY_T; c.alignment = Alignment(horizontal="right", vertical="center")
            c.font = D.f(10, kk == "H", "ink2" if kk == "D" else "ink")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1 if kk == "I" else 0)
            c.font = D.f(10) if kk == "C" else D.f(9, False, "teal")
        c.border = Border(bottom=D.side("hair"))
    sign_rule(yr, f"F{r}"); sign_rule(yr, f"H{r}")
    yr.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f'AND(ISNUMBER($F{r}),$F{r}<0)'], font=Font(color=D.T["accent"]), stopIfTrue=True))
    yr.merge_cells(f"I{r}:J{r}")
    D.h(yr, r, D.GRID_ROW)
    r += 1
YL = r - 1
c = yr[f"C{r}"]; c.value = "Year"; c.font = D.f(10, True); c.alignment = Alignment(vertical="center")
for kk, v in (("D", f'=SUMPRODUCT((E{Y1}:E{YL}<>"")*D{Y1}:D{YL})'), ("E", f"=SUM(E{Y1}:E{YL})"), ("F", f"=SUM(F{Y1}:F{YL})"),
              ("G", f"=SUM(G{Y1}:G{YL})"), ("H", f"=SUM(H{Y1}:H{YL})")):
    cc = yr[f"{kk}{r}"]; cc.value = v; cc.number_format = USD0; cc.font = D.f(10, True); cc.alignment = Alignment(horizontal="right", vertical="center")
sign_rule(yr, f"F{r}"); sign_rule(yr, f"H{r}")
for k in D.cols("C", "J"):
    yr[f"{k}{r}"].border = Border(top=D.side("ink"), bottom=D.side("hair")); yr[f"{k}{r}"].fill = D.fill("total_tint")
D.h(yr, r, D.GRID_ROW + 3)
Y_TOT = r
r += 1
c = yr[f"C{r}"]
c.value = "Planned out in the Year row counts only the months that have entries, so it compares like with like."
c.font = D.f(9, False, "muted"); yr.merge_cells(f"C{r}:J{r}"); D.h(yr, r, D.GRID_ROW)
r += 1
D.h(yr, r, 12)
Y_END = r
D.frame(yr, Y_TOP, Y_END, "B", "K")
OVERM = f'COUNTIF($F${Y1}:$F${YL},"<0")'
D.stat_strip(yr, 7, [
    ("C:D", "Income logged", f"=$G${Y_TOT}", USD0),
    ("E:F", "Total out", f"=$E${Y_TOT}", USD0),
    ("G:H", "Left over", f"=$H${Y_TOT}", USD0),
    ("I:J", "Months over plan", f'={OVERM}', INT),
])
D.h(yr, 9, 10)
D.frame(yr, 6, 9, "B", "K")
sign_rule(yr, "G8")
D.status_rule(yr, "I8", "$I$8>0", fill_tint=False)
D.h(yr, 10, 14)
YN_TOP = Y_END + 2
yn_end = notes_card(yr, YN_TOP, "B", "K", "C", "J", [
    ("Left of plan", "Planned out minus actual out for the month. A minus means that month went over the plan."),
    ("Left over", "Income logged minus everything out for the month. Savings count as out, so a left over of zero can still "
     "be a month where you saved."),
])
YR_LAST = yn_end + 1
D.h(yr, YR_LAST, 14)
fill_heights(yr, YR_LAST)
D.paint_canvas(yr, YR_LAST, "Z")
S.finish_sheet(yr, PRODUCT, YR_LAST, span=("A", "L"), freeze="A11", tab_color="blue")

# =====================================================================  Start Here and Terms
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            mm = re.match(r"^([A-Z][A-Za-z ,'0-9]{1,40})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if mm:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


def money(x):
    return ("-" if x < 0 else "") + f"${abs(x):,.2f}"


EXM = EXAMPLE[EX_MONTH]
sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, PRODUCT, "Start here. Plan once on Setup, log the month on Actuals, read the Month tab.", AS_OF,
              span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "A plain monthly budget planner. Type your take-home pay, your bills with their due days, your everyday spending and what "
    "you save on the Setup tab. Each month, type what came in and what went out on the Actuals tab. The Month tab shows what "
    "is left to spend, every line against its plan, and your bills in due-date order. The Year tab lines up all twelve months.",
    "para")])
r = sh_card(sh, r, "Three steps", [
    (1, "Setup tab. Type each paycheck and how often it comes, then your bills, spending and savings. Left to plan shows what has no job yet.", "num"),
    (2, "Actuals tab. As the month goes, type what came in and what went out under that month. Blank means nothing yet.", "num"),
    (3, "Month tab. Pick the month (blank shows this month). Read Left to spend in the dark tile and check the lines marked Over plan.", "num"),
])
r = sh_card(sh, r, "How to read the cells", [
    ((1950, MONEY), "Yellow cells with blue numbers are yours to type in. Each one shows a hint when you select it.", "chip-in"),
    ((7275, MONEY), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((EXM["left_to_spend"], MONEY),
     "Dark tiles are your answers. The main one sits at the top of the Month tab and stays on screen as you scroll.", "chip-key"),
    (None, "Nothing is locked for good: each tab is protected without a password, so you can still change anything. To edit a "
           "formula cell, use Review, Unprotect Sheet in Excel, or Data, Protect sheets and ranges in Google Sheets.", "note"),
])
S.page_break_before(sh, r)
EX1 = EXAMPLE[1]
ex_text = (f"Budget-Planner-EXAMPLE.xlsx holds January to September 2026 for a made-up household. Two paychecks ($1,950.00 "
           f"every 2 weeks and $1,400.00 twice a month) and $250.00 of side income come to {money(EX1['inc_plan'])} a month. Bills, "
           f"everyday spending and savings are planned at {money(EX1['out_plan'])}, which leaves {money(EX1['left_to_plan'])} to plan. "
           f"In September {money(EXM['out_act'])} went out, so the Month tab shows {money(EXM['left_to_spend'])} left to spend and "
           f"{money(EXM['left_after'])} of income left after everything out. January and July each hold a third paycheck.")
r = sh_card(sh, r, "The example", [(None, ex_text, "para")])
r = sh_card(sh, r, "Which file to open", [
    (None, "Excel on a computer. Double-click Budget-Planner.xlsx. If Excel shows a yellow Protected View bar, click Enable Editing.", "para"),
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. No macros, no add-ons, no sign-up.", "para"),
    (None, "Mac, iPad or phone. Unzip on a computer first (phones often can't open a .zip), then upload the .xlsx to Google Drive "
           "and open it in the Google Sheets app, or open it in Excel for Mac or iPad.", "para"),
    (None, "See it filled in first. Open Budget-Planner-EXAMPLE.xlsx. Budget-Planner.xlsx has the same line names with every "
           "amount empty, so there is no sample data to clear.", "para"),
])
r = sh_card(sh, r, "Good to know", [
    (None, "Any pay schedule. Paid weekly, every 2 weeks, twice a month or monthly? Pick it on Setup and the plan uses the monthly "
           "amount. In a month with a third paycheck, type what really came in on Actuals.", "para"),
    (None, "Your own lines. Rename, clear or add any line on Setup: 4 income lines, 12 bills, 15 spending lines and 6 savings lines. "
           "The other tabs follow.", "para"),
    (None, "Yearly costs. Give a yearly bill such as car registration the Yearly setting and the plan holds a twelfth of it each month.", "para"),
    (None, "Start a fresh copy each year. Save the file under a new name in January and clear the Actuals tab.", "para"),
    (None, "Related tools. For a budget built around each paycheck, see the shop's paycheck budget; for paying debts down, the debt payoff tracker.", "para"),
])
r = sh_card(sh, r, "Before you rely on a number", [(None, S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab"), "note")])
SH_LAST = r - 1
D.paint_canvas(sh, SH_LAST, "Z")
S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
               footer_notice="Estimates only. Not professional advice. See the Terms tab or LICENSE-AND-DISCLAIMER.txt.")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.", AS_OF,
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More tools from the shop", [
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this planner helps you run your month, a short review on Etsy helps other people find it. Thank you.", "para"),
    (None, "Message me on Etsy if anything doesn't open or looks wrong, and it will be fixed.", "para"),
])
TM_LAST = r - 1
D.paint_canvas(tm, TM_LAST, "Z")
S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb)]
if problems:
    print("AUDIT:", *problems[:40], sep="\n  ")
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
print({"setup": ROWS, "sums": SUMS, "r_out": r_out, "r_left": r_left, "r_year": r_year, "act": AROW, "atot": ATOT, "A_OUT": A_OUT,
       "A_LEFT": A_LEFT, "month": {"r_m": r_m, "r_op": r_op, "r_oa": r_oa, "r_ia": r_ia, "L_TOP": L_TOP, "bills": BILL_ROWS[:2]},
       "year": (Y1, Y_TOT)})
print("example Sep:", EXM, "Jan:", EXAMPLE[1]["inc_act"], "year:", YEAR_INC, YEAR_OUT)
