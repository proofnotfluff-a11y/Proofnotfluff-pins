#!/usr/bin/env python3
"""#18 One-Tab Business Expense Tracker, version 2 (v3 dashboard design, Oct 7, 2026).

Same inputs, categories, example and maths as version 1 (Oct 5, 2026): one Entries tab feeds a
Summary of money in, inventory and materials, Schedule C expense lines and profit before the
inventory count, month by month for the year shown, plus the same four checks. Rebuilt to
design/WORKBOOK_STANDARD.md section 0; tools/compare_xlsx.py with compare_map.json proves the
answers match.

Fixes in version 2 (each named in compare_map.json or the build notes):
- Equipment's Schedule C reference said "Lines 13 and 30"; line 30 is business use of your home.
  It now says "Line 13 (depreciation), for your preparer".
- A category name cleared on the Categories tab showed as 0 on the Summary; it now shows blank
  and its row stays 0 (v1 also stayed 0).
- "Net income" (v1's label for sales and other income less refunds) is now "Income after refunds",
  since net income usually means after expenses.
- Every tab is protected without a password (v1 left Entries, Start Here and Terms open).

The product ships two workbooks, as v1 did:
  python3 build_xlsx.py Business-Expense-Tracker-EXAMPLE.xlsx            (the worked example)
  python3 build_xlsx.py Business-Expense-Tracker.xlsx --blank           (blank, year follows today)
"""
import datetime as dt
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402

PRODUCT = "One-Tab Business Expense Tracker"
VERSION = "2"
CHECKED = "Oct 7, 2026"
CHECKED_LONG = "October 7, 2026"
AS_OF = f"Lines checked {CHECKED}"
SOURCES = "the IRS 2025 Schedule C (Form 1040) and its 2025 instructions"
HERE = os.path.dirname(os.path.abspath(__file__))
args = sys.argv[1:]
BLANK = "--blank" in args
args = [a for a in args if a != "--blank"]
OUT = args[0] if args else "Business-Expense-Tracker-EXAMPLE.xlsx"
USD0, INT, DATE = S.FMT["usd0"], S.FMT["int"], S.FMT["date"]
USD2 = S.FMT["usd2"]
MONEY_T = '"$"#,##0;-"$"#,##0;"-"'   # table cells: whole dollars, a dash for zero
YEAR_FMT = "0"

# ---------------------------------------------------------------- categories (v1 order and names)
# (counts as, category, Schedule C reference, what goes here)
CATS = [
    ("Money in", "Sales and services", "Line 1, Gross receipts or sales",
     "Money customers paid you for products or services, before fees."),
    ("Money in", "Other business income", "Line 6, Other income",
     "Business income that is not a sale, such as a fee refund paid back to you."),
    ("Refund, subtracts", "Refunds given to customers", "Line 2, Returns and allowances",
     "Money you paid back to customers. It is subtracted from income."),
    ("Inventory", "Inventory and materials for resale", "Part III, lines 36 and 38 (feeds line 4)",
     "Stock and materials that become the things you sell. Your preparer works out cost of goods sold from this and your inventory counts."),
    ("Expense", "Advertising", "Line 8, Advertising", "Ads, promoted listings, printed flyers, business cards."),
    ("Expense", "Car and truck expenses", "Line 9, Car and truck expenses",
     "Business driving costs. Keep a mileage log too; your preparer picks the method."),
    ("Expense", "Commissions and fees", "Line 10, Commissions and fees", "Sales commissions and selling fees charged on your sales."),
    ("Expense", "Contract labor", "Line 11, Contract labor", "People you paid who are not employees."),
    ("Expense", "Employee benefit programs", "Line 14, Employee benefit programs", "Benefit programs for employees (not your own)."),
    ("Expense", "Insurance (other than health)", "Line 15, Insurance (other than health)", "Business liability, product or property insurance."),
    ("Expense", "Interest", "Lines 16a and 16b, Interest", "Interest on business loans and business credit cards."),
    ("Expense", "Legal and professional services", "Line 17, Legal and professional services", "Accountant, tax preparer, lawyer, bookkeeper."),
    ("Expense", "Office expense", "Line 18, Office expense", "Postage, printer ink, small office items."),
    ("Expense", "Rent or lease: vehicles and equipment", "Line 20a, Vehicles, machinery, and equipment", "Leased vehicles, machines or equipment."),
    ("Expense", "Rent or lease: other property", "Line 20b, Other business property", "Studio, booth, storage unit or shop rent."),
    ("Expense", "Repairs and maintenance", "Line 21, Repairs and maintenance", "Fixing business equipment or space."),
    ("Expense", "Supplies", "Line 22, Supplies (not included in Part III)", "Supplies used up in the business that are not part of what you sell."),
    ("Expense", "Taxes and licenses", "Line 23, Taxes and licenses", "Business licenses, permits, business taxes."),
    ("Expense", "Travel", "Line 24a, Travel", "Overnight business trips: airfare, lodging."),
    ("Expense", "Deductible meals (enter the full bill)", "Line 24b, Deductible meals",
     "Business meals. Enter the full bill; the IRS instructions say in most cases only 50% is deductible, and your preparer applies the limit."),
    ("Expense", "Utilities", "Line 25, Utilities", "Business phone line, internet, power for a business space."),
    ("Expense", "Wages", "Line 26, Wages (less employment credits)", "Pay to employees (not to yourself)."),
    ("Expense", "Other expenses", "Line 27b, Other expenses", "Software, subscriptions, bank fees, shipping labels and anything else."),
    ("Not in profit", "Equipment over a few hundred dollars", "Line 13 (depreciation), for your preparer",
     "Computers, machines, furniture. Logged here so your preparer can decide on depreciation. Not counted in profit."),
    ("Not in profit", "Owner draw or personal", "Not a business expense",
     "Money you took out for yourself, or a personal purchase made by mistake."),
    ("Not in profit", "Transfer between accounts", "Not income or expense",
     "Moving money between your own accounts or paying off a business card."),
]
INCOME_IDX, INV_IDX, EXP_IDX, OFF_IDX = (0, 1, 2), (3,), tuple(range(4, 23)), (23, 24, 25)

# ---------------------------------------------------------------- example (v1, unchanged): a made-up one-person candle shop
d = dt.datetime
ROWS = [  # date, what it was, paid to or received from, category, amount, paid with, receipt
    (d(2026, 1, 3), 'Wax and wicks, 40 lb order', 'Craft supplier', 'Inventory and materials for resale', 412.6, 'Business card', 'Yes'),
    (d(2026, 1, 5), 'Online shop sales, week 1', 'Customers', 'Sales and services', 486, 'Deposit', None),
    (d(2026, 1, 5), 'Selling fees, week 1', 'Marketplace', 'Commissions and fees', 46.17, 'Deducted', 'Yes'),
    (d(2026, 1, 8), 'Shipping labels', 'Postal service', 'Other expenses', 58.4, 'Business card', 'Yes'),
    (d(2026, 1, 12), 'Online shop sales, week 2', 'Customers', 'Sales and services', 372, 'Deposit', None),
    (d(2026, 1, 12), 'Selling fees, week 2', 'Marketplace', 'Commissions and fees', 35.34, 'Deducted', 'Yes'),
    (d(2026, 1, 15), 'Studio rent, January', 'Landlord', 'Rent or lease: other property', 350, 'Bank transfer', 'Yes'),
    (d(2026, 1, 20), 'Jar labels and boxes', 'Packaging supplier', 'Supplies', 96.25, 'Business card', 'Yes'),
    (d(2026, 1, 22), 'Refund, broken jar', 'Customer', 'Refunds given to customers', 24, 'Refund', None),
    (d(2026, 1, 28), 'Paid myself', 'Owner', 'Owner draw or personal', 500, 'Bank transfer', None),
    (d(2026, 2, 2), 'Online shop sales, February', 'Customers', 'Sales and services', 1240, 'Deposit', None),
    (d(2026, 2, 2), 'Selling fees, February', 'Marketplace', 'Commissions and fees', 117.8, 'Deducted', 'Yes'),
    (d(2026, 2, 6), 'Promoted listings', 'Marketplace', 'Advertising', 45, 'Business card', 'Yes'),
    (d(2026, 2, 10), 'Valentine pop-up market sales', 'Customers', 'Sales and services', 615, 'Cash and card', None),
    (d(2026, 2, 10), 'Market booth fee', 'Market organizer', 'Rent or lease: other property', 75, 'Business card', 'Yes'),
    (d(2026, 2, 10), 'Parking and toll, market day', 'City garage', 'Car and truck expenses', 14, 'Cash', 'Yes'),
    (d(2026, 2, 14), 'Fragrance oils', 'Craft supplier', 'Inventory and materials for resale', 188.9, 'Business card', 'Yes'),
    (d(2026, 2, 15), 'Studio rent, February', 'Landlord', 'Rent or lease: other property', 350, 'Bank transfer', 'Yes'),
    (d(2026, 2, 18), 'Shipping labels', 'Postal service', 'Other expenses', 71.2, 'Business card', 'Yes'),
    (d(2026, 2, 21), 'City business license', 'City clerk', 'Taxes and licenses', 50, 'Business card', 'Yes'),
    (d(2026, 2, 24), 'New pouring pot set', 'Craft supplier', 'Equipment over a few hundred dollars', 329, 'Business card', 'Yes'),
    (d(2026, 3, 2), 'Online shop sales, March', 'Customers', 'Sales and services', 980, 'Deposit', None),
    (d(2026, 3, 2), 'Selling fees, March', 'Marketplace', 'Commissions and fees', 93.1, 'Deducted', 'Yes'),
    (d(2026, 3, 5), 'Wholesale order, gift shop', 'Gift shop', 'Sales and services', 540, 'Check', None),
    (d(2026, 3, 9), 'Product liability insurance', 'Insurer', 'Insurance (other than health)', 31, 'Business card', 'Yes'),
    (d(2026, 3, 12), 'Design software subscription', 'Software company', 'Other expenses', 12.99, 'Business card', 'Yes'),
    (d(2026, 3, 15), 'Studio rent, March', 'Landlord', 'Rent or lease: other property', 350, 'Bank transfer', 'Yes'),
    (d(2026, 3, 18), 'Lunch with gift shop buyer', 'Cafe', 'Deductible meals (enter the full bill)', 38.5, 'Business card', 'Yes'),
    (d(2026, 3, 20), 'Wax restock', 'Craft supplier', 'Inventory and materials for resale', 236.4, 'Business card', 'Yes'),
    (d(2026, 3, 25), 'Tax prep consult', 'Accountant', 'Legal and professional services', 150, 'Business card', 'Yes'),
    (d(2026, 3, 28), 'Paid off business card', 'Card company', 'Transfer between accounts', 1200, 'Bank transfer', None),
]
EX_YEAR = 2026


def _sum(cats, month=None):
    names = [CATS[i][1] for i in cats]
    return sum(r[4] for r in ROWS if r[3] in names and (month is None or r[0].month == month))


def _money(x):
    from decimal import Decimal, ROUND_HALF_UP
    v = int(Decimal(str(abs(x))).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return ("-" if x < 0 else "") + f"${v:,}"


EXAMPLE = dict(
    jan_sales=_sum((0, 1), 1), jan_profit=_sum((0, 1), 1) - _sum((2,), 1) - _sum(INV_IDX, 1) - _sum(EXP_IDX, 1),
    jan_inv=_sum(INV_IDX, 1), rent=350,
    income=_sum((0, 1)) - _sum((2,)), inv=_sum(INV_IDX), exp=_sum(EXP_IDX))
EXAMPLE["profit"] = EXAMPLE["income"] - EXAMPLE["inv"] - EXAMPLE["exp"]
# the Start Here example sentence quotes these; fail the build if the data ever drifts
assert (_money(EXAMPLE["jan_profit"]), _money(EXAMPLE["jan_sales"]), _money(EXAMPLE["jan_inv"])) == ("-$165", "$858", "$413"), EXAMPLE
assert (_money(EXAMPLE["income"]), _money(EXAMPLE["profit"])) == ("$4,209", "$1,386"), EXAMPLE
if BLANK:
    ROWS = []

# ---------------------------------------------------------------- fixed addresses
ENT, SUM_, CAT = "Entries", "Summary", "Categories"
N_ROWS = 1000                 # styled entry rows (v1 validated rows 5 to 1035)
E_TOP = 11
E_FIRST = E_TOP + 5
E_LAST = E_FIRST + N_ROWS - 1
E_CHECK_LAST = E_FIRST + 4999  # the checks read the first 5,000 entry rows, as in v1
C_TOP = 11
C_FIRST = C_TOP + 5
C_LAST = C_FIRST + len(CATS) - 1
CAT_LIST = f"{CAT}!$D${C_FIRST}:$D${C_LAST}"
Y_IN, Y_EFF = 13, 14          # Summary: year input row, year used row (column D)
YEAR = f"{SUM_}!$D${Y_EFF}"


def sign_rule(ws, ref):
    """Accent below zero, teal above, plain at zero (a dash), so empty months stay quiet."""
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{ref}<0"], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{ref}>0"], font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))


def ecol(c, last=None):
    return f"{ENT}!${c}${E_FIRST}:${c}${last or E_CHECK_LAST}"


wb = S.new_workbook(PRODUCT, version=VERSION)

# =====================================================================  Summary (built first: other tabs point at it)
sm = wb.create_sheet(SUM_)
MCOLS = D.cols("D", "O")
S.set_widths(sm, {"A": 3, "B": 2, "C": 36, **{k: 11.5 for k in MCOLS}, "P": 12, "Q": 38, "R": 2, "S": 3})
D.page_header(sm, "Summary", "Fills itself in from the Entries tab: every month and the year, by Schedule C line.",
              AS_OF, span=("B", "R"), side_cols=4)

# ---- year card (left) and headline strip (right), rows 11 to 16
D.h(sm, 11, 12); D.h(sm, 12, 18); D.h(sm, 13, 34); D.h(sm, 14, 21); D.h(sm, 15, 21); D.h(sm, 16, 12)
t = sm["C12"]; t.value = "YEAR SHOWN"; t.font = D.f(8, True, "muted"); t.alignment = Alignment(vertical="bottom")
lab = sm[f"C{Y_IN}"]; lab.value = "Year to show"; lab.font = D.f(10); lab.alignment = Alignment(vertical="center")
yc = S.input_cell(sm, f"D{Y_IN}", None if BLANK else EX_YEAR, YEAR_FMT, name="YearShown", wb=wb,
                  validation=("whole", 2000, 2100), allow_blank=True, prompt_title="Year to show",
                  prompt="Leave blank to show this year (it changes on its own on January 1). Type a year, for example 2025, to look back.",
                  error="Type a year from 2000 to 2100, or leave the cell blank for this year.")
yc.fill = D.fill("input_fill"); yc.font = D.f(11, True, "input_text")
yc.alignment = Alignment(horizontal="center", vertical="center")
yc.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
lab = sm[f"C{Y_EFF}"]; lab.value = "Showing"; lab.font = D.f(10); lab.alignment = Alignment(vertical="center")
ye = sm[f"D{Y_EFF}"]; ye.value = f'=IF($D${Y_IN}="",YEAR(TODAY()),$D${Y_IN})'; ye.number_format = YEAR_FMT
ye.font = D.f(10, True); ye.alignment = Alignment(horizontal="center", vertical="center")
for k in D.cols("C", "G"):
    sm[f"{k}{Y_EFF}"].border = Border(bottom=D.side("hair"), top=D.side("hair"))
n = sm["C15"]; n.value = "Blank shows this year. Type a year to look back."; n.font = D.f(9, False, "muted")
n.alignment = Alignment(vertical="center"); sm.merge_cells("C15:G15")
D.frame(sm, 11, 16, "B", "H")

# ---- the month-by-month table, from row 18
T_TOP = 18
r = T_TOP
D.h(sm, r, 12); r += 1
t = sm[f"C{r}"]; t.value = "Month by month"; t.font = D.f(12, True); sm.merge_cells(f"C{r}:Q{r}"); D.h(sm, r, 24); r += 1
t = sm[f"C{r}"]; t.value = ("Every row is a category from the Categories tab. Refunds subtract from income; the last section is logged "
                            "but kept out of profit. Estimates only, see the Terms tab.")
t.font = D.f(9, False, "muted"); t.alignment = Alignment(vertical="top"); sm.merge_cells(f"C{r}:Q{r}"); D.h(sm, r, 18); r += 1
D.h(sm, r, 6); r += 1
HEAD = r
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
for k, txt in zip(["C", *MCOLS, "P", "Q"], ["Category", *MONTHS, "Year", "Schedule C reference"]):
    c = sm[f"{k}{HEAD}"]; c.value = txt.upper(); c.font = D.f(8, True, "muted")
    c.alignment = Alignment(horizontal="left" if k in ("C", "Q") else "right", vertical="bottom", indent=1 if k == "Q" else 0)
for k in D.cols("C", "Q"):
    sm[f"{k}{HEAD}"].border = Border(bottom=D.side("ink"))
D.h(sm, HEAD, 24)
r += 1
ROW_OF = {}


def t_eyebrow(text):
    global r
    c = sm[f"C{r}"]; c.value = text; c.font = D.f(8, True, "muted"); c.alignment = Alignment(vertical="bottom")
    for k in D.cols("C", "Q"):
        sm[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    D.h(sm, r, D.GRID_ROW); r += 1


def t_cat(i):
    global r
    cr = C_FIRST + i
    c = sm[f"C{r}"]; c.value = f'=IF({CAT}!$D${cr}="","",{CAT}!$D${cr})'; c.font = D.f(10)
    c.alignment = Alignment(vertical="center")
    for m, k in enumerate(MCOLS, 1):
        v = sm[f"{k}{r}"]
        v.value = (f'=IF($C{r}="",0,SUMIFS({ENT}!$G:$G,{ENT}!$F:$F,$C{r},{ENT}!$C:$C,">="&DATE({YEAR},{m},1),'
                   f'{ENT}!$C:$C,"<"&DATE({YEAR},{m + 1},1)))')
        v.number_format = MONEY_T; v.font = D.f(10); v.alignment = Alignment(horizontal="right", vertical="center")
    y = sm[f"P{r}"]; y.value = f"=SUM(D{r}:O{r})"; y.number_format = MONEY_T; y.font = D.f(10, True)
    y.alignment = Alignment(horizontal="right", vertical="center")
    q = sm[f"Q{r}"]; q.value = f"={CAT}!$E${cr}"; q.font = D.f(9, False, "muted")
    q.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for k in D.cols("C", "Q"):
        sm[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    D.h(sm, r, D.GRID_ROW)
    ROW_OF[i] = r
    r += 1


def t_total(label, fn, tint=None, size=10):
    global r
    c = sm[f"C{r}"]; c.value = label; c.font = D.f(size, True); c.alignment = Alignment(vertical="center")
    for k in [*MCOLS, "P"]:
        v = sm[f"{k}{r}"]; v.value = fn(k); v.number_format = MONEY_T; v.font = D.f(size, True)
        v.alignment = Alignment(horizontal="right", vertical="center")
    for k in D.cols("C", "Q"):
        cell = sm[f"{k}{r}"]
        cell.border = Border(top=D.side("ink"), bottom=D.side("hair"))
        if tint:
            cell.fill = D.fill(tint)
    D.h(sm, r, D.GRID_ROW + (3 if tint else 0))
    out = r
    r += 1
    return out


t_eyebrow("MONEY IN")
for i in INCOME_IDX:
    t_cat(i)
a, b, c_ = (ROW_OF[i] for i in INCOME_IDX)
R_INC = t_total("Income after refunds", lambda k: f"={k}{a}+{k}{b}-{k}{c_}")
t_eyebrow("INVENTORY AND MATERIALS FOR RESALE")
t_cat(INV_IDX[0])
R_INV = ROW_OF[INV_IDX[0]]
t_eyebrow("BUSINESS EXPENSES")
for i in EXP_IDX:
    t_cat(i)
e0, e1 = ROW_OF[EXP_IDX[0]], ROW_OF[EXP_IDX[-1]]
R_EXP = t_total("Total expenses", lambda k: f"=SUM({k}{e0}:{k}{e1})")
R_PRO = t_total("Profit before inventory count", lambda k: f"={k}{R_INC}-{k}{R_INV}-{k}{R_EXP}", tint="total_tint")
for k in [*MCOLS, "P"]:
    sign_rule(sm, f"{k}{R_PRO}")
nt = sm[f"C{r}"]
nt.value = "Your preparer adjusts this for inventory left on hand at year end (cost of goods sold, Schedule C Part III)."
nt.font = D.f(9, False, "muted"); nt.alignment = Alignment(vertical="center"); sm.merge_cells(f"C{r}:Q{r}")
D.h(sm, r, D.GRID_ROW); r += 1
t_eyebrow("LOGGED FOR YOUR RECORDS, NOT IN PROFIT")
for i in OFF_IDX:
    t_cat(i)
D.h(sm, r, 12)
T_END = r
D.frame(sm, T_TOP, T_END, "B", "R")
S.page_break_before(sm, T_TOP)  # print: answers on page 1, the whole month table on page 2

# ---- checks (left) and profit by month (right)
K_TOP = T_END + 2
K = D.Card(sm, K_TOP, "B", "C", "D", "H", wb=wb)
K.title("Checks", "The first three should read 0. Rows in other years are fine.")
EA = f"{ENT}!$G${E_FIRST}:$G${E_CHECK_LAST}"
ED = f"{ENT}!$F${E_FIRST}:$F${E_CHECK_LAST}"
EC = f"{ENT}!$C${E_FIRST}:$C${E_CHECK_LAST}"
k1 = K.calc("Amount but no category", f'=COUNTIFS({EA},"<>",{ED},"")', INT)
k2 = K.calc("Amount but no date", f'=COUNTIFS({EA},"<>",{EC},"")', INT)
k3 = K.calc("Category not on the list", f'=SUMPRODUCT(({ED}<>"")*(COUNTIF({CAT_LIST},{ED})=0))', INT)
k4 = K.calc("Dated outside the year shown", f'=COUNT({EC})-COUNTIFS({EC},">="&DATE({YEAR},1,1),{EC},"<"&DATE({YEAR}+1,1,1))', INT)
k5 = K.calc("Problems to fix", f"=D{k1}+D{k2}+D{k3}", INT, total=True)
for rr in (k1, k2, k3, k5):
    D.status_rule(sm, f"D{rr}", f"$D${rr}>0", fill_tint=False)
pill = K.text(f'=IF(COUNT({EA})=0,"No entries yet",IF(D{k5}>0,D{k5}&IF(D{k5}=1," problem"," problems")&" to fix before the totals are complete",'
              f'IF(D{k4}>0,"Every row is counted. "&D{k4}&" dated "&IF(D{k4}=1,"row is","rows are")&" in other years, which is fine","Every row is counted")))',
              size=10, color="ink", bold=True, height=D.GRID_ROW * 2)
sm[f"C{pill}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
K.text("The checks read the first 5,000 entry rows. Totals read every row.", size=9, color="muted")
K.close()
D.status_rule(sm, f"C{pill}:D{pill}", f"$D${k5}>0")

P = D.Card(sm, K_TOP, "J", "K", "N", "R", extra="O", wb=wb)
P.title("Profit by month", "Profit before inventory count. Accent bars are losses.")
p_first = P.r
PR = f"$N${p_first}:$N${p_first + 11}"
MONTH_NAMES = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
               "November", "December"]
for m, k in enumerate(MCOLS):
    rr = P.calc(MONTH_NAMES[m], f"={k}{R_PRO}", MONEY_T)
    b = sm[f"O{rr}"]
    b.value = f'=IFERROR(REPT("{D.BAR_CHAR}",ROUND(ABS(N{rr})/MAX(MAX({PR}),-MIN({PR}))*14,0)),"")'
    b.font = D.f(9, False, "teal"); b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    sm.conditional_formatting.add(f"O{rr}", FormulaRule(formula=[f"$N{rr}<0"], font=Font(color=D.T["accent"]), stopIfTrue=True))
    sign_rule(sm, f"N{rr}")
    for kk in ("P", "Q"):
        sm[f"{kk}{rr}"].border = Border(bottom=D.side("hair"))
p_last = p_first + 11
LOSS = f'COUNTIF({PR},"<0")'
pt = P.text(f'=IF(COUNTIF({PR},"<>0")=0,"No entries in the year shown yet",IF({LOSS}=0,"No month shows a loss",'
            f'{LOSS}&IF({LOSS}=1," month shows"," months show")&" a loss before the inventory count"))', size=9, color="ink2")
P.close()
S.page_break_before(sm, K_TOP)

# ---- hero tiles (rows 6 to 9)
FIX = f"$D${k5}"
D.tile(sm, 6, "B", "H", f'="PROFIT BEFORE INVENTORY COUNT, "&{YEAR}', f"=$P${R_PRO}", USD0,
       f'="Income after refunds "&TEXT($P${R_INC},"$#,##0")&", less "&TEXT($P${R_INV},"$#,##0")&" of stock and "'
       f'&TEXT($P${R_EXP},"$#,##0")&" of expenses"', dark=True)
v, sub, _ = D.tile(sm, 6, "J", "R", f'="INCOME AFTER REFUNDS, "&{YEAR}', f"=$P${R_INC}", USD0,
                   f'=IF(COUNT({EA})=0,"No entries yet",IF({FIX}>0,{FIX}&IF({FIX}=1," problem"," problems")&" to fix: see Checks below","Every entry is counted"))')
sm.conditional_formatting.add(sub.coordinate, FormulaRule(formula=[f"{FIX}>0"], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
D.h(sm, 10, 14)

# ---- headline strip in the right card, rows 11 to 16
D.stat_strip(sm, 12, [
    ("K:M", "Inventory and materials", f"=$P${R_INV}", USD0),
    ("N:P", "Business expenses", f"=$P${R_EXP}", USD0),
    ("Q:Q", "Problems to fix", f"={FIX}", INT),
])
D.status_rule(sm, "Q13", "$Q$13>0", fill_tint=False)
s14 = sm["K15"]; s14.value = "For the year shown. Equipment, draws and transfers are logged but kept out of profit."
s14.font = D.f(9, False, "muted"); s14.alignment = Alignment(vertical="center"); sm.merge_cells("K15:Q15")
D.frame(sm, 11, 16, "J", "R")

# ---- notes and sources
N_TOP = max(K.end, P.end) + 2
NOTES = [
    ("Categories", "The names follow the expense lines on the IRS 2025 Schedule C (Form 1040), irs.gov/pub/irs-pdf/f1040sc.pdf, "
     f"checked {CHECKED}. They are for organizing, not tax advice; your preparer decides the final line for each item. "
     "The IRS draft 2026 Schedule C (dated May 15, 2026) splits Interest into 16a Mortgage, 16b Vehicle loan and 16c Other; "
     "the Interest row stays one line here."),
    ("Business meals", "Enter the full bill. The 2025 Schedule C instructions (irs.gov/pub/irs-pdf/i1040sc.pdf, revised Jan 2, 2026) say "
     f"that in most cases only 50% of business meals is deductible; your preparer applies the limit. Checked {CHECKED}."),
    ("Not calculated here", "Ask your tax preparer about line 12 Depletion, line 13 Depreciation and section 179, line 19 Pension and "
     "profit-sharing plans, line 27a Energy efficient commercial buildings and line 30 Business use of your home, and about your "
     "inventory count at year end (cost of goods sold)."),
    ("Any currency", "The maths works in any currency. Select the money cells and pick your symbol under Format, Number."),
    ("Before you rely on a number", S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab") + " "
     + S.TAX_TIER_CLAUSE.format(date=CHECKED_LONG, sources=SOURCES)),
]
r = N_TOP
D.h(sm, r, 12); r += 1
t = sm[f"C{r}"]; t.value = "Notes and sources"; t.font = D.f(12, True); sm.merge_cells(f"C{r}:Q{r}"); D.h(sm, r, 24); r += 1
D.h(sm, r, 6); r += 1
width = sum(S.width_of(sm, k) for k in D.cols("C", "Q"))
for head, body in NOTES:
    cell = sm[f"C{r}"]
    cell.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=9, b=True, color=D.T["ink"]), head + ". "),
                              TextBlock(InlineFont(rFont=D.FONT, sz=9, color=D.T["ink2"]), body))
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    sm.merge_cells(f"C{r}:Q{r}")
    D.h(sm, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
D.h(sm, r, 12)
N_END = r
D.frame(sm, N_TOP, N_END, "B", "R")
SM_LAST = N_END + 1
D.h(sm, SM_LAST, 14)
for rr in range(1, SM_LAST):
    if sm.row_dimensions[rr].height is None:
        D.h(sm, rr, D.GRID_ROW)
D.paint_canvas(sm, SM_LAST, "Z")
S.finish_sheet(sm, PRODUCT, SM_LAST, span=("A", "S"), freeze="A11", tab_color="accent")
sm.page_setup.orientation = "landscape"

# =====================================================================  Entries
en = wb.create_sheet(ENT, 0)
S.set_widths(en, {"A": 3, "B": 2, "C": 13, "D": 30, "E": 22, "F": 36, "G": 12, "H": 15, "I": 10, "J": 22, "K": 16,
                  "L": 2, "M": 3})
D.page_header(en, "Entries", "One row per money in or money out. Type amounts as positive numbers; the category decides the sign.",
              "Example rows: type over them" if not BLANK else "Type your first row in row 16", span=("B", "L"), side_cols=3)
D.h(en, 6, 10)
D.stat_strip(en, 7, [
    ("C:D", "Income after refunds, year shown", f"={SUM_}!$P${R_INC}", USD0),
    ("E:F", "Spent on stock and expenses", f"={SUM_}!$P${R_INV}+{SUM_}!$P${R_EXP}", USD0),
    ("G:I", "Profit before inventory count", f"={SUM_}!$P${R_PRO}", USD0),
    ("J:K", "Problems to fix", f"={SUM_}!$D${k5}", INT),
])
D.h(en, 9, 10)
D.frame(en, 6, 9, "B", "L")
sign_rule(en, "G8")
D.status_rule(en, "J8", "$J$8>0", fill_tint=False)
D.h(en, 10, 14)


def counts_as(r):
    return (f'=IF(G{r}="","",IF(F{r}="","Needs a category",IF(C{r}="","Needs a date",'
            f'IFERROR(INDEX({CAT}!$C${C_FIRST}:$C${C_LAST},MATCH(F{r},{CAT_LIST},0)),"Not on the list"))))')


cols_en = [
    dict(col="C", head="Date", kind="input", fmt=DATE, values=[x[0] for x in ROWS],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The day the money moved, for example 3/14/2026. It shows the month by name so you can check it.",
         error="Type a real date between 2000 and 2100, for example 3/14/2026."),
    dict(col="D", head="What it was", kind="input", values=[x[1] for x in ROWS], validation=("textLength", 0, 80),
         prompt="A few words you will recognise later, for example Wax restock.", error="Up to 80 characters."),
    dict(col="E", head="Paid to or from", kind="input", values=[x[2] for x in ROWS], validation=("textLength", 0, 60),
         prompt="Who you paid, or who paid you. Optional.", error="Up to 60 characters."),
    dict(col="F", head="Category (pick from list)", kind="input", values=[x[3] for x in ROWS], validation=("list_range", CAT_LIST),
         prompt="Choose a category from the list. You can rename categories on the Categories tab.",
         error="Choose a category from the list. To add a name, rename one on the Categories tab first."),
    dict(col="G", head="Amount", kind="input", fmt=USD2, align="right", values=[x[4] for x in ROWS],
         validation=("decimal", 0, None), prompt="Type the amount as a positive number. The category decides if it adds or subtracts.",
         error="Type 0 or more, as a positive number. The category decides if it adds or subtracts."),
    dict(col="H", head="Paid with", kind="input", values=[x[5] for x in ROWS], validation=("textLength", 0, 30),
         prompt="Optional: card, cash, bank transfer, deposit.", error="Up to 30 characters."),
    dict(col="I", head="Receipt?", kind="input", values=[x[6] for x in ROWS], validation=("list", ["Yes", "No"]),
         prompt="Optional: do you have the receipt saved?", error="Pick Yes or No, or leave it blank."),
    dict(col="J", head="Notes", kind="input", values=[], validation=("textLength", 0, 120),
         prompt="Anything worth remembering about this row. Optional.", error="Up to 120 characters."),
    dict(col="K", head="Counts as", kind="calc", formula=counts_as),
]
f1, f2, en_end = D.table_card(en, E_TOP, "B", "L", cols_en, N_ROWS, title="Your entries",
                              sub=f"Yellow columns are yours. Counts as fills in on its own. {N_ROWS:,} rows; the totals read every row.")
assert (f1, f2) == (E_FIRST, E_LAST), (f1, f2)
for rr in range(E_FIRST, E_LAST + 1):
    en[f"K{rr}"].font = D.f(9, False, "ink2"); en[f"K{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for k in ("D", "E", "F", "H", "J"):
        en[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    en[f"I{rr}"].alignment = Alignment(horizontal="center", vertical="center")
en.conditional_formatting.add(f"C{E_FIRST}:F{E_LAST}", FormulaRule(
    formula=[f'AND($G{E_FIRST}<>"",OR($C{E_FIRST}="",$F{E_FIRST}=""))'], fill=PatternFill("solid", bgColor=D.T["accent_tint"])))
for word, color in (("Needs", "accent"), ("Not on", "accent"), ("Not in", "muted")):
    en.conditional_formatting.add(f"K{E_FIRST}:K{E_LAST}", FormulaRule(
        formula=[f'LEFT($K{E_FIRST},{len(word)})="{word}"'], font=Font(color=D.T[color], bold=color == "accent"), stopIfTrue=True))
D.paint_canvas(en, en_end + 1, "Z")
S.finish_sheet(en, PRODUCT, en_end + 1, span=("A", "M"), freeze=f"A{E_FIRST}", tab_color="gold")
en.page_setup.orientation = "landscape"
en.print_title_rows = f"{E_FIRST - 1}:{E_FIRST - 1}"
en.print_area = f"A1:M{E_FIRST + 99}"  # print the first 100 rows, not 1,000 bare ones

# =====================================================================  Categories
ct = wb.create_sheet(CAT)
S.set_widths(ct, {"A": 3, "B": 2, "C": 16, "D": 36, "E": 36, "F": 60, "G": 2, "H": 3})
D.page_header(ct, "Categories", "Rename any yellow cell. The Entries dropdown and the Summary follow.", AS_OF,
              span=("B", "G"), side_cols=2)
D.h(ct, 6, 10)
_used = (f'SUMPRODUCT(({CAT}!$D${C_FIRST}:$D${C_LAST}<>"")*(COUNTIFS({ENT}!$F:$F,{CAT}!$D${C_FIRST}:$D${C_LAST},'
         f'{ENT}!$C:$C,">="&DATE({YEAR},1,1),{ENT}!$C:$C,"<"&DATE({YEAR}+1,1,1))>0))')
D.stat_strip(ct, 7, [
    ("C:C", "On the list", f'=COUNTA({CAT_LIST})', INT),
    ("D:D", "Used in the year shown", f"={_used}", INT),
    ("E:F", "Entries with a name not on the list", f"={SUM_}!$D${k3}", INT),
])
D.h(ct, 9, 10)
D.frame(ct, 6, 9, "B", "G")
D.status_rule(ct, "E8", "$E$8>0", fill_tint=False)
D.h(ct, 10, 14)
cols_ct = [
    dict(col="C", head="Counts as", kind="text", values=[c[0] for c in CATS]),
    dict(col="D", head="Category", kind="input", values=[c[1] for c in CATS], validation=("textLength", 0, 40),
         prompt="Rename this category if you like (up to 40 characters). Rows already typed with the old name show up in Checks so you can update them.",
         error="Up to 40 characters."),
    dict(col="E", head="Schedule C reference", kind="text", values=[c[2] for c in CATS]),
    dict(col="F", head="What goes here", kind="text", values=[c[3] for c in CATS]),
]
g1, g2, ct_end = D.table_card(ct, C_TOP, "B", "G", cols_ct, len(CATS), title="Your categories",
                              sub="Grouped as the Summary groups them. Names follow the IRS 2025 Schedule C expense lines.")
assert (g1, g2) == (C_FIRST, C_LAST), (g1, g2)
for i, rr in enumerate(range(C_FIRST, C_LAST + 1)):
    for k in ("C", "E", "F"):
        ct[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ct[f"C{rr}"].font = D.f(9, True, "muted" if CATS[i][0] == "Not in profit" else "ink2")
    ct[f"E{rr}"].font = D.f(9, False, "ink2")
    ct[f"D{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    D.h(ct, rr, max(D.GRID_ROW, S.text_height(CATS[i][3], S.width_of(ct, "F"), 10) + 6,
                    S.text_height(CATS[i][2], S.width_of(ct, "E"), 9) + 6))
CN_TOP = ct_end + 2
r = CN_TOP
D.h(ct, r, 12); r += 1
t = ct[f"C{r}"]; t.value = "Notes and sources"; t.font = D.f(12, True); ct.merge_cells(f"C{r}:F{r}"); D.h(ct, r, 24); r += 1
cw = sum(S.width_of(ct, k) for k in "CDEF")
for head, body in [
    ("Source", f"IRS 2025 Schedule C (Form 1040), irs.gov/pub/irs-pdf/f1040sc.pdf, and its instructions, irs.gov/pub/irs-pdf/i1040sc.pdf "
     f"(revised Jan 2, 2026), checked {CHECKED}. For organizing, not tax advice; your preparer decides the final line for each item."),
    ("Not calculated in this tracker", "Line 12 Depletion, line 13 Depreciation and section 179, line 19 Pension and profit-sharing "
     "plans, line 27a Energy efficient commercial buildings, line 30 Business use of your home. Ask your tax preparer."),
    ("Renaming", "Rename a yellow cell and the Entries dropdown and the Summary follow. Rows already typed with the old name are "
     "counted in Checks until you pick the new name. To keep the Summary in order, rename rather than reorder."),
]:
    cell = ct[f"C{r}"]
    cell.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=9, b=True, color=D.T["ink"]), head + ". "),
                              TextBlock(InlineFont(rFont=D.FONT, sz=9, color=D.T["ink2"]), body))
    cell.alignment = Alignment(vertical="top", wrap_text=True); ct.merge_cells(f"C{r}:F{r}")
    D.h(ct, r, S.text_height(head + ". " + body, cw, 9) + 4); r += 1
D.h(ct, r, 12)
D.frame(ct, CN_TOP, r, "B", "G")
CT_LAST = r + 1
D.h(ct, CT_LAST, 14)
D.paint_canvas(ct, CT_LAST, "Z")
S.finish_sheet(ct, PRODUCT, CT_LAST, span=("A", "H"), tab_color="teal")
ct.page_setup.orientation = "landscape"


# =====================================================================  Start Here and Terms (card style)
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            mm = re.match(r"^([A-Z][A-Za-z ,]{1,34})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if mm:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, PRODUCT, "Start here. You type in one tab; the Summary fills itself in.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "A small business expense tracker you can learn in five minutes. Type one row for each money in or money out on the Entries "
    "tab. The Summary tab adds it up by month and for the year: income after refunds, inventory and materials, expenses by "
    "Schedule C line, and profit before the inventory count, plus a Checks box that catches rows missing something.", "para")])
r = sh_card(sh, r, "Three steps", [
    (1, "Summary tab. Check the year in the yellow box. Leave it blank to show this year, or type a year to look back.", "num"),
    (2, "Entries tab. One row per money in or money out: date, what it was, who, category from the list, amount as a positive number.", "num"),
    (3, "Summary tab. Read your profit in the dark tile at the top and each month below it. Checks should all read 0.", "num"),
])
r = sh_card(sh, r, "How to read the cells", [
    ((412.6, USD2), "Yellow cells with blue numbers are yours to type in. Each one shows a hint when you select it.", "chip-in"),
    ((1240, USD0), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((round(EXAMPLE["profit"]), USD0),
     "Dark tiles are your answers. The main one sits at the top of the Summary tab and stays on screen as you scroll.", "chip-key"),
    (None, "Each tab is protected without a password, so you can still change anything. To edit a locked cell, use Review, "
           "Unprotect Sheet in Excel, or Data, Protect sheets and ranges in Google Sheets.", "note"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "The example", [(None,
    "Business-Expense-Tracker-EXAMPLE.xlsx holds three months of a made-up one-person candle shop. January shows a $165 loss "
    "on $858 of sales, because a $413 wax order and $350 studio rent land in the same month. January to March shows $4,209 of "
    "income after refunds and $1,386 of profit before the inventory count. The $329 pouring pot set, the $500 the owner paid "
    "themselves and the $1,200 card payoff are logged but kept out of profit.", "para")])
r = sh_card(sh, r, "Which file to open", [
    (None, "Excel on a computer. Double-click Business-Expense-Tracker.xlsx. If Excel shows a yellow Protected View bar, click Enable Editing.", "para"),
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. No macros, no add-ons, no sign-up.", "para"),
    (None, "On a phone. Unzip on a computer first, then upload the .xlsx to Google Drive and open it in the Google Sheets app. "
           "Phones often can't open a .zip file.", "para"),
    (None, "See it filled in first. Open Business-Expense-Tracker-EXAMPLE.xlsx, then use the blank file for your own business.", "para"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "Good to know", [
    (None, f"Room for a busy year. The Entries tab has {N_ROWS:,} rows and every total reads the whole column. Printing the "
           "Entries tab prints the first 100 rows; to print more, select the rows you want and print the selection.", "para"),
    (None, "Renaming categories. Rename any yellow cell on the Categories tab and the dropdown and the Summary follow. Rows "
           "already typed with the old name show up in Checks so you can update them.", "para"),
    (None, "Kept out of profit. Equipment, owner draws and transfers between your own accounts are logged but kept out of profit. "
           "Depreciation and home office use are not calculated here; hand those to your tax preparer.", "para"),
    (None, "Back it up. Save a copy each month (File, Save a copy) so you always have a backup.", "para"),
])
r = sh_card(sh, r, "Before you rely on a number", [(None, S.SHORT_NOTICE.format(date=CHECKED) + " "
           + S.TAX_TIER_CLAUSE.format(date=CHECKED_LONG, sources=SOURCES), "note")])
SH_LAST = r - 1
D.paint_canvas(sh, SH_LAST, "Z")
S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
               footer_notice="Estimates only. Not professional advice. See the Terms tab or LICENSE-AND-DISCLAIMER.txt.")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More tools from the shop", [
    (None, ("Freelance Income Tracker: the income side, with set-aside by IRS payment period", "https://www.etsy.com/listing/4589995309"), "link"),
    (None, ("Hourly Rate Quick Calculator: the rate that pays the take-home you want", "https://www.etsy.com/listing/4589695837"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this tracker saves you time, a short review on Etsy helps other small business owners find it. Thank you.", "para"),
    (None, "If a file won't open or something looks wrong, message the shop on Etsy and it will be fixed.", "para"),
])
TM_LAST = r - 1
D.paint_canvas(tm, TM_LAST, "Z")
S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb) if "input without validation" not in p]
if problems:
    print("AUDIT:", *problems, sep="\n  ")
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
print({"year_in": f"Summary!D{Y_IN}", "inc": R_INC, "inv": R_INV, "exp": R_EXP, "profit": R_PRO, "rows": ROW_OF,
       "checks": (k1, k2, k3, k4, k5), "pbm": (p_first, p_last), "E_FIRST": E_FIRST, "C_FIRST": C_FIRST})
print("example:", {k: round(v, 2) for k, v in EXAMPLE.items()})
