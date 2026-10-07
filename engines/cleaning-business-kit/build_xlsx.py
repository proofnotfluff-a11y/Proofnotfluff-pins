#!/usr/bin/env python3
"""#14 Cleaning Business Starter Kit, workbook version 3 (dashboard design, Oct 7, 2026).

Same inputs, example and formulas as v2 (Oct 6, 2026, the Legal desk build), rebuilt to
design/WORKBOOK_STANDARD.md section 0 with tools/pnf_dash.py. tools/compare_xlsx.py with
compare_map.json proves the answers match.

Tab changes (standard: at most 6 working tabs): 1 Business Info is now the "Your business
details" card on 1 Your Numbers; 8 Money by Month sits at the top of 6 Job Log; Thank You
moved into Terms. Price Sheet comes before Quote Builder so the forms PDF's "tab 4 Quote
Builder" still holds. Fixes in v3: every tab protected without a password (v2 had none, so a
stray keystroke could overwrite a formula); the frequency names are locked (renaming "One time"
silently broke the one-time logic); the quote check asks for the home first instead of running
on an empty home; a blank quote date no longer prints "dated Dec 30, 1899" in the quote; the price guide title no longer reads
" price guide" when the business name is blank.

Usage: python3 engines/cleaning-business-kit/build_xlsx.py out.xlsx
"""
import datetime as dt
import json
import math
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

PRODUCT = "Cleaning Business Starter Kit"
VERSION = "3"
CHECKED = "Oct 7, 2026"
CHECKED_LONG = "October 7, 2026"
TAX_SOURCES = ("IRS Tax Topic 554 (Self-employment tax), IRS Tax Topic 751 (Social Security and Medicare withholding rates), "
               "the IRS Standard Mileage Rates page (Internal Revenue Bulletin 2026-29) and the SSA Contribution and Benefit Base page")
FORMS_CLAUSE = ("These are sample forms to adapt, not legal documents written for your state or business. Have a local attorney "
                "review any agreement before you use it with clients. If you collect client information, you're responsible for "
                "keeping it secure and following the privacy laws that apply to you.")
AS_OF = f"Figures checked {CHECKED}"
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else "Cleaning-Business-Pricing-and-Planner.xlsx"
USD0, USD2, PCT, INT, NUM1, NUM2 = (S.FMT[k] for k in ("usd0", "usd2", "pct", "int", "num1", "num2"))
PCT1, PCT2, DATE = "0.0%", "0.00%", S.FMT["date"]
MONTH = "mmm yyyy"
X2 = '0.00"x"'

# ---------------------------------------------------------------- example data (v2, unchanged)
EX = dict(take_home=48000, weeks=48, hours=45, share=0.7, se=0.153, se_base=0.9235, income=0.12, margin=0.15,
          per_mile=0.76, mph=25, helper_pay=18, payroll=0.12)
COSTS = [("Insurance and bond", 110, "Liability insurance and bond, per month."),
         ("Cleaning products, paper", 90, "Cleaning products and paper goods you keep in stock, per month."),
         ("Vacuums, mops, cloths", 60, "What you spend replacing vacuums, mops and cloths, spread per month."),
         ("Laundry for cloths, mops", 20, "Laundry for cloths and mop heads, per month."),
         ("Phone and internet", 70, "The business share of your phone and internet, per month."),
         ("Booking or invoicing app", 30, "Booking, scheduling or invoicing app, per month."),
         ("Marketing, ads, website", 75, "Marketing, ads and website, per month."),
         ("Bank and card fees", 40, "Bank and card processing fees, per month."),
         ("Licenses and permits", 10, "Licenses and permits, spread per month."),
         ("Vehicle payment", 0, "Only if you set Vehicle cost per mile to 0. Otherwise the per-mile rate already covers the vehicle."),
         ("Other", 0, "Any other monthly business cost. Type 0 if none.")]
BIZ = [("Business name", "Example Home Cleaning", "Shows on the quote and the price sheet."),
       ("Your name", "Your Name", "Your name, for your own records."),
       ("Phone", "(555) 010-0100", "Shows on the quote and the price sheet. 555-01xx numbers are reserved for examples."),
       ("Email", "hello@example.com", "Shows on the quote and the price sheet."),
       ("Website or social page", "example.com", "Shows on the quote."),
       ("Service area", "Example City and nearby towns", "Where you clean, for your price sheet or posts."),
       ("Payment methods you accept", "Card, bank transfer, check", "Shows on the quote."),
       ("Supplies line on the quote", "We bring all supplies and equipment", "Or: Client provides supplies. Shows on the quote and the price sheet.")]
POLICY = dict(notice=24, fee=40, valid=30)
SQFT_MIN = 90
ROOMS = [("Bedroom", 20), ("Full bathroom", 30), ("Half bathroom", 15), ("Kitchen", 40), ("Living or family area", 20),
         ("Other room (office, dining)", 15)]
PET_MIN = 10
SERVICES = [("Standard clean", 1, 8), ("Deep clean", 1.75, 15), ("Move-in or move-out (empty home)", 2, 20), ("Light touch-up", 0.75, 5)]
CONDITIONS = [("Well kept", 0.9), ("Average", 1), ("Needs extra work", 1.25), ("Heavy buildup", 1.5)]
FREQS = [("One time", 0, "=1"), ("Monthly", 0.05, "=1"), ("Every 2 weeks", 0.1, "=26/12"), ("Weekly", 0.15, "=52/12")]
MIN_CHARGE, TYP_MILES = 120, 16
ADDONS = [("Inside oven", 30, 2), ("Inside fridge", 30, 1), ("Inside kitchen cabinets", 45, 1), ("Interior windows, per window", 5, 0),
          ("Baseboards, whole home", 40, 1), ("Change bed linens, per bed", 10, 0), ("Laundry, wash and fold, per load", 15, 1),
          ("Blinds, per window", 6, 0), (None, None, None), (None, None, None)]
HOME = dict(client="Example: Rivera", date=dt.datetime(2026, 10, 5), sqft=1800, rooms=[3, 2, 1, 1, 1, 1], pets=1,
            condition="Average", service="Standard clean", freq="Every 2 weeks", cleaners=2, helpers=1, miles=16)
QTY = [None, 1, None, 8, None, None, None, None, None, None]
SIZES = [1000, 1500, 2000, 2500, 3000, 3500]
CLIENTS = [("Example: Rivera", "North side", "(555) 010-0111", "Every 2 weeks", "Tue", 215, "Client home, garage door", "1 dog",
            dt.datetime(2026, 10, 7), "Use their vacuum upstairs"),
           ("Example: Chen", "Downtown", "(555) 010-0122", "Monthly", "Fri", 175, "Lockbox, ask by text", "None",
            dt.datetime(2026, 10, 9), None)]
JOBS = [(dt.datetime(2026, 10, 7), "Example: Rivera", "Deep clean", 410, 7, 3.5, 1.3, 16, 18),
        (dt.datetime(2026, 10, 9), "Example: Chen", "Standard clean", 175, 2.5, 0, 0.5, 12, 8)]
YEAR = 2026
N_CLIENTS, N_JOBS = 100, 200

# ---------------------------------------------------------------- the example's answers (for the Start Here text)
_profit = EX["take_home"] / (1 - EX["se"] * EX["se_base"] - EX["income"])
_costs_y = sum(c[1] for c in COSTS) * 12
_bill = EX["weeks"] * EX["hours"] * EX["share"]
BE_X = (_profit + _costs_y) / _bill
TARGET_X = BE_X / (1 - EX["margin"])
GUESS_X = EX["take_home"] / (EX["weeks"] * EX["hours"])
_hc = EX["helper_pay"] * (1 + EX["payroll"])
_base = max(HOME["sqft"] / 1000 * SQFT_MIN / 60, sum(n * m for n, (_, m) in zip(HOME["rooms"], ROOMS)) / 60) + HOME["pets"] * PET_MIN / 60
_addh = sum((q or 0) * a[1] for q, a in zip(QTY, ADDONS) if a[0]) / 60
_adds = sum((q or 0) * a[2] for q, a in zip(QTY, ADDONS) if a[0])
_drive = HOME["miles"] / EX["mph"] * HOME["cleaners"]
_rate = ((HOME["cleaners"] - HOME["helpers"]) * BE_X + HOME["helpers"] * _hc) / HOME["cleaners"]
_veh = HOME["miles"] * EX["per_mile"]


def _floor(mult, sup):
    return (_base * 1 * mult + _addh + _drive) * _rate + _veh + sup + _adds


def ceil5(x):
    return math.ceil(x / 5 - 1e-9) * 5


FLOOR_X = _floor(1, 8)
ONE_X = max(MIN_CHARGE, ceil5(FLOOR_X / (1 - EX["margin"])))
VISIT_X = max(MIN_CHARGE, ceil5(ONE_X * (1 - 0.1)))
DEEP_X = max(MIN_CHARGE, ceil5(_floor(1.75, 15) / (1 - EX["margin"])))


def money(x, cents=True):
    return f"${x:,.2f}" if cents else f"${x:,.0f}"


wb = S.new_workbook(PRODUCT, version=VERSION)
GRID = {"A": 3, "B": 2, "C": 31, "D": 14, "E": 2, "F": 3, "G": 2, "H": 24, "I": 12, "J": 17, "K": 2, "L": 3}
CELLS = {}  # addresses for compare_map and the listing kit


def rich(head, body, size=9, head_color="ink", body_color="ink2"):
    return CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=size, b=True, color=D.T[head_color]), head + ". "),
                        TextBlock(InlineFont(rFont=D.FONT, sz=size, color=D.T[body_color]), body))


def notes_card(ws, top, notes, span=("C", "J"), frame_span=("B", "K"), title="Notes and sources"):
    r = top
    D.h(ws, r, 12); r += 1
    t = ws[f"{span[0]}{r}"]; t.value = title; t.font = D.f(12, True)
    ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}"); D.h(ws, r, 24); r += 1
    D.h(ws, r, 6); r += 1
    width = sum(S.width_of(ws, k) for k in D.cols(*span))
    for head, body in notes:
        cell = ws[f"{span[0]}{r}"]
        cell.value = rich(head, body)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}")
        D.h(ws, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
    D.h(ws, r, 12)
    D.frame(ws, top, r, *frame_span)
    return r


def wide_input(ws, r, label, value, fmt, prompt, title=None, validation=("textLength", 0, 80), span=("D", "J"),
               error="Up to 80 characters.", lst=None, height=26, size=10, align="left"):
    """Label in C, one wide input merged across span (the reference job-picker pattern, for text and long list items)."""
    lab = ws[f"C{r}"]; lab.value = label; lab.font = D.f(10); lab.alignment = Alignment(vertical="center")
    c = ws[f"{span[0]}{r}"]; c.value = value; c.number_format = fmt; c.font = D.f(size, True, "input_text")
    c.alignment = Alignment(horizontal=align, vertical="center", indent=1)
    c.protection = c.protection.copy(locked=False)
    cs = D.cols(*span)
    for k in cs:
        cc = ws[f"{k}{r}"]; cc.fill = D.fill("input_fill")
        cc.border = Border(top=D.side("input_line"), bottom=D.side("input_line"),
                           left=D.side("input_line") if k == cs[0] else None, right=D.side("input_line") if k == cs[-1] else None)
    ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}"); D.h(ws, r, height)
    if lst:
        dv = DataValidation(type="list", formula1=lst, allow_blank=False)
        dv.error = error
    else:
        kind, lo, hi = validation
        dv = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
        dv.error = error
    dv.showErrorMessage = True; dv.errorStyle = "stop"; dv.errorTitle = "Check this cell"
    dv.showInputMessage = True; dv.promptTitle = (title or label)[:32]; dv.prompt = prompt[:255]
    ws.add_data_validation(dv); dv.add(f"{span[0]}{r}")
    return r


def fill_heights(ws, last):
    for rr in range(1, last + 1):
        if ws.row_dimensions[rr].height is None:
            D.h(ws, rr, D.GRID_ROW)


# =====================================================================  1 Your Numbers
yn = wb.create_sheet("1 Your Numbers")
S.set_widths(yn, GRID)
D.page_header(yn, "Your Numbers", "Your pay, hours, costs and taxes, turned into your rate per cleaning hour.", AS_OF)
TOP = 11

A = D.Card(yn, TOP, "B", "C", "D", "E", wb=wb)
A.title("Pay and time", "You, the owner. Type over the yellow cells.")
TH = A.input("Take-home pay you want a year", EX["take_home"], USD0, name="TakeHome", validation=("decimal", 0, None),
             prompt=f"What you want to keep for yourself in a year, after taxes, in dollars. Example: {EX['take_home']}.",
             prompt_title="Take-home pay per year")
WK = A.input("Weeks you work per year", EX["weeks"], INT, name="WeeksPerYear", validation=("decimal", 1, 52),
             prompt=f"52 minus vacation, holidays and sick weeks. From 1 to 52. Example: {EX['weeks']}.")
HW = A.input("Hours you work per week, all of it", EX["hours"], INT, name="HoursPerWeek", validation=("decimal", 1, 100),
             prompt="Include driving, walkthroughs, texts, invoicing, laundry and supply runs. From 1 to 100. Example: 45.",
             prompt_title="Hours per week, all of it")
SH = A.input("Share cleaning or driving to jobs", EX["share"], PCT, name="BillableShare", validation=("decimal", 0.05, 1),
             prompt="Share of those hours spent cleaning or driving to jobs (every quote charges for the drive). Not sure? Try 65% to 75%.",
             prompt_title="Billable share of hours")
BH = A.calc("Billable hours per year", f"=D{WK}*D{HW}*D{SH}", INT, bold=True)
A.close()


B = D.Card(yn, TOP, "G", "H", "I", "K", extra="J", wb=wb)
B.title("How your rate is built", "Your pay grossed up for tax, plus costs, over billable hours.")
b_th = B.calc("Take-home you want", f"=$D${TH}", USD0)
b_tax = B.calc("Plus tax (SE and income)", "", USD0)
b_profit = B.calc("Profit you need before tax", "", USD0, total=True)
b_cost = B.calc("Plus business costs", "", USD0)
b_rev = B.calc("Revenue you need a year", f"=I{b_profit}+I{b_cost}", USD0, total=True)
B.eyebrow("SPREAD OVER BILLABLE HOURS")
b_bh = B.calc("Billable hours per year", f"=$D${BH}", INT)
b_be = B.calc("Break-even per clean hour", f"=IFERROR(I{b_rev}/I{b_bh},0)", USD2, total=True)
b_mg = B.calc("Plus your profit margin", "", USD2)
b_tg = B.calc("Target per cleaning hour", "", USD2, total=True, tint="total_tint")
B.eyebrow("WHERE EACH CLEANING HOUR GOES")
w1 = B.r
wmx = f"MAX($I${w1}:$I${w1 + 3})"
B.bar("Your pay", f"=IFERROR($I${b_th}/$I${b_bh},0)", USD2, wmx, "teal")
B.bar("Tax", f"=IFERROR($I${b_tax}/$I${b_bh},0)", USD2, wmx, "accent")
B.bar("Business costs", f"=IFERROR($I${b_cost}/$I${b_bh},0)", USD2, wmx, "gold")
B.bar("Profit margin", f"=$I${b_mg}", USD2, wmx, "blue")
B.eyebrow("THE USUAL GUESS")
b_guess = B.calc("Pay divided by all hours", f"=IFERROR($D${TH}/($D${WK}*$D${HW}),0)", USD2)
b_gap = B.calc("Guess underprices by", f"=I{b_tg}-I{b_guess}", USD2, bold=True)
B.close()

T2 = D.Card(yn, A.end + 2, "B", "C", "D", "E", wb=wb)
T2.title("Taxes, profit, travel, helpers", "Select a cell to see its source.")
SE = T2.input("Self-employment tax rate", EX["se"], PCT1, name="SETaxRate", validation=("decimal", 0, 0.2),
              prompt=f"15.3%: 12.4% Social Security plus 2.9% Medicare. IRS Tax Topic 554, checked {CHECKED}. From 0% to 20%.")
SB = T2.input("SE tax applies to", EX["se_base"], PCT2, name="SETaxBase", validation=("decimal", 0, 1),
              prompt="IRS Topic 554: self-employment tax generally applies to 92.35% of net earnings. From 0% to 100%.")
IT = T2.input("Income tax rate, estimate", EX["income"], PCT1, name="IncomeTaxRate", validation=("decimal", 0, 0.5),
              prompt="Federal plus state, as an effective rate. An estimate: ask your tax preparer for yours. From 0% to 50%.")
MG = T2.input("Profit margin in quotes", EX["margin"], PCT, name="ProfitMargin", validation=("decimal", 0, 0.9),
              prompt="Built into every quote for slow months, new equipment and growth. From 0% to 90%. Example: 15%.")
CPM = T2.input("Vehicle cost per mile", EX["per_mile"], USD2, name="CostPerMile", validation=("decimal", 0, 5),
               prompt=f"IRS business rate 76 cents a mile for Jul 1 to Dec 31, 2026 (IRS Bulletin 2026-29), checked {CHECKED}. Or your own cost.")
SPD = T2.input("Speed between jobs, mph", EX["mph"], INT, name="DriveSpeed", validation=("decimal", 1, 90),
               prompt="Average driving speed between jobs. Turns miles into drive time. From 1 to 90. Example: 25.")
HP = T2.input("Helper pay per hour", EX["helper_pay"], USD2, name="HelperPay", validation=("decimal", 0, 200),
              prompt="What you pay a helper per hour. Leave as is if you clean alone. From $0 to $200.")
PR = T2.input("Payroll cost on helper pay", EX["payroll"], PCT, name="PayrollCost", validation=("decimal", 0, 1),
              prompt="At least 7.65% employer Social Security and Medicare (IRS Topic 751), plus unemployment tax and workers' comp. Ask your payroll provider.",
              prompt_title="Payroll cost on top of pay")
HC = T2.calc("Helper cost per hour", f"=D{HP}*(1+D{PR})", USD2, bold=True)
T2.close()
COSTS_TOP = max(T2.end, B.end) + 2
C2 = D.Card(yn, COSTS_TOP, "B", "C", "D", "E", wb=wb)
C2.title("Business costs per month", "Costs that run every month, busy or slow.")
cost_rows = [C2.input(lab, val, USD0, validation=("decimal", 0, None), prompt=pr, prompt_title=lab[:32]) for lab, val, pr in COSTS]
CM = C2.calc("Total per month", f"=SUM(D{cost_rows[0]}:D{cost_rows[-1]})", USD0, total=True)
CY = C2.calc("Total per year", f"=D{CM}*12", USD0, bold=True)
C2.close()
yn[f"I{b_cost}"] = f"=$D${CY}"

yn[f"I{b_profit}"] = f"=IFERROR($D${TH}/(1-$D${SE}*$D${SB}-$D${IT}),0)"
yn[f"I{b_tax}"] = f"=I{b_profit}-I{b_th}"
yn[f"I{b_tg}"] = f"=IFERROR(I{b_be}/(1-$D${MG}),0)"
yn[f"I{b_mg}"] = f"=I{b_tg}-I{b_be}"

# business details, full width under both columns
bd_top = max(C2.end, T2.end) + 2
r = bd_top
D.h(yn, r, 12); r += 1
t = yn[f"C{r}"]; t.value = "Your business details"; t.font = D.f(12, True); yn.merge_cells(f"C{r}:J{r}"); D.h(yn, r, 24); r += 1
t = yn[f"C{r}"]; t.value = "Type these once. The quote and the price sheet fill in from here."; t.font = D.f(9, False, "muted")
t.alignment = Alignment(vertical="top"); yn.merge_cells(f"C{r}:J{r}"); D.h(yn, r, 18); r += 1
biz_rows = {}
for lab, val, pr in BIZ:
    biz_rows[lab] = wide_input(yn, r, lab, val, "@", pr)
    D.h(yn, r + 1, 6); r += 2


def policy_row(r, label, value, fmt, prompt, kind, lo, hi, error):
    lab = yn[f"C{r}"]; lab.value = label; lab.font = D.f(10); lab.alignment = Alignment(vertical="center")
    c = S.input_cell(yn, f"D{r}", value, fmt, validation=(kind, lo, hi), prompt=prompt, prompt_title=label[:32], error=error)
    c.fill = D.fill("input_fill"); c.font = D.f(10, True, "input_text")
    c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    c.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
    D.h(yn, r, 24)
    return r


NOTICE = policy_row(r, "Cancellation notice, hours", POLICY["notice"], INT, "Hours of notice you ask for to cancel or reschedule. Whole hours, 0 to 168.",
                    "whole", 0, 168, "Enter whole hours from 0 to 168."); D.h(yn, r + 1, 6); r += 2
FEE = policy_row(r, "Late cancel or lockout fee", POLICY["fee"], USD0, "Your policy, your amount, in dollars. From $0 to $1,000.",
                 "decimal", 0, 1000, "Enter dollars from 0 to 1,000."); D.h(yn, r + 1, 6); r += 2
VALID = policy_row(r, "Quotes are valid for, days", POLICY["valid"], INT, "Days a quote stays valid. Whole days, 1 to 365.",
                   "whole", 1, 365, "Enter whole days from 1 to 365."); r += 1
D.h(yn, r, 12)
D.frame(yn, bd_top, r, "B", "K")
BIZ_REF = {k: f"'1 Your Numbers'!$D${v}" for k, v in biz_rows.items()}

YN = "'1 Your Numbers'!"
BE, TG = (f"{YN}$I${x}" for x in (b_be, b_tg))
CPMR, SPDR, MGR, HCR, SER, SBR, ITR = (f"{YN}$D${x}" for x in (CPM, SPD, MG, HC, SE, SB, IT))

D.tile(yn, 6, "B", "E", "YOUR TARGET RATE PER CLEANING HOUR", f"=$I${b_tg}", USD2,
       f'="Break-even "&TEXT($I${b_be},"$#,##0.00")&". The Quote Builder prices from this."', dark=True)
v, _, _ = D.tile(yn, 6, "G", "K", "THE USUAL GUESS UNDERPRICES EACH HOUR BY", f"=$I${b_gap}", USD2,
                 f'="The guess: "&TEXT($I${b_guess},"$#,##0.00")&" an hour, your pay over all hours"')
D.status_rule(yn, v.coordinate, f"$I${b_gap}>0", fill_tint=False)
D.h(yn, 10, 14)

n_top = r + 2
NOTES = [
    ("Self-employment tax", "15.3% (12.4% Social Security plus 2.9% Medicare), generally on 92.35% of net earnings, per IRS Tax Topic 554 "
     "(irs.gov/taxtopics/tc554). The Social Security part stops at the 2026 wage base of $184,500 (SSA Contribution and Benefit Base "
     "page, ssa.gov/oact/cola/cbb.html), so at high profits your real tax is lower. The planner does not take the income tax deduction "
     f"for one-half of self-employment tax, so it slightly overstates the income tax part and stays on the safe side. Checked {CHECKED}."),
    ("Vehicle cost per mile", "Defaults to the IRS business standard mileage rate of 76 cents a mile for July 1 to December 31, 2026 "
     "(Internal Revenue Bulletin 2026-29, irs.gov/tax-professionals/standard-mileage-rates); it was 72.5 cents from January 1 to June 30, 2026. "
     f"Checked {CHECKED}. Use your own cost per mile if you track it."),
    ("Helper payroll", "Employers pay 6.2% Social Security and 1.45% Medicare on wages (7.65%), plus unemployment taxes and any "
     f"workers' compensation, per IRS Tax Topic 751 (irs.gov/taxtopics/tc751). Checked {CHECKED}. The 12% example is a starting point; "
     "ask your payroll provider for your real rate."),
    ("Income tax", "Your own estimate of your effective federal plus state rate. Ask your tax preparer. Rates change; check them each year."),
    ("Before you rely on a number", "For general information and planning only. This is not legal, tax, financial, medical or other "
     "professional advice, and using it doesn't create a professional relationship. Results are estimates based on the numbers you enter. "
     f"Figures were checked on {CHECKED} and can change. See the Terms tab. " + S.TAX_TIER_CLAUSE.format(date=CHECKED_LONG, sources=TAX_SOURCES)),
]
n_end = notes_card(yn, n_top, NOTES)
S.page_break_before(yn, COSTS_TOP)
S.page_break_before(yn, bd_top)
S.page_break_before(yn, n_top)
D.h(yn, n_end + 1, 14)
fill_heights(yn, n_end)
D.paint_canvas(yn, n_end + 1, "Z")
S.finish_sheet(yn, PRODUCT, n_end + 1, span=("A", "L"), freeze="A11", tab_color="accent")
CELLS["yn"] = dict(TH=TH, WK=WK, HW=HW, SH=SH, BH=BH, CM=CM, CY=CY, costs=cost_rows, SE=SE, SB=SB, IT=IT, MG=MG, CPM=CPM, SPD=SPD,
                   HP=HP, PR=PR, HC=HC, profit=b_profit, rev=b_rev, be=b_be, tg=b_tg, guess=b_guess, gap=b_gap,
                   biz=biz_rows, NOTICE=NOTICE, FEE=FEE, VALID=VALID)

# =====================================================================  2 Cleaning Rates
cr = wb.create_sheet("2 Cleaning Rates")
S.set_widths(cr, {"A": 3, "B": 2, "C": 34, "D": 15, "E": 15, "F": 15, "G": 2, "H": 3})
D.page_header(cr, "Cleaning Rates", "Example starting points, not market prices. Time your own cleans.",
              "Your times, your prices", span=("B", "G"), side_cols=2)
K1 = D.Card(cr, 6, "B", "C", "D", "G", extra="F", wb=wb)
K1.title("Time per home", "Two estimates; the Quote Builder uses the larger. Pet time is added to either.")
K1.eyebrow("BY SQUARE FEET")
SQ = K1.input("Person-minutes per 1,000 sq ft", SQFT_MIN, INT, name="MinutesPer1000SqFt", validation=("decimal", 0, 1000),
              prompt="Standard clean, one person. Example: 1,800 sq ft x 90 / 1,000 = 162 person-minutes (2.7 hours). From 0 to 1,000.",
              prompt_title="Minutes per 1,000 sq ft")
K1.eyebrow("BY ROOM, PERSON-MINUTES, STANDARD CLEAN")
room_rows = [K1.input(lab, m, INT, validation=("decimal", 0, 600), prompt=f"Person-minutes for one {lab.lower().split(' (')[0]} on a standard clean. From 0 to 600.",
                      prompt_title=lab[:32]) for lab, m in ROOMS]
PET = K1.input("Per pet in the home", PET_MIN, INT, validation=("decimal", 0, 600),
               prompt="Extra person-minutes per pet, added on top of either estimate. From 0 to 600.")
K1.close()

tsv_cols = [
    dict(col="C", head="Service type", kind="input", values=[s[0] for s in SERVICES], validation=("textLength", 1, 40),
         prompt="The name shows in the Quote Builder and Job Log lists. Up to 40 characters.", error="Type a name of 1 to 40 characters."),
    dict(col="D", head="Time multiplier", kind="input", fmt=X2, align="right", values=[s[1] for s in SERVICES],
         validation=("decimal", 0, 10), prompt="Times a standard clean. Deep: 1.75 means 75% longer. From 0 to 10.", error="Enter a multiplier from 0 to 10."),
    dict(col="E", head="Supplies per visit", kind="input", fmt=USD0, align="right", values=[s[2] for s in SERVICES],
         validation=("decimal", 0, 1000), prompt="Supplies used up on one visit, in dollars. From $0 to $1,000.", error="Enter dollars from 0 to 1,000."),
]
sv_first, sv_last, sv_end = D.table_card(cr, K1.end + 2, "B", "G", tsv_cols, 4, title="Service types",
                                         sub="Deep adds baseboards and fixtures. Move-out adds inside cabinets.")
cond_cols = [
    dict(col="C", head="Condition", kind="input", values=[c[0] for c in CONDITIONS], validation=("textLength", 1, 40),
         prompt="The name shows in the Quote Builder list. Up to 40 characters.", error="Type a name of 1 to 40 characters."),
    dict(col="D", head="Time multiplier", kind="input", fmt=X2, align="right", values=[c[1] for c in CONDITIONS],
         validation=("decimal", 0, 10), prompt="Times the time for an average home. From 0 to 10.", error="Enter a multiplier from 0 to 10."),
]
cd_first, cd_last, cd_end = D.table_card(cr, sv_end + 2, "B", "G", cond_cols, 4, title="Condition of the home",
                                         sub="How the home looks at the walkthrough.")
freq_cols = [
    dict(col="C", head="How often", kind="text", values=[f[0] for f in FREQS]),
    dict(col="D", head="Discount", kind="input", fmt=PCT, align="right", values=[f[1] for f in FREQS],
         validation=("decimal", 0, 0.9), prompt="Comes off the one-time price. Type a percent, for example 10%. From 0% to 90%.",
         error="Enter a percent from 0% to 90%."),
    dict(col="E", head="Visits per month", kind="calc", fmt=NUM2, align="right", formula=lambda rr: FREQS[rr - fr_first_holder[0]][2]),
]
fr_first_holder = [cd_end + 2 + 5]
fr_first, fr_last, fr_end = D.table_card(cr, cd_end + 2, "B", "G", freq_cols, 4, title="How often",
                                         sub="The Quote Builder warns you when a discount drops a visit below your floor.")
assert fr_first == fr_first_holder[0]
for rr in range(fr_first, fr_last + 1):
    cr[f"C{rr}"].font = D.f(10, False, "ink")

K2 = D.Card(cr, fr_end + 2, "B", "C", "D", "G", extra="F", wb=wb)
K2.title("Minimums and travel", "Every quote is at least the minimum charge.")
MINC = K2.input("Minimum charge per visit", MIN_CHARGE, USD0, name="MinimumCharge", validation=("decimal", 0, 5000),
                prompt="No quote goes below this, in dollars. From $0 to $5,000. Example: 120.")
TMI = K2.input("Typical round-trip miles", TYP_MILES, INT, name="TypicalMiles", validation=("decimal", 0, 500),
               prompt="Used for the Price Sheet only. Each quote takes its own miles. From 0 to 500.")
K2.close()


def addon_price(rr):
    return f'=IF(C{rr}="","",CEILING(D{rr}/60*{TG}+E{rr},1))'


ad_cols = [
    dict(col="C", head="Add-on", kind="input", values=[a[0] for a in ADDONS], validation=("textLength", 0, 40),
         prompt="Name an extra job. It shows in the Quote Builder and the Price Sheet. Up to 40 characters.", error="Up to 40 characters."),
    dict(col="D", head="Person-minutes each", kind="input", fmt=INT, align="right", values=[a[1] for a in ADDONS],
         validation=("decimal", 0, 600), prompt="Minutes for one person, for one of this add-on. From 0 to 600.", error="Enter minutes from 0 to 600."),
    dict(col="E", head="Supplies each", kind="input", fmt=USD0, align="right", values=[a[2] for a in ADDONS],
         validation=("decimal", 0, 1000), prompt="Supplies for one of this add-on, in dollars. From $0 to $1,000.", error="Enter dollars from 0 to 1,000."),
    dict(col="F", head="Price each", kind="calc", fmt=USD0, align="right", bold=True, formula=addon_price),
]
ad_first, ad_last, ad_end = D.table_card(cr, K2.end + 2, "B", "G", ad_cols, len(ADDONS), title="Add-ons",
                                         sub="Minutes at your target rate plus supplies, rounded up to the dollar. Two empty rows for yours.")
for rr in range(ad_first, ad_last + 1):
    for k in "CDE":
        if cr[f"{k}{rr}"].value is None:
            cr[f"{k}{rr}"].value = None
CR_LAST = ad_end + 1
D.h(cr, CR_LAST, 14)
S.page_break_before(cr, sv_end + 2)
S.page_break_before(cr, K2.end + 2)
fill_heights(cr, CR_LAST)
D.paint_canvas(cr, CR_LAST, "Z")
S.finish_sheet(cr, PRODUCT, CR_LAST, span=("A", "H"), tab_color="teal")
CRS = "'2 Cleaning Rates'!"
SV_NAMES, SV_MULT, SV_SUP = (f"{CRS}${k}${sv_first}:${k}${sv_last}" for k in "CDE")
CD_NAMES, CD_MULT = (f"{CRS}${k}${cd_first}:${k}${cd_last}" for k in "CD")
FR_NAMES, FR_DISC, FR_VISITS = (f"{CRS}${k}${fr_first}:${k}${fr_last}" for k in "CDE")
AD_NAMES, AD_MIN, AD_SUP, AD_PRICE = (f"{CRS}${k}${ad_first}:${k}${ad_last}" for k in "CDEF")
SQR, PETR, MINR, TMIR = (f"{CRS}$D${x}" for x in (SQ, PET, MINC, TMI))
ROOMR = [f"{CRS}$D${x}" for x in room_rows]
SV_STD_MULT, SV_DEEP_MULT, SV_MOVE_MULT = (f"{CRS}$D${sv_first + i}" for i in range(3))
SV_STD_SUP, SV_DEEP_SUP, SV_MOVE_SUP = (f"{CRS}$E${sv_first + i}" for i in range(3))
FR_2W_DISC, FR_W_DISC = f"{CRS}$D${fr_first + 2}", f"{CRS}$D${fr_first + 3}"
CELLS["cr"] = dict(SQ=SQ, rooms=room_rows, PET=PET, sv_first=sv_first, cd_first=cd_first, fr_first=fr_first, MINC=MINC, TMI=TMI,
                   ad_first=ad_first)

# =====================================================================  3 Price Sheet
ps = wb.create_sheet("3 Price Sheet")
S.set_widths(ps, {"A": 3, "B": 2, "C": 30, "D": 14, "E": 14, "F": 14, "G": 14, "H": 16, "I": 2, "J": 3})
D.page_header(ps, "Price Sheet", "A price menu by home size for your website or posts. Quote each home in tab 4.", AS_OF,
              span=("B", "I"), side_cols=3)
PS_SUB = "One cleaner, average home, no pets, typical miles, your margin, up to $5. Estimates only, see Terms."


def ps_price(mult, sup):
    def fn(rr):
        return (f"=IF(C{rr}=\"\",\"\",MAX({MINR},CEILING(((C{rr}/1000*{SQR}/60*{mult}+IFERROR({TMIR}/{SPDR},0))*{BE}"
                f"+{TMIR}*{CPMR}+{sup})/(1-{MGR}),5)))")
    return fn


def ps_disc(disc):
    return lambda rr: f'=IF(C{rr}="","",MAX({MINR},CEILING(D{rr}*(1-{disc}),5)))'


ps_cols = [
    dict(col="C", head="Home size, sq ft", kind="input", fmt=INT, align="right", values=SIZES, validation=("decimal", 0, 100000),
         prompt="Square feet for this row of the menu. From 0 to 100,000. Leave blank to hide the row.", error="Enter square feet from 0 to 100,000."),
    dict(col="D", head="One-time standard", kind="calc", fmt=USD0, align="right", bold=True, formula=ps_price(SV_STD_MULT, SV_STD_SUP)),
    dict(col="E", head="Every 2 weeks", kind="calc", fmt=USD0, align="right", formula=ps_disc(FR_2W_DISC)),
    dict(col="F", head="Weekly", kind="calc", fmt=USD0, align="right", formula=ps_disc(FR_W_DISC)),
    dict(col="G", head="Deep clean", kind="calc", fmt=USD0, align="right", formula=ps_price(SV_DEEP_MULT, SV_DEEP_SUP)),
    dict(col="H", head="Move-out", kind="calc", fmt=USD0, align="right", formula=ps_price(SV_MOVE_MULT, SV_MOVE_SUP)),
]
GUIDE_TITLE = f'=IF({BIZ_REF["Business name"]}="","Your price guide",{BIZ_REF["Business name"]}&" price guide")'
p_first, p_last, p_end = D.table_card(ps, 6, "B", "I", ps_cols, len(SIZES), title=GUIDE_TITLE, sub=PS_SUB)
pa_cols = [
    dict(col="C", head="Add-on", kind="calc", formula=lambda rr: f'=IF({CRS}C{rr - pa_first_h[0] + ad_first}="","",{CRS}C{rr - pa_first_h[0] + ad_first})'),
    dict(col="D", head="Price each", kind="calc", fmt=USD0, align="right", formula=lambda rr: f"={CRS}F{rr - pa_first_h[0] + ad_first}"),
]
pa_first_h = [p_end + 2 + 5]
pa_first, pa_last, pa_end = D.table_card(ps, p_end + 2, "B", "I", pa_cols, len(ADDONS), title="Add-ons",
                                         sub="Priced per item from tab 2.")
assert pa_first == pa_first_h[0]
r = pa_end + 2
D.h(ps, r, 12); top_f = r; r += 1
fl = ps[f"C{r}"]
fl.value = (f'="Minimum visit $"&TEXT({MINR},"#,##0")&".  "&{BIZ_REF["Supplies line on the quote"]}&".  "&'
            f'{BIZ_REF["Phone"]}&"  |  "&{BIZ_REF["Email"]}')
fl.font = D.f(10, True); fl.alignment = Alignment(vertical="center", wrap_text=True)
ps.merge_cells(f"C{r}:H{r}"); D.h(ps, r, 30); PS_FOOT = r; r += 1
D.h(ps, r, 12)
D.frame(ps, top_f, r, "B", "I")
PS_LAST = r + 1
D.h(ps, PS_LAST, 14)
fill_heights(ps, PS_LAST)
D.paint_canvas(ps, PS_LAST, "Z")
S.finish_sheet(ps, PRODUCT, PS_LAST, span=("A", "J"), freeze=None, tab_color="blue")
ps.page_setup.fitToHeight = 1
CELLS["ps"] = dict(first=p_first, add_first=pa_first, foot=PS_FOOT, title_row=7)

# =====================================================================  4 Quote Builder
qb = wb.create_sheet("4 Quote Builder")
S.set_widths(qb, GRID)
D.page_header(qb, "Quote Builder", "Enter the home from your walkthrough and read your prices. Estimates only, see the Terms tab.", AS_OF)
r = 11
D.h(qb, r, 12); r += 1
t = qb[f"C{r}"]; t.value = "The service"; t.font = D.f(12, True); qb.merge_cells(f"C{r}:J{r}"); D.h(qb, r, 24); r += 1
t = qb[f"C{r}"]; t.value = "Pick from the lists. They come from the Cleaning Rates tab."; t.font = D.f(9, False, "muted")
t.alignment = Alignment(vertical="top"); qb.merge_cells(f"C{r}:J{r}"); D.h(qb, r, 18); r += 1
Q_SV = wide_input(qb, r, "Service type", HOME["service"], "@", "Pick from the list. It comes from the Service types table on tab 2.",
                  lst=SV_NAMES, error="Pick a service type from the list."); D.h(qb, r + 1, 6); r += 2
Q_FR = wide_input(qb, r, "How often", HOME["freq"], "@", "Pick from the list. It comes from the How often table on tab 2.",
                  lst=FR_NAMES, error="Pick how often from the list."); D.h(qb, r + 1, 6); r += 2
Q_CD = wide_input(qb, r, "Condition of the home", HOME["condition"], "@", "Pick from the list. It comes from the Condition table on tab 2.",
                  lst=CD_NAMES, error="Pick a condition from the list."); r += 1
D.h(qb, r, 12)
D.frame(qb, 11, r, "B", "K")
CARD_TOP = r + 2

L1 = D.Card(qb, CARD_TOP, "B", "C", "D", "E", wb=wb)
L1.title("The home", "From your walkthrough. Blank counts as 0.")
Q_SQ = L1.input("Square feet", HOME["sqft"], INT, name="SquareFeet", validation=("decimal", 0, 100000),
                prompt="The home's square feet. From 0 to 100,000. Example: 1800.")
q_rooms = []
for (lab, _), n in zip(ROOMS, HOME["rooms"]):
    plural = {"Bedroom": "Bedrooms", "Full bathroom": "Full bathrooms", "Half bathroom": "Half bathrooms", "Kitchen": "Kitchens",
              "Living or family area": "Living or family areas", "Other room (office, dining)": "Other rooms"}[lab]
    q_rooms.append(L1.input(plural, n, INT, validation=("whole", 0, 50), prompt=f"How many. Whole number, 0 to 50.",
                            prompt_title=plural[:32]))
Q_PET = L1.input("Pets", HOME["pets"], INT, validation=("whole", 0, 20), prompt="Pets in the home. Whole number, 0 to 20.")
L1.eyebrow("THE TEAM AND THE TRIP")
Q_CL = L1.input("Cleaners on this job", HOME["cleaners"], INT, validation=("whole", 1, 20),
                prompt="You plus helpers. On-site clock time is shared by the team. Whole number, 1 to 20.")
Q_HE = L1.input("Of those, helpers paid hourly", HOME["helpers"], INT, validation=("whole", 0, 20),
                prompt="Helpers cost the helper rate from tab 1; your own hours cost your break-even rate. Whole number, 0 to 20.",
                prompt_title="Helpers paid hourly")
Q_MI = L1.input("Round-trip miles", HOME["miles"], INT, validation=("decimal", 0, 1000),
                prompt="Miles there and back for this home. From 0 to 1,000.")
L1.close()

L2 = D.Card(qb, L1.end + 2, "B", "C", "D", "E", wb=wb)
L2.title("Add-ons for this home", "Type a quantity next to any add-on you quote.")
q_add = []
for i in range(len(ADDONS)):
    rr = L2.r
    q_add.append(L2.input("", QTY[i], INT, validation=("decimal", 0, 100), allow_blank=True,
                          prompt="How many of this add-on for this home. Leave blank for none. From 0 to 100.", prompt_title="Quantity"))
    src = ad_first + i
    lab = qb[f"C{rr}"]
    lab.value = f'=IF({CRS}C{src}="","",{CRS}C{src}&" ($"&TEXT({CRS}F{src},"#,##0")&")")'
    lab.font = D.f(9)
QA0, QA1 = q_add[0], q_add[-1]
L2.close()

L3 = D.Card(qb, L2.end + 2, "B", "C", "D", "E", wb=wb)
L3.title("Check", "Price per visit against your floor and target.")
pill = L3.text("", size=10, color="ink", bold=True)
pill2 = L3._row()
qb.unmerge_cells(f"C{pill}:D{pill}")
qb.merge_cells(f"C{pill}:D{pill2}")
qb[f"C{pill}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
L3.close()

R1 = D.Card(qb, CARD_TOP, "G", "H", "I", "K", extra="J", wb=wb)
R1.title("How this quote is built", "Person-hours at your rates, plus the trip and supplies.")
hdr = R1._row()
for k, txt in (("I", "THIS SERVICE"), ("J", "FIRST VISIT, DEEP")):
    c = qb[f"{k}{hdr}"]; c.value = txt; c.font = D.f(8, True, "muted"); c.alignment = Alignment(horizontal="right", vertical="bottom")
for k in ("H", "I", "J"):
    qb[f"{k}{hdr}"].border = Border(bottom=D.side("ink"))
SVC, FRQ, CND = f"$D${Q_SV}", f"$D${Q_FR}", f"$D${Q_CD}"
ROOMSUM = "+".join(f"D{q}*{rr}" for q, rr in zip(q_rooms, ROOMR))


def both(label, f_i, f_j, fmt, **kw):
    rr = R1.calc(label, f_i, fmt, **kw)
    c = qb[f"J{rr}"]; c.value = f_j; c.number_format = fmt
    c.font = D.f(10, kw.get("bold", False) or kw.get("total", False)); c.alignment = Alignment(horizontal="right", vertical="center")
    return rr


R1.eyebrow("PERSON-HOURS")
x_sq = R1.calc("Hours by square feet", f"=D{Q_SQ}/1000*{SQR}/60", NUM2)
x_rm = R1.calc("Hours by room count", f"=({ROOMSUM})/60", NUM2)
x_bs = R1.calc("Base hours, plus pets", f"=MAX(I{x_sq},I{x_rm})+D{Q_PET}*{PETR}/60", NUM2)
x_cm = R1.calc("Condition multiplier", f"=IFERROR(INDEX({CD_MULT},MATCH({CND},{CD_NAMES},0)),1)", X2)
x_sm = both("Service multiplier", f"=IFERROR(INDEX({SV_MULT},MATCH({SVC},{SV_NAMES},0)),1)", f"={SV_DEEP_MULT}", X2)
x_ah = both("Add-on hours", f"=SUMPRODUCT(D{QA0}:D{QA1},{AD_MIN})/60", "", NUM2)
qb[f"J{x_ah}"] = f"=I{x_ah}"
x_tp = both("Total person-hours", f"=$I${x_bs}*$I${x_cm}*I{x_sm}+I{x_ah}", f"=$I${x_bs}*$I${x_cm}*J{x_sm}+J{x_ah}", NUM2, total=True)
x_ck = both("On-site clock time, hours", f"=IFERROR(I{x_tp}/$D${Q_CL},0)", f"=IFERROR(J{x_tp}/$D${Q_CL},0)", NUM2)
x_dr = both("Drive person-hours", f"=IFERROR($D${Q_MI}/{SPDR},0)*$D${Q_CL}", "", NUM2)
qb[f"J{x_dr}"] = f"=I{x_dr}"
R1.eyebrow("COST AT BREAK-EVEN, ZERO PROFIT")
x_lr = both("Labor rate, blended", f"=IFERROR(((MAX(0,$D${Q_CL}-$D${Q_HE}))*{BE}+$D${Q_HE}*{HCR})/$D${Q_CL},0)", "", USD2)
qb[f"J{x_lr}"] = f"=I{x_lr}"
x_lc = both("Labor cost", f"=(I{x_tp}+I{x_dr})*$I${x_lr}", f"=(J{x_tp}+J{x_dr})*$I${x_lr}", USD2)
x_vc = both("Vehicle", f"=$D${Q_MI}*{CPMR}", "", USD2)
qb[f"J{x_vc}"] = f"=I{x_vc}"
x_su = both("Supplies", f"=IFERROR(INDEX({SV_SUP},MATCH({SVC},{SV_NAMES},0)),0)+SUMPRODUCT(D{QA0}:D{QA1},{AD_SUP})",
            f"={SV_DEEP_SUP}+SUMPRODUCT(D{QA0}:D{QA1},{AD_SUP})", USD2)
x_fl = both("Price floor", f"=SUM(I{x_lc}:I{x_su})", f"=SUM(J{x_lc}:J{x_su})", USD2, total=True)
R1.eyebrow("YOUR PRICES")
x_one = R1.calc("One-time price, up to $5", f"=MAX({MINR},CEILING(I{x_fl}/(1-{MGR}),5))", USD0)
x_dis = R1.calc("Discount, this frequency", f"=IFERROR(INDEX({FR_DISC},MATCH({FRQ},{FR_NAMES},0)),0)", PCT1)
x_vis = R1.calc("Price per visit", f"=MAX({MINR},CEILING(I{x_one}*(1-I{x_dis}),5))", USD0, total=True, tint="total_tint")
x_dp = R1.calc("First visit as deep clean", f'=IF({FRQ}="One time","Not needed",MAX({MINR},CEILING(J{x_fl}/(1-{MGR}),5)))', USD0)
x_vm = R1.calc("Visits per month", f"=IFERROR(INDEX({FR_VISITS},MATCH({FRQ},{FR_NAMES},0)),1)", NUM2)
x_rv = R1.calc("Revenue per month", f'=IF({FRQ}="One time",I{x_vis},I{x_vis}*I{x_vm})', USD0)
x_pv = R1.calc("Profit per visit", f"=I{x_vis}-I{x_fl}", USD2, bold=True)
x_eh = R1.calc("Earned per person-hour", f"=IFERROR((I{x_vis}-I{x_vc}-I{x_su})/(I{x_tp}+I{x_dr}),0)", USD2)
R1.close()
for rr, color in ((x_lc, "teal"), (x_vc, "gold"), (x_su, "gold")):
    pass
D.status_rule(qb, f"I{x_pv}", f"$I${x_pv}<0", fill_tint=False)

# the check pill (v2 wording; asks for the home first when nothing is entered)
qb[f"C{pill}"] = (f'=IF(AND(N($D${Q_SQ})=0,SUM($D${q_rooms[0]}:$D${q_rooms[-1]})=0),"Enter the home above to see the check",'
                  f'IF($I${x_vis}<$I${x_fl},"Below your floor: this discount loses money on every visit. Lower the discount in tab 2.",'
                  f'IF($I${x_vis}<$I${x_fl}/(1-{MGR})*0.999,"Above your floor but under your target margin at this frequency",'
                  f'"At or above your target")))')
for formula, color, tint in ((f'LEFT($C${pill},5)="Below"', "accent", "accent_tint"),
                             (f'LEFT($C${pill},5)="Above"', "gold", "input_fill"),
                             (f'LEFT($C${pill},5)="At or"', "teal", "teal_tint")):
    qb.conditional_formatting.add(f"C{pill}:D{pill2}", FormulaRule(formula=[formula], font=Font(color=D.T[color], bold=True),
                                                                  fill=PatternFill("solid", bgColor=D.T[tint]), stopIfTrue=True))

# the quote block, full width
qt_top = max(L3.end, R1.end) + 2
r = qt_top
D.h(qb, r, 12); r += 1
t = qb[f"C{r}"]; t.value = "Quote to send"; t.font = D.f(12, True); qb.merge_cells(f"C{r}:J{r}"); D.h(qb, r, 24); r += 1
t = qb[f"C{r}"]; t.value = "Copy these lines, print this card, or save the tab as a PDF."; t.font = D.f(9, False, "muted")
t.alignment = Alignment(vertical="top"); qb.merge_cells(f"C{r}:J{r}"); D.h(qb, r, 18); r += 1
Q_CLIENT = wide_input(qb, r, "Client name", HOME["client"], "@", "Who the quote is for. Up to 60 characters.",
                      validation=("textLength", 0, 60), error="Up to 60 characters."); D.h(qb, r + 1, 6); r += 2
lab = qb[f"C{r}"]; lab.value = "Quote date"; lab.font = D.f(10); lab.alignment = Alignment(vertical="center")
c = S.input_cell(qb, f"D{r}", HOME["date"], DATE, validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
                 prompt="The day you send the quote, for example 10/5/2026.", prompt_title="Quote date",
                 error="Enter a date, for example 10/5/2026.")
c.fill = D.fill("input_fill"); c.font = D.f(10, True, "input_text"); c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
c.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
D.h(qb, r, 24); Q_DATE = r; r += 1
D.h(qb, r, 10); r += 1
for k in D.cols("C", "J"):
    qb[f"{k}{r - 1}"].border = Border(bottom=D.side("hair"))
B_ = BIZ_REF
QR = {}
lines = [
    ("name", f"={B_['Business name']}", 13, True),
    ("contact", f'={B_["Phone"]}&"  |  "&{B_["Email"]}&"  |  "&{B_["Website or social page"]}', 10, False),
    ("for", f'="Cleaning quote for "&$D${Q_CLIENT}&IF($D${Q_DATE}="","",", dated "&TEXT($D${Q_DATE},"mmm d, yyyy"))', 10, True),
    ("home", f'=TEXT($D${Q_SQ},"#,##0")&" sq ft, "&$D${q_rooms[0]}&IF($D${q_rooms[0]}=1," bedroom, "," bedrooms, ")&$D${q_rooms[1]}'
             f'&" full and "&$D${q_rooms[2]}&" half bathrooms, "&IF($D${Q_PET}=0,"no pets",$D${Q_PET}&IF($D${Q_PET}=1," pet"," pets"))'
             f'&". Service: "&{SVC}&", "&LOWER({FRQ})&"."', 10, False),
    ("price", f'=IF({FRQ}="One time","Price: $"&TEXT($I${x_vis},"#,##0"),"Price per visit: $"&TEXT($I${x_vis},"#,##0")'
              f'&IF(ISNUMBER($I${x_dp}),". First visit (deep clean): $"&TEXT($I${x_dp},"#,##0"),""))&"."', 11, True),
    ("addons", '=IF(SUM($D${a}:$D${b})=0,"No add-ons.","Add-ons included: "&MID('.format(a=QA0, b=QA1)
               + "&".join(f'IF($D${q}>0,"; "&{CRS}C{ad_first + i}&" x"&$D${q},"")' for i, q in enumerate(q_add)) + ',3,999)&".")', 10, False),
    ("supplies", f'={B_["Supplies line on the quote"]}&". Payment: "&{B_["Payment methods you accept"]}&"."', 10, False),
    ("cancel", f'="Please give "&{YN}$D${NOTICE}&" hours notice to cancel or reschedule; late cancellations and lockouts are $"&{YN}$D${FEE}&"."', 10, False),
    ("valid", f'="This quote is valid for "&{YN}$D${VALID}&" days and is based on the home as described. Prices may change if the home differs at the first visit."', 10, False),
]
for key, formula, size, bold in lines:
    c = qb[f"C{r}"]; c.value = formula; c.font = D.f(size, bold); c.alignment = Alignment(vertical="center", wrap_text=True)
    qb.merge_cells(f"C{r}:J{r}"); D.h(qb, r, 22 if size <= 11 else 26); QR[key] = r; r += 1
D.h(qb, r, 8); r += 1
c = qb[f"C{r}"]
c.value = ("Sample quote wording to adapt. Have a local attorney review your service terms (cancellation, fees, price changes) "
           "before you use them with clients.")
c.font = D.f(9, False, "muted"); c.alignment = Alignment(vertical="center", wrap_text=True); qb.merge_cells(f"C{r}:J{r}"); D.h(qb, r, 28)
QR["sample"] = r; r += 1
D.h(qb, r, 12)
D.frame(qb, qt_top, r, "B", "K")
QB_LAST = r + 1
D.h(qb, QB_LAST, 14)
D.tile(qb, 6, "B", "E", "PRICE PER VISIT", f"=$I${x_vis}", USD0,
       f'=IF({FRQ}="One time","One-time clean, no discount",IF(ISNUMBER($I${x_dp}),"First visit as a deep clean: "&TEXT($I${x_dp},"$#,##0"),""))', dark=True)
v, _, _ = D.tile(qb, 6, "G", "K", "PROFIT PER VISIT BEYOND YOUR FLOOR", f"=$I${x_pv}", USD2,
                 f'="Earned "&TEXT($I${x_eh},"$#,##0.00")&" per person-hour after vehicle and supplies"')
D.status_rule(qb, v.coordinate, f"$I${x_pv}<0", fill_tint=False)
D.h(qb, 10, 14)
S.page_break_before(qb, CARD_TOP - 1)
S.page_break_before(qb, qt_top)
fill_heights(qb, QB_LAST)
D.paint_canvas(qb, QB_LAST, "Z")
S.finish_sheet(qb, PRODUCT, QB_LAST, span=("A", "L"), freeze="A11", tab_color="accent")
qb.sheet_properties.pageSetUpPr.fitToPage = False
qb.page_setup.scale = 72
CELLS["qb"] = dict(SV=Q_SV, FR=Q_FR, CD=Q_CD, SQ=Q_SQ, rooms=q_rooms, PET=Q_PET, CL=Q_CL, HE=Q_HE, MI=Q_MI, add=q_add,
                   x_sq=x_sq, x_rm=x_rm, x_bs=x_bs, x_cm=x_cm, x_sm=x_sm, x_ah=x_ah, x_tp=x_tp, x_ck=x_ck, x_dr=x_dr, x_lr=x_lr,
                   x_lc=x_lc, x_vc=x_vc, x_su=x_su, x_fl=x_fl, x_one=x_one, x_dis=x_dis, x_vis=x_vis, x_dp=x_dp, x_vm=x_vm,
                   x_rv=x_rv, x_pv=x_pv, x_eh=x_eh, pill=pill, CLIENT=Q_CLIENT, DATE=Q_DATE, lines=QR)

# =====================================================================  5 Client List
cl = wb.create_sheet("5 Client List")
S.set_widths(cl, {"A": 3, "B": 2, "C": 18, "D": 14, "E": 15, "F": 14, "G": 9, "H": 10, "I": 9, "J": 11, "K": 27, "L": 10,
                  "M": 12, "N": 24, "O": 2, "P": 3})
D.page_header(cl, "Client List", "One row per client. Two example rows show the format; replace them.", "Keep door and alarm codes elsewhere",
              span=("B", "O"), side_cols=4)
CL_FIRST = 11 + 5
CL_LAST = CL_FIRST + N_CLIENTS - 1
D.h(cl, 6, 10)
D.stat_strip(cl, 7, [
    ("C:E", "Recurring revenue per month", f"=SUM(J{CL_FIRST}:J{CL_LAST})", USD0),
    ("F:H", "Recurring clients", f'=COUNTIF(J{CL_FIRST}:J{CL_LAST},">0")', INT),
    ("K:N", "Privacy", "Do not store alarm or door codes here", "@"),
])
cl["K8"].font = D.f(10, True, "ink2"); cl["K8"].alignment = Alignment(vertical="center", wrap_text=True)
D.h(cl, 9, 10)
D.frame(cl, 6, 9, "B", "O")
D.h(cl, 10, 14)


def cl_visits(rr):
    return f'=IF(F{rr}="","",IFERROR(INDEX({FR_VISITS},MATCH(F{rr},{FR_NAMES},0)),""))'


cl_cols = [
    dict(col="C", head="Client", kind="input", values=[c[0] for c in CLIENTS], validation=("textLength", 0, 60),
         prompt="A name or a label you will recognise. Up to 60 characters.", error="Up to 60 characters."),
    dict(col="D", head="Area or address", kind="input", values=[c[1] for c in CLIENTS], validation=("textLength", 0, 80),
         prompt="Area or address. Up to 80 characters.", error="Up to 80 characters."),
    dict(col="E", head="Phone", kind="input", values=[c[2] for c in CLIENTS], validation=("textLength", 0, 30),
         prompt="Phone number. Up to 30 characters.", error="Up to 30 characters."),
    dict(col="F", head="How often", kind="input", values=[c[3] for c in CLIENTS], validation=("list_range", FR_NAMES),
         prompt="Pick from the list. It comes from tab 2.", error="Pick how often from the list."),
    dict(col="G", head="Usual day", kind="input", values=[c[4] for c in CLIENTS], validation=("textLength", 0, 12),
         prompt="For example Tue. Up to 12 characters.", error="Up to 12 characters."),
    dict(col="H", head="Price per visit", kind="input", fmt=USD0, align="right", values=[c[5] for c in CLIENTS],
         validation=("decimal", 0, 100000), prompt="What this client pays per visit, in dollars. From $0 to $100,000.", error="Enter dollars from 0 to 100,000."),
    dict(col="I", head="Visits a month", kind="calc", fmt=NUM2, align="right", formula=cl_visits),
    dict(col="J", head="Revenue a month", kind="calc", fmt=USD0, align="right",
         formula=lambda rr: f'=IF(OR(F{rr}="",F{rr}="One time",H{rr}=""),"",H{rr}*I{rr})'),
    dict(col="K", head="How you get in", kind="input", values=[c[6] for c in CLIENTS], validation=("textLength", 0, 60),
         prompt="How you get in, without any codes. Keep alarm and door codes somewhere locked. Up to 60 characters.", error="Up to 60 characters."),
    dict(col="L", head="Pets", kind="input", values=[c[7] for c in CLIENTS], validation=("textLength", 0, 30),
         prompt="Pets in the home. Up to 30 characters.", error="Up to 30 characters."),
    dict(col="M", head="Start date", kind="input", fmt=DATE, align="right", values=[c[8] for c in CLIENTS],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"), prompt="First clean, for example 10/7/2026.", error="Enter a date, for example 10/7/2026."),
    dict(col="N", head="Notes", kind="input", values=[c[9] for c in CLIENTS], validation=("textLength", 0, 120),
         prompt="Anything worth remembering. Up to 120 characters.", error="Up to 120 characters."),
]
f1, f2, cl_end = D.table_card(cl, 11, "B", "O", cl_cols, N_CLIENTS, title="Your clients",
                              sub="Yellow columns are yours. Visits and revenue a month come from the How often table on tab 2.")
assert (f1, f2) == (CL_FIRST, CL_LAST)
D.paint_canvas(cl, cl_end + 1, "Z")
S.finish_sheet(cl, PRODUCT, cl_end + 1, span=("A", "P"), freeze=f"A{CL_FIRST}", tab_color="purple")
cl.page_setup.orientation = "landscape"
cl.print_title_rows = f"{CL_FIRST - 1}:{CL_FIRST - 1}"
CELLS["cl"] = dict(first=CL_FIRST, rev="C8", count="F8")

# =====================================================================  6 Job Log (Money by Month at the top)
jl = wb.create_sheet("6 Job Log")
S.set_widths(jl, {"A": 3, "B": 2, "C": 12, "D": 16, "E": 16, "F": 11, "G": 11, "H": 11, "I": 11, "J": 11, "K": 11, "L": 12,
                  "M": 12, "N": 12, "O": 13, "P": 11, "Q": 2, "R": 3})
D.page_header(jl, "Job Log", "Log each finished visit, see what it really earned and what to set aside each month.", "Set-asides are estimates",
              span=("B", "Q"), side_cols=4)
# rows: strip 6-9, money table from 11 (title, sub, gap, head, 12 months + year row), note, then the log
MM_TOP = 11
mm_first = MM_TOP + 6   # pad, title, sub, year picker, gap, head
mm_last = mm_first + 11
mm_tot = mm_last + 1
JL_TOP = mm_tot + 4
jl_first = JL_TOP + 5
jl_last = jl_first + N_JOBS - 1
R = {k: f"{k}{jl_first}:{k}{jl_last}" for k in "ABCDEFGHIJKLMNOP"}
D.h(jl, 6, 10)
D.stat_strip(jl, 7, [
    ("C:D", "Visits logged", f'=COUNTIF({R["F"]},">0")', INT),
    ("E:F", "Revenue logged", f"=SUM({R['F']})", USD0),
    ("G:J", "Profit beyond break-even", f"=SUM({R['M']})", USD0),
    ("K:M", "Jobs below floor", f'=COUNTIF({R["O"]},"Below floor")', INT),
    ("N:P", "Set aside, the year", f"=L{mm_tot}+M{mm_tot}", USD0),
])
D.h(jl, 9, 10)
D.frame(jl, 6, 9, "B", "Q")
D.status_rule(jl, "G8", "G8<0", fill_tint=False)
D.status_rule(jl, "K8", "K8>0", fill_tint=False)
D.h(jl, 10, 14)

# money by month card (built by hand: a year picker row, 12 months and a year total)
r = MM_TOP
D.h(jl, r, 12); r += 1
t = jl[f"C{r}"]; t.value = "Money by month"; t.font = D.f(12, True); jl.merge_cells(f"C{r}:P{r}"); D.h(jl, r, 24); r += 1
t = jl[f"C{r}"]; t.value = ("Revenue, miles, supplies and helper hours come from the log below. Type your other expenses each month. "
                            "Set-asides use your tax rates from tab 1 and are estimates.")
t.font = D.f(9, False, "muted"); t.alignment = Alignment(vertical="top", wrap_text=True); jl.merge_cells(f"C{r}:P{r}"); D.h(jl, r, 18); r += 1
lab = jl[f"C{r}"]; lab.value = "Year"; lab.font = D.f(10); lab.alignment = Alignment(vertical="center")
yc = S.input_cell(jl, f"D{r}", YEAR, "0", validation=("whole", 2000, 2100), prompt="The year to show, for example 2026. Whole number, 2000 to 2100.",
                  prompt_title="Year", error="Enter a year from 2000 to 2100.")
yc.fill = D.fill("input_fill"); yc.font = D.f(10, True, "input_text"); yc.alignment = Alignment(horizontal="right", vertical="center", indent=1)
yc.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
D.h(jl, r, 24); MM_YEAR = r; r += 1
D.h(jl, r, 8); r += 1
head = r
MM_HEADS = [("C", "Month"), ("D", "Jobs"), ("E", "Revenue"), ("F", "Miles"), ("G", "Vehicle cost"), ("H", "Job supplies"),
            ("I", "Helper pay and payroll"), ("J", "Other expenses"), ("K", "Profit estimate"), ("L", "SE tax set-aside"),
            ("M", "Income tax set-aside"), ("N", "Total to set aside")]
for k, txt in MM_HEADS:
    c = jl[f"{k}{head}"]; c.value = txt.upper(); c.font = D.f(8, True, "muted")
    c.alignment = Alignment(horizontal="left" if k == "C" else "right", vertical="bottom", wrap_text=True,
                            indent=1 if k == "J" else 0)
for k in D.cols("C", "P"):
    jl[f"{k}{head}"].border = Border(bottom=D.side("ink"))
D.h(jl, head, 30)
assert head + 1 == mm_first, (head, mm_first)
JN, JD, JH, JI, JF = (f"$" + k + f"${jl_first}:$" + k + f"${jl_last}" for k in "PFJKH")
for i in range(12):
    rr = mm_first + i
    vals = {
        "C": (f"=DATE($D${MM_YEAR},{i + 1},1)", MONTH),
        "D": (f'=COUNTIFS({JN},C{rr},{JD},">0")', INT),
        "E": (f"=SUMIFS({JD},{JN},C{rr})", USD0),
        "F": (f"=SUMIFS({JH},{JN},C{rr})", INT),
        "G": (f"=F{rr}*{CPMR}", USD0),
        "H": (f"=SUMIFS({JI},{JN},C{rr})", USD0),
        "I": (f"=SUMIFS({JF},{JN},C{rr})*{HCR}", USD0),
        "K": (f"=E{rr}-G{rr}-H{rr}-I{rr}-J{rr}", USD0),
        "L": (f"=MAX(0,K{rr}*{SBR}*{SER})", USD0),
        "M": (f"=MAX(0,K{rr}*{ITR})", USD0),
        "N": (f"=L{rr}+M{rr}", USD0),
    }
    for k, (fm, fmt) in vals.items():
        c = jl[f"{k}{rr}"]; c.value = fm; c.number_format = fmt; c.font = D.f(10, k == "N")
        c.alignment = Alignment(horizontal="left" if k == "C" else "right", vertical="center")
    c = S.input_cell(jl, f"J{rr}", 0 if i else 0, USD0, validation=("decimal", 0, 1000000),
                     prompt="Other business expenses this month (insurance, products, phone, ads), in dollars. From $0 to $1,000,000.",
                     prompt_title="Other expenses", error="Enter dollars from 0 to 1,000,000.")
    c.fill = D.fill("input_fill"); c.font = D.f(10, True, "input_text"); c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    c.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
    for k in D.cols("C", "P"):
        if k != "J":
            jl[f"{k}{rr}"].border = Border(bottom=D.side("hair"))
    D.h(jl, rr, D.GRID_ROW)
for k in D.cols("C", "P"):
    c = jl[f"{k}{mm_tot}"]
    if k == "C":
        c.value = "Year"
    elif "D" <= k <= "N":
        c.value = f"=SUM({k}{mm_first}:{k}{mm_last})"
        c.number_format = INT if k in "DF" else USD0
    c.font = D.f(10, True); c.alignment = Alignment(horizontal="left" if k == "C" else "right", vertical="center")
    c.border = Border(top=D.side("ink")); c.fill = D.fill("total_tint")
D.h(jl, mm_tot, 24)
r = mm_tot + 1
c = jl[f"C{r}"]
c.value = ("Vehicle cost here uses your rate per mile from tab 1. If you deduct mileage, the IRS rate changed on July 1, 2026; your tax "
           "preparer will use the rate for each date. Information only, not tax advice.")
c.font = D.f(9, False, "muted"); c.alignment = Alignment(vertical="center", wrap_text=True); jl.merge_cells(f"C{r}:P{r}"); D.h(jl, r, 30)
MM_NOTE = r; r += 1
D.h(jl, r, 12)
D.frame(jl, MM_TOP, r, "B", "Q")
assert r + 2 == JL_TOP, (r, JL_TOP)


def jl_calc(kind):
    def fn(rr):
        if kind == "cost":
            return (f'=IF(F{rr}="","",((G{rr}-H{rr})+IFERROR(I{rr}*(G{rr}-H{rr})/G{rr},I{rr}))*{BE}'
                    f'+(H{rr}+IFERROR(I{rr}*H{rr}/G{rr},0))*{HCR}+J{rr}*{CPMR}+K{rr})')
        if kind == "profit":
            return f'=IF(F{rr}="","",F{rr}-L{rr})'
        if kind == "hour":
            return f'=IF(F{rr}="","",IFERROR((F{rr}-J{rr}*{CPMR}-K{rr})/(G{rr}+I{rr}),0))'
        if kind == "check":
            return f'=IF(F{rr}="","",IF(M{rr}<0,"Below floor",IF(F{rr}<L{rr}/(1-{MGR})*0.999,"Under target","On target")))'
        if kind == "month":
            return f'=IF(C{rr}="","",DATE(YEAR(C{rr}),MONTH(C{rr}),1))'
    return fn


jl_cols = [
    dict(col="C", head="Date", kind="input", fmt=DATE, values=[j[0] for j in JOBS], validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The day of the visit, for example 10/7/2026.", error="Enter a date, for example 10/7/2026."),
    dict(col="D", head="Client", kind="input", values=[j[1] for j in JOBS], validation=("textLength", 0, 60),
         prompt="A name or a label you will recognise. Up to 60 characters.", error="Up to 60 characters."),
    dict(col="E", head="Service type", kind="input", values=[j[2] for j in JOBS], validation=("list_range", SV_NAMES),
         prompt="Pick from the list. It comes from tab 2.", error="Pick a service type from the list."),
    dict(col="F", head="Price charged", kind="input", fmt=USD0, align="right", values=[j[3] for j in JOBS],
         validation=("decimal", 0, 100000), prompt="What the client paid for this visit, before any tip. From $0 to $100,000.", error="Enter dollars from 0 to 100,000."),
    dict(col="G", head="Person-hours on site", kind="input", fmt=NUM1, align="right", values=[j[4] for j in JOBS],
         validation=("decimal", 0, 1000), prompt="Hours on site for everyone added up. Two people for 3.5 hours is 7. From 0 to 1,000.", error="Enter hours from 0 to 1,000."),
    dict(col="H", head="Of those, helper hours", kind="input", fmt=NUM1, align="right", values=[j[5] for j in JOBS],
         validation=("decimal", 0, 1000), prompt="The helper part of the on-site hours. 0 if you cleaned alone. From 0 to 1,000.", error="Enter hours from 0 to 1,000."),
    dict(col="I", head="Team drive hours", kind="input", fmt=NUM1, align="right", values=[j[6] for j in JOBS],
         validation=("decimal", 0, 1000), prompt="Driving time for everyone added up. From 0 to 1,000.", error="Enter hours from 0 to 1,000."),
    dict(col="J", head="Miles", kind="input", fmt=INT, align="right", values=[j[7] for j in JOBS],
         validation=("decimal", 0, 10000), prompt="Round-trip miles driven. From 0 to 10,000.", error="Enter miles from 0 to 10,000."),
    dict(col="K", head="Supplies", kind="input", fmt=USD0, align="right", values=[j[8] for j in JOBS],
         validation=("decimal", 0, 10000), prompt="Supplies used on this visit, in dollars. From $0 to $10,000.", error="Enter dollars from 0 to 10,000."),
    dict(col="L", head="Cost at break-even", kind="calc", fmt=USD0, align="right", formula=jl_calc("cost")),
    dict(col="M", head="Profit beyond break-even", kind="calc", fmt=USD0, align="right", formula=jl_calc("profit")),
    dict(col="N", head="Earned per person-hour", kind="calc", fmt=USD2, align="right", formula=jl_calc("hour")),
    dict(col="O", head="Check", kind="calc", formula=jl_calc("check")),
    dict(col="P", head="Month", kind="calc", fmt=MONTH, align="right", formula=jl_calc("month")),
]
g1, g2, jl_end = D.table_card(jl, JL_TOP, "B", "Q", jl_cols, N_JOBS, title="Finished visits",
                              sub="Yellow columns are yours. Cost uses your break-even rate for your hours and the helper cost for helper hours.")
assert (g1, g2) == (jl_first, jl_last), (g1, g2, jl_first, jl_last)
for rr in range(jl_first, jl_last + 1):
    jl[f"O{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    jl[f"P{rr}"].font = D.f(10, False, "muted")
for word, color, tint in (("Below floor", "accent", "accent_tint"), ("Under target", "gold", "input_fill"), ("On target", "teal", "teal_tint")):
    jl.conditional_formatting.add(f"O{jl_first}:O{jl_last}", FormulaRule(formula=[f'$O{jl_first}="{word}"'],
                                  font=Font(color=D.T[color], bold=True), fill=PatternFill("solid", bgColor=D.T[tint]), stopIfTrue=True))
S.page_break_before(jl, JL_TOP)
fill_heights(jl, jl_end)
D.paint_canvas(jl, jl_end + 1, "Z")
S.finish_sheet(jl, PRODUCT, jl_end + 1, span=("A", "R"), freeze="A10", tab_color="gold")
jl.page_setup.orientation = "landscape"
jl.print_title_rows = f"{jl_first - 1}:{jl_first - 1}"
CELLS["jl"] = dict(first=jl_first, mm_first=mm_first, mm_tot=mm_tot, year=MM_YEAR, strip_below="K8", strip_rev="E8")

# =====================================================================  Start Here and Terms (card style, from the references)
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            m = re.match(r"^([A-Z][A-Za-z0-9 ,]{1,40})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if m:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), m.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), m.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    D.h(sh, r + 1, D.GRID_ROW)
    return r + 2


sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, PRODUCT, "Start here. Six working tabs, about fifteen minutes to set up.", f"Checked {CHECKED}", span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "Most cleaners set a price by dividing the pay they want by the hours they work. That leaves out the hours nobody pays for "
    "(walkthroughs, texts, laundry, supply runs), the costs that run every month, and self-employment tax. This workbook adds all "
    "three back, gives you the rate every quote starts from, prices each home from its square feet and rooms, shows your real "
    "recurring discount floor, and tracks clients, visits and what to set aside each month. The PDF forms in the download run "
    "the visit itself.", "para")])
r = sh_card(sh, r, "Set up once, then quote every home", [
    (1, "Your Numbers. Type the pay you want, the hours you really work, your monthly costs and your business details. Your target rate per cleaning hour is the dark tile at the top.", "num"),
    (2, "Cleaning Rates. Your minutes per 1,000 square feet and per room, service types, discounts for weekly, every 2 weeks and monthly, and add-ons.", "num"),
    (3, "Price Sheet. A price menu by home size for your website or posts. It updates when your numbers change.", "num"),
    (4, "Quote Builder. After the walkthrough (the Walkthrough page in the PDF forms has the same boxes), enter the home and read the price per visit, the first deep clean and the check. Send the quote card at the bottom.", "num"),
    (5, "Client List and Job Log. Who you clean for, what each visit really earned, and what to set aside for taxes each month.", "num"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "How to read the cells", [
    ((EX["take_home"], USD0), "Yellow cells with blue numbers are yours to change. Type over the example. Each one shows a hint when you select it. Type percentages with the % sign, for example 10%.", "chip-in"),
    ((round(_bill), INT), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((round(TARGET_X, 2), USD2), "Dark tiles are your answers. They sit at the top of Your Numbers and Quote Builder and stay on screen as you scroll.", "chip-key"),
    (None, "Each tab is protected without a password. To change anything else, use Review, Unprotect Sheet in Excel, or Data, "
           "Protect sheets and ranges in Google Sheets.", "note"),
])
r = sh_card(sh, r, "The example", [(None,
    f"A {money(EX['take_home'], False)} take-home goal, with {EX['se'] * 100:.1f}% self-employment tax on {EX['se_base'] * 100:.2f}% of profit "
    f"and a {EX['income'] * 100:.0f}% income tax estimate, needs {money(_profit, False)} of profit before tax. Add {money(_costs_y, False)} "
    f"a year of business costs and you need {money(_profit + _costs_y, False)} of revenue. At {EX['weeks']} weeks of {EX['hours']} hours "
    f"with {EX['share'] * 100:.0f}% spent cleaning or driving to jobs, that is {_bill:,.0f} billable hours: a break-even rate of "
    f"{money(BE_X)} and a target of {money(TARGET_X)} per cleaning hour with a {EX['margin'] * 100:.0f}% margin. Pay divided by all "
    f"hours would say {money(GUESS_X)}. The example home, {HOME['sqft']:,} square feet cleaned every 2 weeks by you and one helper, "
    f"quotes {money(VISIT_X, False)} a visit, with a first deep clean at {money(DEEP_X, False)}. These are example numbers; yours will differ.", "para")])
r = sh_card(sh, r, "Good to know", [
    (None, "Two time estimates. The Quote Builder times each home by square feet and by room count and uses the larger, so a small "
           "home with many bathrooms or a large open home with few rooms is not underquoted.", "para"),
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. Every formula is a plain "
           "formula that works in Excel and Google Sheets. No macros, no add-ons, no sign-up and no outside links.", "para"),
    (None, "Example numbers. Minutes, multipliers, discounts and add-on times are examples, not market prices. Time a few of your "
           "own cleans and replace them.", "para"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "Before you rely on a number", [(None,
    "For general information and planning only. This is not legal, tax, financial, medical or other professional advice, and using it "
    "doesn't create a professional relationship. Results are estimates based on the numbers you enter. Figures were checked "
    f"on {CHECKED} and can change. See the Terms tab before you rely on anything here. "
    + S.TAX_TIER_CLAUSE.format(date=CHECKED_LONG, sources=TAX_SOURCES) + " " + FORMS_CLAUSE, "note")])
D.paint_canvas(sh, r - 1, "Z")
S.finish_sheet(sh, PRODUCT, r - 1, span=("A", "G"), tab_color="teal")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "Thank you", [
    (None, "Thank you for buying from ProofNotFluff. If anything doesn't open or a number looks wrong, message the shop on Etsy and it will be fixed.", "para"),
    (None, ("Service Pricing Calculator: price any service job", "https://www.etsy.com/listing/4587962930"), "link"),
    (None, ("Hourly Rate Quick Calculator: the one-tab version", "https://www.etsy.com/listing/4589695837"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this kit helped you price your cleans, a short review on Etsy helps other cleaning business owners find it. Thank you.", "para"),
])
D.paint_canvas(tm, r - 1, "Z")
S.finish_sheet(tm, PRODUCT, r - 1, span=("A", "G"), tab_color="note")

wb.active = 0
problems = S.audit(wb)
if problems:
    print("AUDIT:", *problems[:40], sep="\n  ")
wb.save(OUT)
json.dump(CELLS, open(os.path.join(os.path.dirname(os.path.abspath(OUT)), "cells.json"), "w"), indent=1)
print("saved", OUT)
print("example:", dict(be=round(BE_X, 2), target=round(TARGET_X, 2), guess=round(GUESS_X, 2), floor=round(FLOOR_X, 2),
                       one=ONE_X, visit=VISIT_X, deep=DEEP_X))
