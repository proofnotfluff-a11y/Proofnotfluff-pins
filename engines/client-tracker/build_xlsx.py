#!/usr/bin/env python3
"""Client Tracker for Service Businesses, version 1 (v3 dashboard design, Oct 9, 2026).

Four working tabs:
  1 Today      the dashboard: follow-ups due on the as-of date (most overdue first, with the service
               each client last asked about), open quotes by value, this month's jobs, revenue, unpaid
               balance and new leads, top clients by revenue and where clients come from.
  2 Clients    one row per client: name, type (Lead, Active, Past), source, contact, first and next
               contact dates, a status note. Jobs done, revenue, unpaid, last job, last asked about,
               days to next contact and a follow-up flag in words fill in on their own.
  3 Jobs       one row per job or quote: date, client (picked from the Clients list), service (picked
               from the Services tab), amount, status (Quoted, Booked, Done, Lost), paid, note.
  4 Services   what you offer and what you usually charge; feeds the Jobs dropdown and lookup.

Every total reads fixed row ranges (Clients rows 16 to 315, Jobs rows 16 to 1,015, Services rows 16
to 45), never a whole column. Formulas sit only in locked white columns, so clearing the yellow
example cells never removes one. The as-of date is typed, never read from the clock (no TODAY()),
so the file shows the same numbers on any day and in any app. No TEXT() format strings, no macros,
no tables, no dynamic arrays: the same results in Excel and Google Sheets.

Built to design/WORKBOOK_STANDARD.md section 0 with tools/pnf_dash.py and tools/wb_style.py.

  python3 build_xlsx.py Client-Tracker-EXAMPLE.xlsx           (the worked example)
  python3 build_xlsx.py Client-Tracker.xlsx --blank           (blank)
PNF_JOB_ROWS and PNF_CLIENT_ROWS shorten the tables for listing shots only; labels always say the shipped size.
"""
import datetime as dt
import json
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402

PRODUCT = "Client Tracker"
LONG_NAME = "Client Tracker for Service Businesses"
VERSION = "1"
CHECKED = "Oct 9, 2026"
CHECKED_LONG = "October 9, 2026"
AS_OF = f"Version {VERSION}, {CHECKED}"
HERE = os.path.dirname(os.path.abspath(__file__))
args = sys.argv[1:]
BLANK = "--blank" in args
args = [a for a in args if a != "--blank"]
OUT = args[0] if args else "Client-Tracker-EXAMPLE.xlsx"
USD0, USD2, INT, DATE, NUM1 = S.FMT["usd0"], S.FMT["usd2"], S.FMT["int"], S.FMT["date"], S.FMT["num1"]
DAYS = '#,##0;-#,##0;0'
MONTH = 'mmm yyyy'
EXAMPLE_AS_OF = dt.datetime(2026, 10, 9)

TYPES = ["Lead", "Active", "Past"]
SOURCES = ["Referral", "Google", "Facebook", "Nextdoor", "Yard sign", "Repeat", "Other"]
STATUSES = ["Quoted", "Booked", "Done", "Lost"]

# ---------------------------------------------------------------- example: Brightside Cleaning, a made-up two-person house cleaner
d = dt.datetime
SERVICES = [  # service, standard price, typical hours
    ("Standard clean", 140.00, 2.5),
    ("Deep clean", 260.00, 5.0),
    ("Move-out clean", 320.00, 6.0),
    ("Window add-on", 60.00, 1.0),
    ("Fridge and oven", 45.00, 1.0),
    ("Recurring biweekly", 125.00, 2.0),
]
# name, type, source, contact (blank in the example), first contact, next contact, status note
CLIENTS = [
    ("M. Alvarez", "Active", "Referral", None, d(2026, 7, 12), d(2026, 10, 9), "Biweekly, Thursdays"),
    ("J. Okafor", "Active", "Google", None, d(2026, 7, 15), d(2026, 10, 23), "Monthly standard clean"),
    ("R. Lindqvist", "Active", "Nextdoor", None, d(2026, 7, 20), d(2026, 10, 6), "Ask about a fall deep clean"),
    ("P. Nakamura", "Active", "Referral", None, d(2026, 7, 28), d(2026, 10, 16), "Monthly, first Saturday"),
    ("S. Dubois", "Lead", "Google", None, d(2026, 10, 2), d(2026, 10, 9), "Quoted a deep clean, waiting"),
    ("T. Brennan", "Active", "Referral", None, d(2026, 8, 3), d(2026, 11, 2), "Every six weeks"),
    ("K. Haddad", "Past", "Nextdoor", None, d(2026, 8, 6), None, "Moved away after the move-out"),
    ("L. Moreau", "Active", "Google", None, d(2026, 8, 11), d(2026, 10, 30), "Monthly standard clean"),
    ("D. Petrov", "Lead", "Nextdoor", None, d(2026, 10, 5), d(2026, 10, 7), "Quoted a move-out clean"),
    ("A. Castillo", "Active", "Referral", None, d(2026, 8, 18), d(2026, 10, 9), "Collect the September balance"),
    ("E. Whitfield", "Lead", "Google", None, d(2026, 9, 29), d(2026, 10, 14), "Quoted a standard clean"),
    ("N. Ferreira", "Active", "Referral", None, d(2026, 8, 25), d(2026, 11, 10), "Monthly, no deep clean for now"),
    ("G. Sato", "Past", "Google", None, d(2026, 9, 1), None, "Went with another company"),
    ("B. Oyelaran", "Active", "Referral", None, d(2026, 9, 8), d(2026, 10, 20), "New, ask about a monthly plan"),
]
# date, client, service, amount, status, paid, note
JOBS = [
    (d(2026, 7, 16), "M. Alvarez", "Deep clean", 260, "Done", "Yes", "First visit"),
    (d(2026, 7, 18), "J. Okafor", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 7, 24), "R. Lindqvist", "Deep clean", 260, "Done", "Yes", None),
    (d(2026, 7, 24), "R. Lindqvist", "Window add-on", 60, "Done", "Yes", "Inside only"),
    (d(2026, 7, 30), "M. Alvarez", "Recurring biweekly", 125, "Done", "Yes", None),
    (d(2026, 8, 1), "P. Nakamura", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 8, 1), "P. Nakamura", "Window add-on", 60, "Done", "Yes", None),
    (d(2026, 8, 7), "T. Brennan", "Deep clean", 260, "Done", "Yes", None),
    (d(2026, 8, 7), "T. Brennan", "Fridge and oven", 45, "Done", "Yes", None),
    (d(2026, 8, 9), "K. Haddad", "Move-out clean", 320, "Done", "Yes", "Keys left with the agent"),
    (d(2026, 8, 13), "M. Alvarez", "Recurring biweekly", 125, "Done", "Yes", None),
    (d(2026, 8, 14), "L. Moreau", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 8, 15), "J. Okafor", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 8, 20), "A. Castillo", "Deep clean", 260, "Done", "Yes", None),
    (d(2026, 8, 21), "R. Lindqvist", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 8, 27), "M. Alvarez", "Recurring biweekly", 125, "Done", "Yes", None),
    (d(2026, 8, 28), "N. Ferreira", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 9, 3), "A. Castillo", "Recurring biweekly", 125, "Done", "Yes", None),
    (d(2026, 9, 4), "G. Sato", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 9, 5), "P. Nakamura", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 9, 10), "M. Alvarez", "Recurring biweekly", 125, "Done", "Yes", None),
    (d(2026, 9, 10), "B. Oyelaran", "Deep clean", 260, "Done", "Yes", "Referred by the Alvarez family"),
    (d(2026, 9, 11), "T. Brennan", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 9, 12), "J. Okafor", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 9, 12), "L. Moreau", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 9, 15), "N. Ferreira", "Deep clean", 260, "Lost", None, "Price too high for now"),
    (d(2026, 9, 17), "A. Castillo", "Recurring biweekly", 125, "Done", "No", "Invoice sent"),
    (d(2026, 9, 18), "R. Lindqvist", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 9, 20), "G. Sato", "Deep clean", 260, "Lost", None, "Went with another company"),
    (d(2026, 9, 24), "M. Alvarez", "Recurring biweekly", 125, "Done", "Yes", None),
    (d(2026, 9, 26), "N. Ferreira", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 9, 30), "E. Whitfield", "Standard clean", 140, "Quoted", None, "Wants a weekday morning"),
    (d(2026, 10, 1), "A. Castillo", "Recurring biweekly", 125, "Done", "No", "Second unpaid visit"),
    (d(2026, 10, 2), "P. Nakamura", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 10, 2), "S. Dubois", "Deep clean", 260, "Quoted", None, "Three bedrooms, two baths"),
    (d(2026, 10, 3), "J. Okafor", "Standard clean", 140, "Done", "Yes", None),
    (d(2026, 10, 6), "D. Petrov", "Move-out clean", 320, "Quoted", None, "Lease ends Oct 31"),
    (d(2026, 10, 7), "B. Oyelaran", "Standard clean", 140, "Done", "No", "Pays by check next visit"),
    (d(2026, 10, 8), "M. Alvarez", "Recurring biweekly", 125, "Done", "No", "Card declined, retry Friday"),
    (d(2026, 10, 9), "L. Moreau", "Standard clean", 140, "Booked", None, "Booked for 2 pm"),
]
JOBS.sort(key=lambda x: (x[0], x[1]))


def expected(as_of=EXAMPLE_AS_OF):
    out = {}
    for name, typ, src, _, first, nxt, _ in CLIENTS:
        mine = [j for j in JOBS if j[1] == name]
        done = [j for j in mine if j[4] == "Done"]
        last_any = max(mine, key=lambda j: (j[0], JOBS.index(j))) if mine else None
        days = (nxt - as_of).days if nxt else None
        if nxt is None or typ == "Past":
            flag = ""
        elif days == 0:
            flag = "Due today"
        elif days < 0:
            flag = f"Overdue {-days} day" + ("" if days == -1 else "s")
        else:
            flag = "Not yet"
        out[name] = dict(type=typ, source=src, jobs=len(done), revenue=sum(j[3] for j in done),
                         unpaid=sum(j[3] for j in done if j[5] == "No"), last=max((j[0] for j in done), default=None),
                         asked=last_any[2] if last_any else "", days=days, flag=flag, due=flag.startswith(("Due", "Overdue")))
    ms = as_of.replace(day=1)
    me = (ms.replace(month=ms.month % 12 + 1, year=ms.year + ms.month // 12) - dt.timedelta(days=1))
    month_done = [j for j in JOBS if j[4] == "Done" and ms <= j[0] <= me]
    quotes = [j for j in JOBS if j[4] == "Quoted"]
    tot = dict(due=sum(1 for c in out.values() if c["due"]), overdue=sum(1 for c in out.values() if c["flag"].startswith("Overdue")),
               today=sum(1 for c in out.values() if c["flag"] == "Due today"),
               quotes_value=sum(j[3] for j in quotes), quotes=len(quotes),
               month_jobs=len(month_done), month_revenue=sum(j[3] for j in month_done),
               month_unpaid=sum(j[3] for j in month_done if j[5] == "No"),
               month_leads=sum(1 for c in CLIENTS if ms <= c[4] <= me),
               clients=len(CLIENTS), jobs_done=sum(c["jobs"] for c in out.values()),
               revenue=sum(c["revenue"] for c in out.values()), unpaid=sum(c["unpaid"] for c in out.values()),
               jobs=len(JOBS), lost=sum(1 for j in JOBS if j[4] == "Lost"),
               sources={s: sum(1 for c in CLIENTS if c[2] == s) for s in SOURCES if any(c[2] == s for c in CLIENTS)})
    top = sorted(out.items(), key=lambda kv: -kv[1]["revenue"])
    tot["top"] = [(k, v["revenue"]) for k, v in top[:8]]
    tot["due_list"] = sorted([(k, -v["days"], v["asked"]) for k, v in out.items() if v["due"]], key=lambda x: (-x[1], list(out).index(x[0])))
    return out, tot


EX_CLIENTS, EX = expected()
if BLANK:
    CLIENTS, JOBS, SERVICES = [], [], []

# ---------------------------------------------------------------- fixed addresses
TD_N, CL_N, JB_N, SV_N = "1 Today", "2 Clients", "3 Jobs", "4 Services"
TD, CL, JB, SV = (f"'{n}'" for n in (TD_N, CL_N, JB_N, SV_N))
SHIP_CLIENTS, SHIP_JOBS, N_SERVICES = 300, 1000, 30
N_CLIENTS = int(os.environ.get("PNF_CLIENT_ROWS", str(SHIP_CLIENTS)))   # short copies for listing shots only
N_JOBS = int(os.environ.get("PNF_JOB_ROWS", str(SHIP_JOBS)))
C_TOP = J_TOP = S_TOP = 11
C_FIRST = J_FIRST = S_FIRST = 16
C_LAST, J_LAST, S_LAST = C_FIRST + N_CLIENTS - 1, J_FIRST + N_JOBS - 1, S_FIRST + N_SERVICES - 1


def crng(col):
    return f"{CL}!${col}${C_FIRST}:${col}${C_LAST}"


def jrng(col):
    return f"{JB}!${col}${J_FIRST}:${col}${J_LAST}"


def srng(col):
    return f"{SV}!${col}${S_FIRST}:${col}${S_LAST}"


# Clients: C name, D type, E source, F contact, G first, H next, I note | J jobs, K revenue, L unpaid, M last job,
#          N last asked about, O days to next, P follow-up
CNAMES, CTYPES, CSRC, CFIRST, CNEXT = crng("C"), crng("D"), crng("E"), crng("G"), crng("H")
CJOBS, CREV, CUNPAID, CASKED, CFLAG = crng("J"), crng("K"), crng("L"), crng("N"), crng("P")
# Jobs: C date, D client, E service, F amount, G status, H paid, I note | J usual price, K outstanding, L check
JDATE, JCLIENT, JSERV, JAMT, JSTATUS, JPAID, JOUT, JCHECK = (jrng(c) for c in "CDEFGHKL")
# Services: C name, D price, E hours | F times done, G revenue
SNAMES, SPRICE = srng("C"), srng("D")
AS_OF_CELL = None  # set when the Today tab is built: f"{TD}!$J$<row>"


def rich(head, body, size=9):
    return CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=size, b=True, color=D.T["ink"]), head + ". "),
                        TextBlock(InlineFont(rFont=D.FONT, sz=size, color=D.T["ink2"]), body))


def left_cells(ws, first, last, cols_, size=10, color="ink", bold=False):
    for rr in range(first, last + 1):
        for k in cols_:
            c = ws[f"{k}{rr}"]
            c.font = D.f(size, bold, color)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)


def word_rules(ws, rng, first, col, words):
    """Color a text column by its first word(s): words = [(prefix, color, bold, tint)]."""
    for word, color, bold, tint in words:
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'LEFT(${col}{first},{len(word)})="{word}"'], font=Font(color=D.T[color], bold=bold),
            fill=PatternFill("solid", bgColor=D.T["accent_tint"]) if tint else None, stopIfTrue=True))


wb = S.new_workbook(PRODUCT, version=VERSION)
td = wb.create_sheet(TD_N)   # created first so it sits first; filled after the tables exist
cl = wb.create_sheet(CL_N)
jb = wb.create_sheet(JB_N)
sv = wb.create_sheet(SV_N)

# the as-of date lives on 1 Today at a fixed cell (right column, first input row)
TD_ASOF_ROW = 14
AS_OF_CELL = f"{TD}!$J${TD_ASOF_ROW}"
ASOF = AS_OF_CELL

# =====================================================================  2 Clients
S.set_widths(cl, {"A": 3, "B": 2, "C": 20, "D": 9, "E": 11, "F": 16, "G": 14, "H": 14, "I": 32, "J": 9, "K": 13,
                  "L": 11, "M": 14, "N": 20, "O": 10, "P": 17, "Q": 2, "R": 3})
D.page_header(cl, "Clients", "One row per client. Jobs, revenue, unpaid, last job and the follow-up flag fill in on their own.",
              "Example rows: type over them" if not BLANK else f"Type your first client in row {C_FIRST}", span=("B", "Q"), side_cols=4)
D.h(cl, 6, 10)
D.stat_strip(cl, 7, [
    ("C:D", "Clients listed", f'=COUNTIF({CNAMES},"?*")', INT),
    ("E:F", "Leads", f'=COUNTIF({CTYPES},"Lead")', INT),
    ("G:H", "Active clients", f'=COUNTIF({CTYPES},"Active")', INT),
    ("I:J", "Follow-ups due", "=0", INT),  # filled from the Today tab's due keys once they exist
    ("K:M", "Unpaid, done jobs", f"=SUM({CUNPAID})", USD2),
    ("N:P", "Rows to fix", f'=COUNTIF({CFLAG},"Needs*")+COUNTIF({CFLAG},"Duplicate*")', INT),
])
D.h(cl, 9, 10)
D.frame(cl, 6, 9, "B", "Q")
D.status_rule(cl, "I8", "$I$8>0", fill_tint=False)
D.status_rule(cl, "N8", "$N$8>0", fill_tint=False)
D.h(cl, 10, 14)


def f_cjobs(r):
    return f'=IF($C{r}="","",COUNTIFS({JCLIENT},$C{r},{JSTATUS},"Done"))'


def f_crev(r):
    return f'=IF($C{r}="","",SUMIFS({JAMT},{JCLIENT},$C{r},{JSTATUS},"Done"))'


def f_cunpaid(r):
    return f'=IF($C{r}="","",SUMIFS({JAMT},{JCLIENT},$C{r},{JSTATUS},"Done",{JPAID},"No"))'


def f_clast(r):
    return (f'=IF($C{r}="","",IF(SUMPRODUCT(MAX(({JCLIENT}=$C{r})*({JSTATUS}="Done")*{JDATE}))=0,"",'
            f'SUMPRODUCT(MAX(({JCLIENT}=$C{r})*({JSTATUS}="Done")*{JDATE}))))')


def f_casked(r):
    # the service on the client's most recent row in the log (latest date, then the lowest row on that date)
    last = f'SUMPRODUCT(MAX(({JCLIENT}=$C{r})*{JDATE}))'
    idx = f'SUMPRODUCT(MAX(({JCLIENT}=$C{r})*({JDATE}={last})*(ROW({JDATE})-{J_FIRST - 1})))'
    return f'=IF($C{r}="","",IFERROR(IF({idx}=0,"",INDEX({JSERV},{idx})&""),""))'


def f_cdays(r):
    return f'=IF(OR($C{r}="",$H{r}="",{ASOF}=""),"",$H{r}-{ASOF})'


def f_cflag(r):
    return (f'=IF($C{r}="",IF(COUNTA($D{r}:$I{r})>0,"Needs a name",""),IF(COUNTIF({CNAMES},$C{r})>1,"Duplicate name",'
            f'IF(OR($O{r}="",$D{r}="Past"),"",IF($O{r}=0,"Due today",IF($O{r}<0,"Overdue "&(-$O{r})&IF($O{r}=-1," day"," days"),"Not yet")))))')


def cvals(i):
    return [x[i] for x in CLIENTS]


date_v = ("date", "DATE(2000,1,1)", "DATE(2100,12,31)")
cols_cl = [
    dict(col="C", head="Client", kind="input", values=cvals(0), validation=("textLength", 1, 40),
         prompt="The client's name as you want to see it, for example M. Alvarez. The Jobs tab picks from this list, so each name must be different.",
         error="A name is 1 to 40 characters, and each client needs a different one."),
    dict(col="D", head="Type (pick)", kind="input", values=cvals(1), validation=("list", TYPES),
         prompt="Lead: not a client yet. Active: a client you work for. Past: finished, never listed as due.",
         error="Pick Lead, Active or Past."),
    dict(col="E", head="Source (pick)", kind="input", values=cvals(2), validation=("list", SOURCES),
         prompt="Where this client came from. The Today tab counts clients by source.", error="Pick a source from the list."),
    dict(col="F", head="Contact", kind="input", values=cvals(3), validation=("textLength", 0, 60),
         prompt="Phone or email, however you reach them. Optional; the example leaves it blank.", error="Up to 60 characters."),
    dict(col="G", head="First contact", kind="input", fmt=DATE, align="right", values=cvals(4), validation=date_v,
         prompt="The day they first got in touch, for example 7/12/2026. The Today tab counts new leads by this date.",
         error="Type a real date between 2000 and 2100, for example 7/12/2026."),
    dict(col="H", head="Next contact", kind="input", fmt=DATE, align="right", values=cvals(5), validation=date_v,
         prompt="When to call, text or visit next, for example 10/9/2026. On or before the as-of date puts them on the Today list. Leave blank for none.",
         error="Type a real date between 2000 and 2100, for example 10/9/2026, or leave blank."),
    dict(col="I", head="Status note", kind="input", values=cvals(6), validation=("textLength", 0, 80),
         prompt="A short note on where things stand, for example Quoted a deep clean, waiting. Optional.", error="Up to 80 characters."),
    dict(col="J", head="Jobs done", kind="calc", fmt=INT, align="right", formula=f_cjobs),
    dict(col="K", head="Revenue to date", kind="calc", fmt=USD2, align="right", formula=f_crev),
    dict(col="L", head="Unpaid", kind="calc", fmt=USD2, align="right", formula=f_cunpaid),
    dict(col="M", head="Last job", kind="calc", fmt=DATE, align="right", formula=f_clast),
    dict(col="N", head="Last asked about", kind="calc", formula=f_casked),
    dict(col="O", head="Days to next", kind="calc", fmt=DAYS, align="right", formula=f_cdays),
    dict(col="P", head="Follow-up", kind="calc", formula=f_cflag, bold=True),
]
f1, f2, cl_end = D.table_card(cl, C_TOP, "B", "Q", cols_cl, N_CLIENTS, title="Your clients",
                              sub=f"Yellow columns are yours; white columns are formulas and stay put when you clear the yellow cells. "
                                  f"{SHIP_CLIENTS} rows. Days to next and the follow-up flag count from the as-of date on the Today tab.")
assert (f1, f2) == (C_FIRST, C_LAST), (f1, f2)
left_cells(cl, C_FIRST, C_LAST, ("C", "D", "E", "F", "I"), color="input_text", bold=True)
left_cells(cl, C_FIRST, C_LAST, ("N",), size=9, color="ink2")
left_cells(cl, C_FIRST, C_LAST, ("P",), size=9, color="ink2", bold=True)
cl.conditional_formatting.add(f"L{C_FIRST}:L{C_LAST}", FormulaRule(formula=[f'AND(L{C_FIRST}<>"",L{C_FIRST}>0)'],
                              font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
cl.conditional_formatting.add(f"O{C_FIRST}:O{C_LAST}", FormulaRule(formula=[f'AND(O{C_FIRST}<>"",O{C_FIRST}<0,$D{C_FIRST}<>"Past")'],
                              font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
word_rules(cl, f"P{C_FIRST}:P{C_LAST}", C_FIRST, "P", [("Overdue", "accent", True, True), ("Due today", "accent", True, False),
                                                      ("Needs", "accent", True, True), ("Duplicate", "accent", True, True),
                                                      ("Not yet", "muted", False, False)])
D.paint_canvas(cl, cl_end + 1, "Z")
S.finish_sheet(cl, PRODUCT, cl_end + 1, span=("A", "R"), freeze=f"D{C_FIRST}", tab_color="gold")
cl.page_setup.orientation = "landscape"
cl.print_title_rows = f"{C_FIRST - 1}:{C_FIRST - 1}"
cl.print_area = f"A1:R{C_FIRST + 39}"
for k in ("N", "P"):
    cl[f"{k}{C_FIRST - 1}"].alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True, indent=1)

# =====================================================================  3 Jobs
S.set_widths(jb, {"A": 3, "B": 2, "C": 14, "D": 18, "E": 20, "F": 12, "G": 11, "H": 10, "I": 36, "J": 12, "K": 13,
                  "L": 20, "M": 2, "N": 3})
D.page_header(jb, "Jobs and quotes", "One row per job or quote. Pick the client and the service; the usual price and the balance fill in.",
              "Example rows: type over them" if not BLANK else f"Type your first job in row {J_FIRST}", span=("B", "M"), side_cols=3)
D.h(jb, 6, 10)
D.stat_strip(jb, 7, [
    ("C:D", "Jobs and quotes logged", f'=COUNTIF({JCLIENT},"?*")', INT),
    ("E:E", "Done", f'=COUNTIF({JSTATUS},"Done")', INT),
    ("F:G", "Revenue, done jobs", f'=SUMIFS({JAMT},{JSTATUS},"Done")', USD2),
    ("H:I", "Unpaid", f"=SUM({JOUT})", USD2),
    ("J:K", "Open quotes", f'=SUMIFS({JAMT},{JSTATUS},"Quoted")', USD2),
    ("L:L", "Rows to fix", f'=COUNTIF({JCHECK},"Needs*")+COUNTIF({JCHECK},"*not on*")+COUNTIF({JCHECK},"Paid:*")', INT),
])
D.h(jb, 9, 10)
D.frame(jb, 6, 9, "B", "M")
D.status_rule(jb, "H8", "$H$8>0", fill_tint=False)
D.status_rule(jb, "L8", "$L$8>0", fill_tint=False)
D.h(jb, 10, 14)


def f_usual(r):
    look = f"INDEX({SPRICE},MATCH($E{r},{SNAMES},0))"
    return f'=IF($E{r}="","",IFERROR(IF({look}="","",{look}),""))'


def f_out(r):
    return f'=IF(AND($G{r}="Done",$H{r}="No",$F{r}<>""),$F{r},"")'


def f_jcheck(r):
    return (f'=IF(COUNTA($C{r}:$I{r})=0,"",IF($C{r}="","Needs a date",IF($D{r}="","Needs a client",'
            f'IF(ISNA(MATCH($D{r},{CNAMES},0)),"Client not on list",IF($G{r}="","Needs a status",'
            f'IF(AND($E{r}<>"",ISNA(MATCH($E{r},{SNAMES},0))),"Service not on list",'
            f'IF(AND($G{r}="Done",$H{r}=""),"Paid: Yes or No?",IF(AND($F{r}="",OR($G{r}="Quoted",$G{r}="Done")),"Needs an amount","OK"))))))))')


money_v = ("decimal", 0, 1000000)
cols_jb = [
    dict(col="C", head="Date", kind="input", fmt=DATE, align="right", values=[x[0] for x in JOBS], validation=date_v,
         prompt="The job date, or the day you sent the quote, for example 10/2/2026.", error="Type a real date between 2000 and 2100, for example 10/2/2026."),
    dict(col="D", head="Client (pick)", kind="input", values=[x[1] for x in JOBS], validation=("list_range", CNAMES),
         prompt="Pick the client from the list. New clients go on the Clients tab first.",
         error="Pick a client from the list. Add new clients on the Clients tab first."),
    dict(col="E", head="Service (pick)", kind="input", values=[x[2] for x in JOBS], validation=("list_range", SNAMES),
         prompt="Pick the service from the list. Add new services on the Services tab first.",
         error="Pick a service from the list. Add new services on the Services tab first."),
    dict(col="F", head="Amount", kind="input", fmt=USD2, align="right", values=[x[3] for x in JOBS], validation=money_v,
         prompt="What you quoted or charged for this job, for example 140. Your usual price for the service shows at the right.",
         error="Type an amount from 0 to 1,000,000."),
    dict(col="G", head="Status (pick)", kind="input", values=[x[4] for x in JOBS], validation=("list", STATUSES),
         prompt="Quoted: waiting on the client. Booked: scheduled. Done: finished (counts as revenue). Lost: they said no.",
         error="Pick Quoted, Booked, Done or Lost."),
    dict(col="H", head="Paid (pick)", kind="input", values=[x[5] for x in JOBS], validation=("list", ["Yes", "No"]),
         prompt="For Done jobs: Yes or No. No puts the amount in Unpaid until you change it.", error="Pick Yes or No."),
    dict(col="I", head="Note", kind="input", values=[x[6] for x in JOBS], validation=("textLength", 0, 120),
         prompt="Anything worth remembering about this job or quote. Optional.", error="Up to 120 characters."),
    dict(col="J", head="Usual price", kind="calc", fmt=USD2, align="right", formula=f_usual),
    dict(col="K", head="Unpaid", kind="calc", fmt=USD2, align="right", formula=f_out),
    dict(col="L", head="Check", kind="calc", formula=f_jcheck),
]
g1, g2, jb_end = D.table_card(jb, J_TOP, "B", "M", cols_jb, N_JOBS, title="Your jobs and quotes",
                              sub=f"Yellow columns are yours; white columns fill in on their own. {SHIP_JOBS:,} rows, all read by the totals.")
assert (g1, g2) == (J_FIRST, J_LAST), (g1, g2)
left_cells(jb, J_FIRST, J_LAST, ("D", "E", "G", "H", "I"), color="input_text", bold=True)
left_cells(jb, J_FIRST, J_LAST, ("J",), size=9, color="ink2")
jb.conditional_formatting.add(f"K{J_FIRST}:K{J_LAST}", FormulaRule(formula=[f'AND(K{J_FIRST}<>"",K{J_FIRST}>0)'],
                              font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
left_cells(jb, J_FIRST, J_LAST, ("L",), size=9, color="ink2")
word_rules(jb, f"L{J_FIRST}:L{J_LAST}", J_FIRST, "L", [("Needs", "accent", True, False), ("Client not", "accent", True, False),
                                                      ("Service not", "accent", True, False), ("Paid:", "accent", True, False),
                                                      ("OK", "muted", False, False)])
D.paint_canvas(jb, jb_end + 1, "Z")
S.finish_sheet(jb, PRODUCT, jb_end + 1, span=("A", "N"), freeze=f"A{J_FIRST}", tab_color="blue")
jb.page_setup.orientation = "landscape"
jb.print_title_rows = f"{J_FIRST - 1}:{J_FIRST - 1}"
jb.print_area = f"A1:N{J_FIRST + 59}"
jb[f"L{J_FIRST - 1}"].alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True, indent=1)

# =====================================================================  4 Services
S.set_widths(sv, {"A": 3, "B": 2, "C": 26, "D": 14, "E": 13, "F": 12, "G": 15, "H": 2, "I": 3})
D.page_header(sv, "Services", "What you offer and what you usually charge. Feeds the Jobs tab.",
              "Example rows: type over them" if not BLANK else f"Type your first service in row {S_FIRST}", span=("B", "H"), side_cols=3)
D.h(sv, 6, 10)
D.stat_strip(sv, 7, [
    ("C:C", "Services listed", f'=COUNTIF({SNAMES},"?*")', INT),
    ("D:E", "Revenue, done jobs", f"=SUM({srng('G')})", USD2),
    ("F:G", "Jobs done", f"=SUM({srng('F')})", INT),
])
D.h(sv, 9, 10)
D.frame(sv, 6, 9, "B", "H")
D.h(sv, 10, 14)


def f_stimes(r):
    return f'=IF($C{r}="","",COUNTIFS({JSERV},$C{r},{JSTATUS},"Done"))'


def f_srev(r):
    return f'=IF($C{r}="","",SUMIFS({JAMT},{JSERV},$C{r},{JSTATUS},"Done"))'


cols_sv = [
    dict(col="C", head="Service", kind="input", values=[x[0] for x in SERVICES], validation=("textLength", 1, 30),
         prompt="A short name, for example Deep clean. The Jobs tab picks from this list, so each name must be different.",
         error="A service name is 1 to 30 characters."),
    dict(col="D", head="Standard price", kind="input", fmt=USD2, align="right", values=[x[1] for x in SERVICES], validation=money_v,
         prompt="What you usually charge for it, for example 140. The Jobs tab shows it beside each job as the usual price.",
         error="Type an amount from 0 to 1,000,000."),
    dict(col="E", head="Typical hours", kind="input", fmt=NUM1, align="right", values=[x[2] for x in SERVICES], validation=("decimal", 0, 1000),
         prompt="About how long it takes, in hours, for example 2.5. Optional; for your own planning.", error="Type hours from 0 to 1,000, or leave blank."),
    dict(col="F", head="Times done", kind="calc", fmt=INT, align="right", formula=f_stimes),
    dict(col="G", head="Revenue to date", kind="calc", fmt=USD2, align="right", formula=f_srev),
]
s1, s2, sv_end = D.table_card(sv, S_TOP, "B", "H", cols_sv, N_SERVICES, title="Your services",
                              sub=f"Yellow columns are yours; white columns fill in on their own. {N_SERVICES} rows.")
assert (s1, s2) == (S_FIRST, S_LAST), (s1, s2)
left_cells(sv, S_FIRST, S_LAST, ("C",), color="input_text", bold=True)
D.paint_canvas(sv, sv_end + 1, "Z")
S.finish_sheet(sv, PRODUCT, sv_end + 1, span=("A", "I"), freeze=f"A{S_FIRST}", tab_color="purple")
sv.print_area = f"A1:I{sv_end + 1}"

# =====================================================================  1 Today
# grid: the left column is the wide one (name | status | service, or name | value | bar); the right column holds
# the as-of date, the month and the open quotes list.
S.set_widths(td, {"A": 3, "B": 2, "C": 24, "D": 17, "E": 19, "F": 2, "G": 3, "H": 2, "I": 31, "J": 14, "K": 2, "L": 3})
D.page_header(td, "Today", "Who to follow up, what is quoted and unpaid, who brings the revenue.",
              AS_OF, span=("B", "K"), side_cols=3)
H_TOP = 120  # helper block for the sorted lists, below the printed page
HR = lambda i: H_TOP + 3 + i  # noqa: E731
H_FIRST, H_LAST = HR(0), HR(N_CLIENTS - 1)
KEY_DUE, KEY_REV = (f"${c}${H_FIRST}:${c}${H_LAST}" for c in ("C", "D"))
N_DUE, N_TOP, N_QUOTES = 10, 8, 8  # 10 due rows make the follow-up card the same height as the month card beside it
C_TOP_TD = 11

# ---- right column: the as-of date, this month, all time
R = D.Card(td, C_TOP_TD, "H", "I", "J", "K", wb=wb)
R.title("Your date and this month", "Type the date you are working from. The month follows it.")
r_asof = R.input("As-of date (usually today)", EXAMPLE_AS_OF if not BLANK else None, DATE, name="AsOfDate", validation=date_v,
                 allow_blank=True, prompt_title="As-of date",
                 error="Type a real date between 2000 and 2100, for example 10/9/2026, or leave blank.",
                 prompt="The date the follow-up list and the month count from, for example 10/9/2026. The workbook never reads your computer's clock, so type today's date each time you open it.")
assert r_asof == TD_ASOF_ROW, r_asof
MS = f"DATE(YEAR({ASOF}),MONTH({ASOF}),1)"
ME = f"DATE(YEAR({ASOF}),MONTH({ASOF})+1,0)"
IN_MONTH_J = f'{JDATE},">="&{MS},{JDATE},"<="&{ME}'
R.eyebrow("THIS MONTH")
r_month = R.calc("Month", f'=IF({ASOF}="","",{MS})', MONTH, color="ink2")
r_mjobs = R.calc("Jobs done", f'=IF({ASOF}="",0,COUNTIFS({JSTATUS},"Done",{IN_MONTH_J}))', INT)
r_mrev = R.calc("Revenue, done jobs", f'=IF({ASOF}="",0,SUMIFS({JAMT},{JSTATUS},"Done",{IN_MONTH_J}))', USD2, bold=True)
r_munp = R.calc("Unpaid from this month", f'=IF({ASOF}="",0,SUMIFS({JAMT},{JSTATUS},"Done",{JPAID},"No",{IN_MONTH_J}))', USD2)
r_mlead = R.calc("New leads (first contact)", f'=IF({ASOF}="",0,COUNTIFS({CNAMES},"?*",{CFIRST},">="&{MS},{CFIRST},"<="&{ME}))', INT)
R.eyebrow("ALL TIME")
r_acl = R.calc("Clients on the list", f'=COUNTIF({CNAMES},"?*")', INT)
r_ajobs = R.calc("Jobs done", f'=COUNTIF({JSTATUS},"Done")', INT)
r_arev = R.calc("Revenue, done jobs", f'=SUMIFS({JAMT},{JSTATUS},"Done")', USD2, bold=True)
r_aunp = R.calc("Unpaid balance", f'=SUM({JOUT})', USD2, total=True, tint="total_tint")
R.close()
for rr in (r_munp, r_aunp):  # unpaid: accent when > 0, plain ink at 0
    td.conditional_formatting.add(f"J{rr}", FormulaRule(formula=[f"$J${rr}>0"], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))

# ---- right column: open quotes by value (ranked straight from the Jobs tab, no helper rows)
Q = D.Card(td, R.end + 2, "H", "I", "J", "K", wb=wb)
Q.title("Open quotes by value", "Waiting on an answer, biggest first.")
QKEY = f'(({JSTATUS}="Quoted")*({JAMT}+({J_LAST + 1}-ROW({JAMT}))/1000000))'
N_Q = f'COUNTIF({JSTATUS},"Quoted")'
Q_ROWS = []
for k in range(1, N_QUOTES + 1):
    v = f"SUMPRODUCT(LARGE({QKEY},{k}))"
    idx = f"SUMPRODUCT(MAX(({QKEY}={v})*(ROW({JAMT})-{J_FIRST - 1})))"
    show = f"{N_Q}>={k}"
    rr = Q.calc(f'=IF({show},IFERROR(INDEX({JCLIENT},{idx})&IF(INDEX({JSERV},{idx})="","",": "&INDEX({JSERV},{idx})),""),"")',
                f'=IF({show},IFERROR(INDEX({JAMT},{idx})+0,""),"")', USD2)
    Q_ROWS.append(rr)
Q.text(f'=IF({N_Q}=0,"No open quotes",IF({N_Q}>{N_QUOTES},"Plus "&({N_Q}-{N_QUOTES})&" more on the Jobs tab (Status column)",'
       f'"Mark each one Booked, Done or Lost on the Jobs tab"))', size=9, color="ink2")
Q.close()
RIGHT_END = Q.end

# ---- left column: follow up today (name | status | last asked about)
L = D.Card(td, C_TOP_TD, "B", "C", "D", "F", extra="E", wb=wb)
L.title("Follow up today", "Next contact date on or before the as-of date, most overdue first.")
r_eye = L.eyebrow("CLIENT")
td.unmerge_cells(f"C{r_eye}:E{r_eye}")
for k, txt, al in (("D", "STATUS", "right"), ("E", "LAST ASKED ABOUT", "left")):
    c = td[f"{k}{r_eye}"]; c.value = txt; c.font = D.f(8, True, "muted"); c.alignment = Alignment(horizontal=al, vertical="bottom", indent=1 if al == "left" else 0)
    c.border = Border(bottom=D.side("hair"))
N_DUE_F = f'COUNTIF({KEY_DUE},">0")'
N_OVER_F = f'COUNTIF({KEY_DUE},">=2")'
DUE_ROWS = []
for k in range(1, N_DUE + 1):
    idx = f"MATCH(LARGE({KEY_DUE},{k}),{KEY_DUE},0)"
    show = f"{N_DUE_F}>={k}"
    over = f"({ASOF}-INDEX({CNEXT},{idx}))"
    rr = L.calc(f'=IF({show},INDEX({CNAMES},{idx}),"")',
                f'=IF({show},IF({over}=0,"Due today","Overdue "&{over}&IF({over}=1," day"," days")),"")', "@")
    e = td[f"E{rr}"]; e.value = f'=IF({show},INDEX({CASKED},{idx})&"","")'
    e.font = D.f(9, False, "ink2"); e.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    DUE_ROWS.append(rr)
L.text(f'=IF(COUNTIF({CNAMES},"?*")=0,"Add your clients on the Clients tab",IF({ASOF}="","Type the as-of date at the right to see who is due",'
       f'IF({N_DUE_F}=0,"Nobody is due. Check the Clients tab for the next dates coming up",'
       f'IF({N_DUE_F}>{N_DUE},"Plus "&({N_DUE_F}-{N_DUE})&" more on the Clients tab (Follow-up column)",'
       f'"Everyone due is listed. Next contact dates are on the Clients tab"))))', size=9, color="ink2")
L.close()
for rr in DUE_ROWS:
    td[f"D{rr}"].font = D.f(10, True, "ink2")
    td.conditional_formatting.add(f"D{rr}", FormulaRule(formula=[f'LEFT($D${rr},7)="Overdue"'], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))

# ---- left column: top clients by revenue (bars)
T = D.Card(td, L.end + 2, "B", "C", "D", "F", extra="E", wb=wb)
T.title("Top clients by revenue", "Done jobs, every row in the log. Bars compare each one to the top client.")
TOP_ROWS = []
t_first = T.r
for k in range(1, N_TOP + 1):
    idx = f"MATCH(LARGE({KEY_REV},{k}),{KEY_REV},0)"
    show = f'AND(COUNT({KEY_REV})>={k},IFERROR(INDEX({CREV},{idx}),0)>0)'
    rr = T.bar(f'=IF({show},INDEX({CNAMES},{idx}),"")', f'=IF({show},INDEX({CREV},{idx}),"")', USD2, f"$D${t_first}", "teal")
    TOP_ROWS.append(rr)
T.text(f'=IF(D{t_first}="","No done jobs yet",IFERROR("Top client: "&ROUND(D{t_first}/SUMIFS({JAMT},{JSTATUS},"Done")*100,0)&"% of revenue to date",""))',
       size=9, color="ink2")
T.close()

# ---- left column: where clients come from (bars)
W = D.Card(td, T.end + 2, "B", "C", "D", "F", extra="E", wb=wb)
W.title("Where clients come from", "Every client on the list, by the source you picked.")
w_first = W.r
w_last = w_first + len(SOURCES) - 1
WMAX = f"MAX($D${w_first}:$D${w_last})"
SRC_ROWS = []
for s in SOURCES:
    rr = W.bar(s, f'=COUNTIF({CSRC},"{s}")', INT, WMAX, "blue")
    SRC_ROWS.append(rr)
NO_SRC = f'(COUNTIF({CNAMES},"?*")-SUM($D${w_first}:$D${w_last}))'
W.text(f'=IF(COUNTIF({CNAMES},"?*")=0,"Add your clients on the Clients tab",IF({NO_SRC}>0,{NO_SRC}&IF({NO_SRC}=1," client has"," clients have")&" no source yet",'
       f'"Every client has a source"))', size=9, color="ink2")
W.close()
LEFT_END = W.end
CARDS_END = max(LEFT_END, RIGHT_END)

# ---- tiles (rows 6 to 9)
D.tile(td, 6, "B", "F", "FOLLOW-UPS DUE", f"={N_DUE_F}", INT,
       f'=IF(COUNTIF({CNAMES},"?*")=0,"Add your clients on the Clients tab",IF({ASOF}="","Type the as-of date to see who is due",'
       f'IF({N_DUE_F}=0,"Nobody is due on the as-of date",{N_OVER_F}&" overdue, "&({N_DUE_F}-{N_OVER_F})&" due today")))', dark=True)
v2, sub2, _ = D.tile(td, 6, "H", "K", "OPEN QUOTES", f'=SUMIFS({JAMT},{JSTATUS},"Quoted")', USD2,
                     f'=IF({N_Q}=0,"No quotes waiting on an answer",{N_Q}&IF({N_Q}=1," quote"," quotes")&" waiting on an answer")', dark=False)
td.conditional_formatting.add(v2.coordinate, FormulaRule(formula=[f"{v2.coordinate}>0"], font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))
D.h(td, 10, 14)

# ---- notes and sources (full width)
N_TOP_TD = CARDS_END + 2
NOTES = [
    ("Follow up today", "A lead or active client is due when the next contact date on the Clients tab is on or before the as-of "
     "date. Past clients are never listed. The list shows the 10 most overdue; the Clients tab flags every one."),
    ("As-of date", "You type it, so the workbook never reads your computer's clock and shows the same numbers on any day and "
     "in any app. Type today's date each time you open the file."),
    ("This month", "Jobs dated in the calendar month of the as-of date with status Done, their amounts, the unpaid ones among "
     "them, and clients whose first contact falls in that month."),
    ("Revenue and unpaid", "Only Done rows count as revenue. A Done row with Paid set to No is unpaid until you change it. "
     "Quoted and Booked rows wait; Lost rows are kept for the record and count nowhere."),
    ("Last asked about", "The service on the client's most recent row on the Jobs tab, whatever its status, so you know what "
     "to talk about when you call."),
    ("Any currency", "The maths works in any currency. Select the money cells and pick your symbol under Format, Number."),
    ("Before you rely on a number", S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab")),
]
r = N_TOP_TD
D.h(td, r, 12); r += 1
t = td[f"C{r}"]; t.value = "Notes"; t.font = D.f(12, True); td.merge_cells(f"C{r}:J{r}"); D.h(td, r, 24); r += 1
width = sum(S.width_of(td, k) for k in D.cols("C", "J"))
for head, body in NOTES:
    cell = td[f"C{r}"]; cell.value = rich(head, body)
    cell.alignment = Alignment(vertical="top", wrap_text=True); td.merge_cells(f"C{r}:J{r}")
    D.h(td, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
D.h(td, r, 12)
D.frame(td, N_TOP_TD, r, "B", "K")
TD_LAST = r + 1
D.h(td, TD_LAST, 14)
assert TD_LAST < H_TOP - 2, (TD_LAST, H_TOP)

# ---- helper block (below the printed page): sort keys for the lists above
r = H_TOP
c = td[f"C{r}"]; c.value = "Sorting helpers for the lists above. Formulas only, nothing to type."
c.font = D.f(9, True, "muted"); td.merge_cells(f"C{r}:J{r}")
r += 2
for k, txt in (("C", "DUE KEY"), ("D", "REVENUE KEY")):
    cc = td[f"{k}{r}"]; cc.value = txt; cc.font = D.f(8, True, "muted"); cc.alignment = Alignment(horizontal="right")
for i in range(N_CLIENTS):
    hr, cr = HR(i), C_FIRST + i
    tie = f"{N_CLIENTS - i}/10000000"
    name, typ, nxt, rev = (f"{CL}!${c}${cr}" for c in ("C", "D", "H", "K"))
    # due key: days overdue plus one (so a client due today scores above zero), most overdue first
    td[f"C{hr}"] = (f'=IF(OR({name}="",{typ}="Past",{nxt}="",{ASOF}=""),0,IF({nxt}<={ASOF},{ASOF}-{nxt}+1+{tie},0))')
    td[f"D{hr}"] = f'=IF({name}="","",N({rev})+{tie})'
    for k in ("C", "D"):
        td[f"{k}{hr}"].font = D.f(8, False, "muted"); td[f"{k}{hr}"].number_format = "0.000"
for rr in range(1, H_LAST + 1):
    if td.row_dimensions[rr].height is None:
        D.h(td, rr, D.GRID_ROW if rr < H_TOP else 14)
D.paint_canvas(td, H_LAST, "Z")
S.finish_sheet(td, PRODUCT, TD_LAST, span=("A", "L"), freeze="A11", tab_color="accent")
td.print_area = f"A1:L{TD_LAST}"
S.page_break_before(td, N_TOP_TD)
_pt = sum(td.row_dimensions[rr].height or D.GRID_ROW for rr in range(1, N_TOP_TD))
td.sheet_properties.pageSetUpPr.fitToPage = False
td.page_setup.scale = max(40, min(88, int(680 / _pt * 100)))

# the Clients strip counts follow-ups the same way the Today tab does (its due keys)
cl["I8"].value = f"=COUNTIF({TD}!{KEY_DUE},\">0\")"

# =====================================================================  Start Here and Terms
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            mm = re.match(r"^([A-Z][A-Za-z ,]{1,34})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if mm:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        Lc = sh[f"C{r}"]
        if kind == "num":
            Lc.value = left; Lc.font = D.f(18, True, "accent"); Lc.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            Lc.value = left[0]; Lc.number_format = left[1]; Lc.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                Lc.fill = D.fill("input_fill"); Lc.font = D.f(9, True, "input_text")
                Lc.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                Lc.font = D.f(9, False, "ink"); Lc.border = Border(bottom=D.side("hair"))
            else:
                Lc.fill = D.fill("ink"); Lc.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


def money(x, cents=True):
    return ("-" if x < 0 else "") + (f"${abs(x):,.2f}" if cents else f"${abs(x):,.0f}")


top_name, top_rev = EX["top"][0]
due_names = [n for n, _, _ in EX["due_list"]]
sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, "Client Tracker", "Start here. List your clients, log each job and quote, and open Today each morning.", f"Version {VERSION}, {CHECKED}",
              span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "A client tracker for one or two person service businesses: cleaning, lawn care, pressure washing, detailing, handyman "
    "work, pet sitting, photography. List each client once, log every job and quote on one tab, and the Today tab tells "
    "you who to follow up on the date you type, what is quoted and waiting, what is done and unpaid, which clients bring "
    "the revenue and where they come from. No pipeline to set up first.", "para")])
r = sh_card(sh, r, "Three steps", [
    (1, "Clients tab. One row per client: name, type (Lead, Active, Past), source, first contact and the next date to get in touch.", "num"),
    (2, "Jobs tab. One row per job or quote: date, client and service from the lists, amount, status (Quoted, Booked, Done, Lost) and paid.", "num"),
    (3, "Today tab. Type the as-of date. Read who is due, the open quotes, this month's revenue and unpaid balance, and your top clients.", "num"),
])
r = sh_card(sh, r, "How to read the cells", [
    ((140, USD2), "Yellow cells with blue numbers are yours to type in. Each one shows a hint when you select it.", "chip-in"),
    ((EX["unpaid"], USD2), "Numbers on white are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((EX["due"], INT), "Dark tiles are your answers. Follow-ups due sits at the top of the Today tab and stays on screen as you scroll.", "chip-key"),
    (None, "Each tab is protected without a password, so you can still change anything. To edit a locked cell, use Review, "
           "Unprotect Sheet in Excel, or Data, Protect sheets and ranges in Google Sheets.", "note"),
])
S.page_break_before(sh, r)
ex_text = (f"Client-Tracker-EXAMPLE.xlsx holds July 12 to October 9, 2026 of Brightside Cleaning, a made-up two-person house "
           f"cleaning business with {EX['clients']} clients and {EX['jobs']} jobs and quotes. With the as-of date set to October 9, "
           f"{EX['due']} clients are due for a follow-up ({EX['overdue']} overdue, {EX['today']} due today): " + ", ".join(due_names) + ". "
           f"{EX['quotes']} quotes worth {money(EX['quotes_value'], False)} are waiting on an answer, {EX['month_jobs']} jobs are done so far in October "
           f"for {money(EX['month_revenue'], False)}, {money(EX['unpaid'], False)} is done and unpaid, and the top client, {top_name}, has brought "
           f"{money(top_rev, False)} of the {money(EX['revenue'], False)} revenue to date. Every name is invented.")
r = sh_card(sh, r, "The example", [(None, ex_text, "para")])
r = sh_card(sh, r, "Which file to open", [
    (None, "Excel on a computer. Double-click Client-Tracker.xlsx. If Excel shows a yellow Protected View bar, click Enable Editing.", "para"),
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. No macros, no add-ons, no sign-up.", "para"),
    (None, "Mac, iPad or phone. Unzip on a computer first (phones often can't open a .zip), then upload the .xlsx to Google Drive "
           "and open it in the Google Sheets app, or open it in Excel for Mac or iPad.", "para"),
])
r = sh_card(sh, r, "Good to know", [
    (None, "Your date. The Today tab counts from the date you type, never from your computer's clock, so the file shows "
           "the same numbers on any day and in any app. Type today's date each time you open it.", "para"),
    (None, "Clearing the example. In the example file, select the yellow cells on Clients, Jobs and Services and press Delete. "
           "Formulas live only in the white columns, so every total keeps working. Or start from the blank file.", "para"),
    (None, f"Room to grow. {SHIP_CLIENTS} clients, {SHIP_JOBS:,} jobs and quotes, {N_SERVICES} services, every row read by the totals.", "para"),
    (None, "Quotes and payment. Log a quote the day you send it with status Quoted. When they say yes, change it to Booked, then "
           "Done when the job is finished. Only Done rows count as revenue, and a Done job with Paid set to No stays in Unpaid "
           "until you change it to Yes.", "para"),
])
r = sh_card(sh, r, "Before you rely on a number", [(None, S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab"), "note")])
SH_LAST = r - 1
D.paint_canvas(sh, SH_LAST, "Z")
S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
               footer_notice="Estimates only. Not professional advice. See the Terms tab or LICENSE-AND-DISCLAIMER.txt.")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.", f"Version {VERSION}, {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "Pricing tools from the shop", [
    (None, ("Service Pricing Calculator: the hourly rate and the job quote that pay you", "https://www.etsy.com/listing/4587962930"), "link"),
    (None, ("Hourly Rate Quick Calculator: your minimum rate in six numbers", "https://www.etsy.com/listing/4589695837"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this tracker saves you time, a short review on Etsy helps other small businesses find it. Thank you.", "para"),
    (None, "If a file won't open or something looks wrong, message me on Etsy and it will be fixed.", "para"),
])
TM_LAST = r - 1
D.paint_canvas(tm, TM_LAST, "Z")
S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb) if "input without validation" not in p]
if problems:
    print("AUDIT:", *problems[:20], sep="\n  ")
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
print(json.dumps({"asof": r_asof, "month": (r_month, r_mjobs, r_mrev, r_munp, r_mlead), "all": (r_acl, r_ajobs, r_arev, r_aunp),
                  "quotes": Q_ROWS, "due": DUE_ROWS, "top": TOP_ROWS, "sources": SRC_ROWS, "notes_top": N_TOP_TD, "last": TD_LAST}))
print("example:", json.dumps({k: v for k, v in EX.items() if k not in ("due_list",)}), json.dumps(EX["due_list"]))
