#!/usr/bin/env python3
"""#12 Craft Fair Profit Calculator and Booth Fee Break-Even Workbook, version 3 (dashboard design).

Same inputs, formulas and worked example as version 1 (Oct 1, 2026) as corrected by the shop run on
Oct 6, 2026 (cost per mile 76 cents for Jul 1 to Dec 31, 2026, so the example show costs $315.60).
Rebuilt to design/WORKBOOK_STANDARD.md section 0; tools/compare_xlsx.py with compare_map.json proves
the answers match. Fixes in v3 (declared in compare_map.json "intended"):
  - Price check no longer says "below formula" for a row with a price but no product name.
  - Mix check names the Mix column (v1 said "column F") and asks for a mix when none is typed.
  - Show Log example "All other costs" follows the 76 cent mileage ($165.60, was $163.50).
  - Sheets protected without a password; every input validated with a message; full Terms tab.

Usage: python3 engines/craft-fair-profit/build_xlsx.py out.xlsx [--cells cells.json]
"""
import datetime as dt
import json
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402

PRODUCT = "Craft Fair Profit Calculator"
LONG = "Craft Fair Profit Calculator and Booth Fee Break-Even Workbook"
VERSION = "3"
CHECKED = "Oct 6, 2026"
AS_OF = f"Mileage checked {CHECKED}"
HERE = os.path.dirname(os.path.abspath(__file__))
args = sys.argv[1:]
CELLS_OUT = None
if "--cells" in args:
    i = args.index("--cells"); CELLS_OUT = args[i + 1]; del args[i:i + 2]
OUT = args[0] if args else "Craft-Fair-Profit-Calculator.xlsx"
USD0, USD2, PCT, INT, NUM1 = (S.FMT[k] for k in ("usd0", "usd2", "pct", "int", "num1"))
PCT1, PCT2, DATE = "0.0%", "0.00%", S.FMT["date"]
MULT = '0.00"x"'
TAX_SOURCES = "the IRS Standard Mileage Rates page (IR-2026-29)"

# ---------------------------------------------------------------- example (v1, mileage per the Oct 6 fix)
PRODUCTS = [  # name, materials, minutes, packaging, price, mix
    ("Cold-process soap bar, 4.5 oz", 1.6, 6, 0.25, 8, 0.40),
    ("Soy candle, 8 oz tin", 6.5, 15, 0.6, 22, 0.25),
    ("Wax melt clamshell", 2.1, 6, 0.4, 10, 0.15),
    ("Lip balm tube", 0.9, 3, 0.15, 5, 0.12),
    ("Gift box: candle + 2 soaps", 10.2, 30, 2.5, 38, 0.08),
]
N_PROD = 15
N_LOG = 12
FEES = [50, 100, 150, 250, 400, 600]

wb = S.new_workbook(LONG, version=VERSION)
GRID = {"A": 3, "B": 2, "C": 31, "D": 14, "E": 2, "F": 3, "G": 2, "H": 24, "I": 12, "J": 17, "K": 2, "L": 3}
CELLS = {}


def rich(ws, ref, head, body, size=9, body_color="ink2"):
    ws[ref].value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=size, b=True, color=D.T["ink"]), head + ". "),
                                 TextBlock(InlineFont(rFont=D.FONT, sz=size, color=D.T[body_color]), body))


def notes_card(ws, top, notes, span=("C", "J"), frame_span=("B", "K")):
    r = top
    D.h(ws, r, 12); r += 1
    t = ws[f"{span[0]}{r}"]; t.value = "Notes and sources"; t.font = D.f(12, True)
    ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}"); D.h(ws, r, 24); r += 1
    D.h(ws, r, 6); r += 1
    width = sum(S.width_of(ws, k) for k in D.cols(*span))
    for head, body in notes:
        rich(ws, f"{span[0]}{r}", head, body)
        ws[f"{span[0]}{r}"].alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}")
        D.h(ws, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
    D.h(ws, r, 12)
    D.frame(ws, top, r, *frame_span)
    return r


def fill_heights(ws, last):
    for rr in range(1, last + 1):
        if ws.row_dimensions[rr].height is None:
            D.h(ws, rr, D.GRID_ROW)


NOTICE = ("For general information and planning only. This is not legal, tax, financial, medical or other professional "
          "advice, and using it doesn't create a professional relationship. Results are estimates based on the numbers you "
          f"enter. Figures were checked on the dates shown and can change. See the Terms tab before you rely on anything here. Calculations are estimates for "
          "planning and depend on your inputs. They aren't accounting, tax or financial advice. Check fees, rates and interest "
          "against your own statements and the provider's current terms.")

# =====================================================================  1 Show Costs
sc = wb.create_sheet("1 Show Costs")
S.set_widths(sc, GRID)
D.page_header(sc, "Show Costs", "Everything the day costs you before the first sale.", AS_OF)
TOP = 11
A = D.Card(sc, TOP, "B", "C", "D", "E", wb=wb)
A.title("Your show", "Type over the yellow cells. Every total updates.")
A.eyebrow("SHOW NAME AND DATE")
r = A._row()
c = S.input_cell(sc, f"C{r}", "Riverside Fall Market", "@", name="ShowName", wb=wb,
                 validation=("textLength", 0, 40), prompt_title="Show name",
                 prompt="The fair, market or festival. Up to 40 characters.", error="Up to 40 characters.", align="left")
c.fill = D.fill("input_fill"); c.font = D.f(10, True, "input_text")
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
c.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
c = S.input_cell(sc, f"D{r}", dt.date(2026, 10, 17), DATE, name="ShowDate", wb=wb,
                 validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"), prompt_title="Show date",
                 prompt="The show day, for example 10/17/2026.", error="Enter a date, for example 10/17/2026.")
c.fill = D.fill("input_fill"); c.font = D.f(10, True, "input_text")
c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
c.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
SN_ROW = r
HRS = A.input("Selling hours, doors to close", 7, NUM1, name="SellingHours", validation=("decimal", 0, 24),
              prompt="Hours the doors are open, for example 7 or 6.5. Used for sales per hour.", prompt_title="Selling hours")
A.eyebrow("FIXED COSTS OF THE SHOW")
BF = A.input("Booth fee", 150, USD2, name="BoothFee", validation=("decimal", 0, None),
             prompt="From the show's application. Typical ranges are on the Start Here tab. Example: 150.")
OX = A.input("Organizer extras", 25, USD2, name="OrganizerExtras", validation=("decimal", 0, None),
             prompt="Electricity, tent or table rental, jury or application fee: anything the organizer charges on top of the booth.")
MI = A.input("Round-trip miles", 60, INT, name="Miles", validation=("decimal", 0, None),
             prompt="Home to the show and back, in miles. Example: 60.")
CPM = A.input("Cost per mile", 0.76, USD2, name="CostPerMile", validation=("decimal", 0, 5),
              prompt="Default is the IRS business standard mileage rate, 76 cents a mile for Jul 1 to Dec 31, 2026 (IR-2026-29, "
                     "checked Oct 6, 2026). A cost estimate, not tax advice. Use your own cost per mile if you track it.")
PK = A.input("Parking and tolls", 10, USD2, name="Parking", validation=("decimal", 0, None),
             prompt="Parking, tolls and ferry fees for the day, in dollars.")
LG = A.input("Lodging", 0, USD2, name="Lodging", validation=("decimal", 0, None),
             prompt="Hotel or campground for an away show, in dollars. 0 if you sleep at home.")
ML = A.input("Meals and coffee", 20, USD2, name="Meals", validation=("decimal", 0, None),
             prompt="Food and drinks for you and any helper on the day, in dollars.")
HP = A.input("Helper pay", 0, USD2, name="HelperPay", validation=("decimal", 0, None),
             prompt="What you pay a helper for the day, in dollars. 0 if you work alone.")
DS = A.input("Display spent this year", 300, USD2, name="DisplaySpend", validation=("decimal", 0, None),
             prompt="Total spent this year on things you reuse: tent, tables, racks, banner, tablecloths.")
NS = A.input("Shows you'll do this year", 6, INT, name="ShowsPerYear", validation=("whole", 0, 365),
             prompt="Whole number. Spreads the display cost across the season. Type 0 to leave display out.",
             prompt_title="Shows this year")
OT = A.input("Other: bags, permit, samples", 15, USD2, name="OtherCosts", validation=("decimal", 0, None),
             prompt="Bags, tissue, samples, and a temporary seller's permit fee if your state charges one.")
A.close()

B = D.Card(sc, TOP, "G", "H", "I", "K", extra="J", wb=wb)
B.title("How the day adds up", "Where the money goes before your first sale.")
b1 = B.r; b9 = B.r + 8
bmx = f"MAX($I${b1}:$I${b9})"
B.bar("Booth fee", f"=N($D${BF})", USD2, bmx, "accent")
B.bar("Organizer extras", f"=N($D${OX})", USD2, bmx, "gold")
b_tr = B.bar("Travel, miles x cost", f"=N($D${MI})*N($D${CPM})", USD2, bmx, "blue")
B.bar("Parking and tolls", f"=N($D${PK})", USD2, bmx, "blue")
B.bar("Lodging", f"=N($D${LG})", USD2, bmx, "blue")
B.bar("Meals and coffee", f"=N($D${ML})", USD2, bmx, "gold")
B.bar("Helper pay", f"=N($D${HP})", USD2, bmx, "gold")
b_ds = B.bar("Display share", f"=IF(N($D${NS})>0,N($D${DS})/$D${NS},0)", USD2, bmx, "teal")
B.bar("Other", f"=N($D${OT})", USD2, bmx, "gold")
b_tot = B.calc("Total fixed show cost", f"=SUM(I{b1}:I{b9})", USD2, total=True, tint="total_tint")
B.close()

C = D.Card(sc, B.end + 2, "G", "H", "I", "K", extra="J", wb=wb)
C.title("Card reader fee", "A cost that grows with every card sale.")
CP = C.input("Card fee percent", 0.026, PCT1, name="CardFeePct", validation=("decimal", 0, 0.2),
             prompt="Type a percent, for example 2.6%. Square's in-person rate on its free plan is 2.6% + 15 cents (read Oct 1, 2026). Match your reader.")
CF = C.input("Fee per card sale", 0.15, USD2, name="CardFeeEach", validation=("decimal", 0, 5),
             prompt="The flat part of the fee on each card payment, in dollars. Example: 0.15.")
CS = C.input("Share of sales by card", 0.7, PCT, name="CardShare", validation=("decimal", 0, 1),
             prompt="Type a percent. Most fairs run 60% to 80% card now; cash carries no fee. Example: 70%.")
AV = C.input("Average sale per customer", 30, USD2, name="AverageSale", validation=("decimal", 0, None),
             prompt="What a typical customer spends, in dollars. Turns the flat fee into a cost per dollar. Example: 30.")
CC = C.calc("Card cost per $1 of sales", f'=N(I{CS})*(N(I{CP})+IF(N(I{AV})>0,N(I{CF})/I{AV},0))', PCT2, total=True,
            tint="total_tint")
C.close()

T = D.Card(sc, A.end + 2, "B", "C", "D", "E", wb=wb)
T.title("Sales tax (information only)", "Tax you collect is not yours. Prices here are pre-tax.")
TXR = T.input("Sales tax rate at the show", 0.085, PCT1, name="SalesTaxRate", validation=("decimal", 0, 0.2),
              prompt="Type a percent, for example 8.5%. Most states require a seller's permit at fairs; check yours before the show.")
PTS = T.calc("Pre-tax part of each $1", f"=1/(1+N(D{TXR}))", '"$"0.0000')
T.text("Prices that include tax? Multiply by this first.", size=9)
T.close()

TOTAL = f"$I${b_tot}"
D.tile(sc, 6, "B", "E", "WHAT THIS SHOW REALLY COSTS", f"={TOTAL}", USD2,
       f'=IF(OR({TOTAL}=0,N($D${BF})=0),"",'
       f'"The "&TEXT($D${BF},"$#,##0")&" booth fee is "&TEXT($D${BF}/{TOTAL},"0%")&" of it")', dark=True)
HID = f"({TOTAL}-N($D${BF}))"
v, sub, _ = D.tile(sc, 6, "G", "K", "COSTS ON TOP OF THE BOOTH FEE", f"={HID}", USD2,
                   f'=IF({TOTAL}=0,"",IF({HID}>N($D${BF}),"More than the booth fee itself","Less than the booth fee"))')
D.status_rule(sc, v.coordinate, f"{HID}>N($D${BF})", fill_tint=False)
D.h(sc, 10, 14)

n_top = max(T.end, C.end) + 2
SC_NOTES = [
    ("Cost per mile", "Defaults to the IRS business standard mileage rate of 76 cents a mile for July 1 to December 31, 2026 "
     "(news release IR-2026-29, irs.gov/tax-professionals/standard-mileage-rates); it was 72.5 cents from January 1 to June 30, "
     "2026 (IR-2025-128). Checked Oct 6, 2026. Here it is a cost estimate for the drive, not tax advice."),
    ("Card fee", "Square's in-person rate on the free plan is 2.6% + 15 cents per tap, dip or swipe; manual entry is 3.5% + 15 cents "
     "(squareup.com/us/en/fees, read Oct 1, 2026). Set your own reader's numbers. Card cost per $1 = card share x (percent fee "
     "+ flat fee / average sale), and it flows into every product's margin."),
    ("Sales tax", "Most states make a fair seller hold a seller's permit or a temporary vendor permit and collect and remit sales "
     "tax (Avalara, Craft fairs and sales tax: a state-by-state guide, updated Aug 3, 2026). This workbook treats every price as "
     "pre-tax. Check your state before the show."),
    ("Before you rely on a number", NOTICE),
]
n_end = notes_card(sc, n_top, SC_NOTES)
S.page_break_before(sc, n_top)
fill_heights(sc, n_end)
D.h(sc, n_end + 1, 14)
D.paint_canvas(sc, n_end + 1, "Z")
S.finish_sheet(sc, LONG, n_end + 1, span=("A", "L"), freeze="A11", tab_color="accent")
SCR = "'1 Show Costs'!"
R_TOTAL, R_CC, R_HRS, R_BF = f"{SCR}$I${b_tot}", f"{SCR}$I${CC}", f"{SCR}$D${HRS}", f"{SCR}$D${BF}"
R_CP, R_CF, R_AV = f"{SCR}$I${CP}", f"{SCR}$I${CF}", f"{SCR}$I${AV}"
CELLS["sc"] = dict(hours=f"D{HRS}", booth=f"D{BF}", extras=f"D{OX}", miles=f"D{MI}", cpm=f"D{CPM}", parking=f"D{PK}",
                   lodging=f"D{LG}", meals=f"D{ML}", helper=f"D{HP}", display=f"D{DS}", shows=f"D{NS}", other=f"D{OT}",
                   card_pct=f"I{CP}", card_fee=f"I{CF}", card_share=f"I{CS}", avg_sale=f"I{AV}", tax=f"D{TXR}",
                   travel=f"I{b_tr}", display_share=f"I{b_ds}", total=f"I{b_tot}", card_cost=f"I{CC}", pretax=f"D{PTS}")

# =====================================================================  2 Products
pr = wb.create_sheet("2 Products")
S.set_widths(pr, {"A": 3, "B": 2, "C": 30, "D": 11, "E": 10, "F": 11, "G": 10, "H": 9, "I": 10, "J": 10, "K": 10,
                  "L": 13, "M": 9, "N": 11, "O": 17, "P": 2, "Q": 3})
D.page_header(pr, "Products", "What each piece earns after materials, your time and the card fee.",
              "Example products are invented", span=("B", "P"), side_cols=4)
PT = 19
P1 = PT + 5
PN = P1 + N_PROD - 1
rng = {k: f"${k}${P1}:${k}${PN}" for k in "CDEFGHIJKLMNO"}
HR_ROW = 14  # set by the card below; checked with an assert
SUMMIX = f"SUM({rng['H']})"
BPRICE = f"IF({SUMMIX}>0,SUMPRODUCT({rng['H']},{rng['G']})/{SUMMIX},0)"
D.h(pr, 6, 10)
D.stat_strip(pr, 7, [
    ("C:C", "Blended price per piece", f"={BPRICE}", USD2),
    ("D:G", "Blended unit cost", f"=IF({SUMMIX}>0,(SUMPRODUCT({rng['H']},{rng['D']})+SUMPRODUCT({rng['H']},{rng['F']})"
                                 f"+SUMPRODUCT({rng['H']},{rng['E']})/60*$D${HR_ROW})/{SUMMIX},0)", USD2),
    ("H:L", "Contribution per piece", f"=C8-E8-C8*{R_CC}", USD2),
    ("M:O", "Priced below the x2 check", f'=COUNTIF({rng["O"]},"below formula")', INT),
])
# stat_strip writes each value in the first column of its span
pr["H8"].value = f"=C8-D8-C8*{R_CC}"
D.h(pr, 9, 10)
D.frame(pr, 6, 9, "B", "P")
pr.conditional_formatting.add("H8", FormulaRule(formula=[f"{SUMMIX}=0"], font=Font(color=D.T["ink"], bold=True), stopIfTrue=True))
D.status_rule(pr, "H8", f"AND({SUMMIX}>0,H8<=0)", fill_tint=False)
D.status_rule(pr, "M8", "M8>0", fill_tint=False)
D.h(pr, 10, 14)

K = D.Card(pr, 11, "B", "C", "D", "P", wb=wb)
K.title("Your making time", "Labor in every unit cost uses this rate.")
HR = K.input("Your hourly rate for making", 20, USD2, name="HourlyRate", validation=("decimal", 0, None),
             prompt="What an hour of your making time is worth, in dollars. $15 to $30 is common for hobby-to-business makers.")
assert HR == HR_ROW, HR
MIXT = K.calc("Mix typed so far", f"={SUMMIX}", PCT)
MIXC = K.text(f'=IF({SUMMIX}=0,"Type a Mix share for each product",'
              f'IF(ABS({SUMMIX}-1)<0.005,"Mix totals 100%, good","Mix is not 100%: fix the Mix column"))',
              size=10, color="ink", bold=True)
K.close()
D.status_rule(pr, f"C{MIXC}:D{MIXC}", f"ABS({SUMMIX}-1)>=0.005", fill_tint=False)
HRR = f"$D${HR}"
assert K.end + 2 == PT, (K.end, PT)


def pcalc(kind):
    def fn(r):
        return {
            "labor": f'=IF(E{r}="","",E{r}/60*{HRR})',
            "unit": f'=IF(C{r}="","",N(D{r})+N(F{r})+N(I{r}))',
            "card": f'=IF(G{r}="","",G{r}*{R_CC})',
            "contrib": f'=IF(G{r}="","",G{r}-N(J{r})-N(K{r}))',
            "margin": f'=IF(OR(G{r}="",G{r}=0),"",L{r}/G{r})',
            "formula": f'=IF(C{r}="","",(N(D{r})+N(I{r}))*2)',
            "check": f'=IF(OR(G{r}="",N{r}=""),"",IF(G{r}<N{r},"below formula","ok"))',
        }[kind]
    return fn


cols_pr = [
    dict(col="C", head="Product", kind="input", values=[p[0] for p in PRODUCTS], validation=("textLength", 0, 40),
         prompt="What you sell, up to 40 characters. A row counts once it has a name.", error="Up to 40 characters."),
    dict(col="D", head="Materials", kind="input", fmt=USD2, align="right", values=[p[1] for p in PRODUCTS],
         validation=("decimal", 0, None), prompt="Materials for one piece, in dollars.", error="Enter 0 or more dollars."),
    dict(col="E", head="Minutes to make", kind="input", fmt=INT, align="right", values=[p[2] for p in PRODUCTS],
         validation=("decimal", 0, 6000), prompt="Your making time for one piece, in minutes.", error="Enter minutes from 0 to 6,000."),
    dict(col="F", head="Packaging", kind="input", fmt=USD2, align="right", values=[p[3] for p in PRODUCTS],
         validation=("decimal", 0, None), prompt="Bag, box, label or tissue for one piece, in dollars.", error="Enter 0 or more dollars."),
    dict(col="G", head="Your price", kind="input", fmt=USD2, align="right", values=[p[4] for p in PRODUCTS],
         validation=("decimal", 0, None), prompt="Your table price before sales tax, in dollars.", error="Enter 0 or more dollars."),
    dict(col="H", head="Mix", kind="input", fmt=PCT, align="right", values=[p[5] for p in PRODUCTS],
         validation=("decimal", 0, 1), prompt="Share of the pieces you sell that are this one. Type a percent; the column should total 100%.",
         error="Enter a percent from 0% to 100%."),
    dict(col="I", head="Labor", kind="calc", fmt=USD2, align="right", formula=pcalc("labor")),
    dict(col="J", head="Unit cost", kind="calc", fmt=USD2, align="right", formula=pcalc("unit")),
    dict(col="K", head="Card cost", kind="calc", fmt=USD2, align="right", formula=pcalc("card")),
    dict(col="L", head="Contribution", kind="calc", fmt=USD2, align="right", formula=pcalc("contrib"), bold=True),
    dict(col="M", head="Margin", kind="calc", fmt=PCT, align="right", formula=pcalc("margin")),
    dict(col="N", head="x2 price", kind="calc", fmt=USD2, align="right", formula=pcalc("formula")),
    dict(col="O", head="Price check", kind="calc", formula=pcalc("check")),
]
f1, f2, p_end = D.table_card(pr, PT, "B", "P", cols_pr, N_PROD, title="Your products",
                             sub="Yellow columns are yours, up to 15 products. x2 price is (materials + labor) x 2.")
assert (f1, f2) == (P1, PN), (f1, f2, P1, PN)
pr["M4"].value = f'=IF(COUNTA(C{P1}:C{PN})=0,"","Example products are invented")'
pr[f"O{P1 - 1}"].alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True, indent=1)
for r in range(P1, PN + 1):
    pr[f"O{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    if pr[f"C{r}"].value in ("",):
        pr[f"C{r}"].value = None
for word, color, tint in (("below formula", "accent", "accent_tint"), ("ok", "teal", "teal_tint")):
    pr.conditional_formatting.add(f"O{P1}:O{PN}", FormulaRule(formula=[f'$O{P1}="{word}"'],
                                  font=Font(color=D.T[color], bold=True), fill=PatternFill("solid", bgColor=D.T[tint]), stopIfTrue=True))
D.status_rule(pr, f"L{P1}:L{PN}", f"L{P1}<0", fill_tint=False)
pn_top = p_end + 2
PR_NOTES = [
    ("x2 price", "(Materials + Labor) x 2 = Retail Price, from CraftProfessional.com, Craft Pricing Formula (updated Apr 11, "
     "2025). It is a check, not a rule: a price under it means your time is paying for part of the buyer's purchase."),
    ("Blended figures", "Weighted by the Mix column, so a best seller counts for more than a slow one. Card cost per piece comes "
     "from the card reader fee on tab 1. The example products are invented to show the maths; type over them."),
]
pn_end = notes_card(pr, pn_top, PR_NOTES, span=("C", "O"), frame_span=("B", "P"))
fill_heights(pr, pn_end)
D.h(pr, pn_end + 1, 14)
D.paint_canvas(pr, pn_end + 1, "Z")
S.finish_sheet(pr, LONG, pn_end + 1, span=("A", "Q"), freeze="A10", tab_color="teal")
pr.page_setup.orientation = "landscape"
pr.page_setup.fitToHeight = 1
pr.print_title_rows = f"{P1 - 1}:{P1 - 1}"
PRR = "'2 Products'!"
R_BPRICE, R_BCOST, R_BCONT = f"{PRR}$C$8", f"{PRR}$D$8", f"{PRR}$H$8"
CELLS["pr"] = dict(rate=f"D{HR}", first=P1, last=PN, price="C8", cost="D8", contrib="H8", below="M8", mix=f"D{MIXT}",
                   mixcheck=f"C{MIXC}")

# =====================================================================  3 Break-Even
be = wb.create_sheet("3 Break-Even")
S.set_widths(be, GRID)
D.page_header(be, "Break-Even", "How many pieces pay for the day, and how many make it worth doing.", AS_OF)
G = D.Card(be, TOP, "B", "C", "D", "E", wb=wb)
G.title("Your goal for the day", "Two numbers you choose. The rest comes from tabs 1 and 2.")
TP = G.input("Profit you want on top", 300, USD0, name="TargetProfit", validation=("decimal", 0, None),
             prompt="Profit for the day after costs and after your making time is paid, in dollars. Example: 300.",
             prompt_title="Profit you want")
CU = G.input("Stock cushion", 1.5, '0.0"x"', name="StockCushion", validation=("decimal", 1, 5),
             prompt="How many pieces to bring per piece you plan to sell. 1.5 means half again, so the table never looks empty.")
G.eyebrow("FROM YOUR OTHER TABS")
g_tot = G.calc("Total fixed show cost", f"={R_TOTAL}", USD2)
g_con = G.calc("Contribution per piece", f"={R_BCONT}", USD2)
g_pri = G.calc("Blended price per piece", f"={R_BPRICE}", USD2)
g_hrs = G.calc("Selling hours", f"=N({R_HRS})", NUM1)
g_bf = G.calc("Booth fee", f"=N({R_BF})", USD2)
G.close()

E = D.Card(be, TOP, "G", "H", "I", "K", extra="J", wb=wb)
E.title("Break-even: the day pays for itself", "Pieces and sales that cover every cost of the day.")
e_u = E.calc("Pieces you must sell", f'=IF(D{g_con}>0,ROUNDUP(D{g_tot}/D{g_con},0),"")', INT, total=True, tint="total_tint")
e_s = E.calc("Sales that means", f'=IF(I{e_u}="","",I{e_u}*D{g_pri})', USD2)
e_h = E.calc("Pieces per selling hour", f'=IF(OR(I{e_u}="",D{g_hrs}=0),"",I{e_u}/D{g_hrs})', NUM1)
e_m = E.calc("Minutes between sales", f'=IF(OR(I{e_h}="",I{e_h}=0),"",60/I{e_h})', NUM1)
E.eyebrow("THE 3X RULE")
e_x = E.calc("Break-even vs booth fee", f'=IF(OR(I{e_s}="",D{g_bf}=0),"",I{e_s}/D{g_bf})', MULT)
e_3 = E.calc("3x rule target", f"=3*D{g_bf}", USD2)
e_3u = E.calc("Pieces to hit 3x", f'=IF(D{g_pri}>0,ROUNDUP(I{e_3}/D{g_pri},0),"")', INT)
E.close()

W = D.Card(be, G.end + 2, "B", "C", "D", "E", wb=wb)
W.title("Make the day worth doing", "Pieces for your profit, on top of your making time.")
w_u = W.calc("Pieces for your profit", f'=IF(D{g_con}>0,ROUNDUP((D{g_tot}+N(D{TP}))/D{g_con},0),"")', INT, total=True,
             tint="total_tint")
w_s = W.calc("Sales that means", f'=IF(D{w_u}="","",D{w_u}*D{g_pri})', USD2)
w_h = W.calc("Pieces per selling hour", f'=IF(OR(D{w_u}="",D{g_hrs}=0),"",D{w_u}/D{g_hrs})', NUM1)
w_b = W.calc("Pieces to bring", f'=IF(D{w_u}="","",ROUNDUP(D{w_u}*N(D{CU}),0))', INT, bold=True)
W.close()

RD = D.Card(be, E.end + 2, "G", "H", "I", "K", extra="J", wb=wb)
RD.title("Read it this way", "The rule of thumb against your own numbers.")
rr = RD._row(); RD._row(); RD._row()
c = be[f"H{rr}"]
c.value = (f'=IF(I{e_s}="",IF(AND(D{g_pri}>0,D{g_con}<=0),"Each piece costs more than it earns. Raise prices on tab 2 or cut unit costs before you book this show.","Add products with a mix on tab 2 and your costs on tab 1 to see this."),'
           f'"The 3x rule says "&TEXT(I{e_3},"$#,##0")&" is a good day. Your break-even needs "&TEXT(I{e_s},"$#,##0")'
           f'&" just to cover costs and pay your making time."&IF(I{e_s}>I{e_3}," That gap is why a busy show can still lose money.",""))')
c.font = D.f(10, False, "ink"); c.alignment = Alignment(vertical="top", wrap_text=True)
be.merge_cells(f"H{rr}:J{rr + 2}")
RD.close()

# what-if table, full width
q_top = max(W.end, RD.end) + 2
r = q_top
D.h(be, r, 12); r += 1
t = be[f"C{r}"]; t.value = "What if the booth fee were different?"; t.font = D.f(12, True); be.merge_cells(f"C{r}:J{r}"); D.h(be, r, 24); r += 1
t = be[f"C{r}"]; t.value = "Everything else held. Type any fee in the yellow column; the row matching your booth fee is tinted."
t.font = D.f(9, False, "muted"); be.merge_cells(f"C{r}:J{r}"); D.h(be, r, 18); r += 1
D.h(be, r, 8); r += 1
heads = {"C": ("BOOTH FEE", "left"), "D": ("TOTAL SHOW COST", "right"), "H": ("PIECES TO BREAK EVEN", "right"),
         "I": ("SALES NEEDED", "right"), "J": ("", "left")}
for k, (txt, al) in heads.items():
    cc = be[f"{k}{r}"]; cc.value = txt; cc.font = D.f(8, True, "muted"); cc.alignment = Alignment(horizontal=al, vertical="bottom", wrap_text=True)
for k in D.cols("C", "J"):
    be[f"{k}{r}"].border = Border(bottom=D.side("ink"))
D.h(be, r, 30); r += 1
wf1 = r
wfn = r + len(FEES) - 1
for fee in FEES:
    D.h(be, r, D.GRID_ROW)
    cc = S.input_cell(be, f"C{r}", fee, USD0, validation=("decimal", 0, None), prompt_title="Booth fee to test",
                      prompt="Any booth fee you are weighing, in dollars.", error="Enter 0 or more dollars.")
    cc.fill = D.fill("input_fill"); cc.font = D.f(10, True, "input_text")
    cc.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    cc.border = Border(bottom=D.side("FFFFFF"), right=D.side("FFFFFF"))
    be[f"D{r}"] = f"=$D${g_tot}-$D${g_bf}+N(C{r})"; be[f"D{r}"].number_format = USD2
    be[f"H{r}"] = f'=IF($D${g_con}>0,ROUNDUP(D{r}/$D${g_con},0),"")'; be[f"H{r}"].number_format = INT
    be[f"I{r}"] = f'=IF(H{r}="","",H{r}*$D${g_pri})'; be[f"I{r}"].number_format = USD2
    be[f"J{r}"] = f'=IFERROR(REPT("{D.BAR_CHAR}",MAX(0,ROUND(H{r}/MAX($H${wf1}:$H${wfn})*11,0))),"")'
    for k in "DHI":
        be[f"{k}{r}"].font = D.f(10); be[f"{k}{r}"].alignment = Alignment(horizontal="right", vertical="center")
    be[f"J{r}"].font = D.f(9, False, "blue"); be[f"J{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for k in D.cols("C", "J"):
        if k != "C":
            be[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    r += 1
D.highlight_rule(be, f"D{wf1}:I{wfn}", f"AND($D${g_bf}>0,$C{wf1}=$D${g_bf})")
D.h(be, r, 14)
q_end = r
D.frame(be, q_top, q_end, "B", "K")
S.page_break_before(be, q_top)

BU = f"$I${e_u}"
D.tile(be, 6, "B", "E", "PIECES TO SELL TO BREAK EVEN", f"={BU}", INT,
       f'=IF({BU}="",IF(AND({R_BPRICE}>0,$D${g_con}<=0),"Each piece loses money after costs, so nothing breaks even","Add products with a mix on tab 2 to see your number"),'
       f'TEXT($I${e_s},"$#,##0.00")&" of sales, about "&TEXT($I${e_h},"0.0")&" pieces an hour")', dark=True)
XV = f"$I${e_x}"
v, sub, _ = D.tile(be, 6, "G", "K", f'=IF($D${g_bf}>0,"BREAK-EVEN SALES VS THE "&TEXT($D${g_bf},"$#,##0")&" BOOTH FEE","BREAK-EVEN SALES VS THE BOOTH FEE")', f"={XV}", MULT,
                   f'=IF({XV}="","",IF({XV}>3,"Above the 3x rule: an expensive show for what it is",'
                   f'"Within the 3x rule: the fee is in proportion"))')
D.status_rule(be, v.coordinate, f"{XV}>3", fill_tint=False)
D.h(be, 10, 14)

bn_top = q_end + 2
BE_NOTES = [
    ("The 3x rule", "Aim to gross at least three times the booth fee (The Craft Map, Craft Fair Booth Fees Guide, Feb 13, 2026). "
     "A show whose break-even already sits above 3x the fee is expensive for what it is."),
    ("Profit on top", "Your making time is already paid inside unit cost on tab 2, so the profit you type here is extra."),
    ("Rounding", "Pieces round up to whole pieces, so sales needed can sit a little above the exact break-even."),
]
bn_end = notes_card(be, bn_top, BE_NOTES)
fill_heights(be, bn_end)
D.h(be, bn_end + 1, 14)
D.paint_canvas(be, bn_end + 1, "Z")
S.finish_sheet(be, LONG, bn_end + 1, span=("A", "L"), freeze="A11", tab_color="accent")
CELLS["be"] = dict(target=f"D{TP}", cushion=f"D{CU}", total=f"D{g_tot}", contrib=f"D{g_con}", price=f"D{g_pri}",
                   hours=f"D{g_hrs}", units=f"I{e_u}", sales=f"I{e_s}", per_hour=f"I{e_h}", minutes=f"I{e_m}",
                   multiple=f"I{e_x}", three_x=f"I{e_3}", units_3x=f"I{e_3u}", t_units=f"D{w_u}", t_sales=f"D{w_s}",
                   t_per_hour=f"D{w_h}", bring=f"D{w_b}", wf1=wf1)
BER = "'3 Break-Even'!"

# =====================================================================  4 Show Log
sl = wb.create_sheet("4 Show Log")
S.set_widths(sl, {"A": 3, "B": 2, "C": 13, "D": 30, "E": 11, "F": 12, "G": 8, "H": 9, "I": 11, "J": 11, "K": 10, "L": 11,
                  "M": 11, "N": 11, "O": 11, "P": 2, "Q": 3})
D.page_header(sl, "Show Log", "One row per show. See which shows to apply for again.", "First row is an invented example",
              span=("B", "P"), side_cols=4)
LT = 11
L1 = LT + 5
LN = L1 + N_LOG - 1
lr = {k: f"{k}{L1}:{k}{LN}" for k in "CDEFGHIJKLMNO"}
D.h(sl, 6, 10)
D.stat_strip(sl, 7, [
    ("C:D", "Season gross sales", f"=SUM({lr['I']})", USD2),
    ("E:H", "Season profit", f"=SUM({lr['M']})", USD2),
    ("I:L", "Profit per hour", f'=IF(SUM({lr["G"]})>0,E8/SUM({lr["G"]}),"")', USD2),
    ("M:O", "Shows below 3x", f'=COUNTIF({lr["O"]},"below 3x")', INT),
])
D.h(sl, 9, 10)
D.frame(sl, 6, 9, "B", "P")
D.status_rule(sl, "E8", "E8<0", fill_tint=False)
D.status_rule(sl, "I8", "I8<0", fill_tint=False)
D.status_rule(sl, "M8", "M8>0", fill_tint=False)
D.h(sl, 10, 14)


def lcalc(kind):
    def fn(r):
        return {
            "fees": f'=IF(J{r}="","",J{r}*{R_CP}+IF({R_AV}>0,J{r}/{R_AV},0)*{R_CF})',
            "cogs": f'=IF(H{r}="","",H{r}*{R_BCOST})',
            "profit": f'=IF(I{r}="","",I{r}-N(E{r})-N(F{r})-N(K{r})-N(L{r}))',
            "hour": f'=IF(OR(M{r}="",N(G{r})=0),"",M{r}/G{r})',
            "check": f'=IF(OR(I{r}="",N(E{r})=0),"",IF(I{r}>=3*E{r},"ok","below 3x"))',
        }[kind]
    return fn


EXL = [dt.date(2026, 10, 17), "Riverside Fall Market (example)", 150, 165.6, 7, 60, 790, 560]
cols_sl = [
    dict(col="C", head="Date", kind="input", fmt=DATE, values=[EXL[0]], validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The show day, for example 10/17/2026.", error="Enter a date, for example 10/17/2026."),
    dict(col="D", head="Show", kind="input", values=[EXL[1]], validation=("textLength", 0, 40),
         prompt="The show's name, up to 40 characters.", error="Up to 40 characters."),
    dict(col="E", head="Booth fee", kind="input", fmt=USD2, align="right", values=[EXL[2]], validation=("decimal", 0, None),
         prompt="What the booth cost, in dollars.", error="Enter 0 or more dollars."),
    dict(col="F", head="All other costs", kind="input", fmt=USD2, align="right", values=[EXL[3]], validation=("decimal", 0, None),
         prompt="Everything on tab 1 except the booth fee: extras, travel, parking, food, helper, display share, other.",
         error="Enter 0 or more dollars."),
    dict(col="G", head="Hours", kind="input", fmt=NUM1, align="right", values=[EXL[4]], validation=("decimal", 0, 24),
         prompt="Selling hours that day.", error="Enter hours from 0 to 24."),
    dict(col="H", head="Pieces sold", kind="input", fmt=INT, align="right", values=[EXL[5]], validation=("decimal", 0, None),
         prompt="How many pieces you sold.", error="Enter 0 or more pieces."),
    dict(col="I", head="Gross sales", kind="input", fmt=USD2, align="right", values=[EXL[6]], validation=("decimal", 0, None),
         prompt="All sales before card fees, pre-tax, in dollars.", error="Enter 0 or more dollars."),
    dict(col="J", head="Card sales", kind="input", fmt=USD2, align="right", values=[EXL[7]], validation=("decimal", 0, None),
         prompt="The part of gross sales paid by card, in dollars.", error="Enter 0 or more dollars."),
    dict(col="K", head="Card fees", kind="calc", fmt=USD2, align="right", formula=lcalc("fees")),
    dict(col="L", head="Cost of goods", kind="calc", fmt=USD2, align="right", formula=lcalc("cogs")),
    dict(col="M", head="Profit", kind="calc", fmt=USD2, align="right", formula=lcalc("profit"), bold=True),
    dict(col="N", head="Profit per hour", kind="calc", fmt=USD2, align="right", formula=lcalc("hour")),
    dict(col="O", head="3x check", kind="calc", formula=lcalc("check")),
]
g1, g2, l_end = D.table_card(sl, LT, "B", "P", cols_sl, N_LOG, title="Your shows",
                             sub="Yellow columns are yours. Type each show after it ends; the other columns calculate.")
assert (g1, g2) == (L1, LN), (g1, g2, L1, LN)
sl["M4"].value = f'=IF(COUNTA(D{L1}:D{LN})=0,"","First row is an invented example")'
sl[f"O{L1 - 1}"].alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True, indent=1)
for r in range(L1, LN + 1):
    sl[f"O{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for word, color, tint in (("below 3x", "accent", "accent_tint"), ("ok", "teal", "teal_tint")):
    sl.conditional_formatting.add(f"O{L1}:O{LN}", FormulaRule(formula=[f'$O{L1}="{word}"'],
                                  font=Font(color=D.T[color], bold=True), fill=PatternFill("solid", bgColor=D.T[tint]), stopIfTrue=True))
D.status_rule(sl, f"M{L1}:M{LN}", f"M{L1}<0", fill_tint=False)
ln_top = l_end + 2
SL_NOTES = [
    ("How it is worked out", "Cost of goods uses the blended unit cost from tab 2 (materials, packaging and your making time). "
     "Card fees use the reader fee on tab 1. All other costs is everything on tab 1 except the booth fee."),
    ("The example row", "Invented: 60 pieces and $790 of sales clear the 51-piece break-even and the 3x line, and still net only "
     "a little after the maker's time is paid, which is the point of logging every show."),
]
ln_end = notes_card(sl, ln_top, SL_NOTES, span=("C", "O"), frame_span=("B", "P"))
fill_heights(sl, ln_end)
D.h(sl, ln_end + 1, 14)
D.paint_canvas(sl, ln_end + 1, "Z")
S.finish_sheet(sl, LONG, ln_end + 1, span=("A", "Q"), freeze="A10", tab_color="gold")
sl.page_setup.orientation = "landscape"
sl.print_title_rows = f"{L1 - 1}:{L1 - 1}"
CELLS["sl"] = dict(first=L1, gross="C8", profit="E8", per_hour="I8", below="M8")

# =====================================================================  Start Here and Terms
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    Wd = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            m = re.match(r"^([A-Z][A-Za-z0-9 ,+()-]{1,40})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if m:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), m.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), m.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, Wd + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, Wd, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, PRODUCT, "Know how many pieces you have to sell before the booth fee is paid.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "Type a show's costs and what you sell. The workbook adds up what the day really costs (booth fee, organizer extras, miles, "
    "parking, meals, a share of your tent and display, and your card reader's cut), works out what each piece really earns, and "
    "gives you the number of pieces that pays for the day, the number that makes it worth doing, and how the show compares with "
    "the 3x rule. Decide before the application fee leaves your account.", "para")])
r = sh_card(sh, r, "Three steps", [
    (1, "Show costs. On tab 1, type the booth fee, miles, parking, meals, display spend and your card reader's fee.", "num"),
    (2, "Products. On tab 2, list what you sell: materials, minutes to make, packaging, your price and the mix.", "num"),
    (3, "Break-even. Tab 3 shows the pieces to sell, the 3x check and the pieces for your profit. After the show, log it on tab 4.", "num"),
])
r = sh_card(sh, r, "How to read the cells", [
    ((150, USD0), "Yellow cells with blue numbers are yours to change. Type over the example. Each shows a hint when selected.", "chip-in"),
    ((315.6, USD2), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((51, INT), "Dark tiles are your answers. Each working tab has its answer at the top, and it stays on screen as you scroll.", "chip-key"),
    (None, "Each tab is protected without a password. To change anything else, use Review, Unprotect Sheet in Excel, or Data, "
           "Protect sheets and ranges in Google Sheets.", "note"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "The example", [(None,
    "Riverside Fall Market is an invented show. A $150 booth fee plus $25 of organizer extras, 60 miles at 76 cents, $10 parking, "
    "$20 of meals, $15 of other costs and $50 of display ($300 of setup spread over 6 shows) comes to $315.60, so the booth fee is "
    "less than half of what the day costs. With 70% of sales by card at 2.6% plus 15 cents and a $30 average sale, the card reader "
    "takes 2.17% of every sales dollar. Five invented products blend to $13.84 a piece and $6.24 of contribution after materials, "
    "packaging, making time at $20 an hour and the card fee. That takes 51 pieces, $705.84 of sales, to break even, and 99 pieces "
    "for $300 of profit on top. The 3x rule says $450 is a good day; this show needs $706 just to cover its costs.", "para")])
r = sh_card(sh, r, "Figures used, and where they come from", [
    (None, "IRS mileage rate. 76 cents a mile for business driving from July 1 to December 31, 2026 (IRS Standard Mileage Rates "
           "page, news release IR-2026-29, checked Oct 6, 2026); 72.5 cents from January 1 to June 30, 2026. Used as the default "
           "cost of the drive. A cost estimate, not tax advice.", "para"),
    (None, "Square card fee. 2.6% + 15 cents per tap, dip or swipe on the free plan; manual entry 3.5% + 15 cents "
           "(squareup.com/us/en/fees, read Oct 1, 2026). Change it to match your reader.", "para"),
    (None, "Typical booth fees. Small community fairs $25 to $75; mid-size regional $75 to $200; large juried shows $200 to $500 and "
           "up; multi-day festivals $300 to $1,000 and up; holiday markets $100 to $400 (The Craft Map, Craft Fair Booth Fees "
           "Guide, Feb 13, 2026).", "para"),
    (None, "The 3x rule. Aim to gross at least three times the booth fee (The Craft Map, same guide).", "para"),
    (None, "Pricing formula. (Materials + Labor) x 2 = Retail Price, shown beside your price as a check, not a rule "
           "(CraftProfessional.com, Craft Pricing Formula, updated Apr 11, 2025).", "para"),
    (None, "Sales tax at fairs. Most states require a seller's permit or temporary vendor permit and make the vendor collect and "
           "remit sales tax (Avalara, Craft fairs and sales tax: a state-by-state guide, updated Aug 3, 2026). Prices here are "
           "pre-tax.", "para"),
])
r = sh_card(sh, r, "Good to know", [
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. Every formula is a plain "
           "formula that works in Excel, Google Sheets and LibreOffice. No macros, no add-ons, no sign-up.", "para"),
    (None, "Any currency. The maths works in any currency. Select the money cells and pick your symbol under Format, Number.", "para"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "Before you rely on a number", [(None, NOTICE, "note")])
D.paint_canvas(sh, r - 1, "Z")
S.finish_sheet(sh, LONG, r - 1, span=("A", "G"), tab_color="teal")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More from the shop", [
    (None, ("Hourly Rate Quick Calculator: what an hour of your time has to earn", "https://www.etsy.com/listing/4589695837"), "link"),
    (None, ("Everything in the shop: proofnotfluff.etsy.com", "https://proofnotfluff.etsy.com"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this workbook helped you decide on a show, a short review on Etsy helps other makers find it. Thank you.", "para"),
    (None, "If a file doesn't open or something looks wrong, message me on Etsy and it will be fixed.", "para"),
])
D.paint_canvas(tm, r - 1, "Z")
S.finish_sheet(tm, LONG, r - 1, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb) if "input without validation" not in p]
if problems:
    print("AUDIT:", *problems, sep="\n  ")
wb.save(OUT)
print("saved", OUT)
if CELLS_OUT:
    json.dump(CELLS, open(CELLS_OUT, "w"), indent=1)
print(json.dumps(CELLS))
