#!/usr/bin/env python3
"""#16 Debt Payoff Tracker with a Real Payment Log, version 3 (v3 dashboard design, Oct 7, 2026).

Same inputs, same five-debt example, same payment log and the same maths as version 2 (Oct 5, 2026):
balances today come from each statement balance and date plus the payments logged after it, with
APR / 12 added at each month end (a partial month for the statement month); Snowball (smallest
balance first) and Avalanche (highest APR first) run 240 months, pay every minimum and roll a
paid-off debt's money into the next one; minimums only uses NPER per debt. Rebuilt to
design/WORKBOOK_STANDARD.md section 0; tools/compare_xlsx.py with compare_map.json and
compare_map_avalanche.json proves the answers match.

Changes in version 3:
- The monthly amount moves from the Debts tab to the Dashboard ("Your plan" card), named MonthlyAmount.
- New "Plan you follow" picker (Avalanche or Snowball). The hero tile, the next-month table and
  "Owed over time" follow it; both plans stay side by side in "Snowball vs Avalanche". v2 showed
  both plans in one table; the numbers are the same.
- The line chart is replaced by in-cell bars ("Owed over time"), per the standard.
- The CFPB debt article moved to consumerfinance.gov/archive/blog/how-reduce-your-debt/ (checked
  Oct 7, 2026); the citation follows it.
- The "Leave a Review" tab is folded into the Terms tab. Every tab is protected without a password.
- "Paid off in" on the next-month table reads "Update as-of" (v2: "Done") for a debt whose balance
  can't be worked out, and stays blank for a row with a name but no numbers.
- Deliberate exception to section 0: the month-by-month grids on tabs 4, 5 and 6 (80 columns of
  formulas) sit on the canvas under a carded summary, not inside a card.

Usage: python3 engines/debt-payoff-tracker/build_xlsx.py out.xlsx [--blank]
"""
import datetime as dt
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font  # noqa: E402
from openpyxl.utils import get_column_letter as GL, column_index_from_string as CI  # noqa: E402

PRODUCT = "Debt Payoff Tracker with a Real Payment Log"
SHORT = "Debt Payoff Tracker"
VERSION = "3"
CHECKED = "Oct 7, 2026"
CHECKED_LONG = "October 7, 2026"
AS_OF = f"Sources checked {CHECKED}"
HERE = os.path.dirname(os.path.abspath(__file__))
args = sys.argv[1:]
BLANK = "--blank" in args
args = [a for a in args if a != "--blank"]
OUT = args[0] if args else "Debt-Payoff-Tracker.xlsx"
USD0, USD2, INT, DATE = S.FMT["usd0"], S.FMT["usd2"], S.FMT["int"], S.FMT["date"]
PCT2 = "0.00%"
USD2D = '"$"#,##0.00;-"$"#,##0.00;"-"'
MONTH_FMT = "mmmm yyyy"
MON_FMT = "mmm yyyy"

# ---------------------------------------------------------------- the v2 example, unchanged
d = dt.datetime
DEBTS_EX = [  # name, type, statement balance, as-of date, APR, minimum
    ("Visa card", "Credit card", 4850, d(2026, 8, 1), 0.2499, 145),
    ("Store card", "Store card", 1240, d(2026, 8, 1), 0.2699, 40),
    ("Car loan", "Car loan", 11600, d(2026, 8, 1), 0.074, 329),
    ("Student loan", "Student loan", 18900, d(2026, 8, 1), 0.055, 205),
    ("Personal loan", "Personal loan", 3300, d(2026, 8, 1), 0.129, 112),
]
MONTHLY_EX = 1100
TYPES = ["Credit card", "Store card", "Car loan", "Student loan", "Personal loan", "Mortgage", "Other"]
N_DEBTS = 25
N_LOG = 1000
N_MONTHS = 240
N_HIST = 120


def _load_log_example():
    """The v2 sample log (date, debt, amount), kept in log_example.json beside this script."""
    import json
    rows = json.load(open(os.path.join(HERE, "log_example.json"), encoding="utf-8"))
    return [(dt.datetime.fromisoformat(a), b, c, n) for a, b, c, n in rows]


LOG_EX = _load_log_example()
if BLANK:
    DEBTS_EX, LOG_EX, MONTHLY_EX = [], [], None

# ---------------------------------------------------------------- sheet names and fixed addresses
DASH, DEB, LOG, SNO, AVA, HIS = ("1 Dashboard", "2 Debts", "3 Payment Log", "4 Snowball Plan",
                                 "5 Avalanche Plan", "6 Balance History")


def q(name):
    return f"'{name}'"


# 2 Debts: table rows and columns
D_FIRST = 17
D_LAST = D_FIRST + N_DEBTS - 1
DROWS = list(range(D_FIRST, D_LAST + 1))
DC = dict(num="C", name="D", type="E", bal="F", asof="G", apr="H", min="I", paid="J", intr="K", today="L",
          status="M", snow="N", aval="O", nmon="P", nint="Q", skey="R", akey="S")
TD = f"{q(DEB)}!$Q$8"            # today's date (stat strip)
MINSUM = f"{q(DEB)}!$F$8"        # sum of minimums on debts with a balance
OWED = f"{q(DEB)}!$C$8"          # owed today, all debts
PAIDSUM = f"{q(DEB)}!$I$8"       # paid since statements
INTSUM = f"{q(DEB)}!$L$8"        # interest since statements
PAIDOFF = f"{q(DEB)}!$N$8"       # "2 of 5"

# 3 Payment Log
L_FIRST = 17
L_LAST = L_FIRST + N_LOG - 1
LG = dict(date="C", debt="D", amt="E", note="F", ok="G")
LOGCOUNTED = f"{q(LOG)}!$C$8"
LOGFIX = f"{q(LOG)}!$F$8"
LOGROWS = f"{q(LOG)}!$E$8"

# 4 and 5 plans: per-debt blocks
P_NAME, P_PAIDOFF, P_RANK, P_APR, P_MIN, P_START, P_HEAD = 13, 14, 15, 16, 17, 18, 20
P_FIRST = P_HEAD + 1
P_LAST = P_FIRST + N_MONTHS - 1
DUE0, ROOM0, LEFT0 = CI("J"), CI("J") + N_DEBTS, CI("J") + 2 * N_DEBTS


def due(i):
    return GL(DUE0 + i)


def room(i):
    return GL(ROOM0 + i)


def left(i):
    return GL(LEFT0 + i)


P_MONTHS = "$C$8"   # debt-free in month (stat strip)
P_INT = "$F$8"      # total interest
P_DATE = "$H$8"     # debt-free month as a date

# 6 Balance History
H_START = f"{q(HIS)}!$C$8"
H_HEAD = 12
H_FIRST = H_HEAD + 1
H_LAST = H_FIRST + N_HIST - 1
BAL0, INT0 = CI("D"), CI("D") + N_DEBTS


def hbal(i):
    return GL(BAL0 + i)


def hint(i):
    return GL(INT0 + i)


MONTH1 = "DATE(YEAR(TODAY()),MONTH(TODAY()),1)"

wb = S.new_workbook(PRODUCT, version=VERSION)


def header_cells(ws, row, cells, height=30):
    for col, text, align in cells:
        c = ws[f"{col}{row}"]; c.value = text.upper(); c.font = D.f(8, True, "muted")
        c.alignment = Alignment(horizontal=align, vertical="bottom", wrap_text=True)
        c.border = Border(bottom=D.side("ink"))
    D.h(ws, row, height)


def body(c, fmt=None, bold=False, color="ink", align="right", size=10):
    if fmt:
        c.number_format = fmt
    c.font = D.f(size, bold, color)
    c.alignment = Alignment(horizontal=align, vertical="center")


def strip_card(ws, top, stats, sub, span):
    """Rows top..top+4: pad, stat eyebrows, stat numbers, a sub line, pad. Returns the row after."""
    D.h(ws, top, 10)
    D.stat_strip(ws, top + 1, stats)
    s = ws[f"{span[0]}{top + 3}"]; s.value = sub; s.font = D.f(9, False, "muted"); s.alignment = Alignment(vertical="center")
    ws.merge_cells(f"{span[0]}{top + 3}:{span[1]}{top + 3}"); D.h(ws, top + 3, 20)
    D.h(ws, top + 4, 10)
    return top + 5


def notes_card(ws, top, notes, span, frame_span):
    r = top
    D.h(ws, r, 12); r += 1
    t = ws[f"{span[0]}{r}"]; t.value = "Notes and sources"; t.font = D.f(12, True)
    ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}"); D.h(ws, r, 24); r += 1
    D.h(ws, r, 6); r += 1
    width = sum(S.width_of(ws, k) for k in D.cols(*span))
    for head, text in notes:
        cell = ws[f"{span[0]}{r}"]
        cell.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=9, b=True, color=D.T["ink"]), head + ". "),
                                  TextBlock(InlineFont(rFont=D.FONT, sz=9, color=D.T["ink2"]), text))
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}")
        D.h(ws, r, S.text_height(head + ". " + text, width, 9) + 4); r += 1
    D.h(ws, r, 12)
    D.frame(ws, top, r, *frame_span)
    return r


CFPB_DEBT = ("CFPB, \"How to reduce your debt\" (published July 16, 2019, last modified June 25, 2026), "
             "consumerfinance.gov/archive/blog/how-reduce-your-debt/, checked " + CHECKED)
CFPB_APR = ("CFPB, \"What is a credit card interest rate? What does APR mean?\" (reviewed August 28, 2023), "
            "consumerfinance.gov/ask-cfpb/what-is-a-credit-card-interest-rate-what-does-apr-mean-en-44/, checked " + CHECKED)

# =====================================================================  1 Dashboard (built first; later tabs point at it)
ds = wb.create_sheet(DASH)
S.set_widths(ds, {"A": 3, "B": 2, "C": 31, "D": 14, "E": 2, "F": 3, "G": 2, "H": 24, "I": 12, "J": 17, "K": 2, "L": 3})
D.page_header(ds, "Debt Payoff Dashboard", "Your debt-free month, what you owe today and what to pay each debt next month.",
              AS_OF, span=("B", "K"), side_cols=3)
TOP = 11
A = D.Card(ds, TOP, "B", "C", "D", "E", wb=wb)
A.title("Your plan", "Type over the yellow cells. Your debts go on tab 2.")
A.eyebrow("EACH MONTH")
MA = A.input("Monthly amount for all debts", MONTHLY_EX, USD0, name="MonthlyAmount", validation=("decimal", 0, None),
             allow_blank=True, prompt_title="Monthly amount for all debts",
             prompt="The total you can put toward all debts each month, for example 1100. It must cover every minimum.")
MONTHLY = f"{q(DASH)}!$D${MA}"
a_min = A.calc("Sum of your minimums", f"={MINSUM}", USD2)
a_pill = A.text(f'=IF(D{MA}="","Type your monthly amount above",IF(NOT(ISNUMBER(D{MA})),"Type a number (no words)",'
                f'IF(D{MA}<D{a_min},"Below your minimums: raise it to at least "&TEXT(D{a_min},"$#,##0.00"),'
                f'"Extra each month: "&TEXT(D{MA}-D{a_min},"$#,##0.00"))))', size=10, color="ink", bold=True)
ds[f"C{a_pill}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
A.eyebrow("PLAN YOU FOLLOW")
PL = A.input("Snowball or Avalanche", "Avalanche", "@", name="PlanChoice", validation=("list", ["Avalanche", "Snowball"]),
             prompt_title="Plan you follow",
             prompt="Avalanche pays the highest APR first; Snowball pays the smallest balance first. The tiles and the next-month table follow this.")
A.close()
D.status_rule(ds, f"C{a_pill}:D{a_pill}", f'LEFT($C${a_pill},5)<>"Extra"')
PLAN = f"$D${PL}"
IS_SNOW = f'{PLAN}="Snowball"'
PLAN_NAME = f'IF({IS_SNOW},"Snowball","Avalanche")'

C2 = D.Card(ds, A.end + 2, "B", "C", "D", "E", wb=wb)
C2.title("Progress so far", "From your statements and the payments you logged.")
c_owed = C2.calc("Owed today (estimate)", f"={OWED}", USD2, bold=True)
c_paid = C2.calc("Paid since statements", f"={LOGCOUNTED}", USD2)
c_int = C2.calc("Interest added (est.)", f"={INTSUM}", USD2)
c_off = C2.calc("Debts paid off", f"={PAIDOFF}", "@")
c_fix = C2.calc("Log rows to fix", f"={LOGFIX}", INT)
C2.close()
D.status_rule(ds, f"D{c_fix}", f"$D${c_fix}>0", fill_tint=False)


def plan_date(sheet):
    return f'IF(ISNUMBER({q(sheet)}!{P_MONTHS}),EDATE({MONTH1},{q(sheet)}!{P_MONTHS}),{q(sheet)}!{P_MONTHS})'


B = D.Card(ds, TOP, "G", "H", "I", "K", extra="J", wb=wb)
B.title("Snowball vs Avalanche", "Both pay every minimum; a paid-off debt's money rolls on.")
INT_ROWS = []
PLAN_ROWS = {}
for sheet, eyebrow in ((SNO, "SNOWBALL: SMALLEST BALANCE FIRST"), (AVA, "AVALANCHE: HIGHEST APR FIRST")):
    B.eyebrow(eyebrow)
    r_dt = B.calc("Debt-free month", f"={plan_date(sheet)}", MON_FMT, bold=True)
    ds[f"J{r_dt}"] = f'=IF(ISNUMBER({q(sheet)}!{P_MONTHS}),"in "&{q(sheet)}!{P_MONTHS}&" months","")'
    body(ds[f"J{r_dt}"], color="ink2", align="left", size=9); ds[f"J{r_dt}"].alignment = Alignment(vertical="center", indent=1)
    r_in = B.calc("Interest still to pay (est.)", f"={q(sheet)}!{P_INT}", USD2)
    INT_ROWS.append(r_in)
    r_sv = B.calc("Saved vs minimums only", "", USD2)
    PLAN_ROWS[sheet] = (r_dt, r_in, r_sv)
B.eyebrow("MINIMUMS ONLY: NO EXTRA, NO ROLL-OVER")
MIN_NEVER = f'COUNTIF({q(DEB)}!${DC["nmon"]}${D_FIRST}:${DC["nmon"]}${D_LAST},"Never")>0'
MIN_COUNT = f'COUNT({q(DEB)}!${DC["nmon"]}${D_FIRST}:${DC["nmon"]}${D_LAST})'
MIN_MAX = f'MAX({q(DEB)}!${DC["nmon"]}${D_FIRST}:${DC["nmon"]}${D_LAST})'
INT_NEVER = f'COUNTIF({q(DEB)}!${DC["nint"]}${D_FIRST}:${DC["nint"]}${D_LAST},"Never")>0'
INT_COUNT = f'COUNT({q(DEB)}!${DC["nint"]}${D_FIRST}:${DC["nint"]}${D_LAST})'
INT_SUM = f'SUM({q(DEB)}!${DC["nint"]}${D_FIRST}:${DC["nint"]}${D_LAST})'
r_mdt = B.calc("Debt-free month", f'=IF({MIN_NEVER},"Never",IF({MIN_COUNT}=0,"",EDATE({MONTH1},{MIN_MAX})))', MON_FMT, bold=True)
ds[f"J{r_mdt}"] = f'=IF({MIN_NEVER},"",IF({MIN_COUNT}=0,"","in "&{MIN_MAX}&" months"))'
body(ds[f"J{r_mdt}"], color="ink2", align="left", size=9); ds[f"J{r_mdt}"].alignment = Alignment(vertical="center", indent=1)
r_min = B.calc("Interest still to pay (est.)", f'=IF({INT_NEVER},"Never",IF({INT_COUNT}=0,"",{INT_SUM}))', USD2)
INT_ROWS.append(r_min)
for sheet, (r_dt, r_in, r_sv) in PLAN_ROWS.items():
    ds[f"I{r_sv}"] = f'=IF(AND(ISNUMBER(I{r_in}),ISNUMBER(I{r_min})),I{r_min}-I{r_in},"")'
IMAX = "MAX(" + ",".join(f"$I${x}" for x in INT_ROWS) + ")"
for rr, color in zip(INT_ROWS, ("blue", "teal", "accent")):
    b = ds[f"J{rr}"]
    b.value = f'=IF(ISNUMBER(I{rr}),IFERROR(REPT("{D.BAR_CHAR}",MAX(0,ROUND(I{rr}/{IMAX}*11,0))),""),"")'
    b.font = D.f(9, False, color); b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
r_sa, r_si = PLAN_ROWS[SNO][1], PLAN_ROWS[AVA][1]
B.text(f'=IF(AND(ISNUMBER(I{r_sa}),ISNUMBER(I{r_si})),IF(ABS(I{r_sa}-I{r_si})<0.005,"Both plans cost the same interest.",'
       f'IF(I{r_si}<I{r_sa},"Avalanche","Snowball")&" saves "&TEXT(ABS(I{r_sa}-I{r_si}),"$#,##0")&" of interest over "'
       f'&IF(I{r_si}<I{r_sa},"Snowball","Avalanche")&"."),"")', size=9, color="ink2")
B.close()

# owed over time (right column, under the comparison)
O = D.Card(ds, B.end + 2, "G", "H", "I", "K", extra="J", wb=wb)
O.title("Owed over time", "Your plan's balance after payments, one row a year.")
o_first = O.r
O_ROWS = []
for k in range(6):
    m = 12 * k
    cell_s = f"{q(SNO)}!$G${P_START}" if m == 0 else f"{q(SNO)}!$G${P_HEAD + m}"
    cell_a = f"{q(AVA)}!$G${P_START}" if m == 0 else f"{q(AVA)}!$G${P_HEAD + m}"
    rr = O.calc("", f"=IF({IS_SNOW},{cell_s},{cell_a})", USD0)
    ds[f"H{rr}"] = f'="Today"' if m == 0 else f'=TEXT(EDATE({MONTH1},{m}),"mmm yyyy")'
    O_ROWS.append(rr)
OMAX = f"MAX($I${O_ROWS[0]}:$I${O_ROWS[-1]})"
for rr in O_ROWS:
    b = ds[f"J{rr}"]
    b.value = f'=IFERROR(REPT("{D.BAR_CHAR}",MAX(0,ROUND(I{rr}/{OMAX}*11,0))),"")'
    b.font = D.f(9, False, "teal"); b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
O.close()

# left column filler so both columns end together
L_END, R_END = C2.end, O.end

# ---- full width: what to pay next month
NM_TOP = max(L_END, R_END) + 2


def nm_pay(r):
    i = r - NM_FIRST
    dr = DROWS[i]
    return (f'=IF({q(DEB)}!${DC["name"]}${dr}="","",IF({IS_SNOW},{q(SNO)}!{due(i)}{P_FIRST}-{q(SNO)}!{left(i)}{P_FIRST},'
            f'{q(AVA)}!{due(i)}{P_FIRST}-{q(AVA)}!{left(i)}{P_FIRST}))')


def nm_off(r):
    i = r - NM_FIRST
    dr = DROWS[i]
    po = f'IF({IS_SNOW},{q(SNO)}!{due(i)}{P_PAIDOFF},{q(AVA)}!{due(i)}{P_PAIDOFF})'
    lt = f'{q(DEB)}!${DC["today"]}${dr}'
    return (f'=IF({q(DEB)}!${DC["name"]}${dr}="","",IF(NOT(ISNUMBER({lt})),IF({lt}="","","Update as-of"),'
            f'IF({po}="","Done",IF(ISNUMBER({po}),EDATE({MONTH1},{po}),{po}))))')


NM_FIRST = NM_TOP + 5
nm_cols = [
    dict(col="C", head="Debt", kind="calc", formula=lambda r: f'=IF({q(DEB)}!${DC["name"]}${DROWS[r - NM_FIRST]}="","",{q(DEB)}!${DC["name"]}${DROWS[r - NM_FIRST]})', align="left"),
    dict(col="D", head="Balance today", kind="calc", fmt=USD2, align="right",
         formula=lambda r: f'=IF({q(DEB)}!${DC["name"]}${DROWS[r - NM_FIRST]}="","",{q(DEB)}!${DC["today"]}${DROWS[r - NM_FIRST]})'),
    dict(col="H", head="Pay next month", kind="calc", fmt=USD2, align="right", formula=nm_pay),
    dict(col="I", head="Paid off in", kind="calc", fmt=MON_FMT, align="right", formula=nm_off),
    dict(col="J", head="Pay order", kind="calc", fmt=INT, align="center",
         formula=lambda r: f'=IF({q(DEB)}!${DC["name"]}${DROWS[r - NM_FIRST]}="","",IF({IS_SNOW},{q(DEB)}!${DC["snow"]}${DROWS[r - NM_FIRST]},{q(DEB)}!${DC["aval"]}${DROWS[r - NM_FIRST]}))'),
]
f1, f2, nm_end = D.table_card(ds, NM_TOP, "B", "K", nm_cols, N_DEBTS, title="What to pay next month",
                              sub="Follows the plan you picked. Every minimum is paid; the extra goes to the debt marked 1.", wb=wb)
assert f1 == NM_FIRST, (f1, NM_FIRST)
for rr in range(f1, f2 + 1):
    for k in ("E", "F", "G"):
        ds[f"{k}{rr}"].border = Border(bottom=D.side("hair"))
ds[f"C{NM_TOP + 1}"] = f'="What to pay in "&TEXT(EDATE({MONTH1},1),"mmmm yyyy")&" ("&{PLAN_NAME}&")"'
# a totals row in place of the bottom pad
NM_TOT = nm_end
ds[f"C{NM_TOT}"] = "Total"
ds[f"D{NM_TOT}"] = f"=SUM(D{f1}:D{f2})"
ds[f"H{NM_TOT}"] = f"=SUM(H{f1}:H{f2})"
for k in ("E", "F", "G"):
    ds[f"{k}{NM_TOT}"].border = Border(top=D.side("ink"))
for k in ("C", "D", "H", "I", "J"):
    c = ds[f"{k}{NM_TOT}"]
    body(c, USD2 if k in ("D", "H") else None, bold=True, align="left" if k == "C" else "right")
    c.border = Border(top=D.side("ink"))
ds[f"B{NM_TOT}"].border = Border(); ds[f"K{NM_TOT}"].border = Border()
D.h(ds, NM_TOT, D.GRID_ROW)
D.h(ds, NM_TOT + 1, 12)
D.frame(ds, NM_TOP, NM_TOT + 1, "B", "K")
ds.conditional_formatting.add(f"C{f1}:J{f2}", FormulaRule(formula=[f'$J{f1}=1'], fill=D.fill("teal_tint"), stopIfTrue=True))
S.page_break_before(ds, NM_TOP)

# ---- notes
NOTES = [
    ("Snowball and Avalanche", "Snowball pays the smallest balance first; Avalanche (the highest interest rate method) pays the "
     "highest APR first. Both keep paying every minimum and move a paid-off debt's money to the next one. Source: " + CFPB_DEBT + "."),
    ("How interest is estimated", "APR is a yearly rate (" + CFPB_APR + "). The tracker adds APR / 12 of the balance at each month "
     "end, after that month's payments, and a partial month for the month of your statement date. Lenders work out interest daily, "
     "so a statement can differ by a few dollars: type the new statement balance and date on tab 2 and the tracker resets to it."),
    ("Plans", "Plans start from today's balances, begin next month, assume no new charges and run up to 240 months. Minimums only "
     "pays each debt its own minimum with no extra and no roll-over. This month's payments count once you log them."),
    ("Before you rely on a number", S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab")
     + " Talk to your lender before you change a payment."),
]
N_TOP = NM_TOT + 3
n_end = notes_card(ds, N_TOP, NOTES, ("C", "J"), ("B", "K"))
S.page_break_before(ds, N_TOP)
DS_LAST = n_end + 1
D.h(ds, DS_LAST, 14)

# ---- hero tiles
STALE = f'COUNTIF({q(DEB)}!${DC["status"]}${D_FIRST}:${DC["status"]}${D_LAST},"Update as-of date")'
NAMED = f'COUNTA({q(DEB)}!${DC["name"]}${D_FIRST}:${DC["name"]}${D_LAST})'
P_MON_CH = f'IF({IS_SNOW},{q(SNO)}!{P_MONTHS},{q(AVA)}!{P_MONTHS})'
P_INT_CH = f'IF({IS_SNOW},{q(SNO)}!{P_INT},{q(AVA)}!{P_INT})'
D.tile(ds, 6, "B", "E", f'="DEBT-FREE MONTH WITH "&UPPER({PLAN_NAME})',
       f'=IF({P_MON_CH}="",IF({STALE}>0,"Update a statement",IF({NAMED}>0,"Nothing owed","Add your debts")),'
       f'IF(ISNUMBER({P_MON_CH}),EDATE({MONTH1},{P_MON_CH}),"Over 20 years"))', MONTH_FMT,
       f'=IF(ISNUMBER({P_MON_CH}),{P_MON_CH}&" months from now. Interest still to pay: "&TEXT({P_INT_CH},"$#,##0")&".",'
       f'IF({P_MON_CH}="",IF({STALE}>0,"A statement date is over ten years back: see Status on tab 2.",'
       f'IF({NAMED}>0,"Every debt on tab 2 shows paid off.","Type each debt on tab 2 and your monthly amount below.")),'
       f'"Raise the monthly amount: it does not clear the debts in 20 years."))',
       dark=True, value_size=26)
v, sub, _ = D.tile(ds, 6, "G", "K", "OWED TODAY (ESTIMATE)", f"={OWED}", USD2,
                   f'=IF(COUNT({q(LOG)}!${LG["date"]}${L_FIRST}:${LG["date"]}${L_LAST})=0,"No payments logged yet",'
                   f'IF({LOGFIX}>0,{LOGFIX}&IF({LOGFIX}=1," log row"," log rows")&" to fix on tab 3",'
                   f'"Paid "&TEXT({LOGCOUNTED},"$#,##0")&" since your statements. "&{PAIDOFF}&" debts paid off."))')
ds.conditional_formatting.add(sub.coordinate, FormulaRule(formula=[f"{LOGFIX}>0"], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
D.h(ds, 10, 14)
for rr in range(1, DS_LAST):
    if ds.row_dimensions[rr].height is None:
        D.h(ds, rr, D.GRID_ROW)
D.paint_canvas(ds, DS_LAST, "Z")
S.finish_sheet(ds, PRODUCT, DS_LAST, span=("A", "L"), freeze="A11", tab_color="accent")

# =====================================================================  2 Debts
db = wb.create_sheet(DEB)
S.set_widths(db, {"A": 3, "B": 2, "C": 5, "D": 20, "E": 14, "F": 13, "G": 13, "H": 9, "I": 11, "J": 12, "K": 12, "L": 13,
                  "M": 18, "N": 10, "O": 10, "P": 11, "Q": 12, "R": 10, "S": 10, "T": 2, "U": 3})
D.page_header(db, "Debts", "One row per debt, from its latest statement. New statement? Type its balance and date over the old ones.",
              AS_OF, span=("B", "T"), side_cols=4)
tr = f"${DC['today']}${D_FIRST}:${DC['today']}${D_LAST}"
strip_card(db, 6, [
    ("C:E", "Owed today (estimate)", f"=SUM({tr})", USD2),
    ("F:H", "Sum of minimums", f'=SUMIFS(${DC["min"]}${D_FIRST}:${DC["min"]}${D_LAST},{tr},">0.005")', USD2),
    ("I:K", "Paid since statements", f"=SUM(${DC['paid']}${D_FIRST}:${DC['paid']}${D_LAST})", USD2),
    ("L:M", "Interest since (est.)", f"=SUM(${DC['intr']}${D_FIRST}:${DC['intr']}${D_LAST})", USD2),
    ("N:P", "Debts paid off", f'=COUNTIF(${DC["status"]}${D_FIRST}:${DC["status"]}${D_LAST},"Paid off")&" of "&COUNTA(${DC["name"]}${D_FIRST}:${DC["name"]}${D_LAST})', "@"),
    ("Q:S", "Today", "=TODAY()", DATE),
], "Minimums count only for debts that still have a balance. Balances today come from your statements plus the payments on tab 3.",
    ("C", "S"))
D.frame(db, 6, 10, "B", "T")
assert db["Q8"].value == "=TODAY()" and db["C8"].value.startswith("=SUM(") and db["F8"].value.startswith("=SUMIFS(")
assert db["I8"].value.startswith("=SUM(") and db["L8"].value.startswith("=SUM(") and db["N8"].value.startswith("=COUNTIF(")
VALID = lambda r: f'AND({DC["name"]}{r}<>"",ISNUMBER({DC["bal"]}{r}),ISNUMBER({DC["asof"]}{r}))'  # noqa: E731


def f_paid(r):
    return (f'=IF({VALID(r)},SUMIFS({q(LOG)}!${LG["amt"]}:${LG["amt"]},{q(LOG)}!${LG["debt"]}:${LG["debt"]},{DC["name"]}{r},'
            f'{q(LOG)}!${LG["date"]}:${LG["date"]},">"&{DC["asof"]}{r},{q(LOG)}!${LG["date"]}:${LG["date"]},"<="&{TD}),"")')


def f_int(r):
    i = r - D_FIRST
    return (f'=IF({VALID(r)},SUMIFS({q(HIS)}!${hint(i)}${H_FIRST}:${hint(i)}${H_LAST},'
            f'{q(HIS)}!$C${H_FIRST}:$C${H_LAST},"<="&{TD}),"")')


def f_today(r):
    i = r - D_FIRST
    return (f'=IF({VALID(r)},IF({DC["asof"]}{r}>{TD},{DC["bal"]}{r},IFERROR(INDEX({q(HIS)}!${hbal(i)}${H_FIRST}:${hbal(i)}${H_LAST},'
            f'MATCH(DATE(YEAR({TD}),MONTH({TD}),1),{q(HIS)}!$C${H_FIRST}:$C${H_LAST},0)),"Update as-of")),"")')


def f_status(r):
    j = f"{DC['today']}{r}"
    return (f'=IF({j}="","",IF(NOT(ISNUMBER({j})),"Update as-of date",IF({j}<0.005,"Paid off",'
            f'IF({DC["asof"]}{r}>{TD},"As-of is in the future","Active"))))')


def f_order(key):
    def f(r):
        j = f"{DC['today']}{r}"
        return f'=IF(AND(ISNUMBER({j}),{j}>0.005),COUNTIF(${DC[key]}${D_FIRST}:${DC[key]}${D_LAST},"<"&{DC[key]}{r})+1,"")'
    return f


def f_nmon(r):
    j, g, fr = f"{DC['today']}{r}", f"{DC['min']}{r}", f"{DC['apr']}{r}"
    return (f'=IF(AND(ISNUMBER({j}),{j}>0.005,N({g})>0),IF(N({g})<={j}*N({fr})/12,"Never",'
            f'ROUNDUP(IF(N({fr})=0,{j}/{g},NPER(N({fr})/12,-{g},{j})),0)),"")')


def f_nint(r):
    j, g, fr, n = f"{DC['today']}{r}", f"{DC['min']}{r}", f"{DC['apr']}{r}", f"{DC['nmon']}{r}"
    return (f'=IF(ISNUMBER({n}),IF(N({fr})=0,0,MAX(0,{g}*NPER(N({fr})/12,-{g},{j})-{j})),IF({n}="Never","Never",""))')


def f_skey(r):
    j, a = f"{DC['today']}{r}", f"{DC['num']}{r}"
    return f'=IF(AND(ISNUMBER({j}),{j}>0.005),{j}+{a}/1000000,1000000000000+{a})'


def f_akey(r):
    j, a, fr = f"{DC['today']}{r}", f"{DC['num']}{r}", f"{DC['apr']}{r}"
    return f'=IF(AND(ISNUMBER({j}),{j}>0.005),-N({fr})+{a}/1000000000,1000000000000+{a})'


ex = DEBTS_EX
db_cols = [
    dict(col="C", head="#", kind="muted", values=list(range(1, N_DEBTS + 1)), fmt="0", align="center"),
    dict(col="D", head="Debt name", kind="input", values=[x[0] for x in ex], validation=("textLength", 0, 40),
         prompt="A short name you will pick in the Payment Log, for example Visa card. Each name once.",
         error="Up to 40 characters. Use each name once."),
    dict(col="E", head="Type", kind="input", values=[x[1] for x in ex], validation=("list", TYPES),
         prompt="Pick the kind of debt. It is for your reference; the maths does not use it.", error="Pick a type from the list."),
    dict(col="F", head="Statement balance", kind="input", fmt=USD2, align="right", values=[x[2] for x in ex],
         validation=("decimal", 0, None), prompt="The balance on your latest statement, for example 4850.",
         error="Type a balance of 0 or more, in dollars."),
    dict(col="G", head="Statement date (as-of)", kind="input", fmt=DATE, align="right", values=[x[3] for x in ex],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The date that balance is from, for example 8/1/2026. Payments after this date count against it.",
         error="Type a real date from 2000 to 2100, for example 8/1/2026."),
    dict(col="H", head="APR", kind="input", fmt=PCT2, align="right", values=[x[4] for x in ex],
         validation=("decimal", 0, 1), prompt="The yearly rate from the statement, typed with the % sign, for example 24.99%.",
         error="Type a rate from 0% to 100%, with the % sign."),
    dict(col="I", head="Minimum payment", kind="input", fmt=USD2, align="right", values=[x[5] for x in ex],
         validation=("decimal", 0, None), prompt="The minimum due each month, for example 145.",
         error="Type a minimum of 0 or more, in dollars."),
    dict(col="J", head="Paid since", kind="calc", fmt=USD2, align="right", formula=f_paid),
    dict(col="K", head="Interest", kind="calc", fmt=USD2, align="right", formula=f_int),
    dict(col="L", head="Balance today", kind="calc", fmt=USD2, align="right", formula=f_today, bold=True),
    dict(col="M", head="Status", kind="calc", align="left", formula=f_status),
    dict(col="N", head="Snowball order", kind="calc", fmt=INT, align="center", formula=f_order("skey")),
    dict(col="O", head="Avalanche order", kind="calc", fmt=INT, align="center", formula=f_order("akey")),
    dict(col="P", head="Months on minimums", kind="calc", fmt=INT, align="right", formula=f_nmon),
    dict(col="Q", head="Interest on minimums", kind="calc", fmt=USD2, align="right", formula=f_nint),
    dict(col="R", head="Sort key: Snowball", kind="muted", fmt="0.000", align="right", formula=f_skey),
    dict(col="S", head="Sort key: Avalanche", kind="muted", fmt="0.000000", align="right", formula=f_akey),
]
for c in db_cols:
    if c["col"] in ("R", "S"):
        c["kind"] = "calc"
g1, g2, db_end = D.table_card(db, 12, "B", "T", db_cols, N_DEBTS, title="Your debts",
                              sub="Yellow columns are yours (up to 25 debts). Everything to the right calculates. The two sort keys order the plans; leave them.")
assert (g1, g2) == (D_FIRST, D_LAST), (g1, g2)
for rr in DROWS:
    for k in ("R", "S"):
        db[f"{k}{rr}"].font = D.f(9, False, "muted")
    db[f"M{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=2)
    db[f"D{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    db[f"E{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
db[f"M{D_FIRST - 1}"].alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True, indent=2)
db.conditional_formatting.add(f"R{D_FIRST}:S{D_LAST}", FormulaRule(formula=[f'R{D_FIRST}>=1000000000000'],
                              font=Font(color="FFFFFF"), stopIfTrue=True))
db.conditional_formatting.add(f"M{D_FIRST}:M{D_LAST}", FormulaRule(formula=[f'M{D_FIRST}="Paid off"'],
                              font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))
db.conditional_formatting.add(f"M{D_FIRST}:M{D_LAST}", FormulaRule(
    formula=[f'OR(M{D_FIRST}="Update as-of date",M{D_FIRST}="As-of is in the future")'], font=Font(color=D.T["accent"], bold=True),
    stopIfTrue=True))
DB_NOTES = [
    ("New statement", "Type the new balance and its date over the old ones. Payments dated after that date keep counting; older "
     "ones stop, so nothing is counted twice. Tab 3 shows which rows count."),
    ("Ten-year window", "Balance History tracks 120 months from your earliest statement date. Update every debt's statement "
     "balance and date, or clear paid-off rows, at least once in that window. The window starts at the oldest date, so one date "
     "more than ten years back makes every debt read Update as-of date until you update it."),
    ("Orders", "Snowball order ranks today's balances, smallest first; Avalanche order ranks APR, highest first. A tie goes to the "
     "higher row. Paid-off debts drop out of both."),
    ("Interest", "APR / 12 at each month end, with a partial month for the statement month. " + CFPB_APR + "."),
]
db_n = notes_card(db, db_end + 2, DB_NOTES, ("C", "S"), ("B", "T"))
DB_LAST = db_n + 1
D.h(db, DB_LAST, 14)
D.paint_canvas(db, DB_LAST, "Z")
S.finish_sheet(db, PRODUCT, DB_LAST, span=("A", "U"), freeze=f"E{D_FIRST}", tab_color="gold")
db.page_setup.orientation = "landscape"
S.page_break_before(db, db_end + 2)

# =====================================================================  3 Payment Log
lg = wb.create_sheet(LOG)
S.set_widths(lg, {"A": 3, "B": 2, "C": 14, "D": 22, "E": 14, "F": 30, "G": 46, "H": 2, "I": 3})
D.page_header(lg, "Payment Log", "One row per payment you actually made: date, debt, amount. That is all you type from now on.",
              "Example rows: type over them" if not BLANK else "Type your first payment in row 17", span=("B", "H"), side_cols=2)
strip_card(lg, 6, [
    ("C:D", "Counted payments", f'=SUMIFS(${LG["amt"]}:${LG["amt"]},${LG["ok"]}:${LG["ok"]},"Yes")', USD2),
    ("E:E", "Rows counted", f'=COUNTIF(${LG["ok"]}:${LG["ok"]},"Yes")', INT),
    ("F:F", "Rows needing a fix", f'=COUNTIF(${LG["ok"]}:${LG["ok"]},"Fill in*")+COUNTIF(${LG["ok"]}:${LG["ok"]},"Debt name not*")'
                                  f'+COUNTIF(${LG["ok"]}:${LG["ok"]},"Add an as-of*")+COUNTIF(${LG["ok"]}:${LG["ok"]},"Amount must*")', INT),
], "Totals read the whole columns, so rows you copy below the last one count too.", ("C", "G"))
D.frame(lg, 6, 10, "B", "H")
D.status_rule(lg, "F8", "$F$8>0", fill_tint=False)
assert lg["C8"].value.startswith("=SUMIFS(") and lg["F8"].value.startswith("=COUNTIF(")
DN = f"{q(DEB)}!${DC['name']}${D_FIRST}:${DC['name']}${D_LAST}"
DA = f"{q(DEB)}!${DC['asof']}${D_FIRST}:${DC['asof']}${D_LAST}"


def f_ok(r):
    a, b, c = f"{LG['date']}{r}", f"{LG['debt']}{r}", f"{LG['amt']}{r}"
    return (f'=IF(AND({a}="",{b}="",{c}=""),"",IF(OR({a}="",{b}="",{c}=""),"Fill in date, debt and amount",'
            f'IF(ISNA(MATCH({b},{DN},0)),"Debt name not on Debts tab",'
            f'IF(NOT(ISNUMBER(INDEX({DA},MATCH({b},{DN},0)))),"Add an as-of date for this debt on the Debts tab",'
            f'IF({a}>TODAY(),"Future date: counts on that day",'
            f'IF({a}<=INDEX({DA},MATCH({b},{DN},0)),"On or before the as-of date: already in the balance",'
            f'IF(NOT(ISNUMBER({c})),"Amount must be a number","Yes")))))))')


lg_cols = [
    dict(col="C", head="Payment date", kind="input", fmt=DATE, align="right", values=[x[0] for x in LOG_EX],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The day you paid, for example 10/3/2026.", error="Type a real date from 2000 to 2100, for example 10/3/2026."),
    dict(col="D", head="Debt", kind="input", values=[x[1] for x in LOG_EX], validation=("list_range", DN),
         prompt="Pick the debt from the list. The names come from tab 2.",
         error="Pick a debt from the list. To add one, type it on the Debts tab first."),
    dict(col="E", head="Amount paid", kind="input", fmt=USD2, align="right", values=[x[2] for x in LOG_EX],
         validation=("decimal", 0, None), prompt="What you paid, for example 329.", error="Type an amount of 0 or more, in dollars."),
    dict(col="F", head="Note (optional)", kind="input", values=[x[3] for x in LOG_EX], validation=("textLength", 0, 120),
         prompt="Anything worth remembering, for example extra payment. Optional.", error="Up to 120 characters."),
    dict(col="G", head="Counted?", kind="calc", align="left", formula=f_ok),
]
h1, h2, lg_end = D.table_card(lg, 12, "B", "H", lg_cols, N_LOG, title="Your payments",
                              sub=f"Yellow columns are yours. Counted? says Yes when a row counts. {N_LOG:,} rows ready; copy the last row down for more.")
assert (h1, h2) == (L_FIRST, L_LAST), (h1, h2)
for rr in range(L_FIRST, L_LAST + 1):
    lg[f"G{rr}"].font = D.f(9, False, "ink2"); lg[f"G{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for k in ("D", "F"):
        lg[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
lg.conditional_formatting.add(f"G{L_FIRST}:G{L_LAST}", FormulaRule(formula=[f'G{L_FIRST}="Yes"'], font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))
lg.conditional_formatting.add(f"G{L_FIRST}:G{L_LAST}", FormulaRule(formula=[f'AND(G{L_FIRST}<>"",G{L_FIRST}<>"Yes")'],
                              font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
D.paint_canvas(lg, lg_end + 1, "Z")
S.finish_sheet(lg, PRODUCT, lg_end + 1, span=("A", "I"), freeze=f"A{L_FIRST}", tab_color="gold")
lg.print_title_rows = f"{L_FIRST - 1}:{L_FIRST - 1}"
lg.print_area = f"A1:I{L_FIRST + 99}"   # the first 100 rows, not 1,000 bare ones


# =====================================================================  4 and 5 plans
def build_plan(name, key, tagline, order_word):
    ws = wb.create_sheet(name)
    widths = {"A": 3, "B": 2, "C": 7, "D": 11, "E": 11, "F": 13, "G": 14, "H": 12, "I": 13}
    for i in range(N_DEBTS):
        widths[due(i)] = 11; widths[room(i)] = 11; widths[left(i)] = 11
    widths[GL(LEFT0 + N_DEBTS)] = 2
    S.set_widths(ws, widths)
    D.page_header(ws, name[2:], tagline, "Formulas only", span=("B", "I"), side_cols=2)
    D.h(ws, 6, 10)
    D.stat_strip(ws, 7, [
        ("C:E", "Debt-free in month", f'=IF(G{P_START}<0.005,"",IF(COUNTIF(G{P_FIRST}:G{P_LAST},">0.005")>={N_MONTHS},'
                                      f'"Not within 20 years",COUNTIF(G{P_FIRST}:G{P_LAST},">0.005")+1))', INT),
        ("F:G", "Total interest (est.)", f'=IF(ISNUMBER(C8),SUM(H{P_FIRST}:H{P_LAST}),"")', USD2),
        ("H:I", "Debt-free month", f'=IF(ISNUMBER(C8),EDATE({MONTH1},C8),"")', MON_FMT),
    ])
    s = ws["C9"]; s.value = ("Month 1 is next month. Each month: interest added (APR / 12), every minimum paid, the extra goes to "
                             f"the debts in {order_word} order, and a paid-off debt's money rolls forward.")
    s.font = D.f(9, False, "muted"); s.alignment = Alignment(vertical="center", wrap_text=True); ws.merge_cells("C9:I9"); D.h(ws, 9, 30)
    D.h(ws, 10, 10)
    D.frame(ws, 6, 10, "B", "I")
    D.h(ws, 11, 14)
    D.h(ws, 12, 18)
    e = ws["C12"]; e.value = "PER DEBT, FROM TAB 2 (COLUMNS J ONWARD)"; e.font = D.f(8, True, "muted")
    for row, lab in ((P_NAME, "Debt"), (P_PAIDOFF, "Paid off in month"), (P_RANK, "Pay order (see Room columns)"),
                     (P_APR, "APR"), (P_MIN, "Minimum"), (P_START, "Start: balance today")):
        c = ws[f"I{row}"]; c.value = lab; body(c, color="ink2", align="right", size=9); D.h(ws, row, D.GRID_ROW)
    D.h(ws, 19, 10)
    for i in range(N_DEBTS):
        dr = DROWS[i]
        ws[f"{due(i)}{P_NAME}"] = f'=IF({q(DEB)}!{DC["name"]}{dr}="","(empty)",{q(DEB)}!{DC["name"]}{dr})'
        body(ws[f"{due(i)}{P_NAME}"], bold=True, size=9)
        ws[f"{due(i)}{P_PAIDOFF}"] = (f'=IF({left(i)}{P_START}<0.005,"",IF(COUNTIF({left(i)}{P_FIRST}:{left(i)}{P_LAST},">0.005")>={N_MONTHS},'
                                      f'"20+ yrs",COUNTIF({left(i)}{P_FIRST}:{left(i)}{P_LAST},">0.005")+1))')
        body(ws[f"{due(i)}{P_PAIDOFF}"], INT, size=9)
        ws[f"{room(i)}{P_RANK}"] = f"={q(DEB)}!{DC[key]}{dr}"
        body(ws[f"{room(i)}{P_RANK}"], INT, size=9)
        ws[f"{due(i)}{P_APR}"] = f"=N({q(DEB)}!{DC['apr']}{dr})"
        body(ws[f"{due(i)}{P_APR}"], PCT2, size=9)
        ws[f"{due(i)}{P_MIN}"] = f"=N({q(DEB)}!{DC['min']}{dr})"
        body(ws[f"{due(i)}{P_MIN}"], USD2, size=9)
        ws[f"{left(i)}{P_START}"] = f'=IF(AND(ISNUMBER({q(DEB)}!{DC["today"]}{dr}),{q(DEB)}!{DC["today"]}{dr}>0.005),{q(DEB)}!{DC["today"]}{dr},0)'
        body(ws[f"{left(i)}{P_START}"], USD2, size=9)
    ws[f"G{P_START}"] = f"=SUM({left(0)}{P_START}:{left(N_DEBTS - 1)}{P_START})"
    body(ws[f"G{P_START}"], USD2, bold=True)
    for row in (P_NAME, P_PAIDOFF, P_RANK, P_APR, P_MIN, P_START):
        for k in D.cols("C", left(N_DEBTS - 1)):
            ws[f"{k}{row}"].border = Border(bottom=D.side("hair"))
    heads = [("C", "Month", "center"), ("D", "Date", "right"), ("E", "Extra pool", "right"), ("F", "Owed plus interest", "right"),
             ("G", "Balance after payments", "right"), ("H", "Interest this month", "right"), ("I", "Paid this month", "right")]
    heads += [(due(i), f"Due: {i + 1}", "right") for i in range(N_DEBTS)]
    heads += [(room(i), f"Room: {i + 1}", "right") for i in range(N_DEBTS)]
    heads += [(left(i), f"Left: {i + 1}", "right") for i in range(N_DEBTS)]
    header_cells(ws, P_HEAD, heads)
    RK = f"${room(0)}${P_RANK}:${room(N_DEBTS - 1)}${P_RANK}"
    for m in range(1, N_MONTHS + 1):
        R = P_HEAD + m
        prev = P_START if m == 1 else R - 1
        ws[f"C{R}"] = m; body(ws[f"C{R}"], "0", align="center", color="ink2")
        ws[f"D{R}"] = f"=EDATE({MONTH1},{m})"; body(ws[f"D{R}"], MON_FMT)
        ws[f"E{R}"] = f"=MAX(0,N({MONTHLY})-(SUM({due(0)}{R}:{due(N_DEBTS - 1)}{R})-SUM({room(0)}{R}:{room(N_DEBTS - 1)}{R})))"
        ws[f"F{R}"] = f"=SUM({due(0)}{R}:{due(N_DEBTS - 1)}{R})"
        ws[f"G{R}"] = f"=SUM({left(0)}{R}:{left(N_DEBTS - 1)}{R})"
        ws[f"H{R}"] = f"=F{R}-G{prev}"
        ws[f"I{R}"] = f"=F{R}-G{R}"
        for k in "EFGHI":
            body(ws[f"{k}{R}"], USD2, bold=(k == "G"))
        for i in range(N_DEBTS):
            dc, rc, lc = due(i), room(i), left(i)
            ws[f"{dc}{R}"] = f"=ROUND({lc}{prev}*(1+{dc}${P_APR}/12),2)"
            ws[f"{rc}{R}"] = f"={dc}{R}-MIN({dc}{R},{dc}${P_MIN})"
            ws[f"{lc}{R}"] = (f'=IF({rc}${P_RANK}="",{rc}{R},ROUND({rc}{R}-MIN({rc}{R},MAX(0,$E{R}-SUMIF({RK},"<"&{rc}${P_RANK},'
                              f'${room(0)}{R}:${room(N_DEBTS - 1)}{R}))),2))')
            for k in (dc, rc, lc):
                body(ws[f"{k}{R}"], USD2D, color="ink2", size=9)
        for k in D.cols("C", left(N_DEBTS - 1)):
            ws[f"{k}{R}"].border = Border(bottom=D.side("hair"))
        D.h(ws, R, 18)
    last = P_LAST + 1
    D.h(ws, last, 14)
    D.paint_canvas(ws, last, GL(LEFT0 + N_DEBTS + 1))
    S.finish_sheet(ws, PRODUCT, last, span=("A", "I"), freeze=f"D{P_FIRST}", tab_color="note",
                   canvas_cols=GL(LEFT0 + N_DEBTS + 1))
    ws.print_title_rows = f"{P_HEAD}:{P_HEAD}"
    return ws


build_plan(SNO, "snow", "Smallest balance first, ranked by today's balance.", "Snowball")
build_plan(AVA, "aval", "Highest APR first; a tie goes to the higher row.", "Avalanche")

# =====================================================================  6 Balance History
hs = wb.create_sheet(HIS)
widths = {"A": 3, "B": 2, "C": 11}
for i in range(N_DEBTS):
    widths[hbal(i)] = 12; widths[hint(i)] = 12
widths[GL(INT0 + N_DEBTS)] = 2
S.set_widths(hs, widths)
D.page_header(hs, "Balance History", "Month-end balance and interest for each debt, up to today.",
              "Formulas only", span=("B", "I"), side_cols=2)
D.h(hs, 6, 10)
D.stat_strip(hs, 7, [
    ("C:D", "Start month", f'=IF(COUNT({DA})=0,{MONTH1},IFERROR(DATE(YEAR(MIN({DA})),MONTH(MIN({DA})),1),{MONTH1}))', MON_FMT),
    ("E:F", "Months tracked", f"={N_HIST}", INT),
    ("G:I", "Interest since statements", f"={INTSUM}", USD2),
])
s = hs["C9"]; s.value = ("The current month shows today's balance; interest is added at month end. The 120 months start "
                         "at your earliest statement date.")
s.font = D.f(9, False, "muted"); s.alignment = Alignment(vertical="center", wrap_text=True); hs.merge_cells("C9:I9"); D.h(hs, 9, 30)
D.h(hs, 10, 10)
D.frame(hs, 6, 10, "B", "I")
D.h(hs, 11, 14)
assert H_START.endswith("$C$8")
heads = [("C", "Month", "left")]
for i in range(N_DEBTS):
    dr = DROWS[i]
    heads.append((hbal(i), "", "right"))
    heads.append((hint(i), "", "right"))
header_cells(hs, H_HEAD, heads, height=30)
for i in range(N_DEBTS):
    dr = DROWS[i]
    hs[f"{hbal(i)}{H_HEAD}"] = f'=IF({q(DEB)}!{DC["name"]}{dr}="","(empty)",{q(DEB)}!{DC["name"]}{dr})'
    hs[f"{hint(i)}{H_HEAD}"] = f'=IF({q(DEB)}!{DC["name"]}{dr}="","(empty)","Interest: "&{q(DEB)}!{DC["name"]}{dr})'
LD, LB, LA = (f"{q(LOG)}!${LG['date']}:${LG['date']}", f"{q(LOG)}!${LG['debt']}:${LG['debt']}", f"{q(LOG)}!${LG['amt']}:${LG['amt']}")
for k in range(N_HIST):
    R = H_FIRST + k
    hs[f"C{R}"] = "=$C$8" if k == 0 else f"=EDATE(C{R - 1},1)"
    body(hs[f"C{R}"], MON_FMT, align="left", color="ink2")
    for i in range(N_DEBTS):
        dr = DROWS[i]
        nm, bal, asof, apr = (f"{q(DEB)}!${DC[x]}${dr}" for x in ("name", "bal", "asof", "apr"))
        prev = '""' if k == 0 else f"{hbal(i)}{R - 1}"
        valid = f'AND({nm}<>"",ISNUMBER({bal}),ISNUMBER({asof}))'
        inmonth = f"AND({asof}>=$C{R},{asof}<=EOMONTH($C{R},0))"
        paid = (f'SUMIFS({LA},{LB},{nm},{LD},">"&{asof},{LD},">="&$C{R},{LD},"<="&MIN(EOMONTH($C{R},0),TODAY()))')
        base = f"MAX(0,IF({inmonth},{bal},{prev})-{paid})"
        frac = f"IF({inmonth},(EOMONTH($C{R},0)-{asof})/DAY(EOMONTH($C{R},0)),1)"
        hs[f"{hbal(i)}{R}"] = (f'=IF(OR(NOT({valid}),EOMONTH($C{R},0)<{asof},$C{R}>TODAY()),"",'
                               f'IF(EOMONTH($C{R},0)<TODAY(),ROUND({base}*(1+N({apr})/12*{frac}),2),{base}))')
        hs[f"{hint(i)}{R}"] = (f'=IF(OR(NOT({valid}),EOMONTH($C{R},0)<{asof},EOMONTH($C{R},0)>=TODAY()),"",'
                               f'ROUND({base}*N({apr})/12*{frac},2))')
        body(hs[f"{hbal(i)}{R}"], USD2D, size=9)
        body(hs[f"{hint(i)}{R}"], USD2D, size=9, color="ink2")
    for kk in D.cols("C", hint(N_DEBTS - 1)):
        hs[f"{kk}{R}"].border = Border(bottom=D.side("hair"))
    D.h(hs, R, 18)
HS_LAST = H_LAST + 1
D.h(hs, HS_LAST, 14)
D.paint_canvas(hs, HS_LAST, GL(INT0 + N_DEBTS + 1))
S.finish_sheet(hs, PRODUCT, HS_LAST, span=("A", "I"), freeze=f"D{H_FIRST}", tab_color="note", canvas_cols=GL(INT0 + N_DEBTS + 1))
hs.print_title_rows = f"{H_HEAD}:{H_HEAD}"

# =====================================================================  Start Here and Terms (card style)
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for lft, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); bodytext = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); bodytext = text
            mm = re.match(r"^([A-Z][A-Za-z0-9 ,\-]{1,40})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if mm:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(bodytext, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(bodytext, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = lft; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = lft[0]; L.number_format = lft[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, SHORT, "Start here. Log the payment you made; every balance and payoff date follows.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "A debt payoff tracker built around the payment you actually made. Type each debt once from its latest statement, then "
    "log every payment. Today's balance for every debt, the interest added, your debt-free month with Snowball and with "
    "Avalanche, and exactly what to pay each debt next month all update from that log.", "para")])
r = sh_card(sh, r, "Three steps", [
    (1, "Tab 2 Debts. Type over the example rows (or select D17:I41 and press Delete): each debt's name, type, statement "
        "balance, statement date, APR and minimum. Up to 25 debts.", "num"),
    (2, "Tab 1 Dashboard. Type the total you can put toward all debts each month, and pick Avalanche or Snowball.", "num"),
    (3, "Tab 3 Payment Log. One row per payment you make: date, the debt from the list, the amount. Clear the example rows "
        "first (select C17:F1016 and press Delete).", "num"),
])
r = sh_card(sh, r, "How to read the cells", [
    ((MONTHLY_EX or 1100, USD0), "Yellow cells with blue numbers are yours to type in. Each one shows a hint when you select it.", "chip-in"),
    ((831, USD2), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((dt.datetime(2030, 2, 1), MON_FMT), "Dark tiles are your answers. The Dashboard's main tile stays on screen as you scroll.", "chip-key"),
    (None, "Tabs 4, 5 and 6 are the maths: the Snowball and Avalanche plans month by month and each debt's balance history. "
           "You never type there. Each tab is protected without a password: use Review, Unprotect Sheet in Excel, or Data, "
           "Protect sheets and ranges in Google Sheets, to change anything.", "note"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "The example", [(None,
    "Five made-up debts from August 1, 2026 statements: a Visa card ($4,850 at 24.99%), a store card ($1,240 at 26.99%), a car "
    "loan ($11,600 at 7.4%), a student loan ($18,900 at 5.5%) and a personal loan ($3,300 at 12.9%), with $831 of minimums and "
    "$1,100 a month to spend. Twelve payments are logged for August to October. The plans count from the month you open the "
    "file, so the debt-free month and interest change as the months pass. On October 7, 2026 the file showed $38,154.87 owed, "
    "February 2030 with either plan, $4,965.45 of interest still to pay with Avalanche against $5,233.73 with Snowball, and "
    "$11,092.68 on minimums only.", "para")])
r = sh_card(sh, r, "Good to know", [
    (None, "New statement. Type its balance and date over the old ones on tab 2. Payments dated after that date keep counting; "
           "older ones stop, so nothing is counted twice. Counted? on tab 3 shows which rows count.", "para"),
    (None, "Room for years of payments. The Payment Log has 1,000 ready rows and the totals read whole columns. Need more? Copy "
           "the last row down as far as you like.", "para"),
    (None, "Ten-year window. Balance History tracks 120 months from your earliest statement date. Update each debt's statement "
           "balance and date at least once in that window and the tracker keeps going.", "para"),
    (None, "Which file to open. Excel 2016 or newer on Windows or Mac: double-click Debt-Payoff-Tracker.xlsx and click Enable "
           "Editing if asked. Google Sheets: upload the .xlsx to Google Drive and open it. On a phone, unzip on a computer first. "
           "No macros, no add-ons.", "para"),
    (None, "Printables. Debt-Payoff-Printables-Letter.pdf and -A4.pdf hold a paper debt list, payment log and a progress chart "
           "to color in.", "para"),
])
r = sh_card(sh, r, "Before you rely on a number", [(None, S.SHORT_NOTICE.format(date=CHECKED) + " Interest is an estimate "
           "(APR / 12 at each month end); your lender's statement is the official number. Talk to your lender before you change "
           "a payment.", "note")])
SH_LAST = r - 1
D.paint_canvas(sh, SH_LAST, "Z")
S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
               footer_notice="Estimates only. Not professional advice. See the Terms tab or LICENSE-AND-DISCLAIMER.txt.")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{SHORT}, version {VERSION}. Read before you rely on a number.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[2:]
r = 6
r = sh_card(tm, r, "More tools from the shop", [
    (None, ("Simple Paycheck Budget: give every paycheck a job before it lands", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this tracker helps, a short review on Etsy helps other people paying off debt find it. Thank you.", "para"),
    (None, "If a file won't open or a formula looks wrong, message me on Etsy and it will be fixed.", "para"),
])
TM_LAST = r - 1
D.paint_canvas(tm, TM_LAST, "Z")
S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb)]
if problems:
    print("AUDIT:", *problems[:40], sep="\n  ")
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
print({"MA": MA, "PL": PL, "a_min": a_min, "c_rows": (c_owed, c_paid, c_int, c_off, c_fix), "plan_rows": PLAN_ROWS,
       "min_rows": (r_mdt, r_min), "owed_rows": O_ROWS, "NM_FIRST": NM_FIRST, "NM_TOT": NM_TOT})
