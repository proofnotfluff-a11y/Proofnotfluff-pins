#!/usr/bin/env python3
"""Builds the Fair Share Chore Chart workbook (ProofNotFluff).
Usage: build_xlsx.py out.xlsx [--names "Alex,Jordan"] [--shares "60,40"] [--nofixed]
Default is the shipped sample: Alex and Jordan, equal shares, two fixed jobs."""
import sys, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

args = sys.argv[1:]
OUT = args[0]
def opt(k, d):
    return args[args.index(k) + 1] if k in args else d
NAMES = [n for n in opt("--names", "Alex,Jordan").split(",")] if opt("--names", "Alex,Jordan") != "" else []
SHARES = [float(s) if s else None for s in opt("--shares", "").split(",")] if opt("--shares", "") else []
NOFIXED = "--nofixed" in args
HERE = os.path.dirname(os.path.abspath(__file__))

CHECKED = "October 5, 2026"
INK, PAPER, ACC, LINE, SOFT, GREY = "1D2433", "F5F2EC", "C8502F", "D9D3C7", "FBF9F5", "5A6070"
INPUT = "FFF4D6"
F = "Calibri"
def font(sz=11, b=False, c=INK, i=False): return Font(name=F, size=sz, bold=b, color=c, italic=i)
def fill(c): return PatternFill("solid", fgColor=c)
thin = Side(style="thin", color=LINE); box = Border(left=thin, right=thin, top=thin, bottom=thin)
NOTICE = "Estimates only. Not professional advice. See the Terms tab."
UNLOCK = Protection(locked=False)
CAP = 70                      # chore rows 6..75
C0, C1 = 6, 6 + CAP - 1        # chore rows
M0, M1 = 6, 6 + CAP - 1        # math rows (header row 4, blank row 5)
CHART_ROWS = 44

FREQ = [("Twice a day", "=14"), ("Daily", "=7"), ("5x a week", "=5"), ("4x a week", "=4"), ("3x a week", "=3"),
        ("2x a week", "=2"), ("Weekly", "=1"), ("Every 2 weeks", "=1/2"), ("Monthly", "=12/52"), ("Every 3 months", "=4/52")]
FR0 = 24; FR1 = FR0 + len(FREQ) - 1
FREQ_RNG = f"Setup!$A${FR0}:$B${FR1}"
NM = "Setup!$B$7:$B$10"; RK = "Setup!$F$7:$F$10"; NPEOPLE = "Setup!$B$16"; ROT = "Setup!$E$13"; WEIGHT = "Setup!$B$14"; START = "Setup!$B$12"

# chore, area, minutes each time, how often, include, nobody wants it, always done by
CHORES = [
    ("Wash dishes or load the dishwasher", "Kitchen", 15, "Daily", "Y", "", ""),
    ("Unload the dishwasher and put dishes away", "Kitchen", 10, "5x a week", "Y", "", ""),
    ("Wipe counters, table and stovetop", "Kitchen", 5, "Daily", "Y", "", ""),
    ("Cook dinner", "Kitchen", 40, "5x a week", "Y", "", "P2"),
    ("Plan meals and make the grocery list", "Kitchen", 20, "Weekly", "Y", "", ""),
    ("Grocery shopping", "Errands", 60, "Weekly", "Y", "", ""),
    ("Put groceries away", "Kitchen", 10, "Weekly", "Y", "", ""),
    ("Take out trash and recycling", "Kitchen", 5, "3x a week", "Y", "Y", ""),
    ("Take the bins to the curb", "Outside", 5, "Weekly", "Y", "", ""),
    ("Clean out the fridge", "Kitchen", 20, "Every 2 weeks", "Y", "", ""),
    ("Clean the microwave and oven front", "Kitchen", 10, "Every 2 weeks", "Y", "", ""),
    ("Mop the kitchen floor", "Kitchen", 15, "Weekly", "Y", "", ""),
    ("Deep clean the oven", "Kitchen", 45, "Every 3 months", "Y", "", ""),
    ("Scrub the toilet", "Bathroom", 10, "Weekly", "Y", "Y", ""),
    ("Clean the shower or tub", "Bathroom", 20, "Weekly", "Y", "Y", ""),
    ("Clean bathroom sink, counter and mirror", "Bathroom", 10, "Weekly", "Y", "", ""),
    ("Mop the bathroom floor", "Bathroom", 10, "Weekly", "Y", "", ""),
    ("Restock toilet paper, soap and towels", "Bathroom", 5, "Weekly", "Y", "", ""),
    ("Wash and dry laundry", "Laundry", 15, "3x a week", "Y", "", ""),
    ("Fold and put away laundry", "Laundry", 20, "3x a week", "Y", "", ""),
    ("Change and wash bed sheets", "Laundry", 20, "Weekly", "Y", "", ""),
    ("Wash towels", "Laundry", 10, "Weekly", "Y", "", ""),
    ("Tidy shared rooms (10-minute reset)", "Living areas", 10, "Daily", "Y", "", ""),
    ("Vacuum carpets and rugs", "Living areas", 30, "Weekly", "Y", "", ""),
    ("Sweep or mop hard floors", "Living areas", 20, "Weekly", "Y", "", ""),
    ("Dust surfaces and shelves", "Living areas", 15, "Weekly", "Y", "", ""),
    ("Wipe light switches, handles and remotes", "Living areas", 10, "Every 2 weeks", "Y", "", ""),
    ("Clean windows and glass doors", "Living areas", 30, "Monthly", "Y", "", ""),
    ("Water the houseplants", "Living areas", 10, "Weekly", "Y", "", ""),
    ("Clean the entryway and shoe area", "Living areas", 10, "Weekly", "Y", "", ""),
    ("Pay bills and track shared costs", "Admin", 30, "Monthly", "Y", "", "P1"),
    ("Book repairs and appointments", "Admin", 20, "Monthly", "Y", "", ""),
    ("Run errands (pharmacy, post office, returns)", "Errands", 30, "Weekly", "Y", "", ""),
    ("Home check (filters, smoke alarm batteries)", "Home", 20, "Monthly", "Y", "", ""),
    ("Declutter and do a donation run", "Home", 45, "Monthly", "Y", "", ""),
    ("Mow the lawn", "Outside", 45, "Weekly", "N", "", ""),
    ("Weed and tidy the yard", "Outside", 30, "Every 2 weeks", "N", "", ""),
    ("Rake leaves or shovel snow (in season)", "Outside", 45, "Weekly", "N", "", ""),
    ("Feed the pets", "Pets", 5, "Twice a day", "N", "", ""),
    ("Walk the dog", "Pets", 30, "Daily", "N", "", ""),
    ("Scoop the litter box", "Pets", 5, "Daily", "N", "Y", ""),
    ("Wash the car and clean the inside", "Car", 45, "Monthly", "N", "", ""),
]

wb = Workbook()
wb.properties.creator = "ProofNotFluff"; wb.properties.lastModifiedBy = "ProofNotFluff"
wb.properties.title = "Fair Share Chore Chart"; wb.properties.company = None

def title(ws, text, sub, cols):
    ws.sheet_view.showGridLines = False
    ws["A1"] = text; ws["A1"].font = Font(name=F, size=18, bold=True, color=INK)
    ws["A2"] = sub; ws["A2"].font = font(10, c=GREY)
    ws["A3"] = NOTICE; ws["A3"].font = font(9, i=True, c=ACC)
    for r in (1, 2, 3):
        for c in range(1, cols + 1): ws.cell(r, c).fill = fill(PAPER)
    ws.row_dimensions[1].height = 30

def header(ws, row, labels, col0=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col0 + i, t); c.font = font(10, True, "FFFFFF"); c.fill = fill(INK)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if i else "left"); c.border = box
    ws.row_dimensions[row].height = 32

def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1): ws.column_dimensions[L(i)].width = w

def inp(c):
    c.fill = fill(INPUT); c.protection = UNLOCK; c.border = box

def textblock(ws, row, lines, col=1, width_merge=None):
    for ln in lines:
        if isinstance(ln, tuple): txt, st = ln
        else: txt, st = ln, ""
        c = ws.cell(row, col, txt)
        c.font = font(14 if st == "h" else 11, b=st in ("h", "b"), c=ACC if st == "a" else INK, i=st == "i")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if width_merge: ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + width_merge - 1)
        if st != "h" and txt and len(txt) > 95: ws.row_dimensions[row].height = 15 * (1 + len(txt) // 95)
        row += 1
    return row

# ---------------- Start Here ----------------
sh = wb.active; sh.title = "Start Here"
title(sh, "Fair Share Chore Chart", "For 1 to 4 adults who share a home. Split chores by minutes, not by count.", 2)
widths(sh, [100, 4])
textblock(sh, 5, [
    ("Start here (about 5 minutes)", "h"),
    "1. Setup tab: type each person's name in the yellow cells and the Monday you want to start. That's all you have to type.",
    "2. Chores tab: 42 common household chores are already listed with starting minutes and how often. Set Include to N for anything your home doesn't have, change any minutes that don't match your home, and add your own in the empty rows (set Include to Y for each one you add).",
    "3. Fair Split tab: see each person's minutes per week and how close the split is to fair. It recalculates every time you change a name, a chore or a minute.",
    "4. Weekly Rotation tab: jobs marked 'Nobody wants it' move to the next person every week and the sheet re-balances the rest around them, so nobody keeps the worst jobs and every week stays close. Jobs you pin to one person stay with them.",
    "5. Printable Chart tab: this week's chart for everyone, one page, grouped by person. Print it or keep it open on a phone.",
    "",
    ("How the split works", "h"),
    "Each chore's minutes per week = minutes each time x how many times a week. Jobs marked 'Nobody wants it' count 1.5x (change the 1.5 on Setup), so whoever takes them gets less of everything else.",
    "Each week the sheet places pinned jobs ('Always done by') and that week's rotating jobs first, then hands out the rest biggest first, each to the person who is furthest below their fair share so far.",
    "Want a 60/40 split because one person works longer hours? Type a share for everyone on Setup. Leave the shares blank for an even split.",
    "",
    ("Yellow cells are for typing", "h"),
    "Every other cell is a formula. The formula tabs are locked in Excel against accidental typing (no password: Review > Unprotect Sheet if you want to change them).",
    "Excel: open the .xlsx. Google Sheets: in Google Drive, upload the file, then open it with Google Sheets. Phone: the free Excel or Google Sheets app opens it. The Etsy app can't download files, so download from Etsy in a web browser (Etsy Help, 'Downloading a Digital Item', checked " + CHECKED + ").",
    "",
    ("Why minutes and not a list of names", "h"),
    "On an average day in 2025, 81 percent of people did some household activities and spent about 2 hours on them (U.S. Bureau of Labor Statistics, American Time Use Survey 2025 results, released June 25, 2026, bls.gov/news.release/atus.nr0.htm, checked " + CHECKED + "). Splitting chores by count hides that time: in this sheet's sample, Cook dinner is 200 minutes a week and Take the bins to the curb is 5.",
    "",
    ("Notice", "h"),
    ("For general information and planning only. This is not legal, tax, financial, medical or other professional advice, and using it doesn't create a professional relationship. Results are estimates based on the numbers you enter. Figures were checked on " + CHECKED + " and can change. See the Terms tab before you rely on anything here.", "i"),
])
for r in range(5, 40): sh.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")

# ---------------- Setup ----------------
st = wb.create_sheet("Setup")
title(st, "Setup", "Type in the yellow cells. Names must be different from each other.", 7)
widths(st, [40, 18, 20, 14, 12, 8, 60])
st["A5"] = "Who shares the home (1 to 4 people)"; st["A5"].font = font(12, True)
header(st, 6, ["Person", "Name", "Share of the work (optional, e.g. 60 and 40)", "Share used", "", "Order"])
for i in range(4):
    r = 7 + i
    st.cell(r, 1, f"Person {i+1}").font = font(11, True)
    c = st.cell(r, 2, NAMES[i] if i < len(NAMES) else None); inp(c)
    c = st.cell(r, 3, SHARES[i] if i < len(SHARES) else None); inp(c)
    st.cell(r, 4, f'=IF(B{r}="","",IF(AND($B$16>0,COUNTIFS($B$7:$B$10,"<>",$C$7:$C$10,">0")=$B$16),C{r}/SUMIFS($C$7:$C$10,$B$7:$B$10,"<>"),1/MAX(1,$B$16)))').number_format = "0%"
    st.cell(r, 6, f'=IF(B{r}="","",COUNTIF($B$7:B{r},"<>"))').font = font(9, c=GREY)
    for col in (1, 4, 6): st.cell(r, col).border = box
st["G7"] = '=IF(SUMPRODUCT(($B$7:$B$10<>"")*(COUNTIF($B$7:$B$10,$B$7:$B$10)>1))>0,"Two people have the same name. Make each name different.","")'
st["G7"].font = font(10, True, ACC)
st["A12"] = "Start date (a Monday)"; c = st["B12"]; c.value = dt.date(2026, 10, 12); c.number_format = "mmm d, yyyy"; inp(c)
st["A13"] = "Rotate 'Nobody wants it' jobs weekly?"; c = st["B13"]; c.value = "Yes"; inp(c)
st["D13"] = "In use:"; st["D13"].alignment = Alignment(horizontal="right")
st["E13"] = '=IF(AND($B$13="Yes",$B$16>1),"Yes","No")'; st["E13"].font = font(11, True)
st["A14"] = "Weight for 'Nobody wants it' jobs"; c = st["B14"]; c.value = 1.5; c.number_format = "0.0\"x\""; inp(c)
st["A16"] = "People sharing"; st["B16"] = '=COUNTIF($B$7:$B$10,"<>")'
st["A17"] = "Shares"
st["B17"] = '=IF($B$16=0,"Type at least one name above.",IF(COUNTIFS($B$7:$B$10,"<>",$C$7:$C$10,">0")=$B$16,"Using your shares.",IF(COUNTIFS($B$7:$B$10,"<>",$C$7:$C$10,">0")=0,"Even split.","Type a share for everyone or leave them all blank. Using an even split for now.")))'
st["A18"] = "Rotation"
st["B18"] = '=IF($E$13="Yes","On: each Nobody-wants-it job moves to the next person every week and the other jobs are re-balanced around it. Pinned jobs stay put.",IF($B$13<>"Yes","Off: the same split every week.","Off: rotation needs 2 or more people."))'
for a in ("A12", "A13", "A14", "A16", "A17", "A18", "D13"): st[a].font = font(11, True)
st["B17"].font = font(11, c=GREY); st["B18"].font = font(11, c=GREY)
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False); st.add_data_validation(dv); dv.add("B13")
dvd = DataValidation(type="date", operator="greaterThan", formula1="36526", showErrorMessage=True, error="Type a date, like 10/12/2026"); st.add_data_validation(dvd); dvd.add("B12")
dvw = DataValidation(type="decimal", operator="between", formula1="1", formula2="5", showErrorMessage=True, error="Use a number from 1 to 5"); st.add_data_validation(dvw); dvw.add("B14")
dvs = DataValidation(type="decimal", operator="greaterThan", formula1="0", showErrorMessage=True, error="Type a number above 0, or leave it blank"); st.add_data_validation(dvs); dvs.add("C7:C10")
st[f"A{FR0-2}"] = "How often (used by the Chores tab dropdown)"; st[f"A{FR0-2}"].font = font(12, True)
header(st, FR0 - 1, ["How often", "Times a week"])
for i, (n, v) in enumerate(FREQ):
    r = FR0 + i
    st.cell(r, 1, n).border = box; c = st.cell(r, 2, v); c.number_format = "0.00"; c.border = box
st.protection.sheet = True

# ---------------- Chores ----------------
ch = wb.create_sheet("Chores")
title(ch, "Chores", "Set Include to N for jobs your home doesn't have. Change minutes to match your home. Add your own in the empty rows.", 13)
widths(ch, [5, 44, 14, 11, 15, 9, 16, 16, 11, 11, 14, 12, 22])
header(ch, 5, ["#", "Chore", "Area", "Minutes each time", "How often", "Include? (Y/N)", "Nobody wants it? (Y = extra weight, rotates)", "Always done by (optional)", "Minutes per week", "Weighted minutes", "This week", "Sort key (no typing)", "Check"])
fx = {"P1": NAMES[0] if len(NAMES) > 0 else "", "P2": NAMES[1] if len(NAMES) > 1 else ""}
for i in range(CAP):
    r = C0 + i
    ch.cell(r, 1, i + 1).font = font(9, c=GREY)
    row = CHORES[i] if i < len(CHORES) else (None, None, None, None, None, None, "")
    vals = list(row)
    if vals[6] in fx: vals[6] = "" if NOFIXED else fx[vals[6]]
    for j, v in enumerate(vals):
        c = ch.cell(r, 2 + j, v if v != "" else None); inp(c)
    ch.cell(r, 9, f'=IF(OR($B{r}="",$F{r}<>"Y",$D{r}=""),0,$D{r}*IFERROR(VLOOKUP($E{r},{FREQ_RNG},2,FALSE),0))').number_format = "0"
    ch.cell(r, 10, f'=$I{r}*IF($G{r}="Y",IF(ISNUMBER({WEIGHT}),MAX(1,{WEIGHT}),1),1)').number_format = "0"
    ch.cell(r, 12, f'=IF(AND($J{r}>0,{NPEOPLE}>0),IF(OR(AND($H{r}<>"",COUNTIF({NM},$H{r})>0),AND($G{r}="Y",{ROT}="Yes")),1000000,0)+$J{r}+ROW()/10000000,-ROW())').font = font(8, c="A0A4AE")
    ch.cell(r, 11, f'=IF(ISNA(MATCH({i+1},Math!$C${M0}:$C${M1},0)),"",INDEX(Math!$AE${M0}:$AE${M1},MATCH({i+1},Math!$C${M0}:$C${M1},0)))')
    ch.cell(r, 11).font = font(11, True)
    ch.cell(r, 13, f'=IF($B{r}="","",IF(AND($F{r}<>"Y",$F{r}<>"N"),"Set Include to Y or N",IF(AND($E{r}<>"",ISNA(VLOOKUP($E{r},{FREQ_RNG},2,FALSE))),"Pick How often from the list",IF(AND($H{r}<>"",COUNTIF({NM},$H{r})=0),"Name not on Setup",IF(AND($F{r}="Y",OR($D{r}="",$E{r}="")),"Add minutes and How often","")))))')
    ch.cell(r, 13).font = font(9, True, ACC)
    for col in (1, 9, 10, 11, 12, 13): ch.cell(r, col).border = box
    for col in (4, 6, 7, 9, 10): ch.cell(r, col).alignment = Alignment(horizontal="center")
tr = C1 + 2
ch.cell(tr, 2, "Total for the home").font = font(11, True)
ch.cell(tr, 9, f"=SUM(I{C0}:I{C1})").number_format = "0"; ch.cell(tr, 10, f"=SUM(J{C0}:J{C1})").number_format = "0"
ch.cell(tr + 1, 2, "Hours per week").font = font(11, True); ch.cell(tr + 1, 9, f"=I{tr}/60").number_format = "0.0"
for c in (9, 10): ch.cell(tr, c).font = font(11, True)
ch.freeze_panes = "C6"; ch.row_dimensions[5].height = 48
dvf = DataValidation(type="list", formula1=f"={FREQ_RNG.split(':')[0].replace('$B','$A')}:$A${FR1}", allow_blank=True); ch.add_data_validation(dvf); dvf.add(f"E{C0}:E{C1}")
dvy = DataValidation(type="list", formula1='"Y,N"', allow_blank=True); ch.add_data_validation(dvy); dvy.add(f"F{C0}:F{C1}")
dvx = DataValidation(type="list", formula1='"Y"', allow_blank=True); ch.add_data_validation(dvx); dvx.add(f"G{C0}:G{C1}")
dvn = DataValidation(type="list", formula1=f"={NM}", allow_blank=True); ch.add_data_validation(dvn); dvn.add(f"H{C0}:H{C1}")
dvm = DataValidation(type="decimal", operator="between", formula1="0", formula2="600", showErrorMessage=True, error="Minutes from 0 to 600"); ch.add_data_validation(dvm); dvm.add(f"D{C0}:D{C1}")
ch.protection.sheet = True; ch.protection.formatColumns = False; ch.protection.formatRows = False

# ---------------- Math ----------------
# Columns: A rank, B key, C row, D in?, E chore, F weighted, G real, H pinned to, I rotates?, J rotate seq,
# then per week w (0..3) four load columns and an owner column: week 1 K:N + O, week 2 P:S + T, week 3 U:X + Y, week 4 Z:AC + AD,
# AE print-week owner, AF slot, AG chart key, AH times a week, AI check boxes
OWN = ["O", "T", "Y", "AD"]
mt = wb.create_sheet("Math")
title(mt, "Math (no typing)", "Each week: pinned jobs and rotating 'Nobody wants it' jobs are placed first, then the biggest jobs go to whoever is furthest below their fair share so far.", 35)
widths(mt, [5, 12, 5, 4, 34, 8, 8, 10, 6, 6] + [7, 7, 7, 7, 11] * 4 + [11, 5, 7, 7, 16])
labels = ["Rank", "Sort key", "Row", "In?", "Chore", "Weighted min", "Real min", "Pinned to", "Rotates?", "Rotate #"]
for w in range(4): labels += [f"Wk{w+1} load P1", f"Wk{w+1} load P2", f"Wk{w+1} load P3", f"Wk{w+1} load P4", f"Week {w+1}"]
labels += ["Print week", "Slot", "Chart key", "Times a week", "Check boxes"]
header(mt, M0 - 2, labels)
for k in range(CAP):
    r = M0 + k
    mt.cell(r, 1, k + 1)
    mt.cell(r, 2, f"=LARGE(Chores!$L${C0}:$L${C1},A{r})")
    mt.cell(r, 3, f"=MATCH(B{r},Chores!$L${C0}:$L${C1},0)")
    mt.cell(r, 4, f"=IF(B{r}>0,1,0)")
    mt.cell(r, 5, f'=IF(D{r}=1,INDEX(Chores!$B${C0}:$B${C1},C{r}),"")')
    mt.cell(r, 6, f"=IF(D{r}=1,INDEX(Chores!$J${C0}:$J${C1},C{r}),0)").number_format = "0.0"
    mt.cell(r, 7, f"=IF(D{r}=1,INDEX(Chores!$I${C0}:$I${C1},C{r}),0)").number_format = "0.0"
    mt.cell(r, 8, f'=IF(D{r}=0,"",IF(AND(INDEX(Chores!$H${C0}:$H${C1},C{r})<>"",COUNTIF({NM},INDEX(Chores!$H${C0}:$H${C1},C{r}))>0),INDEX(Chores!$H${C0}:$H${C1},C{r}),""))')
    mt.cell(r, 9, f'=IF(AND(D{r}=1,H{r}="",{ROT}="Yes",INDEX(Chores!$G${C0}:$G${C1},C{r})="Y"),1,0)')
    mt.cell(r, 10, f'=IF(I{r}=1,COUNTIF(I${M0}:I{r},1),"")')
    for w in range(4):
        b = 11 + 5 * w; oc = L(b + 4)
        for p in range(4):
            mt.cell(r, b + p, f'=IF(Setup!$B${7+p}="",9E+99,SUMIFS($F${M0-1}:$F{r-1},{oc}${M0-1}:{oc}{r-1},Setup!$B${7+p})/MAX(Setup!$D${7+p},0.0001))').number_format = "0.0"
        mt.cell(r, b + 4, f'=IF(D{r}=0,"",IF(H{r}<>"",H{r},IF(I{r}=1,INDEX({NM},MATCH(MOD(J{r}-1+{w},{NPEOPLE})+1,{RK},0)),INDEX({NM},MATCH(MIN({L(b)}{r}:{L(b+3)}{r}),{L(b)}{r}:{L(b+3)}{r},0)))))')
    mt.cell(r, 31, f"=CHOOSE('Printable Chart'!$Q$4,O{r},T{r},Y{r},AD{r})")
    mt.cell(r, 32, f'=IF(AE{r}="","",MATCH(AE{r},{NM},0))')
    mt.cell(r, 33, f'=IF(AF{r}="","",AF{r}*1000+COUNTIF(AF${M0}:AF{r},AF{r}))')
    mt.cell(r, 34, f'=IF(D{r}=1,IFERROR(VLOOKUP(INDEX(Chores!$E${C0}:$E${C1},C{r}),{FREQ_RNG},2,FALSE),0),"")').number_format = "0.00"
    mt.cell(r, 35, f'=IF(D{r}=0,"",IF(AH{r}>7,"2 a day: "&REPT("□",7),IF(AH{r}>=1,REPT("□ ",ROUND(AH{r},0)),INDEX(Chores!$E${C0}:$E${C1},C{r}))))')
    for col in range(1, 36): mt.cell(r, col).font = font(9)
mt.freeze_panes = "F6"
mt.protection.sheet = True

# ---------------- Fair Split ----------------
fs = wb.create_sheet("Fair Split", 3)
title(fs, "Fair Split", "Each person's share of the chores for the week shown on Printable Chart (this week unless you pick one).", 9)
widths(fs, [18, 12, 14, 14, 14, 14, 12, 8, 44])
header(fs, 6, ["Person", "Fair share", "Fair target (weighted min)", "Assigned (weighted min)", "Over (+) or under (-) target", "Real minutes a week", "Hours a week", "Jobs", "Real minutes, compared"])
for p in range(4):
    r = 7 + p; s = 7 + p
    fs.cell(r, 1, f'=IF(Setup!$B${s}="","",Setup!$B${s})').font = font(11, True)
    fs.cell(r, 2, f'=IF(A{r}="","",Setup!$D${s})').number_format = "0%"
    fs.cell(r, 3, f'=IF(A{r}="","",B{r}*SUM(Math!$F${M0}:$F${M1}))').number_format = "0"
    fs.cell(r, 4, f'=IF(A{r}="","",SUMIFS(Math!$F${M0}:$F${M1},Math!$AE${M0}:$AE${M1},A{r}))').number_format = "0"
    fs.cell(r, 5, f'=IF(A{r}="","",D{r}-C{r})').number_format = '+0;-0;0'
    fs.cell(r, 6, f'=IF(A{r}="","",SUMIFS(Math!$G${M0}:$G${M1},Math!$AE${M0}:$AE${M1},A{r}))').number_format = "0"
    fs.cell(r, 7, f'=IF(A{r}="","",F{r}/60)').number_format = "0.0"
    fs.cell(r, 8, f'=IF(A{r}="","",COUNTIF(Math!$AE${M0}:$AE${M1},A{r}))')
    fs.cell(r, 9, f'=IF(OR(A{r}="",MAX($F$7:$F$10)=0),"",REPT("█",ROUND(F{r}/MAX($F$7:$F$10)*30,0)))').font = Font(name=F, size=11, color=ACC)
    for col in range(1, 10): fs.cell(r, col).border = box
fs.cell(11, 1, "Whole home").font = font(11, True)
fs.cell(11, 3, "=SUM(C7:C10)").number_format = "0"; fs.cell(11, 4, "=SUM(D7:D10)").number_format = "0"
fs.cell(11, 6, "=SUM(F7:F10)").number_format = "0"; fs.cell(11, 7, "=F11/60").number_format = "0.0"; fs.cell(11, 8, "=SUM(H7:H10)")
for col in range(1, 9): fs.cell(11, col).font = font(11, True)
fs["A13"] = "Fair check"; fs["A13"].font = font(12, True)
fs["B13"] = ('=IF(Setup!$B$16=0,"Type names on the Setup tab.",IF(Setup!$B$16=1,"One person: every job is on their list.",'
             '"Furthest over target: "&TEXT(ROUND(MAX($E$7:$E$10),0),"+0;-0;0")&" weighted min a week. Furthest under: "&TEXT(ROUND(MIN($E$7:$E$10),0),"+0;-0;0")&". The biggest single job is "&TEXT(MAX(Math!$F$' + str(M0) + ':$F$' + str(M1) + '),"0")&" weighted min."))')
fs["B13"].font = font(11, True, ACC)
fs["B14"] = "To even it out more, split a big job into two rows (for example Cook dinner Mon to Wed and Cook dinner Thu and Fri), or pin a job to the person who is under."
fs["B15"] = "Weighted minutes count 'Nobody wants it' jobs at the weight on Setup, so the person who takes them does fewer minutes of everything else. Real minutes are the clock time."
fs["B16"] = "Shares and the weight are yours to set. The sheet does the math; the people in the home decide what is fair."
for a in ("B14", "B15", "B16"): fs[a].font = font(10, c=GREY)
for rr in (13, 14, 15, 16):
    fs.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=9); fs.cell(rr, 2).alignment = Alignment(wrap_text=True, vertical="top"); fs.row_dimensions[rr].height = 30
fs.protection.sheet = True

# ---------------- Weekly Rotation ----------------
ro = wb.create_sheet("Weekly Rotation", 4)
title(ro, "Weekly Rotation", "Who does each job in weeks 1 to 4. The pattern repeats every N weeks, where N is the number of people.", 7)
widths(ro, [42, 11, 9, 16, 16, 16, 16])
ro["A5"] = "=Setup!B18"; ro["A5"].font = font(10, True, ACC)
hdr = ["Chore (pinned and rotating first, then biggest first)", "Real min a week", "Pinned or rotates?"]
header(ro, 6, hdr + ["", "", "", ""])
for w in range(4):
    c = ro.cell(6, 4 + w, f'="Week {w+1}, "&TEXT({START}+{7*w},"mmm d")'); c.font = font(10, True, "FFFFFF")
for k in range(CAP):
    r = 7 + k; m = M0 + k
    ro.cell(r, 1, f'=IF(Math!$D${m}=1,Math!$E${m},"")')
    ro.cell(r, 2, f'=IF(Math!$D${m}=1,Math!$G${m},"")').number_format = "0"
    ro.cell(r, 3, f'=IF(Math!$H${m}<>"","Pinned",IF(Math!$I${m}=1,"Rotates",""))')
    for w in range(4): ro.cell(r, 4 + w, f"=Math!{OWN[w]}{m}")
    for col in range(1, 8): ro.cell(r, col).border = box
    ro.cell(r, 3).font = font(9, c=GREY)
tb = 7 + CAP + 1
ro.cell(tb, 1, "Real minutes each week").font = font(12, True)
header(ro, tb + 1, ["Person", "", "", "Week 1", "Week 2", "Week 3", "Week 4"])
for p in range(4):
    r = tb + 2 + p
    ro.cell(r, 1, f'=IF(Setup!$B${7+p}="","",Setup!$B${7+p})').font = font(11, True)
    for w in range(4):
        ro.cell(r, 4 + w, f'=IF($A{r}="","",SUMIFS(Math!$G${M0}:$G${M1},Math!{OWN[w]}${M0}:{OWN[w]}${M1},$A{r}))').number_format = "0"
    for col in range(1, 8): ro.cell(r, col).border = box
ro.freeze_panes = "A7"
ro.protection.sheet = True

# ---------------- Printable Chart ----------------
# One list grouped by person (Job | Who | Done | Min), portrait, one page for up to CHART_ROWS jobs, whatever the number of people.
from openpyxl.formatting.rule import FormulaRule
pc = wb.create_sheet("Printable Chart", 5)
pc.sheet_view.showGridLines = False
widths(pc, [46, 16, 18, 8])
pc.column_dimensions["P"].width = 8; pc.column_dimensions["Q"].width = 8
pc["A1"] = "Chore Chart"; pc["A1"].font = Font(name=F, size=22, bold=True, color=INK)
pc["A2"] = f'="Week of "&TEXT({START}+7*($P$4-1),"mmmm d")&" to "&TEXT({START}+7*($P$4-1)+6,"mmmm d, yyyy")'
pc["A2"].font = font(13, True, ACC)
pc["A4"] = "Week number to print (leave blank for this week):"; pc["A4"].font = font(10, c=GREY)
pc["A4"].alignment = Alignment(horizontal="right")
c = pc["B4"]; inp(c); c.alignment = Alignment(horizontal="center")
pc["C4"] = '="Calendar week "&P4&", rotation week "&Q4'; pc["C4"].font = font(9, c=GREY)
pc["P4"] = f'=IF(ISNUMBER(B4),MAX(1,INT(B4)),MAX(1,INT((TODAY()-{START})/7)+1))'
dvwk = DataValidation(type="whole", operator="between", formula1="1", formula2="520", allow_blank=True, showErrorMessage=True, errorStyle="stop", error="Type a week number, like 2, or leave it blank."); pc.add_data_validation(dvwk); dvwk.add("B4")
pc["Q4"] = f"=MOD(P4-1,MAX(1,{NPEOPLE}))+1"
for a in ("P4", "Q4"): pc[a].font = font(8, c="A0A4AE")
header(pc, 6, ["Minutes this week", "Minutes", "Jobs"])
for p in range(4):
    r = 7 + p
    pc.cell(r, 1, f'=IF(Setup!$B${7+p}="","",Setup!$B${7+p})').font = font(11, True)
    pc.cell(r, 2, f'=IF(A{r}="","",SUMIFS(Math!$G${M0}:$G${M1},Math!$AF${M0}:$AF${M1},{p+1}))').number_format = "0"
    pc.cell(r, 3, f'=IF(A{r}="","",COUNTIF(Math!$AF${M0}:$AF${M1},{p+1}))')
    for col in (2, 3): pc.cell(r, col).alignment = Alignment(horizontal="center")
H0 = 12; J0 = H0 + 1; J1 = H0 + CHART_ROWS
header(pc, H0, ["Job", "Who", "Done", "Min"])
for i in range(CHART_ROWS):
    r = J0 + i
    key = f"SMALL(Math!$AG${M0}:$AG${M1},{i+1})"
    pos = f"MATCH({key},Math!$AG${M0}:$AG${M1},0)"
    pc.cell(r, 1, f'=IFERROR(INDEX(Math!$E${M0}:$E${M1},{pos}),"")').font = font(11)
    pc.cell(r, 2, f'=IFERROR(INDEX({NM},INT({key}/1000)),"")').font = font(11, True)
    pc.cell(r, 3, f'=IFERROR(INDEX(Math!$AI${M0}:$AI${M1},{pos}),"")').font = font(11, c=GREY)
    c = pc.cell(r, 4, f'=IFERROR(INDEX(Math!$G${M0}:$G${M1},{pos}),"")'); c.number_format = "0"; c.font = font(10, c=GREY)
    c.alignment = Alignment(horizontal="right")
pc.conditional_formatting.add(f"A{J0}:D{J1}", FormulaRule(formula=[f'AND($A{J0}<>"",$B{J0}<>$B{J0-1})'], border=Border(top=Side(style="medium", color=INK))))
pc.conditional_formatting.add(f"A{J0}:D{J1}", FormulaRule(formula=[f'$A{J0}<>""'], border=Border(bottom=Side(style="hair", color=LINE))))
pc.cell(J1 + 1, 1, f'=IF(COUNT(Math!$AG${M0}:$AG${M1})>{CHART_ROWS},"+"&(COUNT(Math!$AG${M0}:$AG${M1})-{CHART_ROWS})&" more jobs: see the Weekly Rotation tab","")').font = font(10, True, ACC)
pc.cell(J1 + 2, 1, "One box for each time a job is due this week. Monthly and every-2-weeks jobs show how often instead. Minutes are starting estimates. " + NOTICE).font = font(8, i=True, c=GREY)
pc.merge_cells(start_row=J1 + 2, start_column=1, end_row=J1 + 2, end_column=4)
pc.cell(J1 + 2, 1).alignment = Alignment(wrap_text=True); pc.row_dimensions[J1 + 2].height = 24
pc.print_area = f"A1:D{J1 + 2}"
pc.page_setup.orientation = "portrait"; pc.page_setup.fitToWidth = 1; pc.page_setup.fitToHeight = 1
pc.sheet_properties.pageSetUpPr.fitToPage = True
pc.print_options.horizontalCentered = True
pc.page_margins.left = pc.page_margins.right = 0.5; pc.page_margins.top = pc.page_margins.bottom = 0.4
pc.protection.sheet = True

# ---------------- Terms ----------------
te = wb.create_sheet("Terms")
te.sheet_view.showGridLines = False; widths(te, [110])
with open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt")) as f: lines = f.read().strip().split("\n")
r = 1
for ln in lines:
    c = te.cell(r, 1, ln); c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = font(16 if r == 1 else 11, b=r <= 2)
    if len(ln) > 100: te.row_dimensions[r].height = 15 * (1 + len(ln) // 100)
    r += 1

# ---------------- Thank You ----------------
ty = wb.create_sheet("Thank You")
ty.sheet_view.showGridLines = False; widths(ty, [100])
ty["A1"] = "Thank you"; ty["A1"].font = Font(name=F, size=18, bold=True, color=INK)
ty["A3"] = "If the chart helps your home, a short review on Etsy helps other people find it."
ty["A4"] = "If anything doesn't open or looks wrong, message me on Etsy and I'll fix it."
ty["A6"] = "ProofNotFluff"; ty["A6"].font = font(11, True, ACC)
for a in ("A3", "A4"): ty[a].font = font(12)

for ws in wb.worksheets:
    if ws.title not in ('Printable Chart', 'Math'):
        ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
wb.active = 0
wb.save(OUT)
print("wrote", OUT)
