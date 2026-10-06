#!/usr/bin/env python3
"""Renders the Start Here guide (Letter, A4) and five 2000x2000 listing images from the recalculated workbooks in out/recalc."""
import openpyxl, html, os
from playwright.sync_api import sync_playwright
HERE = os.path.dirname(os.path.abspath(__file__))
R = lambda f: openpyxl.load_workbook(os.path.join(HERE, "out/recalc", f), data_only=True)
CHECKED = "October 5, 2026"
E = html.escape

def load(f):
    wb = R(f); m = wb["Math"]; fs = wb["Fair Split"]; ro = wb["Weekly Rotation"]
    jobs = []
    for r in range(6, 76):
        if m.cell(r, 4).value == 1:
            jobs.append(dict(chore=m.cell(r, 5).value, w=m.cell(r, 6).value, real=m.cell(r, 7).value, pinned=m.cell(r, 8).value,
                             rot=m.cell(r, 9).value, weeks=[m.cell(r, c).value for c in (15, 20, 25, 30)], boxes=m.cell(r, 35).value))
    people = []
    for r in range(7, 11):
        if fs.cell(r, 1).value:
            people.append(dict(name=fs.cell(r, 1).value, share=fs.cell(r, 2).value, target=fs.cell(r, 3).value, assigned=fs.cell(r, 4).value,
                               diff=fs.cell(r, 5).value, real=fs.cell(r, 6).value, hours=fs.cell(r, 7).value, jobs=fs.cell(r, 8).value))
    weeks = {}
    for r in range(80, 84):
        if ro.cell(r, 1).value: weeks[ro.cell(r, 1).value] = [ro.cell(r, c).value for c in range(4, 8)]
    tot = fs["F11"].value
    return dict(jobs=jobs, people=people, weeks=weeks, total=tot, check=fs["B13"].value, nchores=len(jobs))

S = load("sample.xlsx"); T3 = load("three.xlsx"); SH = load("shares.xlsx")
ch = R("sample.xlsx")["Chores"]
NLIST = sum(1 for r in range(6, 76) if ch.cell(r, 2).value)
assert NLIST == 42, NLIST
cook = next(j for j in S["jobs"] if j["chore"] == "Cook dinner")["real"]
bins = next(j for j in S["jobs"] if j["chore"] == "Take the bins to the curb")["real"]
r0 = lambda v: f"{v:,.0f}"

FONTS = """
@font-face{font-family:Poppins;font-weight:400;src:url(file:///usr/share/fonts/truetype/google-fonts/Poppins-Regular.ttf)}
@font-face{font-family:Poppins;font-weight:600;src:url(file:///usr/share/fonts/truetype/google-fonts/Poppins-Medium.ttf)}
@font-face{font-family:Poppins;font-weight:700;src:url(file:///usr/share/fonts/truetype/google-fonts/Poppins-Bold.ttf)}
@font-face{font-family:Carlito;font-weight:400;src:url(file:///usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf)}
@font-face{font-family:Carlito;font-weight:700;src:url(file:///usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf)}
:root{--paper:#F5F2EC;--ink:#1D2433;--acc:#C8502F;--line:#D9D3C7;--soft:#FBF9F5;--grey:#5A6070}
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:Carlito,sans-serif;color:var(--ink)}
h1,h2,h3,.p{font-family:Poppins,sans-serif}
table{border-collapse:collapse;width:100%}
"""

def split_table(d, big=False):
    rows = "".join(f"<tr><td><b>{E(p['name'])}</b></td><td>{p['share']:.0%}</td><td>{r0(p['target'])}</td><td>{r0(p['assigned'])}</td><td>{round(p['diff']):+d}</td><td><b>{r0(p['real'])}</b></td><td>{p['hours']:.1f}</td><td>{p['jobs']}</td></tr>" for p in d["people"])
    return f"<table class='t'><tr><th>Person</th><th>Fair share</th><th>Target (weighted min)</th><th>Assigned (weighted min)</th><th>Over or under</th><th>Real min a week</th><th>Hours</th><th>Jobs</th></tr>{rows}</table>"

def weeks_table(d):
    rows = "".join(f"<tr><td><b>{E(n)}</b></td>" + "".join(f"<td>{r0(v)}</td>" for v in ws) + "</tr>" for n, ws in d["weeks"].items())
    return f"<table class='t'><tr><th>Real minutes</th><th>Week 1</th><th>Week 2</th><th>Week 3</th><th>Week 4</th></tr>{rows}</table>"

def rot_table(d):
    rows = "".join(f"<tr><td>{E(j['chore'])}</td>" + "".join(f"<td>{E(w)}</td>" for w in j["weeks"]) + "</tr>" for j in d["jobs"] if j["rot"] == 1)
    return f"<table class='t'><tr><th>Nobody wants it</th><th>Week 1</th><th>Week 2</th><th>Week 3</th><th>Week 4</th></tr>{rows}</table>"

def chart(d, week=0, maxrows=19):
    cols = ""
    for p in d["people"]:
        mine = [j for j in d["jobs"] if j["weeks"][week] == p["name"]]
        rows = "".join(f"<tr><td>{E(j['chore'])}</td><td class='bx'>{E(j['boxes'])}</td><td class='mn'>{r0(j['real'])}</td></tr>" for j in mine[:maxrows])
        tot = sum(j["real"] for j in mine)
        more = f"<tr><td colspan=3 style='color:var(--grey);font-style:italic'>+ {len(mine)-maxrows} more jobs</td></tr>" if len(mine) > maxrows else ""
        cols += f"<div class='col'><div class='nm'>{E(p['name'])}</div><table class='ct'><tr><th>Job</th><th>Done</th><th>Min</th></tr>{rows}{more}<tr class='tot'><td>Minutes this week, all {len(mine)} jobs</td><td></td><td>{r0(tot)}</td></tr></table></div>"
    return f"<div class='chart'>{cols}</div>"


def chart_list(d, week=0):
    summ = "".join(f"<tr><td><b>{E(p['name'])}</b></td><td>{r0(sum(j['real'] for j in d['jobs'] if j['weeks'][week]==p['name']))}</td><td>{sum(1 for j in d['jobs'] if j['weeks'][week]==p['name'])}</td></tr>" for p in d["people"])
    rows = ""
    for p in d["people"]:
        first = True
        for j in [j for j in d["jobs"] if j["weeks"][week] == p["name"]]:
            rows += f"<tr class='{'grp' if first else ''}'><td>{E(j['chore'])}</td><td><b>{E(p['name'])}</b></td><td class='bx'>{E(j['boxes'])}</td><td class='mn'>{r0(j['real'])}</td></tr>"; first = False
    return f"<table class='t sm'><tr><th>Minutes this week</th><th>Minutes</th><th>Jobs</th></tr>{summ}</table><table class='lt'><tr><th>Job</th><th>Who</th><th>Done</th><th>Min</th></tr>{rows}</table>"

CSS_DOC = FONTS + """
@page{margin:0}
.page{width:100%;min-height:100vh;padding:0.6in 0.65in;page-break-after:always;background:#fff;position:relative}
.band{background:var(--paper);margin:-0.6in -0.65in 0.25in;padding:0.45in 0.65in 0.3in}
h1{font-size:30px;line-height:1.1} h2{font-size:17px;margin:16px 0 6px;color:var(--ink)} .sub{color:var(--acc);font-family:Poppins;font-weight:600;font-size:14px;margin-top:6px}
p,li{font-size:12.5px;line-height:1.45;margin:4px 0} ol,ul{padding-left:20px}
.note{font-size:10.5px;color:var(--grey);font-style:italic;border-left:3px solid var(--acc);padding:6px 10px;background:var(--soft);margin-top:10px}
.t{font-size:11.5px;margin:6px 0} .t th{background:var(--ink);color:#fff;font-weight:700;padding:5px 6px;text-align:left} .t td{border-bottom:1px solid var(--line);padding:4px 6px}
.big{font-family:Poppins;font-weight:700;color:var(--acc);font-size:22px}
.kpis{display:flex;gap:12px;margin:12px 0}.kpi{flex:1;background:var(--soft);border:1px solid var(--line);padding:10px 12px}.kpi b{display:block;font-family:Poppins;font-size:20px;color:var(--acc)}.kpi span{font-size:11px;color:var(--grey)}
.chart{display:flex;gap:14px}.col{flex:1}.nm{background:var(--ink);color:#fff;font-family:Poppins;font-weight:600;padding:5px 8px;font-size:13px}
.ct{font-size:10.5px}.ct th{text-align:left;color:var(--grey);font-size:9px;border-bottom:1px solid var(--ink);padding:3px 4px}.ct td{border-bottom:1px solid var(--line);padding:3px 4px}.bx{color:var(--grey);white-space:nowrap}.mn{text-align:right;color:var(--grey)}.tot td{font-weight:700;border-bottom:none}
.sm{width:55%}.lt{font-size:10px;margin-top:8px}.lt th{background:var(--ink);color:#fff;text-align:left;padding:3px 5px}.lt td{border-bottom:1px solid var(--line);padding:2px 5px}.lt .grp td{border-top:2px solid var(--ink)}
.foot{position:absolute;bottom:0.35in;left:0.65in;right:0.65in;font-size:9px;color:var(--grey);border-top:1px solid var(--line);padding-top:5px}
"""
NOTICE = f"For general information and planning only. This is not legal, tax, financial, medical or other professional advice, and using it doesn't create a professional relationship. Results are estimates based on the numbers you enter. Figures were checked on {CHECKED} and can change. See the Terms of Use page before you rely on anything here."
a, b = S["people"]
FOOT = "<div class='foot'>ProofNotFluff | Fair Share Chore Chart | Version 1, October 5, 2026 | Estimates only. Not professional advice.</div>"

doc = f"""<html><head><meta charset='utf-8'><style>{CSS_DOC}</style></head><body>
<div class='page'><div class='band'><h1>Fair Share Chore Chart</h1><div class='sub'>Start Here. Split the chores by minutes, not by count.</div></div>
<p>A chore list for 1 to 4 adults who share a home: couples, roommates, adult family. You type the names. The workbook splits {NLIST} prefilled chores by minutes per week, re-balances every week, rotates the jobs nobody wants, and prints one chart for everyone.</p>
<div class='kpis'><div class='kpi'><b>{NLIST}</b><span>chores already listed, with starting minutes and how often</span></div><div class='kpi'><b>{r0(cook)} vs {r0(bins)}</b><span>minutes a week: Cook dinner vs Take the bins to the curb, in the sample</span></div><div class='kpi'><b>{r0(a['real'])} / {r0(b['real'])}</b><span>real minutes a week for {E(a['name'])} and {E(b['name'])} in week 1 of the sample</span></div></div>
<h2>What's inside (4 files, no Canva)</h2>
<ul><li><b>Fair-Share-Chore-Chart.xlsx</b>: one workbook for Excel or Google Sheets. Tabs: Start Here, Setup, Chores, Fair Split, Weekly Rotation, Printable Chart, Math, Terms, Thank You.</li>
<li><b>START-HERE-Guide</b> PDF in US Letter and A4 (this guide).</li><li><b>LICENSE-AND-DISCLAIMER.txt</b>: the Terms of Use.</li></ul>
<h2>Why minutes</h2>
<p>A list that gives each person 15 chores can still be lopsided, because chores are not the same size. In the sample, Cook dinner is {r0(cook)} minutes a week and Take the bins to the curb is {r0(bins)}. The sheet adds up minutes, so the split is about time, not the number of jobs.</p>
<p>On an average day in 2025, 81 percent of people did some household activities and spent about 2 hours on them (U.S. Bureau of Labor Statistics, American Time Use Survey 2025 results, released June 25, 2026, bls.gov/news.release/atus.nr0.htm, checked {CHECKED}).</p>
<div class='note'>{NOTICE}</div>{FOOT}</div>

<div class='page'><h1 style='font-size:24px'>Set it up in about 5 minutes</h1>
<ol><li><b>Download from Etsy in a web browser.</b> Etsy's help page says you can't download a digital purchase through the Etsy app; sign in on a computer or a phone's browser instead (Etsy Help, "Downloading a Digital Item", checked {CHECKED}).</li>
<li><b>Open the workbook.</b> Excel: open the .xlsx. Google Sheets: upload it to Google Drive, then open it with Google Sheets. Phone: the free Excel or Google Sheets app opens it.</li>
<li><b>Setup tab.</b> Type each person's name in the yellow cells and the Monday you want to start. Names must be different from each other.</li>
<li><b>Chores tab.</b> Set Include to N for anything your home doesn't have (yard, pets and car start as N). Change minutes that don't match your home. Add your own chores in the empty rows and set Include to Y for each; there is room for 70.</li>
<li><b>Mark the jobs nobody wants</b> with Y. They count 1.5x (you can change the 1.5 on Setup) and move to the next person every week.</li>
<li><b>Pin a job</b> to one person with "Always done by" if they want it, like the sample's Cook dinner.</li>
<li><b>Print the Printable Chart tab</b> (one page, grouped by person) or keep it open on a phone. It shows this week unless you type a week number.</li></ol>
<h2>How the split works</h2>
<p>Minutes per week = minutes each time x times a week (Monthly counts as 12 out of 52 weeks). Each week the sheet places pinned jobs and that week's rotating jobs first, then hands out the rest biggest first, each to the person furthest below their fair share so far. Change any name, chore, minute or share and everything recalculates.</p>
<p><b>Uneven shares.</b> If one person works longer hours, type a share for everyone on Setup (for example 60 and 40). Leave all shares blank for an even split.</p>
<p><b>Yellow cells are for typing.</b> Everything else is a formula. The formula tabs are locked in Excel against accidental typing, with no password (Review, Unprotect Sheet).</p>
<h2>The tabs</h2>
<table class='t'><tr><th>Tab</th><th>What it does</th></tr>
<tr><td>Setup</td><td>Names, optional shares, start date, rotation on or off, the weight for jobs nobody wants</td></tr>
<tr><td>Chores</td><td>{NLIST} chores with area, minutes, how often, include, nobody wants it, always done by, and who has it this week</td></tr>
<tr><td>Fair Split</td><td>Each person's fair target, assigned minutes, hours and number of jobs, with a fair check line</td></tr>
<tr><td>Weekly Rotation</td><td>Who does each job in weeks 1 to 4, and each person's minutes for each week</td></tr>
<tr><td>Printable Chart</td><td>One page per week with check boxes for each time a job is due</td></tr>
<tr><td>Math</td><td>The working, shown openly so you can see how each job was placed</td></tr></table>
{FOOT}</div>

<div class='page'><h1 style='font-size:24px'>Worked examples from the file</h1>
<h2>A couple: {E(a['name'])} and {E(b['name'])}, even split, week 1</h2>
<p>{S['nchores']} chores included, {r0(S['total'])} real minutes a week for the home. {E(b['name'])} is pinned to Cook dinner and {E(a['name'])} to Pay bills.</p>{split_table(S)}
<p>Each week stays close, even as the jobs nobody wants change hands:</p>{weeks_table(S)}{rot_table(S)}
<h2>Three roommates: the same chores, one more name</h2>
<p>Typing a third name on Setup re-splits everything. Real minutes in each of the first four weeks:</p>{weeks_table(T3)}
<h2>Uneven shares: 60 and 40</h2>
<p>With shares of 60 and 40 typed on Setup:</p>{split_table(SH)}
<div class='note'>Example names and numbers only. The minutes are starting estimates, not measurements; time a chore once and change its minutes to match your home.</div>{FOOT}</div>

<div class='page'><h1 style='font-size:24px'>The Printable Chart, sample week 1</h1><p style='margin-bottom:8px'>Week of October 12 to October 18, 2026. One page, grouped by person. One box for each time the job is due that week; monthly and every-2-weeks jobs show how often instead.</p>
{chart_list(S, 0)}{FOOT}</div>

<div class='page'><h1 style='font-size:24px'>Sources, terms and thank you</h1>
<h2>Sources checked {CHECKED}</h2>
<ul><li>U.S. Bureau of Labor Statistics, American Time Use Survey 2025 results, released June 25, 2026: bls.gov/news.release/atus.nr0.htm</li>
<li>Etsy Help, "Downloading a Digital Item": help.etsy.com</li></ul>
<p>The minutes for each chore are ProofNotFluff's starting estimates for planning, not survey figures.</p>
<h2>Terms of Use</h2>
<p>Your purchase gives you a personal license for yourself or one business you own. No resale, sharing or redistribution. Information and planning only, not professional advice. Results are estimates based on your inputs. Provided "as is"; to the extent the law allows, liability is limited to the price you paid. Microsoft Excel, Google Sheets, Canva and Etsy are trademarks of their owners. ProofNotFluff isn't affiliated with, sponsored or endorsed by them. The full Terms are in LICENSE-AND-DISCLAIMER.txt and on the Terms tab. Made with AI assistance and reviewed and tested by the shop owner.</p>
<h2>Thank you</h2>
<p class='big'>If the chart helps your home, a short review on Etsy helps other people find it.</p>
<p style='font-size:15px'>If anything doesn't open or looks wrong, message me on Etsy and I'll fix it.</p>
<p style='margin-top:14px;font-size:11px;color:var(--grey)'>Copyright 2026 ProofNotFluff. All rights reserved.</p>{FOOT}</div>
</body></html>"""

# ---------------- listing images ----------------
CSS_IMG = FONTS + """
.v{width:2000px;height:2000px;background:var(--paper);padding:110px 120px;position:relative;overflow:hidden}
.v h1{font-size:118px;line-height:1.02;letter-spacing:-2px}.v h2{font-size:64px;line-height:1.1}.k{font-family:Poppins;font-weight:600;color:var(--acc);font-size:44px;margin-top:26px}
.card{background:#fff;border:2px solid var(--line);box-shadow:0 18px 50px rgba(29,36,51,.12);padding:46px;margin-top:60px}
.pill{display:inline-block;background:var(--ink);color:#fff;font-family:Poppins;font-weight:600;font-size:34px;padding:16px 30px;margin:0 16px 16px 0}
.t{font-size:34px}.t th{background:var(--ink);color:#fff;padding:16px 18px;text-align:left;font-family:Poppins;font-weight:600;font-size:28px}.t td{border-bottom:2px solid var(--line);padding:16px 18px}
.chart{display:flex;gap:36px}.col{flex:1}.nm{background:var(--ink);color:#fff;font-family:Poppins;font-weight:600;padding:14px 20px;font-size:40px}
.ct{font-size:27px}.ct th{text-align:left;color:var(--grey);font-size:22px;border-bottom:2px solid var(--ink);padding:8px}.ct td{border-bottom:1px solid var(--line);padding:8px}.bx{color:var(--grey);white-space:nowrap}.mn{text-align:right;color:var(--grey)}.tot td{font-weight:700;border-bottom:none}
.sm{width:60%;font-size:26px}.sm th{font-size:22px;padding:8px 14px}.sm td{padding:6px 14px}.lt{font-size:22px;margin-top:20px}.lt th{background:var(--ink);color:#fff;text-align:left;padding:6px 12px;font-family:Poppins;font-weight:600;font-size:20px}.lt td{border-bottom:1px solid var(--line);padding:3px 12px}.lt .grp td{border-top:3px solid var(--ink)}
.bar{height:56px;background:var(--acc);display:inline-block;vertical-align:middle}
.foot{position:absolute;left:120px;right:120px;bottom:70px;font-size:28px;color:var(--grey)}
ul.big{list-style:none;margin-top:40px}ul.big li{font-size:54px;padding:34px 0;border-bottom:2px solid var(--line)}ul.big li b{font-family:Poppins;color:var(--acc)}
"""
maxreal = max(p["real"] for p in S["people"])
bars = "".join(f"<div style='margin:30px 0;font-size:44px'><b style='display:inline-block;width:260px'>{E(p['name'])}</b><span class='bar' style='width:{int(900*p['real']/maxreal)}px'></span> <b style='margin-left:20px'>{r0(p['real'])} min</b><div style='color:var(--grey);font-size:34px;margin-left:260px'>{p['hours']:.1f} hours a week, {p['jobs']} jobs</div></div>" for p in S["people"])
imgs = f"""<html><head><meta charset='utf-8'><style>{CSS_IMG}</style></head><body>
<div class='v' id='v1'><div class='k' style='margin-top:0'>FOR COUPLES AND ROOMMATES</div><h1 style='margin-top:20px'>Fair Share<br>Chore Chart</h1>
<h2 style='margin-top:40px;font-weight:400'>Split chores by <b style='color:var(--acc)'>minutes</b>,<br>not by count.</h2>
<div class='card' style='padding:34px'>{chart(S, 0, 11)}</div>
<div style='margin-top:50px'><span class='pill'>{NLIST} chores prefilled</span><span class='pill'>1 to 4 adults</span><span class='pill'>Excel + Google Sheets</span></div>
<div class='foot'>Sample names and numbers from the workbook. Digital download.</div></div>

<div class='v' id='v2'><div class='k' style='margin-top:0'>FAIR SPLIT TAB</div><h2 style='margin-top:20px'>Type the names.<br>It balances the minutes.</h2>
<div class='card'>{bars}<p style='font-size:34px;color:var(--grey);margin-top:20px'>{S['nchores']} chores, {r0(S['total'])} real minutes a week in the sample home. Change any name, chore or minute and it recalculates.</p></div>
<div class='card' style='margin-top:40px'>{split_table(S)}</div>
<div class='foot'>Sample week 1. Minutes are starting estimates you can edit.</div></div>

<div class='v' id='v3'><div class='k' style='margin-top:0'>WEEKLY ROTATION TAB</div><h2 style='margin-top:20px'>Nobody keeps the<br>worst jobs.</h2>
<p style='font-size:42px;margin-top:30px;line-height:1.35'>Mark a job "Nobody wants it". It counts 1.5x and moves to the next person every week. Everything else re-balances around it.</p>
<div class='card'>{rot_table(T3)}</div><div class='card' style='margin-top:40px'>{weeks_table(T3)}</div>
<div class='foot'>Three-roommate example: real minutes for each person in weeks 1 to 4.</div></div>

<div class='v' id='v4' style='background:#fff;padding-top:80px'><div class='k' style='margin-top:0'>PRINTABLE CHART TAB</div><h2 style='margin-top:10px;font-size:56px'>One page for everyone, every week.</h2>
<div style='margin-top:30px'>{chart_list(S, 0)}</div>
<div class='foot'>Sample week of October 12, 2026. One box per time a job is due.</div></div>

<div class='v' id='v5'><div class='k' style='margin-top:0'>WHAT'S INSIDE</div><h2 style='margin-top:20px'>4 files, not 40.<br>No design app needed.</h2>
<ul class='big'><li><b>1 workbook</b> for Excel and Google Sheets, 9 tabs</li><li><b>{NLIST} chores</b> already listed with minutes and how often</li><li><b>Up to 4 people</b>, even or custom shares like 60/40</li><li><b>Weekly rotation</b> for the jobs nobody wants</li><li><b>Printable chart</b>, one page, grouped by person</li><li><b>Start Here guide</b> PDF, US Letter and A4</li><li><b>Terms of Use</b> text file</li></ul>
<div class='foot'>Digital download. No physical item ships. Made with AI assistance and reviewed and tested by the shop owner.</div></div>
</body></html>"""

os.makedirs(os.path.join(HERE, "pkg"), exist_ok=True); os.makedirs(os.path.join(HERE, "img"), exist_ok=True)
with open(os.path.join(HERE, "out/guide.html"), "w") as f: f.write(doc)
with open(os.path.join(HERE, "out/imgs.html"), "w") as f: f.write(imgs)
with sync_playwright() as p:
    b = p.chromium.launch(); pg = b.new_page()
    pg.goto("file://" + os.path.join(HERE, "out/guide.html")); pg.wait_for_timeout(600)
    pg.pdf(path=os.path.join(HERE, "pkg", "START-HERE-Guide-Letter.pdf"), width="8.5in", height="11in", print_background=True)
    pg.pdf(path=os.path.join(HERE, "pkg", "START-HERE-Guide-A4.pdf"), format="A4", print_background=True)
    pg2 = b.new_page(viewport={"width": 2000, "height": 2000})
    pg2.goto("file://" + os.path.join(HERE, "out/imgs.html")); pg2.wait_for_timeout(600)
    for i in range(1, 6):
        pg2.query_selector(f"#v{i}").screenshot(path=os.path.join(HERE, "img", f"0{i}.png"))
    b.close()
print("rendered")
