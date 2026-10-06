#!/usr/bin/env python3
"""Freelance Income Tracker, version 3 dashboard design (built Oct 6, 2026).

One entry tab (Income Log: date, client, what it was for, amount, method, paid on) feeds a
Dashboard (income received in the tax year, still owed, set-aside at the buyer's own rate,
month by month, IRS estimated-tax payment periods, top clients, the owed-to-you list) and a
Clients tab (per-client totals and the editable client and payment-method lists).

Income counts in the month the money arrived (the Paid on date). A blank Paid on means owed.

Usage: python3 engines/freelance-income-tracker/build_xlsx.py out.xlsx
"""
import datetime as dt
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402

PRODUCT = "Freelance Income Tracker"
AS_OF = "Figures checked Oct 6, 2026"
OUT = sys.argv[1] if len(sys.argv) > 1 else "Freelance-Income-Tracker.xlsx"
USD0, USD2, PCT, INT, DATE = S.FMT["usd0"], S.FMT["usd2"], S.FMT["pct"], S.FMT["int"], S.FMT["date"]
YEAR_FMT = "0"
HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- example data (a freelance designer, 2026)
d = dt.datetime
CLIENTS = ["Alder Street Bakery", "Juniper Dental Studio", "Bright Path Tutoring", "Kestrel Outdoor Co",
           "Larkspur Realty", "Fernwood Yoga"]
METHODS = ["Bank transfer", "Check", "Card", "Payment app", "Platform payout", "Cash", "Other"]
ROWS = [  # date, client, what, amount, method, paid on (None = still owed)
    (d(2025, 12, 12), "Fernwood Yoga", "Gift card design", 300, "Payment app", d(2026, 1, 6)),
    (d(2026, 1, 8), "Alder Street Bakery", "Menu redesign", 850, "Bank transfer", d(2026, 1, 20)),
    (d(2026, 1, 15), "Juniper Dental Studio", "Logo refresh", 1200, "Check", d(2026, 2, 2)),
    (d(2026, 1, 28), "Bright Path Tutoring", "Flyer set", 320, "Payment app", d(2026, 1, 30)),
    (d(2026, 2, 10), "Kestrel Outdoor Co", "Product photo retouching", 640, "Bank transfer", d(2026, 2, 24)),
    (d(2026, 2, 18), "Alder Street Bakery", "Social graphics, February", 400, "Bank transfer", d(2026, 3, 3)),
    (d(2026, 3, 5), "Larkspur Realty", "Listing brochure template", 950, "Card", d(2026, 3, 12)),
    (d(2026, 3, 20), "Alder Street Bakery", "Social graphics, March", 400, "Bank transfer", d(2026, 4, 1)),
    (d(2026, 3, 26), "Juniper Dental Studio", "Business cards", 280, "Check", d(2026, 4, 9)),
    (d(2026, 4, 14), "Kestrel Outdoor Co", "Spring catalog layout", 2400, "Bank transfer", d(2026, 5, 6)),
    (d(2026, 4, 22), "Alder Street Bakery", "Social graphics, April", 400, "Bank transfer", d(2026, 5, 4)),
    (d(2026, 5, 7), "Fernwood Yoga", "Class schedule posters", 360, "Payment app", d(2026, 5, 9)),
    (d(2026, 5, 20), "Bright Path Tutoring", "Website banner set", 540, "Payment app", d(2026, 6, 3)),
    (d(2026, 6, 2), "Larkspur Realty", "Open house signs", 720, "Card", d(2026, 6, 10)),
    (d(2026, 6, 18), "Alder Street Bakery", "Social graphics, June", 400, "Bank transfer", d(2026, 7, 1)),
    (d(2026, 7, 9), "Kestrel Outdoor Co", "Summer sale graphics", 880, "Bank transfer", d(2026, 7, 27)),
    (d(2026, 7, 23), "Juniper Dental Studio", "Patient forms redesign", 1100, "Check", d(2026, 8, 12)),
    (d(2026, 8, 6), "Fernwood Yoga", "Logo and brand sheet", 1500, "Bank transfer", d(2026, 8, 21)),
    (d(2026, 8, 19), "Alder Street Bakery", "Social graphics, August", 400, "Bank transfer", d(2026, 9, 2)),
    (d(2026, 8, 28), "Kestrel Outdoor Co", "Holiday catalog draft", 1800, "Bank transfer", None),
    (d(2026, 9, 3), "Larkspur Realty", "Fall mailer", 1150, "Card", d(2026, 9, 15)),
    (d(2026, 9, 9), "Bright Path Tutoring", "Enrollment flyers", 380, "Payment app", d(2026, 9, 11)),
    (d(2026, 9, 22), "Juniper Dental Studio", "Waiting room posters", 650, "Check", None),
    (d(2026, 10, 1), "Alder Street Bakery", "Social graphics, October", 400, "Bank transfer", None),
]
EX_YEAR, EX_RATE, EX_GOAL, EX_LATE = 2026, 0.25, 2000, 30
EX_REC = sum(a for *_, a, m, p in ROWS if p and p.year == EX_YEAR)
EX_OWED = sum(a for *_, a, m, p in ROWS if p is None)
EX_OWED_N = sum(1 for r in ROWS if r[5] is None)
EX_MONTHS = len({p.month for *_, p in ROWS if p and p.year == EX_YEAR})
EX_JAN = sum(a for *_, a, m, p in ROWS if p and p.year == EX_YEAR and p.month == 1)

N_LOG = 300
N_CLIENTS = 25
N_METHODS = 8

wb = S.new_workbook(PRODUCT, version="3")

# ---------------------------------------------------------------- fixed addresses, worked out before building
DB = "Dashboard"
LOG = "'Income Log'"
CL = "Clients"
# Dashboard settings rows (Card A: pad 11, title 12, sub 13, eyebrow 14, inputs 15 to 18)
Y_ROW, RT_ROW, GL_ROW, LT_ROW = 15, 16, 17, 18
YEAR, RATE, GOAL, LATE = (f"{DB}!$D${Y_ROW}", f"{DB}!$D${RT_ROW}", f"{DB}!$D${GL_ROW}", f"{DB}!$D${LT_ROW}")
# Income Log table (table_card at row 11 with title and sub: head row 15, body 16 to 315)
LG_TOP = 11
LG_FIRST = LG_TOP + 5
LG_LAST = LG_FIRST + N_LOG - 1


def lg(col):
    return f"{LOG}!${col}${LG_FIRST}:${col}${LG_LAST}"


# Clients tab tables
CL_TOP = 11
CL_FIRST = CL_TOP + 5
CL_LAST = CL_FIRST + N_CLIENTS - 1
CLIENT_LIST = f"{CL}!$C${CL_FIRST}:$C${CL_LAST}"
MT_TOP = CL_LAST + 3  # one canvas row after the clients card ends (end row = last + 1)
MT_FIRST = MT_TOP + 5
MT_LAST = MT_FIRST + N_METHODS - 1
METHOD_LIST = f"{CL}!$C${MT_FIRST}:$C${MT_LAST}"


def paid_between(start, end, amount_col="F"):
    """Income whose Paid on date falls in [start, end)."""
    return f'SUMIFS({lg(amount_col)},{lg("H")},">="&{start},{lg("H")},"<"&{end})'


YS = f"DATE({YEAR},1,1)"
YE = f"DATE({YEAR}+1,1,1)"
REC = f'IF({YEAR}="",0,{paid_between(YS, YE)})'

# =====================================================================  Income Log
lgws = wb.create_sheet("Income Log")
S.set_widths(lgws, {"A": 3, "B": 2, "C": 14, "D": 22, "E": 28, "F": 12, "G": 15, "H": 14, "I": 18, "J": 12, "K": 9,
                    "L": 2, "M": 3})
D.page_header(lgws, "Income Log", "One row per invoice or payment. Leave Paid on blank until the money arrives.",
              "Example rows: type over them", span=("B", "L"), side_cols=4)
# stat strip (rows 6 to 9)
D.h(lgws, 6, 10)
FIX = (f'SUMPRODUCT(((((({lg("C")}<>"")+({lg("D")}<>"")+({lg("E")}<>"")+({lg("H")}<>""))>0)*({lg("F")}=""))'
       f'+(({lg("F")}<>"")*({lg("D")}=""))+(({lg("F")}<>"")*({lg("C")}="")*({lg("H")}=""))>0)*1)'
       f'+SUMPRODUCT(({lg("F")}<>"")*({lg("D")}<>""))-SUM({CL}!$H${CL_FIRST}:$H${CL_LAST})')
D.stat_strip(lgws, 7, [
    ("C:D", "Received in the tax year", f"={REC}", USD0),
    ("E:E", "Still owed to you", f'=SUMPRODUCT({lg("F")}*({lg("H")}=""))', USD0),
    ("F:H", "Set aside at your rate", f"=IFERROR({REC}*{RATE},0)", USD0),
    ("I:K", "Rows to fix", f"={FIX}", INT),
])
D.h(lgws, 9, 10)
D.frame(lgws, 6, 9, "B", "L")
D.status_rule(lgws, "I8", "I8>0", fill_tint=False)
D.h(lgws, 10, 14)


def lg_calc(kind):
    def fn(r):
        if kind == "status":
            days = f'TEXT(MAX(0,TODAY()-C{r}),"0")'
            return (f'=IF(F{r}="","",IF(H{r}<>"","Paid",IF(C{r}="","Owed, add a date",'
                    f'IF(TODAY()-C{r}>{LATE},"Late, "&{days}&" days","Owed, "&{days}&" days"))))')
        if kind == "setaside":
            return f'=IF(OR(F{r}="",H{r}=""),"",F{r}*{RATE})'
        if kind == "owedno":
            return f'=IF(AND(F{r}<>"",H{r}=""),MAX(K${LG_FIRST - 1}:K{r - 1})+1,"")'
    return fn


cols_lg = [
    dict(col="C", head="Date", kind="input", fmt=DATE, values=[x[0] for x in ROWS],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The invoice date, or the day you did the work, for example Oct 1, 2026. It shows the month by name so you can check it.",
         error="Enter a real date between 2000 and 2100, for example Oct 1, 2026."),
    dict(col="D", head="Client", kind="input", values=[x[1] for x in ROWS], validation=("list_range", CLIENT_LIST),
         prompt="Pick from the list. Add or rename clients on the Clients tab.",
         error="Pick a client from the list. To add one, type the name on the Clients tab first."),
    dict(col="E", head="What it was for", kind="input", values=[x[2] for x in ROWS], validation=("textLength", 0, 60),
         prompt="A few words you will recognise later, for example Logo refresh.", error="Up to 60 characters."),
    dict(col="F", head="Amount", kind="input", fmt=USD2, align="right", values=[x[3] for x in ROWS],
         validation=("decimal", 0, None), prompt="What the client pays you for this row, in dollars. Example: 850.",
         error="Enter 0 or more dollars."),
    dict(col="G", head="Method", kind="input", values=[x[4] for x in ROWS], validation=("list_range", METHOD_LIST),
         prompt="How you were paid. Optional. Edit the list on the Clients tab.",
         error="Pick a method from the list, or add one on the Clients tab."),
    dict(col="H", head="Paid on", kind="input", fmt=DATE, values=[x[5] for x in ROWS],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The day the money arrived. Leave blank while it is still owed. Income counts in this month.",
         error="Enter a date between 2000 and 2100, or leave it blank while the money is owed."),
    dict(col="I", head="Status", kind="calc", formula=lg_calc("status")),
    dict(col="J", head="Set aside", kind="calc", fmt=USD2, align="right", formula=lg_calc("setaside")),
    dict(col="K", head="Owed no.", kind="muted", formula=lg_calc("owedno"), fmt="0", align="right"),
]
cols_lg[-1]["kind"] = "calc"
f1, f2, lg_end = D.table_card(lgws, LG_TOP, "B", "L", cols_lg, N_LOG, title="Income log",
                              sub="Yellow columns are yours. Status, Set aside and Owed no. fill in on their own. "
                                  f"{N_LOG} rows.")
assert (f1, f2) == (LG_FIRST, LG_LAST), (f1, f2)
for r in range(LG_FIRST, LG_LAST + 1):
    lgws[f"I{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    lgws[f"K{r}"].font = D.f(9, False, "muted")
    for k in ("D", "E", "G"):
        lgws[f"{k}{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for word, color, tint in (("Late", "accent", "accent_tint"), ("Owed", "gold", "input_fill"), ("Paid", "teal", "teal_tint")):
    lgws.conditional_formatting.add(f"I{LG_FIRST}:I{LG_LAST}", FormulaRule(
        formula=[f'LEFT($I{LG_FIRST},4)="{word}"'], font=Font(color=D.T[color], bold=True),
        fill=PatternFill("solid", bgColor=D.T[tint]), stopIfTrue=True))
D.paint_canvas(lgws, lg_end + 1, "Z")
S.finish_sheet(lgws, PRODUCT, lg_end + 1, span=("A", "M"), freeze=f"A{LG_FIRST}", tab_color="gold")
lgws.page_setup.orientation = "landscape"
lgws.print_title_rows = f"{LG_FIRST - 1}:{LG_FIRST - 1}"

# =====================================================================  Clients
cl = wb.create_sheet("Clients")
S.set_widths(cl, {"A": 3, "B": 2, "C": 30, "D": 15, "E": 15, "F": 10, "G": 20, "H": 9, "I": 7, "J": 2, "K": 3})
D.page_header(cl, "Clients", "Your client list and payment methods. Totals use the tax year on the Dashboard.",
              AS_OF.replace("Figures checked", "Built"), span=("B", "J"), side_cols=3)
D.h(cl, 6, 10)
D.stat_strip(cl, 7, [
    ("C:C", "Clients on your list", f'=COUNTA({CL}!$C${CL_FIRST}:$C${CL_LAST})', INT),
    ("D:E", "Paid you this tax year", f'=COUNTIF({CL}!$D${CL_FIRST}:$D${CL_LAST},">0")', INT),
    ("F:H", "Your biggest client's share", f'=IFERROR(MAX({CL}!$F${CL_FIRST}:$F${CL_LAST}),0)', PCT),
])
D.h(cl, 9, 10)
D.frame(cl, 6, 9, "B", "J")
D.h(cl, 10, 14)


def cl_calc(kind):
    def fn(r):
        if kind == "rec":
            return f'=IF(OR(C{r}="",{YEAR}=""),"",SUMIFS({lg("F")},{lg("D")},C{r},{lg("H")},">="&{YS},{lg("H")},"<"&{YE}))'
        if kind == "owed":
            return f'=IF(C{r}="","",SUMPRODUCT({lg("F")}*({lg("D")}=C{r})*({lg("H")}="")))'
        if kind == "share":
            return f'=IF(C{r}="","",IFERROR(D{r}/{REC},0))'
        if kind == "bar":
            return (f'=IF(OR(C{r}="",N(D{r})<=0),"",IFERROR(REPT("{D.BAR_CHAR}",MAX(0,ROUND(D{r}/MAX($D${CL_FIRST}:$D${CL_LAST})*12,0))),""))')
        if kind == "rows":
            return f'=IF(C{r}="","",IF(COUNTIF($C${CL_FIRST}:C{r},C{r})>1,0,COUNTIFS({lg("D")},C{r},{lg("F")},"<>")))'
        if kind == "rank":
            return (f'=IF(OR(C{r}="",N(D{r})<=0),"",COUNTIF($D${CL_FIRST}:$D${CL_LAST},">"&D{r})'
                    f'+COUNTIF($D${CL_FIRST}:D{r},D{r}))')
    return fn


cols_cl = [
    dict(col="C", head="Client", kind="input", values=CLIENTS, validation=("textLength", 0, 40),
         prompt="A client name or a short label. Up to 40 characters. A rename does not change old log rows; Rows to fix counts them until you pick the new name in the log.",
         error="Up to 40 characters."),
    dict(col="D", head="Received this tax year", kind="calc", fmt=USD0, align="right", formula=cl_calc("rec")),
    dict(col="E", head="Still owed", kind="calc", fmt=USD0, align="right", formula=cl_calc("owed")),
    dict(col="F", head="Share", kind="calc", fmt=PCT, align="right", formula=cl_calc("share")),
    dict(col="G", head="", kind="calc", formula=cl_calc("bar")),
    dict(col="H", head="Log rows", kind="calc", fmt=INT, align="right", formula=cl_calc("rows")),
    dict(col="I", head="Rank", kind="calc", fmt="0", align="right", formula=cl_calc("rank")),
]
c1, c2, cl_end = D.table_card(cl, CL_TOP, "B", "J", cols_cl, N_CLIENTS, title="Your clients",
                              sub=f"Type up to {N_CLIENTS} names. The Client column on the Income Log offers this list.")
assert (c1, c2) == (CL_FIRST, CL_LAST), (c1, c2)
for r in range(CL_FIRST, CL_LAST + 1):
    cl[f"G{r}"].font = D.f(9, False, "teal"); cl[f"G{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cl[f"I{r}"].font = D.f(9, False, "muted")
    cl[f"C{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)


def mt_calc(kind):
    def fn(r):
        if kind == "count":
            return f'=IF(C{r}="","",COUNTIF({lg("G")},C{r}))'
        if kind == "rec":
            return f'=IF(OR(C{r}="",{YEAR}=""),"",SUMIFS({lg("F")},{lg("G")},C{r},{lg("H")},">="&{YS},{lg("H")},"<"&{YE}))'
    return fn


cols_mt = [
    dict(col="C", head="Payment method", kind="input", values=METHODS, validation=("textLength", 0, 30),
         prompt="A way clients pay you. Up to 30 characters.", error="Up to 30 characters."),
    dict(col="D", head="Rows in the log", kind="calc", fmt=INT, align="right", formula=mt_calc("count")),
    dict(col="E", head="Received this tax year", kind="calc", fmt=USD0, align="right", formula=mt_calc("rec")),
]
m1, m2, mt_end = D.table_card(cl, MT_TOP, "B", "J", cols_mt, N_METHODS, title="Payment methods",
                              sub="Edit or add methods. The Income Log offers this list.")
assert (m1, m2) == (MT_FIRST, MT_LAST), (m1, m2)
for r in range(MT_FIRST, MT_LAST + 1):
    cl[f"C{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
S.page_break_before(cl, MT_TOP)
D.paint_canvas(cl, mt_end + 1, "Z")
S.finish_sheet(cl, PRODUCT, mt_end + 1, span=("A", "K"), freeze=None, tab_color="teal")

# =====================================================================  Dashboard
ws = wb.create_sheet(DB, 0)
S.set_widths(ws, {"A": 3, "B": 2, "C": 31, "D": 14, "E": 2, "F": 3, "G": 2, "H": 24, "I": 12, "J": 17, "K": 2, "L": 3})
D.page_header(ws, PRODUCT, "Log each payment once. See what came in, what is owed and what to set aside.", AS_OF)
TOP = 11

A = D.Card(ws, TOP, "B", "C", "D", "E", wb=wb)
A.title("Your settings", "Type over the yellow cells. Every tab follows them.")
A.eyebrow("YOUR YEAR")
y = A.input("Tax year to show", EX_YEAR, YEAR_FMT, name="TaxYear", validation=("whole", 2000, 2100),
            prompt="The calendar year to total, for example 2026. Old rows stay in the log.", prompt_title="Tax year to show")
rt = A.input("Share you set aside for tax", EX_RATE, PCT, name="SetAsideRate", validation=("decimal", 0, 0.9),
             prompt="Your own rate, for example 25%. This file does not tell you what rate to use. See Notes and sources.")
gl = A.input("Monthly income goal", EX_GOAL, USD0, name="MonthlyGoal", validation=("decimal", 0, None),
             prompt="The income you want each month, in dollars. Example: 2000. Type 0 if you have no goal.")
lt = A.input("Days before an invoice is late", EX_LATE, INT, name="LateAfterDays", validation=("whole", 1, 365),
             prompt="A whole number of days, 1 to 365. Owed rows older than this show as Late. Example: 30.")
assert (y, rt, gl, lt) == (Y_ROW, RT_ROW, GL_ROW, LT_ROW), (y, rt, gl, lt)
A.eyebrow("THIS TAX YEAR SO FAR")
a_rec = A.calc("Income received", f"={REC}", USD0)
a_owed = A.calc("Still owed to you", f'=SUMPRODUCT({lg("F")}*({lg("H")}=""))', USD0)
a_set = A.calc("Set aside at your rate", f"=IFERROR(D{a_rec}*$D${RT_ROW},0)", USD0)
a_keep = A.calc("Yours after the set-aside", f"=D{a_rec}-D{a_set}", USD0, total=True, tint="total_tint")
a_rows = A.calc("Rows in the log", f'=COUNT({lg("F")})', INT)
a_other = A.calc("Paid in other years", f'=IF({YEAR}="",0,COUNT({lg("H")})-COUNTIFS({lg("H")},">="&{YS},{lg("H")},"<"&{YE}))', INT)
ISSUES = f"{LOG}!$I$8"
a_pill = A.text(f'=IF({ISSUES}>0,{ISSUES}&" log "&IF({ISSUES}=1,"row needs","rows need")&" a fix: see Rows to fix",'
                f'"Every row in the log is counted")', size=10, color="ink", bold=True)
ws[f"C{a_pill}"].alignment = Alignment(horizontal="center", vertical="center")
A.close()
D.status_rule(ws, f"C{a_pill}:D{a_pill}", f"{ISSUES}>0")

B = D.Card(ws, TOP, "G", "H", "I", "K", extra="J", wb=wb)
B.title("Month by month", "Income received, counted in the month the money arrived.")
m_first = B.r
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
mmax = f"MAX($I${m_first}:$I${m_first + 11})"
for i, name in enumerate(MONTHS, 1):
    rr = B.bar(name, f'=IF({YEAR}="",0,{paid_between(f"DATE({YEAR},{i},1)", f"DATE({YEAR},{i + 1},1)")})', USD0, mmax, "teal")
    ws.conditional_formatting.add(f"I{rr}", FormulaRule(formula=[f'AND($D${GL_ROW}>0,I{rr}>=$D${GL_ROW})'],
                                                         font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))
m_last = m_first + 11
MRANGE = f"$I${m_first}:$I${m_last}"
MWITH = f'COUNTIF({MRANGE},">0")'
B.text(f'=IF({MWITH}=0,"No income received in this tax year",IF(N($D${GL_ROW})<=0,"Months with income: "&{MWITH},"Months at or above your "&TEXT($D${GL_ROW},"$#,##0")'
       f'&" goal: "&COUNTIF({MRANGE},">="&$D${GL_ROW})&" of "&{MWITH}&" with income"))', size=9, color="ink2")
B.close()

# second row of cards
row2 = max(A.end, B.end) + 2
C = D.Card(ws, row2, "B", "C", "D", "E", wb=wb)
C.title("Set aside by IRS payment period", "Your set-aside rate times income received in each period.")
C.eyebrow("ESTIMATED TAX PERIODS")
PERIODS = [  # start month, end month (exclusive, may roll into next year), label start, due month/day, due year offset
    (1, 4, "Jan 1 to Mar 31", 4, 15, 0),
    (4, 6, "Apr 1 to May 31", 6, 15, 0),
    (6, 9, "Jun 1 to Aug 31", 9, 15, 0),
    (9, 13, "Sep 1 to Dec 31", 1, 15, 1),
]
p_rows = []
for sm, em, lab, dm, dd, off in PERIODS:
    rr = C.calc(f'=IFERROR("{lab}, due "&TEXT(DATE({YEAR}+{off},{dm},{dd}),"mmm d"),"{lab}")',
                f'=IF({YEAR}="",0,{paid_between(f"DATE({YEAR},{sm},1)", f"DATE({YEAR},{em},1)")})*$D${RT_ROW}', USD0)
    p_rows.append(rr)
c_tot = C.calc("Set aside for the year", f"=SUM(D{p_rows[0]}:D{p_rows[-1]})", USD0, total=True, tint="total_tint")
C.text("A due date on a weekend or legal holiday moves to the next business day (IRS, checked Oct 6, 2026).",
       size=9, color="muted", height=D.GRID_ROW * 2)
C.close()

E = D.Card(ws, row2, "G", "H", "I", "K", extra="J", wb=wb)
E.title("Top clients this tax year", "Income received. The full list is on the Clients tab.")
t_first = E.r
tmax = f"MAX($I${t_first}:$I${t_first + 4})"
CR = f"{CL}!$I${CL_FIRST}:$I${CL_LAST}"
for k in range(1, 6):
    rr = E.bar(f'=IFERROR(INDEX({CLIENT_LIST},MATCH({k},{CR},0)),"")',
               f'=IFERROR(INDEX({CL}!$D${CL_FIRST}:$D${CL_LAST},MATCH({k},{CR},0)),"")', USD0, tmax, "blue")
t_last = t_first + 4
E.text(f'=IF(N($D${a_rec})<=0,"",IF(COUNT($I${t_first}:$I${t_last})=1,"One client brought in all of it",'
       f'"Your top "&COUNT($I${t_first}:$I${t_last})&" clients brought in "&TEXT(SUM($I${t_first}:$I${t_last})/$D${a_rec},"0%")&" of it"))',
       size=9, color="ink2")
E.close()

# ---- hero tiles (rows 6 to 9)
D.tile(ws, 6, "B", "E", f'=IF({YEAR}="","INCOME RECEIVED THIS TAX YEAR","INCOME RECEIVED IN "&{YEAR})', f"=$D${a_rec}", USD0,
       f'=IF(N($D${a_owed})<=0,"Nothing is owed to you right now","Plus "&TEXT($D${a_owed},"$#,##0")&" owed to you (any year) on "'
       f'&COUNT({lg("K")})&IF(COUNT({lg("K")})=1," invoice"," invoices"))', dark=True)
AVG = f"IFERROR($D${a_rec}/{MWITH},0)"
GAPG = f"({AVG}-$D${GL_ROW})"
v, sub, _ = D.tile(ws, 6, "G", "K", "AVERAGE MONTH WITH INCOME", f"={AVG}", USD0,
                   f'=IF(OR(N($D${GL_ROW})<=0,{MWITH}=0),"",IF({GAPG}<0,TEXT(-{GAPG},"$#,##0")&" a month below your "'
                   f'&TEXT($D${GL_ROW},"$#,##0")&" goal",TEXT({GAPG},"$#,##0")&" a month above your "&TEXT($D${GL_ROW},"$#,##0")&" goal"))')
ws.conditional_formatting.add(v.coordinate, FormulaRule(formula=[f'AND(N($D${GL_ROW})>0,{MWITH}>0,{GAPG}<0)'],
                                                        font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
ws.conditional_formatting.add(v.coordinate, FormulaRule(formula=[f'AND(N($D${GL_ROW})>0,{MWITH}>0,{GAPG}>=0)'],
                                                        font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))
D.h(ws, 10, 14)

# ---- full-width card: Owed to you
o_top = max(C.end, E.end) + 2
N_OWED = 8
r = o_top
D.h(ws, r, 12); r += 1
t = ws[f"C{r}"]; t.value = "Owed to you"; t.font = D.f(12, True); ws.merge_cells(f"C{r}:J{r}"); D.h(ws, r, 24); r += 1
t = ws[f"C{r}"]; t.value = "Every row on the Income Log with no Paid on date, in log order. Type the date when the money arrives."
t.font = D.f(9, False, "muted"); t.alignment = Alignment(vertical="top"); ws.merge_cells(f"C{r}:J{r}"); D.h(ws, r, 18); r += 1
D.h(ws, r, 6); r += 1
heads = {"C": ("Client", "left"), "D": ("Amount", "right"), "H": ("What it was for", "left"), "I": ("Dated", "right"),
         "J": ("Status", "left")}
for k in D.cols("C", "J"):
    ws[f"{k}{r}"].border = Border(bottom=D.side("ink"))
for k, (txt, al) in heads.items():
    c = ws[f"{k}{r}"]; c.value = txt.upper(); c.font = D.f(8, True, "muted")
    c.alignment = Alignment(horizontal=al, vertical="bottom", indent=1 if k == "J" else 0)
D.h(ws, r, 22); r += 1
o_first = r
OWN = lg("K")
for k in range(1, N_OWED + 1):
    m = f"MATCH({k},{OWN},0)"
    vals = {"C": (f'=IFERROR(IF(INDEX({lg("D")},{m})="","(no client)",INDEX({lg("D")},{m})),"")', None, "left"),
            "D": (f'=IFERROR(INDEX({lg("F")},{m}),"")', USD0, "right"),
            "H": (f'=IFERROR(IF(INDEX({lg("E")},{m})="","",INDEX({lg("E")},{m})),"")', None, "left"),
            "I": (f'=IFERROR(IF(INDEX({lg("C")},{m})="","",INDEX({lg("C")},{m})),"")', DATE, "right"),
            "J": (f'=IFERROR(INDEX({lg("I")},{m}),"")', None, "left")}
    for col, (fml, fmt, al) in vals.items():
        c = ws[f"{col}{r}"]; c.value = fml
        if fmt:
            c.number_format = fmt
        c.font = D.f(10, col == "D"); c.alignment = Alignment(horizontal=al, vertical="center", indent=1 if col == "J" else 0)
    for kk in D.cols("C", "J"):
        ws[f"{kk}{r}"].border = Border(bottom=D.side("hair"))
    D.h(ws, r, D.GRID_ROW); r += 1
o_last = r - 1
for word, color in (("Late", "accent"), ("Owed", "gold")):
    ws.conditional_formatting.add(f"J{o_first}:J{o_last}", FormulaRule(formula=[f'LEFT($J{o_first},4)="{word}"'],
                                  font=Font(color=D.T[color], bold=True), stopIfTrue=True))
c = ws[f"C{r}"]
c.value = (f'=IF(COUNT({OWN})=0,"Nothing owed. Nice.",IF(COUNT({OWN})>{N_OWED},COUNT({OWN})-{N_OWED}&" more on the Income Log. Total owed: "'
           f'&TEXT($D${a_owed},"$#,##0"),"Total owed: "&TEXT($D${a_owed},"$#,##0")&" on "&COUNT({OWN})&IF(COUNT({OWN})=1," invoice"," invoices")))')
c.font = D.f(9, True, "ink2"); c.alignment = Alignment(vertical="center"); ws.merge_cells(f"C{r}:J{r}")
D.h(ws, r, D.GRID_ROW); r += 1
D.h(ws, r, 12)
o_end = r
D.frame(ws, o_top, o_end, "B", "K")
S.page_break_before(ws, o_top)

# ---- full-width card: Notes and sources
n_top = o_end + 2
NOTES = [
    ("When income counts", "Each row counts in the month of its Paid on date, the day the money arrived. A December invoice "
     "paid in January counts in January. Rows with no Paid on date are owed and are not in the totals."),
    ("Set-aside rate", "You choose it; this file does not tell you what rate to use. For reference, US self-employment tax "
     "alone is 15.3% (12.4% Social Security plus 2.9% Medicare), generally on 92.35% of net earnings, and you usually owe it "
     "once net earnings reach $400, per IRS Topic 554 (irs.gov/taxtopics/tc554, page updated Sep 24, 2026, checked Oct 6, 2026). "
     "Income tax comes on top. The set-aside here is on income received, before business expenses."),
    ("Payment periods", "The IRS estimated tax periods are Jan 1 to Mar 31 (due Apr 15), Apr 1 to May 31 (due Jun 15), "
     "Jun 1 to Aug 31 (due Sep 15) and Sep 1 to Dec 31 (due Jan 15 of the next year). A due date on a Saturday, Sunday or legal "
     "holiday moves to the next business day. Source: IRS, Pay as you go, so you won't owe (irs.gov, page updated "
     "Sep 25, 2026, checked Oct 6, 2026)."),
    ("Expenses", "This file tracks income only. Pair it with an expense tracker to see profit."),
    ("Any currency", "The maths works in any currency. Select the money cells and pick your symbol under Format, Number."),
    ("Before you rely on a number", S.SHORT_NOTICE.format(date="Oct 6, 2026").replace("Terms of Use page", "Terms tab")),
]
r = n_top
D.h(ws, r, 12); r += 1
t = ws[f"C{r}"]; t.value = "Notes and sources"; t.font = D.f(12, True); ws.merge_cells(f"C{r}:J{r}"); D.h(ws, r, 24); r += 1
D.h(ws, r, 6); r += 1
width = sum(S.width_of(ws, k) for k in "CDEFGHIJ")
for head, body in NOTES:
    cell = ws[f"C{r}"]
    cell.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=9, b=True, color=D.T["ink"]), head + ". "),
                              TextBlock(InlineFont(rFont=D.FONT, sz=9, color=D.T["ink2"]), body))
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(f"C{r}:J{r}")
    D.h(ws, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
D.h(ws, r, 12)
n_end = r
D.frame(ws, n_top, n_end, "B", "K")
LAST = n_end + 1
D.h(ws, LAST, 14)
D.paint_canvas(ws, LAST, "Z")
S.finish_sheet(ws, PRODUCT, LAST, span=("A", "L"), freeze="A11", tab_color="accent")

# =====================================================================  Start Here
sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3})
D.page_header(sh, PRODUCT, "Start here. One log to fill in, four settings, nothing to set up.", "Checked Oct 6, 2026",
              span=("B", "F"), side_cols=2)


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            mm = re.match(r"^([A-Z][A-Za-z ]{1,28})\. (.*)$", text, re.S) if kind == "para" else None
            if mm:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


def money(x):
    from decimal import Decimal, ROUND_HALF_UP
    return "${:,}".format(int(Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)))


r = 6
r = sh_card(sh, r, "What it does", [(None,
    "A simple income tracker for freelancers and side hustles. Type one row each time you send an invoice or get paid. "
    "The Dashboard shows what came in this year, month by month, what clients still owe you, your top clients, and how much "
    "to set aside at a rate you choose, split into the IRS estimated tax periods.", "para")])
r = sh_card(sh, r, "Three steps", [
    (1, "Clients tab: type your client names over the examples. Payment methods are there too.", "num"),
    (2, "Income Log tab: type one row per invoice or payment. Leave Paid on blank until the money arrives, then type the date.", "num"),
    (3, "Dashboard tab: set your tax year, set-aside rate and monthly goal. Read your totals at the top.", "num"),
])
r = sh_card(sh, r, "How to read the cells", [
    ((850, USD0), "Yellow cells with blue numbers are yours to change. Type over the examples. Each one shows a hint when you select it.", "chip-in"),
    ((24, INT), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((EX_REC, USD0), "Dark tiles are your answers. They sit at the top of the Dashboard and stay on screen as you scroll.", "chip-key"),
    (None, "Each tab is protected without a password, so you can still change anything. To edit a locked cell, use Review, "
           "Unprotect Sheet in Excel, or Data, Protect sheets and ranges in Google Sheets.", "note"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "The example", [(None,
    f"A freelance designer with six clients logged {len(ROWS)} rows. {money(EX_REC)} arrived in {EX_YEAR} across {EX_MONTHS} months, "
    f"including a December 2025 invoice paid in January, which counts in January. Three invoices with no Paid on date add up to "
    f"{money(EX_OWED)} still owed. At a 25% set-aside rate that is {money(EX_REC * EX_RATE)} to set aside. Clear the example rows "
    "and type your own.", "para")])
r = sh_card(sh, r, "Good to know", [
    (None, "Formulas keep working. Every total reads all 300 log rows, so a row typed anywhere in the log is counted. "
           "Rows to fix on the Income Log counts any row with an amount but no client or date, a date but no amount, or a client "
           "that is not on the Clients tab (for example after a rename). Printing the log prints all 300 rows.", "para"),
    (None, "Dates show the month by name. Date columns accept only real dates and show them as Oct 1, 2026, so you can "
           "see at a glance that each date was read the way you meant.", "para"),
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. No macros, no add-ons, "
           "no sign-up. Made for Excel and Google Sheets; it also opens in LibreOffice.", "para"),
    (None, "Clearing the example. Select the yellow cells in the log and press Delete. Keep the header rows.", "para"),
])
r = sh_card(sh, r, "Before you rely on a number", [(None, S.SHORT_NOTICE.format(date="Oct 6, 2026"), "note")])
SH_LAST = r - 1
D.paint_canvas(sh, SH_LAST, "Z")
S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
                footer_notice="Estimates only. Not professional advice. See the Terms tab or LICENSE-AND-DISCLAIMER.txt.")

# =====================================================================  Terms
tm = wb.create_sheet("Terms")
S.set_widths(tm, {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3})
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version 3. Read this before you rely on a number.", "Checked Oct 6, 2026",
              span=("B", "F"), side_cols=2)
LICENSE = os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt")
paras = open(LICENSE, encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More tools from the shop", [
    (None, ("One-Tab Business Expense Tracker: the expense side, with Schedule C categories", "https://www.etsy.com/listing/4589291725"), "link"),
    (None, ("Hourly Rate Quick Calculator: the rate that pays the take-home you want", "https://www.etsy.com/listing/4589695837"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this tracker saves you time, a review on Etsy helps other freelancers find it. Thank you.", "para"),
    (None, "If a file won't open or something looks wrong, message the shop on Etsy and it will be fixed.", "para"),
])
TM_LAST = r - 1
D.paint_canvas(tm, TM_LAST, "Z")
S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb) if "input without validation" not in p]
if problems:
    print("AUDIT:", *problems, sep="\n  ")
wb.save(OUT)
print("saved", OUT, {"rec": EX_REC, "owed": EX_OWED, "owed_n": EX_OWED_N, "months": EX_MONTHS, "jan": EX_JAN,
                     "rows": {"rec": a_rec, "owed": a_owed, "set": a_set, "keep": a_keep, "months": (m_first, m_last),
                              "periods": p_rows, "top": (t_first, t_last), "owed_list": (o_first, o_last)}})
