#!/usr/bin/env python3
"""Rebuild #20 Hourly Rate Quick Calculator to the ProofNotFluff workbook standard.
Same inputs, same formulas, same worked example as the original; new layout and styling."""
import sys
import os; sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S
from openpyxl.styles import Alignment

PRODUCT = "Hourly Rate Quick Calculator"
AS_OF = "Figures checked Oct 6, 2026"
OUT = sys.argv[1] if len(sys.argv) > 1 else "Hourly-Rate-Quick-Calculator-v2.xlsx"

wb = S.new_workbook(PRODUCT, version="2")

# ---------------------------------------------------------------- Calculator
ws = wb.create_sheet("Calculator")
S.set_widths(ws, {"A": 2, "B": 40, "C": 14, "D": 2, "E": 52, "F": 2})
r = S.title_block(ws, PRODUCT, "Type in the yellow cells. Everything else updates on its own.", as_of=AS_OF)

# Row map (filled as we go so notes can name real cells)
R = {}
card_row = r            # the result card takes rows card_row..card_row+2; it is written once the rows below are known
r = card_row + 4
FREEZE = f"A{r}"


def row_input(r, text, value, fmt, name, validation, prompt, note, prompt_title=None):
    S.label(ws, f"B{r}", text)
    S.input_cell(ws, f"C{r}", value, fmt, name=name, wb=wb, validation=validation, prompt=prompt, prompt_title=prompt_title or text[:32])
    S.note(ws, f"E{r}", note)
    return r + 1


def row_out(r, text, formula, fmt, note, bold=False, subtotal=False):
    if subtotal:
        S.subtotal_label(ws, f"B{r}", text)
    else:
        S.label(ws, f"B{r}", text, bold=bold)
    S.output_cell(ws, f"C{r}", formula, fmt, bold=bold, subtotal=subtotal)
    S.note(ws, f"E{r}", note)
    return r + 1


def row_key(r, text, formula, fmt, note, tone="ink"):
    S.label(ws, f"B{r}", text, bold=True)
    S.key_cell(ws, f"C{r}", formula, fmt, tone=tone)
    S.note(ws, f"E{r}", note)
    return r + 1


# 1. What you want to take home
r = S.section_header(ws, r, "1. What you want to take home")
R["take"] = r
r = row_input(r, "Take-home pay you want per year (after tax)", 60000, S.FMT["usd0"], "TakeHome", ("decimal", 0, None),
              "The pay you want to keep for the year, after tax. Whole dollars, for example 60000.",
              "What you want left in your pocket for the year.", "Take-home pay")
R["tax"] = r
r = row_input(r, "Tax set-aside rate (share of profit)", 0.25, S.FMT["pct"], "TaxRate", ("decimal", 0, 0.9),
              "Type a percent, for example 25%. US self-employment tax alone is 15.3%; income tax comes on top.",
              "US self-employment tax alone is 15.3% (IRS Topic 554, checked Oct 6, 2026); income tax comes on top. 25% is a starting point, not advice. Change it to fit you.",
              "Tax set-aside rate")
R["exp"] = r
r = row_input(r, "Yearly business expenses", 6000, S.FMT["usd0"], "Expenses", ("decimal", 0, None),
              "Everything the business pays for in a year, in whole dollars.",
              "Software, insurance, equipment, phone, fees, mileage, training: everything the business pays for in a year.",
              "Yearly expenses")
r = S.spacer(ws, r)

# 2. Hours you can bill
r = S.section_header(ws, r, "2. Hours you can bill")
R["weeks"] = r
r = row_input(r, "Weeks you work per year", 48, S.FMT["int"], "WeeksPerYear", ("decimal", 1, 52),
              "Weeks you actually work: 52 minus holidays, sick days and time off.",
              "52 minus holidays, sick days and time off.", "Weeks per year")
R["hours"] = r
r = row_input(r, "Hours you work per week", 40, S.FMT["int"], "HoursPerWeek", ("decimal", 1, 100),
              "All hours, including admin, quoting, marketing and travel.",
              "All hours, including admin, quoting, marketing and travel.", "Hours per week")
R["share"] = r
r = row_input(r, "Share of your hours you can bill", 0.6, S.FMT["pct"], "BillableShare", ("decimal", 0.05, 1),
              "Type a percent, for example 60%. Only hours a client pays for.",
              "Only hours a client pays for. Admin, quotes, marketing and travel are usually 30% to 50% of a week.",
              "Billable share")
r = S.spacer(ws, r)

# 3. Your minimum hourly rate
r = S.section_header(ws, r, "3. Your minimum hourly rate")
R["worked"] = r
r = row_out(r, "Hours worked per year", f"=C{R['weeks']}*C{R['hours']}", S.FMT["int"], "Weeks times hours.")
R["billable"] = r
r = row_out(r, "Billable hours per year", f"=C{R['worked']}*C{R['share']}", S.FMT["int"],
            "Hours worked times billable share. Every cost below has to fit into these hours.")
R["profit"] = r
r = row_out(r, "Profit before tax you need", f"=C{R['take']}/(1-C{R['tax']})", S.FMT["usd0"],
            "Take-home divided by (1 minus the tax set-aside rate).")
R["revenue"] = r
r = row_out(r, "Revenue you need for the year", f"=C{R['profit']}+C{R['exp']}", S.FMT["usd0"],
            "Profit before tax plus business expenses.")
R["rate"] = r
r = row_out(r, "Minimum hourly rate", f'=IFERROR(C{R["revenue"]}/C{R["billable"]},"")', S.FMT["usd2"],
            "Revenue needed divided by billable hours. Charge less than this and you will not reach your take-home target.",
            subtotal=True)
R["rounded"] = r
r = row_out(r, "Rounded up to the next $5", f'=IFERROR(CEILING(C{R["rate"]},5),"")', S.FMT["usd0"],
            "A clean number for your rate card.", bold=True)
r = row_out(r, "Of each billable hour: take-home", f'=IFERROR(C{R["take"]}/C{R["billable"]},"")', S.FMT["usd2"],
            "What one billable hour at the minimum rate puts in your pocket.")
r = row_out(r, "Of each billable hour: tax set-aside", f'=IFERROR((C{R["profit"]}-C{R["take"]})/C{R["billable"]},"")', S.FMT["usd2"],
            "What one billable hour sets aside for tax.")
r = row_out(r, "Of each billable hour: expenses", f'=IFERROR(C{R["exp"]}/C{R["billable"]},"")', S.FMT["usd2"],
            "What one billable hour pays toward your yearly expenses.")
r = S.spacer(ws, r)

# the card, now that the rows are known
S.result_card(ws, card_row, "YOUR MINIMUM HOURLY RATE", f"=C{R['rate']}", S.FMT["usd2"],
              sub="Charge less than this and you will not reach your take-home target. Section 3 shows how it is built.",
              side_title="ROUNDED UP TO THE NEXT $5", side_formula=f"=C{R['rounded']}", side_fmt=S.FMT["usd0"])

# 4. Check a rate you are thinking of
r = S.section_header(ws, r, "4. Check a rate you are thinking of")
R["check"] = r
r = row_input(r, "Rate to check", 40, S.FMT["usd0"], "RateToCheck", ("decimal", 0, None),
              "Any hourly rate you are thinking of charging, for example 40.",
              "Type any hourly rate. The rows below show what it really pays you on the hours above.", "Rate to check")
R["crev"] = r
r = row_out(r, "Revenue at that rate", f"=C{R['check']}*C{R['billable']}", S.FMT["usd0"], "Rate times billable hours.")
R["cafter"] = r
r = row_out(r, "After business expenses", f"=C{R['crev']}-C{R['exp']}", S.FMT["usd0"], "Revenue minus your yearly expenses.")
R["ctax"] = r
r = row_out(r, "Tax set-aside at your rate", f"=C{R['cafter']}*C{R['tax']}", S.FMT["usd0"],
            "Expenses are paid before tax, so the set-aside applies to the line above.")
R["ctake"] = r
r = row_out(r, "Take-home at that rate", f"=C{R['cafter']}-C{R['ctax']}", S.FMT["usd0"],
            "After-expenses amount minus the tax set-aside.", subtotal=True)
R["cperhour"] = r
r = row_key(r, "Take-home per hour actually worked", f'=IFERROR(C{R["ctake"]}/C{R["worked"]},"")', S.FMT["usd2"],
            "Take-home divided by ALL hours worked, not just billable ones. This is what the rate really pays you.")
R["gap"] = r
r = row_out(r, "Gap against your take-home target", f"=C{R['ctake']}-C{R['take']}", S.FMT["usd0"],
            f"Negative means the rate is too low for the target in cell C{R['take']}.", bold=True)
S.negative_rule(ws, f"C{R['gap']}")
r = S.spacer(ws, r)

# 5. How your billable share moves the rate
S.page_break_before(ws, r)
r = S.section_header(ws, r, "5. How your billable share moves the rate")
r = S.table_header(ws, r, {"B": "Billable share", "C": "Rate needed"}, span=("B", "C"))
SHARE_NOTES = ["Same take-home, tax and expenses as above.",
               "Billing one more hour a day lowers the rate you need.",
               "It does not lower what you take home."]
for i, share in enumerate((0.5, 0.6, 0.7, 0.8)):
    if i < len(SHARE_NOTES): S.note(ws, f"E{r}", SHARE_NOTES[i])
    c = ws[f"B{r}"]; c.value = share; c.number_format = '0% "of your hours billable"'
    c.font = S.font(S.SIZE["label"]); c.alignment = Alignment(horizontal="left", vertical="center")
    S.output_cell(ws, f"C{r}", f'=IFERROR($C${R["revenue"]}/($C${R["worked"]}*B{r}),"")', S.FMT["usd2"])
    S.set_height(ws, r, S.ROW["line"]); r += 1
r = S.spacer(ws, r)

# 6. Quick quote at your rounded rate
r = S.section_header(ws, r, "6. Quick quote at your rounded rate")
R["jhours"] = r
r = row_input(r, "Estimated hours for the job", 12, S.FMT["num1"], "JobHours", ("decimal", 0, None),
              "Your best estimate of billable hours on this job, for example 12 or 2.5.",
              "Your best estimate of billable hours on this job.", "Hours for the job")
R["jmat"] = r
r = row_input(r, "Materials and other job costs", 150, S.FMT["usd0"], "JobMaterials", ("decimal", 0, None),
              "Parts, supplies, subcontractors, permits and travel for this job, in whole dollars.",
              "Parts, supplies, subcontractors, permits, travel you want the client to cover.", "Job costs")
R["markup"] = r
r = row_input(r, "Markup on materials", 0.15, S.FMT["pct"], "MaterialsMarkup", ("decimal", 0, 0.9),
              "Type a percent, for example 15%. Set 0% to pass costs through.",
              "Covers the time and cash tied up buying and handling materials. Set to 0% to pass costs through.", "Markup on materials")
R["labor"] = r
r = row_out(r, "Labor at your rounded rate", f'=IFERROR(C{R["jhours"]}*C{R["rounded"]},"")', S.FMT["usd2"],
            f"Hours times the rate in cell C{R['rounded']}.")
R["matmk"] = r
r = row_out(r, "Materials with markup", f"=C{R['jmat']}*(1+C{R['markup']})", S.FMT["usd2"], "Materials times (1 plus the markup).")
R["quote"] = r
r = row_key(r, "Quote total", f'=IFERROR(C{R["labor"]}+C{R["matmk"]},"")', S.FMT["usd2"],
            "Labor plus materials. Add sales tax where it applies to your work.")
r = S.spacer(ws, r)

# sources and notice
S.paragraph(ws, f"B{r}", "Sources: IRS Topic 554, Self-Employment Tax (irs.gov/taxtopics/tc554), page updated Sep 24, 2026, checked Oct 6, 2026: "
                         "rate 15.3% (12.4% Social Security, 2.9% Medicare), generally on 92.35% of net earnings. Figures are subject to change.",
            span=("B", "E"), size=S.SIZE["note"], color="note"); r += 1
S.paragraph(ws, f"B{r}", "Information only, not tax, legal or financial advice. See the Terms tab.", span=("B", "E"), size=S.SIZE["note"], color="note"); r += 1
S.paragraph(ws, f"B{r}", "Copyright ProofNotFluff. Personal and single-business use. Information only, not professional advice. "
                         "Full terms on the Terms tab and in LICENSE-AND-DISCLAIMER.txt.", span=("B", "E"), size=S.SIZE["note"], color="note")
last = r
S.finish_sheet(ws, PRODUCT, last, span=("A", "F"), freeze=FREEZE, tab_color="accent")

# ---------------------------------------------------------------- Start Here
S.start_here_tab(
    wb, PRODUCT,
    tagline="One tab. Ten yellow cells. The hourly rate you must charge to take home what you want.",
    what_it_does="Type the take-home pay you want, a tax set-aside rate, your yearly business expenses and the hours you really work. "
                 "The Calculator tab turns them into the minimum hourly rate you must charge, shows what any rate you are considering "
                 "really pays you per hour worked, and turns the rate into a quick quote.",
    steps=[
        "Open the Calculator tab. The example is already filled in, so you can see how every number is made.",
        "Type over the yellow cells with your own numbers: the take-home pay you want, a tax set-aside rate, yearly business expenses, "
        "weeks and hours you work, and the share of your hours you can bill.",
        "Read your minimum hourly rate in the dark card at the top. Section 4 shows what any rate you are considering really pays you "
        "per hour worked. Section 6 turns the rate into a quick quote.",
    ],
    as_of=AS_OF,
    notice=S.SHORT_NOTICE.format(date="Oct 6, 2026"),
    sections=[
        ("What the example shows", [
            "Target take-home $60,000 a year, 25% set aside for tax, $6,000 of expenses, 48 weeks of 40 hours, 60% of hours billable: "
            "1,152 billable hours, so the minimum rate is $74.65 an hour, or $75 rounded up.",
            "Check a $40 rate on the same hours: $46,080 of revenue, $40,080 after expenses, $30,060 after the tax set-aside. "
            "That is $15.66 for every hour actually worked, $29,940 short of the target.",
        ]),
        ("Google Sheets", [
            "Upload Hourly-Rate-Quick-Calculator.xlsx to Google Drive, open it, then File, Save as Google Sheets. Every formula in this "
            "file is a plain formula that works in Excel, Google Sheets, Numbers and LibreOffice. No macros, no add-ons, no links to click.",
        ]),
        ("Any currency", [
            "The maths does not care about the symbol. Type your numbers in your own currency and read the results the same way. "
            "To change the displayed symbol, select the money cells and pick your currency under Format, Number.",
        ]),
        ("Where the tax rate comes from", [
            "Self-employment tax in the US is 15.3% (12.4% Social Security plus 2.9% Medicare), generally applied to 92.35% of net earnings, "
            "per IRS Topic 554 (page updated Sep 24, 2026, checked Oct 6, 2026). Federal and state income tax come on top of that, so the "
            "example uses 25% as a starting point. It is not advice: set the yellow tax cell to match your own situation or ask a tax professional.",
        ]),
        ("Want more than one tab?", [
            "This is the lite version of two bigger ProofNotFluff tools. The Terms tab has the links: the Service Business Pricing "
            "Calculator (quotes, overhead by category, presets for 7 trades) and the Cleaning Business Starter Kit.",
        ]),
    ],
)

# ---------------------------------------------------------------- Terms (last)
S.terms_tab(wb, PRODUCT, [
    ("The full versions of this calculator", [
        ("link", "Service Business Pricing Calculator: hourly rate, overhead by category, job quotes, presets for 7 trades", "https://www.etsy.com/listing/4587962930"),
        ("link", "Cleaning Business Starter Kit: pricing calculator, client intake, checklists", "https://www.etsy.com/listing/4588330888"),
        ("link", "Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"),
    ]),
    ("Terms of Use and disclaimer (Version 2, October 6, 2026; also in LICENSE-AND-DISCLAIMER.txt)", [
        'License. Your purchase gives you a personal license to use this product for yourself or one business you own. You may edit it and print it for that use. You may not resell, share, give away, upload or redistribute it, in original or edited form, or use it to make a competing product.',
        "Ownership. The design, text and formulas belong to ProofNotFluff. Buying the product doesn't transfer ownership.",
        "Information only. This product shares general information and tools for planning. It is not legal, tax, accounting, financial, employment, medical or other professional advice, and it doesn't create a client relationship of any kind. Your situation may differ. For decisions that matter, check with a qualified professional and the official source.",
        "Accuracy. Rates, rules and figures were checked against the sources named in the product on the date above, including the US self-employment tax rate in IRS Topic 554. Laws, prices and platform rules change, and state and local rules vary. You're responsible for confirming anything you rely on.",
        'Estimates. Calculations depend on the numbers you enter and simplify real-world conditions. They are estimates, not promises of any result, income, savings, approval or outcome.',
        'No warranty. The product is provided "as is", without warranties of any kind, to the extent the law allows.',
        "Limit of liability. To the extent the law allows, ProofNotFluff is not liable for any loss or damage from using or relying on this product, and its total liability for any claim is limited to the price you paid for it. Nothing here limits rights you have under consumer protection laws that can't be waived.",
        "Trademarks. Etsy, Microsoft Excel, Google Sheets, Apple Numbers and LibreOffice are trademarks of their owners. ProofNotFluff isn't affiliated with, sponsored or endorsed by them.",
        'AI assistance. This product was made with AI assistance and reviewed and tested by the shop owner.',
        "Help. If a file won't open or something looks wrong, send a message through Etsy and it will be fixed. This is a digital download; no physical item ships.",
    ]),
    ("A small ask", [
        "If this calculator earned its $3, a review on Etsy helps other freelancers find it. Thank you.",
        "Message me on Etsy if anything doesn't open.",
    ]),
])

wb.active = 0
# Keep the "Any currency" section on the second printed page of Start Here instead of orphaning its header.
_sh = wb["Start Here"]
for _r in range(1, _sh.max_row + 1):
    if any(str(_sh.cell(_r, _c).value or "").strip() == "Any currency" for _c in range(2, 6)):
        S.page_break_before(_sh, _r)
        break
problems = S.audit(wb)
if problems:
    print("AUDIT:", *problems, sep="\n  ")
wb.save(OUT)
print("saved", OUT, "rows:", R)
