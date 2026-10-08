#!/usr/bin/env python3
"""Small Business Inventory Tracker with Unit Cost and Profit, version 1 (v3 dashboard design, Oct 8, 2026).

Three working tabs:
  1 Items      one row per product: SKU, unit cost split into materials, labor and packaging,
               price, fee rate, starting stock, reorder point. Unit cost, profit per unit, units
               sold, units on hand, stock value and a plain-words status fill in on their own.
  2 Stock Log  one row per stock move (made or received, sold, customer return, damaged or lost,
               count fixes). Item name, stock change, sale value and profit fill in on their own.
  3 Dashboard  stock value at cost, items to reorder, sales and profit for any date range, the
               reorder list, top sellers, best profit per unit and the money tied up in stock.

Every total reads fixed row ranges (Items rows 16 to 165, Log rows 16 to 2,015), never a whole
column, so nothing on a tab can loop back into its own totals. Formulas sit only in locked white
columns: clearing the yellow example cells never removes a formula (a counted competitor complaint).
No TEXT() format strings, no macros, no tables, no dynamic arrays: same results in Excel and Google Sheets.

Built to design/WORKBOOK_STANDARD.md section 0 with tools/pnf_dash.py and tools/wb_style.py.

  python3 build_xlsx.py Inventory-Tracker-EXAMPLE.xlsx           (the worked example)
  python3 build_xlsx.py Inventory-Tracker.xlsx --blank           (blank)
"""
import datetime as dt
import json
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402

PRODUCT = "Inventory Tracker"
LONG_NAME = "Small Business Inventory Tracker with Unit Cost and Profit"
VERSION = "1"
CHECKED = "Oct 8, 2026"
CHECKED_LONG = "October 8, 2026"
AS_OF = f"Version {VERSION}, {CHECKED}"
HERE = os.path.dirname(os.path.abspath(__file__))
args = sys.argv[1:]
BLANK = "--blank" in args
args = [a for a in args if a != "--blank"]
OUT = args[0] if args else "Inventory-Tracker-EXAMPLE.xlsx"
USD2, INT, DATE = S.FMT["usd2"], S.FMT["int"], S.FMT["date"]
MONEY = '"$"#,##0.00;-"$"#,##0.00;"$"0.00'
QTY = '#,##0.##;-#,##0.##;0'
PCTIN = '0.0%'
PCT = '0.0%;-0.0%;0.0%'

# ---------------------------------------------------------------- stock move types (sign, counts as a sale)
TYPES = [("Made or received", 1), ("Sold", -1), ("Customer return", 1), ("Damaged or lost", -1),
         ("Count fix: add", 1), ("Count fix: remove", -1)]
TYPE_ARR = "{" + ",".join(f'"{t}"' for t, _ in TYPES) + "}"
SIGN_ARR = "{" + ",".join(str(s) for _, s in TYPES) + "}"

# ---------------------------------------------------------------- example: a made-up candle and soap maker
# sku, name, materials, labor, packaging and other, price, fee rate, starting stock, reorder point, notes
ITEMS = [
    ("CN-LAV-8", "Lavender soy candle, 8 oz", 4.20, 2.00, 1.30, 24.00, 0.10, 20, 12, "Wax from the bulk supplier"),
    ("CN-CED-8", "Cedar soy candle, 8 oz", 4.40, 2.00, 1.30, 24.00, 0.10, 18, 12, None),
    ("CN-CIT-16", "Citrus candle, 16 oz", 7.60, 3.00, 1.90, 38.00, 0.10, 10, 6, None),
    ("SP-OAT", "Oat and honey soap bar", 1.80, 1.00, 0.45, 9.00, 0.10, 40, 25, "Shelf B"),
    ("SP-CHAR", "Charcoal soap bar", 2.10, 1.00, 0.45, 9.00, 0.10, 30, 25, "Shelf B"),
    ("SP-ROSE", "Rose clay soap bar", 2.30, 1.00, 0.45, 10.00, 0.10, 30, 20, "Shelf B"),
    ("WM-TRIO", "Wax melt trio", 2.50, 1.00, 0.80, 12.00, 0.10, 25, 15, None),
    ("GS-MINI", "Mini gift set", 9.50, 4.00, 3.50, 45.00, 0.10, 8, 5, "Holiday item"),
    ("MT-TIN", "Travel tin candle", 2.60, 1.50, 0.90, 14.00, 0.10, 30, 15, None),
    ("LB-LIP", "Lip balm", 0.90, 0.50, 0.30, 5.00, 0.10, 60, 30, None),
]
# weekly sales pattern per item (13 weeks, Jul 3 to Sep 25, 2026); batch size; stop restocking after week n
WEEKLY = {
    "CN-LAV-8": ([4, 5, 3, 6, 4, 5, 6, 4, 5, 6, 5, 6, 7], 24, 99),
    "CN-CED-8": ([3, 3, 4, 2, 4, 3, 3, 4, 3, 4, 3, 4, 4], 24, 9),
    "CN-CIT-16": ([1, 2, 1, 2, 2, 1, 2, 2, 1, 2, 2, 2, 3], 12, 99),
    "SP-OAT": ([6, 7, 5, 8, 6, 7, 8, 6, 7, 8, 7, 8, 9], 40, 99),
    "SP-CHAR": ([4, 3, 5, 4, 3, 4, 5, 4, 4, 5, 4, 5, 5], 36, 8),
    "SP-ROSE": ([3, 4, 3, 4, 4, 3, 4, 4, 3, 4, 4, 4, 5], 36, 99),
    "WM-TRIO": ([2, 3, 2, 3, 3, 2, 3, 3, 2, 3, 3, 3, 4], 24, 99),
    "GS-MINI": ([0, 1, 0, 1, 1, 0, 1, 1, 1, 1, 2, 2, 2], 6, 99),
    "MT-TIN": ([3, 2, 3, 3, 2, 3, 3, 2, 3, 3, 3, 3, 4], 24, 99),
    "LB-LIP": ([6, 5, 7, 6, 5, 6, 7, 6, 5, 6, 7, 6, 8], 48, 99),
}
d = dt.datetime
LOG = []   # date, sku, type, qty, price each (None = list price), note
stock = {it[0]: it[7] for it in ITEMS}
reorder = {it[0]: it[8] for it in ITEMS}
pending = {}
for w in range(13):
    friday = d(2026, 7, 3) + dt.timedelta(days=7 * w)
    monday = friday - dt.timedelta(days=4)
    for sku, (sales, batch, stop) in WEEKLY.items():
        if pending.get(sku):
            LOG.append((monday, sku, "Made or received", batch, None, "New batch"))
            stock[sku] += batch; pending[sku] = False
        n = min(sales[w], stock[sku])
        if n:
            LOG.append((friday, sku, "Sold", n, None, None))
            stock[sku] -= n
        if stock[sku] <= reorder[sku] and w < stop:
            pending[sku] = True
LOG += [
    (d(2026, 8, 12), "CN-LAV-8", "Customer return", 1, None, "Arrived chipped, refunded"),
    (d(2026, 8, 12), "CN-LAV-8", "Damaged or lost", 1, None, "The returned candle, not resellable"),
    (d(2026, 9, 3), "SP-ROSE", "Damaged or lost", 2, None, "Cracked in the mould"),
    (d(2026, 9, 18), "SP-OAT", "Made or received", 12, None, "Extra batch for the wholesale order"),
    (d(2026, 9, 18), "SP-OAT", "Sold", 12, 6.00, "Wholesale order, $6 each"),
    (d(2026, 9, 30), "LB-LIP", "Count fix: remove", 1, None, "Month-end count was 1 short"),
]
LOG.sort(key=lambda x: (x[0], x[1]))


def expected():
    it = {i[0]: i for i in ITEMS}
    sign = dict(TYPES)
    out = {}
    for sku, name, m, l, p, price, fee, start, rp, _ in ITEMS:
        uc = round(m + l + p, 2)
        sold = sum(r[3] for r in LOG if r[1] == sku and r[2] == "Sold") - sum(r[3] for r in LOG if r[1] == sku and r[2] == "Customer return")
        oh = start + sum(r[3] * sign[r[2]] for r in LOG if r[1] == sku)
        out[sku] = dict(name=name, uc=uc, ppu=round(price - uc - price * fee, 4), sold=sold, oh=oh, val=round(max(0, oh) * uc, 2),
                        status="Out of stock" if oh <= 0 else ("Reorder now" if oh <= rp else "OK"))
    value = sales = profit = 0.0
    for dte, sku, typ, q, pr, _ in LOG:
        if typ in ("Sold", "Customer return"):
            s = 1 if typ == "Sold" else -1
            i = it[sku]; p = pr if pr is not None else i[5]
            v = s * q * p
            value += v
            profit += v - s * q * out[sku]["uc"] - v * i[6]
    tot = dict(stock_value=round(sum(o["val"] for o in out.values()), 2), on_hand=sum(o["oh"] for o in out.values()),
               reorder=sum(1 for o in out.values() if o["status"] != "OK"), sales=round(value, 2), profit=round(profit, 2),
               units_sold=sum(o["sold"] for o in out.values()))
    return out, tot


EX_ITEMS, EX = expected()
if BLANK:
    ITEMS, LOG = [], []

# ---------------------------------------------------------------- fixed addresses
IT_N, LG_N, DB_N = "1 Items", "2 Stock Log", "3 Dashboard"
IT, LG, DB = (f"'{n}'" for n in (IT_N, LG_N, DB_N))
SHIP_LOG = 2000  # rows in the shipped file (shot copies may be shorter; labels always say this)
N_ITEMS, N_LOG = 150, int(os.environ.get("PNF_LOG_ROWS", "2000"))  # PNF_LOG_ROWS: short copies for listing shots only
I_TOP, L_TOP = 11, 11
I_FIRST, L_FIRST = I_TOP + 5, L_TOP + 5
I_LAST, L_LAST = I_FIRST + N_ITEMS - 1, L_FIRST + N_LOG - 1


def irng(col):
    return f"{IT}!${col}${I_FIRST}:${col}${I_LAST}"


def lrng(col):
    return f"{LG}!${col}${L_FIRST}:${col}${L_LAST}"


SKUS, NAMES, PRICES, FEES, UCOST = irng("C"), irng("D"), irng("H"), irng("I"), irng("L")
STATUS, VALUE = irng("Q"), irng("P")


def sign_rule(ws, ref):
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{ref}<0"], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{ref}>0"], font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))


def rich(head, body, size=9):
    return CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=size, b=True, color=D.T["ink"]), head + ". "),
                        TextBlock(InlineFont(rFont=D.FONT, sz=size, color=D.T["ink2"]), body))


wb = S.new_workbook(PRODUCT, version=VERSION)

# =====================================================================  1 Items
it = wb.create_sheet(IT_N)
S.set_widths(it, {"A": 3, "B": 2, "C": 12, "D": 28, "E": 11, "F": 10, "G": 11, "H": 10, "I": 9, "J": 10, "K": 10,
                  "L": 10, "M": 11, "N": 10, "O": 10, "P": 13, "Q": 16, "R": 26, "S": 2, "T": 3})
D.page_header(it, "Items", "One row per product. Unit cost, profit per unit, stock on hand and status fill in on their own.",
              "Example rows: type over them" if not BLANK else f"Type your first item in row {I_FIRST}", span=("B", "S"), side_cols=4)
D.h(it, 6, 10)
D.stat_strip(it, 7, [
    ("C:D", "Items on the list", f'=COUNTIF({SKUS},"?*")', INT),
    ("E:H", "Stock value at cost", f"=SUM({VALUE})", USD2),
    ("I:L", "Units on hand", f"=SUM({irng('O')})", QTY),
    ("M:P", "Items to reorder", f'=COUNTIF({STATUS},"Reorder now")+COUNTIF({STATUS},"Out of stock")+COUNTIF({STATUS},"Below zero*")', INT),
    ("Q:R", "Rows to fix", "=0", INT),
])
D.h(it, 9, 10)
D.frame(it, 6, 9, "B", "S")
D.status_rule(it, "M8", "$M$8>0", fill_tint=False)
D.status_rule(it, "Q8", "$Q$8>0", fill_tint=False)
D.h(it, 10, 14)


def L_ref(col, r):
    return f"{LG}!${col}${L_FIRST}:${col}${L_LAST}"


def f_unit_cost(r):
    return f'=IF($C{r}="","",E{r}+F{r}+G{r})'


def f_ppu(r):
    return f'=IF(OR($C{r}="",H{r}=""),"",H{r}-L{r}-H{r}*I{r})'


def f_sold(r):
    return (f'=IF($C{r}="","",SUMIFS({lrng("G")},{lrng("D")},$C{r},{lrng("F")},"Sold")'
            f'-SUMIFS({lrng("G")},{lrng("D")},$C{r},{lrng("F")},"Customer return"))')


def f_on_hand(r):
    return f'=IF($C{r}="","",J{r}+SUMIFS({lrng("I")},{lrng("D")},$C{r}))'


def f_value(r):
    return f'=IF($C{r}="","",MAX(0,O{r})*L{r})'


def f_status(r):
    return (f'=IF($C{r}="",IF(COUNTA(D{r}:K{r})>0,"Needs a SKU",""),IF(COUNTIF({SKUS},$C{r})>1,"Duplicate SKU",'
            f'IF(O{r}<0,"Below zero: check log",IF(O{r}=0,"Out of stock",IF(AND(K{r}<>"",O{r}<=K{r}),"Reorder now","OK")))))')


def col_vals(i):
    return [x[i] for x in ITEMS]


money_v = ("decimal", 0, 1000000)
cols_it = [
    dict(col="C", head="SKU", kind="input", values=col_vals(0), validation=("textLength", 1, 20),
         prompt="A short code that is different for every item, for example CN-LAV-8. The Stock Log uses it.",
         error="A SKU is 1 to 20 characters. Each item needs its own."),
    dict(col="D", head="Item name", kind="input", values=col_vals(1), validation=("textLength", 0, 60),
         prompt="What the item is called, for example Lavender soy candle, 8 oz.", error="Up to 60 characters."),
    dict(col="E", head="Materials per unit", kind="input", fmt=USD2, align="right", values=col_vals(2), validation=money_v,
         prompt="What the materials in one unit cost you, for example 4.20.", error="Type an amount from 0 to 1,000,000."),
    dict(col="F", head="Labor per unit", kind="input", fmt=USD2, align="right", values=col_vals(3), validation=money_v,
         prompt="What making one unit costs in labor, yours or a helper's, for example 2.00. Type 0 to leave it out.",
         error="Type an amount from 0 to 1,000,000."),
    dict(col="G", head="Packaging and other", kind="input", fmt=USD2, align="right", values=col_vals(4), validation=money_v,
         prompt="Box, label, insert and any other cost per unit, for example 1.30.", error="Type an amount from 0 to 1,000,000."),
    dict(col="H", head="Selling price", kind="input", fmt=USD2, align="right", values=col_vals(5), validation=money_v,
         prompt="Your usual price for one unit, for example 24.00. A different price on one sale goes in the Stock Log.",
         error="Type an amount from 0 to 1,000,000."),
    dict(col="I", head="Fees, % of price", kind="input", fmt=PCTIN, align="right", values=col_vals(6), validation=("decimal", 0, 1),
         prompt="Your selling and payment fees as a share of the price, for example 10%. Type 0% if none.",
         error="Type a percent from 0% to 100%, for example 10%."),
    dict(col="J", head="Starting stock", kind="input", fmt=QTY, align="right", values=col_vals(7), validation=("decimal", 0, 10000000),
         prompt="Units on hand on the day you start logging, for example 20.", error="Type 0 or more units."),
    dict(col="K", head="Reorder at", kind="input", fmt=QTY, align="right", values=col_vals(8), validation=("decimal", 0, 10000000),
         prompt="When units on hand fall to this number, the status says Reorder now. For example 12. Leave blank for no alert.",
         error="Type 0 or more units, or leave blank."),
    dict(col="L", head="Unit cost", kind="calc", fmt=USD2, align="right", formula=f_unit_cost),
    dict(col="M", head="Profit per unit", kind="calc", fmt=USD2, align="right", formula=f_ppu),
    dict(col="N", head="Units sold", kind="calc", fmt=QTY, align="right", formula=f_sold),
    dict(col="O", head="On hand", kind="calc", fmt=QTY, align="right", formula=f_on_hand, bold=True),
    dict(col="P", head="Stock value", kind="calc", fmt=USD2, align="right", formula=f_value),
    dict(col="Q", head="Status", kind="calc", formula=f_status),
    dict(col="R", head="Your notes (supplier, shelf)", kind="input", values=col_vals(9), validation=("textLength", 0, 120),
         prompt="Anything you want to keep with the item: supplier, shelf, colour. Optional.", error="Up to 120 characters."),
]
f1, f2, it_end = D.table_card(it, I_TOP, "B", "S", cols_it, N_ITEMS, title="Your items",
                              sub=f"Yellow columns are yours; white columns are formulas and stay put when you clear the yellow cells. "
                                  f"{N_ITEMS} rows.")
assert (f1, f2) == (I_FIRST, I_LAST), (f1, f2)
for rr in range(I_FIRST, I_LAST + 1):
    for k in ("C", "D", "R"):
        it[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    it[f"Q{rr}"].font = D.f(9, True, "ink2"); it[f"Q{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
it.conditional_formatting.add(f"M{I_FIRST}:M{I_LAST}", FormulaRule(formula=[f'AND(M{I_FIRST}<>"",M{I_FIRST}<0)'],
                              font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
for word, color in (("Reorder", "accent"), ("Out of", "accent"), ("Below", "accent"), ("Duplicate", "accent"), ("Needs", "accent"), ("OK", "teal")):
    it.conditional_formatting.add(f"Q{I_FIRST}:Q{I_LAST}", FormulaRule(
        formula=[f'LEFT($Q{I_FIRST},{len(word)})="{word}"'], font=Font(color=D.T[color], bold=True),
        fill=PatternFill("solid", bgColor=D.T["accent_tint"]) if color == "accent" else None, stopIfTrue=True))
D.paint_canvas(it, it_end + 1, "Z")
S.finish_sheet(it, PRODUCT, it_end + 1, span=("A", "T"), freeze=f"D{I_FIRST}", tab_color="gold")
it.page_setup.orientation = "landscape"
it.print_title_rows = f"{I_FIRST - 1}:{I_FIRST - 1}"
it.print_area = f"A1:T{I_FIRST + 29}"
it[f"Q{I_FIRST - 1}"].alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True, indent=1)

# =====================================================================  2 Stock Log
lg = wb.create_sheet(LG_N)
S.set_widths(lg, {"A": 3, "B": 2, "C": 13, "D": 12, "E": 28, "F": 20, "G": 10, "H": 12, "I": 10, "J": 12, "K": 12,
                  "L": 32, "M": 18, "N": 2, "O": 3})
D.page_header(lg, "Stock log", "One row per stock move. Quantities are positive; the type decides if stock goes up or down.",
              "Example rows: type over them" if not BLANK else f"Type your first move in row {L_FIRST}", span=("B", "N"), side_cols=3)
D.h(lg, 6, 10)
D.stat_strip(lg, 7, [
    ("C:D", "Moves logged", f'=COUNT({lrng("G")})', INT),
    ("E:E", "Units sold, net", f"=SUM({irng('N')})", QTY),
    ("F:H", "Sales value", f'=SUM({lrng("J")})', USD2),
    ("I:K", "Profit from sales", f'=SUM({lrng("K")})', USD2),
    ("L:M", "Rows to fix", "=0", INT),
])
D.h(lg, 9, 10)
D.frame(lg, 6, 9, "B", "N")
sign_rule(lg, "I8")
D.status_rule(lg, "L8", "$L$8>0", fill_tint=False)
D.h(lg, 10, 14)


def look(col_rng, r):
    return f"INDEX({col_rng},MATCH($D{r},{SKUS},0))"


def f_item(r):
    return f'=IF($D{r}="","",IFERROR({look(NAMES, r)},"Not on Items"))'


def f_change(r):
    return f'=IF(OR($D{r}="",$G{r}=""),"",$G{r}*IFERROR(INDEX({SIGN_ARR},MATCH($F{r},{TYPE_ARR},0)),0))'


def f_sale(r):
    return (f'=IF(OR($G{r}="",AND($F{r}<>"Sold",$F{r}<>"Customer return")),"",IFERROR(IF($F{r}="Sold",1,-1)*$G{r}*'
            f'IF($H{r}="",{look(PRICES, r)},$H{r}),""))')


def f_profit(r):
    return (f'=IF(J{r}="","",IFERROR(J{r}-IF($F{r}="Sold",1,-1)*$G{r}*{look(UCOST, r)}-J{r}*{look(FEES, r)},""))')


def f_check(r):
    return (f'=IF(COUNTA($C{r}:$D{r},$F{r}:$H{r})=0,"",IF($C{r}="","Needs a date",IF($D{r}="","Needs a SKU",'
            f'IF(ISNA(MATCH($D{r},{SKUS},0)),"SKU not on Items",IF($F{r}="","Needs a type",IF($G{r}="","Needs a quantity","OK"))))))')


cols_lg = [
    dict(col="C", head="Date", kind="input", fmt=DATE, values=[x[0] for x in LOG], validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The day the stock moved, for example 9/18/2026.", error="Type a real date between 2000 and 2100, for example 9/18/2026."),
    dict(col="D", head="SKU (pick)", kind="input", values=[x[1] for x in LOG], validation=("list_range", SKUS),
         prompt="Pick the item's SKU from the list. New items go on the Items tab first.",
         error="Pick a SKU from the list. Add new items on the Items tab first."),
    dict(col="E", head="Item", kind="calc", formula=f_item),
    dict(col="F", head="Type (pick)", kind="input", values=[x[2] for x in LOG], validation=("list", [t for t, _ in TYPES]),
         prompt="Made or received, Sold and Customer return change stock the way they say. Count fixes match your shelf count.",
         error="Pick a type from the list."),
    dict(col="G", head="Quantity", kind="input", fmt=QTY, align="right", values=[x[3] for x in LOG], validation=("decimal", 0, 10000000),
         prompt="How many units, as a positive number, for example 5. The type decides if stock goes up or down.",
         error="Type 0 or more units, as a positive number."),
    dict(col="H", head="Price each (optional)", kind="input", fmt=USD2, align="right", values=[x[4] for x in LOG], validation=money_v,
         prompt="Only for a sale at a different price, for example 6.00 for a wholesale order. Blank uses the price on the Items tab.",
         error="Type an amount from 0 to 1,000,000, or leave blank."),
    dict(col="I", head="Stock change", kind="calc", fmt=QTY, align="right", formula=f_change),
    dict(col="J", head="Sale value", kind="calc", fmt=USD2, align="right", formula=f_sale),
    dict(col="K", head="Profit", kind="calc", fmt=USD2, align="right", formula=f_profit),
    dict(col="L", head="Note", kind="input", values=[x[5] for x in LOG], validation=("textLength", 0, 120),
         prompt="Anything worth remembering, for example the order number. Optional.", error="Up to 120 characters."),
    dict(col="M", head="Check", kind="calc", formula=f_check),
]
g1, g2, lg_end = D.table_card(lg, L_TOP, "B", "N", cols_lg, N_LOG, title="Your stock moves",
                              sub=f"Yellow columns are yours; white columns fill in on their own. {SHIP_LOG:,} rows, all read by the totals.")
assert (g1, g2) == (L_FIRST, L_LAST), (g1, g2)
for rr in range(L_FIRST, L_LAST + 1):
    for k in ("D", "F", "L"):
        lg[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for k in ("E", "M"):
        lg[f"{k}{rr}"].font = D.f(9, False, "ink2"); lg[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
lg.conditional_formatting.add(f"I{L_FIRST}:I{L_LAST}", FormulaRule(formula=[f'AND(I{L_FIRST}<>"",I{L_FIRST}<0)'],
                              font=Font(color=D.T["accent"]), stopIfTrue=True))
for word, color in (("Needs", "accent"), ("SKU not", "accent"), ("Not on", "accent"), ("OK", "muted")):
    lg.conditional_formatting.add(f"M{L_FIRST}:M{L_LAST}", FormulaRule(
        formula=[f'LEFT($M{L_FIRST},{len(word)})="{word}"'], font=Font(color=D.T[color], bold=color == "accent"), stopIfTrue=True))
    if word == "Not on":
        lg.conditional_formatting.add(f"E{L_FIRST}:E{L_LAST}", FormulaRule(
            formula=[f'$E{L_FIRST}="Not on Items"'], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
D.paint_canvas(lg, lg_end + 1, "Z")
S.finish_sheet(lg, PRODUCT, lg_end + 1, span=("A", "O"), freeze=f"A{L_FIRST}", tab_color="blue")
lg.page_setup.orientation = "landscape"
lg.print_title_rows = f"{L_FIRST - 1}:{L_FIRST - 1}"
lg.print_area = f"A1:O{L_FIRST + 99}"
lg[f"L{L_FIRST - 1}"].alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True, indent=1)

# =====================================================================  3 Dashboard
db = wb.create_sheet(DB_N)
S.set_widths(db, {"A": 3, "B": 2, "C": 31, "D": 14, "E": 2, "F": 3, "G": 2, "H": 24, "I": 12, "J": 17, "K": 2, "L": 3})
D.page_header(db, "Inventory dashboard", "Reads every row of the Items tab and the Stock Log. Nothing to type here except the dates.",
              AS_OF, span=("B", "K"), side_cols=3)
H_TOP = 200  # helper block for the sorted lists, below the printed page
# helper block rows (one per Items row)
HR = lambda i: H_TOP + 3 + i  # noqa: E731
H_FIRST, H_LAST = HR(0), HR(N_ITEMS - 1)
KEY_SOLD, KEY_PPU, KEY_VAL, SEQ = (f"${c}${H_FIRST}:${c}${H_LAST}" for c in ("D", "H", "I", "C"))

# ---- left column: sales card
C_TOP = 11
A = D.Card(db, C_TOP, "B", "C", "D", "E", wb=wb)
A.title("Sales in the log", "Leave both dates blank to count every row.")
r_from = A.input("From (optional)", None, DATE, name="FromDate", validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
                 allow_blank=True, prompt_title="From date (optional)",
                 prompt="First day to count, for example 9/1/2026. Leave blank to start from the first row.")
r_to = A.input("To (optional)", None, DATE, name="ToDate", validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
               allow_blank=True, prompt_title="To date (optional)",
               prompt="Last day to count, for example 9/30/2026. Leave blank to count up to the last row.")
FROM = f'IF($D${r_from}="",0,$D${r_from})'
TO = f'IF($D${r_to}="",2958465,$D${r_to})'
DT = f'{lrng("C")},">="&{FROM},{lrng("C")},"<="&{TO}'
A.eyebrow("WHAT SOLD")
r_units = A.calc("Units sold, net of returns",
                 f'=SUMIFS({lrng("G")},{lrng("F")},"Sold",{DT})-SUMIFS({lrng("G")},{lrng("F")},"Customer return",{DT})', QTY)
r_sales = A.calc("Sales value", f'=SUMIFS({lrng("J")},{DT})', USD2)
r_cost = A.calc("Unit costs and fees", f"=D{r_sales}-D{r_sales + 2}", USD2)
r_prof = A.calc("Profit from sales", f'=SUMIFS({lrng("K")},{DT})', USD2, total=True, tint="total_tint")
r_marg = A.calc("Profit margin", f'=IF(D{r_sales}=0,"",D{r_prof}/D{r_sales})', PCT, color="ink2", label_color="ink2")
r_dpill = A.text(f'=IF(AND($D${r_from}<>"",$D${r_to}<>"",$D${r_to}<$D${r_from}),"The To date is before the From date",'
                 f'IF(AND($D${r_from}="",$D${r_to}=""),"Counting every row in the log","Counting the dates you picked"))',
                 size=9, color="ink2", bold=True)
A.close()
assert r_sales + 2 == r_prof
sign_rule(db, f"D{r_prof}")
D.status_rule(db, f"C{r_dpill}", f'AND($D${r_from}<>"",$D${r_to}<>"",$D${r_to}<$D${r_from})', fill_tint=False)

# ---- left column: reorder list
N_RE = 8
B = D.Card(db, A.end + 2, "B", "C", "D", "E", wb=wb)
B.title("Reorder now", "Items at or under their reorder point. Units on hand at the right.")
RE_ROWS = []
for k in range(1, N_RE + 1):
    idx = f"MATCH({k},{SEQ},0)"
    rr = B.calc(f'=IFERROR(INDEX({NAMES},{idx})&"  (reorder at "&INDEX({irng("K")},{idx})&")","")',
                f'=IFERROR(INDEX({irng("O")},{idx}),"")', QTY)
    RE_ROWS.append(rr)
N_FLAG = f"MAX(0,MAX({SEQ}))"
B.text(f'=IF({N_FLAG}=0,IF(COUNTIF({SKUS},"?*")=0,"Add items on the Items tab","Every item is above its reorder point"),'
       f'IF({N_FLAG}>{N_RE},"Plus "&({N_FLAG}-{N_RE})&" more on the Items tab (Status column)",'
       f'{N_FLAG}&IF({N_FLAG}=1," item needs"," items need")&" restocking"))', size=9, color="ink2", bold=True)
B.close()
for rr in RE_ROWS:
    db.conditional_formatting.add(f"D{rr}", FormulaRule(formula=[f'AND(D{rr}<>"",D{rr}<=0)'], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
LEFT_END = B.end

# ---- right column: three top-5 lists with bars


def top_card(top, title, sub, key, val_col, fmt, color, empty):
    K = D.Card(db, top, "G", "H", "I", "K", extra="J", wb=wb)
    K.title(title, sub)
    first = K.r
    rows = []
    for k in range(1, 6):
        idx = f"MATCH(LARGE({key},{k}),{key},0)"
        show = f'AND(COUNT({key})>={k},IFERROR(INDEX({irng(val_col)},{idx}),0)>0)'
        rr = K.bar(f'=IF({show},INDEX({NAMES},{idx}),"")', f'=IF({show},INDEX({irng(val_col)},{idx}),"")', fmt,
                   f"$I${first}", color)
        rows.append(rr)
    K.text(f'=IF(I{first}="","{empty}","")', size=9, color="ink2")
    K.close()
    return K, rows


R1, top_rows = top_card(C_TOP, "Top sellers", "Units sold, net of returns, every row in the log.", KEY_SOLD, "N", QTY, "teal",
                        "No sales logged yet")
R2, ppu_rows = top_card(R1.end + 2, "Best profit per unit", "Price less unit cost and fees, from the Items tab.", KEY_PPU, "M", USD2, "gold",
                        "Add prices on the Items tab")
R3, val_rows = top_card(R2.end + 2, "Most money in stock", "Units on hand times unit cost.", KEY_VAL, "P", USD2, "blue",
                        "No stock on hand yet")
RIGHT_END = R3.end
CARDS_END = max(LEFT_END, RIGHT_END)

# ---- tiles (rows 6 to 9)
D.tile(db, 6, "B", "E", "STOCK VALUE AT COST", f"=SUM({VALUE})", USD2,
       f'=IF(COUNTIF({SKUS},"?*")=0,"Add your items on the Items tab",SUM({irng("O")})&" units on hand across "&COUNTIF({SKUS},"?*")'
       f'&IF(COUNTIF({SKUS},"?*")=1," item"," items"))', dark=True)
REORDER_N = f'(COUNTIF({STATUS},"Reorder now")+COUNTIF({STATUS},"Out of stock")+COUNTIF({STATUS},"Below zero*"))'
v2, sub2, _ = D.tile(db, 6, "G", "K", "ITEMS TO REORDER", f"={REORDER_N}", INT,
                     f'=IF(COUNTIF({SKUS},"?*")=0,"Nothing to check yet",IF({REORDER_N}=0,"Every item is above its reorder point",'
                     f'"Start with "&INDEX({NAMES},MATCH(1,{SEQ},0))))', dark=False)
D.status_rule(db, v2.coordinate, f"{v2.coordinate}>0", fill_tint=False)
D.h(db, 10, 14)

# ---- checks (the Items and Stock Log strips read these two cells)
CK_TOP = CARDS_END + 2
K = D.Card(db, CK_TOP, "B", "C", "D", "E", wb=wb)
K.title("Checks", "Both should read 0.")
k1 = K.calc("Item rows to fix", f'=COUNTIF({STATUS},"Needs*")+COUNTIF({STATUS},"Duplicate*")+COUNTIF({STATUS},"Below*")', INT)
k2 = K.calc("Log rows to fix", f'=COUNTIF({lrng("M")},"Needs*")+COUNTIF({lrng("M")},"SKU not*")', INT)
K.close()
for rr in (k1, k2):
    D.status_rule(db, f"D{rr}", f"$D${rr}>0", fill_tint=False)
NOTE_RIGHT = D.Card(db, CK_TOP, "G", "H", "I", "K", extra="J", wb=wb)
NOTE_RIGHT.title("Clearing the example", "Formulas live only in white cells.")
NOTE_RIGHT.text("Select the yellow cells on Items and the Stock Log and press Delete. Every formula stays.", size=9, color="ink2")
NOTE_RIGHT.close()

# ---- notes and sources
N_TOP = max(K.end, NOTE_RIGHT.end) + 2
NOTES = [
    ("Unit cost", "Materials plus labor plus packaging and other costs, per unit, as you type them on the Items tab. "
     "Profit per unit is the selling price less unit cost less fees (price times your fee rate)."),
    ("Stock on hand", "Starting stock plus every Made or received, Customer return and Count fix: add row, less every Sold, "
     "Damaged or lost and Count fix: remove row for that SKU."),
    ("Sales and profit", "A Sold row counts its quantity at the Items price, or at the price you type on that row. A Customer "
     "return reverses one sale at the same price: the sale value, unit cost and fees come back out of the totals."),
    ("Reorder point", "You set it per item. A common rule of thumb is the units you sell while waiting for a restock, plus a "
     "few spare. The status flags the item when units on hand reach that number."),
    ("Any currency", "The maths works in any currency. Select the money cells and pick your symbol under Format, Number."),
    ("Before you rely on a number", S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab")),
]
r = N_TOP
D.h(db, r, 12); r += 1
t = db[f"C{r}"]; t.value = "Notes"; t.font = D.f(12, True); db.merge_cells(f"C{r}:J{r}"); D.h(db, r, 24); r += 1
width = sum(S.width_of(db, k) for k in D.cols("C", "J"))
for head, body in NOTES:
    cell = db[f"C{r}"]; cell.value = rich(head, body)
    cell.alignment = Alignment(vertical="top", wrap_text=True); db.merge_cells(f"C{r}:J{r}")
    D.h(db, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
D.h(db, r, 12)
D.frame(db, N_TOP, r, "B", "K")
DB_LAST = r + 1
D.h(db, DB_LAST, 14)
assert DB_LAST < H_TOP - 2, DB_LAST

# ---- helper block (below the printed page): sort keys for the lists above
r = H_TOP
c = db[f"C{r}"]; c.value = "Sorting helpers for the lists above. Formulas only, nothing to type."
c.font = D.f(9, True, "muted"); db.merge_cells(f"C{r}:J{r}")
r += 2
for k, txt in (("C", "REORDER ORDER"), ("D", "UNITS SOLD KEY"), ("H", "PROFIT PER UNIT KEY"), ("I", "STOCK VALUE KEY")):
    cc = db[f"{k}{r}"]; cc.value = txt; cc.font = D.f(8, True, "muted"); cc.alignment = Alignment(horizontal="right")
for i in range(N_ITEMS):
    hr, ir = HR(i), I_FIRST + i
    tie = f"{N_ITEMS - i}/10000000"
    db[f"C{hr}"] = (f'=IF(OR({IT}!$Q${ir}="Reorder now",{IT}!$Q${ir}="Out of stock",LEFT({IT}!$Q${ir},10)="Below zero"),'
                    f'MAX($C${H_TOP + 2}:C{hr - 1})+1,"")')
    db[f"D{hr}"] = f'=IF({IT}!$C${ir}="","",N({IT}!$N${ir})+{tie})'
    db[f"H{hr}"] = f'=IF(OR({IT}!$C${ir}="",{IT}!$M${ir}=""),"",{IT}!$M${ir}+{tie})'
    db[f"I{hr}"] = f'=IF({IT}!$C${ir}="","",N({IT}!$P${ir})+{tie})'
    for k in ("C", "D", "H", "I"):
        db[f"{k}{hr}"].font = D.f(8, False, "muted"); db[f"{k}{hr}"].number_format = "0.000"
for rr in range(1, H_LAST + 1):
    if db.row_dimensions[rr].height is None:
        D.h(db, rr, D.GRID_ROW if rr < H_TOP else 14)
D.paint_canvas(db, H_LAST, "Z")
S.finish_sheet(db, PRODUCT, DB_LAST, span=("A", "L"), freeze="A11", tab_color="accent")
db.print_area = f"A1:L{DB_LAST}"
S.page_break_before(db, CK_TOP)
# print: tiles and all cards on page 1 (break before Checks)
_pt = sum(db.row_dimensions[rr].height or D.GRID_ROW for rr in range(1, CK_TOP))
db.sheet_properties.pageSetUpPr.fitToPage = False
db.page_setup.scale = max(40, min(88, int(680 / _pt * 100)))

# the Items and Stock Log strips show the Checks total
for ws, cell in ((it, "Q8"), (lg, "L8")):
    ws[cell].value = f"={DB}!$D${k1}+{DB}!$D${k2}"

# =====================================================================  Start Here and Terms
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            mm = re.match(r"^([A-Z][A-Za-z ,]{1,34})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if mm:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


def money(x):
    return ("-" if x < 0 else "") + f"${abs(x):,.2f}"


lav = EX_ITEMS["CN-LAV-8"]
reorder_names = [v["name"] for v in EX_ITEMS.values() if v["status"] != "OK"]
sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, "Inventory Tracker", "Start here. List your items once, log each stock move, and read the dashboard.", f"Version {VERSION}, {CHECKED}",
              span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "An inventory tracker for small product businesses: makers, resellers and small shops. List each item once with what one "
    "unit costs you in materials, labor and packaging, your price and your fee rate. Log every stock move on one tab. The "
    "workbook keeps units on hand, flags items to reorder at the point you set, and shows the stock value at cost, the profit "
    "per unit and the profit from your sales.", "para")])
r = sh_card(sh, r, "Three steps", [
    (1, "Items tab. One row per product: SKU, name, the three unit costs, price, fee rate, starting stock and your reorder point.", "num"),
    (2, "Stock Log tab. One row per move: date, SKU from the list, type, quantity. Sold, made, returned, damaged or a count fix.", "num"),
    (3, "Dashboard tab. Read the stock value and the reorder count in the tiles, then the reorder list and the top sellers.", "num"),
])
r = sh_card(sh, r, "How to read the cells", [
    ((24, USD2), "Yellow cells with blue numbers are yours to type in. Each one shows a hint when you select it.", "chip-in"),
    ((lav["uc"], USD2), "Numbers on white are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((EX["stock_value"], USD2),
     "Dark tiles are your answers. The main one sits at the top of the Dashboard and stays on screen as you scroll.", "chip-key"),
    (None, "Each tab is protected without a password, so you can still change anything. To edit a locked cell, use Review, "
           "Unprotect Sheet in Excel, or Data, Protect sheets and ranges in Google Sheets.", "note"),
])
S.page_break_before(sh, r)
ex_text = (f"Inventory-Tracker-EXAMPLE.xlsx holds July to September 2026 of a made-up candle and soap maker with 10 items. "
           f"The lavender candle costs {money(lav['uc'])} a unit to make ($4.20 materials, $2.00 labor, $1.30 packaging) and sells "
           f"for $24.00, so after 10% fees it makes {money(lav['ppu'])} a unit. Across all 10 items the stock on hand is worth "
           f"{money(EX['stock_value'])} at cost, {EX['units_sold']:,} units sold for {money(EX['sales'])} and "
           f"{money(EX['profit'])} of profit, and {EX['reorder']} items are at or under their reorder point: "
           + "; ".join(reorder_names) + ".")
r = sh_card(sh, r, "The example", [(None, ex_text, "para")])
r = sh_card(sh, r, "Which file to open", [
    (None, "Excel on a computer. Double-click Inventory-Tracker.xlsx. If Excel shows a yellow Protected View bar, click Enable Editing.", "para"),
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. No macros, no add-ons, no sign-up.", "para"),
    (None, "Mac, iPad or phone. Unzip on a computer first (phones often can't open a .zip), then upload the .xlsx to Google Drive "
           "and open it in the Google Sheets app, or open it in Excel for Mac or iPad.", "para"),
    (None, "See it filled in first. Open Inventory-Tracker-EXAMPLE.xlsx, then use the blank Inventory-Tracker.xlsx for your own items.", "para"),
])
r = sh_card(sh, r, "Good to know", [
    (None, "Clearing the example. In the example file, select the yellow cells on Items and the Stock Log and press Delete. "
           "Formulas live only in the white columns, so every total keeps working. Or start from the blank file.", "para"),
    (None, f"Room to grow. {N_ITEMS} items and {N_LOG:,} stock moves. Every total reads every row, including the last ones.", "para"),
    (None, "Your own columns. The notes column on Items holds supplier, shelf or anything else. The Stock Log has a note on every row.", "para"),
    (None, "Special prices. A wholesale or sale price goes in Price each on that Stock Log row; the profit uses it.", "para"),
    (None, "Count day. When a shelf count differs from On hand, log the difference as Count fix: add or Count fix: remove.", "para"),
    (None, "Any language setting. The formulas use no text-format codes, so they work the same in Excel set to other languages "
           "and in Google Sheets.", "para"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "Before you rely on a number", [(None, S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab"), "note")])
SH_LAST = r - 1
D.paint_canvas(sh, SH_LAST, "Z")
S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
               footer_notice="Estimates only. Not professional advice. See the Terms tab or LICENSE-AND-DISCLAIMER.txt.")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.", f"Version {VERSION}, {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More tools from the shop", [
    (None, ("Cash-Aware Reorder Planner: how much to reorder when cash is tight", "https://www.etsy.com/listing/4585004344"), "link"),
    (None, ("Etsy True Profit System: profit per sale after every Etsy fee", "https://www.etsy.com/listing/4584721212"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this tracker saves you time, a short review on Etsy helps other small shops find it. Thank you.", "para"),
    (None, "If a file won't open or something looks wrong, message the shop on Etsy and it will be fixed.", "para"),
])
TM_LAST = r - 1
D.paint_canvas(tm, TM_LAST, "Z")
S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb) if "input without validation" not in p]
if problems:
    print("AUDIT:", *problems[:20], sep="\n  ")
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
print(json.dumps({"sales": (r_from, r_to, r_units, r_sales, r_cost, r_prof, r_marg), "reorder_rows": RE_ROWS, "top": top_rows,
                  "ppu": ppu_rows, "val": val_rows, "checks": (k1, k2), "log_rows": len(LOG)}))
print("example:", json.dumps(EX), json.dumps({k: (v["oh"], v["status"], v["sold"]) for k, v in EX_ITEMS.items()}))
