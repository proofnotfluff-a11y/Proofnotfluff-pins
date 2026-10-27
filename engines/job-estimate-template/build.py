#!/usr/bin/env python3
"""Free Job Estimate Template (site freebie, built Oct 26, 2026 on Todd's go in a live chat).

Outputs, from one set of example data so the numbers never drift:
  docs/downloads/job-estimate-template.xlsx   Start Here | Estimate | Terms, the worked example filled in
  docs/downloads/job-estimate-template.pdf    page 1 a blank estimate to print and fill by hand, page 2 how to use it and the Terms
  docs/free-job-estimate-template/index.html  the <!--pnf-preview--> block is rewritten with the example estimate

The Estimate tab: your business, the customer and job, up to 8 line items (quantity x rate), one discount, one tax rate
on the total after the discount, and a deposit. Every total is blank-safe: clear every yellow cell and nothing shows an
error. Built to design/WORKBOOK_STANDARD.md with tools/pnf_dash.py and tools/wb_style.py; prints on one Letter page.

  python3 engines/job-estimate-template/build.py                 all three outputs
  python3 engines/job-estimate-template/build.py --xlsx out.xlsx [--blank]   the workbook only (stress tests)
"""
import datetime as dt
import html
import io
import os
import re
import sys
from datetime import datetime, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.styles import Alignment, Border  # noqa: E402

PRODUCT = "Job Estimate Template"
VERSION = "1"
CHECKED = "Oct 26, 2026"
CHECKED_LONG = "October 26, 2026"
SLUG = "free-job-estimate-template"
NAME = "job-estimate-template"
DL = os.path.join(ROOT, "docs", "downloads")
PAGE = os.path.join(ROOT, "docs", SLUG, "index.html")
USD2, DATE = S.FMT["usd2"], S.FMT["date"]
QTY = "General"
N_LINES = 8

# ---------------------------------------------------------------- the worked example (made up; prices are the quotes the free
# pressure washing price calculator gives for its own example business: docs/assets/pnf-calc.js service(), page defaults)
EX = dict(biz="Brightside Exterior (example)", phone="555-0142 (example)", email="hello@example.com", bizid="",
          cust="Sample Customer", addr="123 Example Street", cphone="555-0188 (example)",
          num="1001", date=dt.datetime(2026, 10, 26), days=30, disc=0.10, tax=0.05, dep=0.25,
          terms="Balance due when the job is done.")
LINES = [("House wash", "2-story house, soft wash", 1, "job", 355.00),
         ("Driveway", "2-car concrete, surface clean", 1, "job", 200.00),
         ("Gutters", "Face brightening, front and back", 1, "job", 205.00)]


def r2(x):
    from decimal import Decimal, ROUND_HALF_UP
    return float(Decimal(str(x)).quantize(Decimal("0.01"), ROUND_HALF_UP))


def example_math():
    amounts = [r2(q * rate) for _, _, q, _, rate in LINES]
    sub = r2(sum(amounts)); disc = r2(sub * EX["disc"]); tax = r2((sub - disc) * EX["tax"])
    total = r2(sub - disc + tax); dep = r2(total * EX["dep"]); bal = r2(total - dep)
    return dict(amounts=amounts, sub=sub, disc=disc, after=r2(sub - disc), tax=tax, total=total, dep=dep, bal=bal,
                until=EX["date"] + dt.timedelta(days=EX["days"]))


def money(x):
    return f"${x:,.2f}"


SHORT_NOTICE = S.SHORT_NOTICE.format(date=CHECKED_LONG)
FORMS = ("These are sample forms to adapt, not legal documents written for your state or business. Have a local attorney review "
         "any agreement before you use it with clients. If you collect client information, you're responsible for keeping it "
         "secure and following the privacy laws that apply to you.")
DEPOSITS = ("Deposits and cancelling. Rules on deposits, cancelling and the wording of signed estimates vary by state and city, "
            "especially for work sold at the customer's home. This template doesn't say what any law requires. Have a local attorney "
            "check your deposit, payment and cancellation terms before you use them with clients.")
ACCEPT = "To accept the work, prices and terms above, sign below and return it with the deposit."
SAMPLE_LINE = ("Sample estimate form to adapt. Have a local attorney review your terms before you use them with clients. "
               "Amounts are estimates for the work listed.")


# ================================================================ workbook
def build_xlsx(out, blank=False):
    ex = example_math()
    wb = S.new_workbook(PRODUCT, version=VERSION)
    wb.properties.description = "Copyright ProofNotFluff LLC. Free template for one business. Information only, not professional advice."
    ws = wb.create_sheet("Estimate")
    S.set_widths(ws, {"A": 3, "B": 2, "C": 21, "D": 31, "E": 2, "F": 2, "G": 9, "H": 10, "I": 12, "J": 14, "K": 2, "L": 3})
    D.page_header(ws, "Estimate", "", "", span=("B", "K"), side_cols=3)
    v = (lambda k: None) if blank else (lambda k: EX[k])

    # ---- left card: from and for
    A = D.Card(ws, 11, "B", "C", "D", "E", wb=wb)
    A.eyebrow("FROM")
    r_biz = A.input("Business name", v("biz"), S.FMT["text"], name="BusinessName", validation=("textLength", 0, 60), allow_blank=True,
                    prompt="Your business name as customers know it, for example Brightside Exterior Cleaning.", prompt_title="Business name",
                    error="Up to 60 characters.")
    r_ph = A.input("Phone", v("phone"), S.FMT["text"], validation=("textLength", 0, 40), allow_blank=True,
                   prompt="The number customers call or text.", prompt_title="Phone", error="Up to 40 characters.")
    r_em = A.input("Email or website", v("email"), S.FMT["text"], validation=("textLength", 0, 60), allow_blank=True,
                   prompt="Where the customer replies.", prompt_title="Email or website", error="Up to 60 characters.")
    r_id = A.input("License or ID", v("bizid"), S.FMT["text"], validation=("textLength", 0, 40), allow_blank=True,
                   prompt="A license, registration or insurance number if you show one. Leave blank if not.", prompt_title="License or ID",
                   error="Up to 40 characters.")
    A.eyebrow("FOR")
    r_cu = A.input("Customer", v("cust"), S.FMT["text"], name="CustomerName", validation=("textLength", 0, 60), allow_blank=True,
                   prompt="The customer's name.", prompt_title="Customer", error="Up to 60 characters.")
    r_ad = A.input("Job address", v("addr"), S.FMT["text"], validation=("textLength", 0, 80), allow_blank=True,
                   prompt="Where the work happens.", prompt_title="Job address", error="Up to 80 characters.")
    r_cp = A.input("Customer contact", v("cphone"), S.FMT["text"], validation=("textLength", 0, 60), allow_blank=True,
                   prompt="How you reach the customer. Optional.", prompt_title="Customer phone or email", error="Up to 60 characters.")
    A.close()
    for rr in (r_biz, r_ph, r_em, r_id, r_cu, r_ad, r_cp):
        ws[f"D{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1, shrink_to_fit=True)

    # ---- right card: estimate details and pricing terms
    B = D.Card(ws, 11, "F", "G", "J", "K", wb=wb)
    B.eyebrow("ESTIMATE")
    r_no = B.input("Estimate number", v("num"), S.FMT["text"], name="EstimateNo", validation=("textLength", 0, 20), allow_blank=True,
                   prompt="Your own number for this estimate, for example 1001.", prompt_title="Estimate number", error="Up to 20 characters.")
    r_dt = B.input("Estimate date", v("date"), DATE, name="EstimateDate", validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
                   allow_blank=True, prompt="The date you send it, for example 10/26/2026.", prompt_title="Estimate date",
                   error="Type a real date between 2000 and 2100, for example 10/26/2026.")
    r_vd = B.input("Good for (days)", v("days"), "0", name="ValidDays", validation=("whole", 0, 365), allow_blank=True,
                   prompt="How many days the prices hold. A whole number, for example 30.", prompt_title="Good for (days)",
                   error="Type a whole number of days from 0 to 365.")
    r_vu = B.calc("Good until", f'=IF(OR($J${r_dt}="",$J${r_vd}=""),"",$J${r_dt}+$J${r_vd})', DATE)
    B.eyebrow("DISCOUNT, TAX AND DEPOSIT")
    r_ds = B.input("Discount", v("disc"), "0%", name="DiscountRate", validation=("decimal", 0, 1), allow_blank=True,
                   prompt="Off the subtotal. Type a percent, for example 10%. 0% or blank for none.", prompt_title="Discount",
                   error="Type a percent from 0% to 100%.")
    r_tx = B.input("Sales tax rate", v("tax"), "0.0#%", name="TaxRate", validation=("decimal", 0, 0.5), allow_blank=True,
                   prompt="Charged on the total after the discount. Type a percent, for example 6.25%. 0% or blank if you don't charge tax on this job.",
                   prompt_title="Sales tax rate", error="Type a percent from 0% to 50%.")
    r_dp = B.input("Deposit to book", v("dep"), "0%", name="DepositRate", validation=("decimal", 0, 1), allow_blank=True,
                   prompt="Share of the total paid up front. Type a percent, for example 25%. 0% or blank for no deposit.",
                   prompt_title="Deposit to book", error="Type a percent from 0% to 100%.")
    B.close()
    for rr in (r_no,):
        ws[f"J{rr}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws[f"J{r_vu}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    CARDS_END = max(A.end, B.end)
    for rr in range(11, CARDS_END + 1):
        if ws.row_dimensions[rr].height is None:
            D.h(ws, rr, D.GRID_ROW)

    # ---- line items
    X_TOP = CARDS_END + 2
    D.h(ws, X_TOP - 1, 14)
    vals = [] if blank else LINES
    col = lambda i: [x[i] for x in vals]

    def f_amt(r):
        return f'=IF(OR($G{r}="",$I{r}=""),"",ROUND($G{r}*$I{r},2))'

    cols_x = [
        dict(col="C", head="Item", kind="input", values=col(0), validation=("textLength", 0, 40),
             prompt="A short name, for example Driveway.", error="Up to 40 characters."),
        dict(col="D", head="Description", kind="input", values=col(1), validation=("textLength", 0, 80),
             prompt="What the work covers, for example 2-car concrete driveway, surface clean.", error="Up to 80 characters."),
        dict(col="G", head="Qty", kind="input", fmt=QTY, align="right", values=col(2), validation=("decimal", 0, 1000000),
             prompt="How many: 1 job, 3 hours, 450 sq ft. A number of 0 or more.", error="Type a quantity of 0 or more."),
        dict(col="H", head="Unit", kind="input", values=col(3), validation=("textLength", 0, 12),
             prompt="What one is, for example job, hour, visit or sq ft. Optional.", error="Up to 12 characters."),
        dict(col="I", head="Rate", kind="input", fmt=USD2, align="right", values=col(4), validation=("decimal", -1000000, 1000000),
             prompt="Price for one unit, in dollars, for example 200. A credit can be negative.", error="Type a price in dollars."),
        dict(col="J", head="Amount", kind="calc", fmt=USD2, align="right", formula=f_amt, bold=True),
    ]
    x1, x2, x_end = D.table_card(ws, X_TOP, "B", "K", cols_x, N_LINES, title="Work and prices",
                                 sub="Quantity times rate. Blank rows are left out.", wb=wb)
    for rr in range(x1, x2 + 1):
        for k in ("E", "F"):   # the gap columns read as part of the description box (locked, not merged)
            ws[f"{k}{rr}"].fill = D.fill("input_fill")
            ws[f"{k}{rr}"].border = Border(bottom=D.side("FFFFFF"))
        for k in ("C", "D", "H"):
            ws[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws[f"I{rr}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
        ws[f"G{rr}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    for k in ("E", "F"):   # the header rule runs across the gap columns too
        ws[f"{k}{x1 - 1}"].border = Border(bottom=D.side("ink"))
    AMT = f"$J${x1}:$J${x2}"

    # ---- bottom: terms and acceptance (left), totals (right)
    T_TOP = x_end + 2
    D.h(ws, T_TOP - 1, 14)
    L = D.Card(ws, T_TOP, "B", "C", "D", "E", wb=wb)
    L.eyebrow("TERMS AND ACCEPTANCE")
    r_tm = L.input("Payment terms", v("terms"), S.FMT["text"], name="PaymentTerms", validation=("textLength", 0, 60), allow_blank=True,
                   prompt="One short line, for example Balance due when the job is done.", prompt_title="Payment terms", error="Up to 60 characters.")
    r_tm2 = L._row()   # two grid rows for the terms (a merge, so the right-hand card keeps its row height)
    box = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
    ws[f"D{r_tm2}"].border = box; ws[f"D{r_tm2}"].fill = D.fill("input_fill")
    ws.merge_cells(f"D{r_tm}:D{r_tm2}"); ws.merge_cells(f"C{r_tm}:C{r_tm2}")
    ws[f"D{r_tm}"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    L.text(ACCEPT, size=9, color="ink2", height=D.GRID_ROW * 2)
    sig_rows = []
    for lab in ("Name (print)", "Signature", "Date"):
        rr = L._row(); L._label(rr, lab)
        ws[f"D{rr}"].border = Border(bottom=D.side("ink"))
        sig_rows.append(rr)
    L.close()

    R = D.Card(ws, T_TOP, "F", "G", "J", "K", wb=wb)
    r_sub = R.calc("Subtotal", f'=IF(COUNT({AMT})=0,"",SUM({AMT}))', USD2)
    r_dsa = R.calc("Less discount", f'=IF($J${r_sub}="","",-ROUND($J${r_sub}*N($J${r_ds}),2))', USD2)
    ws[f"G{r_dsa}"] = f'=IF(N($J${r_ds})>0,"Less discount ("&ROUND($J${r_ds}*100,2)&"%)","Less discount")'
    r_txa = R.calc("Sales tax", f'=IF($J${r_sub}="","",ROUND(($J${r_sub}+$J${r_dsa})*N($J${r_tx}),2))', USD2)
    ws[f"G{r_txa}"] = f'=IF(N($J${r_tx})>0,"Sales tax ("&ROUND($J${r_tx}*100,2)&"%)","Sales tax")'
    r_tot = R.calc("Total", f'=IF($J${r_sub}="","",$J${r_sub}+$J${r_dsa}+$J${r_txa})', USD2, total=True, tint="total_tint")
    ws[f"J{r_tot}"].font = D.f(12, True)
    r_dpa = R.calc("Deposit to book", f'=IF($J${r_tot}="","",ROUND($J${r_tot}*N($J${r_dp}),2))', USD2)
    ws[f"G{r_dpa}"] = f'=IF(N($J${r_dp})>0,"Deposit to book ("&ROUND($J${r_dp}*100,2)&"%)","Deposit to book")'
    r_bal = R.calc("Balance due", f'=IF($J${r_tot}="","",$J${r_tot}-$J${r_dpa})', USD2, bold=True)
    while R.r < L.end:   # both cards end on the same row
        R._row()
    R.close()
    for rr in (r_sub, r_dsa, r_txa, r_tot, r_dpa, r_bal):
        ws[f"J{rr}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    BOT = max(L.end, R.end)
    for rr in range(T_TOP, BOT + 1):
        if ws.row_dimensions[rr].height is None:
            D.h(ws, rr, D.GRID_ROW)

    # ---- the sample line, full width under the cards (Legal README, Forms)
    r_fine = BOT + 1
    c = ws[f"C{r_fine}"]; c.value = SAMPLE_LINE; c.font = D.f(8, False, "ink2")
    c.alignment = Alignment(vertical="center", wrap_text=True); ws.merge_cells(f"C{r_fine}:J{r_fine}")
    D.h(ws, r_fine, 24)

    # ---- page header: the business name as the title, the estimate number and date at the right
    BIZ, NO, DT = f"$D${r_biz}", f"$J${r_no}", f"$J${r_dt}"
    ws["B2"] = f'=IF({BIZ}="","Your business name",{BIZ})'
    ws["I2"] = f'=IF({NO}="","ESTIMATE","ESTIMATE "&{NO})'
    ws["I2"].font = D.f(10, True, "ink2")
    ws["B4"] = (f'=IF(AND($D${r_ph}="",$D${r_em}="",$D${r_id}=""),"",$D${r_ph}&IF(AND($D${r_ph}<>"",$D${r_em}<>"")," | ","")&$D${r_em}'
                f'&IF($D${r_id}<>"",IF(OR($D${r_ph}<>"",$D${r_em}<>"")," | ","")&$D${r_id},""))')
    ws["I4"] = f'=IF({DT}="","",{DT})'
    ws["I4"].number_format = '"Date "mmm d, yyyy'

    # ---- tiles
    TOT, DEP, BAL = f"$J${r_tot}", f"$J${r_dpa}", f"$J${r_bal}"
    D.tile(ws, 6, "B", "E", "ESTIMATE TOTAL", f"={TOT}", USD2,
           f'=IF({TOT}="","Add a line item below to start the estimate",IF(N({DEP})>0,"Deposit to book "&TEXT({DEP},"$#,##0.00")'
           f'&", balance "&TEXT({BAL},"$#,##0.00"),"No deposit. Total due as set out in the terms"))', dark=True)
    D.tile(ws, 6, "F", "K", "FOR", f'=IF($D${r_cu}="","Customer name",$D${r_cu})', S.FMT["text"],
           f'=IF($J${r_vu}="",IF($D${r_ad}="","",$D${r_ad}),IF($D${r_ad}="","",$D${r_ad}&". ")&"Prices good until "&TEXT($J${r_vu},"mmm d, yyyy"))',
           value_size=18)
    D.h(ws, 10, 14)

    LAST = r_fine + 1
    D.h(ws, LAST, 12)
    D.paint_canvas(ws, LAST, "Z")
    S.finish_sheet(ws, PRODUCT, LAST, span=("A", "L"), freeze="A11", tab_color="accent",
                   footer_notice="Sample form. Estimates only, not legal or tax advice. See the Terms tab.")
    ws.page_setup.fitToHeight = 1
    ws.print_area = f"A1:L{LAST}"
    cells = dict(sub=r_sub, disc=r_dsa, tax=r_txa, total=r_tot, dep=r_dpa, bal=r_bal, until=r_vu, first=x1, last=x2,
                 inputs=dict(D=[r_biz, r_ph, r_em, r_id, r_cu, r_ad, r_cp, r_tm], J=[r_no, r_dt, r_vd, r_ds, r_tx, r_dp]))

    # ================================================================ Start Here
    SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}

    def sh_card(sh, top, title, rows):
        W = S.width_of(sh, "D") + S.width_of(sh, "E")
        r = top
        D.h(sh, r, 12); r += 1
        t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
        for left, text, kind in rows:
            full = kind in ("para", "link", "note")
            c = sh[f"C{r}"] if full else sh[f"D{r}"]
            if kind == "link":
                c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
            else:
                c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
                mm = re.match(r"^([A-Z][A-Za-z0-9 ,]{1,34})\. (.*)$", text, re.S) if kind in ("para", "num") else None
                if mm:
                    c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                           TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if full:
                sh.merge_cells(f"C{r}:E{r}")
                ht = S.text_height(body, W + 8, 10) + 8
            else:
                sh.merge_cells(f"D{r}:E{r}")
                ht = max(S.text_height(body, W, 10) + 10, 30)
            Lc = sh[f"C{r}"]
            if kind == "num":
                Lc.value = left; Lc.font = D.f(18, True, "accent"); Lc.alignment = Alignment(horizontal="center", vertical="center")
            elif kind.startswith("chip"):
                Lc.value = left[0]; Lc.number_format = left[1]; Lc.alignment = Alignment(horizontal="center", vertical="center")
                if kind == "chip-in":
                    Lc.fill = D.fill("input_fill"); Lc.font = D.f(9, True, "input_text")
                    Lc.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
                elif kind == "chip-calc":
                    Lc.font = D.f(9, False, "ink"); Lc.border = Border(bottom=D.side("hair"))
                else:
                    Lc.fill = D.fill("ink"); Lc.font = D.f(9, True, "on_ink")
            D.h(sh, r, ht); r += 1
        D.h(sh, r, 12)
        D.frame(sh, top, r, "B", "F")
        return r + 2

    sh = wb.create_sheet("Start Here", 0)
    S.set_widths(sh, SH_GRID)
    D.page_header(sh, PRODUCT, "Start here. One estimate tab, one page to print.", f"Version {VERSION}, {CHECKED}",
                  span=("B", "F"), side_cols=2)
    r = 6
    r = sh_card(sh, r, "What it does", [(None,
        "A one-page estimate for cleaning, pest control, lawn care, pressure washing and other service jobs. Type your business, "
        "the customer and up to 8 lines of work. Each amount, the discount, sales tax, total, deposit and balance add up on their "
        "own. Print it or save it as a PDF to send.", "para")])
    r = sh_card(sh, r, "Three steps", [
        (1, "Your details. On the Estimate tab, type your business name, phone and email once. They print at the top.", "num"),
        (2, "The job. Type the customer, the estimate number and date, then one line per piece of work: item, description, "
            "quantity, unit and rate. Set the discount, sales tax and deposit, or leave them at 0%.", "num"),
        (3, "Send it. File, Print fits the tab on one Letter page, or print to PDF and email it. The customer signs the "
            "acceptance lines and returns it with the deposit.", "num"),
    ])
    r = sh_card(sh, r, "How to read the cells", [
        ((355, USD2), "Yellow cells with blue text are yours to type in. Each one shows a hint when you select it.", "chip-in"),
        ((ex["sub"], USD2), "Numbers on white are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
        ((ex["total"], USD2), "The dark tile is the estimate total. It sits at the top of the Estimate tab and stays on screen as you scroll.", "chip-key"),
        (None, "Each tab is protected without a password. To change a locked cell, use Review, Unprotect Sheet in Excel, or Data, "
               "Protect sheets and ranges in Google Sheets.", "note"),
    ])
    S.page_break_before(sh, r)
    ex_text = (f"The example is a made-up pressure washing business. Three lines of work: a 2-story house wash at {money(355)}, a 2-car "
               f"driveway at {money(200)} and gutter brightening at {money(205)}, a subtotal of {money(ex['sub'])}. A 10% discount takes off "
               f"{money(ex['disc'])}, leaving {money(ex['after'])}. Sales tax at an example 5% adds {money(ex['tax'])}, so the total is "
               f"{money(ex['total'])}. A 25% deposit of {money(ex['dep'])} books the job and {money(ex['bal'])} is due when it is done. "
               f"The prices are the quotes the free pressure washing price calculator at proofnotfluff.com gives for its own example "
               f"business, not market prices for your area.")
    r = sh_card(sh, r, "The example", [(None, ex_text, "para")])
    r = sh_card(sh, r, "Good to know", [
        (None, "Clearing the example. Select the yellow cells and press Delete. Formulas live only in the white cells, so every "
               "total keeps working and a blank estimate shows no errors.", "para"),
        (None, "A new estimate each time. Save a copy of the file for each job (File, Save As), so every estimate you sent is kept.", "para"),
        (None, "Sales tax. Whether a service is taxed, and at what rate, depends on where you work and the kind of work. The tax "
               "rate cell is yours to set; it is charged on the total after the discount.", "para"),
        (None, DEPOSITS, "para"),
        (None, "Pricing each line. The free calculators at proofnotfluff.com price a clean, a lawn visit or a pressure washing job "
               "from your own costs and time. The Terms tab links them.", "para"),
        (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. No macros, no add-ons.", "para"),
    ])
    r = sh_card(sh, r, "Before you rely on a number", [
        (None, SHORT_NOTICE.replace("Terms of Use page", "Terms tab"), "note"),
        (None, FORMS, "note"),
    ])
    SH_LAST = r - 1
    D.paint_canvas(sh, SH_LAST, "Z")
    S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
                   footer_notice="Sample form. Estimates only, not legal or tax advice. See the Terms tab.")

    tm = wb.create_sheet("Terms")
    S.set_widths(tm, SH_GRID)
    D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.",
                  f"Version {VERSION}, {CHECKED}", span=("B", "F"), side_cols=2)
    paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
    r = 6
    r = sh_card(tm, r, "Free calculators to price each line", [
        (None, ("Pressure washing price calculator", "https://proofnotfluff.com/pressure-washing-price-calculator/"), "link"),
        (None, ("House cleaning price calculator", "https://proofnotfluff.com/house-cleaning-price-calculator/"), "link"),
        (None, ("Lawn care price calculator", "https://proofnotfluff.com/lawn-care-price-calculator/"), "link"),
        (None, ("Price increase calculator", "https://proofnotfluff.com/price-increase-calculator/"), "link"),
    ])
    r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
    TM_LAST = r - 1
    D.paint_canvas(tm, TM_LAST, "Z")
    S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    problems = [p for p in S.audit(wb) if "input without validation" not in p]
    if problems:
        print("AUDIT:", *problems, sep="\n  ")
    wb.save(out)
    _set_company(out)
    return cells


def _set_company(path):
    """docProps/app.xml: Company ProofNotFluff LLC, Application ProofNotFluff (author stays ProofNotFluff, as legal_scan wants)."""
    import zipfile
    src = open(path, "rb").read(); zin = zipfile.ZipFile(io.BytesIO(src)); out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "docProps/app.xml":
                s = data.decode("utf-8")
                s = re.sub(r"<Application>[^<]*</Application>", "<Application>ProofNotFluff</Application>", s)
                s = re.sub(r"<Company>[^<]*</Company>|<Company/>", "", s)
                s = s.replace("</Properties>", "<Company>ProofNotFluff LLC</Company></Properties>")
                data = s.encode("utf-8")
            z.writestr(it, data)
    open(path, "wb").write(out.getvalue())


# ================================================================ html: the blank PDF and the page preview
PDF_CSS = """
@page{size:Letter;margin:0}
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:Arial,"Liberation Sans",Helvetica,sans-serif;color:#1d2433;background:#fff;font-size:10pt;line-height:1.38}
.page{width:8.5in;height:11in;padding:.55in .65in .7in;position:relative;page-break-after:always;overflow:hidden}
.page:last-child{page-break-after:auto}
h1,h2{font-family:Poppins,Arial,sans-serif;font-weight:700;letter-spacing:-.01em}
.eyebrow{font-size:8pt;letter-spacing:.14em;text-transform:uppercase;color:#a63e21;font-weight:700}
.bars{display:flex;gap:4px;margin:.07in 0 .05in}.bars i{display:block;width:.26in;height:4px;border-radius:2px}
.bars i:nth-child(1){background:#c8502f}.bars i:nth-child(2){background:#2e6b66}.bars i:nth-child(3){background:#6b4fa0}.bars i:nth-child(4){background:#2f5f9e}.bars i:nth-child(5){background:#a87c1f}
.top{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:2px solid #1d2433;padding-bottom:.1in}
.top h1{font-size:22pt;line-height:1}
.top .right{text-align:right;font-size:9pt;color:#4a5163}
.top .right b{display:block;font-family:Poppins,Arial,sans-serif;font-size:15pt;color:#1d2433;letter-spacing:.06em}
.two{display:grid;grid-template-columns:1fr 1fr;gap:.25in;margin-top:.16in}
.lab{font-size:7.4pt;letter-spacing:.1em;text-transform:uppercase;color:#5c6270;font-weight:700;margin-bottom:.03in}
.ln{border-bottom:1px solid #9aa1ae;height:.3in;display:flex;align-items:flex-end;padding-bottom:.03in;font-size:8.6pt;color:#5c6270}
.ln span{min-width:1.1in}
table{width:100%;border-collapse:collapse;margin-top:.2in;font-size:9pt}
th{text-align:left;font-size:7.4pt;letter-spacing:.08em;text-transform:uppercase;color:#5c6270;font-weight:700;padding:.04in .06in;border-bottom:1.5px solid #1d2433}
th.n{text-align:right}
td{height:.34in;border-bottom:1px solid #c9ced9;padding:0 .06in}
.tot{display:grid;grid-template-columns:1fr 2.6in;gap:.3in;margin-top:.2in}
.sum div{display:flex;justify-content:space-between;border-bottom:1px solid #c9ced9;height:.3in;align-items:flex-end;padding-bottom:.03in;font-size:9.4pt}
.sum div.t{border-bottom:2px solid #1d2433;font-weight:700;font-size:10.5pt;background:#f6f3ee;padding-left:.06in}
.sum span:last-child{min-width:.9in;text-align:right;color:#9aa1ae}
.acc p{font-size:9pt;color:#4a5163;margin-bottom:.04in}
.fine{position:absolute;left:.65in;right:.65in;bottom:.62in;font-size:7.6pt;color:#4a5163}
.foot{position:absolute;left:.65in;right:.65in;bottom:.36in;font-size:7.4pt;color:#5c6270;border-top:1px solid #e2ddd3;padding-top:.06in;display:flex;justify-content:space-between}
h2{font-size:12.5pt;margin:.16in 0 .05in}
ol.steps{list-style:none;counter-reset:s}
ol.steps li{counter-increment:s;position:relative;padding:.045in 0 .045in .34in;border-bottom:1px solid #efebe3;font-size:9.8pt}
ol.steps li::before{content:counter(s);position:absolute;left:0;top:.05in;width:.22in;height:.22in;border-radius:50%;background:#2e6b66;color:#fff;font-weight:700;font-size:8.5pt;line-height:.22in;text-align:center}
.terms p{font-size:8.6pt;margin-bottom:.06in}
.small{font-size:8.2pt;color:#4a5163;margin-top:.1in}
"""


def blank_form_html():
    rows = "".join("<tr><td></td><td></td><td></td><td></td><td></td><td></td></tr>" for _ in range(10))
    ln = lambda lab: f'<div class="ln"><span>{lab}</span></div>'
    return f"""<section class="page">
<div class="top"><div><p class="lab">Business name</p><h1>&nbsp;</h1></div><div class="right"><b>ESTIMATE</b>No. ________ &nbsp; Date ____________</div></div>
<div class="two"><div><p class="lab">From</p>{ln("Phone")}{ln("Email or website")}{ln("License or ID")}</div>
<div><p class="lab">For</p>{ln("Customer")}{ln("Job address")}{ln("Phone or email")}</div></div>
<table><thead><tr><th style="width:1.3in">Item</th><th>Description</th><th class="n" style="width:.55in">Qty</th><th style="width:.65in">Unit</th><th class="n" style="width:.85in">Rate</th><th class="n" style="width:1in">Amount</th></tr></thead><tbody>{rows}</tbody></table>
<div class="tot"><div class="acc"><p class="lab" style="margin-top:.04in">Terms and acceptance</p>{ln("Payment terms")}{ln("Prices good until")}
<p style="margin-top:.12in">{html.escape(ACCEPT)}</p>{ln("Name (print)")}{ln("Signature")}{ln("Date")}</div>
<div class="sum"><div><span>Subtotal</span><span>$</span></div><div><span>Less discount (____%)</span><span>$</span></div>
<div><span>Sales tax (____%)</span><span>$</span></div><div class="t"><span>Total</span><span>$</span></div>
<div><span>Deposit to book (____%)</span><span>$</span></div><div><span>Balance due</span><span>$</span></div></div></div>
<p class="fine">{html.escape(SAMPLE_LINE)}</p>
<div class="foot"><span>Free job estimate template from ProofNotFluff | proofnotfluff.com/{SLUG}/</span><span>Page 1 of 2</span></div>
</section>"""


def guide_html():
    steps = [("Your details.", "Write your business name, phone and email at the top, or type them once in the Excel version so they print every time."),
             ("One line per piece of work.", "Item, what it covers, quantity, unit and rate. Amount is quantity times rate."),
             ("Add it up.", "Subtotal, less any discount, plus sales tax on what is left, is the total. The deposit is a share of the total; the balance is the rest."),
             ("Get it signed.", "The customer signs the acceptance lines and returns the estimate with the deposit. Keep a copy.")]
    dep_lead, dep_body = DEPOSITS.split(". ", 1)
    steps_h = "".join(f"<li><b>{html.escape(a)}</b> {html.escape(b)}</li>" for a, b in steps)
    paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
    terms = "".join("<p><b>{}.</b> {}</p>".format(*map(html.escape, p.split(". ", 1))) for p in paras)
    ex = example_math()
    return f"""<section class="page terms">
<p class="eyebrow">How to use it · Terms</p><div class="bars"><i></i><i></i><i></i><i></i><i></i></div>
<h1 style="font-size:19pt">Job estimate template</h1>
<p style="font-size:9.6pt;color:#4a5163">Print page 1 and fill it by hand, or use the Excel version (proofnotfluff.com/downloads/{NAME}.xlsx), which adds every line up for you. Price each line with the free calculators at proofnotfluff.com.</p>
<ol class="steps" style="margin-top:.08in">{steps_h}</ol>
<p class="small" style="color:#1d2433"><b>{html.escape(dep_lead)}.</b> {html.escape(dep_body)}</p>
<p class="small">Worked example in the Excel file: {money(ex['sub'])} of work, less a 10% discount of {money(ex['disc'])}, plus an example 5% sales tax of {money(ex['tax'])}, is {money(ex['total'])}; a 25% deposit is {money(ex['dep'])} and {money(ex['bal'])} is due when the job is done.</p>
<p class="small">{html.escape(SHORT_NOTICE)}</p>
<h2>Terms of Use and Disclaimer</h2>
<p style="font-size:8.8pt;margin-bottom:.08in;color:#4a5163"><b>ProofNotFluff | {PRODUCT} | Version {VERSION}, {CHECKED_LONG}</b></p>
{terms}
<div class="foot"><span>Free job estimate template from ProofNotFluff | proofnotfluff.com/{SLUG}/</span><span>Page 2 of 2</span></div>
</section>"""


def page_preview():
    ex = example_math()
    rows = "".join(f'<tr><th scope="row"><b>{html.escape(i)}</b><span>{html.escape(d)}</span></th><td>{q:g} {html.escape(u)}</td>'
                   f'<td>{money(rate)}</td><td><b>{money(a)}</b></td></tr>'
                   for (i, d, q, u, rate), a in zip(LINES, ex["amounts"]))
    until = ex["until"]
    return f"""<!--pnf-preview--><div class="pe" role="group" aria-label="Sample estimate from the template, worked example">
<div class="pe-top"><div><p class="pe-biz">{html.escape(EX['biz'])}</p><p class="pe-sub">{html.escape(EX['phone'])} | {html.escape(EX['email'])}</p></div>
<div class="pe-no"><b>ESTIMATE {html.escape(EX['num'])}</b><span>Date Oct 26, 2026</span></div></div>
<div class="pe-tiles"><div class="pe-tile key"><span class="lab">Estimate total</span><b>{money(ex['total'])}</b><span>Deposit to book {money(ex['dep'])}, balance {money(ex['bal'])}</span></div>
<div class="pe-tile"><span class="lab">For</span><b class="sm">{html.escape(EX['cust'])}</b><span>{html.escape(EX['addr'])}. Prices good until {until:%b} {until.day}, {until.year}</span></div></div>
<table class="pe-lines"><caption class="pe-vh">Work and prices</caption><thead><tr><th scope="col">Item</th><th scope="col">Qty</th><th scope="col">Rate</th><th scope="col">Amount</th></tr></thead><tbody>{rows}</tbody></table>
<div class="pe-foot"><div class="pe-acc"><p class="lab">Terms and acceptance</p><p>{html.escape(EX['terms'])} {html.escape(ACCEPT)}</p><p class="pe-sig">Name (print)</p><p class="pe-sig">Signature</p></div>
<dl class="pe-sum"><div><dt>Subtotal</dt><dd>{money(ex['sub'])}</dd></div><div><dt>Less discount (10%)</dt><dd>-{money(ex['disc'])}</dd></div>
<div><dt>Sales tax (5%)</dt><dd>{money(ex['tax'])}</dd></div><div class="t"><dt>Total</dt><dd>{money(ex['total'])}</dd></div>
<div><dt>Deposit to book (25%)</dt><dd>{money(ex['dep'])}</dd></div><div><dt>Balance due</dt><dd>{money(ex['bal'])}</dd></div></dl></div>
<p class="pe-fine">{html.escape(SAMPLE_LINE)}</p></div><!--/pnf-preview-->"""


def render_pdf(html_text, path, title, subject):
    from playwright.sync_api import sync_playwright
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import NameObject, TextStringObject
    tmp = path + ".src.html"
    open(tmp, "w", encoding="utf-8").write(html_text)
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page()
        pg.goto("file://" + tmp, wait_until="load"); pg.evaluate("document.fonts.ready")
        raw = pg.pdf(format="Letter", print_background=True, prefer_css_page_size=True, tagged=True)
        b.close()
    os.remove(tmp)
    r = PdfReader(io.BytesIO(raw)); w = PdfWriter()
    for page in r.pages:
        w.add_page(page)
    stamp = datetime.now(timezone.utc).strftime("D:%Y%m%d%H%M%S+00'00'")
    w.add_metadata({"/Title": title, "/Author": "ProofNotFluff LLC", "/Subject": subject, "/Creator": "ProofNotFluff LLC",
                    "/Producer": "ProofNotFluff LLC", "/Keywords": "ProofNotFluff", "/CreationDate": stamp, "/ModDate": stamp})
    w._root_object[NameObject("/Lang")] = TextStringObject("en-US")
    with open(path, "wb") as f:
        w.write(f)


def main():
    args = sys.argv[1:]
    if "--xlsx" in args:
        out = args[args.index("--xlsx") + 1]
        print(build_xlsx(out, blank="--blank" in args))
        return
    os.makedirs(DL, exist_ok=True)
    cells = build_xlsx(os.path.join(DL, NAME + ".xlsx"))
    doc = (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{PRODUCT}</title><style>{PDF_CSS}</style></head>'
           f"<body>{blank_form_html()}{guide_html()}</body></html>")
    render_pdf(doc, os.path.join(DL, NAME + ".pdf"), PRODUCT, "Free template: a blank job estimate to print, with how to use it and the Terms")
    if os.path.isfile(PAGE):
        t = open(PAGE, encoding="utf-8").read()
        t2 = re.sub(r"<!--pnf-preview-->.*?<!--/pnf-preview-->", lambda m: page_preview(), t, flags=re.S)
        if t2 != t:
            open(PAGE, "w", encoding="utf-8").write(t2)
    print("built", NAME + ".xlsx", NAME + ".pdf", cells)


if __name__ == "__main__":
    main()
