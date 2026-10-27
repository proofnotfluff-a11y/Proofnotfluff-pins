#!/usr/bin/env python3
"""Free post-construction cleaning proposal template (Todd's go, Oct 27, 2026: "Let's go with your recommendation",
coo/DECISIONS-2026-10-27.md). The download behind the Kit form on the free bid calculator page.

One source for the proposal's words and layout: PNF.bidProposal in docs/assets/pnf-calc.js (the same function the calculator
page and the area seats' bid sheet print from), run here with node and every price left as a [bracket]. Outputs:
  docs/downloads/post-construction-cleaning-proposal-template.docx   editable: how-to page, the proposal, Terms
  docs/downloads/post-construction-cleaning-proposal-template.pdf    printable, the same three pages (Chromium, Playwright)
  docs/post-construction-cleaning-bid-calculator/thanks/index.html   the page the Kit form redirects to (engines/thanks_page.py)

Usage: python3 engines/post-construction-proposal/build.py   (after tools/make_site.py has written the calculator page)
"""
import html, io, json, os, re, subprocess, sys, zipfile
from datetime import datetime, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
DL = os.path.join(ROOT, "docs", "downloads")
SLUG = "post-construction-cleaning-bid-calculator"
NAME = "post-construction-cleaning-proposal-template"
PRODUCT = "Post-Construction Cleaning Proposal Template"
CHECKED = "October 27, 2026"
CALC = "proofnotfluff.com/post-construction-cleaning-bid-calculator/"
DOCX_URL = f"proofnotfluff.com/downloads/{NAME}.docx"
FIELD = re.compile(r"\[[^\[\]]+\]")

STEPS = [
    ("Get the size.", "Ask the builder for the square footage from the plans, or measure it on a walk-through. For apartments, count the units, "
                      "their average size, and the hallways, lobby and amenity rooms."),
    ("Time your own crew.", "For each clean, write down the square feet one cleaner does in an hour on your own past jobs. Those are the only "
                            "speeds that fit your crew."),
    ("Price each clean.", f"The free post-construction cleaning bid calculator ({CALC}) prices the rough, final and touch-up cleans from your "
                          "speed, labor cost, supplies, travel and margin, and shows the hours behind each price."),
    ("Fill every bracket.", "In the Word file, search for [ to find each field. Delete any clean that is not in your bid, and change the "
                            "scope lines to match what the builder asked for."),
    ("Add your terms.", "Write your payment terms, how long the price holds, and what is not included."),
    ("Send it as a PDF.", "Save the filled proposal as a PDF and email it to the builder's estimator or project manager with a short note. "
                          "Keep a copy and the date you sent it."),
]
SAMPLE = ("This is a sample form to adapt, not a legal document written for your state or business. Have a local attorney review any "
          "agreement before you use it with clients. If you collect client information, you're responsible for keeping it secure and "
          "following the privacy laws that apply to you.")
SHORT_NOTICE = ("For general information and planning only. This is not legal, tax, financial, medical or other professional advice, "
                "and using it doesn't create a professional relationship. Results are estimates based on the numbers you enter. "
                f"Figures were checked on {CHECKED} and can change. See the Terms of Use page before you rely on anything here.")
TERMS = [
    ("License", "This free template is for your own use or one business you own. You may edit it, print it and send the proposals to "
                "your own customers. You may not resell, upload or redistribute the template, or use it to make a competing product."),
    ("Ownership", "The design and text of the template belong to ProofNotFluff. The proposals you fill in and send are yours."),
    ("Sample form", SAMPLE),
    ("Information only", "This product shares general information and tools for planning. It is not legal, tax, accounting, financial, "
                         "employment, medical or other professional advice, and it doesn't create a client relationship of any kind. Your "
                         "situation may differ. For decisions that matter, check with a qualified professional and the official source."),
    ("Accuracy", "The template uses no outside rates or rules: every price, size and term in a proposal comes from you. Laws, prices and "
                 "rules change, and state and local rules vary. You're responsible for confirming anything you rely on."),
    ("Estimates", "Calculations depend on the numbers you enter and simplify real-world conditions. They are estimates, not promises of "
                  "any result, income, savings, approval or outcome."),
    ("No warranty", "The product is provided \"as is\", without warranties of any kind, to the extent the law allows."),
    ("Limit of liability", "To the extent the law allows, ProofNotFluff is not liable for any loss or damage from using or relying on "
                           "this product. Nothing here limits rights you have under consumer protection laws that can't be waived."),
    ("Trademarks", "Microsoft Word and Google Docs are trademarks of their owners. ProofNotFluff isn't affiliated with, sponsored or "
                   "endorsed by them."),
    ("AI assistance", "This product was made with AI assistance and reviewed and tested by the shop owner."),
    ("Help", "Reply to any ProofNotFluff email if something looks wrong and it will be fixed."),
]


def from_js():
    """The proposal HTML and the scope lines from docs/assets/pnf-calc.js, every price left blank."""
    js = r"""
const P = require(process.argv[1]);
const r = { sqft: 0, measure: 0, price: 0, roughOn: 1, finalOn: 1, touchOn: 1, roughPrice: 0, finalPrice: 0, touchPrice: 0, touchVisits: "[n]" };
const t = { date: "", business: "", contact: "", builder: "", address: "", work: "", terms: "" };
console.log(JSON.stringify({ html: P.bidProposal(r, t), scope: P.bidScope, names: P.bidNames }));"""
    out = subprocess.run(["node", "-e", js, os.path.join(ROOT, "docs", "assets", "pnf-calc.js")], capture_output=True, text=True, check=True).stdout
    d = json.loads(out)
    d["html"] = d["html"].replace(", [n] visits", ', <span class="bid-f">[number of]</span> visits')
    assert "[n]" not in d["html"]
    # the PDF matches the Word file (Legal desk, Oct 27): a Work line and a Not included line, as fields to fill
    proj = '<div><dt>Project</dt><dd><span class="bid-f">[Project address]</span></dd></div>'
    assert d["html"].count(proj) == 1
    d["html"] = d["html"].replace(proj, proj + '<div><dt>Work</dt><dd><span class="bid-f">[Project, for example: 6-story apartment building]</span></dd></div>', 1)
    terms_end = '[Your payment terms, and how long this price holds]</span></p>'
    assert d["html"].count(terms_end) == 1
    d["html"] = d["html"].replace(terms_end, terms_end + '<p class="bid-terms" style="margin-top:6px!important">Not included: <span class="bid-f">[anything outside the scope above]</span></p>', 1)
    return d


def check_text(d):
    alltext = " ".join([a + " " + b for a, b in STEPS + TERMS] + [SAMPLE, SHORT_NOTICE] + list(d["scope"].values()) + [re.sub(r"<[^>]+>", " ", d["html"])])
    assert chr(0x2014) not in alltext and chr(0x2013) not in alltext, "dash"
    assert not re.search(r"guarantee|risk-free|legally binding|ensures? compliance", alltext, re.I)


# ------------------------------------------------------------------ docx
def build_docx(path, d):
    from docx import Document
    from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches

    FIELD_BLUE = RGBColor(0x1F, 0x4E, 0x9E)
    INK, MUTED, ACCENT, TEAL = RGBColor(0x1D, 0x24, 0x33), RGBColor(0x5C, 0x62, 0x70), RGBColor(0xA6, 0x3E, 0x21), RGBColor(0x2E, 0x6B, 0x66)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.8)
    st = doc.styles["Normal"]
    st.font.name = "Arial"; st.font.size = Pt(10.5); st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.12
    for lvl, size in (("Heading 1", 20), ("Heading 2", 13)):
        h = doc.styles[lvl]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = INK
        rf = h.element.rPr.find(qn("w:rFonts"))
        if rf is not None:
            for a in list(rf.attrib):
                del rf.attrib[a]
            rf.set(qn("w:ascii"), "Arial"); rf.set(qn("w:hAnsi"), "Arial")
        h.paragraph_format.space_before = Pt(12 if lvl == "Heading 2" else 0); h.paragraph_format.space_after = Pt(6)

    def shade(run, hexfill):
        rpr = run._r.get_or_add_rPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexfill)
        rpr.append(shd)

    def runs(p, text, size=None, bold=False, color=None):
        """Plain text with [bracket] fields in input blue, bold, on light yellow."""
        last = 0
        for m in FIELD.finditer(text):
            if m.start() > last:
                r = p.add_run(text[last:m.start()]); r.bold = bold
                if size: r.font.size = Pt(size)
                if color is not None: r.font.color.rgb = color
            r = p.add_run(m.group(0)); r.bold = True; r.font.color.rgb = FIELD_BLUE; shade(r, "FDF6DC")
            if size: r.font.size = Pt(size)
            last = m.end()
        if last < len(text):
            r = p.add_run(text[last:]); r.bold = bold
            if size: r.font.size = Pt(size)
            if color is not None: r.font.color.rgb = color
        return p

    def para(text, size=None, color=None, bold=False, after=None, lead=None, align=None):
        p = doc.add_paragraph()
        if lead:
            r = p.add_run(lead + " "); r.bold = True
            if size: r.font.size = Pt(size)
        runs(p, text, size, bold, color)
        if after is not None: p.paragraph_format.space_after = Pt(after)
        if align: p.alignment = align
        return p

    def borders(cell, **sides):
        tcpr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for side, (sz, color) in sides.items():
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), color); b.append(e)
        tcpr.append(b)

    def widths(t, ws):
        t.autofit = False
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), ws):
            gc.set(qn("w:w"), str(int(w.twips)))
        for row in t.rows:
            for c, w in zip(row.cells, ws):
                c.width = w

    def page_break():
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    # page 1: how to price and send it
    para("FREE TEMPLATE FROM PROOFNOTFLUFF", size=8.5, color=ACCENT, bold=True, after=2)
    doc.add_heading("Post-construction cleaning proposal template", level=1)
    para("A one-page proposal for a builder: rough, final and touch-up cleans, each with its scope and your price. The proposal is on "
         "page 2. Fill the yellow fields, delete what is not in your bid, and send it as a PDF.", color=MUTED, after=8)
    doc.add_heading("How to price and send it", level=2)
    for i, (lead, text) in enumerate(STEPS, 1):
        para(text, lead=f"{i}. {lead}", after=5)
    para(SAMPLE, size=9.5, color=MUTED, after=6).paragraph_format.space_before = Pt(8)
    para(SHORT_NOTICE, size=8.5, color=MUTED, after=0)

    # page 2: the proposal
    page_break()
    head = doc.add_table(rows=1, cols=2)
    widths(head, (Inches(4.4), Inches(2.3)))
    a, b = head.rows[0].cells
    runs(a.paragraphs[0], "[Your business name]", size=15, bold=True)
    runs(a.add_paragraph(), "[Phone]  |  [Email]", size=10, color=MUTED)
    pb = b.paragraphs[0]; pb.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = pb.add_run("PROPOSAL"); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = ACCENT
    p2 = b.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT; runs(p2, "[Date]", size=10)
    for c in (a, b):
        borders(c, bottom=(18, "1D2433"))
    para("", after=2)
    h = para("Post-construction cleaning proposal", size=17, bold=True, after=8)
    meta = doc.add_table(rows=2, cols=2)
    widths(meta, (Inches(3.35), Inches(3.35)))
    for cell, (k, v) in zip([c for row in meta.rows for c in row.cells],
                            [("PREPARED FOR", "[Builder or general contractor]"), ("PROJECT", "[Project address]"),
                             ("WORK", "[Project, for example: 6-story apartment building]"), ("SIZE", "[Square feet, or units and their size]")]):
        p = cell.paragraphs[0]; r = p.add_run(k); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = MUTED
        p.paragraph_format.space_after = Pt(0)
        q = cell.add_paragraph(); runs(q, v, size=10.5, bold=True); q.paragraph_format.space_after = Pt(6)
    para("", after=2)
    tab = doc.add_table(rows=1, cols=2)
    widths(tab, (Inches(5.4), Inches(1.3)))
    h1, h2 = tab.rows[0].cells
    for c, txt, al in ((h1, "CLEAN AND SCOPE", None), (h2, "PRICE", WD_ALIGN_PARAGRAPH.RIGHT)):
        r = c.paragraphs[0].add_run(txt); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = MUTED
        if al: c.paragraphs[0].alignment = al
        borders(c, bottom=(12, "1D2433"))
    for k in ("rough", "final", "touch"):
        name = d["names"][k] + (", [number of] visits" if k == "touch" else "")
        c1, c2 = tab.add_row().cells
        runs(c1.paragraphs[0], name, bold=True); c1.paragraphs[0].paragraph_format.space_after = Pt(1)
        q = c1.add_paragraph(); runs(q, d["scope"][k], size=9.5, color=MUTED); q.paragraph_format.space_after = Pt(5)
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT; runs(c2.paragraphs[0], "[price]", bold=True)
        for c in (c1, c2):
            borders(c, bottom=(4, "E2DDD3"))
    c1, c2 = tab.add_row().cells
    runs(c1.paragraphs[0], "Total", size=12.5, bold=True)
    c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT; runs(c2.paragraphs[0], "[total]", size=12.5, bold=True)
    for c in (c1, c2):
        borders(c, bottom=(18, "1D2433"))
    para("", after=2)
    para("[Your payment terms, and how long this price holds]", lead="Terms.", after=3)
    para("Not included: [anything outside the scope above]", after=3)
    para("Prices are for the scope and size above. Work outside that scope is priced separately.", size=9, color=MUTED, after=26)
    sign = doc.add_table(rows=1, cols=5)   # spacer columns keep the three signature lines apart
    widths(sign, (Inches(2.5), Inches(0.25), Inches(2.3), Inches(0.25), Inches(1.4)))
    for c, label in zip(sign.rows[0].cells[::2], ("Accepted by, name and title", "Signature", "Date")):
        borders(c, top=(6, "1D2433"))
        r = c.paragraphs[0].add_run(label); r.font.size = Pt(8.5); r.font.color.rgb = MUTED

    # page 3: Terms (a paragraph after the signature table carries the break, so the table stays on page 2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    page_break()
    para("TERMS", size=8.5, color=ACCENT, bold=True, after=2)
    doc.add_heading("Terms of Use and Disclaimer", level=1)
    para(f"ProofNotFluff | {PRODUCT} | Version 1, {CHECKED}", size=9.5, color=MUTED, bold=True)
    for lead, text in TERMS:
        para(text, lead=lead + ".", size=9.5, after=6)
    para(f"Free calculator to price the three cleans: {CALC}", size=9.5, color=TEAL, after=0)

    fp = sec.footer.paragraphs[0]
    r = fp.add_run(f"{PRODUCT}  |  ProofNotFluff  |  A sample form to adapt. Not legal advice."); r.font.size = Pt(8); r.font.color.rgb = MUTED

    cp = doc.core_properties
    now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    cp.author = "ProofNotFluff"; cp.last_modified_by = "ProofNotFluff"; cp.title = PRODUCT
    cp.subject = "Free template, version 1"; cp.keywords = "ProofNotFluff"; cp.comments = ""; cp.category = ""
    cp.created = now; cp.modified = now; cp.revision = 1
    buf = io.BytesIO(); doc.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO(); zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        if item.filename == "docProps/thumbnail.jpeg":
            continue
        data = zin.read(item.filename)
        if item.filename == "_rels/.rels":
            data = re.sub(rb'<Relationship Id="[^"]+" Type="[^"]+/thumbnail" Target="docProps/thumbnail.jpeg"/>', b"", data)
        if item.filename == "[Content_Types].xml":
            data = data.replace(b'<Default Extension="jpeg" ContentType="image/jpeg"/>', b"")
        if item.filename == "docProps/app.xml":
            data = data.replace(b"<Company/>", b"<Company>ProofNotFluff LLC</Company>")
            data = data.replace(b"<Application>Microsoft Macintosh Word</Application>", b"<Application>ProofNotFluff</Application>")
        zout.writestr(item, data)
    zout.close()
    open(path, "wb").write(out.getvalue())


# ------------------------------------------------------------------ pdf
def pdf_html(d):
    fonts = os.path.join(ROOT, "docs", "assets", "fonts")
    face = "".join(f'@font-face{{font-family:{fam};font-weight:{w};src:url("file://{fonts}/{f}") format("woff2")}}'
                   for fam, w, f in (("Inter", 400, "inter-400.woff2"), ("Inter", 600, "inter-600.woff2"), ("Inter", 700, "inter-700.woff2"),
                                     ("Poppins", 700, "poppins-700.woff2")))
    bid_css = open(os.path.join(ROOT, "docs", "assets", "pnf-bid.css"), encoding="utf-8").read()
    bid_css = bid_css.split("/* print only the proposal")[0]   # the document styles only
    css = face + bid_css + """
@page{size:Letter;margin:0}
*{box-sizing:border-box}
body{margin:0;font-family:Inter,Arial,sans-serif;color:#1d2433;background:#fff;font-size:10.5pt;line-height:1.45}
.page{width:8.5in;height:11in;padding:.7in .8in .75in;position:relative;page-break-after:always;overflow:hidden}
.page:last-child{page-break-after:auto}
.page h1,.page h2.x{font-family:Poppins,Arial,sans-serif;font-weight:700;letter-spacing:-.01em;margin:0}
.eyebrow{font-size:8pt;letter-spacing:.14em;text-transform:uppercase;color:#a63e21;font-weight:700;margin:0}
.page h1{font-size:23pt;line-height:1.1;margin:.06in 0 .1in}
.page h2.x{font-size:13pt;margin:.24in 0 .06in}
.lede{color:#4a5163;font-size:11pt;margin:0}
.bars{display:flex;gap:4px;margin:.08in 0 .04in}.bars i{display:block;width:.26in;height:4px;border-radius:2px}
.bars i:nth-child(1){background:#c8502f}.bars i:nth-child(2){background:#2e6b66}.bars i:nth-child(3){background:#6b4fa0}.bars i:nth-child(4){background:#2f5f9e}.bars i:nth-child(5){background:#a87c1f}
ol.steps{list-style:none;counter-reset:s;margin:0;padding:0}
ol.steps li{counter-increment:s;position:relative;padding:.06in 0 .06in .36in;border-bottom:1px solid #efebe3;font-size:10.2pt}
ol.steps li::before{content:counter(s);position:absolute;left:0;top:.07in;width:.23in;height:.23in;border-radius:50%;background:#2e6b66;color:#fff;font-weight:700;font-size:8.5pt;line-height:.23in;text-align:center}
.note{font-size:9.4pt;color:#4a5163;margin:.16in 0 0}
.small{font-size:8.2pt;color:#4a5163;margin:.12in 0 0}
.calc{margin-top:.16in;background:#e2efec;border-left:4px solid #2e6b66;border-radius:8px;padding:.1in .14in;font-size:10pt}
.terms p{font-size:9pt;margin:0 0 .075in}
.page .bid-doc{padding:0;max-width:none;font-size:11pt}
.foot{position:absolute;left:.8in;right:.8in;bottom:.4in;font-size:7.4pt;color:#5c6270;border-top:1px solid #e2ddd3;padding-top:.06in;display:flex;justify-content:space-between}
"""
    foot = lambda n: f'<div class="foot"><span>{PRODUCT} | ProofNotFluff | proofnotfluff.com</span><span>Page {n} of 3</span></div>'
    steps = "".join(f"<li><b>{html.escape(a)}</b> {html.escape(b)}</li>" for a, b in STEPS)
    terms = "".join(f"<p><b>{html.escape(a)}.</b> {html.escape(b)}</p>" for a, b in TERMS)
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{PRODUCT}</title><style>{css}</style></head><body>
<section class="page"><p class="eyebrow">Free template · ProofNotFluff</p><div class="bars"><i></i><i></i><i></i><i></i><i></i></div>
<h1>Post-construction cleaning proposal template</h1>
<p class="lede">A one-page proposal for a builder: rough, final and touch-up cleans, each with its scope and your price. The proposal is on page 2. To type yours, use the Word file: {DOCX_URL}</p>
<h2 class="x">How to price and send it</h2><ol class="steps">{steps}</ol>
<div class="calc">Price the three cleans first: the free bid calculator at <b>{CALC}</b> prices each clean from your own crew speed and costs and fills this same proposal for you to print.</div>
<p class="note">{html.escape(SAMPLE)}</p><p class="small">{html.escape(SHORT_NOTICE).replace("checked on ", "checked&nbsp;on&nbsp;")}</p>{foot(1)}</section>
<section class="page">{d["html"]}{foot(2)}</section>
<section class="page terms"><p class="eyebrow">Terms</p><h1 style="font-size:20pt">Terms of Use and Disclaimer</h1>
<p style="font-size:9.6pt;margin:0 0 .14in;color:#4a5163"><b>ProofNotFluff | {PRODUCT} | Version 1, {CHECKED}</b></p>{terms}{foot(3)}</section>
</body></html>"""


def render_pdf(html_text, path, title, subject):
    from playwright.sync_api import sync_playwright
    from pypdf import PdfReader, PdfWriter
    tmp = path + ".src.html"
    open(tmp, "w", encoding="utf-8").write(html_text)
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page()
        pg.goto("file://" + tmp, wait_until="load"); pg.evaluate("document.fonts.ready")
        raw = pg.pdf(format="Letter", print_background=True, prefer_css_page_size=True, tagged=True, outline=False)
        b.close()
    os.remove(tmp)
    r = PdfReader(io.BytesIO(raw)); w = PdfWriter()
    for page in r.pages:
        w.add_page(page)
    assert len(r.pages) == 3, f"expected 3 pages, got {len(r.pages)}"
    stamp = datetime.now(timezone.utc).strftime("D:%Y%m%d%H%M%S+00'00'")
    w.add_metadata({"/Title": title, "/Author": "ProofNotFluff LLC", "/Subject": subject, "/Creator": "ProofNotFluff LLC",
                    "/Producer": "ProofNotFluff LLC", "/Keywords": "ProofNotFluff", "/CreationDate": stamp, "/ModDate": stamp})
    if "/Lang" not in w._root_object:
        from pypdf.generic import NameObject, TextStringObject
        w._root_object[NameObject("/Lang")] = TextStringObject("en-US")
    with open(path, "wb") as f:
        w.write(f)


def find_jobs_band():
    """The calculator page's closing band (paid lists that fit a construction cleaner), so the thanks page ends on the same next step."""
    sys.path.insert(0, os.path.join(ROOT, "tools"))
    import make_site
    return make_site.find_customers(SLUG, "../../")


def main():
    d = from_js()
    check_text(d)
    os.makedirs(DL, exist_ok=True)
    build_docx(os.path.join(DL, NAME + ".docx"), d)
    render_pdf(pdf_html(d), os.path.join(DL, NAME + ".pdf"), PRODUCT, "Free template: a one-page post-construction cleaning proposal")
    if os.path.isfile(os.path.join(ROOT, "docs", SLUG, "index.html")):
        sys.path.insert(0, os.path.join(ROOT, "engines"))
        import thanks_page
        thanks_page.write(ROOT, SLUG, "post-construction cleaning proposal template",
                          [(NAME + ".docx", "Download the Word file", "The proposal to type in, with the how-to and Terms, 3 pages"),
                           (NAME + ".pdf", "Download the PDF", "Printable, 3 pages")],
                          "No Microsoft Word? Upload the .docx to Google Drive and open it with Google Docs. It works the same.",
                          (SLUG, "Price the three cleans with the free bid calculator"),
                          note="A sample form to adapt, not a legal document. Have a local attorney review it before you use it with clients.",
                          tail=find_jobs_band(), ui=True, crumb="Bid calculator")
    print("built", NAME + ".docx", NAME + ".pdf")


if __name__ == "__main__":
    main()
