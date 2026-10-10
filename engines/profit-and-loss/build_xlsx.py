#!/usr/bin/env python3
"""Profit and Loss Template for a small business, version 1 (v3 dashboard design, Oct 8, 2026).

One Entries log feeds a printable profit and loss statement (the month you pick and January to that
month), a month-by-month P&L for the year, and a Checks box. Categories follow the IRS 2025
Schedule C lines (organizing only). Cost of goods is what you bought for the products you sell;
two optional stock counts turn the year-to-date column into a true cost of goods sold
(Schedule C Part III: beginning stock + purchases - ending stock).

Built to design/WORKBOOK_STANDARD.md section 0 with tools/pnf_dash.py and tools/wb_style.py.
No TEXT() format strings, no locale-dependent functions (a counted competitor complaint).

  python3 build_xlsx.py Profit-and-Loss-EXAMPLE.xlsx           (the worked example)
  python3 build_xlsx.py Profit-and-Loss-Template.xlsx --blank  (blank)
"""
import datetime as dt
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402

PRODUCT = "Profit and Loss Template"
VERSION = "1"
CHECKED = "Oct 8, 2026"
CHECKED_LONG = "October 8, 2026"
AS_OF = f"Lines checked {CHECKED}"
SOURCES = "the IRS 2025 Schedule C (Form 1040) and its 2025 instructions"
HERE = os.path.dirname(os.path.abspath(__file__))
args = sys.argv[1:]
BLANK = "--blank" in args
args = [a for a in args if a != "--blank"]
OUT = args[0] if args else "Profit-and-Loss-EXAMPLE.xlsx"
USD2, INT, DATE = S.FMT["usd2"], S.FMT["int"], S.FMT["date"]
MONEY = '"$"#,##0.00;-"$"#,##0.00;"$"0.00'      # statement: exact to the cent, zero shows $0.00
MONEY_T = '"$"#,##0.00;-"$"#,##0.00;"-"'        # month table: a dash for zero so empty months stay quiet
PCT = '0.0%;-0.0%;0.0%'
YEAR_FMT = "0"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH_NAMES = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
               "October", "November", "December"]
ARR_MON = "{" + ",".join(f'"{m}"' for m in MONTHS) + "}"
ARR_NAME = "{" + ",".join(f'"{m}"' for m in MONTH_NAMES) + "}"

# ---------------------------------------------------------------- categories
# (counts as, category, Schedule C reference, what goes here)
CATS = [
    ("Revenue", "Sales and services", "Line 1, Gross receipts or sales",
     "Money customers paid you for products or services, before fees."),
    ("Revenue", "Other business income", "Line 6, Other income",
     "Business income that is not a sale, such as a fee refund paid back to you."),
    ("Refund, subtracts", "Refunds given to customers", "Line 2, Returns and allowances",
     "Money you paid back to customers. It is subtracted from revenue."),
    ("Cost of goods", "Stock bought for resale", "Part III, line 36, Purchases",
     "Finished goods or blanks you buy to sell, such as blank shirts or wholesale stock."),
    ("Cost of goods", "Materials for products", "Part III, line 38, Materials and supplies",
     "Materials that become part of what you sell: ink, fabric, wax, packaging that ships with the product."),
    ("Cost of goods", "Product labor and other costs", "Part III, lines 37 and 39",
     "People paid to make the products (never yourself) and other costs of making them."),
    ("Expense", "Advertising", "Line 8, Advertising", "Ads, promoted listings, printed flyers, business cards."),
    ("Expense", "Car and truck expenses", "Line 9, Car and truck expenses",
     "Business driving costs. Keep a mileage log too; your preparer picks the method."),
    ("Expense", "Commissions and fees", "Line 10, Commissions and fees", "Selling fees and commissions charged on your sales."),
    ("Expense", "Contract labor", "Line 11, Contract labor", "People you paid who are not employees, for work other than making products."),
    ("Expense", "Employee benefit programs", "Line 14, Employee benefit programs", "Benefit programs for employees (not your own)."),
    ("Expense", "Insurance (other than health)", "Line 15, Insurance (other than health)", "Business liability, product or property insurance."),
    ("Expense", "Interest", "Lines 16a and 16b, Interest", "Interest on business loans and business credit cards."),
    ("Expense", "Legal and professional services", "Line 17, Legal and professional services", "Accountant, tax preparer, lawyer, bookkeeper."),
    ("Expense", "Office expense", "Line 18, Office expense", "Postage, printer ink, small office items."),
    ("Expense", "Rent or lease: vehicles, equipment", "Line 20a, Vehicles, machinery, and equipment", "Leased vehicles, machines or equipment."),
    ("Expense", "Rent or lease: other property", "Line 20b, Other business property", "Studio, booth, storage unit or shop rent."),
    ("Expense", "Repairs and maintenance", "Line 21, Repairs and maintenance", "Fixing business equipment or space."),
    ("Expense", "Supplies", "Line 22, Supplies (not included in Part III)", "Supplies used up in the business that are not part of what you sell."),
    ("Expense", "Taxes and licenses", "Line 23, Taxes and licenses", "Business licenses, permits, business taxes."),
    ("Expense", "Travel", "Line 24a, Travel", "Overnight business trips: airfare, lodging."),
    ("Expense", "Meals (enter the full bill)", "Line 24b, Deductible meals",
     "Business meals. Enter the full bill; the IRS instructions say in most cases only 50% is deductible, and your preparer applies the limit."),
    ("Expense", "Utilities", "Line 25, Utilities", "Business phone line, internet, power for a business space."),
    ("Expense", "Wages", "Line 26, Wages (less employment credits)", "Pay to employees (not to yourself)."),
    ("Expense", "Other expenses", "Line 27b, Other expenses", "Software, subscriptions, bank fees, shipping labels and anything else."),
    ("Not in profit", "Equipment over a few hundred dollars", "Line 13 (depreciation), for your preparer",
     "Computers, machines, furniture. Logged so your preparer can decide on depreciation. Not counted in profit."),
    ("Not in profit", "Owner draw or personal", "Not a business expense",
     "Money you took out for yourself, or a personal purchase made by mistake."),
    ("Not in profit", "Transfer between accounts", "Not income or expense",
     "Moving money between your own accounts or paying off a business card."),
]
REV_IDX, REF_IDX, COGS_IDX, EXP_IDX, OFF_IDX = (0, 1), (2,), (3, 4, 5), tuple(range(6, 25)), (25, 26, 27)
assert len(CATS) == 28 and CATS[EXP_IDX[-1]][1] == "Other expenses"

# ---------------------------------------------------------------- example: a made-up one-person T-shirt printing shop, Jan to Sep 2026
EX_YEAR, EX_MONTH, EX_STOCK0, EX_STOCK1 = 2026, 9, 1200.0, 1450.0
online = [2400, 2150, 2800, 3100, 3350, 2900, 2700, 3050, 3400]
custom = {3: 1200, 5: 950, 8: 1500}
refunds = [60, 0, 45, 0, 90, 0, 30, 0, 75]
blanks = [900, 600, 1100, 1000, 1300, 900, 800, 1200, 1100]
ink = [180, 0, 220, 150, 0, 240, 0, 200, 160]
fees = [240, 215, 280, 310, 335, 290, 270, 305, 340]
labels = [310, 280, 360, 400, 430, 370, 350, 390, 440]
ads = [50, 50, 75, 75, 100, 50, 50, 75, 100]
office = [0, 22.40, 0, 0, 18.75, 0, 0, 31.10, 0]
supplies = [60, 0, 75, 0, 80, 0, 70, 0, 85]
d = dt.datetime
ROWS = []  # date, what it was, customer or vendor, category, amount, note
for i in range(9):
    m = i + 1
    ROWS.append((d(2026, m, 2), "Blank shirts order", "Garment wholesaler", "Stock bought for resale", blanks[i], None))
    ROWS.append((d(2026, m, 5), "Online shop sales, first half", "Customers", "Sales and services", online[i] / 2, None))
    ROWS.append((d(2026, m, 5), "Selling fees, first half", "Marketplace", "Commissions and fees", fees[i] / 2, None))
    if ink[i]:
        ROWS.append((d(2026, m, 8), "Ink and screens", "Print supplier", "Materials for products", ink[i], None))
    ROWS.append((d(2026, m, 10), "Shipping labels", "Postal service", "Other expenses", labels[i], None))
    ROWS.append((d(2026, m, 12), "Design software", "Software company", "Other expenses", 29.99, "Monthly plan"))
    ROWS.append((d(2026, m, 15), "Studio rent", "Landlord", "Rent or lease: other property", 450, None))
    ROWS.append((d(2026, m, 16), "Promoted listings", "Marketplace", "Advertising", ads[i], None))
    ROWS.append((d(2026, m, 18), "Liability insurance", "Insurer", "Insurance (other than health)", 42, None))
    ROWS.append((d(2026, m, 18), "Phone and internet, business share", "Phone company", "Utilities", 65, None))
    if m in custom:
        ROWS.append((d(2026, m, 19), "Team shirts, local league", "League", "Sales and services", custom[m], "Custom order"))
        ROWS.append((d(2026, m, 19), "Helper for the custom order", "Helper", "Product labor and other costs",
                     {3: 300, 5: 250, 8: 400}[m], None))
    ROWS.append((d(2026, m, 20), "Online shop sales, second half", "Customers", "Sales and services", online[i] / 2, None))
    ROWS.append((d(2026, m, 20), "Selling fees, second half", "Marketplace", "Commissions and fees", fees[i] / 2, None))
    if refunds[i]:
        ROWS.append((d(2026, m, 22), "Refund, misprint", "Customer", "Refunds given to customers", refunds[i], None))
    if office[i]:
        ROWS.append((d(2026, m, 23), "Printer paper and postage", "Office store", "Office expense", office[i], None))
    if supplies[i]:
        ROWS.append((d(2026, m, 24), "Cleaning rags and tape", "Hardware store", "Supplies", supplies[i], None))
    ROWS.append((d(2026, m, 28), "Paid myself", "Owner", "Owner draw or personal", 1500, None))
ROWS += [
    (d(2026, 1, 14), "City business license", "City clerk", "Taxes and licenses", 75, None),
    (d(2026, 3, 25), "Tax prep", "Accountant", "Legal and professional services", 250, None),
    (d(2026, 4, 6), "Heat press", "Equipment dealer", "Equipment over a few hundred dollars", 1850, None),
    (d(2026, 6, 11), "Fee refund from marketplace", "Marketplace", "Other business income", 40, None),
    (d(2026, 6, 30), "Paid off business card", "Card company", "Transfer between accounts", 2000, None),
    (d(2026, 7, 9), "Press repair", "Repair shop", "Repairs and maintenance", 85.50, None),
]
ROWS.sort(key=lambda x: x[0])


def _sum(idx, m0=1, m1=12):
    names = [CATS[i][1] for i in idx]
    return round(sum(r[4] for r in ROWS if r[3] in names and m0 <= r[0].month <= m1), 2)


def ex_pl(m0, m1, stock=False):
    rev = _sum(REV_IDX, m0, m1) - _sum(REF_IDX, m0, m1)
    cogs = _sum(COGS_IDX, m0, m1) + ((EX_STOCK0 - EX_STOCK1) if stock else 0)
    exp = _sum(EXP_IDX, m0, m1)
    return dict(rev=round(rev, 2), cogs=round(cogs, 2), gp=round(rev - cogs, 2), exp=round(exp, 2), np=round(rev - cogs - exp, 2))


EX_MON = ex_pl(9, 9)
EX_YTD = ex_pl(1, 9, stock=True)
EX_AUG = ex_pl(8, 8)
EXAMPLE = dict(mon=EX_MON, ytd=EX_YTD, aug=EX_AUG)
if BLANK:
    ROWS = []

# ---------------------------------------------------------------- fixed addresses
ENT_N, STM_N, MBM_N, CAT_N = "1 Entries", "2 Statement", "3 Month by Month", "4 Categories"
ENT, STM, MBM, CAT = (f"'{n}'" for n in (ENT_N, STM_N, MBM_N, CAT_N))
N_ROWS = 1000
E_TOP = 11
E_FIRST = E_TOP + 5
E_LAST = E_FIRST + N_ROWS - 1
E_CHECK_LAST = E_FIRST + 4999
# bounded entry ranges: whole columns would include Entries C8 (which reads the Statement) and loop
RG_C = f"{ENT}!$C${E_FIRST}:$C${E_CHECK_LAST}"
RG_F = f"{ENT}!$F${E_FIRST}:$F${E_CHECK_LAST}"
RG_G = f"{ENT}!$G${E_FIRST}:$G${E_CHECK_LAST}"
C_TOP = 11
C_FIRST = C_TOP + 5
C_LAST = C_FIRST + len(CATS) - 1
CAT_LIST = f"{CAT}!$D${C_FIRST}:$D${C_LAST}"


def catref(i):
    return f"{CAT}!$D${C_FIRST + i}"


def sign_rule(ws, ref):
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{ref}<0"], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{ref}>0"], font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))


def rich(head, body, size=9):
    return CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=size, b=True, color=D.T["ink"]), head + ". "),
                        TextBlock(InlineFont(rFont=D.FONT, sz=size, color=D.T["ink2"]), body))


wb = S.new_workbook(PRODUCT, version=VERSION)

# =====================================================================  2 Statement
st = wb.create_sheet(STM_N)
S.set_widths(st, {"A": 3, "B": 2, "C": 40, "D": 2, "E": 2, "F": 2, "G": 16, "H": 16, "I": 14, "J": 2, "K": 3})
D.page_header(st, "Profit and loss statement", "Fills itself in from the Entries tab: the month you pick and the year to date.",
              AS_OF, span=("B", "J"), side_cols=3)

# ---- period card (rows 11+)
P = D.Card(st, 11, "B", "C", "G", "J", wb=wb)
P.title("Period and stock counts", "Pick the year and month. Stock counts are optional.")
r_y = P.input("Year", None if BLANK else EX_YEAR, YEAR_FMT, name="YearShown", validation=("whole", 2000, 2100), allow_blank=True,
              prompt="Leave blank for this year. Type a year, for example 2025, to look back.", prompt_title="Year")
r_m = P.input("Month the statement runs to", None if BLANK else MONTHS[EX_MONTH - 1], "@", name="MonthShown",
              validation=("list", MONTHS), allow_blank=True, prompt_title="Month",
              prompt="Pick a month from the list. Leave blank for this month (or December when looking back at an earlier year).")
r_s0 = P.input("Stock on hand on January 1", None if BLANK else EX_STOCK0, USD2, name="StockStart",
               validation=("decimal", 0, None), allow_blank=True, prompt_title="Stock on Jan 1 (optional)",
               prompt="Optional. What your unsold stock and materials cost you, counted at the start of the year, for example 1200.")
r_s1 = P.input("Stock on hand at the end of that month", None if BLANK else EX_STOCK1, USD2, name="StockEnd",
               validation=("decimal", 0, None), allow_blank=True, prompt_title="Stock at month end (optional)",
               prompt="Optional. What your unsold stock and materials cost you, counted at the end of the month shown, for example 1450.")
for rr in (r_y, r_m):
    st[f"G{rr}"].alignment = Alignment(horizontal="center", vertical="center")
Y = f"$H${r_y}"
M = f"$H${r_m}"
hy = st[f"H{r_y}"]; hy.value = f'=IF(ISNUMBER($G${r_y}),$G${r_y},YEAR(TODAY()))'; hy.number_format = YEAR_FMT
hm = st[f"H{r_m}"]
hm.value = f'=IF($G${r_m}="",IF({Y}=YEAR(TODAY()),MONTH(TODAY()),12),IFERROR(MATCH($G${r_m},{ARR_MON},0),IF({Y}=YEAR(TODAY()),MONTH(TODAY()),12)))'
hm.number_format = INT
for c in (hy, hm):
    c.font = D.f(9, False, "muted"); c.alignment = Alignment(horizontal="center", vertical="center")
    c.number_format = ";;;"   # helper numbers the formulas read; the words in column I say the same thing
# I column: plain words beside the two period inputs
iy = st[f"I{r_y}"]; iy.value = f'="Showing "&{Y}'
im = st[f"I{r_m}"]; im.value = f'=INDEX({ARR_NAME},{M})'
for c in (iy, im):
    c.font = D.f(9, False, "muted"); c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
EC = f"{ENT}!$C${E_FIRST}:$C${E_CHECK_LAST}"
HAS_YTD = f'COUNTIFS({EC},">="&DATE({Y},1,1),{EC},"<"&DATE({Y},{M}+1,1))>0'
BOTH = f'AND($G${r_s0}<>"",$G${r_s1}<>"")'
STOCK_ON = f'AND({BOTH},{HAS_YTD})'
STOCK_ONE = f'AND(OR($G${r_s0}<>"",$G${r_s1}<>""),NOT({BOTH}))'
r_ps = P.text(f'=IF({STOCK_ON},"Both counts in: the year-to-date cost of goods uses them",IF({BOTH},"Stock counts wait for your first entry in this period",'
              f'IF({STOCK_ONE},"Add the other stock count too, or the counts are left out",'
              f'"No stock counts: cost of goods is what you bought")))', size=9, color="ink2", bold=True)
P.close()
D.status_rule(st, f"C{r_ps}", f"OR({STOCK_ONE},AND({BOTH},NOT({HAS_YTD})))", fill_tint=False)


def sumifs(i, m_from, m_to_excl):
    return (f'IF({catref(i)}="",0,SUMIFS({RG_G},{RG_F},{catref(i)},'
            f'{RG_C},">="&DATE({Y},{m_from},1),{RG_C},"<"&DATE({Y},{m_to_excl},1)))')


# ---- the statement card
S_TOP = P.end + 2
r = S_TOP
D.h(st, r, 12); r += 1
t = st[f"C{r}"]; t.value = f'="Profit and loss, "&INDEX({ARR_NAME},{M})&" "&{Y}'
t.font = D.f(12, True); st.merge_cells(f"C{r}:I{r}"); D.h(st, r, 24); r += 1
t = st[f"C{r}"]
t.value = "Exact to the cent. Refunds and costs subtract. Estimates only, see the Terms tab."
t.font = D.f(9, False, "muted"); st.merge_cells(f"C{r}:I{r}"); D.h(st, r, 18); r += 1
D.h(st, r, 6); r += 1
HEAD = r
heads = {"C": "", "G": f'=UPPER(INDEX({ARR_NAME},{M}))',
         "H": f'=IF({M}=1,"JANUARY","JAN TO "&UPPER(INDEX({ARR_MON},{M})))', "I": "% OF REVENUE"}
for k, v in heads.items():
    c = st[f"{k}{HEAD}"]; c.value = v; c.font = D.f(8, True, "muted")
    c.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
for k in D.cols("C", "I"):
    st[f"{k}{HEAD}"].border = Border(bottom=D.side("ink"))
D.h(st, HEAD, 24)
r += 1
SROW = {}
TOTALS = {}
PCT_ROWS = []


def s_eyebrow(text):
    global r
    c = st[f"C{r}"]; c.value = text; c.font = D.f(8, True, "muted"); c.alignment = Alignment(vertical="bottom")
    for k in D.cols("C", "I"):
        st[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    D.h(st, r, D.GRID_ROW); r += 1


def s_money_row(label, g, h_, bold=False, total=False, tint=None, indent=0):
    global r
    c = st[f"C{r}"]; c.value = label; c.font = D.f(10, bold or total); c.alignment = Alignment(vertical="center", indent=indent)
    for k, v in (("G", g), ("H", h_)):
        cell = st[f"{k}{r}"]; cell.value = v; cell.number_format = MONEY; cell.font = D.f(10, bold or total)
        cell.alignment = Alignment(horizontal="right", vertical="center")
    PCT_ROWS.append(r)
    for k in D.cols("C", "I"):
        cell = st[f"{k}{r}"]
        cell.border = Border(top=D.side("ink") if total else None, bottom=D.side("hair"))
        if tint:
            cell.fill = D.fill(tint)
    D.h(st, r, D.GRID_ROW)
    out = r
    r += 1
    return out


def s_cat(i, negative=False):
    sign = "-" if negative else ""
    rr = s_money_row(f'=IF({catref(i)}="","",{catref(i)})', f"={sign}{sumifs(i, M, f'{M}+1')}",
                     f"={sign}{sumifs(i, 1, f'{M}+1')}")
    SROW[i] = rr
    return rr


s_eyebrow("REVENUE")
for i in REV_IDX:
    s_cat(i)
s_cat(REF_IDX[0], negative=True)
a, b = SROW[REV_IDX[0]], SROW[REF_IDX[0]]
R_REV = s_money_row("Net revenue", f"=SUM(G{a}:G{b})", f"=SUM(H{a}:H{b})", total=True)
s_eyebrow("COST OF GOODS SOLD")
for i in COGS_IDX:
    s_cat(i)
R_STK = s_money_row("Change in stock on hand (from your counts)", '=""',
                    f'=IF({STOCK_ON},$G${r_s0}-$G${r_s1},0)')
st[f"G{R_STK}"].value = None
gm = st[f"G{R_STK}"]; gm.value = "year only"; gm.font = D.f(9, False, "muted")
gm.alignment = Alignment(horizontal="right", vertical="center")
a = SROW[COGS_IDX[0]]
R_COGS = s_money_row("Cost of goods sold", f"=SUM(G{a}:G{R_STK - 1})", f"=SUM(H{a}:H{R_STK})", total=True)
R_GP = s_money_row("Gross profit", f"=G{R_REV}-G{R_COGS}", f"=H{R_REV}-H{R_COGS}", total=True, tint="total_tint")
s_eyebrow("OPERATING EXPENSES")
for i in EXP_IDX:
    s_cat(i)
a, b = SROW[EXP_IDX[0]], SROW[EXP_IDX[-1]]
R_EXP = s_money_row("Total operating expenses", f"=SUM(G{a}:G{b})", f"=SUM(H{a}:H{b})", total=True)
R_NP = s_money_row("Net profit", f"=G{R_GP}-G{R_EXP}", f"=H{R_GP}-H{R_EXP}", total=True, tint="total_tint")
for k in ("G", "H"):
    sign_rule(st, f"{k}{R_GP}"); sign_rule(st, f"{k}{R_NP}")
for rr in (R_NP,):
    for k in ("C", "G", "H", "I"):
        st[f"{k}{rr}"].font = D.f(11, True)
    D.h(st, rr, D.GRID_ROW + 3)
# margins
MARGINS = []
for label, num in (("Gross margin", R_GP), ("Net margin", R_NP)):
    c = st[f"C{r}"]; c.value = label; c.font = D.f(10, False, "ink2"); c.alignment = Alignment(vertical="center", indent=1)
    for k in ("G", "H"):
        cell = st[f"{k}{r}"]; cell.value = f'=IF({k}{R_REV}=0,"",{k}{num}/{k}{R_REV})'; cell.number_format = PCT
        cell.font = D.f(10, False, "ink2"); cell.alignment = Alignment(horizontal="right", vertical="center")
    for k in D.cols("C", "I"):
        st[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    D.h(st, r, D.GRID_ROW); MARGINS.append(r); r += 1
for rr in PCT_ROWS:
    cell = st[f"I{rr}"]; cell.value = f'=IF(H${R_REV}=0,"",H{rr}/H${R_REV})'; cell.number_format = PCT
    cell.font = D.f(9, False, "muted"); cell.alignment = Alignment(horizontal="right", vertical="center")
st[f"I{R_REV}"].value = f'=IF(H${R_REV}=0,"",1)'
s_eyebrow("LOGGED FOR YOUR RECORDS, NOT IN PROFIT")
for i in OFF_IDX:
    s_cat(i)
    st[f"I{SROW[i]}"].value = None
    for k in ("C", "G", "H"):
        st[f"{k}{SROW[i]}"].font = D.f(10, False, "ink2")
D.h(st, r, 12)
S_END = r
D.frame(st, S_TOP, S_END, "B", "J")
S.page_break_before(st, S_TOP)

# ---- checks card
EA = f"{ENT}!$G${E_FIRST}:$G${E_CHECK_LAST}"
ED = f"{ENT}!$F${E_FIRST}:$F${E_CHECK_LAST}"
EC = f"{ENT}!$C${E_FIRST}:$C${E_CHECK_LAST}"
K = D.Card(st, S_END + 2, "B", "C", "G", "J", wb=wb)
K.title("Checks", "The first three should read 0. Rows in other years are fine.")
k1 = K.calc("Amount but no category", f'=COUNTIFS({EA},"<>",{ED},"")', INT)
k2 = K.calc("Amount but no date", f'=COUNTIFS({EA},"<>",{EC},"")', INT)
k3 = K.calc("Category not on the list", f'=SUMPRODUCT(({ED}<>"")*(COUNTIF({CAT_LIST},{ED})=0))', INT)
k4 = K.calc("Dated in another year", f'=COUNT({EC})-COUNTIFS({EC},">="&DATE({Y},1,1),{EC},"<"&DATE({Y}+1,1,1))', INT)
k5 = K.calc("Problems to fix", f"=G{k1}+G{k2}+G{k3}", INT, total=True)
for rr in (k1, k2, k3, k5):
    D.status_rule(st, f"G{rr}", f"$G${rr}>0", fill_tint=False)
FIX = f"$G${k5}"
pill = K.text(f'=IF(COUNT({EA})=0,"No entries yet",IF({FIX}>0,{FIX}&IF({FIX}=1," problem"," problems")&" to fix before the totals are complete",'
              f'IF(G{k4}>0,"Every row is counted. "&G{k4}&IF(G{k4}=1," row is"," rows are")&" in other years, which is fine","Every row is counted")))',
              size=10, color="ink", bold=True, height=D.GRID_ROW * 2)
st[f"C{pill}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
K.text(f"The checks and totals read entry rows {E_FIRST} to {E_CHECK_LAST:,}.", size=9, color="muted")
K.close()
D.status_rule(st, f"C{pill}:G{pill}", f"{FIX}>0")
S.page_break_before(st, K.top)

# ---- tiles (rows 6 to 9)
NO_ENTRIES = f"COUNT({EA})=0"
D.tile(st, 6, "B", "D",
       f'="NET PROFIT, "&IF({M}=1,"JANUARY","JANUARY TO "&UPPER(INDEX({ARR_NAME},{M})))&" "&{Y}',
       f"=$H${R_NP}", MONEY,
       f'=IF(COUNTIFS({EC},">="&DATE({Y},1,1),{EC},"<"&DATE({Y},{M}+1,1))=0,"No entries for this period yet",IF($H${R_REV}=0,"No revenue logged for this period yet","Net margin "&ROUND($H${MARGINS[1]}*100,1)&"% of net revenue"'
       f'&IF({STOCK_ON},", after your stock counts","")))', dark=True)
PREV = f"INDEX($D${{np}}:$O${{np}},{M}-1)"   # filled after the month tab exists
v2, sub2, _ = D.tile(st, 6, "F", "J", f'="NET PROFIT, "&UPPER(INDEX({ARR_NAME},{M}))&" "&{Y}', f"=$G${R_NP}", MONEY, "", dark=False)
sign_rule(st, v2.coordinate)
D.h(st, 10, 14)

# ---- notes and sources
N_TOP = K.end + 2
NOTES = [
    ("Categories", "The names follow the lines on the IRS 2025 Schedule C (Form 1040), irs.gov/pub/irs-pdf/f1040sc.pdf, "
     f"checked {CHECKED}. They are for organizing, not tax advice; your preparer decides the final line for each item."),
    ("Cost of goods sold", "Schedule C Part III works it out as stock at the start of the year plus what you bought, less stock at "
     "the end. With both counts filled in, the year-to-date column does the same. Without them, cost of goods is simply what you "
     "bought for the products in that period."),
    ("Business meals", "Enter the full bill. The 2025 Schedule C instructions (irs.gov/pub/irs-pdf/i1040sc.pdf) say that in most "
     f"cases only 50% of business meals is deductible; your preparer applies the limit. Checked {CHECKED}."),
    ("Not calculated here", "Depreciation and section 179 (line 13), depletion (line 12), pension plans (line 19) and business use "
     "of your home (line 30). Equipment, owner draws and transfers are logged but kept out of profit."),
    ("Any currency", "The maths works in any currency. Select the money cells and pick your symbol under Format, Number."),
    ("Before you rely on a number", S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab") + " "
     + S.TAX_TIER_CLAUSE.format(date=CHECKED_LONG, sources=SOURCES)),
]
r = N_TOP
D.h(st, r, 12); r += 1
t = st[f"C{r}"]; t.value = "Notes and sources"; t.font = D.f(12, True); st.merge_cells(f"C{r}:I{r}"); D.h(st, r, 24); r += 1
width = sum(S.width_of(st, k) for k in D.cols("C", "I"))
for head, body in NOTES:
    cell = st[f"C{r}"]; cell.value = rich(head, body)
    cell.alignment = Alignment(vertical="top", wrap_text=True); st.merge_cells(f"C{r}:I{r}")
    D.h(st, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
D.h(st, r, 12)
D.frame(st, N_TOP, r, "B", "J")
ST_LAST = r + 1
D.h(st, ST_LAST, 14)
for rr in range(1, ST_LAST):
    if st.row_dimensions[rr].height is None:
        D.h(st, rr, D.GRID_ROW)
D.paint_canvas(st, ST_LAST, "Z")
S.finish_sheet(st, PRODUCT, ST_LAST, span=("A", "K"), freeze="A11", tab_color="accent")
# print: the whole statement card on one Letter or A4 page (page breaks sit before it and before Checks)
_stmt_pt = sum(st.row_dimensions[rr].height or D.GRID_ROW for rr in range(S_TOP, S_END + 1))
st.sheet_properties.pageSetUpPr.fitToPage = False
st.page_setup.scale = max(40, min(100, int(680 / _stmt_pt * 100)))

# =====================================================================  3 Month by Month
mb = wb.create_sheet(MBM_N)
MCOLS = D.cols("D", "O")
S.set_widths(mb, {"A": 3, "B": 2, "C": 34, **{k: 11.5 for k in MCOLS}, "P": 13, "Q": 2, "R": 3})
D.page_header(mb, "Month by month", "Every month of the year on the Statement tab, side by side.", AS_OF, span=("B", "Q"), side_cols=4)
D.h(mb, 6, 10)
T_TOP = 11
r = T_TOP
D.h(mb, r, 12); r += 1
t = mb[f"C{r}"]; t.value = f'="Profit and loss by month, "&{STM}!{Y}'; t.font = D.f(12, True); mb.merge_cells(f"C{r}:P{r}"); D.h(mb, r, 24); r += 1
t = mb[f"C{r}"]; t.value = ("Cost of goods here is what you bought each month; stock counts change only the year-to-date column "
                            "on the Statement tab.")
t.font = D.f(9, False, "muted"); mb.merge_cells(f"C{r}:P{r}"); D.h(mb, r, 18); r += 1
D.h(mb, r, 6); r += 1
MHEAD = r
for k, txt in zip(["C", *MCOLS, "P"], ["", *[m.upper() for m in MONTHS], "YEAR"]):
    c = mb[f"{k}{MHEAD}"]; c.value = txt; c.font = D.f(8, True, "muted")
    c.alignment = Alignment(horizontal="right", vertical="bottom")
for k in D.cols("C", "P"):
    mb[f"{k}{MHEAD}"].border = Border(bottom=D.side("ink"))
D.h(mb, MHEAD, 24)
r += 1
MROW = {}
YM = f"{STM}!{Y}"


def m_eyebrow(text):
    global r
    c = mb[f"C{r}"]; c.value = text; c.font = D.f(8, True, "muted"); c.alignment = Alignment(vertical="bottom")
    for k in D.cols("C", "P"):
        mb[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    D.h(mb, r, D.GRID_ROW); r += 1


def m_row(label, fn, total=False, tint=None, muted=False):
    global r
    c = mb[f"C{r}"]; c.value = label; c.font = D.f(10, total, "ink2" if muted else "ink"); c.alignment = Alignment(vertical="center")
    for m, k in enumerate(MCOLS, 1):
        v = mb[f"{k}{r}"]; v.value = fn(k, m); v.number_format = MONEY_T; v.font = D.f(10, total, "ink2" if muted else "ink")
        v.alignment = Alignment(horizontal="right", vertical="center")
    y = mb[f"P{r}"]; y.value = f"=SUM(D{r}:O{r})"; y.number_format = MONEY_T; y.font = D.f(10, True)
    y.alignment = Alignment(horizontal="right", vertical="center")
    for k in D.cols("C", "P"):
        cell = mb[f"{k}{r}"]
        cell.border = Border(top=D.side("ink") if total else None, bottom=D.side("hair"))
        if tint:
            cell.fill = D.fill(tint)
    D.h(mb, r, D.GRID_ROW + (3 if tint else 0))
    out = r
    r += 1
    return out


def m_cat(i, negative=False, muted=False):
    sign = "-" if negative else ""
    cr = catref(i)
    rr = m_row(f'=IF({cr}="","",{cr})',
               lambda k, m: (f'={sign}IF({cr}="",0,SUMIFS({RG_G},{RG_F},{cr},{RG_C},">="&DATE({YM},{m},1),'
                             f'{RG_C},"<"&DATE({YM},{m + 1},1)))'), muted=muted)
    MROW[i] = rr
    return rr


m_eyebrow("REVENUE")
for i in REV_IDX:
    m_cat(i)
m_cat(REF_IDX[0], negative=True)
a, b = MROW[REV_IDX[0]], MROW[REF_IDX[0]]
M_REV = m_row("Net revenue", lambda k, m: f"=SUM({k}{a}:{k}{b})", total=True)
m_eyebrow("COST OF GOODS (WHAT YOU BOUGHT)")
for i in COGS_IDX:
    m_cat(i)
a, b = MROW[COGS_IDX[0]], MROW[COGS_IDX[-1]]
M_COGS = m_row("Cost of goods", lambda k, m: f"=SUM({k}{a}:{k}{b})", total=True)
M_GP = m_row("Gross profit", lambda k, m: f"={k}{M_REV}-{k}{M_COGS}", total=True, tint="total_tint")
m_eyebrow("OPERATING EXPENSES")
for i in EXP_IDX:
    m_cat(i)
a, b = MROW[EXP_IDX[0]], MROW[EXP_IDX[-1]]
M_EXP = m_row("Total operating expenses", lambda k, m: f"=SUM({k}{a}:{k}{b})", total=True)
M_NP = m_row("Net profit", lambda k, m: f"={k}{M_GP}-{k}{M_EXP}", total=True, tint="total_tint")
for k in [*MCOLS, "P"]:
    sign_rule(mb, f"{k}{M_GP}"); sign_rule(mb, f"{k}{M_NP}")
c = mb[f"C{r}"]; c.value = "Net margin"; c.font = D.f(10, False, "ink2"); c.alignment = Alignment(vertical="center", indent=1)
for k in [*MCOLS, "P"]:
    v = mb[f"{k}{r}"]; v.value = f'=IF({k}{M_REV}=0,"",{k}{M_NP}/{k}{M_REV})'; v.number_format = PCT
    v.font = D.f(9, False, "ink2"); v.alignment = Alignment(horizontal="right", vertical="center")
for k in D.cols("C", "P"):
    mb[f"{k}{r}"].border = Border(bottom=D.side("hair"))
D.h(mb, r, D.GRID_ROW); M_NM = r; r += 1
m_eyebrow("LOGGED FOR YOUR RECORDS, NOT IN PROFIT")
for i in OFF_IDX:
    m_cat(i, muted=True)
D.h(mb, r, 12)
T_END = r
D.frame(mb, T_TOP, T_END, "B", "Q")

# ---- headline strip, rows 6 to 9
MPR = f"$D${M_NP}:$O${M_NP}"
D.stat_strip(mb, 7, [
    ("C:C", "Net revenue, year", f"=$P${M_REV}", USD2),
    ("D:F", "Gross profit, before counts", f"=$P${M_GP}", USD2),
    ("G:I", "Net profit, before counts", f"=$P${M_NP}", USD2),
    ("J:M", "Best month", f'=IF(COUNTIF({MPR},"<>0")=0,"None yet",INDEX({ARR_NAME},MATCH(MAX({MPR}),{MPR},0)))', "@"),
    ("N:P", "Loss months", f'=COUNTIF({MPR},"<0")', INT),
])
D.h(mb, 9, 10)
D.frame(mb, 6, 9, "B", "Q")
sign_rule(mb, "G8")
D.status_rule(mb, "N8", "$N$8>0", fill_tint=False)
D.h(mb, 10, 14)

# ---- net profit by month card with bars
B_TOP = T_END + 2
B = D.Card(mb, B_TOP, "B", "C", "D", "Q", extra="E", wb=wb)
B.title("Net profit by month", "Teal bars are profit, accent bars are losses. Purchases-based cost of goods.")
b_first = B.r
BR = f"$D${b_first}:$D${b_first + 11}"
for m, k in enumerate(MCOLS):
    rr = B.calc(MONTH_NAMES[m], f"={k}{M_NP}", MONEY_T)
    bc = mb[f"E{rr}"]
    bc.value = f'=IFERROR(REPT("{D.BAR_CHAR}",ROUND(ABS(D{rr})/MAX(MAX({BR}),-MIN({BR}))*40,0)),"")'
    bc.font = D.f(9, False, "teal"); bc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    mb.conditional_formatting.add(f"E{rr}", FormulaRule(formula=[f"$D{rr}<0"], font=Font(color=D.T["accent"]), stopIfTrue=True))
    sign_rule(mb, f"D{rr}")
LOSS = f'COUNTIF({BR},"<0")'
B.text(f'=IF(COUNTIF({BR},"<>0")=0,"No entries in this year yet",IF({LOSS}=0,"No month shows a loss",'
       f'{LOSS}&IF({LOSS}=1," month shows"," months show")&" a loss"))', size=9, color="ink2")
B.close()
for rr in range(b_first, b_first + 12):
    mb.merge_cells(f"E{rr}:P{rr}")
S.page_break_before(mb, B_TOP)
MB_LAST = B.end + 1
D.h(mb, MB_LAST, 14)
for rr in range(1, MB_LAST):
    if mb.row_dimensions[rr].height is None:
        D.h(mb, rr, D.GRID_ROW)
D.paint_canvas(mb, MB_LAST, "Z")
S.finish_sheet(mb, PRODUCT, MB_LAST, span=("A", "R"), freeze="A10", tab_color="blue")
mb.page_setup.orientation = "landscape"
mb.print_title_rows = f"{MHEAD}:{MHEAD}"

# right tile sub line on the Statement: compare with the month before
np_ref = f"{MBM}!$D${M_NP}:$O${M_NP}"
prev = f"INDEX({np_ref},{M}-1)"
sub2.value = (f'=IF(COUNTIFS({EC},">="&DATE({Y},{M},1),{EC},"<"&DATE({Y},{M}+1,1))=0,"No entries in "&INDEX({ARR_NAME},{M})&" yet",IF({M}=1,"First month of the year",'
              f'IF($G${R_NP}>{prev},"Up from "&INDEX({ARR_NAME},{M}-1),IF($G${R_NP}<{prev},"Down from "&INDEX({ARR_NAME},{M}-1),'
              f'"Same as "&INDEX({ARR_NAME},{M}-1)))))')

# =====================================================================  1 Entries
en = wb.create_sheet(ENT_N, 0)
S.set_widths(en, {"A": 3, "B": 2, "C": 13, "D": 32, "E": 22, "F": 36, "G": 13, "H": 26, "I": 20, "J": 2, "K": 3})
D.page_header(en, "Entries", "One row per money in or money out. Type amounts as positive numbers; the category decides the sign.",
              "Example rows: type over them" if not BLANK else f"Type your first row in row {E_FIRST}", span=("B", "J"), side_cols=3)
D.h(en, 6, 10)
D.stat_strip(en, 7, [
    ("C:D", "Net revenue, year to date", f"={STM}!$H${R_REV}", USD2),
    ("E:F", "Net profit, year to date", f"={STM}!$H${R_NP}", USD2),
    ("G:H", "Rows logged", f"=COUNT({EA})", INT),
    ("I:I", "Problems to fix", f"={STM}!{FIX}", INT),
])
D.h(en, 9, 10)
D.frame(en, 6, 9, "B", "J")
sign_rule(en, "E8")
D.status_rule(en, "I8", "$I$8>0", fill_tint=False)
D.h(en, 10, 14)


def counts_as(rr):
    return (f'=IF(G{rr}="","",IF(F{rr}="","Needs a category",IF(C{rr}="","Needs a date",'
            f'IFERROR(INDEX({CAT}!$C${C_FIRST}:$C${C_LAST},MATCH(F{rr},{CAT_LIST},0)),"Not on the list"))))')


cols_en = [
    dict(col="C", head="Date", kind="input", fmt=DATE, values=[x[0] for x in ROWS],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The day the money moved, for example 3/14/2026. It shows the month by name so you can check it.",
         error="Type a real date between 2000 and 2100, for example 3/14/2026."),
    dict(col="D", head="What it was", kind="input", values=[x[1] for x in ROWS], validation=("textLength", 0, 80),
         prompt="A few words you will recognise later, for example Blank shirts order.", error="Up to 80 characters."),
    dict(col="E", head="Customer or vendor", kind="input", values=[x[2] for x in ROWS], validation=("textLength", 0, 60),
         prompt="Who paid you, or who you paid. Optional.", error="Up to 60 characters."),
    dict(col="F", head="Category (pick from list)", kind="input", values=[x[3] for x in ROWS], validation=("list_range", CAT_LIST),
         prompt="Choose a category from the list. You can rename categories on the Categories tab.",
         error="Choose a category from the list. To add a name, rename one on the Categories tab first."),
    dict(col="G", head="Amount", kind="input", fmt=USD2, align="right", values=[x[4] for x in ROWS],
         validation=("decimal", 0, None), prompt="Type the amount as a positive number, cents included. The category decides if it adds or subtracts.",
         error="Type 0 or more, as a positive number. The category decides if it adds or subtracts."),
    dict(col="H", head="Notes", kind="input", values=[x[5] for x in ROWS], validation=("textLength", 0, 120),
         prompt="Anything worth remembering about this row. Optional.", error="Up to 120 characters."),
    dict(col="I", head="Counts as", kind="calc", formula=counts_as),
]
f1, f2, en_end = D.table_card(en, E_TOP, "B", "J", cols_en, N_ROWS, title="Your entries",
                              sub=f"Yellow columns are yours. Counts as fills in on its own. {N_ROWS:,} rows; the totals read rows {E_FIRST} to {E_CHECK_LAST:,}.")
assert (f1, f2) == (E_FIRST, E_LAST), (f1, f2)
for rr in range(E_FIRST, E_LAST + 1):
    en[f"I{rr}"].font = D.f(9, False, "ink2"); en[f"I{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for k in ("D", "E", "F", "H"):
        en[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
en.conditional_formatting.add(f"C{E_FIRST}:F{E_LAST}", FormulaRule(
    formula=[f'AND($G{E_FIRST}<>"",OR($C{E_FIRST}="",$F{E_FIRST}=""))'], fill=PatternFill("solid", bgColor=D.T["accent_tint"])))
for word, color in (("Needs", "accent"), ("Not on", "accent"), ("Not in", "muted")):
    en.conditional_formatting.add(f"I{E_FIRST}:I{E_LAST}", FormulaRule(
        formula=[f'LEFT($I{E_FIRST},{len(word)})="{word}"'], font=Font(color=D.T[color], bold=color == "accent"), stopIfTrue=True))
D.paint_canvas(en, en_end + 1, "Z")
S.finish_sheet(en, PRODUCT, en_end + 1, span=("A", "K"), freeze=f"A{E_FIRST}", tab_color="gold")
en.page_setup.orientation = "landscape"
en.print_title_rows = f"{E_FIRST - 1}:{E_FIRST - 1}"
en.print_area = f"A1:K{E_FIRST + 99}"

# =====================================================================  4 Categories
ct = wb.create_sheet(CAT_N)
S.set_widths(ct, {"A": 3, "B": 2, "C": 16, "D": 36, "E": 36, "F": 60, "G": 2, "H": 3})
D.page_header(ct, "Categories", "Rename any yellow cell. The Entries dropdown, the Statement and the month table follow.", AS_OF,
              span=("B", "G"), side_cols=2)
D.h(ct, 6, 10)
D.stat_strip(ct, 7, [
    ("C:C", "On the list", f'=COUNTA({CAT_LIST})', INT),
    ("D:D", "Kept out of profit", f"={len(OFF_IDX)}", INT),
    ("E:F", "Entries with a name not on the list", f"={STM}!$G${k3}", INT),
])
D.h(ct, 9, 10)
D.frame(ct, 6, 9, "B", "G")
D.status_rule(ct, "E8", "$E$8>0", fill_tint=False)
D.h(ct, 10, 14)
cols_ct = [
    dict(col="C", head="Counts as", kind="text", values=[c[0] for c in CATS]),
    dict(col="D", head="Category", kind="input", values=[c[1] for c in CATS], validation=("textLength", 0, 40),
         prompt="Rename this category if you like (up to 40 characters). Rows already typed with the old name show up in Checks so you can update them.",
         error="Up to 40 characters."),
    dict(col="E", head="Schedule C reference", kind="text", values=[c[2] for c in CATS]),
    dict(col="F", head="What goes here", kind="text", values=[c[3] for c in CATS]),
]
g1, g2, ct_end = D.table_card(ct, C_TOP, "B", "G", cols_ct, len(CATS), title="Your categories",
                              sub="Grouped as the Statement groups them. Names follow the IRS 2025 Schedule C lines.")
assert (g1, g2) == (C_FIRST, C_LAST), (g1, g2)
for i, rr in enumerate(range(C_FIRST, C_LAST + 1)):
    for k in ("C", "E", "F"):
        ct[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ct[f"C{rr}"].font = D.f(9, True, "muted" if CATS[i][0] == "Not in profit" else "ink2")
    ct[f"E{rr}"].font = D.f(9, False, "ink2")
    ct[f"D{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    D.h(ct, rr, max(D.GRID_ROW, S.text_height(CATS[i][3], S.width_of(ct, "F"), 10) + 6,
                    S.text_height(CATS[i][2], S.width_of(ct, "E"), 9) + 6))
CN_TOP = ct_end + 2
r = CN_TOP
D.h(ct, r, 12); r += 1
t = ct[f"C{r}"]; t.value = "Notes and sources"; t.font = D.f(12, True); ct.merge_cells(f"C{r}:F{r}"); D.h(ct, r, 24); r += 1
cw = sum(S.width_of(ct, k) for k in "CDEF")
for head, body in [
    ("Source", f"IRS 2025 Schedule C (Form 1040), irs.gov/pub/irs-pdf/f1040sc.pdf, and its instructions, irs.gov/pub/irs-pdf/i1040sc.pdf, "
     f"checked {CHECKED}. For organizing, not tax advice; your preparer decides the final line for each item."),
    ("Renaming", "Rename a yellow cell and the Entries dropdown, the Statement and the month table follow. Rows already typed with the "
     "old name are counted in Checks until you pick the new name. To keep the groups right, rename within a group rather than reorder."),
]:
    cell = ct[f"C{r}"]; cell.value = rich(head, body)
    cell.alignment = Alignment(vertical="top", wrap_text=True); ct.merge_cells(f"C{r}:F{r}")
    D.h(ct, r, S.text_height(head + ". " + body, cw, 9) + 4); r += 1
D.h(ct, r, 12)
D.frame(ct, CN_TOP, r, "B", "G")
CT_LAST = r + 1
D.h(ct, CT_LAST, 14)
D.paint_canvas(ct, CT_LAST, "Z")
S.finish_sheet(ct, PRODUCT, CT_LAST, span=("A", "H"), tab_color="teal")
ct.page_setup.orientation = "landscape"
ct.print_title_rows = f"{C_FIRST - 1}:{C_FIRST - 1}"

# =====================================================================  Start Here and Terms
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            mm = re.match(r"^([A-Z][A-Za-z ,]{1,34})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if mm:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


def money(x):
    return ("-" if x < 0 else "") + f"${abs(x):,.2f}"


sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, PRODUCT, "Start here. You type in one tab; the statement fills itself in.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "A profit and loss template for a small business. Type one row for each money in or money out on the Entries tab. "
    "The Statement tab turns it into a profit and loss statement for the month you pick and for January to that month: net "
    "revenue, cost of goods sold, gross profit, operating expenses by Schedule C line and net profit, exact to the cent. The "
    "Month by Month tab shows the whole year side by side.", "para")])
r = sh_card(sh, r, "Three steps", [
    (1, "Entries tab. One row per money in or money out: date, what it was, category from the list, amount as a positive number.", "num"),
    (2, "Statement tab. Pick the year and month at the top. Leave them blank for this month. Stock counts are optional.", "num"),
    (3, "Statement tab. Read your net profit in the dark tile, the full statement below it, and check that Problems to fix reads 0.", "num"),
])
r = sh_card(sh, r, "How to read the cells", [
    ((1100, USD2), "Yellow cells with blue numbers are yours to type in. Each one shows a hint when you select it.", "chip-in"),
    ((3400, USD2), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((EX_YTD["np"], MONEY),
     "Dark tiles are your answers. The main one sits at the top of the Statement tab and stays on screen as you scroll.", "chip-key"),
    (None, "Each tab is protected without a password, so you can still change anything. To edit a locked cell, use Review, "
           "Unprotect Sheet in Excel, or Data, Protect sheets and ranges in Google Sheets.", "note"),
])
S.page_break_before(sh, r)
ex_text = (f"Profit-and-Loss-EXAMPLE.xlsx holds January to September 2026 of a made-up one-person T-shirt printing shop. "
           f"September shows {money(EX_MON['rev'])} of net revenue and {money(EX_MON['np'])} of net profit. January to September "
           f"shows {money(EX_YTD['rev'])} of net revenue, {money(EX_YTD['cogs'])} cost of goods sold after the two stock counts "
           f"($1,200.00 on January 1 and $1,450.00 at the end of September) and {money(EX_YTD['np'])} of net profit. The "
           "$1,850.00 heat press, the owner's monthly $1,500.00 and the $2,000.00 card payoff are logged but kept out of profit.")
r = sh_card(sh, r, "The example", [(None, ex_text, "para")])
r = sh_card(sh, r, "Which file to open", [
    (None, "Excel on a computer. Double-click Profit-and-Loss-Template.xlsx. If Excel shows a yellow Protected View bar, click Enable Editing.", "para"),
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. No macros, no add-ons, no sign-up.", "para"),
    (None, "Mac, iPad or phone. Unzip on a computer first (phones often can't open a .zip), then upload the .xlsx to Google Drive "
           "and open it in the Google Sheets app, or open it in Excel for Mac or iPad.", "para"),
    (None, "See it filled in first. Open Profit-and-Loss-EXAMPLE.xlsx, then use the blank template for your own business.", "para"),
])
r = sh_card(sh, r, "Good to know", [
    (None, "Exact numbers. Every total shows dollars and cents, so you can check the statement against your bank statement and your preparer's figures.", "para"),
    (None, "Any language setting. The formulas use no text-format codes, so they work the same in Excel set to other languages "
           "and in Google Sheets.", "para"),
    (None, "Print the statement. File, Print on the Statement tab prints the statement on a page of its own, ready to hand to a lender or preparer.", "para"),
    (None, f"Room for a busy year. The Entries tab has {N_ROWS:,} rows and every total reads rows {E_FIRST} to {E_CHECK_LAST:,}, so you "
           "can add rows below them after you unprotect the tab (see How to read the cells). Start a fresh copy each year.", "para"),
    (None, "Kept out of profit. Equipment, owner draws and transfers between your own accounts are logged but kept out of "
           "profit. Depreciation and home office use are not calculated here; hand those to your tax preparer.", "para"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "Before you rely on a number", [(None, S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab") + " "
           + S.TAX_TIER_CLAUSE.format(date=CHECKED_LONG, sources=SOURCES), "note")])
SH_LAST = r - 1
D.paint_canvas(sh, SH_LAST, "Z")
S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
               footer_notice="Estimates only. Not professional advice. See the Terms tab or LICENSE-AND-DISCLAIMER.txt.")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More tools from the shop", [
    (None, ("One-Tab Business Expense Tracker: expenses by Schedule C line, one tab", "https://www.etsy.com/listing/4589291725"), "link"),
    (None, ("Freelance Income Tracker: the income side, with tax set-aside", "https://www.etsy.com/listing/4589995309"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this template saves you time, a short review on Etsy helps other small business owners find it. Thank you.", "para"),
    (None, "If a file won't open or something looks wrong, message the shop on Etsy and it will be fixed.", "para"),
])
TM_LAST = r - 1
D.paint_canvas(tm, TM_LAST, "Z")
S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

wb.move_sheet(STM_N, offset=-(wb.sheetnames.index(STM_N) - 2))
wb.active = 0
problems = [p for p in S.audit(wb) if "input without validation" not in p]
if problems:
    print("AUDIT:", *problems, sep="\n  ")
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
print({"rev": R_REV, "cogs": R_COGS, "gp": R_GP, "exp": R_EXP, "np": R_NP, "stk": R_STK, "margins": MARGINS,
       "period": (r_y, r_m, r_s0, r_s1), "checks": (k1, k2, k3, k4, k5), "mbm": (M_REV, M_GP, M_NP), "E_FIRST": E_FIRST,
       "rows": len(ROWS)})
print("example:", EXAMPLE)
