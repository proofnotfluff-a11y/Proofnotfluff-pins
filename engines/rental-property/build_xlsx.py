#!/usr/bin/env python3
"""#19 Rental Property Spreadsheet for 1 to 10 Doors, version 2 (v3 dashboard design, Oct 7, 2026).

Same inputs, categories, worked example and maths as version 1 (Oct 5, 2026, the file live on Etsy):
one Setup tab (up to 10 properties and 10 doors) and one Log tab (every dollar in or out) feed a
Rent Roll for any month, a Property P&L and cash flow for any year (one column per property on the
IRS Schedule E Part I line names), and a Monthly summary for all properties or one. Rebuilt to
design/WORKBOOK_STANDARD.md section 0; tools/compare_xlsx.py with compare_map.json proves the
answers match.

Fixes in version 2 (each declared in compare_map.json "intended" or listed in the build notes):
- A blank month on the Rent Roll gave error values on every door row (DATE(YEAR(""),...)); a blank
  month now checks this month.
- A blank year on the Property P&L looked for the year 1900 and showed all zeros; a blank year now
  shows this year (the Monthly tab follows it).
- A door with no monthly rent typed showed "Paid"; it now says "No rent set".
- The Monthly tab told buyers to type "All properties", which its own dropdown rejected; the picker
  now ships blank, blank means all properties, and a Showing line says which is shown.
- The Monthly tab used SUMPRODUCT(SUMIFS(...)) with a range of categories as one criterion. Google Sheets reads
  only the first category there, so in Sheets Money in counted Rent only, Expenses counted Advertising only and
  Principal and improvements counted principal only. v2 uses a plain SUMPRODUCT that gives the same answer in
  Excel, LibreOffice and Google Sheets (tested in Google Sheets, Oct 7, 2026).
- Checks added (new outputs, nothing else changes): rent rows with no door (left off the Rent Roll),
  and category names not on the list (after a rename).
- Every input now has a validation rule and an input message; negatives show a leading minus, not
  red parentheses; the Thank You tab is folded into the Terms tab.

Usage: python3 engines/rental-property/build_xlsx.py out.xlsx [--blank]
"""
import datetime as dt
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402

PRODUCT = "Rental Property Spreadsheet for 1 to 10 Doors"
VERSION = "2"
CHECKED = "Oct 7, 2026"
CHECKED_LONG = "October 7, 2026"
AS_OF = f"Lines checked {CHECKED}"
SOURCES = "IRS Schedule E (Form 1040) 2025, the 2025 Instructions for Schedule E and IRS Publication 527 (2025)"
HERE = os.path.dirname(os.path.abspath(__file__))
args = sys.argv[1:]
BLANK = "--blank" in args
args = [a for a in args if a != "--blank"]
OUT = args[0] if args else "Rental-Property-Spreadsheet-1-to-10-Doors.xlsx"
USD0, INT, DATE = S.FMT["usd0"], S.FMT["int"], S.FMT["date"]
MONEY2 = '"$"#,##0.00;-"$"#,##0.00;"-"'   # rent, log amounts, deposits: cents kept, a dash for zero
MONEY0 = '"$"#,##0;-"$"#,##0;"-"'         # yearly and monthly totals: whole dollars, a dash for zero
YEAR_FMT = "0"
MONTH_FMT = "mmmm yyyy"

# ---------------------------------------------------------------- tab names (numbered: more than 3 working tabs)
SET, LOG, RR, PL, MO, CAT = "1 Setup", "2 Log", "3 Rent Roll", "4 Property P&L", "5 Monthly", "6 Categories"


def q(name):
    return f"'{name}'"


# ---------------------------------------------------------------- categories (v1 order, names and Schedule E lines)
# (kind, category, Schedule E line, what goes here)
CATS = [
    ("Income", "Rent", "Line 3, Rents received", "Monthly rent you collected."),
    ("Income", "Other rental income", "Line 3, Rents received", "Late fees, pet rent, a kept deposit. Confirm with your preparer."),
    ("Expense", "Advertising", "Line 5, Advertising", "Listing sites, signs."),
    ("Expense", "Auto and travel", "Line 6, Auto and travel", "Trips to the property. The IRS instructions set the rules."),
    ("Expense", "Cleaning and maintenance", "Line 7, Cleaning and maintenance", "Turnover cleaning, lawn, snow."),
    ("Expense", "Commissions", "Line 8, Commissions", "Leasing agent commissions."),
    ("Expense", "Insurance", "Line 9, Insurance", "Landlord or rental dwelling policy."),
    ("Expense", "Legal and other professional fees", "Line 10, Legal and other professional fees", "Lease review, tax prep share."),
    ("Expense", "Management fees", "Line 11, Management fees", "Property manager."),
    ("Expense", "Mortgage interest", "Line 12, Mortgage interest paid to banks, etc.", "Interest only. Principal goes on its own row."),
    ("Expense", "Other interest", "Line 13, Other interest", "Interest on other loans for the rental."),
    ("Expense", "Repairs", "Line 14, Repairs", "Keeps the place in working order. Upgrades are improvements."),
    ("Expense", "Supplies", "Line 15, Supplies", "Light bulbs, filters, smoke detector batteries."),
    ("Expense", "Taxes", "Line 16, Taxes", "Property tax for the rental."),
    ("Expense", "Utilities", "Line 17, Utilities", "Utilities you pay, not the tenant."),
    ("Expense", "Other expenses", "Line 19, Other (list)", "Bank fees, HOA dues, software. Ask your preparer."),
    ("Cash only", "Mortgage principal", "Not an expense line", "Counts in cash flow, not in profit."),
    ("Cash only", "Capital improvement", "Not an expense line (depreciated)", "New roof, water heater. Your preparer handles depreciation."),
    ("Deposit", "Security deposit received", "Not income if you plan to return it",
     "Held for the tenant. IRS Pub. 527. A deposit meant to be the last month's rent is advance rent and counts as income when received (Pub. 527)."),
    ("Deposit", "Security deposit returned", "Not an expense line", "Money you gave back."),
]
# P&L row labels and their Schedule E text (v1 wording)
PL_LABELS = {0: ("Rent", "Line 3, Rents received"), 1: ("Other rental income", "Line 3, Rents received"),
             16: ("Mortgage principal", "Not on Schedule E"), 17: ("Capital improvement", "Depreciated, ask your preparer"),
             18: ("Security deposits received", "Not income if you plan to return it"), 19: ("Security deposits returned", "")}

# ---------------------------------------------------------------- example (v1, unchanged): a made-up duplex and house
d = dt.datetime
PROPS = [("Maple Duplex", "Sample: 2 units, 1 mortgage"), ("Oak Street House", "Sample: single family")]
DOORS = [  # door, property, tenant, rent, due day, grace days, lease end, deposit
    ("MAPLE-A", "Maple Duplex", "Tenant A", 1150, 1, 5, d(2027, 5, 31), 1150),
    ("MAPLE-B", "Maple Duplex", "Tenant B", 1100, 1, 5, d(2026, 11, 30), 1100),
    ("OAK", "Oak Street House", "Tenant C", 1650, 1, 3, d(2027, 7, 31), 1650),
]
LOGROWS = [  # date, property, door, category, amount, paid to or from, note
    (d(2026, 1, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 1, 1), 'Oak Street House', 'OAK', 'Rent', 1650, 'Tenant', None),
    (d(2026, 1, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 1100, 'Tenant', None),
    (d(2026, 1, 5), 'Maple Duplex', None, 'Mortgage interest', 612, 'Lender', 'From the monthly statement'),
    (d(2026, 1, 5), 'Maple Duplex', None, 'Mortgage principal', 288, 'Lender', 'From the monthly statement'),
    (d(2026, 1, 5), 'Oak Street House', None, 'Mortgage interest', 534, 'Lender', None),
    (d(2026, 1, 5), 'Oak Street House', None, 'Mortgage principal', 316, 'Lender', None),
    (d(2026, 1, 20), 'Maple Duplex', None, 'Utilities', 138, 'City water', 'Water and sewer, both units'),
    (d(2026, 1, 30), 'Maple Duplex', None, 'Legal and other professional fees', 250, 'Tax preparer', 'Rental share of prep fee'),
    (d(2026, 2, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 2, 1), 'Oak Street House', 'OAK', 'Rent', 1650, 'Tenant', None),
    (d(2026, 2, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 1100, 'Tenant', None),
    (d(2026, 2, 5), 'Maple Duplex', None, 'Mortgage interest', 612, 'Lender', 'From the monthly statement'),
    (d(2026, 2, 5), 'Maple Duplex', None, 'Mortgage principal', 288, 'Lender', 'From the monthly statement'),
    (d(2026, 2, 5), 'Oak Street House', None, 'Mortgage interest', 534, 'Lender', None),
    (d(2026, 2, 5), 'Oak Street House', None, 'Mortgage principal', 316, 'Lender', None),
    (d(2026, 2, 11), 'Maple Duplex', 'MAPLE-B', 'Repairs', 185, 'Plumber', 'Kitchen drain'),
    (d(2026, 2, 20), 'Maple Duplex', None, 'Utilities', 138, 'City water', 'Water and sewer, both units'),
    (d(2026, 3, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 3, 1), 'Oak Street House', 'OAK', 'Rent', 1650, 'Tenant', None),
    (d(2026, 3, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 1100, 'Tenant', None),
    (d(2026, 3, 2), 'Maple Duplex', None, 'Insurance', 1380, 'Insurer', 'Annual landlord policy'),
    (d(2026, 3, 5), 'Maple Duplex', None, 'Mortgage interest', 612, 'Lender', 'From the monthly statement'),
    (d(2026, 3, 5), 'Maple Duplex', None, 'Mortgage principal', 288, 'Lender', 'From the monthly statement'),
    (d(2026, 3, 5), 'Oak Street House', None, 'Mortgage interest', 534, 'Lender', None),
    (d(2026, 3, 5), 'Oak Street House', None, 'Mortgage principal', 316, 'Lender', None),
    (d(2026, 3, 18), 'Oak Street House', 'OAK', 'Supplies', 62, 'Hardware store', 'Filters and smoke detector batteries'),
    (d(2026, 3, 20), 'Maple Duplex', None, 'Utilities', 138, 'City water', 'Water and sewer, both units'),
    (d(2026, 4, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 4, 1), 'Oak Street House', 'OAK', 'Rent', 1650, 'Tenant', None),
    (d(2026, 4, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 1100, 'Tenant', None),
    (d(2026, 4, 5), 'Maple Duplex', None, 'Mortgage interest', 612, 'Lender', 'From the monthly statement'),
    (d(2026, 4, 5), 'Maple Duplex', None, 'Mortgage principal', 288, 'Lender', 'From the monthly statement'),
    (d(2026, 4, 5), 'Oak Street House', None, 'Mortgage interest', 534, 'Lender', None),
    (d(2026, 4, 5), 'Oak Street House', None, 'Mortgage principal', 316, 'Lender', None),
    (d(2026, 4, 8), 'Oak Street House', None, 'Auto and travel', 38, None, 'Trips to the property, see IRS rules'),
    (d(2026, 4, 20), 'Maple Duplex', None, 'Utilities', 138, 'City water', 'Water and sewer, both units'),
    (d(2026, 4, 22), 'Oak Street House', 'OAK', 'Repairs', 240, 'Handyman', 'Back door lock and frame'),
    (d(2026, 5, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 5, 1), 'Oak Street House', 'OAK', 'Rent', 1650, 'Tenant', None),
    (d(2026, 5, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 1100, 'Tenant', None),
    (d(2026, 5, 3), 'Maple Duplex', None, 'Cleaning and maintenance', 320, 'Lawn service', 'Spring cleanup'),
    (d(2026, 5, 5), 'Maple Duplex', None, 'Mortgage interest', 612, 'Lender', 'From the monthly statement'),
    (d(2026, 5, 5), 'Maple Duplex', None, 'Mortgage principal', 288, 'Lender', 'From the monthly statement'),
    (d(2026, 5, 5), 'Oak Street House', None, 'Mortgage interest', 534, 'Lender', None),
    (d(2026, 5, 5), 'Oak Street House', None, 'Mortgage principal', 316, 'Lender', None),
    (d(2026, 5, 20), 'Maple Duplex', None, 'Utilities', 138, 'City water', 'Water and sewer, both units'),
    (d(2026, 6, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 6, 1), 'Oak Street House', 'OAK', 'Rent', 1650, 'Tenant', None),
    (d(2026, 6, 1), 'Oak Street House', None, 'Insurance', 1140, 'Insurer', 'Annual landlord policy'),
    (d(2026, 6, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 1100, 'Tenant', None),
    (d(2026, 6, 5), 'Maple Duplex', None, 'Mortgage interest', 612, 'Lender', 'From the monthly statement'),
    (d(2026, 6, 5), 'Maple Duplex', None, 'Mortgage principal', 288, 'Lender', 'From the monthly statement'),
    (d(2026, 6, 5), 'Oak Street House', None, 'Mortgage interest', 534, 'Lender', None),
    (d(2026, 6, 5), 'Oak Street House', None, 'Mortgage principal', 316, 'Lender', None),
    (d(2026, 6, 15), 'Maple Duplex', None, 'Taxes', 2140, 'County treasurer', 'First installment'),
    (d(2026, 6, 15), 'Oak Street House', None, 'Taxes', 1860, 'County treasurer', 'First installment'),
    (d(2026, 6, 20), 'Maple Duplex', None, 'Utilities', 138, 'City water', 'Water and sewer, both units'),
    (d(2026, 7, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 7, 1), 'Oak Street House', 'OAK', 'Rent', 1650, 'Tenant', None),
    (d(2026, 7, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 1100, 'Tenant', None),
    (d(2026, 7, 5), 'Maple Duplex', None, 'Mortgage interest', 612, 'Lender', 'From the monthly statement'),
    (d(2026, 7, 5), 'Maple Duplex', None, 'Mortgage principal', 288, 'Lender', 'From the monthly statement'),
    (d(2026, 7, 5), 'Oak Street House', None, 'Mortgage interest', 534, 'Lender', None),
    (d(2026, 7, 5), 'Oak Street House', None, 'Mortgage principal', 316, 'Lender', None),
    (d(2026, 7, 9), 'Oak Street House', 'OAK', 'Capital improvement', 1450, 'Plumber', 'New water heater'),
    (d(2026, 7, 20), 'Maple Duplex', None, 'Utilities', 138, 'City water', 'Water and sewer, both units'),
    (d(2026, 8, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 8, 1), 'Oak Street House', 'OAK', 'Rent', 1650, 'Tenant', None),
    (d(2026, 8, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 1100, 'Tenant', None),
    (d(2026, 8, 3), 'Maple Duplex', None, 'Cleaning and maintenance', 160, 'Lawn service', 'Summer mowing'),
    (d(2026, 8, 5), 'Maple Duplex', None, 'Mortgage interest', 612, 'Lender', 'From the monthly statement'),
    (d(2026, 8, 5), 'Maple Duplex', None, 'Mortgage principal', 288, 'Lender', 'From the monthly statement'),
    (d(2026, 8, 5), 'Oak Street House', None, 'Mortgage interest', 534, 'Lender', None),
    (d(2026, 8, 5), 'Oak Street House', None, 'Mortgage principal', 316, 'Lender', None),
    (d(2026, 8, 20), 'Maple Duplex', None, 'Utilities', 138, 'City water', 'Water and sewer, both units'),
    (d(2026, 9, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 9, 1), 'Oak Street House', 'OAK', 'Rent', 1650, 'Tenant', None),
    (d(2026, 9, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 1100, 'Tenant', None),
    (d(2026, 9, 2), 'Maple Duplex', 'MAPLE-A', 'Other rental income', 50, 'Tenant', 'Pet rent, September'),
    (d(2026, 9, 5), 'Maple Duplex', None, 'Mortgage interest', 612, 'Lender', 'From the monthly statement'),
    (d(2026, 9, 5), 'Maple Duplex', None, 'Mortgage principal', 288, 'Lender', 'From the monthly statement'),
    (d(2026, 9, 5), 'Oak Street House', None, 'Mortgage interest', 534, 'Lender', None),
    (d(2026, 9, 5), 'Oak Street House', None, 'Mortgage principal', 316, 'Lender', None),
    (d(2026, 9, 14), 'Maple Duplex', None, 'Taxes', 2140, 'County treasurer', 'Second installment'),
    (d(2026, 9, 14), 'Oak Street House', None, 'Taxes', 1860, 'County treasurer', 'Second installment'),
    (d(2026, 9, 20), 'Maple Duplex', None, 'Utilities', 138, 'City water', 'Water and sewer, both units'),
    (d(2026, 10, 1), 'Maple Duplex', 'MAPLE-A', 'Rent', 1150, 'Tenant', None),
    (d(2026, 10, 2), 'Maple Duplex', 'MAPLE-B', 'Rent', 600, 'Tenant', 'Partial, rest promised Friday'),
]
EX_YEAR, EX_MONTH, EX_PICK = 2026, d(2026, 10, 1), "Maple Duplex"


def _sum(prop=None, cats=None, month=None, door=None):
    return sum(x[4] for x in LOGROWS if (prop is None or x[1] == prop) and (cats is None or x[3] in cats)
               and (month is None or x[0].month == month) and (door is None or x[2] == door))


_names = lambda idx: [CATS[i][1] for i in idx]  # noqa: E731
INC, EXP, CASH = _names((0, 1)), _names(range(2, 16)), _names((16, 17))
EX = dict(oak_profit=_sum("Oak Street House", INC) - _sum("Oak Street House", EXP),
          oak_principal=_sum("Oak Street House", ["Mortgage principal"]),
          oak_improve=_sum("Oak Street House", ["Capital improvement"]),
          all_profit=_sum(None, INC) - _sum(None, EXP))
EX["oak_cash"] = EX["oak_profit"] - EX["oak_principal"] - EX["oak_improve"]
EX["all_cash"] = EX["all_profit"] - _sum(None, CASH)
EX["oct_due"] = sum(x[3] for x in DOORS)
EX["oct_in"] = _sum(None, ["Rent"], 10)
# the Start Here example, the listing and the pins quote these; fail the build if the data ever drifts
assert (EX["oak_profit"], EX["oak_principal"], EX["oak_improve"], EX["oak_cash"]) == (4844, 2844, 1450, 550), EX
assert (EX["all_profit"], EX["all_cash"], EX["oct_due"], EX["oct_in"]) == (13569, 6683, 3900, 1750), EX
if BLANK:
    PROPS, DOORS, LOGROWS = [], [], []
    EX_YEAR, EX_MONTH, EX_PICK = None, None, None

# ---------------------------------------------------------------- fixed addresses
NP, ND = 10, 10                       # properties and doors (v1: 10 and 10)
P_FIRST = 16; P_LAST = P_FIRST + NP - 1          # 1 Setup: properties table body
D_FIRST = 33; D_LAST = D_FIRST + ND - 1          # 1 Setup: doors table body
N_ROWS = 1000                                    # styled Log rows (v1: 1,000 typing rows)
L_FIRST = 16; L_LAST = L_FIRST + N_ROWS - 1
L_READ = L_FIRST + 4999                          # formulas read 5,000 Log rows, as in v1
C_FIRST = 16; C_LAST = C_FIRST + len(CATS) - 1   # 6 Categories: names in column D
PROP_LIST = f"{q(SET)}!$D${P_FIRST}:$D${P_LAST}"
DOOR_LIST = f"{q(SET)}!$C${D_FIRST}:$C${D_LAST}"
CAT_LIST = f"{q(CAT)}!$D${C_FIRST}:$D${C_LAST}"


def catn(i):
    return f"{q(CAT)}!$D${C_FIRST + i}"


def lcol(c):
    return f"{q(LOG)}!${c}${L_FIRST}:${c}${L_READ}"


LD, LP, LDOOR, LC, LA = lcol("C"), lcol("D"), lcol("E"), lcol("F"), lcol("G")

# P&L year and pick (4 Property P&L, column D)
PL_YIN, PL_YEFF, PL_PICK = 13, 14, 15
YEAR = f"{q(PL)}!$D${PL_YEFF}"
# Rent Roll month (3 Rent Roll, column D)
RR_MIN, RR_MEFF = 13, 14
MONTH = f"{q(RR)}!$D${RR_MEFF}"


def sign_rule(ws, ref):
    """Accent below zero, teal above, plain at zero (a dash), so empty months stay quiet."""
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{ref.split(':')[0]}<0"], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{ref.split(':')[0]}>0"], font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))


def box(c, size=10):
    c.fill = D.fill("input_fill"); c.font = D.f(size, True, "input_text")
    c.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))


def card_label(ws, cell, text, size=10, bold=False, color="ink"):
    c = ws[cell]; c.value = text; c.font = D.f(size, bold, color); c.alignment = Alignment(vertical="center")
    return c


def hair(ws, row, a, z, top_ink=False, tint=None):
    for k in D.cols(a, z):
        cell = ws[f"{k}{row}"]
        cell.border = Border(top=D.side("ink") if top_ink else None, bottom=D.side("hair"))
        if tint:
            cell.fill = D.fill(tint)


def table_head(ws, row, heads, a, z, height=30):
    """heads: [(col, text, align)]; small caps over an ink rule across a..z."""
    for col, text, align in heads:
        c = ws[f"{col}{row}"]; c.value = text.upper(); c.font = D.f(8, True, "muted")
        c.alignment = Alignment(horizontal=align, vertical="bottom", wrap_text=True)
    for k in D.cols(a, z):
        ws[f"{k}{row}"].border = Border(bottom=D.side("ink"))
    D.h(ws, row, height)


def card_title(ws, row, col, last, title, sub=None):
    D.h(ws, row, 12); row += 1
    t = ws[f"{col}{row}"]; t.value = title; t.font = D.f(12, True); t.alignment = Alignment(vertical="center")
    ws.merge_cells(f"{col}{row}:{last}{row}"); D.h(ws, row, 24); row += 1
    if sub:
        s = ws[f"{col}{row}"]; s.value = sub; s.font = D.f(9, False, "muted"); s.alignment = Alignment(vertical="top")
        ws.merge_cells(f"{col}{row}:{last}{row}"); D.h(ws, row, 18); row += 1
    D.h(ws, row, 6); row += 1
    return row


def notes_card(ws, top, a, z, left, right, notes, title="Notes and sources"):
    r = top
    D.h(ws, r, 12); r += 1
    t = ws[f"{a}{r}"]; t.value = title; t.font = D.f(12, True); ws.merge_cells(f"{a}{r}:{z}{r}"); D.h(ws, r, 24); r += 1
    D.h(ws, r, 6); r += 1
    width = sum(S.width_of(ws, k) for k in D.cols(a, z))
    for head, body in notes:
        cell = ws[f"{a}{r}"]
        cell.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=9, b=True, color=D.T["ink"]), head + ". "),
                                  TextBlock(InlineFont(rFont=D.FONT, sz=9, color=D.T["ink2"]), body))
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(f"{a}{r}:{z}{r}")
        D.h(ws, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
    D.h(ws, r, 12)
    D.frame(ws, top, r, left, right)
    return r


def fill_heights(ws, last):
    for rr in range(1, last + 1):
        if ws.row_dimensions[rr].height is None:
            D.h(ws, rr, D.GRID_ROW)


TAX_NOTE = S.TAX_TIER_CLAUSE.format(date=CHECKED_LONG, sources=SOURCES)
NOTICE = S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab")

wb = S.new_workbook(PRODUCT, version=VERSION)

# =====================================================================  1 Setup
st = wb.create_sheet(SET)
S.set_widths(st, {"A": 3, "B": 2, "C": 14, "D": 22, "E": 28, "F": 14, "G": 11, "H": 11, "I": 15, "J": 14, "K": 2, "L": 3})
D.page_header(st, "Setup", "Type your properties and doors once. Up to 10 properties and 10 doors.",
              "Example rows: type over them" if not BLANK else "Start with your first property", span=("B", "K"), side_cols=3)
D.h(st, 6, 10)
D.stat_strip(st, 7, [
    ("C:C", "Properties", f'=COUNTIF({PROP_LIST},"?*")', INT),
    ("D:D", "Doors", f'=COUNTIF({DOOR_LIST},"?*")', INT),
    ("E:E", "Monthly rent, all doors", f"=SUM(F{D_FIRST}:F{D_LAST})", MONEY2),
    ("F:H", "Deposits held", f"=SUM(J{D_FIRST}:J{D_LAST})", MONEY2),
    ("I:J", "Doors to fix", None, INT),
])
D.h(st, 9, 10)
D.frame(st, 6, 9, "B", "K")
D.h(st, 10, 14)
cols_p = [
    dict(col="C", head="No.", kind="muted", values=list(range(1, NP + 1)), align="left"),
    dict(col="D", head="Property name (short)", kind="input", values=[p[0] for p in PROPS], validation=("textLength", 0, 30),
         prompt="A short name you will pick from dropdowns, for example Maple Duplex. Up to 30 characters.",
         error="Up to 30 characters."),
    dict(col="E", head="Address or note", kind="input", values=[p[1] for p in PROPS], validation=("textLength", 0, 60),
         prompt="Optional: the street address or a note to yourself.", error="Up to 60 characters."),
]
p1, p2, p_end = D.table_card(st, 11, "B", "K", cols_p, NP, title="Properties",
                             sub="A house is one property with one door. A duplex is one property with two doors.")
assert (p1, p2) == (P_FIRST, P_LAST), (p1, p2)
cols_d = [
    dict(col="C", head="Door ID", kind="input", values=[x[0] for x in DOORS], validation=("textLength", 0, 20),
         prompt="Any short name you will recognize in a dropdown, like MAPLE-A or 12-ELM. Up to 20 characters.",
         error="Up to 20 characters."),
    dict(col="D", head="Property", kind="input", values=[x[1] for x in DOORS], validation=("list_range", PROP_LIST),
         prompt="Pick the property this door belongs to. Add properties in the table above first.",
         error="Pick a property from the list. Add it to the Properties table above first."),
    dict(col="E", head="Tenant", kind="input", values=[x[2] for x in DOORS], validation=("textLength", 0, 40),
         prompt="The tenant's name, or a label such as Tenant A. Up to 40 characters.", error="Up to 40 characters."),
    dict(col="F", head="Monthly rent", kind="input", fmt=MONEY2, align="right", values=[x[3] for x in DOORS],
         validation=("decimal", 0, None), prompt="Rent due each month in dollars, for example 1150.",
         error="Type 0 or more, in dollars, for example 1150."),
    dict(col="G", head="Due day", kind="input", fmt=INT, align="right", values=[x[4] for x in DOORS], validation=("whole", 1, 31),
         prompt="Day of the month rent is due, a whole number from 1 to 31. 31 means the last day in shorter months.",
         error="Type a whole number from 1 to 31."),
    dict(col="H", head="Grace days", kind="input", fmt=INT, align="right", values=[x[5] for x in DOORS], validation=("whole", 0, 31),
         prompt="Whole days after the due day before rent shows LATE, for example 5.", error="Type a whole number from 0 to 31."),
    dict(col="I", head="Lease end", kind="input", fmt=DATE, align="right", values=[x[6] for x in DOORS],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The last day of the lease, for example 5/31/2027. The Rent Roll warns 60 days ahead. Leave blank for month to month.",
         error="Type a real date between 2000 and 2100, for example 5/31/2027."),
    dict(col="J", head="Deposit held", kind="input", fmt=MONEY2, align="right", values=[x[7] for x in DOORS],
         validation=("decimal", 0, None), prompt="Security deposit you hold for this door, in dollars, for example 1150.",
         error="Type 0 or more, in dollars."),
]
d1, d2, d_end = D.table_card(st, 28, "B", "K", cols_d, ND, title="Doors",
                             sub="One row per rentable unit. Rent due day 1 with 5 grace days shows LATE from the 7th if not fully paid.")
assert (d1, d2) == (D_FIRST, D_LAST), (d1, d2)
for rr in range(P_FIRST, P_LAST + 1):
    for k in ("D", "E"):
        st[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for rr in range(D_FIRST, D_LAST + 1):
    for k in ("C", "D", "E"):
        st[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
# doors to fix: an ID with no property or a property not on the list, or a duplicate ID
DOOR_FIX = (f'SUMPRODUCT(({q(SET)}!$C${D_FIRST}:$C${D_LAST}<>"")*(COUNTIF({PROP_LIST},{q(SET)}!$D${D_FIRST}:$D${D_LAST})=0))'
            f'+SUMPRODUCT(({q(SET)}!$C${D_FIRST}:$C${D_LAST}<>"")*(COUNTIF({DOOR_LIST},{q(SET)}!$C${D_FIRST}:$C${D_LAST})>1))')
st["I8"].value = f"={DOOR_FIX}"
D.status_rule(st, "I8", "$I$8>0", fill_tint=False)
st.conditional_formatting.add(f"C{D_FIRST}:D{D_LAST}", FormulaRule(
    formula=[f'AND($C{D_FIRST}<>"",OR(COUNTIF({PROP_LIST},$D{D_FIRST})=0,COUNTIF({DOOR_LIST},$C{D_FIRST})>1))'],
    fill=PatternFill("solid", bgColor=D.T["accent_tint"])))
ST_NOTES = [
    ("Doors to fix", "Counts door rows with a property that is not in the Properties table, and door IDs used twice. The row turns "
     "pink until it is fixed. A door's rent only reaches the Rent Roll when its ID matches the Door picked on the 2 Log tab."),
    ("Lease end", "The 3 Rent Roll tab shows a heads-up 60 days before a lease ends, and Lease ended after the date."),
    ("Starting fresh", "Select the example rows here and on the 2 Log tab and press Delete, then clear the Month to check on "
     "3 Rent Roll and the Year and Property on 4 Property P&L. Every formula keeps working on a blank file, and the answer tabs "
     "fill in again as you type."),
    ("Private information", "Tenant names and deposits are personal information. Keep this file somewhere secure and share it "
     "only with people you trust."),
]
st_last = notes_card(st, d_end + 2, "C", "J", "B", "K", ST_NOTES, title="Good to know")
ST_LAST = st_last + 1
D.h(st, ST_LAST, 14)
fill_heights(st, ST_LAST)
D.paint_canvas(st, ST_LAST, "Z")
S.finish_sheet(st, PRODUCT, ST_LAST, span=("A", "L"), freeze="A10", tab_color="gold")
st.page_setup.orientation = "landscape"
S.page_break_before(st, 28)

# =====================================================================  2 Log
lg = wb.create_sheet(LOG)
S.set_widths(lg, {"A": 3, "B": 2, "C": 13, "D": 20, "E": 13, "F": 33, "G": 13, "H": 20, "I": 34, "J": 20, "K": 2, "L": 3})
D.page_header(lg, "Log", "Every dollar in or out, one row each. Amounts are always positive; the category decides where it counts.",
              "Example rows: type over them" if not BLANK else f"Type your first row in row {L_FIRST}", span=("B", "K"), side_cols=2)
D.h(lg, 6, 10)
# the strip points at the P&L "All properties" column; rows are filled in once the P&L exists
D.h(lg, 9, 10)
D.h(lg, 10, 14)


def counts_as(r):
    return (f'=IF(G{r}="","",IF(D{r}="","Needs a property",IF(F{r}="","Needs a category",IF(C{r}="","Needs a date",'
            f'IF(AND(F{r}={catn(0)},E{r}=""),"Rent needs a door",'
            f'IFERROR(INDEX({q(CAT)}!$C${C_FIRST}:$C${C_LAST},MATCH(F{r},{CAT_LIST},0)),"Not on the list"))))))')


cols_l = [
    dict(col="C", head="Date", kind="input", fmt=DATE, values=[x[0] for x in LOGROWS],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The day the money moved, for example 3/14/2026. It shows the month by name so you can check it.",
         error="Type a real date between 2000 and 2100, for example 3/14/2026."),
    dict(col="D", head="Property", kind="input", values=[x[1] for x in LOGROWS], validation=("list_range", PROP_LIST),
         prompt="Pick the property from the list. Rows with no property are not counted anywhere.",
         error="Pick a property from the list. Add it on the 1 Setup tab first."),
    dict(col="E", head="Door (for rent)", kind="input", values=[x[2] for x in LOGROWS], validation=("list_range", DOOR_LIST),
         prompt="Pick the door for rent rows so the Rent Roll can count them. Leave blank for bills.",
         error="Pick a door from the list. Add it on the 1 Setup tab first."),
    dict(col="F", head="Category", kind="input", values=[x[3] for x in LOGROWS], validation=("list_range", CAT_LIST),
         prompt="Pick a category. Each one matches a Schedule E line or a cash item (see the 6 Categories tab).",
         error="Pick a category from the list. To use a new name, rename one on the 6 Categories tab first."),
    dict(col="G", head="Amount", kind="input", fmt=MONEY2, align="right", values=[x[4] for x in LOGROWS],
         validation=("decimal", 0, None), prompt="Type the amount as a positive number, for example 1150. The category decides where it counts.",
         error="Type 0 or more, as a positive number. The category decides where it counts."),
    dict(col="H", head="Paid to or from", kind="input", values=[x[5] for x in LOGROWS], validation=("textLength", 0, 40),
         prompt="Who you paid, or who paid you. Optional.", error="Up to 40 characters."),
    dict(col="I", head="Note", kind="input", values=[x[6] for x in LOGROWS], validation=("textLength", 0, 120),
         prompt="Anything worth remembering about this row. Optional.", error="Up to 120 characters."),
    dict(col="J", head="Counts as", kind="calc", formula=counts_as),
]
l1, l2, l_end = D.table_card(lg, 11, "B", "K", cols_l, N_ROWS, title="Your log",
                             sub=f"Yellow columns are yours. Counts as fills in on its own. {N_ROWS:,} rows here; the totals read 5,000 rows.")
assert (l1, l2) == (L_FIRST, L_LAST), (l1, l2)
for rr in range(L_FIRST, L_LAST + 1):
    lg[f"J{rr}"].font = D.f(9, False, "ink2"); lg[f"J{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for k in ("D", "E", "F", "H", "I"):
        lg[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
lg.conditional_formatting.add(f"C{L_FIRST}:F{L_LAST}", FormulaRule(
    formula=[f'AND($G{L_FIRST}<>"",OR($C{L_FIRST}="",$D{L_FIRST}="",$F{L_FIRST}="",AND($F{L_FIRST}={catn(0)},$E{L_FIRST}="")))'],
    fill=PatternFill("solid", bgColor=D.T["accent_tint"])))
for word, color in (("Needs", "accent"), ("Rent needs", "accent"), ("Not on", "accent"), ("Cash", "muted"), ("Deposit", "muted")):
    lg.conditional_formatting.add(f"J{L_FIRST}:J{L_LAST}", FormulaRule(
        formula=[f'LEFT($J{L_FIRST},{len(word)})="{word}"'], font=Font(color=D.T[color], bold=color == "accent"), stopIfTrue=True))
LG_LAST = l_end + 1
D.paint_canvas(lg, LG_LAST, "Z")
S.finish_sheet(lg, PRODUCT, LG_LAST, span=("A", "L"), freeze=f"A{L_FIRST}", tab_color="gold")
lg.page_setup.orientation = "landscape"
lg.print_title_rows = f"{L_FIRST - 1}:{L_FIRST - 1}"
lg.print_area = f"A1:L{L_FIRST + 99}"  # print the first 100 rows, not 1,000 bare ones

# =====================================================================  6 Categories (built early: everything points at it)
ct = wb.create_sheet(CAT)
S.set_widths(ct, {"A": 3, "B": 2, "C": 12, "D": 32, "E": 38, "F": 62, "G": 2, "H": 3})
D.page_header(ct, "Categories", "The Category dropdown list, matched to the Schedule E Part I lines.",
              AS_OF, span=("B", "G"), side_cols=2)
D.h(ct, 6, 10)
D.stat_strip(ct, 7, [
    ("C:C", "On the list", f'=COUNTIF({CAT_LIST},"?*")', INT),
    ("D:D", "Used in the log", None, INT),
    ("E:F", "Log rows with a name not on the list", None, INT),
])
D.h(ct, 9, 10)
D.frame(ct, 6, 9, "B", "G")
D.h(ct, 10, 14)
cols_c = [
    dict(col="C", head="Kind", kind="text", values=[c[0] for c in CATS]),
    dict(col="D", head="Category", kind="input", values=[c[1] for c in CATS], validation=("textLength", 1, 40),
         prompt="Rename this category if you like (1 to 40 characters). Rows already typed with the old name show up as Not on the list on the 2 Log tab.",
         error="Type 1 to 40 characters. Every row needs a name; rename it rather than clearing it."),
    dict(col="E", head="Schedule E line (tax year 2025 form)", kind="text", values=[c[2] for c in CATS]),
    dict(col="F", head="What goes here", kind="text", values=[c[3] for c in CATS]),
]
c1, c2, c_end = D.table_card(ct, 11, "B", "G", cols_c, len(CATS), title="Your categories",
                             sub="Income, the 14 Schedule E expense lines, cash items kept out of profit, and deposits kept out of income.")
assert (c1, c2) == (C_FIRST, C_LAST), (c1, c2)
for i, rr in enumerate(range(C_FIRST, C_LAST + 1)):
    for k in ("C", "E", "F"):
        ct[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ct[f"C{rr}"].font = D.f(9, True, "muted" if CATS[i][0] in ("Cash only", "Deposit") else "ink2")
    ct[f"E{rr}"].font = D.f(9, False, "ink2")
    ct[f"D{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    D.h(ct, rr, max(D.GRID_ROW, S.text_height(CATS[i][3], S.width_of(ct, "F"), 10) + 6,
                    S.text_height(CATS[i][2], S.width_of(ct, "E"), 9) + 6))
CT_NOTES = [
    ("Sources", f"IRS Schedule E (Form 1040) 2025, Part I lines 3 to 19, irs.gov/pub/irs-pdf/f1040se.pdf; the 2025 Instructions for "
     f"Schedule E, line 14 (repairs are deductible, improvements are capitalized and depreciated), irs.gov/instructions/i1040se; "
     f"IRS Publication 527 (2025), security deposits and advance rent, irs.gov/publications/p527. All checked {CHECKED}."),
    ("For organizing only", "The line names match the 2025 Schedule E form; your tax preparer decides what goes where on your return. "
     "The IRS draft 2026 Schedule E (created May 6, 2026) renames line 13 Interest and splits it into 13a Vehicle loan and "
     "13b Other; the Other interest row stays one line here."),
    ("Depreciation", "Line 18 (depreciation) is left out on purpose. Improvements and the building itself are depreciated over "
     "years, and that math belongs to your preparer."),
    ("Renaming", "Rename a yellow cell and the Log dropdown, the Rent Roll, the Property P&L and the Monthly tab follow. Rows "
     "already typed with the old name show Not on the list on the 2 Log tab until you pick the new name. The order of the rows "
     "decides where each one counts, so rename rather than reorder. The first row is always the rent the Rent Roll counts."),
]
S.page_break_before(ct, c_end + 2)
ct_last = notes_card(ct, c_end + 2, "C", "F", "B", "G", CT_NOTES)
CT_LAST = ct_last + 1
D.h(ct, CT_LAST, 14)
fill_heights(ct, CT_LAST)
D.paint_canvas(ct, CT_LAST, "Z")
S.finish_sheet(ct, PRODUCT, CT_LAST, span=("A", "H"), tab_color="teal")

# =====================================================================  4 Property P&L
pl = wb.create_sheet(PL)
PCOLS = D.cols("F", "O")
S.set_widths(pl, {"A": 3, "B": 2, "C": 30, "D": 20, "E": 14, **{k: 11.5 for k in PCOLS}, "P": 34, "Q": 2, "R": 3})
D.page_header(pl, "Property P&L and cash flow",
              "One column per property for the year shown, on the Schedule E line names. Fills itself in from the 2 Log tab.",
              AS_OF, span=("B", "Q"), side_cols=4)
for rr, ht in zip(range(11, 18), (12, 18, 34, 21, 24, 21, 12)):
    D.h(pl, rr, ht)
e = pl["C12"]; e.value = "YEAR AND PROPERTY"; e.font = D.f(8, True, "muted"); e.alignment = Alignment(vertical="bottom")
card_label(pl, f"C{PL_YIN}", "Year to show")
yc = S.input_cell(pl, f"D{PL_YIN}", EX_YEAR, YEAR_FMT, name="YearShown", wb=wb, validation=("whole", 2000, 2100), allow_blank=True,
                  prompt_title="Year to show", prompt="Leave blank to show this year (it changes on its own on January 1). Type a year, for example 2025, to look back.",
                  error="Type a year from 2000 to 2100, or leave the cell blank for this year.")
box(yc, 11); yc.alignment = Alignment(horizontal="center", vertical="center")
card_label(pl, f"C{PL_YEFF}", "Showing")
ye = pl[f"D{PL_YEFF}"]; ye.value = f'=IF($D${PL_YIN}="",YEAR(TODAY()),$D${PL_YIN})'; ye.number_format = YEAR_FMT
ye.font = D.f(10, True); ye.alignment = Alignment(horizontal="center", vertical="center")
card_label(pl, f"C{PL_PICK}", "Property to check")
pk = pl[f"D{PL_PICK}"]; pk.value = EX_PICK; pk.protection = pk.protection.copy(locked=False)
S.add_validation(pl, f"D{PL_PICK}", "list", options=["x"], allow_blank=True)  # replaced just below with the range list
pl.data_validations.dataValidation[-1].formula1 = PROP_LIST
dvp = pl.data_validations.dataValidation[-1]
dvp.promptTitle = "Property to check"; dvp.showInputMessage = True
dvp.prompt = "Pick one property for the Selected column. Leave it blank and that column stays empty."
dvp.error = "Pick a property from the list on the 1 Setup tab, or leave the cell blank."
box(pk); pk.alignment = Alignment(horizontal="left", vertical="center", indent=1)
S.define_name(wb, "PropertyToCheck", pl, f"D{PL_PICK}")
for rr in (PL_YEFF, PL_PICK):
    for k in D.cols("C", "G"):
        if k != "D":
            pl[f"{k}{rr}"].border = Border(bottom=D.side("hair"))
n = pl["C16"]; n.value = "Blank year shows this year. The Selected column shows the property you pick."
n.font = D.f(9, False, "muted"); n.alignment = Alignment(vertical="center"); pl.merge_cells("C16:G16")
D.frame(pl, 11, 17, "B", "H")

T_TOP = 19
r = card_title(pl, T_TOP, "C", "P", "Profit and cash flow by property",
               "Money in, the Schedule E expense lines, then cash items and deposits kept out of profit. Estimates only, see the Terms tab.")
HDR = r
for col, text, align in [("C", "Category", "left"), ("E", "All properties", "right"), ("P", "Schedule E line", "left")]:
    c = pl[f"{col}{HDR}"]; c.value = text; c.font = D.f(9, True, "ink2"); c.alignment = Alignment(horizontal=align, vertical="bottom", wrap_text=True, indent=1 if col == "P" else 0)
c = pl[f"D{HDR}"]; c.value = f'=IF($D${PL_PICK}="","Selected property","Selected: "&$D${PL_PICK})'
c.font = D.f(9, True, "ink"); c.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
for j, k in enumerate(PCOLS):
    c = pl[f"{k}{HDR}"]; c.value = f'=IF({q(SET)}!$D${P_FIRST + j}="","",{q(SET)}!$D${P_FIRST + j})'
    c.font = D.f(9, True, "ink2"); c.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
for k in D.cols("C", "P"):
    pl[f"{k}{HDR}"].border = Border(bottom=D.side("ink"))
D.h(pl, HDR, 34)
r = HDR + 1
PROW = {}
SEL = lambda rr: f"=SUMIF($F${HDR}:$O${HDR},$D${PL_PICK},F{rr}:O{rr})"  # noqa: E731


def pl_eyebrow(text):
    global r
    c = pl[f"C{r}"]; c.value = text; c.font = D.f(8, True, "muted"); c.alignment = Alignment(vertical="bottom")
    hair(pl, r, "C", "P"); D.h(pl, r, D.GRID_ROW); r += 1


def pl_line(label, cells, sched, total=False, tint=None, bold=False):
    """cells: callable(col) -> formula for each property column."""
    global r
    c = pl[f"C{r}"]; c.value = label; c.font = D.f(10, total or bold); c.alignment = Alignment(vertical="center")
    for k in PCOLS:
        v = pl[f"{k}{r}"]; v.value = cells(k); v.number_format = MONEY0; v.font = D.f(10, total)
        v.alignment = Alignment(horizontal="right", vertical="center")
    a = pl[f"E{r}"]; a.value = f"=SUM(F{r}:O{r})"; a.number_format = MONEY0; a.font = D.f(10, True)
    a.alignment = Alignment(horizontal="right", vertical="center")
    s = pl[f"D{r}"]; s.value = SEL(r); s.number_format = MONEY0; s.font = D.f(10, True)
    s.alignment = Alignment(horizontal="right", vertical="center")
    p = pl[f"P{r}"]; p.value = sched; p.font = D.f(9, False, "muted"); p.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    hair(pl, r, "C", "P", top_ink=total, tint=tint)
    D.h(pl, r, D.GRID_ROW + (3 if tint else 0))
    out = r; r += 1
    return out


def sumifs_cat(i):
    return lambda k: (f'=IF({k}${HDR}="",0,SUMIFS({LA},{LP},{k}${HDR},{LC},{catn(i)},{LD},">="&DATE($D${PL_YEFF},1,1),'
                      f'{LD},"<="&DATE($D${PL_YEFF},12,31)))')


pl_eyebrow("MONEY IN")
for i in (0, 1):
    PROW[i] = pl_line(PL_LABELS[i][0], sumifs_cat(i), PL_LABELS[i][1])
R_IN = pl_line("Total money in", lambda k: f"=SUM({k}{PROW[0]}:{k}{PROW[1]})", "Line 3", total=True)
pl_eyebrow("EXPENSES (SCHEDULE E LINES)")
for i in range(2, 16):
    PROW[i] = pl_line(CATS[i][1], sumifs_cat(i), CATS[i][2])
R_EXP = pl_line("Total expenses", lambda k: f"=SUM({k}{PROW[2]}:{k}{PROW[15]})", "Lines 5 to 19 added, without line 18", total=True)
R_PRO = pl_line("Profit before depreciation", lambda k: f"={k}{R_IN}-{k}{R_EXP}", "Line 3 minus expenses", total=True, tint="total_tint")
pl_eyebrow("CASH ITEMS (NOT EXPENSES)")
for i in (16, 17):
    PROW[i] = pl_line(PL_LABELS[i][0], sumifs_cat(i), PL_LABELS[i][1])
R_CASH = pl_line("Cash flow (what you kept)", lambda k: f"={k}{R_PRO}-{k}{PROW[16]}-{k}{PROW[17]}",
                 "Profit minus principal and improvements", total=True, tint="total_tint")
pl_eyebrow("DEPOSITS (HELD FOR TENANTS)")
for i in (18, 19):
    PROW[i] = pl_line(PL_LABELS[i][0], sumifs_cat(i), PL_LABELS[i][1])
for rr in (R_PRO, R_CASH):
    sign_rule(pl, f"D{rr}:O{rr}")
nt = pl[f"C{r}"]
nt.value = ("Profit here is before depreciation (Schedule E line 18), which your preparer works out. "
            "Deposits are kept out of income. For organizing, not a tax return.")
nt.font = D.f(9, False, "muted"); nt.alignment = Alignment(vertical="center"); pl.merge_cells(f"C{r}:P{r}")
D.h(pl, r, D.GRID_ROW); r += 1
D.h(pl, r, 12)
T_END = r
D.frame(pl, T_TOP, T_END, "B", "Q")

# ---- checks (left) and cash flow by property (right)
K_TOP = T_END + 2
K = D.Card(pl, K_TOP, "B", "C", "D", "H", wb=wb)
K.title("Checks", "All three should read 0. Fix the rows on the 2 Log tab.")
k1 = K.calc("Missing a property or category", f'=COUNTIFS({LA},">0",{LP},"")+COUNTIFS({LA},">0",{LC},"")', INT)
k2 = K.calc("Rent rows with no door", f'=COUNTIFS({LA},">0",{LC},{catn(0)},{LDOOR},"")', INT)
k3 = K.calc("Category not on the list", f'=SUMPRODUCT(({LC}<>"")*(COUNTIF({CAT_LIST},{LC})=0))', INT)
k4 = K.calc("Rows to fix", f"=D{k1}+D{k2}+D{k3}", INT, total=True)
for rr in (k1, k2, k3, k4):
    D.status_rule(pl, f"D{rr}", f"$D${rr}>0", fill_tint=False)
pill = K.text(f'=IF(COUNT({LA})=0,"No log rows yet",IF(D{k4}=0,"Every log row is counted",'
              f'D{k4}&IF(D{k4}=1," row needs"," rows need")&" a fix on the 2 Log tab"))', size=10, color="ink", bold=True)
pl[f"C{pill}"].alignment = Alignment(horizontal="center", vertical="center")
K.text("Rows missing a property or category are not counted anywhere. Rent with no door counts here but not on the Rent Roll.",
       height=D.GRID_ROW * 2)
K.close()
D.status_rule(pl, f"C{pill}:D{pill}", f"$D${k4}>0")

P = D.Card(pl, K_TOP, "J", "K", "O", "Q", extra="P", wb=wb)
P.title("Cash flow by property", "What each property kept in the year shown. Accent bars are negative.")
cf_first = P.r
CFR = f"$O${cf_first}:$O${cf_first + NP - 1}"
for j, k in enumerate(PCOLS):
    rr = P.calc(f"={k}{HDR}", f"={k}{R_CASH}", MONEY0)
    b = pl[f"P{rr}"]
    b.value = f'=IFERROR(REPT("{D.BAR_CHAR}",ROUND(ABS(O{rr})/MAX(MAX({CFR}),-MIN({CFR}))*16,0)),"")'
    b.font = D.f(9, False, "teal"); b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    pl.conditional_formatting.add(f"P{rr}", FormulaRule(formula=[f"$O{rr}<0"], font=Font(color=D.T["accent"]), stopIfTrue=True))
    sign_rule(pl, f"O{rr}")
P.close()
S.page_break_before(pl, K_TOP)

# ---- hero tiles
D.tile(pl, 6, "B", "H", f'="PROFIT BEFORE DEPRECIATION, "&$D${PL_YEFF}&", ALL PROPERTIES"', f"=$E${R_PRO}", USD0,
       f'="Money in "&TEXT($E${R_IN},"$#,##0")&", less "&TEXT($E${R_EXP},"$#,##0")&" of Schedule E expenses"', dark=True)
v, sub, _ = D.tile(pl, 6, "J", "Q", f'="CASH FLOW (WHAT YOU KEPT), "&$D${PL_YEFF}&", ALL PROPERTIES"', f"=$E${R_CASH}", USD0,
                   f'="After "&TEXT($E${PROW[16]}+$E${PROW[17]},"$#,##0")&" of mortgage principal and improvements"')
D.status_rule(pl, v.coordinate, f"$E${R_CASH}<0", fill_tint=False)
D.h(pl, 10, 14)
D.stat_strip(pl, 12, [
    ("K:L", "Money in", f"=$E${R_IN}", USD0),
    ("M:N", "Schedule E expenses", f"=$E${R_EXP}", USD0),
    ("O:P", "Log rows to fix", f"=$D${k4}", INT),
])
D.status_rule(pl, "O13", "$O$13>0", fill_tint=False)
s = pl["K15"]; s.value = "All properties, the year shown. Deposits are kept out of income."
s.font = D.f(9, False, "muted"); s.alignment = Alignment(vertical="center"); pl.merge_cells("K15:P15")
D.frame(pl, 11, 17, "J", "Q")

PL_NOTES = [
    ("Schedule E lines", f"The rows follow Part I of the IRS Schedule E (Form 1040) 2025, irs.gov/pub/irs-pdf/f1040se.pdf, checked {CHECKED}. "
     "Line 20 on the form adds lines 5 to 19, including line 18 depreciation; this tab leaves depreciation to your preparer."),
    ("Mortgage interest and principal", "Mortgage interest is an expense (line 12). Principal and capital improvements are not "
     "expenses; they lower cash flow only. The 2025 Instructions for Schedule E, line 14: amounts paid to improve your property "
     f"must generally be capitalized and depreciated (irs.gov/instructions/i1040se, checked {CHECKED})."),
    ("Security deposits", "IRS Publication 527 (2025): don't include a security deposit in income if you plan to return it; an "
     "amount you keep, or a deposit meant as the final month's rent (advance rent), is income when you keep or receive it. "
     f"Log a kept deposit as Other rental income. irs.gov/publications/p527, checked {CHECKED}."),
    ("Before you rely on a number", NOTICE + " " + TAX_NOTE),
]
pl_last = notes_card(pl, max(K.end, P.end) + 2, "C", "P", "B", "Q", PL_NOTES)
PL_LAST = pl_last + 1
D.h(pl, PL_LAST, 14)
fill_heights(pl, PL_LAST)
D.paint_canvas(pl, PL_LAST, "Z")
S.finish_sheet(pl, PRODUCT, PL_LAST, span=("A", "R"), freeze="A11", tab_color="accent")
pl.page_setup.orientation = "landscape"
S.page_break_before(pl, T_TOP)

# ---- strips on 2 Log and 6 Categories now that the P&L cells exist
D.stat_strip(lg, 7, [
    ("C:D", "Money in, year shown", f"={q(PL)}!$E${R_IN}", USD0),
    ("E:F", "Schedule E expenses", f"={q(PL)}!$E${R_EXP}", USD0),
    ("G:H", "Cash flow, year shown", f"={q(PL)}!$E${R_CASH}", USD0),
    ("I:J", "Rows to fix", f"={q(PL)}!$D${k4}", INT),
])
for rr in range(6, 10):
    for k in D.cols("B", "K"):
        lg[f"{k}{rr}"].fill = D.fill("card")
D.frame(lg, 6, 9, "B", "K")
sign_rule(lg, "G8")
D.status_rule(lg, "I8", "$I$8>0", fill_tint=False)
ct["D8"].value = f'=SUMPRODUCT(({CAT_LIST}<>"")*(COUNTIF({LC},{CAT_LIST})>0))'
ct["E8"].value = f"={q(PL)}!$D${k3}"
D.status_rule(ct, "E8", "$E$8>0", fill_tint=False)

# =====================================================================  3 Rent Roll
rr_ = wb.create_sheet(RR)
ws = rr_
S.set_widths(ws, {"A": 3, "B": 2, "C": 22, "D": 20, "E": 14, "F": 14, "G": 14, "H": 14, "I": 11, "J": 11, "K": 22, "L": 15, "M": 2, "N": 3})
D.page_header(ws, "Rent Roll", "Who has paid this month. LATE uses today's date, each door's due day and its grace days.",
              "Fills itself in from 1 Setup and 2 Log", span=("B", "M"), side_cols=3)
for rr, ht in zip(range(11, 17), (12, 18, 34, 21, 21, 12)):
    D.h(ws, rr, ht)
e = ws["C12"]; e.value = "MONTH TO CHECK"; e.font = D.f(8, True, "muted"); e.alignment = Alignment(vertical="bottom")
card_label(ws, f"C{RR_MIN}", "Month to check")
mc = S.input_cell(ws, f"D{RR_MIN}", EX_MONTH, MONTH_FMT, name="MonthToCheck", wb=wb,
                  validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"), allow_blank=True, prompt_title="Month to check",
                  prompt="Type any date in the month you want, for example 10/1/2026. Leave blank to check this month.",
                  error="Type a real date between 2000 and 2100, for example 10/1/2026, or leave the cell blank.")
box(mc, 11); mc.alignment = Alignment(horizontal="center", vertical="center")
card_label(ws, f"C{RR_MEFF}", "Showing")
me = ws[f"D{RR_MEFF}"]; me.value = f'=IF($D${RR_MIN}="",TODAY(),$D${RR_MIN})'; me.number_format = MONTH_FMT
me.font = D.f(10, True); me.alignment = Alignment(horizontal="center", vertical="center")
card_label(ws, "C15", "Today")
td = ws["D15"]; td.value = "=TODAY()"; td.number_format = DATE; td.font = D.f(10); td.alignment = Alignment(horizontal="center", vertical="center")
for rr in (RR_MEFF, 15):
    for k in ("C", "E"):
        ws[f"{k}{rr}"].border = Border(bottom=D.side("hair"))
    ws[f"D{rr}"].border = Border(bottom=D.side("hair"))
D.frame(ws, 11, 16, "B", "F")

RT_TOP = 18
r = card_title(ws, RT_TOP, "C", "L", "Every door, this month",
               "Rent counts when a 2 Log row has the Rent category and a door picked. Partial payments add up. Late fees go under Other rental income.")
RHEAD = r
table_head(ws, RHEAD, [("C", "Door", "left"), ("D", "Property", "left"), ("E", "Tenant", "left"), ("F", "Rent due", "right"),
                       ("G", "Received this month", "right"), ("H", "Still owed", "right"), ("I", "Due date", "right"),
                       ("J", "Status", "center"), ("K", "Lease", "left"), ("L", "Deposit held", "right")], "C", "L")
R1 = RHEAD + 1
R2 = R1 + ND - 1
M = f"$D${RR_MEFF}"
for i in range(ND):
    k = R1 + i; s_ = D_FIRST + i
    SR = lambda c: f"{q(SET)}!${c}${s_}"  # noqa: E731
    vals = {
        "C": (f'=IF({SR("C")}="","",{SR("C")})', None, "left"),
        "D": (f'=IF($C{k}="","",{SR("D")})', None, "left"),
        "E": (f'=IF($C{k}="","",{SR("E")})', None, "left"),
        "F": (f'=IF($C{k}="","",N({SR("F")}))', MONEY2, "right"),
        "G": (f'=IF($C{k}="","",SUMIFS({LA},{LDOOR},$C{k},{LC},{catn(0)},{LD},">="&DATE(YEAR({M}),MONTH({M}),1),{LD},"<="&EOMONTH({M},0)))', MONEY2, "right"),
        "H": (f'=IF($C{k}="","",MAX(0,F{k}-G{k}))', MONEY2, "right"),
        "I": (f'=IF($C{k}="","",DATE(YEAR({M}),MONTH({M}),MIN(MAX(1,N({SR("G")})),DAY(EOMONTH({M},0)))))', "mmm d", "right"),
        "J": (f'=IF($C{k}="","",IF(AND(F{k}=0,G{k}=0),"No rent set",IF(H{k}<=0,"Paid",IF(TODAY()>I{k}+N({SR("H")}),"LATE",IF(G{k}>0,"Partial","Due")))))', None, "center"),
        "K": (f'=IF(OR($C{k}="",{SR("I")}=""),"",IF({SR("I")}<TODAY(),"Lease ended",IF({SR("I")}-TODAY()<=60,IF({SR("I")}=TODAY(),"Ends today","Ends in "&({SR("I")}-TODAY())&IF({SR("I")}-TODAY()=1," day"," days")),"Ends "&TEXT({SR("I")},"mmm d, yyyy"))))', None, "left"),
        "L": (f'=IF($C{k}="","",N({SR("J")}))', MONEY2, "right"),
    }
    for col, (fml, fmt, align) in vals.items():
        c = ws[f"{col}{k}"]; c.value = fml
        if fmt:
            c.number_format = fmt
        c.font = D.f(10, col in ("C", "J")); c.alignment = Alignment(horizontal=align, vertical="center")
    hair(ws, k, "C", "L"); D.h(ws, k, D.GRID_ROW)
RTOT = R2 + 1
c = ws[f"C{RTOT}"]; c.value = "Totals"; c.font = D.f(10, True); c.alignment = Alignment(vertical="center")
for col in ("F", "G", "H", "L"):
    v = ws[f"{col}{RTOT}"]; v.value = f"=SUM({col}{R1}:{col}{R2})"; v.number_format = MONEY2; v.font = D.f(10, True)
    v.alignment = Alignment(horizontal="right", vertical="center")
hair(ws, RTOT, "C", "L", top_ink=True, tint="total_tint"); D.h(ws, RTOT, D.GRID_ROW + 3)
RSUM = RTOT + 1
c = ws[f"C{RSUM}"]
c.value = (f'=IF(COUNTIF(J{R1}:J{R2},"LATE")=0,"No late rent for this month.",COUNTIF(J{R1}:J{R2},"LATE")'
           f'&" door(s) LATE. Total still owed this month: "&TEXT(H{RTOT},"$#,##0.00"))')
c.font = D.f(10, True); c.alignment = Alignment(vertical="center"); ws.merge_cells(f"C{RSUM}:L{RSUM}")
ws.conditional_formatting.add(f"C{RSUM}", FormulaRule(formula=[f'COUNTIF($J${R1}:$J${R2},"LATE")>0'], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
D.h(ws, RSUM, D.GRID_ROW + 4)
RNOTE = RSUM + 1
c = ws[f"C{RNOTE}"]; c.value = "Status: Paid, Partial, Due, LATE once the due day plus grace days has passed, or No rent set. Lease shows a heads-up 60 days ahead."
c.font = D.f(9, False, "muted"); c.alignment = Alignment(vertical="center"); ws.merge_cells(f"C{RNOTE}:L{RNOTE}"); D.h(ws, RNOTE, D.GRID_ROW)
D.h(ws, RNOTE + 1, 12)
RT_END = RNOTE + 1
D.frame(ws, RT_TOP, RT_END, "B", "M")
for word, color in (("LATE", "accent"), ("Paid", "teal"), ("Partial", "gold"), ("Due", "ink"), ("No rent", "muted")):
    ws.conditional_formatting.add(f"J{R1}:J{R2}", FormulaRule(formula=[f'LEFT($J{R1},{len(word)})="{word}"'],
                                  font=Font(color=D.T[color], bold=color != "muted"), stopIfTrue=True))
ws.conditional_formatting.add(f"K{R1}:K{R2}", FormulaRule(formula=[f'OR(LEFT($K{R1},7)="Ends in",LEFT($K{R1},5)="Lease")'],
                              font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))

# ---- hero tiles
NDOORS = f'COUNTIF($C${R1}:$C${R2},"?*")'
D.tile(ws, 6, "B", "F", f'="STILL OWED FOR "&UPPER(TEXT({M},"mmmm yyyy"))', f"=$H${RTOT}", S.FMT["usd2"],
       f'=IF({NDOORS}=0,"Add your doors on the 1 Setup tab",TEXT($G${RTOT},"$#,##0.00")&" of "&TEXT($F${RTOT},"$#,##0.00")&" received so far")',
       dark=True)
v, sub, _ = D.tile(ws, 6, "H", "M", "DOORS LATE THIS MONTH", f'=COUNTIF($J${R1}:$J${R2},"LATE")', INT,
                   f'=IF({NDOORS}=0,"No doors yet",COUNTIF($J${R1}:$J${R2},"Paid")&" of "&{NDOORS}&IF({NDOORS}=1," door"," doors")&" paid in full")')
D.status_rule(ws, v.coordinate, f'COUNTIF($J${R1}:$J${R2},"LATE")>0', fill_tint=False)
D.h(ws, 10, 14)
D.stat_strip(ws, 12, [
    ("I:J", "Rent due", f"=$F${RTOT}", USD0),
    ("K:K", "Received", f"=$G${RTOT}", USD0),
    ("L:L", "Deposits held", f"=$L${RTOT}", USD0),
])
s = ws["I15"]; s.value = "Blank month checks this month. Change it any time to look back."
s.font = D.f(9, False, "muted"); s.alignment = Alignment(vertical="center"); ws.merge_cells("I15:L15")
D.frame(ws, 11, 16, "H", "M")
RR_NOTES = [
    ("How LATE works", "A door shows LATE when today's date is past its due day plus its grace days and the month is not paid in "
     "full. Due day 31 means the last day of a shorter month. Whether and when you may charge a late fee is set by your lease "
     "and your state and local law."),
    ("Leases", "Lease end comes from the 1 Setup tab. Ends in shows 60 days ahead; Lease ended shows after the date."),
    ("Before you rely on a number", NOTICE),
]
rr_last = notes_card(ws, RT_END + 2, "C", "L", "B", "M", RR_NOTES)
RR_LAST = rr_last + 1
D.h(ws, RR_LAST, 14)
fill_heights(ws, RR_LAST)
D.paint_canvas(ws, RR_LAST, "Z")
S.finish_sheet(ws, PRODUCT, RR_LAST, span=("A", "N"), freeze="A11", tab_color="accent")
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToHeight = 1  # one landscape page: the door table never splits

# =====================================================================  5 Monthly
mo = wb.create_sheet(MO)
S.set_widths(mo, {"A": 3, "B": 2, "C": 26, "D": 24, "E": 13, "F": 13, "G": 14, "H": 13, "I": 14, "J": 18, "K": 2, "L": 3})
D.page_header(mo, "Monthly summary", "Month by month for the year on the 4 Property P&L tab, for all properties or one.",
              AS_OF, span=("B", "K"), side_cols=3)
for rr, ht in zip(range(11, 17), (12, 18, 34, 21, 21, 12)):
    D.h(mo, rr, ht)
e = mo["C12"]; e.value = "PROPERTY AND YEAR"; e.font = D.f(8, True, "muted"); e.alignment = Alignment(vertical="bottom")
card_label(mo, "C13", "Property to show")
mp = mo["D13"]; mp.value = None; mp.protection = mp.protection.copy(locked=False)
S.add_validation(mo, "D13", "list", options=["x"], allow_blank=True)
dvm = mo.data_validations.dataValidation[-1]
dvm.formula1 = PROP_LIST; dvm.promptTitle = "Property to show"; dvm.showInputMessage = True
dvm.prompt = "Pick one property, or leave the cell blank to show all properties."
dvm.error = "Pick a property from the list on the 1 Setup tab, or leave the cell blank for all properties."
box(mp); mp.alignment = Alignment(horizontal="left", vertical="center", indent=1)
S.define_name(wb, "PropertyToShow", mo, "D13")
card_label(mo, "C14", "Showing")
sh_ = mo["D14"]; sh_.value = '=IF(OR($D$13="",$D$13="All properties"),"All properties",$D$13)'; sh_.font = D.f(10, True)
sh_.alignment = Alignment(horizontal="left", vertical="center", indent=1)
card_label(mo, "C15", "Year (from tab 4)")
yy = mo["D15"]; yy.value = f"={YEAR}"; yy.number_format = YEAR_FMT; yy.font = D.f(10, True)
yy.alignment = Alignment(horizontal="left", vertical="center", indent=1)
for rr in (14, 15):
    for k in ("C", "D"):
        mo[f"{k}{rr}"].border = Border(bottom=D.side("hair"))
D.frame(mo, 11, 16, "B", "E")
MT_TOP = 18
r = card_title(mo, MT_TOP, "C", "J", "Month by month",
               "Money in less Schedule E expenses is profit before depreciation; less principal and improvements is cash flow.")
MHEAD = r
table_head(mo, MHEAD, [("C", "Month", "left"), ("D", "Money in", "right"), ("E", "Expenses", "right"),
                       ("F", "Profit before depreciation", "right"), ("G", "Principal and improvements", "right"),
                       ("H", "Cash flow", "right"), ("I", "Running cash flow", "right"), ("J", "", "left")], "C", "J", height=34)
M1 = MHEAD + 1
CRIT = 'IF(OR($D$13="",$D$13="All properties"),"*",$D$13)'
Y = "$D$15"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def msum(a, b, m):
    """Boolean SUMPRODUCT, not SUMPRODUCT(SUMIFS(..., range criteria)): Google Sheets only reads the first
    category of a range criterion inside SUMIFS (v1 bug, tested in Sheets Oct 7, 2026); this form gives the
    same answer in Excel, LibreOffice and Google Sheets."""
    cats = f"{q(CAT)}!$D${C_FIRST + a}:$D${C_FIRST + b}"
    return (f'=SUMPRODUCT({LA}*({LC}<>"")*(COUNTIF({cats},{LC})>0)*({LP}<>"")'
            f'*((({LP}=$D$13)+($D$13="")+($D$13="All properties"))>0)'
            f'*({LD}>=DATE({Y},{m},1))*({LD}<=EOMONTH(DATE({Y},{m},1),0)))')


HR = f"$H${M1}:$H${M1 + 11}"
for m in range(1, 13):
    k = M1 + m - 1
    row = {"C": (MONTHS[m - 1], None), "D": (msum(0, 1, m), MONEY0), "E": (msum(2, 15, m), MONEY0), "F": (f"=D{k}-E{k}", MONEY0),
           "G": (msum(16, 17, m), MONEY0), "H": (f"=F{k}-G{k}", MONEY0), "I": (f"=H{k}" if m == 1 else f"=I{k - 1}+H{k}", MONEY0)}
    for col, (val, fmt) in row.items():
        c = mo[f"{col}{k}"]; c.value = val
        if fmt:
            c.number_format = fmt
        c.font = D.f(10, col in ("H",)); c.alignment = Alignment(horizontal="left" if col == "C" else "right", vertical="center")
    b = mo[f"J{k}"]
    b.value = f'=IFERROR(REPT("{D.BAR_CHAR}",ROUND(ABS(H{k})/MAX(MAX({HR}),-MIN({HR}))*10,0)),"")'
    b.font = D.f(9, False, "teal"); b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    mo.conditional_formatting.add(f"J{k}", FormulaRule(formula=[f"$H{k}<0"], font=Font(color=D.T["accent"]), stopIfTrue=True))
    hair(mo, k, "C", "J"); D.h(mo, k, D.GRID_ROW)
M12 = M1 + 11
MY = M12 + 1
c = mo[f"C{MY}"]; c.value = "Year"; c.font = D.f(10, True); c.alignment = Alignment(vertical="center")
for col in ("D", "E", "F", "G", "H"):
    v = mo[f"{col}{MY}"]; v.value = f"=SUM({col}{M1}:{col}{M12})"; v.number_format = MONEY0; v.font = D.f(10, True)
    v.alignment = Alignment(horizontal="right", vertical="center")
v = mo[f"I{MY}"]; v.value = f"=I{M12}"; v.number_format = MONEY0; v.font = D.f(10, True); v.alignment = Alignment(horizontal="right", vertical="center")
hair(mo, MY, "C", "J", top_ink=True, tint="total_tint"); D.h(mo, MY, D.GRID_ROW + 3)
for col in ("F", "H", "I"):
    sign_rule(mo, f"{col}{M1}:{col}{MY}")
MNOTE = MY + 1
c = mo[f"C{MNOTE}"]; c.value = "Months after today show a dash until you log them, except the running total. Deposits are left out because they belong to the tenant."
c.font = D.f(9, False, "muted"); c.alignment = Alignment(vertical="center"); mo.merge_cells(f"C{MNOTE}:J{MNOTE}"); D.h(mo, MNOTE, D.GRID_ROW)
D.h(mo, MNOTE + 1, 12)
MT_END = MNOTE + 1
D.frame(mo, MT_TOP, MT_END, "B", "K")
D.tile(mo, 6, "B", "E", f'="CASH FLOW, "&$D$15&", "&UPPER($D$14)', f"=$H${MY}", USD0,
       f'="Profit "&TEXT($F${MY},"$#,##0;-$#,##0")&" less "&TEXT($G${MY},"$#,##0")&" principal and improvements"',
       dark=True)
NEG = f'COUNTIF($H${M1}:$H${M12},"<0")'
v, sub, _ = D.tile(mo, 6, "G", "K", "MONTHS THAT LOST CASH", f"={NEG}", INT,
                   f'=IF(COUNTIF($H${M1}:$H${M12},"<>0")=0,"No log rows in the year shown yet",IF({NEG}=0,"Every month with log rows kept money",'
                   f'"Lowest: "&INDEX($C${M1}:$C${M12},MATCH(MIN($H${M1}:$H${M12}),$H${M1}:$H${M12},0))&", "&TEXT(MIN($H${M1}:$H${M12}),"$#,##0;-$#,##0")))')
D.status_rule(mo, v.coordinate, f"{NEG}>0", fill_tint=False)
D.h(mo, 10, 14)
D.stat_strip(mo, 12, [
    ("H:H", "Money in", f"=$D${MY}", USD0),
    ("I:J", "Expenses", f"=$E${MY}", USD0),
])
s = mo["H15"]; s.value = "For the property and year shown."
s.font = D.f(9, False, "muted"); s.alignment = Alignment(vertical="center"); mo.merge_cells("H15:J15")
D.frame(mo, 11, 16, "G", "K")
MO_NOTES = [
    ("Same numbers as the P&L", "Money in, expenses and the cash items use the same categories as the 4 Property P&L tab, so the "
     "Year row matches its All properties column (or one property's column when you pick it here)."),
    ("Before you rely on a number", NOTICE),
]
mo_last = notes_card(mo, MT_END + 2, "C", "J", "B", "K", MO_NOTES)
MO_LAST = mo_last + 1
D.h(mo, MO_LAST, 14)
fill_heights(mo, MO_LAST)
D.paint_canvas(mo, MO_LAST, "Z")
S.finish_sheet(mo, PRODUCT, MO_LAST, span=("A", "L"), freeze="A11", tab_color="accent")

# ---- tab order: 1 Setup, 2 Log, 3 Rent Roll, 4 Property P&L, 5 Monthly, 6 Categories
wb._sheets = [st, lg, rr_, pl, mo, ct]

# =====================================================================  Start Here and Terms (card style)
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            mm = re.match(r"^([A-Z0-9][A-Za-z0-9 ,&]{1,34})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if mm:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, "Rental Property Spreadsheet", "Start here. For 1 to 10 doors: type your doors once, log every dollar, read the answers.",
              f"Checked {CHECKED}", span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "A landlord spreadsheet built for 1 to 10 doors, not 50. Type your properties and doors once on 1 Setup and log every rent "
    "payment and bill on 2 Log. The 3 Rent Roll shows who has paid for any month, with LATE flags and lease heads-ups. The "
    "4 Property P&L shows profit and cash flow for each property on the IRS Schedule E line names, and 5 Monthly shows the year "
    "month by month.", "para")])
r = sh_card(sh, r, "Three steps", [
    (1, "1 Setup. Type each property once, then one row per door: rent, due day, grace days, lease end and deposit.", "num"),
    (2, "2 Log. One row per payment or bill. Pick the property, the door (for rent) and the category from the dropdowns. Amounts are always positive.", "num"),
    (3, "Read the answers. 3 Rent Roll for this month, 4 Property P&L for the year (pick a year and a property), 5 Monthly for the months.", "num"),
])
r = sh_card(sh, r, "Before you rely on a number", [(None, NOTICE + " " + TAX_NOTE, "note")])
S.page_break_before(sh, r)
r = sh_card(sh, r, "How to read the cells", [
    ((1150, MONEY2), "Yellow cells with blue numbers are yours to type in. Each one shows a hint when you select it.", "chip-in"),
    ((36850, USD0), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((13569, USD0), "Dark tiles are your answers. Each answer tab has one at the top that stays on screen as you scroll.", "chip-key"),
    (None, "Each tab is protected without a password, so you can still change anything. To edit a locked cell, use Review, "
           "Unprotect Sheet in Excel, or Data, Protect sheets and ranges in Google Sheets.", "note"),
])
r = sh_card(sh, r, "The example", [(None,
    "The file opens with a made-up duplex and a single-family house, January to October 2026. On the 3 Rent Roll for October, "
    "MAPLE-A has paid, MAPLE-B paid $600 of $1,100 and OAK has paid nothing: $1,750 of $3,900 received, $2,150 still owed. "
    "Status and the lease countdown use today's date, so the unpaid sample doors show LATE once their grace days pass. On the "
    "4 Property P&L, Oak Street House shows $4,844 of profit before depreciation but keeps $550 in cash after $2,844 of mortgage "
    "principal and a $1,450 water heater. To start fresh, delete the example rows on 1 Setup and 2 Log, then clear the Month to "
    "check on 3 Rent Roll and the Year and Property on 4 Property P&L.", "para")])
r = sh_card(sh, r, "Which file to open", [
    (None, "Excel on a computer. Double-click the .xlsx file. If Excel shows a yellow Protected View bar, click Enable Editing.", "para"),
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. It uses only functions "
           "Google Sheets supports. No macros, no add-ons, no sign-up.", "para"),
    (None, "On a phone. The Etsy app can't download files. Sign in at etsy.com in your phone's browser, go to You, Purchases, "
           "Download Files, unzip, then open the .xlsx in the Google Sheets app or the Microsoft Excel app. Without a "
           "qualifying Microsoft 365 plan, the Excel app edits only on screens up to 10.1 inches (Microsoft support, checked "
           f"{CHECKED}). Setup is easier on a computer.", "para"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "Good to know", [
    (None, f"Room for a busy year. 2 Log has {N_ROWS:,} rows ready to type in, and every total reads 5,000 rows. Need more? "
           "Unprotect the tab (no password) and keep typing below. Rows can be in any order.", "para"),
    (None, "Interest, principal, improvements. Mortgage interest is an expense. Principal and capital improvements are not "
           "expenses; they lower cash flow only. Depreciation (Schedule E line 18) is left to your tax preparer.", "para"),
    (None, "Security deposits. They are tracked on their own rows and left out of income. IRS Publication 527: don't include a "
           "deposit in income if you plan to return it. A deposit meant to be the last month's rent is advance rent and counts as "
           "income when received. How a deposit must be held, how fast it must be returned and what you may keep is set by state "
           "and local landlord-tenant law; check yours.", "para"),
    (None, "Private information. Tenant names and deposits are personal information, so keep this file somewhere secure. "
           "Save a copy each month as a backup.", "para"),
    (None, "Tabs. Start Here, 1 Setup, 2 Log, 3 Rent Roll, 4 Property P&L, 5 Monthly, 6 Categories, Terms.", "para"),
])
SH_LAST = r - 1
D.paint_canvas(sh, SH_LAST, "Z")
S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
               footer_notice="Estimates only. Not professional advice. See the Terms tab or LICENSE-AND-DISCLAIMER.txt.")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"Version {VERSION}. Read this before you rely on a number.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More tools from the shop", [
    (None, ("One-Tab Business Expense Tracker: a simple log for a side business, on Schedule C lines", "https://www.etsy.com/listing/4589291725"), "link"),
    (None, ("Pricing Calculator for Airbnb and VRBO Hosts: a nightly rate that covers your costs", "https://www.etsy.com/listing/4585252476"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this spreadsheet helps you keep track of your rentals, a short review on Etsy helps other landlords find it. Thank you.", "para"),
    (None, "If a file won't open or something looks wrong, message the shop on Etsy and it will be fixed.", "para"),
])
TM_LAST = r - 1
D.paint_canvas(tm, TM_LAST, "Z")
S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb) if "input without validation" not in p]
if problems:
    print("AUDIT:", *problems, sep="\n  ")
wb.save(OUT)
ADDR = {"P_FIRST": P_FIRST, "D_FIRST": D_FIRST, "L_FIRST": L_FIRST, "C_FIRST": C_FIRST, "PL_HDR": HDR, "PROW": PROW,
        "R_IN": R_IN, "R_EXP": R_EXP, "R_PRO": R_PRO, "R_CASH": R_CASH, "checks": [k1, k2, k3, k4], "RR_R1": R1, "RTOT": RTOT,
        "RSUM": RSUM, "M1": M1, "MY": MY, "PL_YIN": PL_YIN, "PL_PICK": PL_PICK, "RR_MIN": RR_MIN, "cf_first": cf_first}
print("saved", OUT, wb.sheetnames)
import json  # noqa: E402
json.dump(ADDR, open(os.path.join(HERE, "addresses.json"), "w"), indent=1)
print(ADDR)
