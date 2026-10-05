#!/usr/bin/env python3
"""Renders START-HERE PDFs (Letter, A4) and five 2000x2000 listing images from the recalculated sample workbook."""
import openpyxl, datetime as dt, html, os
from playwright.sync_api import sync_playwright
HERE = os.path.dirname(os.path.abspath(__file__))
wb = openpyxl.load_workbook(os.path.join(HERE, "recalc/sample.xlsx"), data_only=True)
rr, pl, mo, lg = wb["Rent Roll"], wb["Property P&L"], wb["Monthly"], wb["Log"]
def m(v, d=0):
    if v is None or v == "": return ""
    if isinstance(v, (int, float)):
        s = f"${abs(v):,.{d}f}"; return f"({s})" if v < 0 else ("-" if v == 0 else s)
    return html.escape(str(v))
FONTS = """
@font-face{font-family:Poppins;font-weight:400;src:url(file:///usr/share/fonts/truetype/google-fonts/Poppins-Regular.ttf)}
@font-face{font-family:Poppins;font-weight:600;src:url(file:///usr/share/fonts/truetype/google-fonts/Poppins-Medium.ttf)}
@font-face{font-family:Poppins;font-weight:700;src:url(file:///usr/share/fonts/truetype/google-fonts/Poppins-Bold.ttf)}
@font-face{font-family:Carlito;font-weight:400;src:url(file:///usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf)}
@font-face{font-family:Carlito;font-weight:700;src:url(file:///usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf)}
:root{--paper:#F5F2EC;--ink:#1D2433;--acc:#C8502F;--line:#D9D3C7;--soft:#FBF9F5}
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:Carlito,sans-serif;color:var(--ink)}
h1,h2,h3,.p{font-family:Poppins,sans-serif}
"""
# ---------- tables from the real recalculated file ----------
def rent_roll_table(big=False):
    rows = ""
    for r in range(8, 11):
        v = [rr.cell(r, c).value for c in range(1, 11)]
        st = v[7]; cls = {"LATE": "late", "Paid": "paid", "Partial": "part"}.get(st, "")
        lease = html.escape(v[8]); lcls = "acc" if lease.startswith("Ends in") else ""
        rows += f"<tr><td><b>{v[0]}</b></td><td>{v[1]}</td><td>{m(v[3])}</td><td>{m(v[4])}</td><td>{m(v[5])}</td><td><span class='st {cls}'>{st}</span></td><td class='{lcls}'>{lease}</td></tr>"
    t = [rr.cell(18, c).value for c in range(1, 11)]
    rows += f"<tr class='tot'><td>Totals</td><td></td><td>{m(t[3])}</td><td>{m(t[4])}</td><td>{m(t[5])}</td><td></td><td></td></tr>"
    return f"<table class='sheet'><tr><th>Door</th><th>Property</th><th>Rent due</th><th>Received</th><th>Still owed</th><th>Status</th><th>Lease</th></tr>{rows}</table><div class='flag'>{html.escape(rr['A20'].value)}</div>"
def pl_table(keep=None):
    rows = ""
    for r in range(8, 36):
        a = pl.cell(r, 1).value
        if a is None: continue
        vals = [pl.cell(r, c).value for c in (3, 4, 13)]
        if keep and a not in keep and vals[0] is not None: continue
        if vals[0] is None and vals[2] is None:
            if not keep: rows += f"<tr class='sec'><td colspan=5>{html.escape(a)}</td></tr>"
            continue
        if keep and a == "Total expenses":
            rows += "<tr><td colspan=5 class='ln' style='font-style:italic'>plus 11 more expense rows in the file</td></tr>"
        b = a in ("Total money in", "Total expenses", "Profit before depreciation", "Cash flow (what you kept)")
        rows += f"<tr class='{'tot' if b else ''}'><td>{html.escape(a)}</td><td class='ln'>{html.escape(pl.cell(r,2).value or '')}</td>" + "".join(f"<td>{m(x)}</td>" for x in vals) + "</tr>"
    return f"<table class='sheet pl'><tr><th>Category</th><th>Schedule E line</th><th>Maple Duplex</th><th>Oak Street House</th><th>All properties</th></tr>{rows}</table>"
def monthly_table():
    cols = list(range(2, 12)) + [14]
    hdr = "".join(f"<th>{mo.cell(8,c).value}</th>" for c in cols)
    rows = ""
    for r in range(9, 15):
        a = mo.cell(r, 1).value; b = a in ("Profit before depreciation", "Cash flow")
        cells = ""
        for c in cols:
            v = mo.cell(r, c).value; cells += f"<td class='{'neg' if isinstance(v,(int,float)) and v<0 else ''}'>{m(v)}</td>"
        rows += f"<tr class='{'tot' if b else ''}'><td>{a}</td>{cells}</tr>"
    return f"<table class='sheet mo'><tr><th></th>{hdr}</tr>{rows}</table>"
def log_table(n=9):
    rows = ""; cnt = 0
    for r in range(6, 200):
        d = lg.cell(r, 1).value
        if d is None: break
        if d.month < 9: continue
        v = [lg.cell(r, c).value for c in range(1, 8)]
        rows += f"<tr><td>{d:%m/%d/%Y}</td><td>{v[1]}</td><td>{v[2] or ''}</td><td>{v[3]}</td><td>{m(v[4],2)}</td><td>{html.escape(v[6] or '')}</td></tr>"
        cnt += 1
        if cnt >= n: break
    return f"<table class='sheet'><tr><th>Date</th><th>Property</th><th>Door</th><th>Category</th><th>Amount</th><th>Note</th></tr>{rows}</table>"
TABLE_CSS = """
table.sheet{border-collapse:collapse;width:100%;background:#fff;font-family:Carlito,sans-serif}
.sheet th{background:var(--ink);color:#fff;text-align:left;padding:.55em .7em;font-weight:700}
.sheet td{padding:.5em .7em;border-bottom:1px solid var(--line)}
.sheet td:first-child,.sheet th{white-space:nowrap}
.sheet tr.tot td{font-weight:700;border-top:2px solid var(--ink)}
.sheet tr.sec td{color:var(--acc);font-weight:700;background:var(--soft)}
.sheet td.ln{color:#6a7080;font-size:.85em}
.sheet td.neg{color:var(--acc);font-weight:700}
.st{display:inline-block;padding:.15em .7em;border-radius:6px;font-weight:700}
.st.late{background:var(--acc);color:#fff}.st.paid{background:#DCEBDD}.st.part{background:#FBE3B0}
td.acc{color:var(--acc);font-weight:700}
.flag{margin-top:.8em;color:var(--acc);font-weight:700;font-family:Poppins}
"""
# ---------- guide ----------
CHECK = "October 5, 2026"
guide = f"""<html><head><meta charset=utf-8><style>{FONTS}{TABLE_CSS}
@page{{margin:0}}
.page{{width:100%;height:100vh;padding:0.6in 0.7in;page-break-after:always;position:relative;background:var(--paper);font-size:11.5pt;line-height:1.45}}
.page:last-child{{page-break-after:auto}}
h1{{font-size:24pt;line-height:1.15;margin-bottom:.15in}} h2{{font-size:14pt;color:var(--acc);margin:.22in 0 .08in}}
.kick{{font-family:Poppins;font-weight:600;color:var(--acc);letter-spacing:.08em;font-size:9.5pt;text-transform:uppercase}}
.note{{font-size:9pt;color:var(--acc);border-left:3px solid var(--acc);padding-left:.12in;margin:.12in 0}}
ol,ul{{margin-left:.25in}} li{{margin:.05in 0}}
.box{{background:#fff;border:1px solid var(--line);padding:.14in .18in;margin:.1in 0;border-radius:6px}}
.foot{{position:absolute;bottom:.4in;left:.7in;right:.7in;font-size:8.5pt;color:#6a7080;border-top:1px solid var(--line);padding-top:.06in}}
table.sheet{{font-size:9.5pt}}
</style></head><body>
<div class=page>
<div class=kick>Start here</div>
<h1>Rental Property Spreadsheet<br>for 1 to 10 Doors</h1>
<p>One workbook for rent, expenses, late rent and profit per property, sized for a small landlord instead of a 50-unit company. Version 1, {CHECK}.</p>
<div class=note>For general information and planning only. This is not legal, tax, financial, medical or other professional advice, and using it doesn't create a professional relationship. Results are estimates based on the numbers you enter. Figures were checked on {CHECK} and can change. See the Terms tab before you rely on anything here. Calculations are estimates for planning and depend on your inputs. They aren't accounting, tax or financial advice.</div>
<h2>Which file to open</h2>
<div class=box><b>The only file you need is Rental-Property-Spreadsheet-1-to-10-Doors.xlsx.</b> The other files are this guide and the terms.</div>
<ul>
<li><b>Computer with Excel:</b> unzip the download, then double-click the .xlsx.</li>
<li><b>Computer without Excel:</b> go to sheets.google.com, open a blank sheet, then File &gt; Import &gt; Upload, pick the .xlsx and choose Replace spreadsheet. The workbook uses only functions Google Sheets supports, and the dropdowns import with the file.</li>
<li><b>Phone or tablet:</b> a computer is easier for setup. If you only have a phone, save the zip to your Files app, tap it to unzip, then open the .xlsx in the free Microsoft Excel or Google Sheets app.</li>
<li><b>Can't find the download?</b> The Etsy app can't download files. Sign in at etsy.com in a browser (computer or phone), go to Your account &gt; Purchases, and select Download Files next to this order.</li>
</ul>
<h2>What's inside the workbook</h2>
<ul>
<li><b>Setup:</b> up to 10 properties and 10 doors, typed once.</li>
<li><b>Log:</b> one row per payment or bill, with dropdowns.</li>
<li><b>Rent Roll:</b> Paid, Partial, Due or LATE for each door, plus leases ending within 60 days.</li>
<li><b>Property P&amp;L:</b> one column per property, expenses on the IRS Schedule E line names, profit and cash flow.</li>
<li><b>Monthly:</b> money in, expenses and cash flow by month.</li>
<li><b>Categories, Terms, Thank You.</b></li>
</ul>
<div class=foot>ProofNotFluff | Rental Property Spreadsheet for 1 to 10 Doors | Page 1 of 3</div>
</div>
<div class=page>
<div class=kick>Five-minute setup</div>
<h2 style="margin-top:.05in">1. Setup tab</h2>
<p>Type each property once with a short name. Then one row per door: Door ID (any short name, like MAPLE-A), property, tenant, monthly rent, due day, grace days, lease end and deposit held. A duplex is one property with two doors.</p>
<h2>2. Log tab</h2>
<p>One row for every dollar in or out. Pick property, door and category from the dropdowns. Amounts are always positive; the category says which way the money moved. A rent row with no door turns pink so you notice it.</p>
{log_table(7)}
<h2>3. Rent Roll tab</h2>
<p>Type any date in the month you want to check. A door shows LATE once today is past its due day plus grace days and it isn't fully paid. Status and the lease countdown use today's date, so the sample may show more LATE doors when you open it.</p>
{rent_roll_table()}
<h2>4. Property P&amp;L and Monthly</h2>
<p>Type the year. Everything else fills itself. Mortgage interest is an expense; principal and improvements lower cash flow only. Depreciation (Schedule E line 18) is left to your tax preparer.</p>
<div class=foot>The tables on this page are the sample data in the file as of October 5, 2026: a duplex and a single-family house, January to October 2026. Delete the sample rows on Setup and Log to start fresh. Page 2 of 3</div>
</div>
<div class=page>
<div class=kick>Good to know</div>
<h2 style="margin-top:.05in">How the categories were set</h2>
<p>The 14 expense categories use the names of the IRS Schedule E (Form 1040) Part I expense lines, so totals are easy to hand to a tax preparer. They are for organizing; your preparer decides what goes where on your return.</p>
<ul>
<li>IRS, Schedule E (Form 1040) 2025, Part I lines 3 to 19, checked {CHECK}: irs.gov/pub/irs-pdf/f1040se.pdf</li>
<li>IRS, Instructions for Schedule E, line 14: repairs keep property in working condition; improvements better, restore or adapt it. Checked {CHECK}: irs.gov/instructions/i1040se</li>
<li>IRS, Publication 527: don't include a security deposit in income if you plan to return it to the tenant. Checked {CHECK}: irs.gov/publications/p527</li>
</ul>
<h2>Small things that save time</h2>
<ul>
<li>Yellow cells, including every row of the Log, are for typing. The Rent Roll, Property P&amp;L and Monthly tabs are locked against accidental typing (Review &gt; Unprotect Sheet, no password).</li>
<li>Formulas read the Log down to row 5,000. Keep adding rows under the last one, and sort by date whenever you like.</li>
<li>The Property P&amp;L shows a count of Log rows with an amount but no property or category, so nothing goes missing.</li>
<li>Late fees, pet rent or a kept deposit go under Other rental income.</li>
</ul>
<h2>Terms</h2>
<p>Personal license for you or one business you own. No resale or sharing. Full terms are in LICENSE-AND-DISCLAIMER.txt and on the Terms tab. Made with AI assistance and reviewed and tested by the shop owner.</p>
<div class=box style="margin-top:.3in">
<b class=p>Thank you for buying from ProofNotFluff.</b><br>
If this spreadsheet helps you keep track of your rentals, a short review on Etsy helps other landlords find it.<br>
Message me on Etsy if anything doesn't open, and I'll fix it.
</div>
<div class=foot>ProofNotFluff | Copyright 2026 ProofNotFluff | Page 3 of 3</div>
</div>
</body></html>"""
open(os.path.join(HERE, "guide.html"), "w").write(guide)

# ---------- listing images ----------
IMG = f"""<html><head><meta charset=utf-8><style>{FONTS}{TABLE_CSS}
body{{background:#999}}
.v{{width:2000px;height:2000px;background:var(--paper);position:relative;overflow:hidden;padding:130px 120px}}
.v.ink{{background:var(--ink);color:#fff}}
.kick{{font-family:Poppins;font-weight:600;color:var(--acc);letter-spacing:.1em;font-size:44px;text-transform:uppercase}}
h1{{font-size:118px;line-height:1.05;margin:30px 0}} h2{{font-size:84px;line-height:1.1;margin:24px 0 50px}}
.sub{{font-size:52px;line-height:1.35;color:#4a5060}} .ink .sub{{color:#d8d4cc}}
.card{{background:#fff;border-radius:28px;box-shadow:0 30px 80px rgba(29,36,51,.18);padding:50px;font-size:40px}}
.card table.sheet{{font-size:1em}}
.foot{{position:absolute;left:120px;right:120px;bottom:90px;font-family:Poppins;font-size:38px;color:#6a7080;display:flex;justify-content:space-between}}
.ink .foot{{color:#aab}}
.big{{font-family:Poppins;font-weight:700;color:var(--acc)}}
.chips span{{display:inline-block;background:#fff;color:var(--ink);border:3px solid var(--ink);border-radius:60px;padding:16px 38px;margin:0 18px 22px 0;font-family:Poppins;font-weight:600;font-size:40px}}
ul.inside{{list-style:none;font-size:56px;line-height:1.3}} ul.inside li{{margin:0 0 52px;padding-left:70px;position:relative}}
ul.inside li:before{{content:"";position:absolute;left:0;top:18px;width:34px;height:34px;background:var(--acc);border-radius:8px}}
ul.inside b{{font-family:Poppins}}
.tilt{{transform:rotate(-1.5deg)}}
</style></head><body>
<div class=v id=v1>
<div class=kick>Excel and Google Sheets</div>
<h1>Rental Property<br>Spreadsheet<br><span class=big>for 1 to 10 doors</span></h1>
<div class=sub style="max-width:1500px">Rent roll with late-rent flags, a P&amp;L for each property, and cash flow after principal and improvements.</div>
<div class="card tilt" style="position:absolute;left:120px;right:120px;top:930px;font-size:31px">{rent_roll_table()}</div>
<div class=foot><span>ProofNotFluff</span><span>Instant digital download</span></div>
</div>
<div class=v id=v2>
<div class=kick>Rent Roll tab</div>
<h2>See who's late the moment<br>you open the file</h2>
<div class=card style="font-size:40px">{rent_roll_table()}</div>
<div class=sub style="margin-top:60px">Paid, Partial, Due or LATE for every door, from your own due day and grace days. Leases ending within 60 days show in orange.</div>
<div class=foot><span>Sample data shown as of October 5, 2026</span><span>ProofNotFluff</span></div>
</div>
<div class=v id=v3>
<div class=kick>Property P&amp;L tab</div>
<h2>Profit vs. what you kept,<br>per property</h2>
<div class=card style="font-size:34px;padding:30px 40px">{pl_table(keep={"Total money in","Mortgage interest","Taxes","Repairs","Total expenses","Profit before depreciation","Mortgage principal","Capital improvement","Cash flow (what you kept)"})}</div>
<div class=sub style="margin-top:60px">Sample: Oak Street House shows $4,844 profit but keeps $550 after mortgage principal and a $1,450 water heater.</div>
<div class=foot><span>Expenses on the IRS Schedule E line names</span><span>Sample data, Jan to Oct 2026</span></div>
</div>
<div class=v id=v4>
<div class=kick>Monthly tab</div>
<h2>Tax months and repair months<br>stop being a surprise</h2>
<div class=card style="font-size:27px;padding:30px 24px">{monthly_table()}</div>
<div class=sub style="margin-top:50px;font-size:44px">Sample year: June shows ($3,128), the month two property tax installments and an insurance renewal land. July drops to $562 when a $1,450 water heater goes in.</div>
<div class=sub style="margin-top:40px">Money in, expenses and cash flow by month for all properties or just one. Deposits stay out because they belong to the tenant.</div>
<div class=foot><span>Sample data shown</span><span>ProofNotFluff</span></div>
</div>
<div class="v ink" id=v5>
<div class=kick>What's inside</div>
<h2 style="color:#fff">One workbook, five-minute setup</h2>
<ul class=inside>
<li><b>Setup:</b> up to 10 properties and 10 doors, typed once</li>
<li><b>Log:</b> one row per payment or bill, with dropdowns</li>
<li><b>Rent Roll:</b> Paid, Partial, Due, LATE and lease-end alerts</li>
<li><b>Property P&amp;L:</b> 14 expense categories named after Schedule E lines</li>
<li><b>Monthly:</b> money in, expenses, cash flow, running total</li>
<li><b>3-page Start Here PDF</b> (Letter and A4): which file to open on a computer or phone</li>
</ul>
<div class=sub style="margin:20px 0 50px;font-size:46px">Files: one .xlsx workbook, the Start Here PDF in Letter and A4, and LICENSE-AND-DISCLAIMER.txt</div>
<div class=chips><span>.xlsx</span><span>Excel</span><span>Google Sheets</span><span>Worked example included</span></div>
<div class=foot><span>ProofNotFluff</span><span>Digital download. No physical item ships.</span></div>
</div>
</body></html>"""
open(os.path.join(HERE, "images.html"), "w").write(IMG)

os.makedirs(os.path.join(HERE, "pkg"), exist_ok=True); os.makedirs(os.path.join(HERE, "img"), exist_ok=True)
with sync_playwright() as p:
    b = p.chromium.launch(); pg = b.new_page()
    pg.goto("file://" + os.path.join(HERE, "guide.html")); pg.wait_for_timeout(500)
    pg.pdf(path=os.path.join(HERE, "pkg", "START-HERE-Guide-Letter.pdf"), width="8.5in", height="11in", print_background=True)
    pg.pdf(path=os.path.join(HERE, "pkg", "START-HERE-Guide-A4.pdf"), format="A4", print_background=True)
    pg2 = b.new_page(viewport={"width": 2000, "height": 2000})
    pg2.goto("file://" + os.path.join(HERE, "images.html")); pg2.wait_for_timeout(800)
    for i in range(1, 6):
        pg2.query_selector(f"#v{i}").screenshot(path=os.path.join(HERE, "img", f"0{i}.png"))
    b.close()
print("done")
