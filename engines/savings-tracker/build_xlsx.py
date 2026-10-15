#!/usr/bin/env python3
"""Savings Tracker Spreadsheet, version 1 (v3 dashboard design, Oct 14, 2026).

A savings goals and sinking funds tracker for Excel and Google Sheets. You type each deposit once:
  1 Goals      up to 20 goals: name, target, an optional target date, anything already saved.
  2 Deposits   one row per money put in (or taken out, as a minus): date, goal from the list, amount.
  3 Progress   saved, left, percent and a bar for every goal. A goal with a target date shows the amount
               needed each month and whether your pace keeps up; a goal with no date shows the month it
               finishes at your current pace (no forced end date).

Built to design/WORKBOOK_STANDARD.md section 0 with tools/pnf_dash.py and tools/wb_style.py.
Only plain functions (SUMIFS, COUNTIFS, COUNTIF, SUMPRODUCT, DATE, YEAR, MONTH, TODAY, TEXT, CEILING, REPT, IF and simple maths):
no macros, no scripts, no charts, no array formulas, so it works in the Google Sheets phone app.
No interest rates, bank figures or outside facts are in the file.

  python3 build_xlsx.py Savings-Tracker-EXAMPLE.xlsx           (the worked example, as of Oct 15, 2026)
  python3 build_xlsx.py Savings-Tracker.xlsx --blank           (no goals, no deposits)
"""
import datetime as dt
import math
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402

PRODUCT = "Savings Tracker Spreadsheet"
VERSION = "1"
CHECKED = "Oct 14, 2026"
AS_OF = f"Checked {CHECKED}"
HERE = os.path.dirname(os.path.abspath(__file__))
args = sys.argv[1:]
BLANK = "--blank" in args
args = [a for a in args if a != "--blank"]
OUT = args[0] if args else "Savings-Tracker-EXAMPLE.xlsx"
MONEY = '"$"#,##0.00;-"$"#,##0.00;"$"0.00'
MONEY_T = '"$"#,##0.00;-"$"#,##0.00;"-"'
PCT = '0%;-0%;0%'
INT = S.FMT["int"]
DATE = S.FMT["date"]
HIDE = ";;;"
BAR = D.BAR_CHAR
N_GOALS = 20
N_ROWS = int(os.environ.get("PNF_SHOT_ROWS", "1000"))  # PNF_SHOT_ROWS only for listing shots of the Deposits tab
ROWS_TEXT = "1,000"  # what the shipped file has; the text never follows PNF_SHOT_ROWS

# ---------------------------------------------------------------- the example: made-up goals, as of Oct 15, 2026
d = dt.datetime
EX_ASOF = d(2026, 10, 15)
GOALS = [  # name, target, target date or None, already saved
    ("Emergency fund", 6000, None, 1500),
    ("Holiday gifts", 800, d(2026, 12, 1), 0),
    ("Car repairs", 1200, None, 200),
    ("Summer trip 2027", 2400, d(2027, 6, 1), 0),
    ("Car insurance (6 months)", 720, d(2027, 1, 1), 0),
    ("New laptop", 1100, None, 0),
]
DEP = []  # (date, goal, amount, note)
for m in range(5, 11):
    DEP.append((d(2026, m, 3), "Emergency fund", 250, "Payday transfer"))
    DEP.append((d(2026, m, 5), "Car repairs", 75, "Payday transfer"))
    DEP.append((d(2026, m, 10), "Car insurance (6 months)", 120, ""))
for m in range(6, 11):
    DEP.append((d(2026, m, 15), "Holiday gifts", 100, ""))
for m in range(7, 11):
    DEP.append((d(2026, m, 1), "Summer trip 2027", 250, ""))
DEP.append((d(2026, 8, 22), "Car repairs", -180, "Brake pads (taken out)"))
DEP.append((d(2026, 9, 12), "New laptop", 150, ""))
DEP.append((d(2026, 10, 12), "New laptop", 150, ""))
DEP.sort(key=lambda t: (t[0], t[1]))


def model(goals, dep, asof):
    """The workbook's maths in Python, used to check the render and to write the Start Here example."""
    out = []
    for name, target, td, already in goals:
        rows = [t for t in dep if t[1] == name and t[0] <= asof]
        dsum = sum(t[2] for t in rows)
        saved = already + dsum
        left = target - saved
        alld = [t[0] for t in dep if t[1] == name]
        first = min(alld) if alld else None
        months = ((asof.year - first.year) * 12 + asof.month - first.month + 1) if first else 0
        pace = dsum / months if months > 0 else 0
        ml = ((td.year - asof.year) * 12 + td.month - asof.month) if td else None
        needed = (left / max(1, ml) if (ml is not None and ml >= 0) else 0) if (td and left > 0) else 0
        if left <= 0:
            status = "Reached" if left == 0 else f"Reached, ${-left:,.0f} over"
        elif td:
            if ml < 0:
                status = f"Date passed, ${left:,.0f} to go"
            elif pace >= needed - 0.005:
                status = f"On track: ${math.ceil(needed):,} a month"
            else:
                status = f"Behind: needs ${math.ceil(needed):,} a month"
        elif pace <= 0:
            status = "Add a deposit to see a finish date"
        else:
            n = math.ceil(left / pace)
            mm = asof.month + n
            fin = dt.date(asof.year + (mm - 1) // 12, (mm - 1) % 12 + 1, 1)
            status = f"At your pace: done {fin:%b %Y}"
        out.append(dict(name=name, target=target, saved=saved, left=max(0, left), pct=saved / target, pace=pace,
                        needed=needed, status=status))
    month_in = sum(t[2] for t in dep if t[0].year == asof.year and t[0].month == asof.month and t[0] <= asof and t[2] > 0)
    month_out = -sum(t[2] for t in dep if t[0].year == asof.year and t[0].month == asof.month and t[0] <= asof and t[2] < 0)
    tot = dict(target=sum(g["target"] for g in out), saved=sum(g["saved"] for g in out),
               left=sum(g["left"] for g in out), month_in=month_in, month_out=month_out,
               need=sum(math.ceil(g["needed"]) if g["needed"] else 0 for g in out),
               behind=sum(1 for g in out if g["status"].startswith(("Behind", "Date passed"))))
    tot["pct"] = tot["saved"] / tot["target"]
    return out, tot


EXG, EXT = model(GOALS, DEP, EX_ASOF)
# the Start Here example and the listing quote these; fail the build if the data ever drifts
assert (EXT["target"], EXT["saved"], EXT["left"], EXT["month_in"], EXT["need"], EXT["behind"]) == \
    (12220, 5990, 6230, 945, 325, 1), EXT
assert [g["status"] for g in EXG] == ["At your pace: done Oct 2027", "Behind: needs $150 a month",
                                      "At your pace: done Mar 2028", "On track: $175 a month", "Reached",
                                      "At your pace: done Apr 2027"], [g["status"] for g in EXG]
N_DEP = len(DEP)
if BLANK:
    GOALS, DEP = [], []

GO_N, DE_N, PR_N = "1 Goals", "2 Deposits", "3 Progress"
GO, DE, PR = (f"'{n}'" for n in (GO_N, DE_N, PR_N))


def rich(head, body, size=9):
    return CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=size, b=True, color=D.T["ink"]), head + ". "),
                        TextBlock(InlineFont(rFont=D.FONT, sz=size, color=D.T["ink2"]), body))


def sign_rule(ws, ref):
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'AND(ISNUMBER({ref}),{ref}<0)'], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'AND(ISNUMBER({ref}),{ref}>0)'], font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))


def notes_card(ws, top, left, right, c0, c1, notes, title="Notes"):
    r = top
    D.h(ws, r, 12); r += 1
    t = ws[f"{c0}{r}"]; t.value = title; t.font = D.f(12, True); ws.merge_cells(f"{c0}{r}:{c1}{r}"); D.h(ws, r, 24); r += 1
    width = sum(S.width_of(ws, k) for k in D.cols(c0, c1))
    for head, body in notes:
        cell = ws[f"{c0}{r}"]; cell.value = rich(head, body)
        cell.alignment = Alignment(vertical="top", wrap_text=True); ws.merge_cells(f"{c0}{r}:{c1}{r}")
        D.h(ws, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
    D.h(ws, r, 12)
    D.frame(ws, top, r, left, right)
    return r


def fill_heights(ws, last):
    for rr in range(1, last + 1):
        if ws.row_dimensions[rr].height is None:
            D.h(ws, rr, D.GRID_ROW)


def head_row(ws, row, spec, span):
    for k, txt, al in spec:
        c = ws[f"{k}{row}"]; c.value = txt; c.font = D.f(8, True, "muted")
        c.alignment = Alignment(horizontal=al, vertical="bottom", wrap_text=True)
    for k in D.cols(*span):
        ws[f"{k}{row}"].border = Border(bottom=D.side("ink"))
    D.h(ws, row, 24)


wb = S.new_workbook(PRODUCT, version=VERSION)

# fixed addresses on the 3 Progress tab, asserted when that tab is built
ASOF_ROW = 14
ASOF = f"{PR}!$F${ASOF_ROW}"          # hidden helper: the as-of input, or today when it is blank

# =====================================================================  1 Goals
go = wb.create_sheet(GO_N)
S.set_widths(go, {"A": 3, "B": 2, "C": 30, "D": 14, "E": 15, "F": 14, "G": 2, "H": 14, "I": 14, "J": 2, "K": 3})
D.page_header(go, "Your goals", "One line per goal or sinking fund. A target date is optional.", AS_OF,
              span=("B", "J"), side_cols=3)
G_TOP = 11
gv = lambda i, k: GOALS[i][k] if i < len(GOALS) else None  # noqa: E731
cols_g = [
    dict(col="C", head="Goal", kind="input", values=[g[0] for g in GOALS], validation=("textLength", 0, 40),
         prompt="A short name, up to 40 characters, for example Emergency fund. Each name once. Clear it to drop the goal.",
         error="Up to 40 characters."),
    dict(col="D", head="Target", kind="input", fmt=MONEY, align="right", values=[g[1] for g in GOALS],
         validation=("decimal", 0, 100000000), error="Type a target from 0 to 100,000,000, as a positive number.",
         prompt="How much this goal needs in total, for example 6000."),
    dict(col="E", head="Target date (optional)", kind="input", fmt=DATE, align="right", values=[g[2] for g in GOALS],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="Optional. The date you want it saved by, for example 12/1/2026. Leave blank for an open goal: you get a finish month at your pace instead.",
         error="Type a real date between 2000 and 2100, or leave the cell blank."),
    dict(col="F", head="Already saved", kind="input", fmt=MONEY, align="right", values=[g[3] or None for g in GOALS],
         validation=("decimal", 0, 100000000), error="Type an amount from 0 to 100,000,000, or leave it blank.",
         prompt="Optional. What was already in this goal before your first deposit row, for example 1500."),
    dict(col="H", head="Saved now", kind="calc", fmt=MONEY_T, align="right",
         formula=lambda rr: (f'=IF(C{rr}="","",N(F{rr})+SUMIFS({DE}!$E:$E,{DE}!$D:$D,C{rr},{DE}!$C:$C,"<="&{ASOF}))')),
    dict(col="I", head="Left", kind="calc", fmt=MONEY_T, align="right",
         formula=lambda rr: f'=IF(C{rr}="","",MAX(0,N(D{rr})-H{rr}))'),
]
GF, GL, G_END = D.table_card(go, G_TOP, "B", "J", cols_g, N_GOALS, title="Your goals and sinking funds",
                             sub="Yellow cells are yours. Saved now adds what you had already to every deposit on the 2 Deposits tab.")
for rr in range(GF, GL + 1):
    go[f"C{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    go[f"H{rr}"].font = D.f(10, True)
NAMES = f"{GO}!$C${GF}:$C${GL}"
r_dup_txt = f'SUMPRODUCT(({GO}!$C${GF}:$C${GL}<>"")*(COUNTIF({GO}!$C${GF}:$C${GL},{GO}!$C${GF}:$C${GL})>1))'
D.tile(go, 6, "B", "F", "STILL TO SAVE, ALL GOALS", f"=SUM($I${GF}:$I${GL})", MONEY,
       f'=IF(COUNTIF($C${GF}:$C${GL},"?*")=0,"Type your first goal below to start",IF({r_dup_txt}>0,"Two goals share a name. Give each its own so deposits count once.",'
       f'"Across "&COUNTIF($C${GF}:$C${GL},"?*")&IF(COUNTIF($C${GF}:$C${GL},"?*")=1," goal"," goals")&". See the 3 Progress tab for each one."))', dark=True)
v_g, s_g, _ = D.tile(go, 6, "G", "J", "SAVED SO FAR", f"=SUM($H${GF}:$H${GL})", MONEY,
                     f'=IF(SUM($D${GF}:$D${GL})=0,"Add a target to each goal",ROUND(SUM($H${GF}:$H${GL})/SUM($D${GF}:$D${GL})*100,0)&"% of $"&TEXT(SUM($D${GF}:$D${GL}),"#,##0")&" in targets")')
D.h(go, 10, 14)
GN_TOP = G_END + 2
gn_end = notes_card(go, GN_TOP, "B", "J", "C", "I", [
    ("Goals and sinking funds", "A goal is money you are building toward once (an emergency fund, a laptop). A sinking fund is "
     "money you set aside for a bill you know is coming (car insurance, holiday gifts). Both go on this list the same way."),
    ("Target date", "Give a goal a date and the 3 Progress tab shows how much it needs each month to get there. Leave the date "
     "blank and it shows the month you finish at your current pace instead. You can add or clear a date at any time."),
    ("Already saved", "Use it for money that was in the goal before you started the deposit log, so your first month is not "
     "counted as one giant deposit."),
    ("Renaming", "Rename a goal here and the dropdown on the 2 Deposits tab follows. Rows already typed with the old name show "
     "Not on the list until you pick the new name."),
])
GO_LAST = gn_end + 1
D.h(go, GO_LAST, 14)
fill_heights(go, GO_LAST)
D.paint_canvas(go, GO_LAST, "Z")
S.finish_sheet(go, PRODUCT, GO_LAST, span=("A", "K"), freeze="A11", tab_color="gold")
S.page_break_before(go, GN_TOP)

# =====================================================================  2 Deposits
de = wb.create_sheet(DE_N)
S.set_widths(de, {"A": 3, "B": 2, "C": 13, "D": 28, "E": 14, "F": 34, "G": 18, "H": 2, "I": 3})
D.page_header(de, "Deposits", "One row each time money goes in or out of a goal.",
              "Example rows: type over them" if not BLANK else "Type your first row in row 16", span=("B", "H"), side_cols=3)
T_TOP = 11
T_FIRST = T_TOP + 5
T_LAST = T_FIRST + N_ROWS - 1


def counts_as(r):
    return (f'=IF(E{r}="","",IF(D{r}="","Needs a goal",IF(C{r}="","Needs a date",'
            f'IF(COUNTIF({NAMES},D{r})=0,"Not on the list",IF(E{r}<0,"Taken out","Put in")))))')


cols_t = [
    dict(col="C", head="Date", kind="input", fmt=DATE, values=[t[0] for t in DEP],
         validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The day the money moved, for example 10/3/2026.",
         error="Type a real date between 2000 and 2100, for example 10/3/2026."),
    dict(col="D", head="Goal (pick)", kind="input", values=[t[1] for t in DEP], validation=("list_range", NAMES),
         prompt="Choose a goal from the list. The list comes from the 1 Goals tab; add goals there.",
         error="Choose a goal from the list. To add one, type it on the 1 Goals tab first."),
    dict(col="E", head="Amount", kind="input", fmt=MONEY, align="right", values=[t[2] for t in DEP],
         validation=("decimal", -100000000, 100000000),
         prompt="Money put in, for example 250. Money taken out of the goal is a minus amount, for example -180.",
         error="Type an amount from -100,000,000 to 100,000,000. Use a minus for money taken out."),
    dict(col="F", head="Note", kind="input", values=[t[3] or None for t in DEP], validation=("textLength", 0, 80),
         prompt="Optional: where it came from or why it was taken out.", error="Up to 80 characters."),
    dict(col="G", head="Counts as", kind="calc", formula=counts_as),
]
f1, f2, t_end = D.table_card(de, T_TOP, "B", "H", cols_t, N_ROWS, title="Your deposits",
                             sub=f"Yellow columns are yours. Counts as fills in on its own. {ROWS_TEXT} rows, in any order.")
assert (f1, f2) == (T_FIRST, T_LAST), (f1, f2)
for rr in range(T_FIRST, T_LAST + 1):
    de[f"G{rr}"].font = D.f(9, False, "ink2"); de[f"G{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for k in ("D", "F"):
        de[f"{k}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    # hidden helper: this row's date when it is the goal's first deposit (earliest date, first row on a tie)
    tie = f'+COUNTIFS($D${T_FIRST}:D{rr - 1},D{rr},$C${T_FIRST}:C{rr - 1},C{rr})' if rr > T_FIRST else ""
    hc = de[f"H{rr}"]
    hc.value = (f'=IF(OR(G{rr}="Put in",G{rr}="Taken out"),IF(COUNTIFS($D${T_FIRST}:$D${T_LAST},D{rr},$C${T_FIRST}:$C${T_LAST},"<"&C{rr})'
                f'{tie}=0,C{rr},""),"")')
    hc.number_format = HIDE
de.conditional_formatting.add(f"C{T_FIRST}:D{T_LAST}", FormulaRule(
    formula=[f'AND($E{T_FIRST}<>"",OR($C{T_FIRST}="",$D{T_FIRST}=""))'], fill=PatternFill("solid", bgColor=D.T["accent_tint"])))
for word, color in (("Needs", "accent"), ("Not on", "accent"), ("Taken", "gold")):
    de.conditional_formatting.add(f"G{T_FIRST}:G{T_LAST}", FormulaRule(
        formula=[f'LEFT($G{T_FIRST},{len(word)})="{word}"'], font=Font(color=D.T[color], bold=True), stopIfTrue=True))
TA = f"{DE}!$E${T_FIRST}:$E${T_LAST}"
TG = f"{DE}!$D${T_FIRST}:$D${T_LAST}"
TD = f"{DE}!$C${T_FIRST}:$C${T_LAST}"
TS = f"{DE}!$G${T_FIRST}:$G${T_LAST}"
TH = f"{DE}!$H${T_FIRST}:$H${T_LAST}"
FIX = f'(COUNTIF({TS},"Needs*")+COUNTIF({TS},"Not on*"))'
D.h(de, 6, 10)
D.stat_strip(de, 7, [
    ("C:D", "Put in, all time", f'=SUMIFS({TA},{TS},"Put in")', MONEY),
    ("E:F", "Taken out, all time", f'=-SUMIFS({TA},{TS},"Taken out")', MONEY),
    ("G:G", "Rows to fix", f"={FIX}", INT),
])
D.h(de, 9, 10)
D.frame(de, 6, 9, "B", "H")
D.status_rule(de, "G8", "$G$8>0", fill_tint=False)
D.h(de, 10, 14)
D.paint_canvas(de, t_end + 1, "Z")
S.finish_sheet(de, PRODUCT, t_end + 1, span=("A", "I"), freeze=f"A{T_FIRST}", tab_color="teal")
de.print_title_rows = f"{T_FIRST - 1}:{T_FIRST - 1}"
de.print_area = f"A1:I{T_FIRST + 299}"  # print the first 300 rows, not 1,000 bare ones

# =====================================================================  3 Progress
pr = wb.create_sheet(PR_N)
S.set_widths(pr, {"A": 3, "B": 2, "C": 26, "D": 13, "E": 13, "F": 13, "G": 9, "H": 19, "I": 33, "J": 2, "K": 3})
D.page_header(pr, "Your progress", "Every goal: saved, left, and whether your pace gets you there.", AS_OF,
              span=("B", "J"), side_cols=2)
K = D.Card(pr, 11, "B", "C", "E", "F", wb=wb)
K.title("Your totals", "All goals on the 1 Goals tab.")
r_asof = K.input("As of date", None if BLANK else EX_ASOF, DATE, name="AsOfDate", validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
                 allow_blank=True, prompt_title="As of date",
                 prompt="Leave blank to use today. Type a date, for example 10/15/2026, to see your goals as they stood that day.",
                 error="Type a real date between 2000 and 2100, or leave the cell blank for today.")
assert r_asof == ASOF_ROW, r_asof
hz = pr[f"F{r_asof}"]; hz.value = f'=IF($E${r_asof}="",TODAY(),$E${r_asof})'; hz.number_format = HIDE
A = f"$F${r_asof}"
MSTART = f"DATE(YEAR({A}),MONTH({A}),1)"
r_n = K.calc("Goals on the list", f'=COUNTIF({NAMES},"?*")', INT)
r_t = K.calc("Total of targets", f"=SUM({GO}!$D${GF}:$D${GL})", MONEY)
r_s = K.calc("Saved so far", f"=SUM({GO}!$H${GF}:$H${GL})", MONEY)
r_l = K.calc("Still to save", f"=SUM({GO}!$I${GF}:$I${GL})", MONEY, total=True)
r_p = K.calc("Saved of all targets", f'=IF($E${r_t}=0,"",$E${r_s}/$E${r_t})', PCT, bold=True, tint="total_tint")
K.close()

G = D.Card(pr, 11, "G", "H", "I", "K", extra="J", wb=wb)
G.title("This month", "Deposits dated this month, up to the as of date.")
r_mo = G.calc("Month", f'=TEXT({A},"mmmm yyyy")', "@")
r_in = G.calc("Put in", f'=SUMIFS({TA},{TS},"Put in",{TD},">="&{MSTART},{TD},"<="&{A})', MONEY)
r_out = G.calc("Taken out", f'=-SUMIFS({TA},{TS},"Taken out",{TD},">="&{MSTART},{TD},"<="&{A})', MONEY)
r_net = G.calc("Net this month", f"=I{r_in}-I{r_out}", MONEY, total=True)
r_need = G.calc("Dated goals need", "=0", MONEY)      # filled after the table
r_beh = G.calc("Goals behind", "=0", INT, bold=True)  # filled after the table
G.close()
for rr in (r_mo, r_in, r_out, r_net, r_need, r_beh):
    pr[f"I{rr}"].alignment = Alignment(horizontal="right", vertical="center")
assert K.end == G.end, (K.end, G.end)

# ---- goal by goal
L_TOP = K.end + 2
r = L_TOP
D.h(pr, r, 12); r += 1
t = pr[f"C{r}"]; t.value = "Goal by goal"; t.font = D.f(12, True); pr.merge_cells(f"C{r}:I{r}"); D.h(pr, r, 24); r += 1
t = pr[f"C{r}"]; t.value = "Goals with a target date show the monthly amount they need. Open goals show the month they finish at your pace."
t.font = D.f(9, False, "muted"); pr.merge_cells(f"C{r}:I{r}"); D.h(pr, r, 18); r += 1
D.h(pr, r, 6); r += 1
head_row(pr, r, [("C", "GOAL", "left"), ("D", "TARGET", "right"), ("E", "SAVED", "right"), ("F", "LEFT", "right"),
                 ("G", "DONE", "right"), ("H", "", "left"), ("I", "PLAN AND STATUS", "left")], ("C", "I"))
P_HEAD = r
r += 1
L_FIRST = r
# hidden helpers per row, outside the print area: M pace a month, N months to the target month, O needed a month,
# P first deposit date, Q months since the first deposit, R needed a month rounded up to whole dollars
for i in range(N_GOALS):
    g = GF + i
    nm = f"{GO}!$C${g}"
    vals = {
        "C": f'=IF({nm}="","",{nm})',
        "D": f'=IF(C{r}="","",N({GO}!$D${g}))',
        "E": f'=IF(C{r}="","",{GO}!$H${g})',
        "F": f'=IF(C{r}="","",{GO}!$I${g})',
        "G": f'=IF(OR(C{r}="",N(D{r})=0),"",MIN(1,MAX(0,E{r}/D{r})))',
        "H": f'=IF(OR(G{r}="",G{r}=0),"",REPT("{BAR}",MAX(1,ROUND(G{r}*13,0))))',
        "P": f'=IF(C{r}="",0,SUMIFS({TH},{TG},C{r}))',
        "Q": f'=IF(N(P{r})=0,0,(YEAR({A})-YEAR(P{r}))*12+MONTH({A})-MONTH(P{r})+1)',
        "M": f'=IF(N(Q{r})<=0,0,(E{r}-N({GO}!$F${g}))/Q{r})',
        "N": f'=IF(OR(C{r}="",{GO}!$E${g}=""),"",(YEAR({GO}!$E${g})-YEAR({A}))*12+MONTH({GO}!$E${g})-MONTH({A}))',
        "O": f'=IF(OR(N{r}="",N(F{r})<=0),0,IF(N{r}<0,0,F{r}/MAX(1,N{r})))',
        "R": f'=CEILING(O{r},1)',
    }
    left_signed = f"(N(D{r})-E{r})"
    vals["I"] = (f'=IF(C{r}="","",IF(N(D{r})<=0,"Add a target on the 1 Goals tab",IF({left_signed}<=0,'
                 f'IF({left_signed}<0,"Reached, "&TEXT(-{left_signed},"$#,##0")&" over","Reached"),'
                 f'IF(N{r}<>"",IF(N{r}<0,"Date passed, "&TEXT(F{r},"$#,##0")&" to go",'
                 f'IF(M{r}>=O{r}-0.005,"On track: "&TEXT(CEILING(O{r},1),"$#,##0")&" a month",'
                 f'"Behind: needs "&TEXT(CEILING(O{r},1),"$#,##0")&" a month")),'
                 f'IF(M{r}<=0,IF(N(Q{r})>0,"More taken out than put in so far","Add a deposit to see a finish date"),IF(F{r}/M{r}>600,"Over 50 years at this pace",'
                 f'"At your pace: done "&TEXT(DATE(YEAR({A}),MONTH({A})+CEILING(F{r}/M{r},1),1),"mmm yyyy")))))))')
    for k, v in vals.items():
        c = pr[f"{k}{r}"]; c.value = v
        if k in "MNOPQR":
            c.number_format = HIDE
            continue
        if k in "DEF":
            c.number_format = MONEY_T; c.alignment = Alignment(horizontal="right", vertical="center")
            c.font = D.f(10, k == "E", "ink2" if k == "D" else "ink")
        elif k == "G":
            c.number_format = PCT; c.alignment = Alignment(horizontal="right", vertical="center"); c.font = D.f(10, True)
        elif k == "H":
            c.font = D.f(9, False, "teal"); c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        elif k == "I":
            c.font = D.f(9, False, "ink2"); c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        else:
            c.font = D.f(10); c.alignment = Alignment(horizontal="left", vertical="center")
    for k in D.cols("C", "I"):
        pr[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    pr.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f'OR(LEFT($I{r},6)="Behind",LEFT($I{r},4)="Date")'],
                                                      font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
    pr.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f'OR(LEFT($I{r},8)="On track",LEFT($I{r},7)="Reached")'],
                                                      font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))
    pr.conditional_formatting.add(f"G{r}", FormulaRule(formula=[f'AND(ISNUMBER($G{r}),$G{r}>=1)'],
                                                      font=Font(color=D.T["teal"], bold=True), stopIfTrue=True))
    D.h(pr, r, D.GRID_ROW)
    r += 1
L_LAST = r - 1
# totals row
c = pr[f"C{r}"]; c.value = "All goals"; c.font = D.f(10, True); c.alignment = Alignment(vertical="center")
for k, v, fm in (("D", f"=SUM(D{L_FIRST}:D{L_LAST})", MONEY_T), ("E", f"=SUM(E{L_FIRST}:E{L_LAST})", MONEY_T),
                 ("F", f"=SUM(F{L_FIRST}:F{L_LAST})", MONEY_T), ("G", f'=IF(N(D{r})=0,"",MIN(1,E{r}/D{r}))', PCT)):
    cc = pr[f"{k}{r}"]; cc.value = v; cc.number_format = fm; cc.font = D.f(10, True)
    cc.alignment = Alignment(horizontal="right", vertical="center")
cc = pr[f"H{r}"]; cc.value = f'=IF(OR(G{r}="",G{r}=0),"",REPT("{BAR}",MAX(1,ROUND(G{r}*13,0))))'
cc.font = D.f(9, False, "ink"); cc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
for k in D.cols("C", "I"):
    cell = pr[f"{k}{r}"]; cell.border = Border(top=D.side("ink"), bottom=D.side("hair")); cell.fill = D.fill("total_tint")
D.h(pr, r, D.GRID_ROW)
L_TOT = r
r += 1
D.h(pr, r, 12)
L_END = r
D.frame(pr, L_TOP, L_END, "B", "K")
S.page_break_before(pr, L_TOP)

pr[f"I{r_need}"].value = f'=SUM($R${L_FIRST}:$R${L_LAST})'
pr[f"I{r_beh}"].value = f'=COUNTIF($I${L_FIRST}:$I${L_LAST},"Behind*")+COUNTIF($I${L_FIRST}:$I${L_LAST},"Date passed*")'
D.status_rule(pr, f"I{r_beh}", f"$I${r_beh}>0", fill_tint=False)

# ---- the last six months
S6_TOP = L_END + 2
r = S6_TOP
D.h(pr, r, 12); r += 1
t = pr[f"C{r}"]; t.value = "The last six months"; t.font = D.f(12, True); pr.merge_cells(f"C{r}:I{r}"); D.h(pr, r, 24); r += 1
t = pr[f"C{r}"]; t.value = "Everything put in and taken out across all goals, month by month, ending with the as of month."
t.font = D.f(9, False, "muted"); pr.merge_cells(f"C{r}:I{r}"); D.h(pr, r, 18); r += 1
D.h(pr, r, 6); r += 1
head_row(pr, r, [("C", "MONTH", "left"), ("D", "PUT IN", "right"), ("E", "TAKEN OUT", "right"), ("F", "NET", "right"),
                 ("H", "", "left")], ("C", "I"))
r += 1
S6_FIRST = r
SCALE6 = f"MAX(MAX($F${S6_FIRST}:$F${S6_FIRST + 5}),1)"
for k6 in range(6):
    back = 5 - k6
    ms = f"DATE(YEAR({A}),MONTH({A})-{back},1)"
    me = f"DATE(YEAR({A}),MONTH({A})-{back}+1,1)"
    vals = {
        "C": f'=TEXT({ms},"mmmm yyyy")',
        "D": f'=SUMIFS({TA},{TS},"Put in",{TD},">="&{ms},{TD},"<"&{me},{TD},"<="&{A})',
        "E": f'=-SUMIFS({TA},{TS},"Taken out",{TD},">="&{ms},{TD},"<"&{me},{TD},"<="&{A})',
        "F": f"=D{r}-E{r}",
        "H": f'=IF(F{r}<=0,"",REPT("{BAR}",MAX(1,ROUND(F{r}/{SCALE6}*22,0))))',
    }
    for k, v in vals.items():
        cc = pr[f"{k}{r}"]; cc.value = v
        if k in "DEF":
            cc.number_format = MONEY_T; cc.alignment = Alignment(horizontal="right", vertical="center")
            cc.font = D.f(10, k == "F", "ink2" if k != "F" else "ink")
        elif k == "H":
            cc.font = D.f(9, False, "blue"); cc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        else:
            cc.font = D.f(10); cc.alignment = Alignment(horizontal="left", vertical="center")
    pr.merge_cells(f"H{r}:I{r}")
    sign_rule(pr, f"F{r}")
    for k in D.cols("C", "I"):
        pr[f"{k}{r}"].border = Border(bottom=D.side("hair"))
    D.h(pr, r, D.GRID_ROW)
    r += 1
D.h(pr, r, 12)
S6_END = r
D.frame(pr, S6_TOP, S6_END, "B", "K")

# ---- tiles (rows 6 to 9)
D.tile(pr, 6, "B", "F", "SAVED SO FAR, ALL GOALS", f"=$E${r_s}", MONEY,
       f'=IF($E${r_n}=0,"Add your goals on the 1 Goals tab to start",IF($E${r_t}=0,"Add a target to each goal on the 1 Goals tab",'
       f'ROUND($E${r_s}/$E${r_t}*100,0)&"% of $"&TEXT($E${r_t},"#,##0")&" across "&$E${r_n}&IF($E${r_n}=1," goal"," goals")))', dark=True)
v_m, s_m, _ = D.tile(pr, 6, "G", "K", f'="PUT IN THIS MONTH, NET ("&UPPER(TEXT({A},"mmm yyyy"))&")"', f"=$I${r_net}", MONEY,
                     f'=IF($E${r_n}=0,"No goals yet",IF($I${r_beh}>0,$I${r_beh}&IF($I${r_beh}=1," goal is behind."," goals are behind.")&" Dated goals need $"&TEXT($I${r_need},"#,##0")&" a month.",'
                     f'IF($I${r_need}>0,"Dated goals need $"&TEXT($I${r_need},"#,##0")&" a month. All on track.","Every goal on track or open-ended")))')
sign_rule(pr, v_m.coordinate)
pr.conditional_formatting.add(s_m.coordinate, FormulaRule(formula=[f"$I${r_beh}>0"], font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
D.h(pr, 10, 14)

PN_TOP = S6_END + 2
pn_end = notes_card(pr, PN_TOP, "B", "K", "C", "I", [
    ("Saved", "What you had already (1 Goals tab) plus every deposit for that goal on the 2 Deposits tab dated on or before the "
     "as of date. Money taken out is a minus amount, so it comes off."),
    ("Your pace", "Deposits for the goal since its first deposit, divided by the months from that first month to the as of "
     "month, counting both. Saving $250 every month gives a pace of $250 a month."),
    ("Goals with a target date", "Left divided by the months from the as of month to the target month (at least one). On track "
     "means your pace is at least that amount; Behind means it is less. A date that has gone by shows what is still to go."),
    ("Open goals", "No target date: the finish month is the as of month plus the months your pace needs to cover what is left, "
     "rounded up. Add a deposit and the month moves."),
    ("As of date", "Blank shows today and moves on its own. Type a date to see your goals as they stood that day."),
    ("Before you rely on a number", S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab")
     + " " + S.MONEY_TIER_CLAUSE),
], title="How the numbers work")
PR_LAST = pn_end + 1
D.h(pr, PR_LAST, 14)
fill_heights(pr, PR_LAST)
D.paint_canvas(pr, PR_LAST, "Z")
S.finish_sheet(pr, PRODUCT, PR_LAST, span=("A", "K"), freeze="A11", tab_color="accent")
S.page_break_before(pr, S6_TOP)

# =====================================================================  Start Here and Terms
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            mm = re.match(r"^([A-Z][A-Za-z ,'0-9]{1,40})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if mm:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), mm.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), mm.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


def money(x, cents=False):
    return ("-" if x < 0 else "") + (f"${abs(x):,.2f}" if cents else f"${abs(x):,.0f}")


sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, PRODUCT, "Start here. Type each deposit once; the 3 Progress tab does the rest.", AS_OF,
              span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "A savings goals and sinking funds tracker. List up to 20 goals once on the 1 Goals tab, each with a target and, if you "
    "want one, a target date. Type each deposit as one row on the 2 Deposits tab: date, goal, amount. The 3 Progress tab shows "
    "saved, left, percent done and a bar for every goal. A goal with a date shows the amount it needs each month and whether "
    "your pace keeps up. A goal with no date shows the month you finish at your current pace.", "para")])
r = sh_card(sh, r, "Before you rely on a number", [
    (None, S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab") + " " + S.MONEY_TIER_CLAUSE, "para"),
])
r = sh_card(sh, r, "Three steps", [
    (1, "List your goals. On the 1 Goals tab, type a name and a target for each goal. Add a target date only if the goal has one.", "num"),
    (2, "Log each deposit. On the 2 Deposits tab, one row each time money goes in: date, goal from the list, amount. Money taken out is a minus amount.", "num"),
    (3, "Read your progress. The 3 Progress tab shows every goal with its plan and status. Leave the as of date blank to see today.", "num"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "How to read the cells", [
    ((250, MONEY), "Yellow cells with blue numbers are yours to type in. Each one shows a hint when you select it.", "chip-in"),
    ((EXG[0]["saved"], MONEY), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((EXT["saved"], MONEY),
     "Dark tiles are your answers. The main one sits at the top of the 3 Progress tab and stays on screen as you scroll.", "chip-key"),
    (None, "Nothing is locked for good: each tab is protected without a password, so you can still change anything. To edit a "
           "formula cell, use Review, Unprotect Sheet in Excel, or Data, Protect sheets and ranges in Google Sheets.", "note"),
])
g0, g1, g3 = EXG[0], EXG[1], EXG[3]
ex_text = (f"Savings-Tracker-EXAMPLE.xlsx holds six made-up goals and {N_DEP} deposits from May to October 2026, as of "
           f"Oct 15, 2026. {money(EXT['saved'])} is saved of {money(EXT['target'])} in targets. The emergency fund has no "
           f"date: {money(g0['saved'])} saved at {money(g0['pace'])} a month shows At your pace: done Oct 2027. Holiday gifts "
           f"needs {money(g1['left'])} more by December at {money(g1['needed'])} a month, but its pace is {money(g1['pace'])}, "
           f"so it shows Behind. The summer trip needs {money(g3['needed'])} a month and saves $250, so it is On track. A $180 "
           f"brake repair is a minus amount under Car repairs.")
r = sh_card(sh, r, "The example", [(None, ex_text, "para")])
r = sh_card(sh, r, "Which file to open", [
    (None, "Excel on a computer. Double-click Savings-Tracker.xlsx. If Excel shows a yellow Protected View bar, click Enable Editing.", "para"),
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. No macros, no add-ons, no sign-up.", "para"),
    (None, "Mac, iPad or phone. Unzip on a computer first (phones often can't open a .zip), then upload the .xlsx to Google Drive "
           "and open it in the Google Sheets app, or open it in Excel for Mac or iPad.", "para"),
    (None, "See it filled in first. Open Savings-Tracker-EXAMPLE.xlsx. Savings-Tracker.xlsx is the same workbook with no goals "
           "and no deposits, so there is no sample data to clear.", "para"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "Good to know", [
    (None, "Goals without an end date. Leave the target date blank and the goal never shows as late; you see the month it "
           "finishes at your pace instead.", "para"),
    (None, "On a phone. Every formula is a plain sum, count or date with no scripts, charts or add-ons, so balances and "
           "percents work in the Google Sheets app too.", "para"),
    (None, f"Room for years. The 2 Deposits tab has {ROWS_TEXT} rows and every total reads the whole column. Rows can go in "
           "any order; the date decides the month. Printing it prints the first 300 rows; to print more, select the rows "
           "and print the selection.", "para"),
    (None, "Moving money between goals. Type a minus row for the goal it leaves and a plus row for the goal it joins, same date.", "para"),
    (None, "Interest. If your savings account pays interest, log it as a deposit to the goal it belongs to on the day it is paid.", "para"),
])
SH_LAST = r - 1
D.paint_canvas(sh, SH_LAST, "Z")
S.finish_sheet(sh, PRODUCT, SH_LAST, span=("A", "G"), tab_color="teal",
               footer_notice="Estimates only. Not professional advice. See the Terms tab or LICENSE-AND-DISCLAIMER.txt.")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.", AS_OF,
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More tools from the shop", [
    (None, ("Budget Tracker Spreadsheet: log each purchase and see what is left to spend", "https://www.etsy.com/listing/4594768127"), "link"),
    (None, ("Budget Planner Spreadsheet: plan each month by pay schedule and bill due dates", "https://www.etsy.com/listing/4592917337"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this tracker helps you reach a savings goal, a short review on Etsy helps other people find it. Thank you.", "para"),
    (None, "Message me on Etsy if anything doesn't open or looks wrong, and it will be fixed.", "para"),
])
TM_LAST = r - 1
D.paint_canvas(tm, TM_LAST, "Z")
S.finish_sheet(tm, PRODUCT, TM_LAST, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb)]
if problems:
    print("AUDIT:", *problems[:40], sep="\n  ")
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
print({"goals": (GF, GL), "dep": (T_FIRST, T_LAST), "progress": {"asof": r_asof, "n": r_n, "t": r_t, "s": r_s, "l": r_l,
       "p": r_p, "mo": r_mo, "in": r_in, "out": r_out, "net": r_net, "need": r_need, "beh": r_beh, "rows": (L_FIRST, L_LAST),
       "tot": L_TOT, "six": S6_FIRST}})
print("example:", EXT)
for g in EXG:
    print("  ", g["name"], g["saved"], g["left"], round(g["pace"], 2), round(g["needed"], 2), g["status"])
