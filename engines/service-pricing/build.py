from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

INK = "1D2433"; ACC = "C8502F"; PAPER = "F5F2EC"; TEAL = "2E6B66"
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
HEAD_FILL = PatternFill("solid", fgColor=INK)
OUT_FILL = PatternFill("solid", fgColor="E7F0EE")
KEY_FILL = PatternFill("solid", fgColor="FBE3DA")
PAPER_FILL = PatternFill("solid", fgColor=PAPER)
F = "Arial"
thin = Side(style="thin", color="D5D0C6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = Font(name=F, size=18, bold=True, color=INK)
    ws.row_dimensions[1].height = 30
    if sub:
        ws["A2"] = sub; ws["A2"].font = font(10, color="5A5F6B", italic=True)

def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    ws.row_dimensions[row].height = 30

def inp(c, val, fmt=None):
    c.value = val; c.fill = INPUT_FILL; c.font = Font(name=F, size=10, color="1F3FBF"); c.border = BOX
    if fmt: c.number_format = fmt

def out(c, formula, fmt=None, bold=False, fill=OUT_FILL):
    c.value = formula; c.fill = fill; c.font = font(10, bold=bold); c.border = BOX
    if fmt: c.number_format = fmt

def label(c, text, bold=False, size=10):
    c.value = text; c.font = font(size, bold=bold); c.alignment = Alignment(wrap_text=True, vertical="center")

USD = '$#,##0.00'; USD0 = '$#,##0'; PCT = '0.0%'; NUM = '#,##0.0'

wb = Workbook()

# ---------------- Start Here ----------------
ws = wb.active; ws.title = "Start Here"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 100
ws["B2"] = "Service Pricing Calculator"; ws["B2"].font = Font(name=F, size=22, bold=True, color=INK)
ws["B3"] = "Find your true hourly rate, quote every job with confidence, and see which jobs actually make money."
ws["B3"].font = font(11, color="5A5F6B")
lines = [
    ("", None),
    ("Set up in about 10 minutes", "h"),
    ("1. Open the tab 1 Your Numbers. Fill in the yellow cells: the pay you want, the hours you really work, and your monthly business costs.", None),
    ("2. Open 2 Trade Presets. Pick the rows for your trade and change the hours, supplies and miles to match your real jobs. Add your own job types in the empty rows.", None),
    ("3. Open 3 Quote Builder. Choose a job from the dropdown and read your quote. Change any yellow cell for this one job.", None),
    ("4. Open 4 Price Sheet to see every job type priced at once. Print it or use it to build your service menu.", None),
    ("5. Log finished jobs in 5 Job Tracker to see your real profit and effective hourly rate per job.", None),
    ("", None),
    ("How to read the colors", "h"),
    ("Yellow cells with blue text are yours to change. Green cells are formulas and peach cells are your key results: leave both alone and they update on their own. Type percentages with the % sign, for example 25%.", None),
    ("Every number already in the file is an example so you can see how it works. Replace the examples with your own numbers.", None),
    ("", None),
    ("Works in Excel and Google Sheets", "h"),
    ("Excel: open the .xlsx file. Google Sheets: go to sheets.google.com, then File, Import, Upload, and choose this file. No add-ons, no macros, no sign-up, no survey, no outside link needed.", None),
    ("", None),
    ("Why the rate is higher than you expect", "h"),
    ("Most service owners divide the pay they want by the hours they work. That leaves out three things: hours you are not paid for (driving, quoting, invoicing, cleanup), business costs that keep running every month, and self-employment tax. Tab 1 adds all three back.", None),
    ("", None),
    ("Figures used, verified Oct 3, 2026", "h"),
    ("Self-employment tax 15.3% (12.4% Social Security plus 2.9% Medicare) on 92.35% of net profit; Social Security part applies up to the 2026 wage base of $184,500. Sources: SSA contribution and benefit base; IRS Topic 554 and Schedule SE. The calculator applies 15.3% to all profit, so if your profit is above about $199,800 your real tax is a little lower.", None),
    ("Vehicle cost per mile defaults to the IRS business standard mileage rate of 76 cents, effective July 1, 2026 (IRS Announcement 2026-11; it was 72.5 cents from January 1 to June 30, 2026). Change it to your own cost per mile if you track it.", None),
    ("The income tax rate is your own estimate. Ask your tax preparer for your effective rate. Tax rules and rates change; check them each year.", None),
    ("", None),
    ("Important", "h"),
    ("This workbook is a planning tool. It gives information only, not tax, legal or financial advice, and it does not guarantee any result. Example trade times are starting points, not market prices. See LICENSE-AND-DISCLAIMER.txt.", None),
    ("Made with AI assistance and reviewed and tested by the shop owner. Copyright ProofNotFluff.", None),
]
r = 5
for text, kind in lines:
    c = ws.cell(row=r, column=2, value=text if text else None)
    if kind == "h":
        c.font = Font(name=F, size=12, bold=True, color=ACC)
    else:
        c.font = font(10); c.alignment = Alignment(wrap_text=True, vertical="top")
        if text: ws.row_dimensions[r].height = 15 * max(1, (len(text) // 110) + 1)
    r += 1

# ---------------- 1 Your Numbers ----------------
y = wb.create_sheet("1 Your Numbers"); y.sheet_view.showGridLines = False
for col, w in zip("ABCD", [52, 16, 4, 66]):
    y.column_dimensions[col].width = w
title(y, "1  Your Numbers", "Fill in the yellow cells. Everything else calculates.")

def sec(row, text):
    c = y.cell(row=row, column=1, value=text); c.font = Font(name=F, size=12, bold=True, color=ACC)

sec(4, "Pay and time")
label(y["A5"], "Take-home pay you want per year, after taxes"); inp(y["B5"], 52000, USD0)
y["D5"] = "What you want to keep for yourself."
label(y["A6"], "Weeks you work per year"); inp(y["B6"], 48, '0')
y["D6"] = "52 minus vacation, holidays and sick weeks."
label(y["A7"], "Hours you work per week, all of it"); inp(y["B7"], 45, '0')
y["D7"] = "Include driving, quoting, texts, invoicing, shopping for supplies."
label(y["A8"], "Share of those hours you are actually paid for"); inp(y["B8"], 0.65, PCT)
y["D8"] = "Paid hours divided by all hours. Not sure? Try 60% to 70%."
label(y["A9"], "Paid hours per year", bold=True); out(y["B9"], "=B6*B7*B8", '#,##0')

sec(11, "Business costs per month")
header(y, 12, ["Cost", "Per month"])
costs = [("Business insurance", 120), ("Phone and internet (business share)", 90), ("Software, apps and booking tools", 45),
         ("Marketing, ads, website", 100), ("Equipment and tools (replacement, spread per month)", 75),
         ("Shared supplies not tied to one job", 60), ("Accounting, bank and payment fees", 40), ("Licenses and permits", 15),
         ("Vehicle payment, only if you set cost per mile to 0 below", 0), ("Other", 0), ("Other", 0)]
for i, (n, v) in enumerate(costs):
    rr = 13 + i
    c = y.cell(row=rr, column=1, value=n); c.font = font(); c.border = BOX
    inp(y.cell(row=rr, column=2), v, USD0)
label(y["A24"], "Total per month", bold=True); out(y["B24"], "=SUM(B13:B23)", USD0, bold=True)
label(y["A25"], "Total per year", bold=True); out(y["B25"], "=B24*12", USD0, bold=True)

sec(27, "Taxes, profit and travel")
label(y["A28"], "Self-employment tax rate"); inp(y["B28"], 0.153, PCT)
y["D28"] = "12.4% + 2.9%. IRS Topic 554 and Schedule SE, checked Oct 3, 2026."
label(y["A29"], "Share of profit that self-employment tax applies to"); inp(y["B29"], 0.9235, '0.00%')
y["D29"] = "IRS Schedule SE uses 92.35% of net profit."
label(y["A30"], "Your estimated income tax rate, federal plus state"); inp(y["B30"], 0.12, PCT)
y["D30"] = "An estimate. Ask your tax preparer for your effective rate."
label(y["A31"], "Profit margin to build into every quote"); inp(y["B31"], 0.15, PCT)
y["D31"] = "For slow months, new equipment and growth. Type it as a percent."
label(y["A32"], "Vehicle cost per mile"); inp(y["B32"], 0.76, '$0.000')
y["D32"] = "IRS rate from July 1, 2026 (Ann. 2026-11). Or use your own."
label(y["A33"], "Average driving speed between jobs, mph"); inp(y["B33"], 30, '0')
y["D33"] = "Used to turn miles into drive time."

sec(35, "Your results")
label(y["A36"], "Profit you need before taxes"); out(y["B36"], "=IFERROR(B5/(1-B28*B29-B30),0)", USD0)
y["D36"] = "Your pay plus self-employment and income tax (an estimate)."
label(y["A37"], "Revenue you need per year, before job supplies and mileage"); out(y["B37"], "=B36+B25", USD0)
label(y["A38"], "Break-even rate per paid hour", bold=True); out(y["B38"], "=IFERROR(B37/B9,0)", USD, bold=True, fill=KEY_FILL)
y["D38"] = "Charge less than this and you fall short of your pay goal."
label(y["A39"], "Target rate per paid hour, with profit", bold=True); out(y["B39"], "=IFERROR(B38/(1-B31),0)", USD, bold=True, fill=KEY_FILL)
y["D39"] = "Use this to price jobs. The Quote Builder already does."
label(y["A41"], "The rate you would guess (pay divided by all hours)"); out(y["B41"], "=IFERROR(B5/(B6*B7),0)", USD)
label(y["A42"], "How much the guess underprices each paid hour"); out(y["B42"], "=B39-B41", USD)
y["D42"] = "The gap most generic calculators never show you."
for rr in list(range(5, 9)) + list(range(28, 34)) + [36, 38, 39, 42]:
    y.cell(row=rr, column=4).font = font(9, color="5A5F6B", italic=True)

# ---------------- 2 Trade Presets ----------------
p = wb.create_sheet("2 Trade Presets"); p.sheet_view.showGridLines = False
title(p, "2  Trade Presets", "Example starting points, not market prices. Change every number to match your real jobs, and add your own job types in the empty rows.")
for col, w in zip("ABCDEFGH", [46, 22, 30, 14, 10, 14, 14, 44]):
    p.column_dimensions[col].width = w
header(p, 5, ["Job (shows in dropdowns)", "Trade", "Job type", "Hours on site per person", "People", "Supplies per job ($)", "Round-trip miles", "Notes"])
presets = [
    ("House cleaning", "Standard clean, 2 bed", 2.5, 1, 8, 15, ""),
    ("House cleaning", "Deep clean, 2 bed", 5, 1, 15, 15, ""),
    ("House cleaning", "Move-out clean, 3 bed", 7, 1, 20, 15, ""),
    ("Mobile auto detailing", "Interior detail, sedan", 2.5, 1, 12, 20, ""),
    ("Mobile auto detailing", "Full detail, sedan", 4, 1, 20, 20, ""),
    ("Mobile auto detailing", "Full detail, SUV or truck", 5, 1, 25, 20, ""),
    ("Lawn care", "Weekly mow, small yard", 0.75, 1, 3, 10, "Fuel and blades in supplies"),
    ("Lawn care", "Spring cleanup", 4, 2, 20, 10, ""),
    ("Lawn care", "Mulch install, labor only", 5, 2, 0, 10, "Bill mulch separately as an add-on"),
    ("Photography", "Mini session", 3, 1, 5, 20, "1 hour shooting + 2 hours editing"),
    ("Photography", "Family session", 5.5, 1, 10, 20, "1.5 hours shooting + 4 hours editing"),
    ("Photography", "Wedding, 8 hours", 38, 1, 60, 40, "8 hours shooting + 30 hours editing"),
    ("Hair and salon", "Cut and style", 1, 1, 6, 0, ""),
    ("Hair and salon", "Full color", 2.5, 1, 25, 0, ""),
    ("Hair and salon", "Balayage", 3.5, 1, 35, 0, ""),
    ("Handyman", "TV wall mount", 1.5, 1, 10, 15, ""),
    ("Handyman", "Small repair visit", 2, 1, 10, 15, ""),
    ("Handyman", "Drywall patch and paint", 3, 1, 25, 15, ""),
    ("Pressure washing", "Driveway", 2, 1, 10, 20, ""),
    ("Pressure washing", "House wash, 1 story", 4, 1, 30, 20, ""),
]
FIRST, LAST = 6, 45
for i in range(FIRST, LAST + 1):
    vals = presets[i - FIRST] if i - FIRST < len(presets) else (None,) * 7
    k = p.cell(row=i, column=1, value=f'=IF(C{i}="","",B{i}&" - "&C{i})'); k.font = font(10, bold=True); k.fill = OUT_FILL; k.border = BOX
    for j, v in enumerate(vals):
        c = p.cell(row=i, column=2 + j)
        if j == 6:
            c.value = v if v else None; c.font = font(9, italic=True, color="5A5F6B"); c.border = BOX
        else:
            fmt = {2: '0.00', 3: '0', 4: USD0, 5: '0'}.get(j)
            inp(c, v, fmt)

# ---------------- 3 Quote Builder ----------------
q = wb.create_sheet("3 Quote Builder"); q.sheet_view.showGridLines = False
title(q, "3  Quote Builder", "Pick a job, adjust anything for this one customer, and read your price.")
for col, w in zip("ABCD", [44, 20, 4, 58]):
    q.column_dimensions[col].width = w
YN = "'1 Your Numbers'"; TP = "'2 Trade Presets'"
label(q["A4"], "Job", bold=True); inp(q["B4"], "House cleaning - Standard clean, 2 bed"); q.merge_cells("B4:D4")
dv = DataValidation(type="list", formula1=f"={TP}!$A${FIRST}:$A${LAST}", allow_blank=False)
dv.error = "Pick a job from the list, or add one in 2 Trade Presets."; dv.showErrorMessage = True; q.add_data_validation(dv); dv.add("B4")
q["A5"] = "Leave the override cells blank to use the preset. Type a number to change it for this quote only."
q["A5"].font = font(9, italic=True, color="5A5F6B")

def look(col):
    return f'IFERROR(INDEX({TP}!${col}${FIRST}:${col}${LAST},MATCH($B$4,{TP}!$A${FIRST}:$A${LAST},0)),0)'

header(q, 7, ["Item", "Preset", "", "Override for this quote"])
rows = [("Hours on site per person", "D", '0.00'), ("People on the job", "E", '0'), ("Supplies for this job", "F", USD0), ("Round-trip miles", "G", '0')]
for i, (lab, col, fmt) in enumerate(rows):
    rr = 8 + i
    label(q.cell(row=rr, column=1), lab)
    out(q.cell(row=rr, column=2), "=" + look(col), fmt)
    inp(q.cell(row=rr, column=4), None, fmt)
label(q["A12"], "Add-ons billed at cost (materials, parking, fees)"); inp(q["D12"], 0, USD0)
label(q["A13"], "Deposit to collect"); inp(q["D13"], 0.25, PCT)
label(q["A14"], "Your current price for this job, if you have one"); inp(q["D14"], 150, USD0)

q["A16"] = "Working numbers"; q["A16"].font = Font(name=F, size=12, bold=True, color=ACC)
label(q["A17"], "Hours on site per person (used)"); out(q["B17"], '=IF(D8="",B8,D8)', '0.00')
label(q["A18"], "People (used)"); out(q["B18"], '=IF(D9="",B9,D9)', '0')
label(q["A19"], "Supplies (used)"); out(q["B19"], '=IF(D10="",B10,D10)', USD)
label(q["A20"], "Miles (used)"); out(q["B20"], '=IF(D11="",B11,D11)', '0')
label(q["A21"], "Drive time, hours per person"); out(q["B21"], f"=IFERROR(B20/{YN}!$B$33,0)", '0.00')
label(q["A22"], "Labor on site at your break-even rate"); out(q["B22"], f"=B17*B18*{YN}!$B$38", USD)
label(q["A23"], "Drive time at your break-even rate"); out(q["B23"], f"=B21*B18*{YN}!$B$38", USD)
label(q["A24"], "Vehicle cost"); out(q["B24"], f"=B20*{YN}!$B$32", USD)
label(q["A25"], "Supplies"); out(q["B25"], "=B19", USD)
label(q["A26"], "Price floor: covers your pay, costs and taxes, zero profit", bold=True); out(q["B26"], "=SUM(B22:B25)+D12", USD, bold=True)

q["A28"] = "Your quote"; q["A28"].font = Font(name=F, size=12, bold=True, color=ACC)
label(q["A29"], "Quote, rounded up to the next $5", bold=True, size=12)
out(q["B29"], f"=CEILING((B26-D12)/(1-{YN}!$B$31)+D12,5)", USD0, bold=True, fill=KEY_FILL)
q["B29"].font = Font(name=F, size=14, bold=True, color=INK)
label(q["A30"], "Deposit due at booking"); out(q["B30"], "=ROUND(B29*D13,2)", USD)
label(q["A31"], "Profit in this quote"); out(q["B31"], "=B29-B26", USD)
label(q["A32"], "Effective rate per hour worked, including drive time"); out(q["B32"], "=IFERROR((B29-B24-B25-D12)/((B17+B21)*B18),0)", USD)
label(q["A34"], "At your current price you would make"); out(q["B34"], '=IF(D14="","",D14-B26)', USD)
label(q["A35"], "Check"); out(q["B35"], '=IF(D14="","Enter your current price to compare",IF(D14<B26,"Below your floor: this price does not cover your pay goal and costs",IF(D14<B29,"Above the floor but under your target","At or above your target")))')
q.merge_cells("B35:D35")
q["D17"] = "Drive time and vehicle cost are the two costs most quotes leave out."
q["D17"].font = font(9, italic=True, color="5A5F6B")

# ---------------- 4 Price Sheet ----------------
s = wb.create_sheet("4 Price Sheet"); s.sheet_view.showGridLines = False
title(s, "4  Price Sheet", "Every job type in 2 Trade Presets, priced with your numbers. Updates automatically.")
for col, w in zip("ABCDE", [24, 32, 16, 16, 18]):
    s.column_dimensions[col].width = w
header(s, 4, ["Trade", "Job type", "Price floor", "Your price", "Per hour worked"])
for i in range(FIRST, LAST + 1):
    rr = i - FIRST + 5
    h = f"{TP}!D{i}"; n = f"{TP}!E{i}"; sup = f"{TP}!F{i}"; mi = f"{TP}!G{i}"
    drive = f"IFERROR({mi}/{YN}!$B$33,0)"
    floor = f"({h}*{n}+{drive}*{n})*{YN}!$B$38+{mi}*{YN}!$B$32+{sup}"
    out(s.cell(row=rr, column=1), f'=IF({TP}!C{i}="","",{TP}!B{i})', None, fill=PAPER_FILL)
    out(s.cell(row=rr, column=2), f'=IF({TP}!C{i}="","",{TP}!C{i})', None, fill=PAPER_FILL)
    out(s.cell(row=rr, column=3), f'=IF({TP}!C{i}="","",{floor})', USD0)
    out(s.cell(row=rr, column=4), f'=IF({TP}!C{i}="","",CEILING(C{rr}/(1-{YN}!$B$31),5))', USD0, bold=True, fill=KEY_FILL)
    out(s.cell(row=rr, column=5), f'=IF({TP}!C{i}="","",IFERROR((D{rr}-{mi}*{YN}!$B$32-{sup})/(({h}+{drive})*{n}),0))', USD)

# ---------------- 5 Job Tracker ----------------
t = wb.create_sheet("5 Job Tracker"); t.sheet_view.showGridLines = False
title(t, "5  Job Tracker", "Log each finished job. Yellow columns are yours; the rest calculates. Two example rows show the format; replace them. Job is a label: type the real hours, miles and supplies for each visit.")
cols = ["Date", "Client", "Job", "Price charged", "Person-hours on site", "Drive hours", "Miles", "Supplies", "Job cost at break-even", "Profit beyond break-even", "Earned per hour worked", "Check"]
widths = [12, 18, 36, 13, 12, 11, 9, 11, 14, 12, 14, 26]
for i, w in enumerate(widths):
    t.column_dimensions[chr(65 + i)].width = w
header(t, 5, cols)
import datetime
ex = [(datetime.date(2026, 10, 1), "Example: Lee", "House cleaning - Standard clean, 2 bed", 150, 2.5, 0.5, 15, 8),
      (datetime.date(2026, 10, 2), "Example: Ortiz", "Mobile auto detailing - Full detail, sedan", 350, 4, 0.7, 20, 20)]
dvj = DataValidation(type="list", formula1=f"={TP}!$A${FIRST}:$A${LAST}", allow_blank=True); dvj.showErrorMessage = True; dvj.error = "Pick a job from the list in 2 Trade Presets."; t.add_data_validation(dvj)
for r_ in range(6, 106):
    vals = ex[r_ - 6] if r_ - 6 < len(ex) else (None,) * 8
    for j, v in enumerate(vals):
        fmt = ['mm/dd/yyyy', None, None, USD0, '0.00', '0.00', '0', USD0][j]
        inp(t.cell(row=r_, column=j + 1), v, fmt)
    dvj.add(f"C{r_}")
    out(t.cell(row=r_, column=9), f'=IF(D{r_}="","",(E{r_}+F{r_})*{YN}!$B$38+G{r_}*{YN}!$B$32+H{r_})', USD0)
    out(t.cell(row=r_, column=10), f'=IF(D{r_}="","",D{r_}-I{r_})', USD0)
    out(t.cell(row=r_, column=11), f'=IF(D{r_}="","",IFERROR((D{r_}-G{r_}*{YN}!$B$32-H{r_})/(E{r_}+F{r_}),0))', USD)
    out(t.cell(row=r_, column=12), f'=IF(D{r_}="","",IF(J{r_}<0,"Below floor",IF(K{r_}<{YN}!$B$39,"Under target","On target")))')
def tl(cell, text):
    t[cell] = text; t[cell].font = font(10, bold=True); t[cell].alignment = Alignment(horizontal="right")
tl("A3", "Revenue"); out(t["B3"], "=SUM(D6:D105)", USD0, bold=True)
tl("C3", "Profit beyond your break-even"); out(t["D3"], "=SUM(J6:J105)", USD0, bold=True)
tl("A4", "Per hour"); out(t["B4"], '=IFERROR((SUM(D6:D105)-SUM(G6:G105)*\'1 Your Numbers\'!$B$32-SUM(H6:H105))/(SUM(E6:E105)+SUM(F6:F105)),0)', USD, bold=True)
tl("C4", "Jobs below your floor"); out(t["D4"], '=COUNTIF(L6:L105,"Below floor")', '0', bold=True)
t.freeze_panes = "A6"

# ---------------- Thank You ----------------
k = wb.create_sheet("Thank You"); k.sheet_view.showGridLines = False
k.column_dimensions["B"].width = 90
k["B2"] = "Thank you for buying from ProofNotFluff"; k["B2"].font = Font(name=F, size=16, bold=True, color=INK)
k["B4"] = "If this calculator helped you price your work, a short Etsy review helps other service owners find it."; k["B4"].font = font(11)
k["B5"] = "If anything doesn't open or a number looks wrong, message me on Etsy and I'll fix it."; k["B5"].font = font(11)
k["B7"] = "Information only, not tax, legal or financial advice. No guarantee of results. See LICENSE-AND-DISCLAIMER.txt."; k["B7"].font = font(9, italic=True, color="5A5F6B")

for sh in wb.worksheets:
    sh.sheet_properties.tabColor = ACC if sh.title in ("Start Here", "Thank You") else INK
for sh in wb.worksheets:
    sh.page_setup.paperSize = sh.PAPERSIZE_LETTER; sh.page_setup.orientation = "landscape"
    sh.sheet_properties.pageSetUpPr.fitToPage = True; sh.page_setup.fitToWidth = 1; sh.page_setup.fitToHeight = 0
wb.properties.title = "Service Pricing Calculator"; wb.properties.creator = "ProofNotFluff"
wb.save("/home/claude/factory/svc-pricing/Service-Pricing-Calculator.xlsx")
print("saved")
