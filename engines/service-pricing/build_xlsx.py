#!/usr/bin/env python3
"""#13 Service Pricing Calculator, version 3 (dashboard design, Oct 6, 2026).

Same inputs, presets, formulas and worked example as version 1 (Oct 3, 2026), rebuilt to
design/WORKBOOK_STANDARD.md section 0. tools/compare_xlsx.py with compare_map.json proves the
answers match. Fixes in v3: the mileage source now cites IRS news release IR-2026-29 (v1 named
an announcement number the IRS page does not use); sheets are protected without a password;
full Terms of Use tab.

Usage: python3 engines/service-pricing/build_xlsx.py out.xlsx
"""
import math
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
import datetime as dt  # noqa: E402

PRODUCT = "Service Pricing Calculator"
CHECKED = "Oct 6, 2026"
AS_OF = f"Figures checked {CHECKED}"
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else "Service-Pricing-Calculator.xlsx"
USD0, USD2, PCT, INT, NUM1, NUM2 = (S.FMT[k] for k in ("usd0", "usd2", "pct", "int", "num1", "num2"))
PCT1, PCT2, DATE = "0.0%", "0.00%", S.FMT["date"]
HRS2 = '#,##0.00 "hrs"'

# ---------------------------------------------------------------- example data (v1, unchanged)
EX = dict(take_home=52000, weeks=48, hours=45, paid_share=0.65, se=0.153, se_base=0.9235, income=0.12,
          margin=0.15, per_mile=0.76, mph=30)
COSTS = [("Business insurance", 120), ("Phone and internet", 90), ("Software and booking apps", 45),
         ("Marketing, ads and website", 100), ("Equipment and tools", 75),
         ("Shared supplies", 60), ("Accounting and payment fees", 40), ("Licenses and permits", 15),
         ("Vehicle payment", 0), ("Other", 0), ("Other", 0)]
PRESETS = [
    ("House cleaning", "Standard clean, 2 bed", 2.5, 1, 8, 15, None),
    ("House cleaning", "Deep clean, 2 bed", 5, 1, 15, 15, None),
    ("House cleaning", "Move-out clean, 3 bed", 7, 1, 20, 15, None),
    ("Mobile auto detailing", "Interior detail, sedan", 2.5, 1, 12, 20, None),
    ("Mobile auto detailing", "Full detail, sedan", 4, 1, 20, 20, None),
    ("Mobile auto detailing", "Full detail, SUV or truck", 5, 1, 25, 20, None),
    ("Lawn care", "Weekly mow, small yard", 0.75, 1, 3, 10, "Fuel and blades in supplies"),
    ("Lawn care", "Spring cleanup", 4, 2, 20, 10, None),
    ("Lawn care", "Mulch install, labor only", 5, 2, 0, 10, "Bill mulch separately as an add-on"),
    ("Photography", "Mini session", 3, 1, 5, 20, "1 hour shooting + 2 hours editing"),
    ("Photography", "Family session", 5.5, 1, 10, 20, "1.5 hours shooting + 4 hours editing"),
    ("Photography", "Wedding, 8 hours", 38, 1, 60, 40, "8 hours shooting + 30 hours editing"),
    ("Hair and salon", "Cut and style", 1, 1, 6, 0, None),
    ("Hair and salon", "Full color", 2.5, 1, 25, 0, None),
    ("Hair and salon", "Balayage", 3.5, 1, 35, 0, None),
    ("Handyman", "TV wall mount", 1.5, 1, 10, 15, None),
    ("Handyman", "Small repair visit", 2, 1, 10, 15, None),
    ("Handyman", "Drywall patch and paint", 3, 1, 25, 15, None),
    ("Pressure washing", "Driveway", 2, 1, 10, 20, None),
    ("Pressure washing", "House wash, 1 story", 4, 1, 30, 20, None),
]
N_PRESETS = 40
N_JOBS = 100
JOBS = [(dt.datetime(2026, 10, 1), "Example client", "House cleaning - Standard clean, 2 bed", 150, 2.5, 0.5, 15, 8),
        (dt.datetime(2026, 10, 2), "Example client", "Mobile auto detailing - Full detail, sedan", 350, 4, 0.7, 20, 20)]
QUOTE_JOB = "House cleaning - Standard clean, 2 bed"
CURRENT_PRICE = 150
DEPOSIT = 0.25

# the example's answers, computed the same way the workbook does, so every sentence matches the file
_profit = EX["take_home"] / (1 - EX["se"] * EX["se_base"] - EX["income"])
_costs_y = sum(v for _, v in COSTS) * 12
_paid = EX["weeks"] * EX["hours"] * EX["paid_share"]
BE_X = (_profit + _costs_y) / _paid
TARGET_X = BE_X / (1 - EX["margin"])
GUESS_X = EX["take_home"] / (EX["weeks"] * EX["hours"])
_h, _p, _s, _m = 2.5, 1, 8, 15
_drive = _m / EX["mph"]
FLOOR_X = (_h * _p + _drive * _p) * BE_X + _m * EX["per_mile"] + _s
QUOTE_X = math.ceil((FLOOR_X / (1 - EX["margin"])) / 5 - 1e-9) * 5


def money(x, cents=True):
    return f"${x:,.2f}" if cents else f"${x:,.0f}"


wb = S.new_workbook(PRODUCT, version="3")
GRID = {"A": 3, "B": 2, "C": 31, "D": 14, "E": 2, "F": 3, "G": 2, "H": 24, "I": 12, "J": 17, "K": 2, "L": 3}


def notes_card(ws, top, notes, span=("C", "J"), frame_span=("B", "K")):
    r = top
    D.h(ws, r, 12); r += 1
    t = ws[f"{span[0]}{r}"]; t.value = "Notes and sources"; t.font = D.f(12, True)
    ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}"); D.h(ws, r, 24); r += 1
    D.h(ws, r, 6); r += 1
    width = sum(S.width_of(ws, k) for k in D.cols(*span))
    for head, body in notes:
        cell = ws[f"{span[0]}{r}"]
        cell.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=9, b=True, color=D.T["ink"]), head + ". "),
                                  TextBlock(InlineFont(rFont=D.FONT, sz=9, color=D.T["ink2"]), body))
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}")
        D.h(ws, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
    D.h(ws, r, 12)
    D.frame(ws, top, r, *frame_span)
    return r


# =====================================================================  Your Numbers
yn = wb.create_sheet("Your Numbers")
S.set_widths(yn, GRID)
D.page_header(yn, "Your Numbers", "Your pay goal, real hours, costs and taxes, turned into the rate every quote starts from.", AS_OF)
TOP = 11

A = D.Card(yn, TOP, "B", "C", "D", "E", wb=wb)
A.title("Pay and time", "Type over the yellow cells. Every other number updates.")
TH = A.input("Take-home pay you want per year", EX["take_home"], USD0, name="TakeHome", validation=("decimal", 0, None),
             prompt="What you want to keep for yourself in a year, after tax. Example: 52000.")
WK = A.input("Weeks you work per year", EX["weeks"], INT, name="WeeksPerYear", validation=("decimal", 1, 52),
             prompt="52 minus vacation, holidays and sick weeks. Example: 48.")
HW = A.input("Hours you work per week, all of it", EX["hours"], INT, name="HoursPerWeek", validation=("decimal", 1, 100),
             prompt="Include driving, quoting, texts, invoicing and shopping for supplies. Example: 45.")
PS = A.input("Share of those hours you are paid for", EX["paid_share"], PCT, name="PaidShare", validation=("decimal", 0.05, 1),
             prompt="Paid hours divided by all hours. Not sure? Try 60% to 70%.")
PH = A.calc("Paid hours per year", f"=D{WK}*D{HW}*D{PS}", INT, bold=True)
A.close()

C2 = D.Card(yn, A.end + 2, "B", "C", "D", "E", wb=wb)
C2.title("Business costs per month", "Costs that run every month, busy or slow.")
cost_rows = []
for i, (lab, val) in enumerate(COSTS):
    prompt = ("Your monthly cost, in dollars. Equipment: what you spend replacing tools, spread per month. Type 0 if it doesn't apply." if "Vehicle" not in lab else
              "Only if you set Vehicle cost per mile to 0. Otherwise the per-mile rate already covers the vehicle.")
    cost_rows.append(C2.input(lab, val, USD0, validation=("decimal", 0, None), prompt=prompt, prompt_title="Monthly cost"))
CM = C2.calc("Total per month", f"=SUM(D{cost_rows[0]}:D{cost_rows[-1]})", USD0, total=True)
CY = C2.calc("Total per year", f"=D{CM}*12", USD0, bold=True)
C2.close()

# right column: build-up first (its rows feed the tiles), taxes card under it
B = D.Card(yn, TOP, "G", "H", "I", "K", extra="J", wb=wb)
B.title("How your rate is built", "Your pay, grossed up for tax, plus your costs, spread over paid hours.")
b_th = B.calc("Take-home you want", f"=$D${TH}", USD0)
b_tax_row = B.r  # filled once the tax inputs exist
B.calc("Plus tax (SE and income)", "", USD0)
b_profit = B.calc("Profit you need before tax", "", USD0, total=True)
b_cost = B.calc("Plus business costs", f"=$D${CY}", USD0)
b_rev = B.calc("Revenue you need", f"=I{b_profit}+I{b_cost}", USD0, total=True)
B.eyebrow("SPREAD OVER YOUR PAID HOURS")
b_ph = B.calc("Paid hours per year", f"=$D${PH}", INT)
b_be = B.calc("Break-even per paid hour", f"=IFERROR(I{b_rev}/I{b_ph},0)", USD2, total=True)
b_mg_row = B.r
B.calc("Plus your profit margin", "", USD2)
b_tg = B.calc("Target rate per paid hour", "", USD2, total=True, tint="total_tint")
B.eyebrow("WHERE EACH PAID HOUR GOES")
w1 = B.r; w4 = B.r + 3
wmx = f"MAX($I${w1}:$I${w4})"
B.bar("Your pay", f"=IFERROR($I${b_th}/$I${b_ph},0)", USD2, wmx, "teal")
B.bar("Tax", f"=IFERROR($I${b_tax_row}/$I${b_ph},0)", USD2, wmx, "accent")
B.bar("Business costs", f"=IFERROR($I${b_cost}/$I${b_ph},0)", USD2, wmx, "gold")
B.bar("Profit margin", f"=$I${b_mg_row}", USD2, wmx, "blue")
B.close()

T2 = D.Card(yn, B.end + 2, "G", "H", "I", "K", wb=wb)
T2.title("Taxes, profit and travel", "Select a cell to see its source.")
SE = T2.input("Self-employment tax rate", EX["se"], PCT1, name="SETaxRate", validation=("decimal", 0, 0.2),
              prompt="15.3%: 12.4% Social Security plus 2.9% Medicare. IRS Topic 554 and Schedule SE, checked " + CHECKED + ".")
SB = T2.input("SE tax applies to", EX["se_base"], PCT2, name="SETaxBase", validation=("decimal", 0, 1),
              prompt="IRS Schedule SE applies self-employment tax to 92.35% of net profit.")
IT = T2.input("Income tax rate, estimate", EX["income"], PCT, name="IncomeTaxRate", validation=("decimal", 0, 0.5),
              prompt="Federal plus state, as an effective rate. An estimate: ask your tax preparer for yours.")
MG = T2.input("Profit margin in quotes", EX["margin"], PCT, name="ProfitMargin", validation=("decimal", 0, 0.9),
              prompt="Built into every quote for slow months, new equipment and growth. Example: 15%.")
CPM = T2.input("Vehicle cost per mile", EX["per_mile"], USD2, name="CostPerMile", validation=("decimal", 0, 5),
               prompt="IRS business rate 76 cents a mile for Jul 1 to Dec 31, 2026 (IR-2026-29). Or your own cost per mile.")
SPD = T2.input("Speed between jobs, mph", EX["mph"], INT, name="DriveSpeed", validation=("decimal", 1, 90),
               prompt="Used to turn miles into drive time. Example: 30.")
T2.close()

# now that the tax inputs exist, fill the build-up rows (same formulas as v1)
yn[f"I{b_profit}"] = f"=IFERROR($D${TH}/(1-$I${SE}*$I${SB}-$I${IT}),0)"
yn[f"I{b_tax_row}"] = f"=I{b_profit}-I{b_th}"
yn[f"I{b_tg}"] = f"=IFERROR(I{b_be}/(1-$I${MG}),0)"
yn[f"I{b_mg_row}"] = f"=I{b_tg}-I{b_be}"
BE = f"'Your Numbers'!$I${b_be}"
TG = f"'Your Numbers'!$I${b_tg}"
CPMR = f"'Your Numbers'!$I${CPM}"
SPDR = f"'Your Numbers'!$I${SPD}"
MGR = f"'Your Numbers'!$I${MG}"

# tiles
GUESS = f'IFERROR($D${TH}/($D${WK}*$D${HW}),0)'
D.tile(yn, 6, "B", "E", "YOUR TARGET RATE PER PAID HOUR", f"=$I${b_tg}", USD2,
       f'="Break-even "&TEXT($I${b_be},"$#,##0.00")&". Every quote in this workbook starts here."', dark=True)
v, sub, _ = D.tile(yn, 6, "G", "K", "THE USUAL GUESS UNDERPRICES EACH PAID HOUR BY", f"=$I${b_tg}-{GUESS}", USD2,
                   f'="The guess: "&TEXT({GUESS},"$#,##0.00")&" an hour, your pay divided by all hours"')
D.status_rule(yn, v.coordinate, f"($I${b_tg}-{GUESS})>0", fill_tint=False)
D.h(yn, 10, 14)

n_top = max(C2.end, T2.end) + 2
NOTES = [
    ("Self-employment tax", "15.3% (12.4% Social Security plus 2.9% Medicare) on 92.35% of net profit, per IRS Topic 554 and "
     "Schedule SE. The Social Security part stops at the 2026 wage base of $184,500 (SSA, ssa.gov/OACT/COLA/cbb.html). "
     "The calculator applies 15.3% to all profit, so above about $199,800 of profit your real tax is a little lower. "
     f"Checked {CHECKED}."),
    ("Vehicle cost per mile", "Defaults to the IRS business standard mileage rate of 76 cents a mile for July 1 to "
     "December 31, 2026 (news release IR-2026-29, irs.gov/tax-professionals/standard-mileage-rates); it was 72.5 cents "
     f"from January 1 to June 30, 2026. Checked {CHECKED}. Use your own cost per mile if you track it."),
    ("Income tax", "Your own estimate of your effective federal plus state rate. Ask your tax preparer. Rates change; check "
     "them each year."),
    ("Before you rely on a number", "For general information and planning only. This is not legal, tax, financial or other "
     "professional advice, and using it doesn't create a professional relationship. Results are estimates based on the "
     f"numbers you enter. Figures were checked on {CHECKED} and can change. See the Terms tab."),
]
n_end = notes_card(yn, n_top, NOTES)
S.page_break_before(yn, n_top)
D.h(yn, n_end + 1, 14)
D.paint_canvas(yn, n_end + 1, "Z")
S.finish_sheet(yn, PRODUCT, n_end + 1, span=("A", "L"), freeze="A11", tab_color="accent")

# =====================================================================  Trade Presets
tp = wb.create_sheet("Trade Presets")
S.set_widths(tp, {"A": 3, "B": 2, "C": 22, "D": 28, "E": 11, "F": 9, "G": 11, "H": 11, "I": 40, "J": 40, "K": 2, "L": 3})
D.page_header(tp, "Trade Presets", "Example starting points, not market prices. Change every number to match your real jobs.",
              "Add your own in the empty rows", span=("B", "K"), side_cols=2)
cols_tp = [
    dict(col="C", head="Trade", kind="input", values=[p[0] for p in PRESETS], validation=("textLength", 0, 40),
         prompt="Your trade, for example House cleaning. Up to 40 characters."),
    dict(col="D", head="Job type", kind="input", values=[p[1] for p in PRESETS], validation=("textLength", 0, 40),
         prompt="What the customer buys, for example Deep clean, 2 bed. A row shows up in the dropdowns once this is filled."),
    dict(col="E", head="Hours per person", kind="input", fmt=NUM2.replace(".00", ".0#"), align="right",
         values=[p[2] for p in PRESETS], validation=("decimal", 0, 500), prompt="Hours each person spends on the job, editing included."),
    dict(col="F", head="People", kind="input", fmt=INT, align="right", values=[p[3] for p in PRESETS],
         validation=("whole", 1, 50), prompt="How many people work the job."),
    dict(col="G", head="Supplies", kind="input", fmt=USD0, align="right", values=[p[4] for p in PRESETS],
         validation=("decimal", 0, None), prompt="Supplies used up on this one job, in dollars."),
    dict(col="H", head="Miles round trip", kind="input", fmt=INT, align="right", values=[p[5] for p in PRESETS],
         validation=("decimal", 0, None), prompt="Miles there and back. 0 if clients come to you."),
    dict(col="I", head="Notes", kind="input", values=[p[6] or "" for p in PRESETS], validation=("textLength", 0, 80),
         prompt="Anything worth remembering about this job."),
    dict(col="J", head="Name in the dropdowns", kind="muted", formula=lambda r: f'=IF(D{r}="","",C{r}&" - "&D{r})'),
]
cols_tp[-1]["kind"] = "calc"
tp_first, tp_last, tp_end = D.table_card(tp, 6, "B", "K", cols_tp, N_PRESETS, title="Your job types",
                                         sub="Yellow columns are yours. The last column is what the Quote Builder and Job Tracker list.")
for r in range(tp_first, tp_last + 1):
    tp[f"J{r}"].font = D.f(10, False, "muted")
    for k in "CDI":
        tp[f"{k}{r}"].value = tp[f"{k}{r}"].value or None
D.paint_canvas(tp, tp_end + 1, "Z")
S.finish_sheet(tp, PRODUCT, tp_end + 1, span=("A", "L"), freeze=f"A{tp_first}", tab_color="teal")
tp.page_setup.orientation = "landscape"
TP = "'Trade Presets'!"
LIST = f"{TP}$J${tp_first}:$J${tp_last}"


def preset(col, job_cell):
    return f"IFERROR(INDEX({TP}${col}${tp_first}:${col}${tp_last},MATCH({job_cell},{LIST},0)),0)"


# =====================================================================  Quote Builder
qb = wb.create_sheet("Quote Builder")
S.set_widths(qb, GRID)
D.page_header(qb, "Quote Builder", "Pick a job, adjust it for this customer, and read your price.", AS_OF)

# job picker, full width
r = 11
D.h(qb, r, 12); r += 1
t = qb[f"C{r}"]; t.value = "Pick a job"; t.font = D.f(12, True); qb.merge_cells(f"C{r}:J{r}"); D.h(qb, r, 24); r += 1
t = qb[f"C{r}"]; t.value = "The list comes from the Trade Presets tab. Add a job type there and it shows up here."
t.font = D.f(9, False, "muted"); t.alignment = Alignment(vertical="top"); qb.merge_cells(f"C{r}:J{r}"); D.h(qb, r, 18); r += 1
JOB_ROW = r
lab = qb[f"C{r}"]; lab.value = "Job"; lab.font = D.f(10); lab.alignment = Alignment(vertical="center")
jc = qb[f"D{r}"]; jc.value = QUOTE_JOB; jc.font = D.f(11, True, "input_text")
jc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
jc.protection = jc.protection.copy(locked=False)
for k in D.cols("D", "J"):
    c = qb[f"{k}{r}"]; c.fill = D.fill("input_fill")
    c.border = Border(top=D.side("input_line"), bottom=D.side("input_line"),
                      left=D.side("input_line") if k == "D" else None, right=D.side("input_line") if k == "J" else None)
qb.merge_cells(f"D{r}:J{r}"); D.h(qb, r, 26)
dv = DataValidation(type="list", formula1=LIST, allow_blank=False)
dv.showErrorMessage = True; dv.errorStyle = "stop"; dv.errorTitle = "Pick a job"
dv.error = "Pick a job from the list, or add one on the Trade Presets tab."
dv.showInputMessage = True; dv.promptTitle = "Job"; dv.prompt = "Pick from the list. It comes from the Trade Presets tab."
qb.add_data_validation(dv); dv.add(f"D{r}")
r += 1
D.h(qb, r, 12)
D.frame(qb, 11, r, "B", "K")
JOB = f"$D${JOB_ROW}"
CARD_TOP = r + 2

L1 = D.Card(qb, CARD_TOP, "B", "C", "D", "E", wb=wb)
L1.title("Adjust for this customer", "Leave a cell blank to use the preset shown on the right.")
OH = L1.input("Hours on site per person", None, NUM2.replace(".00", ".0#"), validation=("decimal", 0, 500), allow_blank=True,
              prompt="Blank uses the preset. Type a number to change it for this quote only.")
OP = L1.input("People on the job", None, INT, validation=("whole", 1, 50), allow_blank=True,
              prompt="Blank uses the preset. Type a number to change it for this quote only.")
OS = L1.input("Supplies for this job", None, USD0, validation=("decimal", 0, None), allow_blank=True,
              prompt="Blank uses the preset. Type a number to change it for this quote only.")
OM = L1.input("Round-trip miles", None, INT, validation=("decimal", 0, None), allow_blank=True,
              prompt="Blank uses the preset. Type a number to change it for this quote only.")
L1.eyebrow("EXTRAS")
AO = L1.input("Add-ons billed at cost", 0, USD2, name="AddOns", validation=("decimal", 0, None),
              prompt="Materials, parking or fees you pass through at cost, for example mulch. No margin is added.")
DP = L1.input("Deposit to collect", DEPOSIT, PCT, name="Deposit", validation=("decimal", 0, 1),
              prompt="Share of the quote you collect at booking. Example: 25%.")
L1.close()

R1 = D.Card(qb, CARD_TOP, "G", "H", "I", "K", extra="J", wb=wb)
R1.title("How this quote is built", "Your time at your break-even rate, plus the costs of the trip.")
R1.eyebrow("THE JOB")


def used(label, override_row, col, fmt):
    rr = R1.calc(label, f'=IF($D${override_row}="",{preset(col, JOB)},$D${override_row})', fmt)
    n = qb[f"J{rr}"]; n.value = f'=IF($D${override_row}="","preset","your number")'
    n.font = D.f(9, False, "muted"); n.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return rr


U_H = used("Hours on site per person", OH, "E", NUM2.replace(".00", ".0#"))
U_P = used("People", OP, "F", INT)
U_S = used("Supplies", OS, "G", USD0)
U_M = used("Round-trip miles", OM, "H", INT)
U_D = R1.calc("Drive time per person", f"=IFERROR(I{U_M}/{SPDR},0)", HRS2)
R1.eyebrow("YOUR COST AT BREAK-EVEN")
k1 = R1.r
QT_ROW_PLACE = "QUOTE"
c_lab = R1.calc("Labor on site", f"=I{U_H}*I{U_P}*{BE}", USD2)
c_drv = R1.calc("Drive time", f"=I{U_D}*I{U_P}*{BE}", USD2)
c_veh = R1.calc("Vehicle", f"=I{U_M}*{CPMR}", USD2)
c_sup = R1.calc("Supplies", f"=I{U_S}", USD2)
c_add = R1.calc("Add-ons at cost", f"=$D${AO}", USD2)
c_flr = R1.calc("Price floor, zero profit", f"=SUM(I{c_lab}:I{c_add})", USD2, total=True)
R1.eyebrow("YOUR QUOTE")
c_q = R1.calc("Quote, rounded up to $5", f"=CEILING((I{c_flr}-I{c_add})/(1-{MGR})+I{c_add},5)", USD2,
              total=True, tint="total_tint")
c_dep = R1.calc("Deposit at booking", f"=ROUND(I{c_q}*$D${DP},2)", USD2)
c_pf = R1.calc("Profit in this quote", f"=I{c_q}-I{c_flr}", USD2)
c_eff = R1.calc("Per hour worked, with driving", f"=IFERROR((I{c_q}-I{c_veh}-I{c_sup}-I{c_add})/((I{U_H}+I{U_D})*I{U_P}),0)", USD2)
R1.close()
# bars: where the quote goes (labor, drive, vehicle, supplies, add-ons, profit), scaled to the quote
for rr, color in ((c_lab, "teal"), (c_drv, "teal"), (c_veh, "gold"), (c_sup, "gold"), (c_add, "gold"), (c_pf, "blue")):
    b = qb[f"J{rr}"]
    b.value = f'=IFERROR(REPT("{D.BAR_CHAR}",MAX(0,ROUND(I{rr}/$I${c_q}*11,0))),"")'
    b.font = D.f(9, False, color); b.alignment = Alignment(horizontal="left", vertical="center", indent=1)

L2 = D.Card(qb, L1.end + 2, "B", "C", "D", "E", wb=wb)
L2.title("Check your current price", "Today's price against your floor and quote.")
CP = L2.input("Your current price for this job", CURRENT_PRICE, USD0, name="CurrentPrice", validation=("decimal", 0, None),
              allow_blank=True, prompt="What you charge now for this job. Leave blank to skip the check.")
c_now = L2.calc("Profit at your current price", f'=IF($D${CP}="","",$D${CP}-$I${c_flr})', USD2, bold=True)
pill = L2.text(f'=IF($D${CP}="","Enter your current price to compare",IF($D${CP}<$I${c_flr},"Below your floor: this price does not '
               f'cover your pay goal and costs",IF($D${CP}<$I${c_q},"Above the floor but under your target","At or above your target")))',
               size=10, color="ink", bold=True)
pill2 = L2._row()
qb.unmerge_cells(f"C{pill}:D{pill}")
qb.merge_cells(f"C{pill}:D{pill2}")
qb[f"C{pill}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
L2.close()
anchor = f"$D${CP}"
for formula, color, tint in ((f'AND({anchor}<>"",{anchor}<$I${c_flr})', "accent", "accent_tint"),
                             (f'AND({anchor}<>"",{anchor}>=$I${c_flr},{anchor}<$I${c_q})', "gold", "input_fill"),
                             (f'AND({anchor}<>"",{anchor}>=$I${c_q})', "teal", "teal_tint")):
    qb.conditional_formatting.add(f"C{pill}:D{pill2}", FormulaRule(formula=[formula], font=Font(color=D.T[color], bold=True),
                                                                  fill=PatternFill("solid", bgColor=D.T[tint]), stopIfTrue=True))

D.status_rule(qb, f"D{c_now}", f"$D${c_now}<0", fill_tint=False)
D.tile(qb, 6, "B", "E", "YOUR QUOTE FOR THIS JOB", f"=$I${c_q}", USD2,
       f'="Deposit at booking: "&TEXT($I${c_dep},"$#,##0.00")', dark=True)
D.tile(qb, 6, "G", "K", "PROFIT IN THIS QUOTE", f"=$I${c_pf}", USD2,
       f'="You earn "&TEXT($I${c_eff},"$#,##0.00")&" per hour worked, drive time included"')
D.h(qb, 10, 14)
QB_LAST = max(L2.end, R1.end) + 1
D.h(qb, QB_LAST, 14)
D.paint_canvas(qb, QB_LAST, "Z")
S.finish_sheet(qb, PRODUCT, QB_LAST, span=("A", "L"), freeze="A11", tab_color="accent")

# =====================================================================  Price Sheet
ps = wb.create_sheet("Price Sheet")
S.set_widths(ps, {"A": 3, "B": 2, "C": 22, "D": 34, "E": 14, "F": 14, "G": 16, "H": 2, "I": 3})
D.page_header(ps, "Price Sheet", "Every job in Trade Presets, priced with your numbers. Print it as your service menu.", AS_OF,
              span=("B", "H"), side_cols=3)


def tp_ref(col, r):
    return f"{TP}{col}{r - ps_first + tp_first}"


ps_first_holder = {}


def ps_formula(kind):
    def fn(r):
        i = r - ps_first_holder["first"] + tp_first
        Dd, Ee, Gg, Hh, Ff = (f"{TP}{c}{i}" for c in "DEGHF")
        cond = f'{TP}D{i}=""'
        if kind == "trade":
            return f'=IF({cond},"",{TP}C{i})'
        if kind == "job":
            return f'=IF({cond},"",{Dd})'
        if kind == "floor":
            return (f'=IF({cond},"",({Ee}*{Ff}+IFERROR({Hh}/{SPDR},0)*{Ff})*{BE}+{Hh}*{CPMR}+{Gg})')
        if kind == "price":
            return f'=IF({cond},"",CEILING(E{r}/(1-{MGR}),5))'
        if kind == "hour":
            return (f'=IF({cond},"",IFERROR((F{r}-{Hh}*{CPMR}-{Gg})/(({Ee}+IFERROR({Hh}/{SPDR},0))*{Ff}),0))')
        if kind == "bar":
            return f'=IF({cond},"",IFERROR(REPT("{D.BAR_CHAR}",MAX(0,ROUND(F{r}/MAX($F${ps_first_holder["first"]}:$F${ps_first_holder["first"] + N_PRESETS - 1})*12,0))),""))'
    return fn


# table_card writes header at top+4 when a title and sub are given: first body row = top + 5
PS_TOP = 6
ps_first_holder["first"] = PS_TOP + 5
cols_ps = [
    dict(col="C", head="Trade", kind="calc", formula=ps_formula("trade")),
    dict(col="D", head="Job type", kind="calc", formula=ps_formula("job")),
    dict(col="E", head="Price floor", kind="calc", fmt=USD2, align="right", formula=ps_formula("floor")),
    dict(col="F", head="Your price", kind="calc", fmt=USD0, align="right", bold=True, formula=ps_formula("price")),
    dict(col="G", head="Per hour worked", kind="calc", fmt=USD2, align="right", formula=ps_formula("hour")),
]
ps_first, ps_last, ps_end = D.table_card(ps, PS_TOP, "B", "H", cols_ps, N_PRESETS, title="Your prices",
                                         sub="Price floor covers your pay, costs and taxes with zero profit. Your price adds your margin, rounded up to $5.")
assert ps_first == ps_first_holder["first"], (ps_first, ps_first_holder)
for r in range(ps_first, ps_last + 1):
    ps[f"C{r}"].font = D.f(10, False, "ink2")
D.paint_canvas(ps, ps_end + 1, "Z")
S.finish_sheet(ps, PRODUCT, ps_end + 1, span=("A", "I"), freeze=f"A{ps_first}", tab_color="blue")

# =====================================================================  Job Tracker
jt = wb.create_sheet("Job Tracker")
S.set_widths(jt, {"A": 3, "B": 2, "C": 13, "D": 16, "E": 44, "F": 11, "G": 10, "H": 9, "I": 8, "J": 10, "K": 13, "L": 14,
                  "M": 12, "N": 15, "O": 2, "P": 3})
D.page_header(jt, "Job Tracker", "Log each finished job and see what it really paid you.", "Two example rows: replace them",
              span=("B", "O"), side_cols=4)
JT_TOP_TABLE = 11
jt_first = JT_TOP_TABLE + 5
jt_last = jt_first + N_JOBS - 1
rngs = {k: f"{k}{jt_first}:{k}{jt_last}" for k in "FGHIJKLMN"}
# summary strip, rows 6 to 9
D.h(jt, 6, 10)
D.stat_strip(jt, 7, [
    ("C:D", "Revenue logged", f"=SUM({rngs['F']})", USD0),
    ("E:E", "Profit beyond break-even", f"=SUM({rngs['L']})", USD0),
    ("F:I", "Earned per hour worked", f"=IFERROR((SUM({rngs['F']})-SUM({rngs['I']})*{CPMR}-SUM({rngs['J']}))/(SUM({rngs['G']})+SUM({rngs['H']})),0)", USD2),
    ("J:N", "Jobs below your floor", f'=COUNTIF({rngs["N"]},"Below floor")', INT),
])
D.h(jt, 9, 10)
D.frame(jt, 6, 9, "B", "O")

jt["E8"].font = D.f(18, True)
D.status_rule(jt, "E8", "E8<0", fill_tint=False)
D.status_rule(jt, "J8", "J8>0", fill_tint=False)
D.h(jt, 10, 14)


def jt_calc(kind):
    def fn(r):
        if kind == "cost":
            return f'=IF(F{r}="","",(G{r}+H{r})*{BE}+I{r}*{CPMR}+J{r})'
        if kind == "profit":
            return f'=IF(F{r}="","",F{r}-K{r})'
        if kind == "hour":
            return f'=IF(F{r}="","",IFERROR((F{r}-I{r}*{CPMR}-J{r})/(G{r}+H{r}),0))'
        if kind == "check":
            return f'=IF(F{r}="","",IF(L{r}<0,"Below floor",IF(M{r}<{TG},"Under target","On target")))'
    return fn


cols_jt = [
    dict(col="C", head="Date", kind="input", fmt=DATE, values=[j[0] for j in JOBS], validation=("date", "DATE(2000,1,1)", "DATE(2100,12,31)"),
         prompt="The day you finished the job, for example 10/1/2026."),
    dict(col="D", head="Client", kind="input", values=[j[1] for j in JOBS], validation=("textLength", 0, 40),
         prompt="A name or a short label you will recognise."),
    dict(col="E", head="Job", kind="input", values=[j[2] for j in JOBS], validation=("list_range", LIST),
         prompt="Pick from the list. It comes from the Trade Presets tab.", error="Pick a job from the list, or add one on the Trade Presets tab."),
    dict(col="F", head="Price charged", kind="input", fmt=USD0, align="right", values=[j[3] for j in JOBS],
         validation=("decimal", 0, None), prompt="What the customer paid, before any tip."),
    dict(col="G", head="Hours on site", kind="input", fmt=NUM1, align="right", values=[j[4] for j in JOBS],
         validation=("decimal", 0, 1000), prompt="Hours on site for everyone added up. Two people for 3 hours is 6."),
    dict(col="H", head="Drive hours", kind="input", fmt=NUM1, align="right", values=[j[5] for j in JOBS],
         validation=("decimal", 0, 1000), prompt="Driving time for everyone added up."),
    dict(col="I", head="Miles", kind="input", fmt=INT, align="right", values=[j[6] for j in JOBS],
         validation=("decimal", 0, None), prompt="Round-trip miles driven."),
    dict(col="J", head="Supplies", kind="input", fmt=USD0, align="right", values=[j[7] for j in JOBS],
         validation=("decimal", 0, None), prompt="Supplies used on this job, in dollars."),
    dict(col="K", head="Cost at break-even", kind="calc", fmt=USD2, align="right", formula=jt_calc("cost")),
    dict(col="L", head="Profit over break-even", kind="calc", fmt=USD2, align="right", formula=jt_calc("profit")),
    dict(col="M", head="Earned per hour", kind="calc", fmt=USD2, align="right", formula=jt_calc("hour")),
    dict(col="N", head="Check", kind="calc", formula=jt_calc("check")),
]
f1, f2, jt_end = D.table_card(jt, JT_TOP_TABLE, "B", "O", cols_jt, N_JOBS, title="Finished jobs",
                              sub="Yellow columns are yours. Cost at break-even uses your break-even rate and cost per mile from Your Numbers.")
assert (f1, f2) == (jt_first, jt_last), (f1, f2, jt_first, jt_last)
for r in range(jt_first, jt_last + 1):
    jt[f"N{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for word, color, tint in (("Below floor", "accent", "accent_tint"), ("Under target", "gold", "input_fill"), ("On target", "teal", "teal_tint")):
    jt.conditional_formatting.add(f"N{jt_first}:N{jt_last}", FormulaRule(formula=[f'$N{jt_first}="{word}"'],
                                  font=Font(color=D.T[color], bold=True), fill=PatternFill("solid", bgColor=D.T[tint]), stopIfTrue=True))
D.paint_canvas(jt, jt_end + 1, "Z")
S.finish_sheet(jt, PRODUCT, jt_end + 1, span=("A", "P"), freeze=f"A{jt_first}", tab_color="gold")
jt.page_setup.orientation = "landscape"

# =====================================================================  Start Here and Terms (card style, from #20)
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            m = re.match(r"^([A-Z][A-Za-z ,]{1,34})\. (.*)$", text, re.S) if kind in ("para", "num") else None
            if m:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), m.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), m.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, PRODUCT, "Start here. Five tabs, about ten minutes to set up.", f"Checked {CHECKED}", span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "Most service owners divide the pay they want by the hours they work. That leaves out the hours nobody pays for "
    "(driving, quoting, invoicing), the costs that run every month, and self-employment tax. This workbook adds all three "
    "back, gives you the rate every quote should start from, prices any job in seconds, and shows which finished jobs "
    "really made money.", "para")])
r = sh_card(sh, r, "Five steps", [
    (1, "Your Numbers. Type your pay goal, the hours you really work and your monthly costs. Your target rate is the dark tile at the top.", "num"),
    (2, "Trade Presets. Change the hours, supplies and miles to match your real jobs, or add your own job types in the empty rows.", "num"),
    (3, "Quote Builder. Pick a job from the list and read your quote. Change any yellow cell for this one customer.", "num"),
    (4, "Price Sheet. Every job type priced at once. Print it or use it to build your service menu.", "num"),
    (5, "Job Tracker. Log finished jobs to see your real profit and what each one paid per hour.", "num"),
])
r = sh_card(sh, r, "How to read the cells", [
    ((52000, USD0), "Yellow cells with blue numbers are yours to change. Type over the example. Each one shows a hint when you select it.", "chip-in"),
    ((1404, INT), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((round(TARGET_X, 2), USD2), "Dark tiles are your answers. They sit at the top of Your Numbers and Quote Builder and stay on screen as you scroll.", "chip-key"),
    (None, "Each tab is protected without a password. To change anything else, use Review, Unprotect Sheet in Excel, or Data, "
           "Protect sheets and ranges in Google Sheets.", "note"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "The example", [(None,
    f"A {money(EX['take_home'], False)} take-home goal, with 15.3% self-employment tax and a 12% income tax estimate, needs "
    f"{money(_profit, False)} of profit before tax. "
    f"Add {money(_costs_y, False)} a year of business costs and you need {money(_profit + _costs_y, False)} of revenue. "
    f"At 48 weeks of 45 hours with 65% paid, that is {_paid:,.0f} paid hours: a break-even rate of {money(BE_X)} and a target "
    f"of {money(TARGET_X)} with a 15% margin. Pay divided by all hours would say {money(GUESS_X)}, "
    f"{money(TARGET_X - GUESS_X)} too low. A standard 2-bed clean then quotes at {money(QUOTE_X, False)}.", "para")])
r = sh_card(sh, r, "Good to know", [
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. Every formula is a plain "
           "formula that works in Excel, Google Sheets, Numbers and LibreOffice. No macros, no add-ons, no sign-up.", "para"),
    (None, "Presets. The trade times are starting points so you can see how it works, not market prices. Replace them with "
           "your own.", "para"),
])
r = sh_card(sh, r, "Before you rely on a number", [(None,
    "For general information and planning only. This is not legal, tax, financial or other professional advice, and using it "
    "doesn't create a professional relationship. Results are estimates based on the numbers you enter. Figures were checked "
    f"on {CHECKED} and can change. See the Terms tab before you rely on anything here.", "note")])
D.paint_canvas(sh, r - 1, "Z")
S.finish_sheet(sh, PRODUCT, r - 1, span=("A", "G"), tab_color="teal")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version 3. Read this before you rely on a number.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More from the shop", [
    (None, ("Cleaning Business Starter Kit: pricing calculator, client intake, checklists", "https://www.etsy.com/listing/4588330888"), "link"),
    (None, ("Hourly Rate Quick Calculator: the one-tab version", "https://www.etsy.com/listing/4589695837"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this calculator helped you price your work, a short review on Etsy helps other service owners find it. Thank you.", "para"),
    (None, "If a file won't open or a number looks wrong, message the shop on Etsy and it will be fixed.", "para"),
])
D.paint_canvas(tm, r - 1, "Z")
S.finish_sheet(tm, PRODUCT, r - 1, span=("A", "G"), tab_color="note")

wb.active = 0
problems = [p for p in S.audit(wb) if "input without validation" not in p]
if problems:
    print("AUDIT:", *problems, sep="\n  ")
wb.save(OUT)
print("saved", OUT)
print({"YN_target": f"I{b_tg}", "YN_be": f"I{b_be}", "TH": f"D{TH}", "quote": f"I{c_q}", "job": f"D{JOB_ROW}",
       "tp_first": tp_first, "ps_first": ps_first, "jt_first": jt_first, "rows": dict(PH=PH, CM=CM, CY=CY, SE=SE, SB=SB, IT=IT, MG=MG,
       CPM=CPM, SPD=SPD, OH=OH, OP=OP, OS=OS, OM=OM, AO=AO, DP=DP, CP=CP, c_flr=c_flr, c_dep=c_dep, c_pf=c_pf, c_eff=c_eff,
       c_now=c_now, pill=pill, b_profit=b_profit, b_rev=b_rev, b_ph=b_ph, cost_rows=cost_rows, WK=WK, HW=HW, PS=PS)})
print("example:", round(BE_X, 2), round(TARGET_X, 2), round(GUESS_X, 2), round(FLOOR_X, 2), QUOTE_X)
