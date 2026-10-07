#!/usr/bin/env python3
"""#9 STR Nightly Pricing Worksheet, version 3 (dashboard design, Oct 6, 2026).

Same inputs, starter curves, formulas and worked example (the fictional Ridgeline Cabin) as the
live version 1 (Sep 29, 2026), rebuilt to design/WORKBOOK_STANDARD.md section 0.
tools/compare_xlsx.py with compare_map.json proves the answers match.

Restructured for the dashboard: the weekly and monthly discounts and the floor moved from the
discounts tab to 1 Your Property (one place for every setting), and the five starter curves moved
from Setup to 2 Seasonality, next to the history they stand in for. Every formula is unchanged.

Fixes in v3 (declared under "intended" in compare_map.json):
  - 5 Fee Reset: with the old rate blank the result sentence said "Raise to $0 to keep $0.00", and
    the reset status said RESET with nothing typed. Both now wait for the numbers.
  - The fee note now matches Airbnb's own pages (Sep 15, 2026 price deadline outside the EEA and
    Switzerland, Oct 13, 2026 inside), checked Oct 6, 2026.
  - 2 Seasonality: with every starter cell and the history cleared, v1 showed #DIV/0! in the
    multipliers, the curve averages and the Rate Table. v3 shows blanks.
  - "What to try" on the old Instructions tab said a $140 floor spreads the weekly breaches to
    March; the file shows May. The new text is computed from the workbook's own maths.
  - Removed: a third-party fee example, an unverifiable forum story, and a mention of a chart the
    file never had. Added: Start Here, full Terms of Use tab, the short notice and tier clauses.

Usage: python3 engines/str-nightly-pricing/build_xlsx.py out.xlsx
"""
import math
import os
import re
import sys
import datetime as dt

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "tools"))
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.cell_range import MultiCellRange  # noqa: E402

PRODUCT = "STR Nightly Pricing Worksheet"
VERSION = "3"
CHECKED = "Oct 6, 2026"
AS_OF = f"Figures checked {CHECKED}"
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else "STR-Nightly-Pricing-Worksheet.xlsx"
USD0, USD2, PCT, INT, DATE = (S.FMT[k] for k in ("usd0", "usd2", "pct", "int", "date"))
PCT1 = "0.0%"
MULT = '0.00"x"'
HIST = "#,##0.##"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
MARKETS = ["beach", "mountain", "urban", "lake", "rural"]
MARKET_LIST = '{"beach","mountain","urban","lake","rural"}'

# ---------------------------------------------------------------- example data (v1, unchanged)
EX = dict(name="Ridgeline Cabin", base=185, cleaning=120, min_stay=2, market="mountain",
          weekly=0.10, monthly=0.20, floor=115,
          fee_old_rate=185, fee_old=0.03, fee_new=0.155, booked_nights=200)
CURVES = {  # one multiplier a month per market, each averaging 1.00 (v1 Setup tab)
    "beach":    [0.70, 0.75, 0.90, 0.95, 1.05, 1.30, 1.45, 1.35, 1.05, 0.85, 0.75, 0.90],
    "mountain": [1.15, 1.20, 0.90, 0.70, 0.75, 1.05, 1.25, 1.20, 1.00, 1.05, 0.60, 1.15],
    "urban":    [0.85, 0.85, 1.00, 1.05, 1.10, 1.10, 1.00, 0.95, 1.10, 1.15, 0.95, 0.90],
    "lake":     [0.60, 0.60, 0.70, 0.85, 1.10, 1.40, 1.55, 1.45, 1.15, 0.90, 0.80, 0.90],
    "rural":    [0.75, 0.75, 0.85, 1.00, 1.10, 1.20, 1.25, 1.20, 1.15, 1.15, 0.80, 0.80],
}
HISTORY = [4800, 5200, 3900, 2600, 3100, 4400, 5600, 5300, 4300, 4700, 2400, 5000]
EVENTS = [("Summer Bluegrass Festival", dt.datetime(2027, 7, 16), dt.datetime(2027, 7, 18), 0.40, None),
          ("Fall Foliage Marathon", dt.datetime(2027, 10, 9), dt.datetime(2027, 10, 10), 0.30, None),
          ("Holiday Week", dt.datetime(2026, 12, 26), dt.datetime(2026, 12, 31), 0.25, 5)]
N_EVENTS = 10


# ---------------------------------------------------------------- the example's answers, the workbook's own way
def model(base=EX["base"], floor=EX["floor"], weekly=EX["weekly"], monthly=EX["monthly"], hist=HISTORY):
    avg = sum(hist) / 12
    mult = [v / avg for v in hist]
    basen = [int(math.floor(base * m + 0.5)) for m in mult]          # ROUND(x,0) for positive x
    night = [max(b, floor) for b in basen]
    at_floor = sum(1 for b in basen if b < floor)
    wk = [n * (1 - weekly) for n in night]
    mo = [n * (1 - monthly) for n in night]
    wk_breach = [MONTHS[i] for i, v in enumerate(wk) if v < floor]
    mo_breach = [MONTHS[i] for i, v in enumerate(mo) if v < floor]
    return dict(mult=mult, base=basen, night=night, at_floor=at_floor, wk=wk, mo=mo,
                wk_breach=wk_breach, mo_breach=mo_breach)


M0 = model()
PEAK_I = M0["night"].index(max(M0["night"]))
LOW_I = M0["night"].index(min(M0["night"]))
EV_RATES = []
for name, d1, d2, prem, booked in EVENTS:
    rate = M0["night"][d1.month - 1]
    ev = int(math.floor(rate * (1 + prem) + 0.5))
    nights = (d2 - d1).days + 1
    priced = nights if booked is None else min(booked, nights)
    EV_RATES.append(dict(name=name, rate=rate, ev=ev, up=ev - rate, nights=nights, priced=priced, total=(ev - rate) * priced))
EV_TOTAL = sum(e["total"] for e in EV_RATES)
FEE_OLD_PAY = EX["fee_old_rate"] * (1 - EX["fee_old"])
FEE_NOW_PAY = EX["fee_old_rate"] * (1 - EX["fee_new"])
FEE_EXACT = FEE_OLD_PAY / (1 - EX["fee_new"])
FEE_SET = math.ceil(FEE_EXACT - 1e-9)
FEE_LOST_Y = (FEE_OLD_PAY - FEE_NOW_PAY) * EX["booked_nights"]
FEE_SHARE = (FEE_OLD_PAY - FEE_NOW_PAY) / FEE_OLD_PAY
M140 = model(floor=140)
M_RESET = model(base=FEE_SET)
FEST2 = EV_RATES[0]["up"] * 2


def money(x, cents=False):
    return f"${x + 1e-9:,.2f}" if cents else f"${x + 1e-9:,.0f}"


def joined(items):
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


wb = S.new_workbook(PRODUCT, version=VERSION)
GRID = {"A": 3, "B": 2, "C": 31, "D": 14, "E": 2, "F": 3, "G": 2, "H": 24, "I": 12, "J": 17, "K": 2, "L": 3}


def notes_card(ws, top, notes, span=("C", "J"), frame_span=("B", "K")):
    r = top
    D.h(ws, r, 12); r += 1
    t = ws[f"{span[0]}{r}"]; t.value = "Notes and sources"; t.font = D.f(12, True)
    ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}"); D.h(ws, r, 24); r += 1
    D.h(ws, r, 6); r += 1
    width = sum(S.width_of(ws, k) for k in D.cols(*span))
    for head, body in notes:
        cell = ws[f"{span[0]}{r}"]
        cell.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=9, b=True, color=D.T["ink"]), head + ". "),
                                  TextBlock(InlineFont(rFont=D.FONT, sz=9, color=D.T["ink2"]), body))
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(f"{span[0]}{r}:{span[1]}{r}")
        D.h(ws, r, S.text_height(head + ". " + body, width, 9) + 4); r += 1
    D.h(ws, r, 12)
    D.frame(ws, top, r, *frame_span)
    return r


def total_row(ws, row, span, label_col, label, cells):
    """Turn the last body row of a table card into its total row: ink rule on top, bold, locked, no inputs."""
    for k in span:
        c = ws[f"{k}{row}"]
        c.value = None
        c.fill = D.fill("total_tint")
        c.font = D.f(10, True)
        c.protection = c.protection.copy(locked=True)
        c.border = Border(top=D.side("ink"), bottom=D.side("hair"))
        c.alignment = Alignment(horizontal="right", vertical="center")
    lab = ws[f"{label_col}{row}"]; lab.value = label; lab.alignment = Alignment(horizontal="left", vertical="center")
    for col, (formula, fmt) in cells.items():
        c = ws[f"{col}{row}"]; c.value = formula; c.number_format = fmt


def trim_validations(ws, last_body):
    """table_card validates every body row; the total row is not an input, so cut the ranges back."""
    for dv in ws.data_validations.dataValidation:
        new = []
        for rng in str(dv.sqref).split():
            a, _, b = rng.partition(":")
            if b:
                col = "".join(ch for ch in b if ch.isalpha())
                row = int("".join(ch for ch in b if ch.isdigit()))
                if row > last_body:
                    b = f"{col}{last_body}"
                new.append(f"{a}:{b}")
            else:
                new.append(rng)
        dv.sqref = MultiCellRange(" ".join(new))


def finish(ws, last, span, freeze, tab_color):
    for rr in range(1, last + 1):
        if ws.row_dimensions[rr].height is None:
            D.h(ws, rr, D.GRID_ROW)
    D.paint_canvas(ws, last, "Z")
    S.finish_sheet(ws, PRODUCT, last, span=span, freeze=freeze, tab_color=tab_color)


# =====================================================================  1 Your Property
yp = wb.create_sheet("1 Your Property")
S.set_widths(yp, GRID)
D.page_header(yp, "Your Property",
              "Your listing, base rate and floor. Every other tab prices the year from these.", AS_OF)
TOP = 11
A = D.Card(yp, TOP, "B", "C", "D", "E", wb=wb)
A.title("Your property", "Type over the yellow cells. Every tab updates.")
A.eyebrow("THE LISTING")
PN = A.input("Property name", EX["name"], "@", name="PropertyName", validation=("textLength", 0, 40),
             prompt="Shows on the 6 Rate Table title. Up to 40 characters. Example: Ridgeline Cabin.", prompt_title="Property name",
             allow_blank=True)
yp[f"D{PN}"].alignment = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
yp[f"D{PN}"].font = D.f(9, True, "input_text")
for _dv in yp.data_validations.dataValidation:
    if str(_dv.sqref) == f"D{PN}":
        _dv.error = "Up to 40 characters."
BR = A.input("Base nightly rate", EX["base"], USD0, name="BaseRate", validation=("decimal", 0, None),
             prompt="Your average-month nightly rate, in dollars. Every month is priced from it. Example: 185.")
CL = A.input("Cleaning fee", EX["cleaning"], USD0, name="CleaningFee", validation=("decimal", 0, None),
             prompt="In dollars. Not part of the nightly maths; it prints on the 6 Rate Table. Example: 120.")
MS = A.input("Minimum stay, nights", EX["min_stay"], INT, name="MinimumStay", validation=("whole", 1, 30),
             prompt="Whole nights, 1 to 30. Prints on the 6 Rate Table. Example: 2.", prompt_title="Minimum stay")
MT = A.input("Market type", EX["market"], "@", name="MarketType", validation=("list", MARKETS),
             prompt="Pick beach, mountain, urban, lake or rural. It loads that starter curve on 2 Seasonality.")
yp[f"D{MT}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
A.eyebrow("LONG-STAY DISCOUNTS AND YOUR FLOOR")
WD = A.input("Weekly discount, 7+ nights", EX["weekly"], PCT1, name="WeeklyDiscount", validation=("decimal", 0, 0.9),
             prompt="Type a percent, 0% to 90%, for example 10%. Applied per night on stays of 7 or more nights. 0% turns it off.",
             prompt_title="Weekly discount")
MD = A.input("Monthly discount, 28+ nights", EX["monthly"], PCT1, name="MonthlyDiscount", validation=("decimal", 0, 0.9),
             prompt="Type a percent, 0% to 90%, for example 20%. Applied per night on stays of 28 or more nights.",
             prompt_title="Monthly discount")
FL = A.input("Floor per night", EX["floor"], USD0, name="FloorRate", validation=("decimal", 0, None),
             prompt="The nightly rate below which a booking isn't worth having (cleaning, utilities, wear, your time). "
                    "No month and no discount prices below it. Example: 115.")
A.close()
P = "'1 Your Property'!"
BRr, FLr, WDr, MDr, MTr, PNr, CLr, MSr = (f"{P}$D${x}" for x in (BR, FL, WD, MD, MT, PN, CL, MS))

# =====================================================================  2 Seasonality
se = wb.create_sheet("2 Seasonality")
S.set_widths(se, {"A": 3, "B": 2, "C": 13, "D": 9, "E": 10, "F": 9, "G": 9, "H": 9, "I": 14, "J": 12, "K": 11,
                  "L": 11, "M": 13, "N": 13, "O": 2, "P": 3})
D.page_header(se, "Seasonality", "Your last 12 months, or a starter curve, turned into a multiplier and a base rate per month.",
              AS_OF, span=("B", "O"), side_cols=3)
SE_TOP = 11
se_first = SE_TOP + 5
se_last = se_first + 11
se_avg = se_last + 1
R = {k: f"${k}${se_first}:${k}${se_last}" for k in "CDEFGHIJKLMN"}
STATUS = "$C$8"   # the stat strip's curve-in-use cell (written below)


def se_formula(kind):
    def fn(r):
        if kind == "starter":
            return f"=IFERROR(INDEX(D{r}:H{r},MATCH({MTr},{MARKET_LIST},0)),\"\")"
        if kind == "used":
            return f'=IF({STATUS}="YOURS",I{r},J{r})'
        if kind == "mult":
            return f'=IFERROR(IF(AND(ISNUMBER(K{r}),COUNT({R["K"]})=12,AVERAGE({R["K"]})>0),K{r}/AVERAGE({R["K"]}),""),"")'
        if kind == "base":
            return f'=IF(ISNUMBER(L{r}),ROUND({BRr}*L{r},0),"")'
        if kind == "bar":
            return f'=IF(ISNUMBER(M{r}),IFERROR(REPT("{D.BAR_CHAR}",MAX(0,ROUND(M{r}/MAX({R["M"]})*10,0))),""),"")'
    return fn


cols_se = [dict(col="C", head="Month", kind="text", values=MONTHS + ["Average"])]
for col, mk in zip("DEFGH", MARKETS):
    cols_se.append(dict(col=col, head=mk.capitalize(), kind="input", fmt=MULT, align="right", values=CURVES[mk] + [None],
                        validation=("decimal", 0.1, 3),
                        prompt=f"Starter multiplier for the {mk} market, 0.1 to 3. Keep the column's average near 1.00x.",
                        error="Enter a multiplier from 0.1 to 3, for example 1.15."))
cols_se += [
    dict(col="I", head="Your last 12 months", kind="input", fmt=HIST, align="right", values=HISTORY + [None],
         validation=("decimal", 0, None),
         prompt="Nightly revenue in dollars, or occupancy as a number, for that month. Fill all 12 to use your own curve; "
                "a month with no bookings needs a small number, not 0.",
         error="Enter 0 or more, for example 4800."),
    dict(col="J", head="Starter, your market", kind="calc", fmt=MULT, align="right", formula=se_formula("starter")),
    dict(col="K", head="Value used", kind="calc", fmt=HIST, align="right", formula=se_formula("used")),
    dict(col="L", head="Multiplier", kind="calc", fmt=MULT, align="right", formula=se_formula("mult")),
    dict(col="M", head="Base nightly rate", kind="calc", fmt=USD0, align="right", bold=True, formula=se_formula("base")),
    dict(col="N", head="", kind="calc", formula=se_formula("bar")),
]
f1, l1, se_end = D.table_card(se, SE_TOP, "B", "O", cols_se, 13, title="Your curve",
                              sub="placeholder")
assert (f1, l1) == (se_first, se_avg), (f1, l1)
for r in range(se_first, se_last + 1):
    se[f"N{r}"].font = D.f(9, False, "teal"); se[f"N{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
total_row(se, se_avg, D.cols("C", "N"), "C", "Average", {
    **{c: (f'=IFERROR(AVERAGE({c}{se_first}:{c}{se_last}),"")', MULT) for c in "DEFGH"},
    "I": (f'=IF(COUNT(I{se_first}:I{se_last})>0,AVERAGE(I{se_first}:I{se_last}),"")', HIST),
    "J": (f'=IFERROR(AVERAGE(J{se_first}:J{se_last}),"")', MULT),
    "K": (f'=IF(COUNT(K{se_first}:K{se_last})=12,AVERAGE(K{se_first}:K{se_last}),"")', HIST),
    "L": (f'=IF(COUNT(L{se_first}:L{se_last})=12,AVERAGE(L{se_first}:L{se_last}),"")', MULT),
    "M": (f'=IF(COUNT(M{se_first}:M{se_last})=12,AVERAGE(M{se_first}:M{se_last}),"")', USD0),
})
trim_validations(se, se_last)
# the sub line says which curve is in use and what to do next
sub_row = SE_TOP + 2
se[f"C{sub_row}"].value = (f'=IF({STATUS}="YOURS","Your last 12 months set the shape of the year. Clear column I to go back to the starter curve.",'
                           f'IF({STATUS}="STARTER","Starter curve for "&IF({MTr}="","no market type yet",{MTr})&". Type all 12 months of your own history in the yellow column I to replace it.",'
                           f'"Some months in column I are blank or 0, so the starter curve is in use. Fill all 12 (a quiet month still needs a small number) or clear the column."))')
# colours: market averages more than 2% off 1.00x, multipliers 15% above or below average
se.conditional_formatting.add(f"D{se_avg}:H{se_avg}", FormulaRule(formula=[f"ABS(D{se_avg}-1)>0.02"],
                              font=Font(color=D.T["accent"], bold=True), fill=PatternFill("solid", bgColor=D.T["accent_tint"])))
se.conditional_formatting.add(f"L{se_first}:L{se_last}", FormulaRule(formula=[f'AND(ISNUMBER(L{se_first}),L{se_first}>1.15)'],
                              font=Font(color=D.T["teal"], bold=True), fill=PatternFill("solid", bgColor=D.T["teal_tint"])))
se.conditional_formatting.add(f"L{se_first}:L{se_last}", FormulaRule(formula=[f'AND(ISNUMBER(L{se_first}),L{se_first}<0.85)'],
                              font=Font(color=D.T["accent"], bold=True), fill=PatternFill("solid", bgColor=D.T["accent_tint"])))
# stat strip, rows 6 to 9
D.h(se, 6, 10)
D.stat_strip(se, 7, [
    ("C:E", "Curve in use", f'=IF(COUNT({R["I"]})=0,"STARTER",IF(AND(COUNT({R["I"]})=12,MIN({R["I"]})>0),"YOURS","INCOMPLETE"))', "@"),
    ("F:I", "Busiest month", f'=IF(COUNT({R["L"]})=12,INDEX({R["C"]},MATCH(MAX({R["L"]}),{R["L"]},0))&"  "&TEXT(MAX({R["L"]}),"0.00")&"x","-")', "@"),
    ("J:L", "Quietest month", f'=IF(COUNT({R["L"]})=12,INDEX({R["C"]},MATCH(MIN({R["L"]}),{R["L"]},0))&"  "&TEXT(MIN({R["L"]}),"0.00")&"x","-")', "@"),
    ("M:N", "Base nightly rate", f"={BRr}", USD0),
])
for word, color in (("YOURS", "teal"), ("STARTER", "ink"), ("INCOMPLETE", "accent")):
    se.conditional_formatting.add("C8", FormulaRule(formula=[f'$C$8="{word}"'], font=Font(color=D.T[color], bold=True), stopIfTrue=True))
D.h(se, 9, 10)
D.frame(se, 6, 9, "B", "O")
D.h(se, 10, 14)
SE_NOTES = [
    ("Revenue or occupancy", "Either works: only the shape matters, because each month is divided by the 12-month average. "
     "Revenue is the better signal if you already priced by season; occupancy is better if you charged one flat rate all year. "
     "Your platform's earnings report by month has both."),
    ("Starter curves", "Generic shapes for a market type, not your town and not market data. Edit any cell if you know better, "
     "and keep each column's average near 1.00x: the average turns orange when it is more than 2% off, because a column above "
     "1.00x quietly raises every rate and one below lowers them. Beach and lake markets swing hardest; urban is flattest."),
    ("Colours", "Multipliers 15% or more above the average are green, 15% or more below are orange."),
]
se_n_end = notes_card(se, se_end + 2, SE_NOTES, span=("C", "N"), frame_span=("B", "O"))
S.page_break_before(se, se_end + 2)
finish(se, se_n_end + 1, ("A", "P"), f"A{se_first}", "teal")
se.page_setup.orientation = "landscape"
SE = "'2 Seasonality'!"

# =====================================================================  4 Stay Discounts (built before 3 Events, which reads it)
sd = wb.create_sheet("4 Stay Discounts")
S.set_widths(sd, {"A": 3, "B": 2, "C": 14, "D": 11, "E": 14, "F": 13, "G": 12, "H": 11, "I": 17, "J": 12, "K": 12,
                  "L": 11, "M": 17, "N": 12, "O": 2, "P": 3})
D.page_header(sd, "Stay Discounts", "Weekly and monthly discounts on every month's rate, with a floor they can't go under.",
              AS_OF, span=("B", "O"), side_cols=3)
SD_TOP = 11
sd_first = SD_TOP + 5
sd_last = sd_first + 11
sd_tot = sd_last + 1


def sd_formula(kind):
    def fn(r):
        i = r - sd_first + se_first
        return {
            "base": f'=IF({SE}M{i}="","",{SE}M{i})',
            "night": f'=IF(ISNUMBER(D{r}),MAX(D{r},{FLr}),"")',
            "vs": f'=IF(ISNUMBER(D{r}),IF(D{r}<{FLr},"AT FLOOR","ok"),"")',
            "wk": f'=IF(ISNUMBER(E{r}),E{r}*(1-{WDr}),"")',
            "wkb": f'=IF(ISNUMBER(G{r}),MAX(G{r},{FLr}),"")',
            "wkf": f'=IF(ISNUMBER(G{r}),IF(G{r}<{FLr},"FLOOR BREACH","ok"),"")',
            "wkt": f'=IF(ISNUMBER(H{r}),H{r}*7,"")',
            "mo": f'=IF(ISNUMBER(E{r}),E{r}*(1-{MDr}),"")',
            "mob": f'=IF(ISNUMBER(K{r}),MAX(K{r},{FLr}),"")',
            "mof": f'=IF(ISNUMBER(K{r}),IF(K{r}<{FLr},"FLOOR BREACH","ok"),"")',
            "mot": f'=IF(ISNUMBER(L{r}),L{r}*28,"")',
        }[kind]
    return fn


cols_sd = [
    dict(col="C", head="Month", kind="text", values=MONTHS + [""]),
    dict(col="D", head="Base nightly rate", kind="calc", fmt=USD0, align="right", formula=sd_formula("base")),
    dict(col="E", head="Nightly rate, floor applied", kind="calc", fmt=USD0, align="right", bold=True, formula=sd_formula("night")),
    dict(col="F", head="Base vs floor", kind="calc", align="left", formula=sd_formula("vs")),
    dict(col="G", head="7+ nights, discounted", kind="calc", fmt=USD2, align="right", formula=sd_formula("wk")),
    dict(col="H", head="7+ nights, billed", kind="calc", fmt=USD2, align="right", bold=True, formula=sd_formula("wkb")),
    dict(col="I", head="7-night check", kind="calc", align="left", formula=sd_formula("wkf")),
    dict(col="J", head="7-night stay total", kind="calc", fmt=USD2, align="right", formula=sd_formula("wkt")),
    dict(col="K", head="28+ nights, discounted", kind="calc", fmt=USD2, align="right", formula=sd_formula("mo")),
    dict(col="L", head="28+ nights, billed", kind="calc", fmt=USD2, align="right", bold=True, formula=sd_formula("mob")),
    dict(col="M", head="28-night check", kind="calc", align="left", formula=sd_formula("mof")),
    dict(col="N", head="28-night stay total", kind="calc", fmt=USD2, align="right", formula=sd_formula("mot")),
]
f1, l1, sd_end = D.table_card(sd, SD_TOP, "B", "O", cols_sd, 13, title="Every month, every stay length",
                              sub="placeholder")
assert (f1, l1) == (sd_first, sd_tot), (f1, l1)
total_row(sd, sd_tot, D.cols("C", "N"), "C", "Times at floor", {
    "F": (f'=COUNTIF(F{sd_first}:F{sd_last},"AT FLOOR")', INT),
    "I": (f'=COUNTIF(I{sd_first}:I{sd_last},"FLOOR BREACH")', INT),
    "M": (f'=COUNTIF(M{sd_first}:M{sd_last},"FLOOR BREACH")', INT),
})
for c in "FIM":
    sd[f"{c}{sd_tot}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for r in range(sd_first, sd_last + 1):
        sd[f"{c}{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
sd[f"C{SD_TOP + 2}"].value = (f'="Weekly "&TEXT({WDr},"0%")&" off, monthly "&TEXT({MDr},"0%")&" off, floor "&TEXT({FLr},"$#,##0")&'
                              f'" a night. Change them on 1 Your Property. Billed is what to expect per night; stay totals are before the cleaning fee."')
for col, word, color, tint in (("F", "AT FLOOR", "gold", "input_fill"), ("I", "FLOOR BREACH", "accent", "accent_tint"),
                               ("M", "FLOOR BREACH", "accent", "accent_tint")):
    sd.conditional_formatting.add(f"{col}{sd_first}:{col}{sd_last}", FormulaRule(
        formula=[f'${col}{sd_first}="{word}"'], font=Font(color=D.T[color], bold=True), fill=PatternFill("solid", bgColor=D.T[tint]), stopIfTrue=True))
    sd.conditional_formatting.add(f"{col}{sd_first}:{col}{sd_last}", FormulaRule(
        formula=[f'${col}{sd_first}="ok"'], font=Font(color=D.T["teal"]), stopIfTrue=True))
D.h(sd, 6, 10)
D.stat_strip(sd, 7, [
    ("C:E", "Months held at the floor", f"=F{sd_tot}", INT),
    ("F:H", "Floor breaches, 7+ nights", f"=I{sd_tot}", INT),
    ("I:K", "Floor breaches, 28+ nights", f"=M{sd_tot}", INT),
    ("L:N", "Your floor per night", f"={FLr}", USD0),
])
for c in ("C8", "F8", "I8"):
    D.status_rule(sd, c, f"{c}>0", bad="accent", good="teal", fill_tint=False)
D.h(sd, 9, 10)
D.frame(sd, 6, 9, "B", "O")
D.h(sd, 10, 14)
SD_NOTES = [
    ("How the floor works", "Each month's nightly rate is the higher of its base rate and your floor (AT FLOOR when the floor wins). "
     "Each discounted rate is checked against the floor too: where a discount would go under it, the floor is charged and the "
     "check reads FLOOR BREACH. The usual way to lose money on a long stay is a weekly discount stacked on a rate that was already "
     "too low; the floor stops that."),
    ("Platform rules", "Airbnb applies a weekly discount to stays of 7 or more nights and a monthly discount to stays of 28 or more "
     f"nights, on the nightly price (Airbnb Help Center, Weekly and monthly discounts for home stays, checked {CHECKED}). "
     "Check your own platform's settings."),
]
sd_n_end = notes_card(sd, sd_end + 2, SD_NOTES, span=("C", "N"), frame_span=("B", "O"))
S.page_break_before(sd, sd_end + 2)
finish(sd, sd_n_end + 1, ("A", "P"), f"A{sd_first}", "blue")
sd.page_setup.orientation = "landscape"
SD = "'4 Stay Discounts'!"

# =====================================================================  3 Events
ev = wb.create_sheet("3 Events", 2)
S.set_widths(ev, {"A": 3, "B": 2, "C": 26, "D": 13, "E": 13, "F": 9, "G": 10, "H": 13, "I": 8, "J": 13, "K": 13,
                  "L": 10, "M": 9, "N": 12, "O": 2, "P": 3})
D.page_header(ev, "Events", "One row per local event: a premium on that month's nightly rate for the nights in the window.",
              AS_OF, span=("B", "O"), side_cols=3)
EV_TOP = 11
ev_first = EV_TOP + 5
ev_last = ev_first + N_EVENTS - 1
ev_tot = ev_last + 1
SDN = f"{SD}$E${sd_first}:$E${sd_last}"


def ev_formula(kind):
    def fn(r):
        return {
            "nights": f'=IF(C{r}="","",IF(OR(D{r}="",E{r}=""),"",IF(E{r}<D{r},"CHECK DATES",E{r}-D{r}+1)))',
            "month": f'=IF(C{r}="","",IF(D{r}="","",MONTH(D{r})))',
            "rate": f'=IF(C{r}="","",IF(ISNUMBER(I{r}),INDEX({SDN},I{r}),""))',
            "ev": f'=IF(C{r}="","",IF(AND(ISNUMBER(J{r}),ISNUMBER(G{r})),ROUND(J{r}*(1+G{r}),0),""))',
            "up": f'=IF(C{r}="","",IF(ISNUMBER(K{r}),K{r}-J{r},""))',
            "priced": f'=IF(C{r}="","",IF(ISNUMBER(F{r}),IF(H{r}="",F{r},MIN(H{r},F{r})),""))',
            "total": f'=IF(C{r}="","",IF(AND(ISNUMBER(L{r}),ISNUMBER(M{r})),L{r}*M{r},""))',
        }[kind]
    return fn


def pad(vals, n):
    return list(vals) + [None] * (n - len(vals))


cols_ev = [
    dict(col="C", head="Event", kind="input", values=pad([e[0] for e in EVENTS], N_EVENTS + 1), validation=("textLength", 0, 40),
         prompt="A short name: festival, race, holiday week. The row prices once a name is in. Up to 40 characters.",
         error="Up to 40 characters."),
    dict(col="D", head="First night", kind="input", fmt=DATE, values=pad([e[1] for e in EVENTS], N_EVENTS + 1),
         validation=("date", "DATE(2020,1,1)", "DATE(2100,12,31)"), prompt="The first night of the event window, for example 7/16/2027.",
         error="Enter a date from 2020 on, for example 7/16/2027."),
    dict(col="E", head="Last night", kind="input", fmt=DATE, values=pad([e[2] for e in EVENTS], N_EVENTS + 1),
         validation=("date", "DATE(2020,1,1)", "DATE(2100,12,31)"), prompt="The last night of the window, on or after the first night.",
         error="Enter a date from 2020 on, on or after the first night."),
    dict(col="F", head="Nights in window", kind="calc", fmt=INT, align="right", formula=ev_formula("nights")),
    dict(col="G", head="Premium", kind="input", fmt=PCT1, align="right", values=pad([e[3] for e in EVENTS], N_EVENTS + 1),
         validation=("decimal", -0.5, 3), prompt="Type a percent, for example 40%. Start near 25% for a holiday week and 40% for "
                                                 "a festival that fills the town. -50% to 300%.",
         error="Enter a percent from -50% to 300%, for example 40%."),
    dict(col="H", head="Expected booked nights", kind="input", fmt=INT, align="right", values=pad([e[4] for e in EVENTS], N_EVENTS + 1),
         validation=("whole", 0, 366), prompt="Optional, whole nights. Blank prices every night in the window.",
         error="Enter whole nights from 0 to 366, or leave it blank."),
    dict(col="I", head="Month", kind="calc", fmt=INT, align="right", formula=ev_formula("month")),
    dict(col="J", head="Rate that month", kind="calc", fmt=USD0, align="right", formula=ev_formula("rate")),
    dict(col="K", head="Event nightly rate", kind="calc", fmt=USD0, align="right", bold=True, formula=ev_formula("ev")),
    dict(col="L", head="Uplift per night", kind="calc", fmt=USD0, align="right", formula=ev_formula("up")),
    dict(col="M", head="Nights priced", kind="calc", fmt=INT, align="right", formula=ev_formula("priced")),
    dict(col="N", head="Uplift for the window", kind="calc", fmt=USD0, align="right", formula=ev_formula("total")),
]
f1, l1, ev_end = D.table_card(ev, EV_TOP, "B", "O", cols_ev, N_EVENTS + 1, title="Your event windows",
                              sub="An event is priced at the rate of the month its first night falls in, floor applied. "
                                  "The uplift is what the premium adds if those nights book: built from your inputs, not a forecast.")
assert (f1, l1) == (ev_first, ev_tot), (f1, l1)
total_row(ev, ev_tot, D.cols("C", "N"), "C", "All events, if those nights book", {
    "M": (f"=SUM(M{ev_first}:M{ev_last})", INT),
    "N": (f"=SUM(N{ev_first}:N{ev_last})", USD0),
})
trim_validations(ev, ev_last)
for r in range(ev_first, ev_last + 1):
    ev[f"C{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ev.conditional_formatting.add(f"F{ev_first}:F{ev_last}", FormulaRule(formula=[f'$F{ev_first}="CHECK DATES"'],
                              font=Font(color=D.T["accent"], bold=True), fill=PatternFill("solid", bgColor=D.T["accent_tint"])))
D.h(ev, 6, 10)
EK = f"K{ev_first}:K{ev_last}"
D.stat_strip(ev, 7, [
    ("C:C", "Events entered", f'=COUNTIF(C{ev_first}:C{ev_last},"?*")', INT),
    ("D:G", "Event nights priced", f"=M{ev_tot}", INT),
    ("H:K", "Top event nightly rate", f'=IF(COUNT({EK})=0,"-",MAX({EK}))', USD0),
    ("L:N", "Event uplift for the year", f"=N{ev_tot}", USD0),
])
D.h(ev, 9, 10)
D.frame(ev, 6, 9, "B", "O")
D.h(ev, 10, 14)
EV_NOTES = [
    ("Where a premium comes from", "What the same dates booked at last year, or the going rate when the event's dates are searched a "
     "few months out. Start near 25% for a holiday week and 40% for a festival that fills the town, then let last year's "
     "bookings correct you. The three example events are fictional."),
    ("Expected booked nights", "Leave it blank and every night in the window is priced at the event rate. Type a number and only "
     "those nights count toward the uplift. A window whose last night is before its first night shows CHECK DATES."),
    ("More than 10 events", "Unprotect the tab (Review, Unprotect Sheet in Excel), copy the last event row with its formulas, "
     "and insert it above the total row; do the same on 6 Rate Table's event windows."),
]
ev_n_end = notes_card(ev, ev_end + 2, EV_NOTES, span=("C", "N"), frame_span=("B", "O"))
S.page_break_before(ev, ev_end + 2)
finish(ev, ev_n_end + 1, ("A", "P"), f"A{ev_first}", "purple")
ev.page_setup.orientation = "landscape"
EV = "'3 Events'!"

# =====================================================================  5 Fee Reset
fr = wb.create_sheet("5 Fee Reset")
S.set_widths(fr, GRID)
D.page_header(fr, "Fee Reset", "What the 15.5% single host fee does to your payout, and the rate that keeps it whole.", AS_OF)
L1 = D.Card(fr, TOP, "B", "C", "D", "E", wb=wb)
L1.title("Your fee change", "Type over the yellow cells. Select one to see its source.")
L1.eyebrow("BEFORE AND AFTER")
OR_ = L1.input("Nightly rate before the change", EX["fee_old_rate"], USD0, name="OldNightlyRate", validation=("decimal", 0, None),
               prompt="What you charged under the split fee, when guests paid their own service fee on top. Example: 185.",
               prompt_title="Rate before the change", allow_blank=True)
OF = L1.input("Host fee under the split fee", EX["fee_old"], PCT1, name="OldHostFee", validation=("decimal", 0, 0.5),
              prompt="Airbnb: most hosts paid 3%, some more; 4% in Brazil and Mexico (Help Center article 1857, checked "
                     + CHECKED + "). Check an old payout statement.", prompt_title="Old host fee")
NF = L1.input("Host fee now, single fee", EX["fee_new"], PCT1, name="NewHostFee", validation=("decimal", 0, 0.5),
              prompt="Airbnb's single fee: 15.5% for most hosts, 16% in Brazil and Mexico (article 1857, checked " + CHECKED
                     + "). Every number on this tab reads this cell.", prompt_title="Host fee now")
L1.eyebrow("FOR THE YEARLY FIGURE")
BN = L1.input("Booked nights per year", EX["booked_nights"], INT, name="BookedNights", validation=("whole", 0, 366),
              prompt="Whole nights, 0 to 366. Last year's count is fine. Used for the yearly figure only.",
              prompt_title="Booked nights per year")
L1.close()

R1 = D.Card(fr, TOP, "G", "H", "I", "K", extra="J", wb=wb)
R1.title("What it does to your payout", "Per night, and what to raise your rate to.")
R1.eyebrow("IF YOU CHANGE NOTHING")
p_old = R1.calc("Payout per night, old fee", f"=$D${OR_}*(1-$D${OF})", USD2)
p_now = R1.calc("Payout now, same price", f"=$D${OR_}*(1-$D${NF})", USD2)
p_lost = R1.calc("Payout lost per night", f"=I{p_old}-I{p_now}", USD2, total=True)
p_share = R1.calc("Share of payout lost", f"=IF(I{p_old}>0,I{p_lost}/I{p_old},0)", PCT1)
p_year = R1.calc("Payout lost per year", f"=I{p_lost}*$D${BN}", USD0, bold=True)
R1.eyebrow("WHAT TO RAISE IT TO")
p_exact = R1.calc("Rate that keeps old payout", f"=IF($D${NF}<1,I{p_old}/(1-$D${NF}),0)", USD2)
p_set = R1.calc("Set your nightly rate to", f"=ROUNDUP(I{p_exact},0)", USD0, total=True, tint="total_tint")
p_at = R1.calc("Payout at that rate", f"=I{p_set}*(1-$D${NF})", USD2)
p_guest = R1.calc("Guest paid before, with fee", f'=TEXT($D${OR_}*1.141,"$#,##0")&" to "&TEXT($D${OR_}*1.165,"$#,##0")', "@")
R1.close()

L2 = D.Card(fr, L1.end + 2, "B", "C", "D", "E", wb=wb)
L2.title("Is your base rate reset?", "Your base rate on 1 Your Property against the raise-to rate.")
k_base = L2.calc("Base nightly rate, 1 Your Property", f"={BRr}", USD0)
k_gap = L2.calc("Gap to the raise-to rate", f"=I{p_exact}-D{k_base}", USD2)
k_stat = L2.calc("Status", f'=IF(OR($D${OR_}="",{BRr}=""),"",IF(D{k_base}>=I{p_exact}-0.005,"RESET","NOT RESET"))', "@", bold=True)
L2.close()
D.status_rule(fr, f"D{k_stat}", f'$D${k_stat}="NOT RESET"', fill_tint=True)
# result sentence, full width under both columns
res_top = max(L2.end, R1.end) + 2
D.h(fr, res_top, 12)
t = fr[f"C{res_top + 1}"]; t.value = "Your result"; t.font = D.f(12, True); fr.merge_cells(f"C{res_top + 1}:J{res_top + 1}"); D.h(fr, res_top + 1, 24)
RES = res_top + 2
c = fr[f"C{RES}"]
c.value = (f'=IF($D${OR_}="","Type your nightly rate before the change to see your result.",'
           f'"Raise to "&TEXT(I{p_set},"$#,##0")&" to keep "&TEXT(I{p_old},"$#,##0.00")&" payout per night. Change nothing and you keep "'
           f'&TEXT(I{p_now},"$#,##0.00")&", which is "&TEXT(I{p_share},"0.0%")&" less on every booking.")')
c.font = D.f(11, True); c.alignment = Alignment(vertical="center", wrap_text=True)
fr.merge_cells(f"C{RES}:J{RES}"); D.h(fr, RES, 34)
D.h(fr, RES + 1, 12)
D.frame(fr, res_top, RES + 1, "B", "K")
k_note_top = RES + 3
FR_NOTES = [
    ("The fee change", "Airbnb is moving every home host to a single service fee paid by the host: 15.5% for most hosts and 16% for "
     "listings in Brazil and Mexico, where guests no longer pay a separate service fee. Under the old split fee most hosts paid 3% "
     "(4% in Brazil and Mexico) and guests paid 14.1% to 16.5% of the booking subtotal on top. Airbnb set the price-adjustment "
     "deadline at September 15, 2026 for hosts outside the European Economic Area, or October 13, 2026 for hosts in the EEA or Switzerland. "
     "Sources: Airbnb Help Center article 1857 (Airbnb service fees) and the Airbnb Resource Center article \"Simplifying service fees "
     f"on Airbnb\" (published Jul 7, 2026, updated Aug 24, 2026). Checked {CHECKED}."),
    ("Divide, don't add", "The raise-to rate is your old payout divided by (1 minus the new fee). Adding 15.5% to your price leaves "
     "you short. Airbnb's own example: a $100 price becomes $115 to keep about $97."),
    ("What guests see", "Under the split fee a guest paid 14.1% to 16.5% on top of your price, so the raise-to rate is close to what "
     "they already paid. The price they see barely moves."),
    ("Your old fee", "The old host fee varied by host, so the 3% here is a default, not a promise. Check an old payout statement."),
]
fr_n_end = notes_card(fr, k_note_top, FR_NOTES)
S.page_break_before(fr, k_note_top)
D.tile(fr, 6, "B", "E", "SET YOUR NIGHTLY RATE TO", f'=IF($D${OR_}="","",I{p_set})', USD0,
       f'=IF($D${OR_}="","Type your rate before the change",'
       f'"Keeps your "&TEXT(I{p_old},"$#,##0.00")&" payout per night at a "&TEXT($D${NF},"0.0%")&" fee")', dark=True)
v, sub, _ = D.tile(fr, 6, "G", "K", "PAYOUT LOST PER YEAR IF YOU CHANGE NOTHING", f'=IF($D${OR_}="","",I{p_year})', USD0,
                   f'=IF($D${OR_}="","Type your rate before the change",'
                   f'TEXT(I{p_lost},"$#,##0.00")&" a night, "&TEXT(I{p_share},"0.0%")&" of every payout")')
D.status_rule(fr, v.coordinate, f"I{p_year}>0", fill_tint=False)
D.h(fr, 10, 14)
finish(fr, fr_n_end + 1, ("A", "L"), "A11", "gold")
FR = "'5 Fee Reset'!"

# =====================================================================  6 Rate Table
rt = wb.create_sheet("6 Rate Table")
S.set_widths(rt, {"A": 3, "B": 2, "C": 26, "D": 13, "E": 12, "F": 14, "G": 12, "H": 16, "I": 16, "J": 12, "K": 2, "L": 3})
D.page_header(rt, "Nightly rates", "x", AS_OF, span=("B", "K"), side_cols=3)
rt["B2"].value = f'="Nightly rates: "&IF({PNr}="","your property",{PNr})'
rt["B4"].value = (f'="Floor "&TEXT({FLr},"$#,##0")&" per night  |  Minimum stay "&IF({MSr}="","-",{MSr}&" nights")&'
                  f'"  |  Cleaning fee "&TEXT({CLr},"$#,##0")')
RT_TOP = 11
rt_first = RT_TOP + 5
rt_last = rt_first + 11


def rt_formula(kind):
    def fn(r):
        i = r - rt_first
        return {
            "mult": f'=IF({SE}L{se_first + i}="","",{SE}L{se_first + i})',
            "base": f'=IF({SE}M{se_first + i}="","",{SE}M{se_first + i})',
            "night": f'=IF({SD}E{sd_first + i}="","",{SD}E{sd_first + i})',
            "flag": f'=IF({SD}F{sd_first + i}="","",{SD}F{sd_first + i})',
            "wk": f'=IF({SD}H{sd_first + i}="","",{SD}H{sd_first + i})',
            "mo": f'=IF({SD}L{sd_first + i}="","",{SD}L{sd_first + i})',
            "events": f"=COUNTIF({EV}$I${ev_first}:$I${ev_last},{i + 1})",
        }[kind]
    return fn


cols_rt = [
    dict(col="C", head="Month", kind="text", values=MONTHS),
    dict(col="D", head="Multiplier", kind="calc", fmt=MULT, align="right", formula=rt_formula("mult")),
    dict(col="E", head="Base nightly rate", kind="calc", fmt=USD0, align="right", formula=rt_formula("base")),
    dict(col="F", head="Nightly rate, floor applied", kind="calc", fmt=USD0, align="right", bold=True, formula=rt_formula("night")),
    dict(col="G", head="Flag", kind="calc", align="left", formula=rt_formula("flag")),
    dict(col="H", head="7+ nights", kind="calc", fmt=USD2, align="right", formula=rt_formula("wk")),
    dict(col="I", head="28+ nights", kind="calc", fmt=USD2, align="right", formula=rt_formula("mo")),
    dict(col="J", head="Events this month", kind="calc", fmt="0;-0;-", align="right", formula=rt_formula("events")),
]
f1, l1, rt_end = D.table_card(rt, RT_TOP, "B", "K", cols_rt, 12, title="Your 12 monthly rates",
                              sub="Nothing to type here. Copy each month's rate into your listing, then set each event window's nights to the event rate.")
assert (f1, l1) == (rt_first, rt_last)
rt_head = rt_first - 1
rt[f"H{rt_head}"].value = f'="7+ NIGHTS ("&TEXT({WDr},"0%")&" OFF), PER NIGHT"'
rt[f"I{rt_head}"].value = f'="28+ NIGHTS ("&TEXT({MDr},"0%")&" OFF), PER NIGHT"'
for r in range(rt_first, rt_last + 1):
    rt[f"G{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
rt.conditional_formatting.add(f"G{rt_first}:G{rt_last}", FormulaRule(formula=[f'$G{rt_first}="AT FLOOR"'],
                              font=Font(color=D.T["gold"], bold=True), fill=PatternFill("solid", bgColor=D.T["input_fill"]), stopIfTrue=True))
rt.conditional_formatting.add(f"G{rt_first}:G{rt_last}", FormulaRule(formula=[f'$G{rt_first}="ok"'],
                              font=Font(color=D.T["teal"]), stopIfTrue=True))
rt.conditional_formatting.add(f"J{rt_first}:J{rt_last}", FormulaRule(formula=[f"$J{rt_first}>0"],
                              font=Font(color=D.T["accent"], bold=True), stopIfTrue=True))
RF = f"F{rt_first}:F{rt_last}"
RC = f"C{rt_first}:C{rt_last}"
# event windows
EW_TOP = rt_end + 2
ew_first = EW_TOP + 5
ew_last = ew_first + N_EVENTS - 1
ew_tot = ew_last + 1


def ew_formula(src):
    def fn(r):
        i = r - ew_first + ev_first
        return f'=IF({EV}C{i}="","",{EV}{src}{i})'
    return fn


cols_ew = [
    dict(col="C", head="Event", kind="calc", formula=ew_formula("C")),
    dict(col="D", head="First night", kind="calc", fmt=DATE, align="right", formula=ew_formula("D")),
    dict(col="E", head="Last night", kind="calc", fmt=DATE, align="right", formula=ew_formula("E")),
    dict(col="F", head="Nights in window", kind="calc", fmt=INT, align="right", formula=ew_formula("F")),
    dict(col="G", head="Premium", kind="calc", fmt=PCT1, align="right", formula=ew_formula("G")),
    dict(col="H", head="Event nightly rate", kind="calc", fmt=USD0, align="right", bold=True, formula=ew_formula("K")),
    dict(col="I", head="Uplift per night", kind="calc", fmt=USD0, align="right", formula=ew_formula("L")),
    dict(col="J", head="Uplift for the window", kind="calc", fmt=USD0, align="right", formula=ew_formula("N")),
]
f1, l1, ew_end = D.table_card(rt, EW_TOP, "B", "K", cols_ew, N_EVENTS + 1, title="Event windows",
                              sub="Event nights override the month's rate. From 3 Events.")
assert (f1, l1) == (ew_first, ew_tot)
total_row(rt, ew_tot, D.cols("C", "J"), "C", "All events, if those nights book", {"J": (f"={EV}N{ev_tot}", USD0)})
D.h(rt, 6, 10)
D.stat_strip(rt, 7, [
    ("C:C", "Peak month", f'=IF(COUNT({RF})=12,INDEX({RC},MATCH(MAX({RF}),{RF},0))&"  "&TEXT(MAX({RF}),"$#,##0"),"-")', "@"),
    ("D:F", "Lowest month", f'=IF(COUNT({RF})=12,INDEX({RC},MATCH(MIN({RF}),{RF},0))&"  "&TEXT(MIN({RF}),"$#,##0"),"-")', "@"),
    ("G:H", "Months held at the floor", f"={SD}F{sd_tot}", INT),
    ("I:J", "Fee reset", f"={FR}D{k_stat}", "@"),
])
D.status_rule(rt, "I8", '$I$8="NOT RESET"', fill_tint=False)
D.h(rt, 9, 10)
D.frame(rt, 6, 9, "B", "K")
D.h(rt, 10, 14)
RT_NOTES = [
    ("How these are built", "The 12 nightly rates are your base rate times each month's multiplier, held at your floor where a month "
     "falls below it. The 7+ and 28+ columns are the per-night prices after your weekly and monthly discounts, also held at the "
     "floor. Event windows override the month's rate for those nights."),
    ("Not a forecast", "Nothing here predicts occupancy or revenue, and there is no live market data or sync with any platform: "
     "it prices the nights you already decided to sell, from your own inputs."),
]
rt_n_end = notes_card(rt, ew_end + 2, RT_NOTES)
S.page_break_before(rt, EW_TOP)
finish(rt, rt_n_end + 1, ("A", "L"), f"A{rt_first}", "accent")
RT = "'6 Rate Table'!"

# =====================================================================  1 Your Property: glance card, tiles, notes
B = D.Card(yp, TOP, "G", "H", "I", "K", extra="J", wb=wb)
B.title("Your year at a glance", "From the other tabs. Change a yellow cell and these move.")
B.eyebrow("NIGHTLY RATES")


def with_month(text, formula, month_formula, fmt=USD0):
    r = B.calc(text, formula, fmt)
    m = yp[f"J{r}"]; m.value = month_formula
    m.font = D.f(9, False, "muted"); m.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return r


RTF = f"{RT}$F${rt_first}:$F${rt_last}"
RTC = f"{RT}$C${rt_first}:$C${rt_last}"
g_peak = with_month("Peak month rate", f'=IF(COUNT({RTF})=12,MAX({RTF}),"")',
                    f'=IF(COUNT({RTF})=12,INDEX({RTC},MATCH(MAX({RTF}),{RTF},0)),"")')
g_low = with_month("Lowest month rate", f'=IF(COUNT({RTF})=12,MIN({RTF}),"")',
                   f'=IF(COUNT({RTF})=12,INDEX({RTC},MATCH(MIN({RTF}),{RTF},0)),"")')
g_floor = B.calc("Months held at floor", f"={SD}F{sd_tot}", INT)
B.eyebrow("LONG STAYS AND EVENTS")
g_wk = B.calc("Breaches, 7+ nights", f"={SD}I{sd_tot}", INT)
g_mo = B.calc("Breaches, 28+ nights", f"={SD}M{sd_tot}", INT)
g_ev = B.calc("Event uplift, if booked", f"={EV}N{ev_tot}", USD0)
B.eyebrow("FEE RESET")
g_set = B.calc("Raise your rate to", f"={FR}I{p_set}", USD0, total=True, tint="total_tint")
B.close()
for rr in (g_floor, g_wk, g_mo):
    D.status_rule(yp, f"I{rr}", f"$I${rr}>0", fill_tint=False)
assert A.end == B.end, (A.end, B.end)
D.tile(yp, 6, "B", "E", "BASE NIGHTLY RATE", f"=$D${BR}", USD0,
       f'=IF($D${BR}="","Type your base nightly rate below","Market "&IF($D${MT}="","not set",$D${MT})&", floor "&TEXT($D${FL},"$#,##0")'
       f'&", minimum stay "&$D${MS}&" nights")', dark=True)
v, sub, _ = D.tile(yp, 6, "G", "K", "HOST FEE RESET", f"={FR}D{k_stat}", "@",
                   f'=IF({FR}D{k_stat}="NOT RESET","Raise your base rate to "&TEXT({FR}I{p_set},"$#,##0")&" to keep your old payout",'
                   f'IF({FR}D{k_stat}="RESET","Your base rate keeps your old payout","Fill in 5 Fee Reset to check"))')
D.status_rule(yp, v.coordinate, f'{v.coordinate}="NOT RESET"', fill_tint=False)
D.h(yp, 10, 14)
YP_NOTES = [
    ("Base nightly rate", "Your average-month rate. Each month on 2 Seasonality is this rate times that month's multiplier, so a "
     "1.25x month is priced 25% above it. If you never repriced for the single host fee, run 5 Fee Reset and come back here."),
    ("Floor", "Your judgment, not a computed cost. No low month and no long-stay discount prices below it."),
    ("Sources", "Airbnb host service fees: Help Center article 1857 and the Resource Center article \"Simplifying service fees on "
     "Airbnb\". Weekly (7+ nights) and monthly (28+ nights) discounts: Help Center article 1233. All checked "
     f"{CHECKED}. Other platforms set their own fees and discount rules."),
    ("Before you rely on a number", S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab")),
]
yp_n_end = notes_card(yp, A.end + 2, YP_NOTES)
S.page_break_before(yp, A.end + 2)
finish(yp, yp_n_end + 1, ("A", "L"), "A11", "accent")

# order the working tabs 1 to 6
wb._sheets = [wb[n] for n in ("1 Your Property", "2 Seasonality", "3 Events", "4 Stay Discounts", "5 Fee Reset", "6 Rate Table")]

# =====================================================================  Start Here and Terms (card style, from #13)
SH_GRID = {"A": 3, "B": 2, "C": 8, "D": 62, "E": 18, "F": 2, "G": 3}


def sh_card(sh, top, title, rows):
    W = S.width_of(sh, "D") + S.width_of(sh, "E")
    r = top
    D.h(sh, r, 12); r += 1
    t = sh[f"C{r}"]; t.value = title; t.font = D.f(12, True); sh.merge_cells(f"C{r}:E{r}"); D.h(sh, r, 24); r += 1
    for left, text, kind in rows:
        full = kind in ("para", "link", "note")
        c = sh[f"C{r}"] if full else sh[f"D{r}"]
        if kind == "link":
            c.value = text[0]; c.hyperlink = text[1]; c.font = D.f(10, False, "blue", underline="single"); body = text[0]
        else:
            c.value = text; c.font = D.f(9 if kind == "note" else 10, False, "muted" if kind == "note" else "ink"); body = text
            m = re.match(r"^([A-Z0-9][A-Za-z0-9 ,+]{1,40})\. (.*)$", text, re.S) if kind in ("para", "num", "bullet") else None
            if m:
                c.value = CellRichText(TextBlock(InlineFont(rFont=D.FONT, sz=10, b=True, color=D.T["ink"]), m.group(1) + ". "),
                                       TextBlock(InlineFont(rFont=D.FONT, sz=10, color=D.T["ink"]), m.group(2)))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if full:
            sh.merge_cells(f"C{r}:E{r}")
            ht = S.text_height(body, W + 8, 10) + 8
        else:
            sh.merge_cells(f"D{r}:E{r}")
            ht = max(S.text_height(body, W, 10) + 10, 30)
        L = sh[f"C{r}"]
        if kind == "num":
            L.value = left; L.font = D.f(18, True, "accent"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind == "bullet":
            L.value = "•"; L.font = D.f(12, True, "teal"); L.alignment = Alignment(horizontal="center", vertical="center")
        elif kind.startswith("chip"):
            L.value = left[0]; L.number_format = left[1]; L.alignment = Alignment(horizontal="center", vertical="center")
            if kind == "chip-in":
                L.fill = D.fill("input_fill"); L.font = D.f(9, True, "input_text")
                L.border = Border(left=D.side("input_line"), right=D.side("input_line"), top=D.side("input_line"), bottom=D.side("input_line"))
            elif kind == "chip-calc":
                L.font = D.f(9, False, "ink"); L.border = Border(bottom=D.side("hair"))
            else:
                L.fill = D.fill("ink"); L.font = D.f(9, True, "on_ink")
        D.h(sh, r, ht); r += 1
    D.h(sh, r, 12)
    D.frame(sh, top, r, "B", "F")
    return r + 2


sh = wb.create_sheet("Start Here", 0)
S.set_widths(sh, SH_GRID)
D.page_header(sh, PRODUCT, "Start here. Price every night of the year in about 30 minutes.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
r = 6
r = sh_card(sh, r, "What it does", [(None,
    "Most hosts set one nightly rate, then stack discounts on top and hope. This workbook turns your last 12 months (or a "
    "starter curve for your market) into a rate for every month, adds a premium for local events, checks every weekly and "
    "monthly discount against a floor you set, and shows the rate that keeps your payout whole under Airbnb's single 15.5% host "
    "fee. The answer is one table you copy into your listing calendar.", "para")])
r = sh_card(sh, r, "Six steps", [
    (1, "Your Property. Base nightly rate, cleaning fee, minimum stay, market type, your discounts and your floor.", "num"),
    (2, "Seasonality. Type your last 12 months of nightly revenue or occupancy, or use the starter curve for your market.", "num"),
    (3, "Events. One row per local event: dates and a premium. The workbook prices the window.", "num"),
    (4, "Stay Discounts. Check every month's weekly and monthly price against your floor. Nothing to type.", "num"),
    (5, "Fee Reset. Your old rate and fee. Read the rate that keeps your payout, and whether your base rate is reset.", "num"),
    (6, "Rate Table. Your 12 monthly rates, long-stay prices and event windows on one page. Copy them into your calendar.", "num"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "How to read the cells", [
    ((EX["base"], USD0), "Yellow cells with blue numbers are yours to change. Type over the example. Each one shows a hint when you select it.", "chip-in"),
    ((round(M0["mult"][PEAK_I], 2), MULT), "Plain numbers are formulas. They update on their own and are locked so a stray keystroke can't break them.", "chip-calc"),
    ((FEE_SET, USD0), "Dark tiles and bold numbers are your answers. The tiles sit at the top of 1 Your Property and 5 Fee Reset.", "chip-key"),
    (None, "Each tab is protected without a password. To change anything else, use Review, Unprotect Sheet in Excel, or Data, "
           "Protect sheets and ranges in Google Sheets.", "note"),
])
wk_b = joined(M0["wk_breach"]); mo_b = joined(M0["mo_breach"])
r = sh_card(sh, r, "The example", [(None,
    f"Ridgeline Cabin, a fictional {EX['market']} cabin, has a {money(EX['base'])} base rate and a {money(EX['floor'])} floor. Its last "
    f"12 months make {MONTHS[PEAK_I]} the peak at {money(M0['night'][PEAK_I])} a night and hold {M0['at_floor']} months at the floor. "
    f"A {EX['weekly']*100:.0f}% weekly discount would go under the floor in {wk_b}, and the {EX['monthly']*100:.0f}% monthly discount "
    f"in {mo_b}: those nights bill at the floor instead. Three events add {money(EV_TOTAL)} if those nights book. At the single "
    f"15.5% fee, the {money(EX['fee_old_rate'])} rate pays {money(FEE_NOW_PAY, True)} a night instead of {money(FEE_OLD_PAY, True)}, "
    f"{money(FEE_LOST_Y)} over {EX['booked_nights']} nights; raising it to {money(FEE_SET)} keeps the old payout.", "para")])
r = sh_card(sh, r, "Try it on the example", [
    (None, f"Floor. On 1 Your Property, raise the floor from {money(EX['floor'])} to $140. Months held at the floor go from "
           f"{M0['at_floor']} to {M140['at_floor']}, and the weekly discount now breaches in {joined(M140['wk_breach'])}. Put it back.", "bullet"),
    (None, "Market. Change the market type to lake and clear the yellow history column on 2 Seasonality. Curve in use turns "
           "to STARTER and the year reshapes around summer. Undo both.", "bullet"),
    (None, f"Fee reset. Change the base nightly rate to {money(FEE_SET)}. 5 Fee Reset turns to RESET and every month on 6 Rate "
           f"Table rises by about {(FEE_SET / EX['base'] - 1) * 100:.0f}%.", "bullet"),
    (None, f"Events. Set the festival's expected booked nights to 2. Its uplift drops from {money(EV_RATES[0]['total'])} to "
           f"{money(FEST2)}, because only nights that book earn the premium.", "bullet"),
])
S.page_break_before(sh, r)
r = sh_card(sh, r, "Good to know", [
    (None, "Google Sheets. Upload the .xlsx to Google Drive, open it, then File, Save as Google Sheets. Every formula is a plain "
           "formula that works in Excel, Google Sheets, Numbers and LibreOffice. No macros, no add-ons, no sign-up.", "para"),
    (None, "Clear the example. Select the yellow cells and press Delete: 1 Your Property, the history column on 2 Seasonality "
           "and the event rows on 3 Events. Don't delete whole rows. The starter curves are yours to keep or edit.", "para"),
    (None, "Twice a year. Refresh the 12 months of history, check that the floor still covers your costs, and re-check the host "
           "fee on 5 Fee Reset against your platform's current terms.", "para"),
])
r = sh_card(sh, r, "Before you rely on a number", [(None,
    S.SHORT_NOTICE.format(date=CHECKED).replace("Terms of Use page", "Terms tab") + " "
    "Calculations are estimates for planning and depend on your inputs. They aren't accounting, tax or financial advice. Check "
    "fees, rates and interest against your own statements and the provider's current terms. Short-term rental rules, permits, "
    "taxes, HOA limits and platform policies vary by location and change often. Check your local requirements and platform terms.",
    "note")])
D.paint_canvas(sh, r - 1, "Z")
S.finish_sheet(sh, PRODUCT, r - 1, span=("A", "G"), tab_color="teal")

tm = wb.create_sheet("Terms")
S.set_widths(tm, SH_GRID)
D.page_header(tm, "Terms of Use", f"{PRODUCT}, version {VERSION}. Read this before you rely on a number.", f"Checked {CHECKED}",
              span=("B", "F"), side_cols=2)
paras = open(os.path.join(HERE, "LICENSE-AND-DISCLAIMER.txt"), encoding="utf-8").read().strip().split("\n\n")[1:]
r = 6
r = sh_card(tm, r, "More from the shop", [
    (None, ("Hourly Rate Quick Calculator: price your own time", "https://www.etsy.com/listing/4589695837"), "link"),
    (None, ("Everything in the shop: etsy.com/shop/ProofNotFluff", "https://www.etsy.com/shop/ProofNotFluff"), "link"),
])
S.page_break_before(tm, r)
r = sh_card(tm, r, "Terms of Use and disclaimer", [(None, p, "para") for p in paras])
r = sh_card(tm, r, "A small ask", [
    (None, "If this worksheet helped you price your nights, a short review on Etsy helps other hosts find it. Thank you.", "para"),
])
D.paint_canvas(tm, r - 1, "Z")
S.finish_sheet(tm, PRODUCT, r - 1, span=("A", "G"), tab_color="note")

wb.active = 0
problems = S.audit(wb)
if problems:
    print("AUDIT:", *problems, sep="\n  ")
wb.save(OUT)
print("saved", OUT)
print({"yp": dict(PN=PN, BR=BR, CL=CL, MS=MS, MT=MT, WD=WD, MD=MD, FL=FL, g_peak=g_peak, g_low=g_low, g_set=g_set),
       "se": dict(first=se_first, avg=se_avg), "ev": dict(first=ev_first, tot=ev_tot), "sd": dict(first=sd_first, tot=sd_tot),
       "fr": dict(OR=OR_, OF=OF, NF=NF, BN=BN, p_old=p_old, p_now=p_now, p_lost=p_lost, p_share=p_share, p_year=p_year,
                  p_exact=p_exact, p_set=p_set, p_at=p_at, p_guest=p_guest, k_base=k_base, k_gap=k_gap, k_stat=k_stat, RES=RES),
       "rt": dict(first=rt_first, ew_first=ew_first, ew_tot=ew_tot)})
print("example:", M0["night"], M0["at_floor"], M0["wk_breach"], M0["mo_breach"], EV_TOTAL, FEE_SET, round(FEE_LOST_Y, 2),
      M140["at_floor"], M140["wk_breach"], FEST2)
