#!/usr/bin/env python3
"""Mechanical legal pre-check for a ProofNotFluff product (the Legal desk's first pass; legal/README.md has the full checklist).

Usage: python3 legal_scan.py <product.zip | folder> [--listing listing.md] [--tier high|medium|low] [--json]

Finds: personal details in text or file metadata, promise words, em dashes, missing Terms page or short notice,
missing listing block lines, hidden sheets, comments, and workbook error cells. It never decides CLEAR:
a human-grade reviewer still reads every page against the checklist. Exit code 1 when anything high is found.
"""
import argparse, io, json, os, re, subprocess, sys, tempfile, zipfile, warnings
warnings.filterwarnings("ignore")

PERSONAL = [r"\btodd\b", r"\bmillen\b", r"outlandia", r"cata-?kor", r"reus research", r"outlandiastays", r"\boutla\b",
            r"t\.millen", r"630[\s.)-]*294", r"edison park"]
PROMISES = [r"guarantee", r"will save you (\$|\d|money)", r"\bget hired\b", r"pass (any|every|all) ats", r"ats[- ]proof", r"never (get )?audited",
            r"\bstay compliant\b", r"100% accurate", r"irs[- ]approved", r"lawyer[- ]approved", r"attorney[- ]approved",
            r"risk[- ]free", r"\bensures? compliance\b", r"\blegally binding\b"]
SOFT = [r"\bcertified\b", r"\bofficial\b", r"\bcompliant\b", r"\bensures?\b", r"\bproven\b"]
EMDASH = chr(0x2014)
BANNED = ["honestly", "genuinely", "straightforward", "delve", "unlock", "elevate", "seamless", "game-changer", "effortless", "supercharge"]
NOTICE = [r"not [^.\n]{0,40}(legal|tax|financial|professional)[^.\n]{0,40}advice"]
TERMS = [r"terms of use", r"as is", r"limit(ed)? (of )?liability|liability is limited|total liability"]
LISTING = ["Made with AI assistance and reviewed and tested by the shop owner.", "Digital download. No physical item ships."]

def text_of(path, findings, root=""):
    show = os.path.relpath(path, root) if root else path
    ext = os.path.splitext(path)[1].lower(); out = ""
    try:
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=False)
            for ws in wb.worksheets:
                if ws.sheet_state != "visible": findings.append(("medium", show, f"hidden sheet '{ws.title}' ({ws.sheet_state}): check what buyers can't see"))
                for row in ws.iter_rows():
                    for c in row:
                        if c.value is not None: out += f"\n{c.value}"
                        if c.comment: findings.append(("low", show, f"cell comment at {ws.title}!{c.coordinate}: {str(c.comment.text)[:80]}"))
            try:
                wbv = openpyxl.load_workbook(path, data_only=True)
                errs = [f"{ws.title}!{c.coordinate}" for ws in wbv.worksheets for row in ws.iter_rows() for c in row
                        if isinstance(c.value, str) and c.value.startswith("#") and c.value.rstrip("!?0/").upper() in ("#DIV", "#REF", "#VALUE", "#NAME", "#N/A", "#NUM", "#NULL")]
                if errs: findings.append(("high", show, f"error values cached in {len(errs)} cells, e.g. {errs[:5]}"))
            except Exception: pass
        elif ext == ".docx":
            import docx
            d = docx.Document(path)
            out = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                for r in t.rows: out += "\n" + " | ".join(c.text for c in r.cells)
            for s in d.sections:
                for p in list(s.header.paragraphs) + list(s.footer.paragraphs): out += "\n" + p.text
        elif ext == ".pdf":
            out = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout
        elif ext in (".txt", ".md", ".csv", ".html", ".htm", ".json"):
            out = open(path, encoding="utf-8", errors="replace").read()
    except Exception as e:
        findings.append(("medium", show, f"could not read: {e}"))
    return out

def meta_of(path, findings, root=""):
    show = os.path.relpath(path, root) if root else path
    ext = os.path.splitext(path)[1].lower(); vals = []
    try:
        if ext in (".xlsx", ".xlsm", ".docx", ".pptx"):
            z = zipfile.ZipFile(path)
            for n in z.namelist():
                if n.startswith("docProps/") or "comments" in n.lower() or n.endswith("people.xml"):
                    vals.append((n, z.read(n).decode("utf-8", "replace")))
                if n == "word/comments.xml" and "<w:comment " in z.read(n).decode("utf-8", "replace"): findings.append(("medium", show, "document has review comments: remove them"))
            if ext == ".docx":
                body = z.read("word/document.xml").decode("utf-8", "replace")
                if "<w:ins " in body or "<w:del " in body: findings.append(("high", show, "tracked changes left in the document"))
        elif ext == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(path); m = r.metadata or {}
            vals.append(("pdf-info", json.dumps({k: str(v) for k, v in m.items()})))
            try:
                x = r.xmp_metadata
                if x is not None: vals.append(("pdf-xmp", str(getattr(x, "dc_creator", "")) + str(getattr(x, "pdf_producer", ""))))
            except Exception: pass
        elif ext in (".png", ".jpg", ".jpeg"):
            from PIL import Image
            im = Image.open(path); info = dict(im.info); ex = im.getexif()
            vals.append(("image", json.dumps({k: str(v)[:200] for k, v in info.items()}) + json.dumps({str(k): str(v)[:200] for k, v in ex.items()})))
    except Exception as e:
        findings.append(("low", show, f"metadata not read: {e}"))
    for where, v in vals:
        for p in PERSONAL:
            if re.search(p, v, re.I): findings.append(("high", show, f"personal detail '{p}' in metadata {where}"))
    if ext in (".xlsx", ".docx", ".pptx"):
        core = next((v for w, v in vals if w == "docProps/core.xml"), "")
        a = re.search(r"<dc:creator>(.*?)</dc:creator>", core); lm = re.search(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core)
        for lab, m in (("author", a), ("last modified by", lm)):
            if m and m.group(1).strip() and m.group(1).strip().lower() not in ("proofnotfluff", "rob marda"):
                v = m.group(1).strip(); generic = v.lower() in ("un-named", "openpyxl", "python-docx", "pptxgenjs", "author", "user", "unknown")
                findings.append(("low" if generic else "medium", show, f"{lab} is '{v}': set it to ProofNotFluff"))

def scan(root, listing=None, tier="medium"):
    import shutil
    findings = []; tmp = tempfile.mkdtemp()
    if os.path.isfile(root) and root.lower().endswith(".zip"): zipfile.ZipFile(root).extractall(tmp)
    elif os.path.isfile(root): shutil.copy(root, tmp)
    else: shutil.copytree(root, tmp, dirs_exist_ok=True)
    root = tmp
    for _ in range(3):  # unpack nested zips (a product folder often holds the delivery zip)
        nested = [os.path.join(d, f) for d, _, fs in os.walk(root) for f in fs if f.lower().endswith(".zip")]
        if not nested: break
        for z in nested:
            zipfile.ZipFile(z).extractall(z[:-4] + "_unzipped"); os.remove(z)
    files = [os.path.join(d, f) for d, _, fs in os.walk(root) for f in fs if not f.startswith(".") and "__MACOSX" not in d]
    alltext = ""
    for f in sorted(files):
        meta_of(f, findings, root); t = text_of(f, findings, root); alltext += "\n" + t
        rel = os.path.relpath(f, root)
        for p in PERSONAL:
            if re.search(p, t, re.I): findings.append(("high", rel, f"personal detail matching '{p}' in the text"))
        for p in PROMISES:
            for m in re.finditer(p, t, re.I):
                ctx = t[max(0, m.start() - 50): m.end() + 50].replace("\n", " ")
                if re.search(r"\b(no|not|nothing|can't|cannot|doesn't|don't|won't|isn't|aren't|without|never|neither|nor)\b[^.]{0,40}" + p, ctx, re.I): continue
                findings.append(("high", rel, f"promise wording '{m.group(0)}': ...{ctx.strip()}..."))
        for p in SOFT:
            for m in re.finditer(p, t, re.I):
                ctx = t[max(0, m.start() - 40): m.end() + 40].replace("\n", " ")
                if re.search(r"professional\s+and\s+the\s+official\s+source", ctx, re.I): continue  # clause library wording
                findings.append(("low", rel, f"check context of '{m.group(0)}': ...{ctx.strip()}..."))
        if EMDASH in t: findings.append(("low", rel, f"{t.count(chr(0x2014))} em dash(es)"))
        for w in BANNED:
            if re.search(r"\b" + re.escape(w) + r"\b", t, re.I): findings.append(("low", rel, f"banned word '{w}'"))
    if not all(re.search(p, alltext, re.I) for p in NOTICE): findings.append(("high", "product", "no short notice found (not legal, tax or financial / professional advice)"))
    if not all(re.search(p, alltext, re.I) for p in TERMS): findings.append(("high", "product", "no full Terms of Use found (terms of use, as is, limit of liability)"))
    if not re.search(r"(checked|verified|as of)[^.\n]{0,40}(20\d\d)", alltext, re.I): findings.append(("medium", "product", "no as-of date found for figures"))
    if tier == "high" and not re.search(r"(attorney|cpa|enrolled agent|licensed)", alltext, re.I):
        findings.append(("high", "product", "High tier product with no 'check with an attorney or CPA' clause"))
    if listing:
        lt = open(listing, encoding="utf-8", errors="replace").read()
        for line in LISTING:
            if line not in lt: findings.append(("high", "listing", f"missing line: {line}"))
        for p in PROMISES:
            for m in re.finditer(p, lt, re.I):
                ctx = lt[max(0, m.start() - 50): m.end() + 50].replace("\n", " ")
                if not re.search(r"\b(no|not|nothing|can't|cannot|doesn't|don't|won't|isn't|aren't|without|never|neither|nor)\b[^.]{0,40}" + p, ctx, re.I):
                    findings.append(("high", "listing", f"promise wording '{m.group(0)}': ...{ctx.strip()}..."))
        if EMDASH in lt: findings.append(("low", "listing", "em dash in listing"))
    return findings

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("path"); ap.add_argument("--listing"); ap.add_argument("--tier", default="medium"); ap.add_argument("--json", action="store_true")
    a = ap.parse_args(); f = scan(a.path, a.listing, a.tier)
    order = {"high": 0, "medium": 1, "low": 2}; f.sort(key=lambda x: order[x[0]])
    if a.json: print(json.dumps([{"severity": s, "where": w, "issue": i} for s, w, i in f], indent=1))
    else:
        for s, w, i in f: print(f"[{s}] {w}: {i}")
        print(f"{sum(1 for x in f if x[0]=='high')} high, {sum(1 for x in f if x[0]=='medium')} medium, {sum(1 for x in f if x[0]=='low')} low")
    sys.exit(1 if any(x[0] == "high" for x in f) else 0)
