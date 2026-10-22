#!/usr/bin/env python3
"""The members' web app for Chicago New Business Licenses (ProofNotFluff lead product 4, Oct 2026).

Same quality bar and the same parts as the restaurant product's app (tools/rest_app.py) and the construction product's
(tools/leads_app.py, whose stylesheet, Leaflet pins and encryption this reuses by import, never by editing it): one
AES-256-GCM box under a key derived with PBKDF2-SHA256 at 250,000 rounds from the members' password (held only as the
LIC_LEADS_PASSWORD Actions secret and on the Gumroad product's post-purchase content), decrypted in the buyer's browser.
The repo holds only ciphertext, and only once the secret exists.

Buyers: commercial cleaning and pest control companies first, and any vendor who sells to a business in its first months.
Their job: pick the trades they sell to, open the businesses the City licensed in the last 30 days, see what the City's
records say (first license, new site, added license; a food license inspection), call or visit.

Drill-in: overview (businesses by trade), a filtered list with a map (#/list?t=<trade>&q=&area=&days=&type=&h=&fi=&new=1
&fu=&sort=&view=map), a business page (#/biz/<license>), the pitch kit (#/pitch?b=clean|pest|any) and the Terms (#/terms,
open before the password). Hash routes, so views can be bookmarked and the back button works. Contacted marks and notes
live in that buyer's browser only (localStorage).

Wording rule: every screen states the City's license type, activity and issue date and the City's disclaimer word for word.
Nothing says a business is open, new to Chicago or buying anything.

Usage: LIC_LEADS_PASSWORD=... python3 tools/lic_app.py FULL.json LEADS.xlsx docs/chicago-new-business-licenses-members/index.html
       python3 tools/lic_app.py --pw-file PATH FULL.json LEADS.xlsx OUT.html
       python3 tools/lic_app.py --selftest
"""
import base64, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from leads_page import decrypt, encrypt  # noqa: E402
import leads_app  # noqa: E402  (stylesheet and the Leaflet pins only)
import lic_xlsx as X  # noqa: E402  (the trades, words, pitch kit and Terms: one source)

FREE_PAGE = "https://proofnotfluff.com/chicago-new-business-licenses/"
SLUG = "chicago-new-business-licenses-members"
SECRET = "LIC_LEADS_PASSWORD"
LEGAL = ("Public City of Chicago business license records, as of __ASOF__. __CITY__ ProofNotFluff is not affiliated with or endorsed by "
         "the City of Chicago, and no business listed has endorsed or works with ProofNotFluff. A listing repeats the City's license "
         "record; it does not say a business is open or buying anything. You make your own calls, visits and emails and are responsible "
         "for following the law when you do. Licensed to one business for its own prospecting; no reselling or republishing. Provided as "
         "is, without warranties of any kind, to the extent the law allows, and ProofNotFluff's total liability for any claim is limited "
         "to the amount you paid for the most recent month of your membership. Made with AI assistance. Read the full Terms of Use (also "
         "on the Terms tab of the Excel file).")
TERMS_LINK = "Terms of Use"


def _esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def base_css():
    """The construction app's stylesheet, read from tools/leads_app.py so the products look like one family."""
    return re.search(r"<style>(.*?)</style>", leads_app.PAGE, re.S).group(1)


def legal_text():
    return LEGAL.replace("__CITY__", X.city_disclaimer())


def payload(full, xlsx_bytes):
    groups = X.by_trade(full)
    rows = []
    for k in X.ORDER:
        for r in groups.get(k, []):
            rows.append({"id": r["license"], "n": r["business"], "a": r["address"], "z": r.get("zip") or "", "ar": r.get("communityArea") or "",
                         "t": k, "ty": r["licenseType"], "ot": r.get("otherTypes") or [], "ac": r.get("activity") or "", "d": r["issued"],
                         "h": r.get("history") or "", "hs": r.get("since") or "", "fi": r.get("foodInsp") or "", "nw": X.is_new(r, full),
                         "lat": r.get("lat"), "lon": r.get("lon")})
    checks = (full.get("counts") or {}).get("checks") or {}
    return {"asOf": X.as_of_text(full), "asOfIso": full["source"]["fetchedAt"][:10], "window": full["window"],
            "trades": [{"key": k, "tab": t, "short": s, "kinds": w} for k, t, s, w in X.TRADES], "rows": rows,
            "checks": {"history": checks.get("history", True) is not False, "food": checks.get("foodInspections", True) is not False},
            "record": X.RECORD, "fileName": f"chicago-new-business-licenses-{full['source']['fetchedAt'][:10]}.xlsx",
            "xlsx": base64.b64encode(xlsx_bytes).decode()}


def legal_html():
    t = legal_text()
    assert t.count(TERMS_LINK) == 1
    return _esc(t).replace(TERMS_LINK, f'<a href="#/terms">{TERMS_LINK}</a>')


def words_json():
    """The shared words, from tools/lic_xlsx.py: they are in the clear because they hold no list data."""
    ct = X.city_terms()
    t = X.terms_text("__ASOF__")
    return json.dumps({"version": t[0], "paras": t[1:], "city": X.city_disclaimer(), "cityTerms": ct["url"], "citySections": ct["sections"],
                       "cityChecked": X.pretty(ct["checked"]), "notice": X.OUR_NOTICE, "leftOut": X.LEFT_OUT, "overlap": X.OVERLAP,
                       "history": X.HISTORY, "historyLong": X.HISTORY_LONG, "kits": X.PITCH, "rules": X.PITCH_RULES,
                       "trades": [{"key": k, "short": s, "kinds": w} for k, _, s, w in X.TRADES]},
                      ensure_ascii=False).replace("</", "<\\/")


def build_html(full, xlsx_bytes, password):
    box = encrypt(password, json.dumps(payload(full, xlsx_bytes), separators=(",", ":")).encode())
    js, jsi = leads_app.LEAFLET_JS
    css, cssi = leads_app.LEAFLET_CSS
    as_of = X.as_of_text(full)
    return (PAGE.replace("__BASECSS__", base_css()).replace("__BOX__", json.dumps(box)).replace("__WORDS__", words_json().replace("__ASOF__", as_of))
            .replace("__LEGAL__", legal_html()).replace("__ASOF__", as_of).replace("__FREE__", FREE_PAGE)
            .replace("__SURVEY__", leads_app.survey_url(X.PRODUCT)).replace("__LJS__", js).replace("__LJSI__", jsi)
            .replace("__LCSS__", css).replace("__LCSSI__", cssi))


def build(full, xlsx_path, out, password):
    html = build_html(full, open(xlsx_path, "rb").read(), password)
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    open(out, "w", encoding="utf-8").write(html)
    g = X.by_trade(full)
    return {k: len(g.get(k, [])) for k in X.ORDER} | {"rows": len(full["records"]), "bytes": len(html)}


PAGE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">
<meta name="robots" content="noindex, nofollow"><meta name="theme-color" content="#1D2433">
<title>Members: Chicago New Business Licenses | ProofNotFluff</title>
<link rel="preload" href="../assets/fonts/inter-400.woff2" as="font" type="font/woff2" crossorigin>
<link rel="preload" href="../assets/fonts/poppins-700.woff2" as="font" type="font/woff2" crossorigin>
<link rel="stylesheet" href="__LCSS__" integrity="__LCSSI__" crossorigin="anonymous">
<script defer src="__LJS__" integrity="__LJSI__" crossorigin="anonymous"></script>
<style>__BASECSS__
/* ---- new business licenses additions ---- */
:root{--first:#A63E21;--first-bg:#FBEAE4;--site:#2F5F9E;--site-bg:#E6EDF6;--add:#3A4152;--add-bg:#ECE6DA}
.hi{display:inline-flex;align-items:center;gap:6px;height:26px;padding:0 10px;border-radius:999px;font-size:12.5px;font-weight:700;white-space:nowrap;line-height:1}
.hi i{width:8px;height:8px;border-radius:50%;background:currentColor;flex:none}
.hi.first{color:var(--first);background:var(--first-bg)}.hi.first i{border-radius:2px}
.hi.site{color:var(--site);background:var(--site-bg)}.hi.added{color:var(--add);background:var(--add-bg)}
.hi.lg{height:32px;padding:0 14px;font-size:14px}
.tch{display:inline-flex;align-items:center;gap:5px;height:26px;padding:0 9px 0 7px;border-radius:7px;font-size:12.5px;font-weight:600;color:var(--ink2);background:#F3F0EA;border:1px solid var(--line2);white-space:nowrap}
.tch .ico{width:15px;height:15px;color:var(--muted)}
.row .mn{font:700 17px/1.25 var(--head);letter-spacing:-.005em}
.row .bd{flex-wrap:wrap;gap:2px 0}
.row .bd .mi{white-space:nowrap;display:inline-flex;align-items:center;gap:7px}
.row .bd .mi .dot{margin:0 0 0 7px}
.row .bd .mi.ty{white-space:normal}
.row .nwt{color:var(--accent-ink);font-weight:700}
.row .ctx{color:var(--teal);font-weight:700;display:inline-flex;align-items:center;gap:4px}
.row .fit{color:var(--muted);font-size:13px}
.row{grid-template-columns:minmax(0,1fr) auto}
.row .rt{grid-row:1 / span 3;flex-direction:column;align-items:flex-end;flex-wrap:nowrap;max-width:220px;gap:6px}
@media(max-width:600px){.row{grid-template-columns:minmax(0,1fr)}.row .rt{grid-row:auto;flex-direction:row;align-items:center;flex-wrap:wrap;max-width:none}}
/* overview: businesses by trade (the signature) */
.trades{display:grid;grid-template-columns:repeat(6,minmax(0,1fr));gap:12px;margin:22px 0 26px}
@media(max-width:1180px){.trades{grid-template-columns:repeat(3,minmax(0,1fr))}}
@media(max-width:700px){.trades{grid-template-columns:minmax(0,1fr);gap:8px}}
.trade{position:relative;display:flex;flex-direction:column;gap:2px;padding:16px 16px 14px;border-radius:var(--r);background:#fff;border:1px solid var(--line);box-shadow:var(--sh);text-decoration:none;color:var(--ink);min-width:0}
.trade .ti{display:grid;place-items:center;width:40px;height:40px;border-radius:10px;background:var(--paper);color:var(--ink2);margin:0 0 10px}
.trade .ti .ico{width:22px;height:22px}
.trade .tn{font:700 34px/1 var(--head);letter-spacing:-.02em;margin:0 0 6px}
.trade .tt{font-weight:600;font-size:15px}
.trade .tm{color:var(--muted);font-size:13.5px}
.trade .tm b{color:var(--accent-ink)}
.trade.top{background:var(--ink);border-color:var(--ink);color:#fff}
.trade.top .ti{background:rgba(255,255,255,.1);color:#fff}.trade.top .tm{color:#C9CFDB}.trade.top .tm b{color:#F7B9A5}
.trade.zero{box-shadow:none;background:#FBFAF7}.trade.zero .tn{color:var(--muted)}
@media(hover:hover){.trade:hover .tt{text-decoration:underline;text-underline-offset:3px}}
@media(max-width:700px){.trade{display:grid;grid-template-columns:40px minmax(0,1fr) auto;column-gap:14px;align-items:center;padding:12px 14px}
 .trade .ti{margin:0;grid-row:1 / span 2}.trade .tn{grid-column:3;grid-row:1 / span 2;font-size:28px;margin:0;text-align:right}.trade .tt{grid-column:2}.trade .tm{grid-column:2}}
.kpis{display:flex;flex-wrap:wrap;gap:8px 22px;margin:14px 0 0;font-size:14.5px;color:var(--ink2)}
.kpis b{font:700 16px var(--head);color:var(--ink)}
/* lists */
.lsw{display:flex;gap:8px;flex-wrap:wrap;margin:0 0 14px}
.lsw a{display:inline-flex;align-items:center;gap:8px;min-height:44px;padding:0 16px;border-radius:999px;border:1px solid var(--line);background:#fff;color:var(--ink);font-weight:600;font-size:14.5px;text-decoration:none}
.lsw a .ico{width:16px;height:16px;color:var(--muted)}
.lsw a .c{color:var(--muted)}
.lsw a[aria-current=page]{background:var(--ink);border-color:var(--ink);color:#fff}.lsw a[aria-current=page] .c,.lsw a[aria-current=page] .ico{color:#C9CFDB}
@media(max-width:860px){.lsw{flex-wrap:nowrap;overflow-x:auto;scrollbar-width:none;margin:0 -16px 12px;padding:0 16px}.lsw::-webkit-scrollbar{display:none}.lsw a{flex:none}}
.ldesc{margin:-6px 0 16px;color:var(--muted);font-size:14.5px;max-width:80ch}
.ovl{margin:-4px 0 16px;padding:12px 16px;border-radius:10px;background:var(--paper);color:var(--ink2);font-size:14px;max-width:86ch;border:1px solid var(--line2)}
.ovl b{color:var(--ink)}
.lg-1{width:12px;height:12px;background:#C8502F;box-shadow:0 0 0 1.5px #fff,0 0 0 2.5px #A63E21;border-radius:2px!important;margin:0 3px}
.lg-2{width:10px;height:10px;background:#2F5F9E;border:2px solid #fff;box-shadow:0 0 0 1px #2F5F9E;margin:0 3px}
.lg-3{width:10px;height:10px;background:#fff;border:2px solid #3A4152;margin:0 3px}
.pin-f{background:#C8502F;border:1.5px solid #fff;box-shadow:0 0 0 1px #A63E21,0 1px 3px rgba(29,36,51,.3);border-radius:2px;box-sizing:border-box}
.pin-f.hot{background:#1D2433;box-shadow:0 0 0 2px #C8502F}
/* business page */
.lic{padding:20px}
.lic h2{font-size:17px;font-weight:600;margin:0 0 14px}
.lic .top{display:flex;flex-wrap:wrap;align-items:center;gap:8px 12px;margin:0 0 6px}
.lic .top .ty{font:700 20px/1.25 var(--head);letter-spacing:-.01em}
.lic .top .dt{color:var(--ink2);font-size:15px;font-weight:600}
.lic .act{margin:0 0 14px;color:var(--ink2);font-size:15px}
.lic .why{margin:0;padding:12px 14px;border-radius:10px;background:var(--paper);font-size:14.5px;color:var(--ink2)}
.lic .why b{color:var(--ink)}
.lic .why+.why{margin-top:10px}
.facts dd.nv{color:var(--muted);font-weight:400}
.facts dd .dn{display:block;font-size:12.5px;font-weight:400;color:var(--muted);margin-top:2px}
.acts .btn{min-height:44px}
/* pitch kit */
.bk{display:flex;gap:8px;flex-wrap:wrap;margin:0 0 18px}
@media(max-width:860px){.bk{flex-wrap:nowrap;overflow-x:auto;scrollbar-width:none;margin:0 -16px 16px;padding:0 16px}.bk::-webkit-scrollbar{display:none}}
.bk a{flex:none;display:inline-flex;align-items:center;gap:8px;min-height:44px;padding:0 14px;border-radius:999px;border:1px solid var(--line);background:#fff;color:var(--ink);font-weight:600;font-size:14.5px;text-decoration:none}
.bk a[aria-current=page]{background:var(--ink);border-color:var(--ink);color:#fff}
.call{font-size:16px;line-height:1.6;color:var(--ink);margin:0;padding:16px 18px;border-radius:10px;background:var(--paper)}
.call .f{color:var(--blue);font-weight:700}
.rules{margin:18px 0 0;padding:0 0 0 20px;color:var(--ink2);font-size:14.5px}.rules li{margin:0 0 6px}
.pk-h{display:flex;align-items:center;justify-content:space-between;gap:12px;margin:0 0 4px}.pk .sub{margin:0 0 14px;font-size:13.5px;color:var(--muted)}
.pk .panel{padding:20px}
.pk h2{font-size:17px;font-weight:600;margin:0}
.terms{max-width:780px}.terms .tx p{max-width:none}.terms .tx h3{font:600 14px var(--body);letter-spacing:.04em;color:var(--ink);margin:18px 0 6px}
.terms .tx h3:first-child{margin-top:0}
#fsheetH:focus{outline:none;box-shadow:none}
.pick{display:flex;flex-wrap:wrap;align-items:flex-end;gap:10px 14px;margin:14px 0 0}
.pick .fld select{max-width:300px}
.pick p{margin:0;color:var(--muted);font-size:14px;line-height:40px}
@media(max-width:860px){.pick p{line-height:44px}}
@media(max-width:600px){.pick{row-gap:8px}.pick p{line-height:1.4;flex-basis:100%}}
</style></head><body>
<a class="skip" href="#/" id="skip">Skip to main content</a>
<header class="bar"><div class="bar-in">
<a class="brand" href="#/" aria-label="ProofNotFluff, Chicago New Business Licenses, overview"><b>ProofNotFluff</b><span class="marks" aria-hidden="true"><i></i><i></i><i></i><i></i><i></i></span><small>Chicago New Business Licenses</small></a>
<nav class="tabs" id="tabs" aria-label="Sections" hidden></nav>
<span class="bar-sp"></span>
<button class="dl" id="dlTop" type="button" hidden><svg class="ico" viewBox="0 0 24 24" aria-hidden="true"><path d="M12 4v11M7 10l5 5 5-5M5 20h14"/></svg><span>Download workbook</span><b class="sr">Download this week's Excel workbook</b></button>
</div></header>
<main id="gateView" class="gate">
<form class="gate-card" id="gate" novalidate>
<span class="chip s2"><svg class="ico s" viewBox="0 0 24 24" aria-hidden="true"><rect x="5" y="11" width="14" height="10" rx="2"/><path d="M8 11V8a4 4 0 0 1 8 0v3"/></svg>Members</span>
<h1>This week's new Chicago business licenses</h1>
<p class="sub">List of __ASOF__. Every new license the City issued in the last 30 days at a Chicago business site, one row per business and site, sorted by trade, for vendors who sell to new businesses.</p>
<label for="pw">Members' password</label>
<div class="pw"><input id="pw" type="password" autocomplete="current-password" autocapitalize="off" spellcheck="false" required aria-describedby="err pwHelp">
<button type="button" class="eye" id="eye" aria-label="Show password" aria-pressed="false"><svg class="ico" viewBox="0 0 24 24" aria-hidden="true"><path d="M2 12s3.6-7 10-7 10 7 10 7-3.6 7-10 7S2 12 2 12z"/><circle cx="12" cy="12" r="3"/><path id="eyeslash" d="M3 3l18 18" hidden/></svg></button></div>
<button class="btn pri" type="submit" id="go">Open this week's list</button>
<p class="err" id="err" role="alert"></p>
<div class="gate-help" id="pwHelp"><p>The password is on your Gumroad purchase page, linked from your receipt email. <a href="https://gumroad.com/library">Open your Gumroad library</a></p>
<p>Not a member yet? <a href="__FREE__">See the free sample</a></p></div>
</form>
</main>
<p class="gate-legal" id="gateLegal">__LEGAL__</p>
<main id="app" class="page" hidden tabindex="-1"></main>
<footer class="foot-legal" id="appLegal" hidden><p>__LEGAL__</p><p>List of __ASOF__. Questions about your membership? Reply to your Gumroad receipt.</p><p class="leave">Thinking of cancelling? <a href="__SURVEY__" target="_blank" rel="noopener">Tell us why first<span class="sr"> (opens in a new tab)</span></a>. It takes a minute and decides what we fix next.</p></footer>
<nav class="bnav" id="bnav" aria-label="Sections" hidden></nav>
<div class="toast" id="toast" role="status" aria-live="polite"></div>
<script>
"use strict";
const BOX=__BOX__;
const W=__WORDS__;
const b64=s=>Uint8Array.from(atob(s),c=>c.charCodeAt(0));
async function openBox(pw){const k0=await crypto.subtle.importKey("raw",new TextEncoder().encode(pw),"PBKDF2",false,["deriveKey"]);
const key=await crypto.subtle.deriveKey({name:"PBKDF2",salt:b64(BOX.salt),iterations:BOX.rounds,hash:"SHA-256"},k0,{name:"AES-GCM",length:256},false,["decrypt"]);
const pt=await crypto.subtle.decrypt({name:"AES-GCM",iv:b64(BOX.iv)},key,b64(BOX.ct));return JSON.parse(new TextDecoder().decode(pt));}
const $=s=>document.querySelector(s);
const esc=s=>String(s??"").replace(/[&<>"']/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]));
const int=n=>Math.round(+n||0).toLocaleString("en-US");
const MON=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"];
function day(s){const m=/^(\d{4})-(\d{2})-(\d{2})/.exec(s||"");return m?MON[+m[2]-1]+" "+(+m[3])+", "+m[1]:"";}
function dayShort(s){const m=/^(\d{4})-(\d{2})-(\d{2})/.exec(s||"");return m?MON[+m[2]-1]+" "+(+m[3]):"";}
const plural=(n,a,b)=>int(n)+" "+(n===1?a:(b||a+"s"));
const I={phone:'<path d="M5 4h4l2 5-2.5 1.5a11 11 0 0 0 5 5L15 13l5 2v4a2 2 0 0 1-2 2A16 16 0 0 1 3 6a2 2 0 0 1 2-2"/>',
copy:'<rect x="9" y="9" width="11" height="11" rx="2"/><path d="M5 15V6a2 2 0 0 1 2-2h8"/>',
pin:'<path d="M12 21s-7-6.2-7-11.5a7 7 0 0 1 14 0C19 14.8 12 21 12 21z"/><circle cx="12" cy="9.5" r="2.5"/>',
search:'<circle cx="11" cy="11" r="7"/><path d="m20 20-3.5-3.5"/>',filter:'<path d="M4 6h16M7 12h10M10 18h4"/>',check:'<path d="m5 12.5 4.5 4.5L19 7.5"/>',
back:'<path d="M15 18l-6-6 6-6"/>',chev:'<path d="m9 6 6 6-6 6"/>',x:'<path d="M6 6l12 12M18 6 6 18"/>',list:'<path d="M8 6h12M8 12h12M8 18h12M4 6h.01M4 12h.01M4 18h.01"/>',
map:'<path d="M9 4 3 6.5v13.5L9 17.5l6 2.5 6-2.5V4l-6 2.5z"/><path d="M9 4v13.5M15 6.5V20"/>',
home:'<path d="M3 10.5 12 4l9 6.5V20a1 1 0 0 1-1 1h-5v-6h-6v6H4a1 1 0 0 1-1-1z"/>',
kit:'<path d="M8 4h8M9 4V3h6v1M6 6h12v15H6z"/><path d="M9 11h6M9 15h6"/>',csv:'<path d="M14 3H6v18h12V7z"/><path d="M14 3v4h4M9 13h6M9 17h6"/>',
doc:'<path d="M14 3H6v18h12V7z"/><path d="M14 3v4h4"/>',all:'<rect x="4" y="4" width="7" height="7" rx="1.5"/><rect x="13" y="4" width="7" height="7" rx="1.5"/><rect x="4" y="13" width="7" height="7" rx="1.5"/><rect x="13" y="13" width="7" height="7" rx="1.5"/>',
food:'<path d="M7 3v8M5 3v4a2 2 0 0 0 4 0V3M7 11v10M17 21V3c-2 0-3.5 2.6-3.5 6s1.5 4.5 3.5 4.5"/>',
retail:'<path d="M4 10 6 4h12l2 6"/><path d="M4 10a2.7 2.7 0 0 0 5.3 0 2.7 2.7 0 0 0 5.4 0 2.7 2.7 0 0 0 5.3 0M5.5 12.5V20h13v-7.5M10 20v-4.5h4V20"/>',
beauty:'<circle cx="6.5" cy="7" r="2.5"/><circle cx="6.5" cy="17" r="2.5"/><path d="M8.6 8.4 20 18M8.6 15.6 20 6"/>',
venue:'<path d="m12 4 2.4 5 5.4.6-4 3.7 1.1 5.4L12 16l-4.9 2.7 1.1-5.4-4-3.7 5.4-.6z"/>',
auto:'<path d="M4 16h16M6 16v-4l2-5h8l2 5v4M6 16v2.5M18 16v2.5"/><path d="M6 12h12"/>',
care:'<rect x="4" y="12.5" width="7" height="7" rx="1.2"/><rect x="13" y="12.5" width="7" height="7" rx="1.2"/><rect x="8.5" y="4" width="7" height="7" rx="1.2"/>',
clean:'<path d="M12 3v4M12 17v4M3 12h4M17 12h4M6.3 6.3l2.4 2.4M15.3 15.3l2.4 2.4M6.3 17.7l2.4-2.4M15.3 8.7l2.4-2.4"/>',
pest:'<ellipse cx="12" cy="13.5" rx="4.2" ry="5.5"/><path d="M12 8V5.5M9.5 4.5 11 6M14.5 4.5 13 6M7.8 11H4.5M7.8 15H4.5M8.4 18.5 6 20.5M16.2 11h3.3M16.2 15h3.3M15.6 18.5l2.4 2"/>',
any:'<path d="M4 10 6 4h12l2 6"/><path d="M4 10a2.7 2.7 0 0 0 5.3 0 2.7 2.7 0 0 0 5.4 0 2.7 2.7 0 0 0 5.3 0M5.5 12.5V20h13v-7.5"/>'};
const ico=(n,c)=>`<svg class="ico${c?" "+c:""}" viewBox="0 0 24 24" aria-hidden="true">${I[n]}</svg>`;
function toast(t){const el=$("#toast");el.textContent=t;el.classList.add("on");clearTimeout(toast.t);toast.t=setTimeout(()=>el.classList.remove("on"),2500);}
function untoast(){clearTimeout(toast.t);$("#toast").classList.remove("on");}
async function copyText(t){try{await navigator.clipboard.writeText(t);return true;}catch(e){const a=document.createElement("textarea");a.value=t;a.setAttribute("readonly","");a.style.position="fixed";a.style.opacity="0";document.body.appendChild(a);a.select();let ok=false;try{ok=document.execCommand("copy");}catch(x){}a.remove();return ok;}}

/* trades and the City record */
const TR={},TORD=W.trades.map(t=>t.key);W.trades.forEach(t=>{TR[t.key]=t;});
const tch=k=>`<span class="tch">${ico(k)}${esc(TR[k].short)}</span>`;
const HCLS=h=>h==="first"?"first":h==="site"?"site":h==="added"?"added":"";
function hiText(r){if(r.h==="site"&&r.hs)return "New site, licensed since "+r.hs;if(r.h==="added"&&r.hs)return "Added license, here since "+r.hs;return W.history[r.h]||"";}
const hi=(r,lg)=>r.h?`<span class="hi ${HCLS(r.h)}${lg?" lg":""}"><i aria-hidden="true"></i>${esc(hiText(r))}</span>`:"";

/* private marks and notes: this browser only */
const SKEY="pnf-lic-v1";let MARKS={},STORE_OK=true;
try{const t="__pnf";localStorage.setItem(t,"1");localStorage.removeItem(t);MARKS=JSON.parse(localStorage.getItem(SKEY)||"{}")||{};}catch(e){STORE_OK=false;MARKS={};}
function saveMarks(){try{localStorage.setItem(SKEY,JSON.stringify(MARKS));return true;}catch(e){STORE_OK=false;return false;}}
const mk=id=>MARKS[id]||{};
function setMark(id,patch){const m=Object.assign({},mk(id),patch);if(!m.c&&!(m.n||"").trim())delete MARKS[id];else MARKS[id]=m;return saveMarks();}
const STORE_LINE=()=>STORE_OK?"Saved in this browser on this device only. Nobody else sees it, including ProofNotFluff. Clearing your browser data removes it.":"This browser is blocking storage, so contacted marks and notes will be lost when you close the page.";

let D,R,BYID;
const STATE={trail:[],map:null,mini:null};
function prep(d){D=d;R=d.rows;BYID={};R.forEach(r=>{BYID[r.id]=r;});}
const inTrade=t=>t?R.filter(r=>r.t===t):R;

/* routing */
function parse(){const h=(location.hash||"#/").slice(1);const i=h.indexOf("?");const path=i<0?h:h.slice(0,i);return{parts:path.split("/").filter(Boolean).map(decodeURIComponent),q:new URLSearchParams(i<0?"":h.slice(i+1))};}
function go(h){location.hash=h;}
function back(fallback){if(STATE.trail.length>1)history.back();else go(fallback);}
function onHash(){const h=location.hash||"#/";const t=STATE.trail;if(t.length>1&&t[t.length-2]===h)t.pop();else if(t[t.length-1]!==h)t.push(h);route(true);}
function setNav(v){const items=[["overview","#/","Overview","home"],["list","#/list","Businesses","list"],["pitch","#/pitch","Pitch kit","kit"]];
const html=items.map(([k,h,t,i])=>`<a href="${h}"${k===v?' aria-current="page"':""}>${ico(i)}${t}</a>`).join("");$("#tabs").innerHTML=html;$("#bnav").innerHTML=html;}
function route(scroll){const {parts,q}=parse();const v=parts[0]||"overview";
untoast();for(const k of ["map","mini"]){const m=STATE[k];if(m){try{m.stop();m.off();m.remove();}catch(e){}STATE[k]=null;}}closeSheet();
const app=$("#app");let nav=v;
if(v==="list"){viewList(q);nav="list";}
else if(v==="biz"&&BYID[parts[1]]){viewBiz(BYID[parts[1]]);nav="list";}
else if(v==="pitch")viewPitch(q);
else if(v==="terms"){viewTerms();nav="";}
else if(v==="biz"){app.innerHTML=`<div class="panel empty"><h3>That business is not on this week's list</h3><p>Businesses leave the list when their license is more than 30 days old.</p><a class="btn ink" href="#/list">See this week's list</a></div>`;nav="list";document.title="Not on this week's list | Chicago New Business Licenses";}
else{viewOverview();nav="overview";}
setNav(nav);app.classList.remove("view-enter");void app.offsetWidth;app.classList.add("view-enter");
if(scroll){window.scrollTo(0,0);const h=app.querySelector("h1");if(h){h.setAttribute("tabindex","-1");h.focus({preventScroll:true});}}}

/* a row in any list */
function row(r,ov){const m=mk(r.id);
const meta=[r.nw?'<span class="nwt">New</span>':"",`<span>Licensed ${dayShort(r.d)}</span>`,`<span>${esc(r.ty)}</span>`,
 r.fi?`<span>License inspection ${dayShort(r.fi)}</span>`:"",m.c?`<span class="ctx">${ico("check","s")}Contacted</span>`:""].filter(Boolean)
 .map((t,i)=>`<span class="mi${i&&/^<span>[^<]{28,}/.test(t)?" ty":""}">${i?'<span class="dot" aria-hidden="true">&middot;</span>':""}${t}</span>`).join("");
return `<li><a class="row" href="#/biz/${encodeURIComponent(r.id)}" data-id="${esc(r.id)}">
<span class="mn">${esc(r.n)}</span>
<span class="ad">${esc(r.a)}${r.ar?`<span class="ar">, ${esc(r.ar)}</span>`:""}</span>
<span class="bd">${meta}</span>${r.ac?`<span class="bd"><span class="fit">${esc(r.ac)}</span></span>`:""}
<span class="rt">${ov?tch(r.t):""}${hi(r)}</span>${ico("chev","chev")}</a></li>`;}
function rowsHtml(list,limit,ov){if(!list.length)return "";const n=Math.min(limit||list.length,list.length);return `<ul class="list-rows">${list.slice(0,n).map(r=>row(r,ov)).join("")}</ul>`;}
const newest=(a,b)=>b.d.localeCompare(a.d)||a.n.localeCompare(b.n);

/* overview */
function viewOverview(){document.title="Overview | Chicago New Business Licenses";
const n=R.length,nw=R.filter(r=>r.nw).length,first=R.filter(r=>r.h==="first").length,food=R.filter(r=>r.fi).length;
const counts=TORD.map(k=>[k,inTrade(k).length]);const topk=counts.slice().sort((a,b)=>b[1]-a[1])[0][0];
const tile=k=>{const rs=inTrade(k);const nn=rs.filter(r=>r.nw).length;
return `<a class="trade${k===topk?" top":""}${rs.length?"":" zero"}" href="#/list?t=${k}"><span class="ti">${ico(k)}</span><span class="tn num">${int(rs.length)}</span><span class="tt">${esc(TR[k].short)}</span><span class="tm num">${nn?`<b>${int(nn)} new</b> this week`:"none new this week"}</span></a>`;};
const ac={};R.forEach(r=>{if(r.ar){const a=ac[r.ar]=ac[r.ar]||[0,0];a[r.h==="first"?0:1]++;}});
const areas=Object.entries(ac).map(([a,v])=>[a,v[0]+v[1],v]).sort((a,b)=>b[1]-a[1]||a[0].localeCompare(b[0])).slice(0,8);const amax=areas.length?areas[0][1]:1;
const marked=Object.keys(MARKS).filter(id=>BYID[id]);const cN=marked.filter(id=>MARKS[id].c).length,nN=marked.filter(id=>(MARKS[id].n||"").trim()).length;
const fresh=R.filter(r=>r.h==="first").slice().sort(newest).slice(0,6);
$("#app").innerHTML=`<section class="hero"><h1>${int(n)} new business licenses in Chicago in the last 30 days, ${int(nw)} of them this week</h1>
<p>From City of Chicago business license records, issued ${esc(dayShort(D.window.listFrom))} to ${esc(day(D.window.to))}. List of ${esc(D.asOf)}. One row per business and site, sorted by trade.</p>
<p class="kpis"><span><b class="num">${int(first)}</b> first Chicago license</span><span><b class="num">${int(n-first-R.filter(r=>!r.h).length)}</b> new site or added license</span>${D.checks.food?`<span><b class="num">${int(food)}</b> with a food license inspection</span>`:""}</p></section>
<nav class="trades" aria-label="Businesses by trade">${TORD.map(tile).join("")}</nav>
<div class="ov"><div class="ov-col">
<section class="panel"><div class="sec-h"><h2>Newest licenses<span class="c num">${int(n)}</span></h2><a href="#/list">See all ${int(n)}</a></div>${rowsHtml(R.slice().sort(newest),8,true)}</section>
${D.checks.history?`<section class="panel"><div class="sec-h"><h2>First Chicago license<span class="c num">${int(first)}</span></h2><a href="#/list?h=first">See all ${int(first)}</a></div><p class="rec" style="padding-top:0">${esc(W.historyLong.first)}</p>${rowsHtml(fresh,0,true)}</section>`:""}
</div><div class="ov-col">
<section class="panel"><div class="sec-h"><h2>Where they are</h2><a href="#/list?view=map">Open the map</a></div>
<p class="blg" aria-hidden="true"><span><i style="background:#C8502F"></i>First license</span><span><i style="background:#2F5F9E"></i>Other new licenses</span></p>
<ul class="bars" aria-label="Businesses on the list by community area"><li class="hd" aria-hidden="true"><span></span><span class="sp2"></span><span>First</span><span>All</span></li>${areas.map(([a,nn,v])=>`<li><a href="#/list?area=${encodeURIComponent(a)}" aria-label="${esc(a)}: ${plural(nn,"business","businesses")}, ${int(v[0])} with a first license"><span>${esc(a)}</span><span class="tr" style="width:${Math.max(3,Math.round(nn/amax*100))}%" aria-hidden="true">${v[0]?`<i style="flex-grow:${v[0]};background:#C8502F"></i>`:""}${v[1]?`<i style="flex-grow:${v[1]};background:#2F5F9E"></i>`:""}</span><span class="nv num${v[0]?"":" z"}" style="color:var(--first)">${int(v[0])}</span><span class="v num">${int(nn)}</span></a></li>`).join("")}</ul></section>
<section class="panel"><div class="sec-h"><h2>Your follow-ups</h2>${cN||nN?`<a href="#/list?fu=mine">Show them</a>`:""}</div>
<div class="mine">${cN||nN?`<p class="mh"><b>${plural(cN,"business","businesses")}</b> marked contacted, <b>${plural(nN,"note")}</b> written.</p>`:`<p class="mh"><b class="blk">No follow-ups yet</b>Mark a business as contacted or add a note on its page, and it shows up here.</p>`}<p>${esc(STORE_LINE())}</p></div></section>
<section class="panel"><div class="sec-h"><h2>What to say</h2><a href="#/pitch">Open the pitch kit</a></div><p class="rec" style="padding-top:0">A call opener and an email for commercial cleaning, pest control and any vendor to new businesses, with the business's details to fill in.</p></section>
</div></div>`;}

/* the list: filters, list and map, in sync */
const SORTS={newest:()=>newest,az:()=>(a,b)=>a.n.localeCompare(b.n),area:()=>(a,b)=>(a.ar||"~").localeCompare(b.ar||"~")||newest(a,b)};
const DAYS=[["","Last 30 days"],["14","Last 14 days"],["7","Last 7 days"]];
function readF(q){return{t:TR[q.get("t")]?q.get("t"):"",q:(q.get("q")||"").trim(),area:q.get("area")||"",days:["7","14"].includes(q.get("days"))?q.get("days"):"",
type:q.get("type")||"",h:["first","site","added"].includes(q.get("h"))?q.get("h"):"",fi:["yes","no"].includes(q.get("fi"))?q.get("fi"):"",nw:q.get("new")==="1",
fu:["todo","mine"].includes(q.get("fu"))?q.get("fu"):"",sort:SORTS[q.get("sort")]?q.get("sort"):"newest",view:q.get("view")==="map"?"map":"list"};}
function writeF(f){const q=new URLSearchParams();for(const k of ["t","q","area","days","type","h","fi"])if(f[k])q.set(k,f[k]);if(f.nw)q.set("new","1");if(f.fu)q.set("fu",f.fu);if(f.sort!=="newest")q.set("sort",f.sort);if(f.view==="map")q.set("view","map");const s=q.toString();return "#/list"+(s?"?"+s:"");}
const norm=s=>String(s||"").toLowerCase().normalize("NFKD").replace(/[^a-z0-9 ]/g," ").replace(/\s+/g," ").trim();
function since(days){const t=new Date(D.window.to+"T12:00:00Z");t.setUTCDate(t.getUTCDate()-(+days));return t.toISOString().slice(0,10);}
function pass(r,f,skip){
if(skip!=="t"&&f.t&&r.t!==f.t)return false;
if(skip!=="q"&&f.q){const t=norm(f.q);const hay=norm([r.n,r.a,r.ar,r.z,r.ty,r.ac,r.id,...r.ot].join(" "));if(!t.split(" ").every(w=>hay.includes(w)))return false;}
if(skip!=="area"&&f.area&&r.ar!==f.area)return false;if(skip!=="days"&&f.days&&r.d<since(f.days))return false;if(skip!=="type"&&f.type&&r.ty!==f.type)return false;
if(skip!=="h"&&f.h&&r.h!==f.h)return false;if(skip!=="fi"&&f.fi&&(!!r.fi)!==(f.fi==="yes"))return false;if(skip!=="new"&&f.nw&&!r.nw)return false;
if(f.fu==="todo"&&mk(r.id).c)return false;if(f.fu==="mine"&&!(mk(r.id).c||(mk(r.id).n||"").trim()))return false;return true;}
function activeCount(f){return [f.area,f.days,f.type,f.h,f.fi,f.nw,f.fu].filter(Boolean).length;}
function viewList(q){const f=readF(q);STATE.f=f;STATE.shown=60;const tr=TR[f.t];document.title=`${tr?tr.short:"All businesses"} | Chicago New Business Licenses`;
STATE.base=R;
STATE.areas=[...new Set(R.map(r=>r.ar).filter(Boolean))].sort();STATE.types=[...new Set(R.map(r=>r.ty).filter(Boolean))].sort();
const foodish=!f.t||f.t==="food";
$("#app").innerHTML=`<div class="ph"><div><h1>${tr?esc(tr.short):"All businesses"}</h1></div></div>
<nav class="lsw" aria-label="Trades"><a href="${writeF(Object.assign({},f,{t:"",type:""}))}"${!f.t?' aria-current="page"':""}>${ico("all")}All<span class="c num">${int(R.length)}</span></a>${TORD.map(k=>`<a href="${writeF(Object.assign({},f,{t:k,type:"",fi:k==="food"?f.fi:""}))}"${k===f.t?' aria-current="page"':""}>${ico(k)}${esc(TR[k].short)}<span class="c num">${int(inTrade(k).length)}</span></a>`).join("")}</nav>
<p class="ldesc">${tr?esc(tr.kinds)+" with a new City license in the last 30 days.":"Every business with a new City license in the last 30 days, at a Chicago business site."} Newest first.</p>
${foodish&&D.checks.food?`<p class="ovl"><b>Also a member of Chicago Restaurant Inspection Leads?</b> ${esc(W.overlap)} <button class="lnk" type="button" id="hideFood">Hide them here</button></p>`:""}
<div class="tool"><label class="search"><span class="sr">Search this list</span>${ico("search")}<input id="fq" type="search" placeholder="Search name, address, license or activity" value="${esc(f.q)}" autocomplete="off" enterkeyhint="search"><button type="button" class="clr" id="fqx" aria-label="Clear search"${f.q?"":" hidden"}>${ico("x","s")}</button></label>
<div class="seg" id="fh" role="group" aria-label="On the City's record"></div>
<button class="btn fbtn" id="fopen" type="button" aria-haspopup="dialog" aria-controls="fsheet">${ico("filter")}<span id="fopenT">Filters</span></button>
<span class="tsp"></span><button class="btn csvb" type="button" id="csv" aria-label="Export these businesses as a CSV file">${ico("csv","s")}<span class="t" aria-hidden="true">Export CSV</span><span class="m" aria-hidden="true">CSV</span></button>
<div class="seg vtog" role="group" aria-label="Show as"><button type="button" data-v="list">${ico("list","s")}List</button><button type="button" data-v="map">${ico("map","s")}Map</button></div></div>
<div class="filters" id="fsheet" aria-label="Filters">
<div class="sheet-h"><h2 id="fsheetH">Filters</h2><button class="btn sm" type="button" id="fclose" aria-label="Close filters">${ico("x","s")}</button></div>
<div class="fld"><label for="fdays">Licensed</label><select id="fdays"></select></div>
<div class="fld"><label for="farea">Community area</label><select id="farea"></select></div>
<div class="fld"><label for="ftype">License type</label><select id="ftype"></select></div>
${D.checks.food?`<div class="fld"><label for="ffi">Food license inspection</label><select id="ffi"></select></div>`:""}
<div class="fld"><label for="fsort">Sort</label><select id="fsort"><option value="newest">Newest first</option><option value="az">Name, A to Z</option><option value="area">Community area</option></select></div>
<div class="fld"><label for="ffu">Follow-up</label><select id="ffu"><option value="">All businesses</option><option value="todo">Not contacted yet</option><option value="mine">Contacted or with a note</option></select></div>
<label class="tog"><input type="checkbox" id="fnew"><span id="fnewT">New this week</span></label>
<div class="sheet-f"><button class="btn" type="button" id="freset">Clear filters</button><button class="btn ink" type="button" id="fdone">Show businesses</button></div>
</div>
<div class="rbar" aria-live="polite"><span id="rcount"></span><button class="lnk" type="button" id="fclear" hidden>Clear filters</button></div>
<div class="split" id="split" data-v="${f.view}"><section class="panel lst" aria-label="Businesses on this list"><div id="plist"></div></section>
<div class="mapbox" id="mapbox"><div class="map" id="map" role="region" aria-label="Map of the businesses in the list. The list has the same businesses."></div>
<div class="legend" aria-hidden="true"><span><i class="lg-1"></i>First license</span><span><i class="lg-2"></i>New site or added license</span><span><i class="lg-3"></i>Not checked</span></div>
<p class="map-note" id="mapnote" hidden></p></div></div>`;
const upd=patch=>{Object.assign(STATE.f,patch);STATE.shown=60;history.replaceState(history.state,"",writeF(STATE.f));STATE.trail[STATE.trail.length-1]=location.hash;refresh(true);};
let tq;$("#fq").addEventListener("input",e=>{$("#fqx").hidden=!e.target.value;clearTimeout(tq);tq=setTimeout(()=>upd({q:e.target.value.trim()}),140);});
$("#fqx").onclick=()=>{$("#fq").value="";$("#fqx").hidden=true;upd({q:""});$("#fq").focus();};
$("#fdays").onchange=e=>upd({days:e.target.value});$("#farea").onchange=e=>upd({area:e.target.value});$("#ftype").onchange=e=>upd({type:e.target.value});
if($("#ffi"))$("#ffi").onchange=e=>upd({fi:e.target.value});
$("#fsort").onchange=e=>upd({sort:e.target.value});$("#fnew").onchange=e=>upd({nw:e.target.checked});$("#ffu").onchange=e=>upd({fu:e.target.value});
const hf=$("#hideFood");if(hf)hf.onclick=()=>{upd({fi:"no"});toast("Businesses with a food license inspection are hidden");};
const clear=()=>{$("#fq").value="";$("#fqx").hidden=true;upd({q:"",area:"",days:"",type:"",h:"",fi:"",nw:false,fu:""});};
$("#fclear").onclick=clear;$("#freset").onclick=clear;$("#fopen").onclick=openSheet;$("#fclose").onclick=closeSheet;$("#fdone").onclick=closeSheet;
document.querySelectorAll(".vtog button").forEach(b=>b.onclick=()=>{upd({view:b.dataset.v});$("#split").dataset.v=b.dataset.v;syncVtog();if(b.dataset.v==="map")showMap();});
$("#csv").onclick=exportCsv;STATE.upd=upd;
refresh(false);syncVtog();
const phone=window.matchMedia("(max-width:860px)").matches;if(!phone)ensureMap();else if(f.view==="map")setTimeout(showMap,60);}
function showMap(){const t=document.querySelector(".tool");if(t&&window.matchMedia("(max-width:860px)").matches){const y=t.getBoundingClientRect().top+window.scrollY-parseInt(getComputedStyle(document.documentElement).getPropertyValue("--bar-h"))-12;window.scrollTo(0,Math.max(0,y));}
setTimeout(()=>{ensureMap();if(STATE.map){STATE.map.invalidateSize();drawMap(true);}},40);}
function syncVtog(){document.querySelectorAll(".vtog button").forEach(b=>b.setAttribute("aria-pressed",String(b.dataset.v===STATE.f.view)));}
function refresh(refit){const f=STATE.f,base=STATE.base;
const facet=(skip,test)=>base.filter(r=>pass(r,f,skip)&&test(r)).length;
const hb=D.checks.history?[["","Any record",()=>true],["first","First license",r=>r.h==="first"],["site","New site",r=>r.h==="site"],["added","Added license",r=>r.h==="added"]]:[];
$("#fh").hidden=!hb.length;
$("#fh").innerHTML=hb.map(([k,t,fn])=>`<button type="button" data-k="${k}" aria-pressed="${f.h===k}">${k?`<span class="hi ${k}" style="height:auto;padding:0;background:none"><i aria-hidden="true"></i></span>`:""}${esc(t)}<span class="c num">${int(facet("h",fn))}</span></button>`).join("");
$("#fh").querySelectorAll("button").forEach(b=>b.onclick=()=>STATE.upd({h:b.dataset.k}));
const opt=(v,t,n,sel)=>`<option value="${esc(v)}"${sel?" selected":""}>${esc(t)}${n==null?"":` (${int(n)})`}</option>`;
$("#fdays").innerHTML=DAYS.map(([v,t])=>opt(v,t,v?facet("days",r=>r.d>=since(v)):null,f.days===v)).join("");
$("#farea").innerHTML=opt("","All community areas",null,!f.area)+STATE.areas.map(a=>opt(a,a,facet("area",r=>r.ar===a),f.area===a)).join("");
$("#ftype").innerHTML=opt("","Any license type",null,!f.type)+STATE.types.map(t=>[t,facet("type",r=>r.ty===t)]).filter(([t,n])=>n||t===f.type).map(([t,n])=>opt(t,t,n,f.type===t)).join("");
if($("#ffi"))$("#ffi").innerHTML=opt("","Any",null,!f.fi)+opt("yes","With one (likely on the restaurant list)",facet("fi",r=>!!r.fi),f.fi==="yes")+opt("no","Without one",facet("fi",r=>!r.fi),f.fi==="no");
$("#fsort").value=f.sort;$("#fnew").checked=f.nw;$("#ffu").value=f.fu;$("#fnewT").textContent=`New this week (${int(facet("new",r=>r.nw))})`;
const ac=activeCount(f);$("#fopenT").innerHTML=ac?`Filters <span class="fcount num">${ac}</span>`:"Filters";
const rows=base.filter(r=>pass(r,f)).sort(SORTS[f.sort]());STATE.rows=rows;const tot=inTrade(f.t).length;
$("#rcount").innerHTML=rows.length===tot?`<b class="num">${int(rows.length)}</b> ${rows.length===1?"business":"businesses"}`:`<b class="num">${int(rows.length)}</b> of <span class="num">${int(tot)}</span> businesses`;
$("#fclear").hidden=!(f.q||f.h||ac);$("#fdone").textContent=rows.length?`Show ${plural(rows.length,"business","businesses")}`:"No businesses match";$("#csv").disabled=!rows.length;
drawList();drawMap(refit);}
function drawList(){const rows=STATE.rows,n=Math.min(STATE.shown,rows.length);const el=$("#plist");
if(!rows.length){const q=STATE.f.q;el.innerHTML=`<div class="empty"><h3>${q?`No businesses match "${esc(q)}"`:"No businesses match these filters"}</h3><p>${q?"Check the spelling, or clear the search":"Try another area or trade, or clear the filters"} to see all ${int(inTrade(STATE.f.t).length)}.</p><button class="btn ink" type="button" id="eclr">${q?"Clear search and filters":"Clear filters"}</button></div>`;$("#eclr").onclick=()=>$("#fclear").click();return;}
el.innerHTML=rowsHtml(rows,n,!STATE.f.t)+(n<rows.length?`<div class="row-more"><button class="btn" type="button" id="more">Show ${int(Math.min(60,rows.length-n))} more</button><span>${int(n)} of ${int(rows.length)} shown</span></div>`:"");
const m=$("#more");if(m)m.onclick=()=>{STATE.shown+=60;drawList();};
el.querySelectorAll(".row").forEach(a=>{a.addEventListener("mouseenter",()=>hot(a.dataset.id,true));a.addEventListener("mouseleave",()=>hot(a.dataset.id,false));});}

/* map */
const MS={site:{radius:6.5,color:"#fff",weight:2,fillColor:"#2F5F9E",fillOpacity:1},added:{radius:6.5,color:"#fff",weight:2,fillColor:"#2F5F9E",fillOpacity:1},
none:{radius:6,color:"#3A4152",weight:2.2,fillColor:"#fff",fillOpacity:1}};
const PK=r=>r.h==="first"?"first":r.h?r.h:"none";
function tiles(m){L.tileLayer("https://tile.openstreetmap.org/{z}/{x}/{y}.png",{maxZoom:19,attribution:'&copy; <a href="https://www.openstreetmap.org/copyright">OpenStreetMap</a> contributors'}).addTo(m);}
function ensureMap(){if(STATE.map||!$("#map"))return;if(!window.L){const box=$("#mapbox");if(box&&!box.querySelector(".map-fail"))box.insertAdjacentHTML("beforeend",'<div class="map-fail"><p>The map did not load on this connection. The list has every business.</p></div>');if(!ensureMap.tries){ensureMap.tries=1;window.addEventListener("load",()=>{if($("#map")&&!STATE.map){const f=$("#mapbox .map-fail");f&&f.remove();ensureMap();}},{once:true});}return;}
const m=L.map("map",{zoomControl:true,scrollWheelZoom:true}).setView([41.85,-87.68],11);
m.attributionControl.setPrefix('<a href="https://leafletjs.com">Leaflet</a>');tiles(m);STATE.map=m;STATE.layer=L.layerGroup().addTo(m);STATE.mk={};drawMap(true);}
function popup(r){return `<div class="pop"><span class="mn" style="font-size:16px">${esc(r.n)}</span><span class="ad">${esc(r.a)}</span><span class="bd">Licensed ${dayShort(r.d)}, ${esc(r.ty)}</span>${hi(r)}<p style="margin:10px 0 0"><a class="btn ink" href="#/biz/${encodeURIComponent(r.id)}">Open business</a></p></div>`;}
function drawMap(refit){const m=STATE.map;const rows=STATE.rows||[];const geo=rows.filter(r=>r.lat&&r.lon);const note=$("#mapnote");
if(note){const miss=rows.length-geo.length;note.hidden=!(miss||!rows.length);note.textContent=!rows.length?"No pins for these filters.":miss?`${plural(miss,"business","businesses")} in the list ${miss===1?"has":"have"} no map location in the City record.`:"";}
if(!m)return;STATE.layer.clearLayers();STATE.mk={};
const rank={none:0,added:1,site:1,first:2};
geo.slice().sort((a,b)=>rank[PK(a)]-rank[PK(b)]).forEach(r=>{const k=PK(r);const c=pin([r.lat,r.lon],k,1).bindPopup(popup(r),{maxWidth:280});c.addTo(STATE.layer);STATE.mk[r.id]=[c,k];});
if(refit&&geo.length){const bb=L.latLngBounds(geo.map(r=>[r.lat,r.lon]));if(geo.length===1)m.setView(bb.getCenter(),15,{animate:false});else m.fitBounds(bb,{padding:[40,40],maxZoom:15,animate:false});}}
/* first-license pins are the legend's square (a div icon); the others stay circles */
function pin(ll,k,scale){if(k==="first"){const z=Math.round(13*scale);return L.marker(ll,{icon:L.divIcon({className:"pin-f",iconSize:[z,z]}),keyboard:false,riseOnHover:true});}
return L.circleMarker(ll,Object.assign({},MS[k],{radius:MS[k].radius*scale}));}
function hot(id,on){const e=STATE.mk&&STATE.mk[id];if(!e)return;const [c,k]=e;
if(c.setStyle){c.setStyle(on?Object.assign({},MS[k],{radius:MS[k].radius+4,weight:3,color:"#1D2433"}):MS[k]);if(on)c.bringToFront();}
else{const el=c.getElement&&c.getElement();if(el)el.classList.toggle("hot",on);c.setZIndexOffset(on?1000:0);}}

/* filters sheet (phone) */
let lastFocus=null;
function openSheet(){const s=$("#fsheet");if(!s)return;lastFocus=document.activeElement;s.classList.add("open");s.setAttribute("role","dialog");s.setAttribute("aria-modal","true");s.setAttribute("aria-labelledby","fsheetH");
const sc=document.createElement("div");sc.className="scrim";sc.id="scrim";sc.onclick=closeSheet;document.body.appendChild(sc);setTimeout(()=>{const h=$("#fsheetH");h.setAttribute("tabindex","-1");h.focus();},60);document.addEventListener("keydown",sheetKeys);}
function closeSheet(){const s=$("#fsheet");const sc=$("#scrim");if(sc)sc.remove();document.removeEventListener("keydown",sheetKeys);if(!s||!s.classList.contains("open"))return;s.classList.remove("open");s.removeAttribute("role");s.removeAttribute("aria-modal");if(lastFocus&&lastFocus.focus)lastFocus.focus();}
function sheetKeys(e){if(e.key==="Escape"){closeSheet();return;}if(e.key!=="Tab")return;const f=[...$("#fsheet").querySelectorAll("select,input,button")].filter(x=>x.offsetParent!==null);if(!f.length)return;const a=f[0],z=f[f.length-1];
if(e.shiftKey&&document.activeElement===a){e.preventDefault();z.focus();}else if(!e.shiftKey&&document.activeElement===z){e.preventDefault();a.focus();}}

/* CSV of the current view */
function cell(v){v=v==null?"":String(v);if(/^[=+\-@\t\r]/.test(v)&&!/^-?\d+(\.\d+)?$/.test(v))v="'"+v;return /[",\r\n]/.test(v)?'"'+v.replace(/"/g,'""')+'"':v;}
function exportCsv(){const rows=STATE.rows||[];if(!rows.length)return;
const head=["Business","Address","ZIP","Community area","Trade","Licensed","License type","Business activity","On the City's record","Food license inspection","Other licenses here","New this week","License","City record","Contacted","Contacted on","Your note","Latitude","Longitude"];
const lines=[head.map(cell).join(",")];
rows.forEach(r=>{const m=mk(r.id);lines.push([r.n,r.a,r.z,r.ar,TR[r.t].short,r.d,r.ty,r.ac,hiText(r),r.fi,r.ot.join("; "),r.nw?"Yes":"",r.id,D.record+r.id,m.c?"Yes":"",m.c||"",(m.n||"").trim(),r.lat||"",r.lon||""].map(cell).join(","));});
const blob=new Blob(["﻿"+lines.join("\r\n")+"\r\n"],{type:"text/csv;charset=utf-8"});
const a=document.createElement("a");a.href=URL.createObjectURL(blob);a.download=`chicago-new-business-licenses-${D.asOfIso}-${STATE.f.t||"all"}-${rows.length}.csv`;document.body.appendChild(a);a.click();a.remove();setTimeout(()=>URL.revokeObjectURL(a.href),4000);
toast(`CSV exported: ${plural(rows.length,"business","businesses")}`);}
function downloadWorkbook(){const blob=new Blob([b64(D.xlsx)],{type:"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"});
const a=document.createElement("a");a.href=URL.createObjectURL(blob);a.download=D.fileName;document.body.appendChild(a);a.click();a.remove();setTimeout(()=>URL.revokeObjectURL(a.href),4000);toast("Workbook downloaded");}

/* business page */
const fv=(v,c,sub)=>v?`<dd${c?` class="${c}"`:""}>${esc(v)}${sub?`<span class="dn">${esc(sub)}</span>`:""}</dd>`:'<dd class="nv">Not on the City record</dd>';
function viewBiz(r){document.title=`${r.n} | Chicago New Business Licenses`;const m=mk(r.id);
const maps="https://www.google.com/maps/search/?api=1&query="+encodeURIComponent(r.n+", "+r.a+", Chicago, IL");
const find="https://www.google.com/search?q="+encodeURIComponent(r.n+" "+r.a+" Chicago");
$("#app").innerHTML=`<a class="back" href="#/list" id="back">${ico("back")}Back</a>
<div class="pg"><section class="pg-head"><div class="ph-chips">${tch(r.t)}${r.nw?'<span class="badge new">New this week</span>':""}${m.c?`<span class="badge done">${ico("check","s")}Contacted ${esc(day(m.c))}</span>`:""}</div>
<div class="p-title"><h1>${esc(r.n)}</h1><p class="sub">${esc(r.a)}, ${esc([r.ar,"Chicago"].filter(Boolean).join(", "))}</p></div>
<div class="acts"><button class="btn" type="button" id="cpy">${ico("copy","s")}Copy address</button><a class="btn" href="${maps}" target="_blank" rel="noopener">${ico("pin","s")}Map and phone<span class="sr"> (opens Google Maps in a new tab)</span></a><a class="btn" href="${find}" target="_blank" rel="noopener">${ico("search","s")}Search the web<span class="sr"> (opens in a new tab)</span></a></div></section>
<div class="pg-main">
<section class="panel lic" aria-labelledby="lh"><h2 id="lh">The City's new license</h2>
<div class="top"><span class="ty">${esc(r.ty)}</span></div><p class="act">Issued ${esc(day(r.d))}${r.ac?`. Business activity: ${esc(r.ac)}`:""}.</p>
${r.h?`<p class="why">${hi(r,true)}<br><span style="display:block;margin-top:8px">${esc(W.historyLong[r.h])}${r.hs?` The earliest such license on the City's record is from ${esc(r.hs)}.`:""}</span></p>`:`<p class="why">${esc(W.historyLong[""])}</p>`}
${r.fi?`<p class="why"><b>City food license inspection on ${esc(day(r.fi))}.</b> Chicago inspects a food business before it issues a food license. ${esc("Chicago Restaurant Inspection Leads lists those inspections as Opening soon, so a member of both products most likely saw this business there first.")}</p>`:""}
${r.ot.length?`<p class="why"><b>Also licensed at this site in the last 30 days:</b> ${esc(r.ot.join("; "))}.</p>`:""}
</section>
<section class="panel pg-facts"><div class="sec-h"><h2>Facts from the City record</h2></div><dl class="facts">
<div><dt>License type</dt>${fv(r.ty)}</div><div><dt>Issued</dt>${fv(day(r.d))}</div><div><dt>Community area</dt>${fv(r.ar)}</div><div><dt>ZIP</dt>${fv(r.z,"num")}</div>
<div><dt>License number</dt>${fv(r.id,"num")}</div><div><dt>Trade</dt>${fv(TR[r.t].short,"",TR[r.t].kinds)}</div><div><dt>On the City's record</dt>${fv(hiText(r)||"")}</div><div><dt>Food license inspection</dt>${r.fi?fv(day(r.fi)):'<dd class="nv">None on file</dd>'}</div></dl>
<p class="rec" style="padding-top:12px"><a href="${esc(D.record+r.id)}" target="_blank" rel="noopener">City record for license ${esc(r.id)}<span class="sr"> (opens in a new tab)</span></a></p>
${r.lat&&r.lon?'<div class="minimap" id="minimap" role="region" aria-label="Map of the business site"></div>':""}</section>
</div>
<aside class="pg-side" aria-label="Your follow-up">
<section class="panel note" aria-labelledby="nh"><h2 id="nh">Your follow-up</h2><p class="hint">${esc(STORE_LINE())}</p>
<button class="btn${m.c?" on":""}" type="button" id="mark" aria-pressed="${!!m.c}">${ico("check","s")}<span id="markT">${m.c?"Contacted "+esc(day(m.c))+" (undo)":"Mark as contacted"}</span></button>
<label for="nt" class="sr">Private note</label><textarea id="nt" placeholder="Who you spoke to, what they need, when to call back" maxlength="2000">${esc(m.n||"")}</textarea><p class="st" id="ntst" aria-live="polite"></p></section>
<section class="panel bc"><h2 style="font-size:17px">What to say</h2><p class="meta" style="margin:6px 0 14px">A call opener and an email for cleaning, pest control and other vendors, with this business's details to fill in.</p><a class="btn" href="#/pitch">${ico("kit","s")}Open the pitch kit</a></section>
</aside></div>`;
$("#back").onclick=e=>{e.preventDefault();back("#/list");};
$("#cpy").onclick=async()=>{toast(await copyText(r.n+", "+r.a+", Chicago, IL")?"Address copied":"Copy did not work. Press and hold the address.");};
$("#mark").onclick=()=>{const now=mk(r.id).c?"":new Date().toISOString().slice(0,10);const ok=setMark(r.id,{c:now});const b=$("#mark");b.classList.toggle("on",!!now);b.setAttribute("aria-pressed",String(!!now));$("#markT").textContent=now?"Contacted "+day(now)+" (undo)":"Mark as contacted";toast(ok?(now?"Marked as contacted":"Contacted mark removed"):"Not saved: this browser blocks saving");};
const grow=t=>{t.style.height="auto";t.style.height=Math.max(96,t.scrollHeight+2)+"px";};grow($("#nt"));
let tn;$("#nt").addEventListener("input",e=>{grow(e.target);clearTimeout(tn);$("#ntst").textContent="Saving...";tn=setTimeout(()=>{const ok=setMark(r.id,{n:e.target.value});$("#ntst").textContent=ok?"Note saved in this browser":"Not saved. This browser blocks saving here; private browsing is the usual cause.";},400);});
if(r.lat&&r.lon){const mm=()=>{if(!window.L||!$("#minimap"))return;const m2=L.map("minimap",{scrollWheelZoom:false,dragging:!L.Browser.mobile,tap:false}).setView([r.lat,r.lon],15);m2.attributionControl.setPrefix('<a href="https://leafletjs.com">Leaflet</a>');tiles(m2);pin([r.lat,r.lon],PK(r),1.6).addTo(m2);STATE.mini=m2;};if(window.L)mm();else window.addEventListener("load",mm,{once:true});}}

/* pitch kit */
function viewPitch(q){const ks=W.kits;const b=ks.find(k=>k.key===q.get("b"))||ks[0];
document.title="Pitch kit | Chicago New Business Licenses";
const fill=s=>esc(s).replace(/\[([^\]]+)\]/g,'<span class="f">[$1]</span>');
$("#app").innerHTML=`<div class="ph"><div><h1>Pitch kit</h1><p>A short call opener and an email for each buyer. Fill in the parts in blue brackets from the business page, and keep every claim true.</p></div></div>
<nav class="bk" aria-label="Buyers">${ks.map(k=>`<a href="#/pitch?b=${k.key}"${k===b?' aria-current="page"':""}>${ico(k.key,"s")}${esc(k.short)}</a>`).join("")}</nav>
<div class="pk"><section class="panel"><div class="pk-h"><h2>Call or visit</h2><button class="btn sm" type="button" data-copy="c">${ico("copy","s")}Copy</button></div><p class="sub">${esc(b.buyer)}. About 20 seconds.</p><p class="call">${fill(b.call)}</p>
<ol class="rules">${W.rules.map(r=>`<li>${esc(r)}</li>`).join("")}</ol></section>
<section class="panel mail"><div class="pk-h"><h2>Email</h2><button class="btn sm" type="button" data-copy="m">${ico("copy","s")}Copy</button></div><p class="sub">Send it the same day as the call.</p>${b.email.map(x=>`<p>${fill(x)}</p>`).join("")}</section></div>`;
document.querySelectorAll("[data-copy]").forEach(x=>x.onclick=async()=>{toast(await copyText(x.dataset.copy==="c"?b.call:b.email.join("\n\n"))?"Copied":"Copy did not work. Select the text.");});
const on=document.querySelector('.bk a[aria-current="page"]');if(on&&on.scrollIntoView)on.scrollIntoView({block:"nearest",inline:"center"});}

/* terms */
function termsHtml(){const para=t=>{const m=/^([A-Z][A-Za-z ]{1,40}?\.) (.+)$/.exec(t);return m?`<p><b>${esc(m[1])}</b> ${esc(m[2])}</p>`:`<p>${esc(t)}</p>`;};
return `<div class="ph"><div><h1>Terms of Use</h1><p>${esc(W.version)}.<br>The same terms are on the Terms tab of the Excel file.</p></div></div>
<div class="terms"><section class="panel" aria-labelledby="th1"><div class="sec-h"><h2 id="th1">City of Chicago Data Terms of Use</h2></div><div class="tx"><p>The City of Chicago's Data Terms of Use, word for word, as the City published them (checked ${esc(W.cityChecked)}). They apply to the City data on this list.</p>
${W.citySections.map(([h,ps])=>`<h3>${esc(h)}</h3>${ps.map(p=>`<p>${esc(p)}</p>`).join("")}`).join("")}
<p><a href="${esc(W.cityTerms)}" target="_blank" rel="noopener">City of Chicago Data Terms of Use<span class="sr"> (opens in a new tab)</span></a></p></div></section>
<section class="panel" aria-labelledby="th2"><div class="sec-h"><h2 id="th2">ProofNotFluff</h2></div><div class="tx"><p>${esc(W.notice)}</p>${W.paras.map(para).join("")}</div></section></div>`;}
function viewTerms(){document.title="Terms of Use | Chicago New Business Licenses";$("#app").innerHTML=`<a class="back" href="#/" id="back">${ico("back")}Back</a>`+termsHtml();$("#back").onclick=e=>{e.preventDefault();back("#/");};}
function gateHash(){if(D)return;const t=(location.hash||"")==="#/terms";$("#gateView").hidden=t;$("#gateLegal").hidden=t;$("#app").hidden=!t;$("#appLegal").hidden=!t;
if(t){document.title="Terms of Use | Chicago New Business Licenses";$("#app").innerHTML=`<a class="back" href="#/">${ico("back")}Back to the password</a>`+termsHtml();window.scrollTo(0,0);const h=$("#app h1");h.setAttribute("tabindex","-1");h.focus({preventScroll:true});}
else{$("#app").innerHTML="";document.title="Members: Chicago New Business Licenses | ProofNotFluff";}}
window.addEventListener("hashchange",gateHash);gateHash();

$("#skip").onclick=e=>{e.preventDefault();const t=$("#app").hidden?$("#pw"):($("#app h1")||$("#app"));if(t){t.setAttribute("tabindex",t.id==="pw"?"0":"-1");t.focus();}};
const pw=$("#pw"),eye=$("#eye");
eye.onclick=()=>{const show=pw.type==="password";pw.type=show?"text":"password";eye.setAttribute("aria-pressed",String(show));eye.setAttribute("aria-label",show?"Hide password":"Show password");$("#eyeslash").hidden=!show;pw.focus();};
$("#gate").onsubmit=async e=>{e.preventDefault();const err=$("#err"),btn=$("#go");
if(!pw.value.trim()){err.textContent="Type the password from your Gumroad purchase page.";pw.focus();return;}
btn.disabled=true;btn.textContent="Opening the list...";err.textContent="";
try{const d=await openBox(pw.value.trim());pw.value="";prep(d);
$("#gateView").hidden=true;$("#gateLegal").hidden=true;$("#app").hidden=false;$("#appLegal").hidden=false;$("#tabs").hidden=false;$("#bnav").hidden=false;$("#dlTop").hidden=false;document.body.classList.add("in");
$("#dlTop").onclick=downloadWorkbook;window.removeEventListener("hashchange",gateHash);window.addEventListener("hashchange",onHash);STATE.trail=[location.hash||"#/"];route(true);}
catch(x){btn.disabled=false;btn.textContent="Open this week's list";err.textContent="That password did not open the list. Check your Gumroad purchase page, or reply to your receipt for help.";pw.select();}};
</script></body></html>
"""


def check(full, html, password, xlsx_path):
    """The rows check (gate D): the app's decrypted rows and every trade tab of the workbook hold exactly the runner's list,
    business for business and field for field. Prints counts only (no list data in the log)."""
    from openpyxl import load_workbook
    box = json.loads(re.search(r"const BOX=(\{.*?\});", html).group(1))
    d = json.loads(decrypt(password, box))
    want = {r["license"]: r for r in full["records"]}
    got = {r["id"]: r for r in d["rows"]}
    assert len(d["rows"]) == len(full["records"]) == len(want) == len(got), (len(d["rows"]), len(full["records"]), len(want), len(got))
    assert set(got) == set(want), f"{len(set(want) ^ set(got))} licenses differ between the runner's list and the app"
    for k, r in want.items():
        g = got[k]
        pairs = [(g["n"], r["business"]), (g["a"], r["address"]), (g["ar"], r.get("communityArea") or ""), (g["ty"], r["licenseType"]),
                 (g["ac"], r.get("activity") or ""), (g["d"], r["issued"]), (g["t"], r["segment"]), (g["h"], r.get("history") or ""),
                 (g["fi"], r.get("foodInsp") or ""), (g["ot"], r.get("otherTypes") or []), (g["z"], r.get("zip") or "")]
        assert all(a == b for a, b in pairs), f"license {k}: app field differs from the runner's list"
    wb = load_workbook(xlsx_path)
    seen = {}
    for key in X.ORDER:
        ws = wb[X.TAB[key]]
        heads = {ws.cell(11, c).value: c for c in range(3, 20) if ws.cell(11, c).value}
        for row in range(12, ws.max_row + 1):
            lic = ws.cell(row, heads["LICENSE"]).value
            if not lic:
                continue
            r = want[str(lic)]
            issued = ws.cell(row, heads["LICENSED"]).value
            assert r["segment"] == key and ws.cell(row, heads["BUSINESS"]).value == r["business"] and \
                ws.cell(row, heads["ADDRESS"]).value == r["address"] and ws.cell(row, heads["LICENSE TYPE"]).value == r["licenseType"] and \
                issued is not None and issued.strftime("%Y-%m-%d") == r["issued"], f"workbook row {X.TAB[key]}!{row} differs"
            seen[str(lic)] = key
    assert set(seen) == set(want), f"{len(set(want) ^ set(seen))} licenses differ between the runner's list and the workbook"
    by = {k: sum(1 for v in seen.values() if v == k) for k in X.ORDER}
    return {"rows": len(want), "app": len(got), "workbook": len(seen), **by}


def selftest():
    import tempfile, shutil, subprocess
    import voice_lint
    full = X._planted()
    out = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    X.build(full, out)
    html = build_html(full, open(out, "rb").read(), "correct horse battery")
    clear = re.sub(r"const BOX=\{.*?\};", "", html)          # base64 ciphertext can spell anything by chance
    for name in ("Luna Tacos", "Bright Kids", "Main Street Phones", "100 W Lake"):
        assert name not in clear, f"list data must stay inside the ciphertext: {name}"
    assert 'name="robots" content="noindex' in html and "List of Oct 22, 2026" in html
    text_only = re.sub(r"<[^>]+>", "", html).replace("&#x27;", "'").replace("&quot;", '"')
    assert text_only.count(legal_text().replace("__ASOF__", "Oct 22, 2026")) == 2, "the legal line word for word on the gate and the footer"
    assert X.city_disclaimer() in legal_text() and html.count('<a href="#/terms">Terms of Use</a>') == 2
    words = json.loads(re.search(r"const W=(\{.*?\});\n", html).group(1))
    assert [words["version"]] + words["paras"] == X.terms_text("Oct 22, 2026") and words["city"] == X.city_disclaimer()
    assert words["citySections"] == X.city_terms()["sections"] and words["overlap"] == X.OVERLAP and words["leftOut"] == X.LEFT_OUT
    assert [k["key"] for k in words["kits"]] == ["clean", "pest", "any"] and words["rules"] == X.PITCH_RULES
    assert base_css()[:200] in html and ".trades{" in html, "the construction app's stylesheet plus ours"
    for u, h in (leads_app.LEAFLET_JS, leads_app.LEAFLET_CSS):
        assert u in html and h in html
    assert "openstreetmap.org/copyright" in html
    box = json.loads(re.search(r"const BOX=(\{.*?\});", html).group(1))
    assert box["rounds"] == 250_000 and box["kdf"] == "PBKDF2-SHA256"
    d = json.loads(decrypt("correct horse battery", box))
    by = {r["n"]: r for r in d["rows"]}
    assert [r["t"] for r in d["rows"]] == ["food", "food", "retail", "beauty", "auto", "care"], "trade order, newest first inside"
    assert by["Luna Tacos"]["fi"] == "2026-10-02" and by["Luna Tacos"]["nw"] and by["Luna Tacos"]["ot"] and not by["Bright Kids Academy"]["nw"]
    assert by["Main Street Phones"]["h"] == "site" and by["Main Street Phones"]["hs"] == "2015" and d["checks"] == {"history": True, "food": True}
    assert d["xlsx"] and d["fileName"] == "chicago-new-business-licenses-2026-10-22.xlsx"
    assert check(full, html, "correct horse battery", out)["rows"] == 6
    bad = json.loads(json.dumps(full)); bad["records"][0]["business"] = "Changed"
    caught = False
    try:
        check(bad, html, "correct horse battery", out)
    except AssertionError as e:
        caught = "license 1" in str(e)
    assert caught, "a row that differs from the runner's list must fail the check"
    try:
        decrypt("wrong password", box)
        raise AssertionError("wrong password decrypted")
    except AssertionError:
        raise
    except Exception:
        pass
    node = shutil.which("node")
    if node:   # browser helpers: dates, the CSV formula guard, the history labels
        fns = "\n".join(re.search(p, PAGE, re.M).group(0) for p in (r"^const MON=.*$", r"^function day\(s\).*$", r"^function dayShort\(s\).*$",
                                                                     r"^function cell\(v\).*$", r"^function hiText\(r\).*$"))
        js = ("const W={history:" + json.dumps(X.HISTORY) + "};\n" + fns + "\nconsole.log(JSON.stringify([day('2026-10-05'),dayShort('2026-10-15'),day(''),"
              "[-87.6,'=SUM(A1)','-1 call','@x','a,b'].map(cell),[{h:'first'},{h:'site',hs:'2015'},{h:'added',hs:'2019'},{h:''}].map(hiText)]))")
        got = json.loads(subprocess.run([node, "-e", js], capture_output=True, text=True, check=True).stdout)
        assert got == ["Oct 5, 2026", "Oct 15", "", ["-87.6", "'=SUM(A1)", "'-1 call", "'@x", '"a,b"'],
                       ["First license", "New site, licensed since 2015", "Added license, here since 2019", "Not checked"]], got
    else:
        print("node not found: browser helpers not checked here")
    text = PAGE + LEGAL + words_json()
    assert "—" not in text and "–" not in text, "em or en dash"
    for w in ("honestly", "genuinely", "straightforward", "delve", "unlock ", "elevate", "seamless", "game-changer", "effortless", "supercharge",
              "guarantee", "struggling"):
        assert w not in text.lower(), w
    tells = voice_lint.tells(re.sub(r"<[^>]+>|\$\{[^}]*\}", " ", PAGE + legal_text()))
    assert not tells, tells
    print("lic_app selftest ok")


if __name__ == "__main__":
    args = sys.argv[1:]
    if args == ["--selftest"]:
        selftest()
    else:
        pw = os.environ.get(SECRET, "")
        verify = "--check" in args
        args = [a for a in args if a != "--check"]
        if "--pw-file" in args:
            i = args.index("--pw-file")
            pw = open(args[i + 1]).read().strip()
            args = args[:i] + args[i + 2:]
        if len(pw) < 12:
            raise SystemExit(f"{SECRET} missing or shorter than 12 characters; nothing written")
        full = json.load(open(args[0], encoding="utf-8"))
        res = build(full, args[1], args[2], pw)
        if verify:
            res["check"] = check(full, open(args[2], encoding="utf-8").read(), pw, args[1])
        print(res)
