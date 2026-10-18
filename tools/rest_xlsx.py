#!/usr/bin/env python3
"""Chicago Restaurant Inspection Leads: the weekly workbook (WORKBOOK_STANDARD v3, the same card language as the
construction product's tools/leads_xlsx.py, whose card helpers this reuses).

Tabs: Start Here, 1 Failed or conditional, 2 Opening soon, 3 Failed more than once, 4 Inspection history, 5 Pitch kit,
Terms. Built from tools/rest_leads.py build() so it matches the members' app number for number. Business-level City
fields only: the business name, address and license as the City lists them, the City's result, date and inspection
type, the City's violation code numbers, and our category for each. Never the inspector's notes.

This file is also the one source of the product's words that two surfaces share: the pitch kit (PITCH) and the Terms
of Use (terms_text), read by the members' app (tools/rest_app.py) so they cannot drift.

Usage: python3 tools/rest_xlsx.py DATA.json OUT.xlsx      (DATA from tools/rest_leads.py)
       python3 tools/rest_xlsx.py --selftest
"""
import datetime as dt, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import wb_style as S  # noqa: E402
import pnf_dash as D  # noqa: E402
from leads_xlsx import CITY_DISCLAIMER, GRID, INNER, TEXT_W, card_sheet, rich  # noqa: E402
from openpyxl.styles import Alignment, Border  # noqa: E402
from openpyxl.utils import get_column_letter as gl  # noqa: E402
from rest_leads import CATEGORIES, CAT_LABEL, CAT_ORDER, PRODUCT, pretty
CAT_SHORT = {k: s for k, _, s, _ in CATEGORIES}  # noqa: E402

TABS = {"flagged": "1 Failed or conditional", "opening": "2 Opening soon", "repeat": "3 Failed more than once",
        "history": "4 Inspection history", "pitch": "5 Pitch kit"}
FOOT = "Public City of Chicago food inspection records. See the Terms tab."
OUR_NOTICE = ("For general business information only. Places appear because their food inspections are public City of Chicago "
              "records. This list repeats the City's recorded result, date and violation codes, and sorts the violations into "
              "categories; it does not say whether any place is clean, safe or open today, and a result can change at the next "
              "inspection. No business listed has endorsed or works with ProofNotFluff. You make your own calls and emails, and "
              "you are responsible for following the law when you do. ProofNotFluff never contacts businesses or anyone else on "
              "your behalf.")
# result chips: one color system with the members' app (text color, fill); the result word is always written out
RESULT_CHIP = {"FAIL": ("A63E21", "FBEAE4"), "PASS W/ CONDITIONS": ("7A5A12", "F6EDD5"), "PASS": ("2E6B66", "E3EFED"),
               "NOT READY": ("2F5F9E", "E6EDF6"), "NO ENTRY": ("3A4152", "ECE6DA")}

# ---------------------------------------------------------------- the pitch kit (one source: the app reads it too)
# [Brackets] are the parts the member fills in. Every opener names the City's public record and the violation category,
# never a judgment about the business.
PITCH = [
    {"key": "pests", "buyer": "Pest control", "short": "Pest control", "list": "flagged", "cat": "pests",
     "call": "Hi, this is [your name] with [company], a licensed pest control company in Chicago. The City's public inspection "
             "record for [business] on [date] lists violation [code], [the City's title for it]. We set restaurants up with a service plan and "
             "the pest control log the City asks to see. Is the owner or manager around for two minutes?",
     "email": ["Subject: Pest control for [business]",
               "Hi [business] team, I'm [your name] with [company], a licensed pest control company working with Chicago restaurants. "
               "The City's public food inspection record for [address] on [date] lists violation [code] ([the City's title for it]).",
               "We can inspect this week, treat, seal entry points and leave a service log for your next inspection. Would [day] "
               "work for a visit?",
               "[Your name], [company], [phone], Illinois structural pest control license [number]"]},
    {"key": "hood", "buyer": "Hood and exhaust cleaning", "short": "Hood cleaning", "list": "flagged", "cat": "hood",
     "call": "Hi, this is [your name] with [company]. We clean kitchen hoods, ducts and exhaust fans for Chicago restaurants. "
             "The City's public inspection record for [business] on [date] lists violation [code], [the City's title for it]. "
             "Could I quote a cleaning before the next inspection? Who handles the kitchen equipment?",
     "email": ["Subject: Hood and exhaust cleaning for [business]",
               "Hi [business] team, I'm [your name] with [company]. We clean hoods, filters, ducts and exhaust fans in Chicago "
               "kitchens and leave a dated service record. The City's public food inspection record for [address] on [date] "
               "lists violation [code] ([the City's title for it]). We can give you a quote and a time this week. Who should I send it to?",
               "[Your name], [company], [phone]"]},
    {"key": "drain", "buyer": "Grease trap and drain service", "short": "Grease trap", "list": "flagged", "cat": "drain",
     "call": "Hi, this is [your name] with [company]. We pump and clean grease traps and jet drains for Chicago restaurants. "
             "The City's public inspection record for [business] on [date] lists violation [code], [the City's title for it]. Can we set "
             "up a service visit? Who looks after the kitchen plumbing?",
     "email": ["Subject: Grease trap and drain service for [business]",
               "Hi [business] team, I'm [your name] with [company]. We pump grease traps and interceptors and clear kitchen "
               "drains across Chicago, on a schedule that fits your volume. The City's public food inspection record for "
               "[address] on [date] lists violation [code] ([the City's title for it]). Can I book a visit or send a quote?",
               "[Your name], [company], [phone]"]},
    {"key": "clean", "buyer": "Commercial kitchen cleaning", "short": "Kitchen cleaning", "list": "flagged", "cat": "clean",
     "call": "Hi, this is [your name] with [company]. We do deep cleaning for Chicago restaurant kitchens, after hours. The "
             "City's public inspection record for [business] on [date] lists violation [code], [the City's title for it]. Could we quote a deep clean "
             "before the next inspection? Who should I speak with?",
     "email": ["Subject: Kitchen deep cleaning for [business]",
               "Hi [business] team, I'm [your name] with [company]. We deep clean restaurant kitchens in Chicago after hours: "
               "equipment, floors, walls and storage areas. The City's public food inspection record for [address] on [date] "
               "lists violation [code] ([the City's title for it]). We can give you a quote and a time this week. Who should I send it to?",
               "[Your name], [company], [phone]"]},
    {"key": "opening", "buyer": "Vendors to new openings (POS, linen, cleaning, pest, hood)", "short": "New openings", "list": "opening", "cat": "",
     "call": "Hi, this is [your name] with [company]. I saw the City's license inspection for [business] at [address] on [date], "
             "so you may be opening soon. We set up new Chicago restaurants with [your service]. Who is putting the "
             "vendors together for the opening?",
     "email": ["Subject: [Your service] for the opening at [address]",
               "Hi [business] team, I'm [your name] with [company]. The City's public records show a license inspection at "
               "[address] on [date], so you may be lining up vendors for the opening. We set up new Chicago restaurants with "
               "[your service], usually within [days] of the first call. Could I send a short quote?",
               "[Your name], [company], [phone]"]},
]
PITCH_RULES = ["Name the City's public record and its date, and stop there. Don't say a place is dirty, unsafe or closed.",
               "Keep every claim about your company true, and never promise a place will pass its next inspection.",
               "One call and one email per place. Honor any request to stop, and follow the law on calls and email."]


def terms_text(date):
    """The Terms of Use (legal/README.md clause library, adapted for a monthly membership like the construction product's).
    One source: the workbook's Terms tab and the members' app's #/terms view are both built from this list."""
    return [
        f"ProofNotFluff | {PRODUCT} | Version 1, Oct 17, 2026 | List of {date}",
        "License. Your membership gives you a license to use this list for one business you own, for its own sales prospecting. You may "
        "edit it and print it for that use. You may not resell, share, give away, upload, publish or redistribute it, in original or "
        "edited form, or use it to make a competing product.",
        "Ownership. The selection, layout, categories, text and formulas belong to ProofNotFluff. The inspection records themselves are "
        "public City of Chicago data. Buying a membership doesn't transfer ownership.",
        "Accuracy. The list is built from City of Chicago food inspection records as of the list date above. City data can be late, "
        "incomplete or wrong. The City says its dataset may contain duplicate reports, and that each result and the violations noted "
        "reflect what the inspector found at the time of that inspection and may not reflect other times. A place's result can change at "
        "its next inspection. The violation categories are ProofNotFluff's reading of the City's violation codes and notes, not the City's "
        "own labels. Inspection types are the City's, except that the City's \"Suspected Food Poisoning\" inspections, which follow a "
        "complaint, are shown as Complaint. Confirm anything you rely on with the business or the City.",
        "What the list does not say. A place on the list has a public inspection record; the list does not say it is clean or unsafe, "
        "open or closed, or in need of any service. Don't present the list, or anything in it, as a judgment about a business.",
        "No promise of results. Counts and categories are figures from City records. Nothing here promises any work, contract, income "
        "or outcome.",
        "Availability. The list depends on the City of Chicago publishing its food inspection data. If the City stops or changes that "
        "data, or asks us to stop using it, the list may pause or end, and future charges will stop. The map uses OpenStreetMap's free "
        "public map service, which can limit or stop access at any time; the list and downloads work without it. Map data © "
        "OpenStreetMap contributors.",
        "No warranty. The product is provided \"as is\", without warranties of any kind, to the extent the law allows.",
        "Limit of liability. To the extent the law allows, ProofNotFluff is not liable for any loss or damage from using or relying on "
        "this product, and its total liability for any claim is limited to the amount you paid for the most recent month of your "
        "membership. Nothing here limits rights you have under consumer protection laws that can't be waived.",
        "Trademarks. Microsoft Excel, Google Sheets, Google Maps, OpenStreetMap and Gumroad are trademarks of their owners, and business "
        "names on the list belong to those businesses. ProofNotFluff isn't affiliated with, sponsored or endorsed by them, by the City "
        "of Chicago, or by any business listed.",
        "AI assistance. This product was made with AI assistance and reviewed and tested by the shop owner.",
        "Help and cancelling. If the file won't open or something looks wrong, reply to your Gumroad receipt and it will be fixed. Cancel "
        "any time from your Gumroad receipt or library; cancelling stops future charges.",
    ]


# ---------------------------------------------------------------- helpers
def day(s):
    try:
        return dt.date.fromisoformat(s[:10])
    except (TypeError, ValueError):
        return None


def chip(c, result):
    fg, bg = RESULT_CHIP.get((result or "").upper(), ("3A4152", "ECE6DA"))
    c.fill = D.fill(bg); c.font = D.f(9, True, fg)
    c.border = Border(bottom=D.side("hair"))


def codes_text(codes):
    return ", ".join(str(c) for c in codes)


def latest_note(p, I, lead):
    """The place's newest City result when it is newer than the inspection that put it on the list ("Pass, Oct 14, 2026")."""
    x = I[p["latestId"]]
    if p["latestId"] == lead or x["d"] <= I[lead]["d"]:
        return ""
    return f"{x['r']}, {pretty(x['d'])}"


WRAP = ("name", "address", "area", "cats", "codes", "latest", "itype", "license", "also")


def lic_text(p):
    """One license, or the first and a count ("3078843 +2"), so the cell never spills (design director, round 1)."""
    ls = p.get("licenses") or ([p["license"]] if p.get("license") else [])
    return "" if not ls else ls[0] if len(ls) == 1 else f"{ls[0]} +{len(ls) - 1}"


def also_text(p, l):
    """Categories cited on an earlier qualifying inspection than the row's own ("Hood 56, Oct 3")."""
    return "; ".join(f"{CAT_SHORT[a['cat']]} {codes_text(a['codes'])}, {pretty(a['d'])[:-6]}" for a in (p.get("alsoBy") or {}).get(l, []))


def data_sheet(wb, key, title, tagline, list_date, tab, cols, rows, stats, freeze_col=1, print_cols=None):
    """A full-width table tab: cols [(key, heading, width, align, fmt)], rows [{key: value}], stats [(eyebrow, formula fn, fmt)]
    where the formula fn takes {key: "C12:C40"} ranges. The first stat is the ink answer tile. print_cols: how many columns
    print (the rest stay on screen), so a wide tab prints at 8 pt or more."""
    ws = wb.create_sheet(TABS[key])
    letters = [gl(i + 3) for i in range(len(cols))]
    pad_r, gut_r = gl(len(cols) + 3), gl(len(cols) + 4)
    S.set_widths(ws, {"A": 3, "B": 2, **{l: c[2] for l, c in zip(letters, cols)}, pad_r: 2, gut_r: 3})
    D.page_header(ws, title, tagline, list_date, span=("B", pad_r), side_cols=3)
    top = 10
    first, last = top + 2, top + 2 + max(len(rows), 1) - 1
    col = {c[0]: l for l, c in zip(letters, cols)}
    rng = {k: f"{l}{first}:{l}{last}" for k, l in col.items()}
    D.h(ws, 6, 10)
    k = len(stats)
    widths = [c[2] for c in cols[:print_cols or len(cols)]]
    total, acc, cuts = sum(widths), 0, []
    for w in widths:
        acc += w
        cuts.append(acc)
    starts = [0] + [min(range(len(cuts)), key=lambda j: abs(cuts[j] - total * i / k)) + 1 for i in range(1, k)]
    lastc = letters[(print_cols or len(cols)) - 1]
    groups = [f"{letters[starts[i]]}:{letters[starts[i + 1] - 1] if i < k - 1 else lastc}" for i in range(k)]
    D.stat_strip(ws, 7, [(g, e, fn(rng), fmt) for g, (e, fn, fmt) in zip(groups, stats)])
    # the answer tile: the first stat on ink, as on Start Here (design director, round 1, fix 36)
    a, z = groups[0].split(":")
    tile_cols = [gl(i) for i in range(2, letters.index(z) + 4)]
    for rr in range(6, 10):
        for c in tile_cols:
            ws[f"{c}{rr}"].fill = D.fill("ink")
    for g in groups[1:]:                       # light stats sit off the ink tile's edge (round 2, B2)
        for rr in (7, 8):
            c = ws[f"{g.split(':')[0]}{rr}"]
            c.alignment = Alignment(horizontal="left", vertical=c.alignment.vertical, indent=1)
    ws[f"{a}7"].font = D.f(8, True, "on_ink_soft")
    ws[f"{a}8"].font = D.f(22, True, "on_ink")
    D.h(ws, 9, 10)
    D.frame(ws, 6, 9, "B", pad_r)
    spec = []
    for l, (ck, head, w, align, fmt) in zip(letters, cols):
        vals = [r.get(ck) for r in rows]
        spec.append({"col": l, "head": head, "kind": "text", "values": [None if v in ("", None) else v for v in vals], "align": align,
                     "fmt": "mmm d, yyyy" if fmt == "date" else (S.FMT[fmt] if fmt else None), "bold": ck == "name"})
    f1, f2, end = D.table_card(ws, top, "B", pad_r, spec, max(len(rows), 1), row_height=30)
    head = f1 - 1
    D.h(ws, head, 30)
    # the column after a result chip starts one indent in, so it never touches the fill (round 2, B6)
    after_chip = {letters[i + 1] for i, c in enumerate(cols[:-1]) if c[0] in ("result", "lastResult")}
    ind = lambda l, align: 1 if align == "right" or l in after_chip else 0
    for l, (ck, head_, w, align, fmt) in zip(letters, cols):
        ws[f"{l}{head}"].alignment = Alignment(horizontal=align, vertical="bottom", wrap_text=True, indent=ind(l, align))
    if not rows:
        ws[f"{letters[0]}{f1}"] = "None this week."
        ws[f"{letters[0]}{f1}"].font = D.f(10, False, "muted")
    for i, r in enumerate(rows):
        rr = f1 + i
        need = 30
        for l, (ck, head_, w, align, fmt) in zip(letters, cols):
            c = ws[f"{l}{rr}"]
            c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=ck in WRAP, indent=ind(l, align))
            if ck in ("license", "area", "type") and r.get(ck) in ("", None, "Not given"):
                c.value = "Not given"; c.font = D.f(10, False, "muted")      # one convention for a missing value (fix 34)
            elif ck in ("result", "lastResult") and r.get(ck):
                chip(c, r.get(ck))
            elif ck in ("address", "area", "type", "itype", "date", "license", "codes", "latest", "lastFail", "lastDate"):
                c.font = D.f(10, False, "ink2")
            elif ck in CAT_ORDER and r.get(ck):
                c.font = D.f(10, True, "teal")
            elif ck == "also" and r.get(ck):
                c.font = D.f(9, False, "ink2")
            elif ck == "new" and r.get(ck):
                c.font = D.f(9, True, "accent")
            elif ck == "record" and r.get("recordUrl"):
                c.hyperlink = r["recordUrl"]; c.font = D.f(10, False, "blue", underline="single")
            if ck in WRAP and c.value:
                need = max(need, S.text_height(str(c.value), w, 10, bold=ck == "name") + 8)   # nothing cut off (fix 24)
        D.h(ws, rr, need)
    if rows:
        ws.auto_filter.ref = f"{letters[0]}{head}:{letters[-1]}{f2}"
    D.paint_canvas(ws, end + 1, gl(len(cols) + 6))
    S.finish_sheet(ws, PRODUCT, end, span=("A", gut_r), freeze=f"{letters[freeze_col]}{f1}", protect=False, footer_notice=FOOT,
                   tab_color=tab, canvas_cols=gl(len(cols) + 6))
    if print_cols:
        ws.print_area = f"A1:{gl(print_cols + 2)}{end}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.print_title_rows = f"{head}:{head}"
    return ws, col


def finish_card(ws, r, tab):
    D.paint_canvas(ws, r, "S")
    S.finish_sheet(ws, PRODUCT, r - 1, span=("A", "Q"), protect=False, canvas_cols="S", footer_notice=FOOT, tab_color=tab)


def blue(line):
    """[brackets] in bold blue, the rest plain."""
    return [(part, part.startswith("["), "input_text" if part.startswith("[") else "ink") for part in re.split(r"(\[[^\]]+\])", line) if part]


def th(text, cols, size=10, bold=False):
    """Row height for text in a merged run of card columns. A column w characters wide is about w*7+5 px in Excel, so a
    merged run is wider than the sum of its widths; measuring against the sum left empty bands under long paragraphs
    (design director, round 2, B4). A 5 percent margin keeps Excel from cutting a last word."""
    usable = (sum(GRID[c] * 7 + 5 for c in cols) * 0.75 - 6) * 0.95
    return S.text_height(text, (usable + 2) / 5.2, size, bold)


def card(ws, top, title, rows, eyebrow=None, bg="card", line="frame"):
    """The one-column card (leads_xlsx.card's layout) with three more row kinds: "step" (an accent number in C and the text
    hanging in E to O, so wrapped lines never run under the number), "gap" (a paragraph space) and "notegrey" (9 pt note
    grey). Every row height comes from text_height()."""
    r = top
    D.h(ws, r, 14); r += 1
    if eyebrow:
        e = ws[f"C{r}"]; e.value = eyebrow; e.font = D.f(8, True, "muted"); ws.merge_cells(f"C{r}:O{r}"); D.h(ws, r, 16); r += 1
    if title:
        t = ws[f"C{r}"]; t.value = title; t.font = D.f(12, True); t.alignment = Alignment(vertical="center")
        ws.merge_cells(f"C{r}:O{r}"); D.h(ws, r, 24); r += 1
        D.h(ws, r, 6); r += 1
    for kind, val in rows:
        if kind == "step":
            n, text = val
            c = ws[f"C{r}"]; c.value = n; c.font = D.f(16, True, "accent"); c.alignment = Alignment(horizontal="right", vertical="top", indent=1)
            t = ws[f"E{r}"]; t.value = text; t.font = D.f(10); t.alignment = Alignment(vertical="top", wrap_text=True)
            ws.merge_cells(f"E{r}:O{r}"); D.h(ws, r, max(th(text, 'EFGHIJKLMNO', 10) + 10, 26))
        elif kind == "gap":
            D.h(ws, r, 8)
        elif kind == "lead":
            lead, text = val
            L = ws[f"C{r}"]; L.value = lead; L.font = D.f(10, True); L.alignment = Alignment(vertical="top", wrap_text=True)
            c = ws[f"E{r}"]; c.value = text if isinstance(text, str) else rich(text)
            c.font = D.f(10); c.alignment = Alignment(vertical="top", wrap_text=True)
            plain = text if isinstance(text, str) else "".join(t for t, _, _ in text)
            ws.merge_cells(f"E{r}:O{r}"); D.h(ws, r, max(th(plain, 'EFGHIJKLMNO', 10) + 6, 22))
            for k in INNER:
                ws[f"{k}{r}"].border = Border(bottom=D.side("hair"))
        elif kind == "link":
            text, url = val
            c = ws[f"C{r}"]; c.value = text; c.hyperlink = url; c.font = D.f(10, False, "blue", underline="single")
            c.alignment = Alignment(vertical="center", wrap_text=True); ws.merge_cells(f"C{r}:O{r}"); D.h(ws, r, 22)
        elif kind == "rich":
            text = "".join(t for t, _, _ in val)
            c = ws[f"C{r}"]; c.value = rich(val); c.alignment = Alignment(vertical="top", wrap_text=True)
            ws.merge_cells(f"C{r}:O{r}"); D.h(ws, r, th(text, INNER, 10) + 6)
        else:
            size, color = (9, "5C6370") if kind == "notegrey" else (9, "muted") if kind == "note" else (10, "ink")
            c = ws[f"C{r}"]; c.value = val; c.font = D.f(size, False, color); c.alignment = Alignment(vertical="top", wrap_text=True)
            ws.merge_cells(f"C{r}:O{r}"); D.h(ws, r, th(val, INNER, size) + 6)
        r += 1
    D.h(ws, r, 14)
    D.frame(ws, top, r, "B", "P", bg=bg, line=line)
    return r + 2


def lead_para(t):
    """A Terms paragraph with its lead-in in bold ("License. Your membership...")."""
    m = re.match(r"^([A-Z][A-Za-z ]{1,40}?\.) (.+)$", t)
    return ("rich", [(m.group(1) + " ", True, "ink"), (m.group(2), False, "ink")]) if m else ("para", t)


# ---------------------------------------------------------------- build
def build(data, out):
    P, I, L = data["places"], data["inspections"], data["lists"]
    list_date = f"List of {data['asOfText']}"
    rec = data["source"]["record"]
    wb = S.new_workbook(PRODUCT)
    catcols = [(k, f"{s} (violation)", 10, "center", None) for k, _, s, _ in CATEGORIES]

    def base(p, x, l=None):
        return {"name": p["name"], "address": p["address"], "area": p["area"], "type": p["type"], "result": x["r"], "date": day(x["d"]),
                "itype": x["t"], "license": lic_text(p), "new": "New" if l and p["newBy"][l] else "",
                "recordUrl": rec + x["id"], "record": "City record"}

    # 1 failed or conditional: one row per place, its newest qualifying inspection. The category columns hold that
    # inspection's own City codes; "Also cited" names categories from an earlier qualifying inspection (holdout, round 1).
    rows1 = []
    for pk in L["flagged"]:
        p = P[pk]
        x = I[p["flaggedIds"][0]]
        r = base(p, x, "flagged")
        r.update({c: codes_text(x["c"][c]) if c in x["c"] else "" for c in CAT_ORDER})
        r["also"] = also_text(p, "flagged")
        r["codes"] = codes_text([v[0] for v in x["v"]])
        r["fails"] = p["fails90"]
        r["latest"] = latest_note(p, I, x["id"])
        rows1.append(r)
    cnt = lambda col: (lambda rg: f'=COUNTIF({rg[col]},"?*")')
    catcnt = lambda col, word: (lambda rg: f'=COUNTIF({rg[col]},"?*")+COUNTIFS({rg[col]},"",{rg["also"]},"*{word}*")')
    # Columns C to O print (13 columns, about 9 pt on legal landscape); "Also cited" and the rest stay on screen
    # (design director, round 2, B1).
    cols1 = ([("name", "Business", 24, "left", None), ("address", "Address", 20, "left", None), ("area", "Community area", 15, "left", None),
              ("result", "City result", 17, "center", None), ("date", "Inspected", 13, "left", "date"), ("itype", "Inspection type", 18, "left", None)]
             + catcols + [("fails", "Fails in 90 days", 10, "right", "int"), ("latest", "Newer City result", 17, "left", None),
                          ("new", "New this week", 9, "center", None), ("also", "Also cited in the last 30 days", 22, "left", None),
                          ("codes", "All violations cited", 18, "left", None), ("type", "Facility type", 14, "left", None),
                          ("license", "License", 14, "left", None), ("record", "City record", 11, "left", None)])
    _, c1 = data_sheet(wb, "flagged", "Failed or conditional",
                       "City result Fail or Pass w/ Conditions in the last 30 days, with a violation a buyer works on. Newest first.",
                       list_date, "accent", cols1, rows1,
                       [("Places", cnt("name"), S.FMT["int"]), ("Failed", lambda rg: f'=COUNTIF({rg["result"]},"Fail")', S.FMT["int"]),
                        ("Pests", catcnt("pests", "Pests"), S.FMT["int"]), ("Hood", catcnt("hood", "Hood"), S.FMT["int"]),
                        ("Grease trap", catcnt("drain", "Grease trap"), S.FMT["int"]), ("Cleaning", catcnt("clean", "Cleaning"), S.FMT["int"])],
                       print_cols=13)

    # 2 opening soon
    rows2 = []
    for pk in L["opening"]:
        p = P[pk]
        x = I[p["openingIds"][0]]
        r = base(p, x, "opening")
        r["cats"] = ", ".join(CAT_SHORT[c] for c in x["c"]) if x["c"] else ""
        r["n"] = len(p["openingIds"])
        rows2.append(r)
    _, c2 = data_sheet(wb, "opening", "Opening soon",
                       "A license inspection in the last 30 days: Chicago inspects a food business before it issues a food license. Newest first.",
                       list_date, "teal",
                       [("name", "Business", 30, "left", None), ("address", "Address", 26, "left", None), ("area", "Community area", 18, "left", None),
                        ("result", "License inspection result", 17, "center", None), ("date", "Inspected", 13, "left", "date"),
                        ("itype", "Inspection type", 20, "left", None), ("n", "License inspections in 30 days", 13, "right", "int"),
                        ("cats", "Violation categories cited", 24, "left", None), ("new", "New this week", 9, "center", None),
                        ("type", "Facility type", 15, "left", None), ("license", "License", 14, "left", None), ("record", "City record", 11, "left", None)],
                       rows2, [("Places", cnt("name"), S.FMT["int"]), ("Passed", lambda rg: f'=COUNTIF({rg["result"]},"Pass*")', S.FMT["int"]),
                               ("Not passed yet", lambda rg: f'=COUNTIF({rg["name"]},"?*")-COUNTIF({rg["result"]},"Pass*")', S.FMT["int"]),
                               ("New this week", cnt("new"), S.FMT["int"])])

    # 3 failed more than once
    rows3 = []
    for pk in L["repeat"]:
        p = P[pk]
        last = I[p["failIds"][0]]
        r = base(p, last, "repeat")
        r["fails"] = p["fails90"]
        r["lastFail"] = day(last["d"])
        r["cats"] = ", ".join(f"{CAT_SHORT[c]} {codes_text(last['c'][c])}" for c in CAT_ORDER if c in last["c"])
        r["also"] = also_text(p, "repeat")
        x = I[p["latestId"]]
        r["lastResult"] = x["r"]
        r["lastDate"] = day(x["d"])
        rows3.append(r)
    data_sheet(wb, "repeat", "Failed more than once",
               "Two or more failed routine, complaint or re-inspections in 90 days, at least one in the last 30. Newest first.",
               list_date, None,
               [("name", "Business", 28, "left", None), ("address", "Address", 24, "left", None), ("area", "Community area", 17, "left", None),
                ("fails", "Fails in 90 days", 14, "right", "int"), ("lastFail", "Last failed", 13, "left", "date"),
                ("cats", "Categories on the last fail", 24, "left", None), ("also", "Also cited in the 90 days", 22, "left", None),
                ("lastResult", "Newest City result", 17, "center", None), ("lastDate", "Result date", 13, "left", "date"),
                ("type", "Facility type", 15, "left", None), ("license", "License", 14, "left", None), ("record", "City record", 11, "left", None)],
               rows3, [("Places", cnt("name"), S.FMT["int"]), ("Fails between them", lambda rg: f'=SUM({rg["fails"]})', S.FMT["int"]),
                       ("Most fails at one place", lambda rg: f'=MAX({rg["fails"]})', S.FMT["int"])])

    # 4 inspection history: every inspection on file for the places on the three lists (duplicate City reports counted once)
    rows4 = []
    order = list(dict.fromkeys(L["flagged"] + L["opening"] + L["repeat"]))
    for pk in sorted(order, key=lambda k: (P[k]["name"].upper(), P[k]["address"])):
        p = P[pk]
        for i in p["history"]:
            x = I[i]
            r = base(p, x)
            r["cats"] = ", ".join(f"{CAT_SHORT[c]} {codes_text(x['c'][c])}" for c in CAT_ORDER if c in x["c"])
            r["codes"] = codes_text([v[0] for v in x["v"]])
            rows4.append(r)
    data_sheet(wb, "history", "Inspection history",
               "Every City inspection on file in the last 24 months for each place on tabs 1 to 3, A to Z, newest first within a place.",
               list_date, None,
               [("name", "Business", 30, "left", None), ("address", "Address", 26, "left", None), ("date", "Inspected", 13, "left", "date"),
                ("itype", "Inspection type", 20, "left", None), ("result", "City result", 17, "center", None),
                ("cats", "Categories and codes", 26, "left", None), ("codes", "Violations cited", 22, "left", None),
                ("license", "License", 14, "left", None), ("record", "City record", 11, "left", None)],
               rows4, [("Inspections", cnt("name"), S.FMT["int"]), ("Failed", lambda rg: f'=COUNTIF({rg["result"]},"Fail")', S.FMT["int"]),
                       ("Passed with conditions", lambda rg: f'=COUNTIF({rg["result"]},"Pass w/ Conditions")', S.FMT["int"])], freeze_col=2)

    pitch_kit(wb, list_date)
    terms(wb, data, list_date)
    start_here(wb, data, list_date, len(rows1), len(rows2), len(rows3), c1["new"], c2["new"])
    wb.active = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "—" in c.value:
                    raise SystemExit(f"em dash in {ws.title}!{c.coordinate}")
    wb.save(out)
    return {"flagged": len(rows1), "opening": len(rows2), "repeat": len(rows3), "history": len(rows4)}


def start_here(wb, data, list_date, n1, n2, n3, new1, new2):
    ws = card_sheet(wb, "Start Here", PRODUCT, "Recent City food inspections, sorted into leads for the trades that serve restaurants.",
                    list_date, "teal", position=0)
    rng = lambda key, col, n: f"'{TABS[key]}'!{col}12:{col}{11 + max(n, 1)}"
    new_f = f'=COUNTIF({rng("flagged", new1, n1)},"New")+COUNTIF({rng("opening", new2, n2)},"New")'
    tiles = [("B", "D", "FAILED OR CONDITIONAL", f'=COUNTIF({rng("flagged", "C", n1)},"?*")', S.FMT["int"], "places, last 30 days", True),
             ("F", "H", "OPENING SOON", f'=COUNTIF({rng("opening", "C", n2)},"?*")', S.FMT["int"], "license inspections, 30 days", False),
             ("J", "L", "FAILED MORE THAN ONCE", f'=COUNTIF({rng("repeat", "C", n3)},"?*")', S.FMT["int"], "2 or more fails in 90 days", False),
             ("N", "P", "NEW ON TABS 1 AND 2", new_f, S.FMT["int"], f"inspected {pretty(data['window']['newFrom'])[:-6]} or later", False)]
    D.h(ws, 6, 10)
    for left, right, eb, fml, fmt, sub, dark in tiles:
        D.tile(ws, 7, left, right, eb, fml, fmt, sub, dark=dark, value_size=28 if dark else 20)
    r = card(ws, 12, "Three steps", [
        ("step", (1, "Find your trade below and open its tab. Filter your category column.")),
        ("step", (2, "Call in a quiet hour, 2 to 4 pm. Name the City's record and its date, nothing more.")),
        ("step", (3, f"Send the matching email from {TABS['pitch']} the same day.")),
    ])
    r = card(ws, r, "Find your trade", [("lead", (x["short"], (
        f"{TABS['flagged']}: filter the {CAT_SHORT[x['cat']]} column." if x["cat"] else f"{TABS['opening']}, for POS, linen, cleaning, pest and hood vendors. Every new license is a new account.")
        + (" " + TABS["repeat"] + " lists places with ongoing problems." if x["key"] == "pests" else ""))) for x in PITCH])
    w = data["window"]
    r = card(ws, r, "How the list is made", [
        ("para", f"From the City of Chicago's Food Inspections data, pulled {data['asOfText']}. The City posts results a few days late, so "
                 f"the window is the 30 days up to the newest inspection on file: {pretty(w['from'])[:-6]} to {data['throughText']}."),
        ("lead", ("Results", "The City's own words. Fail and Pass w/ Conditions are on tab 1. A place on tab 1 may have passed a re-inspection "
                             "since; the Newer City result column says so. Inspection types are the City's too, except that the City's "
                             "\"Suspected Food Poisoning\" type, an inspection that follows a complaint, is shown as Complaint.")),
        ("lead", ("Categories", "Ours, from the City's violation codes: Pests is violation 38 and any note naming pests; Hood is violation 56 "
                                "when it names a hood, filter, exhaust or grease; Grease trap is violation 52 and any grease trap, interceptor "
                                "or clogged drain; Cleaning is violation 49 and cleaning notes under 53, 54, 55 and 60. A 59 or 60 (an earlier "
                                "violation not corrected) counts in the category of the violation its note names. The numbers in each row are the "
                                "codes cited on that row's inspection; Also cited names a category from an earlier inspection in the window.")),
        ("lead", ("Left out", "Schools, colleges, day care and other children's sites, hospitals, senior and nursing sites, mobile vendors, any "
                              "place whose business name is a person's full name, and places the City last marked out of business.")),
        ("lead", ("Duplicates", "The City sometimes files one visit under several license numbers. Those reports count once, with every "
                                "violation they cite.")),
        ("lead", ("Not included", "The inspector's written notes. They are public on the City's site; this list only uses them to set a category.")),
        ("link", ("Source: City of Chicago Data Portal, Food Inspections", data["source"]["url"])),
    ])
    r = card(ws, r, "Before you rely on it", [
        ("notegrey", "For general information and planning only. This is not legal, tax, financial or other professional advice, and using it "
                     f"doesn't create a professional relationship. Figures come from City of Chicago records as of {data['asOfText']} and can change. "
                     "See the Terms tab before you rely on anything here."),
        ("gap", None), ("notegrey", OUR_NOTICE), ("gap", None), ("notegrey", CITY_DISCLAIMER)])
    finish_card(ws, r, "teal")
    return ws


def pitch_kit(wb, list_date):
    ws = card_sheet(wb, TABS["pitch"], "Pitch kit", "A call opener and an email for each trade. Fill in the blue parts.",
                    list_date, "teal")
    r = card(ws, 6, "Three rules", [("step", (i + 1, t)) for i, t in enumerate(PITCH_RULES)])
    for i, x in enumerate(PITCH):
        if i % 2:                                      # two trades per printed page, the rules share page 1 (round 2, B3)
            S.page_break_before(ws, r)
        r = card(ws, r, x["buyer"], [("lead", ("Call opener", blue(x["call"])))],
                 eyebrow=("TAB 2" if x["list"] == "opening" else f"TAB 1, {CAT_SHORT[x['cat']].upper()} COLUMN"))
        lines = []
        for i, line in enumerate(x["email"]):
            if i:
                lines.append(("gap", None))
            lines.append(("rich", blue(line)))
        r = card(ws, r, None, lines, eyebrow="EMAIL: COPY, THEN FILL IN THE BLUE PARTS", bg="total_tint", line="frame")
    finish_card(ws, r, "teal")
    return ws


def terms(wb, data, list_date):
    t = terms_text(data["asOfText"])
    ws = card_sheet(wb, "Terms", "Terms of Use", t[0].replace("ProofNotFluff | ", "").split(" | List of ")[0], list_date, "note")
    r = card(ws, 6, "City of Chicago data", [("para", CITY_DISCLAIMER), ("link", ("City of Chicago Data Terms of Use", data["source"]["terms"]))])
    rows = [("para", OUR_NOTICE)]
    for x in t[1:]:
        rows += [("gap", None), lead_para(x)]
    r = card(ws, r, "ProofNotFluff", rows)
    finish_card(ws, r, "note")
    return ws


def selftest():
    import tempfile
    from openpyxl import load_workbook
    import rest_leads
    rest_leads_data = rest_leads.build(_planted())
    out = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    res = build(rest_leads_data, out)
    wb = load_workbook(out, rich_text=True)
    assert wb.sheetnames == ["Start Here", *TABS.values(), "Terms"], wb.sheetnames
    assert res == {"flagged": 2, "opening": 2, "repeat": 1, "history": 8}, res
    t1 = wb[TABS["flagged"]]
    heads = [t1[f"{gl(i)}11"].value for i in range(3, 3 + 6)]
    assert heads == ["BUSINESS", "ADDRESS", "COMMUNITY AREA", "CITY RESULT", "INSPECTED", "INSPECTION TYPE"], heads
    assert t1["C12"].value == "Acme Tacos" and t1["F12"].value == "Fail" and t1["G12"].number_format == "mmm d, yyyy"
    assert t1["I12"].value == "38" and t1["J12"].value is None, (t1["I12"].value, t1["J12"].value)
    assert t1["C12"].font.name == "Arial" and t1.auto_filter.ref.startswith("C11")
    sh = wb["Start Here"]
    texts = [c.value for row in sh.iter_rows() for c in row if isinstance(c.value, str)]
    assert CITY_DISCLAIMER in texts and OUR_NOTICE in texts
    # the Start Here tiles count the right columns
    assert "'1 Failed or conditional'!C12:C13" in sh["C8"].value and "'1 Failed or conditional'!O12:O13" in [c.value for row in sh.iter_rows() for c in row if isinstance(c.value, str) and "O12" in c.value][0]
    assert wb[TABS["opening"]]["K11"].value == "NEW THIS WEEK", wb[TABS["opening"]]["K11"].value
    tt = [c.value for row in wb["Terms"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("Version 1, Oct 17, 2026" in t for t in tt) and CITY_DISCLAIMER in tt, "version in the subtitle, City text word for word"
    assert not any("Version 1" in t and "List of" in t for t in tt), "the list date is in the stamp only (round 2, B5)"
    assert len(pk_breaks := wb[TABS["pitch"]].row_breaks.brk) == len(PITCH) // 2, "two trades per printed page (round 2, B3)"
    rich_runs = [c.value for row in wb["Terms"].iter_rows() for c in row if type(c.value).__name__ == "CellRichText"]
    assert any(str(v).startswith("License. ") and v[0].font.b for v in rich_runs), "Terms lead-ins in bold"
    # holdout round 1: a row's category codes are its own inspection's; the 30-day union is the labelled Also column
    assert t1["P11"].value == "ALSO CITED IN THE LAST 30 DAYS" and t1["O11"].value == "NEW THIS WEEK", (t1["P11"].value, t1["O11"].value)
    assert t1.print_area and "$O$" in t1.print_area, t1.print_area           # Also cited stays on screen (round 2, B1)
    assert t1["G12"].alignment.indent == 1 and t1["G11"].alignment.indent == 1, "the column after City result is indented"
    assert t1["C7"].alignment.indent == 0 and all(t1[f"{c}7"].alignment.indent == 1 for c in "DEFGHIJKLMNOP" if t1[f"{c}7"].value), "light stats off the ink tile"
    fh = [r for r in range(12, 15) if t1[f"C{r}"].value == "Fry House"][0]
    assert t1[f"J{fh}"].value == "56" and t1[f"N{fh}"].value.startswith("Pass, "), "newer City result shown"
    t2 = wb[TABS["opening"]]
    lic = [t2[f"M{r}"].value for r in range(12, 14)]
    assert "1 +0" not in str(lic)
    pk = wb[TABS["pitch"]]
    blue_runs = [c.value for row in pk.iter_rows() for c in row if type(c.value).__name__ == "CellRichText"]
    assert any(any(getattr(b, "font", None) and b.font.color and str(b.font.color.rgb).endswith("1F4E9E") for b in v) for v in blue_runs), "blue blanks"
    for x in PITCH:
        for s in [x["call"]] + x["email"]:
            assert "—" not in s and not re.search(r"\b(dirty|unsafe|filthy|gross|disgusting)\b", s, re.I), s
    print("rest_xlsx selftest ok", res)


def _planted():
    import rest_leads
    V = rest_leads.V
    base = {"address": "1234 W MAIN ST", "zip": "60601", "latitude": "41.88", "longitude": "-87.63", "_area": "Loop", "facility_type": "Restaurant"}
    R = lambda i, name, lic, typ, res, d, v="", **kw: dict(base, inspection_id=str(i), dba_name=name, license_=lic, inspection_type=typ,
                                                          results=res, inspection_date=d + "T00:00:00.000", violations=v, **kw)
    rows = [R(1, "ACME TACOS", "111", "Canvass", "Fail", "2026-10-15", V(38, "MICE DROPPINGS")),
            R(2, "ACME TACOS", "111", "Canvass Re-Inspection", "Fail", "2026-09-30", V(38, "ROACHES")),
            R(3, "NEW PLACE", "222", "License", "Pass", "2026-10-14"),
            R(4, "FRY HOUSE", "444", "Canvass", "Pass w/ Conditions", "2026-10-13", V(56, "GREASE ON HOOD FILTERS")),
            R(5, "FRY HOUSE", "444", "Canvass Re-Inspection", "Pass", "2026-10-15"),
            R(6, "SECOND SHOP", "555", "License Re-Inspection", "Not Ready", "2026-10-12", address="9 N STATE ST"),
            R(7, "SECOND SHOP", "555", "License", "Fail", "2026-10-02", V(38, "GAP UNDER DOOR"), address="9 N STATE ST"),
            R(8, "ACME TACOS", "111", "Canvass", "Pass", "2025-03-01")]
    return {"source": {"name": "x", "url": rest_leads.DATASET_PAGE, "fetchedAt": "2026-10-18T03:00:00+00:00", "terms": rest_leads.CITY_TERMS},
            "rows": rows[:7], "history": rows[7:]}


if __name__ == "__main__":
    a = sys.argv[1:]
    if a == ["--selftest"]:
        selftest()
    else:
        print(build(json.load(open(a[0])), a[1]))
